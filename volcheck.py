"""Volume check and first use (2.7.0): check final volumes against the effort's list.

Pure functions (no GUI, no files read except where a caller passes paths to write). Inputs are the scan results the
app already has, the Effort List rows, and the volumes in submission order.

Report parts:
  first_use     acronyms only, judged in running text (2.23.0: figures, tables, captions, headers and notes are
                reported in their own column, never as a first use). Per volume (rule "volume"), per section of each volume (rule "section", 2.22.0) or across
                the submission (rule "submission"): where the
                acronym first appears, whether it is written "Spelled Out Words (ACR)", whether the spelled-out words
                match the effort list, later re-definitions.
  missing       terms the volumes use that the effort list does not have.
  not_approved  terms in the list whose status is Proposed or Needs author.
  differs       the volume spells a term out differently from the effort list.
  unused        list entries (not Rejected) that no volume uses: keep them out of the volumes' acronym lists.
  lists         per volume: every acronym, abbreviation and unit it uses, with the list's definition.
Nothing leaves the PC: the report is written only where the user saves it.
"""
import re
import alt_forms as AF

LIST_TYPES = ("Acronym", "Abbreviation", "Unit")
OK, WARN, BAD, INFO = "ok", "warn", "bad", "info"
RULE_TEXT = {"volume": "Restarts in each volume", "section": "Restarts in each section of each volume",
             "submission": "Runs across the whole submission"}
CHECK_MARKS = {OK: "\u2713 OK", BAD: "\u2717 Fix", WARN: "! Check", INFO: "i Note"}
KIND_NAMES = {"table": "Table", "figure": "Figure", "caption": "Caption", "header": "Header or footer",
              "notes": "Comment or notes"}
RULE_SHORT = {"volume": "per volume", "section": "per section of each volume", "submission": "across the whole submission"}


def natural_key(name):
    """Vol 2 before Vol 10."""
    return [int(p) if p.isdigit() else p.lower() for p in re.split(r"(\d+)", name)]


def norm_def(s):
    """Definition comparison: case, hyphens, slashes, punctuation, plural endings and a leading 'the' do not count."""
    s = (s or "").lower().replace("&", " and ")
    s = re.sub(r"[-/–—_]", " ", s)
    s = re.sub(r"[^a-z0-9 ]", "", s)
    words = [w[:-1] if len(w) > 3 and w.endswith("s") and not w.endswith("ss") else w for w in s.split()]
    if words and words[0] == "the":
        words = words[1:]
    return " ".join(words)


def same_def(a, b):
    return bool(a) and bool(b) and norm_def(a) == norm_def(b)


def _placeholder(sd):
    """The app's stand-in texts ("[Not in Database ...]", "MULTIPLE ...") are not definitions."""
    s = (sd or "").lstrip("\u26a0\ufe0f ")
    return not s or s.startswith(("[", "MULTIPLE", "Rescan"))


def kind(category):
    for k in LIST_TYPES:
        if (category or "").startswith(k):
            return k
    return ""


def _uses(d, vol, file_of):
    return [loc for loc in d.get("locations", []) if file_of(loc) == vol]


def build(results, rows, volumes, rule="volume", vol_defs=None, file_of=None, in_picture=None, vol_lists=None):
    """results: the app's scan results {term: data}. rows: EffortList.rows {term: {column: value}}.
    volumes: file names (as in location labels) in submission order. vol_defs: {volume: custdefs.extract() result}.
    vol_lists: {volume: custdefs.acronym_list() result}: each volume's own acronym list (2.25.0).
    Returns a dict with the report parts (lists of dicts) and a summary."""
    vol_defs = vol_defs or {}
    file_of = file_of or (lambda loc: loc)
    in_picture = in_picture or (lambda loc: False)
    used = {v: {} for v in volumes}            # volume -> {term: [locations]}
    for term, d in results.items():
        if d.get("manual") or d.get("file_error") or d.get("gone") or not kind(d.get("category", "")):
            continue
        row = rows.get(term)
        if d.get("excluded") and not (row and row.get("Status") == "Rejected"):
            continue                           # excluded in the database: not a term anywhere
        for v in volumes:
            locs = _uses(d, v, file_of)
            if locs:
                used[v][term] = locs

    def list_def(term):
        r = rows.get(term)
        return (r or {}).get("Definition", "")

    def status(term):
        r = rows.get(term)
        return r["Status"] if r else "Not in effort list"

    rejected = {t for t, r in rows.items() if r.get("Status") == "Rejected"}
    vol_lists = vol_lists or {}
    alts_found = {}                            # primary -> results of its alternate forms (2.26.0)
    for t, d in results.items():
        if d.get("alt_of"):
            alts_found.setdefault(d["alt_of"], []).append(d)

    # ---------------- first use (acronyms only)
    first_use = []
    defined_in = {}                            # term -> first volume where it was defined (submission rule)
    seen_in = {}                               # term -> first volume where it appeared (submission rule)
    no_sections = set()                        # section rule: terms whose saved results predate 2.22.0 (rescan)

    def _first_use_row(first_use, v, sec, term, u, fig=None):
        want = list_def(term)
        lf = u.get("lf", "")
        redefs = ["later in the same passage" if r == u["loc"] else r for r in u.get("redefs", [])]
        row = {"volume": v, "section": sec, "term": term, "loc": u["loc"], "written": lf, "list_definition": want,
               "status": status(term), "also_defined": "; ".join(redefs[:5]), "level": OK, "state": "",
               "in_figures": fig_text(fig)}
        pic = " (first use is in a picture)" if in_picture(u["loc"]) else ""
        earlier = seen_in.get(term) if rule == "submission" else None
        if earlier:
            if u["defined"] is True:
                row["level"], row["state"] = INFO, f"Defined again; first defined in {defined_in.get(term) or earlier}"
                if not defined_in.get(term):
                    row["level"], row["state"] = WARN, f"Defined here, but first used undefined in {earlier}"
                    defined_in[term] = v
            elif defined_in.get(term):
                row["level"], row["state"] = OK, f"Defined earlier in {defined_in[term]}"
            else:
                row["level"], row["state"] = BAD, f"Not defined before this; first used in {earlier}"
                if u.get("later"):
                    defined_in[term] = v
            first_use.append(row)
            return
        if rule == "submission":
            seen_in[term] = v
        if u["defined"] is True:
            if rule == "submission":
                defined_in[term] = v
            if want and lf and not same_def(lf, want):
                row["level"], row["state"] = WARN, "Defined, but the words differ from the effort list"
            elif not lf:
                row["level"], row["state"] = WARN, "Written as (ACR), but the spelled-out words were not found right before it"
            else:
                row["state"] = "Defined at first use"
        elif u["defined"] is None:
            row["level"], row["state"] = WARN, "Check: read with an OCR correction"
        elif u.get("later"):
            row["level"] = BAD
            row["state"] = ("Not defined at first use; spelled out later in the same passage" if u["later"] == u["loc"]
                            else f"Not defined at first use; spelled out later at {u['later']}")
            if rule == "submission":
                defined_in[term] = v
        else:
            row["level"], row["state"] = BAD, {"volume": "Never defined in the text of this volume", "section": "Never defined in the text of this section"}.get(
                rule, "Not defined at first use")
        if redefs and row["level"] == OK:
            row["level"] = INFO
            row["state"] += "; defined again later"
        row["state"] += pic
        first_use.append(row)

    def other_uses(term, v):
        """Uses in the volume besides the first use and the volume's own acronym list (text, tables, figures)."""
        n = ((results.get(term) or {}).get("uses") or {}).get(v)
        if n is None:
            return None                         # results saved before 2.25.0
        n += sum(((x.get("uses") or {}).get(v) or 0) for x in alts_found.get(term, []))   # its alt forms count (2.26.0)
        listed = term in (vol_lists.get(v) or {}).get("entries", {})
        return n - 1 - (1 if listed else 0)

    def fig_text(fig):
        return f"{KIND_NAMES.get(fig['kind'], fig['kind'])}: {fig['loc']}" if fig else ""

    def units(d, v, term):
        """[(section, first use)] for one acronym in one volume: one entry per section under the section rule."""
        if rule == "section":
            secs = (d.get("first_use_sec") or {}).get(v)
            if secs:
                named = {k: u for k, u in secs.items() if k}
                return list((named or secs).items())
            no_sections.add(term)
        u = d.get("first_use", {}).get(v)
        return [("", u)] if u else []

    for v in volumes:
        for term in sorted(used[v], key=str.lower):
            d = results[term]
            if kind(d.get("category", "")) != "Acronym" or term in rejected:
                continue
            if d.get("alt_of") and d["alt_of"] in used[v]:
                continue                        # an alt form: its primary's first use is what counts (2.26.0)
            fig = (d.get("first_fig") or {}).get(v)
            us = units(d, v, term)
            for i, (sec, u) in enumerate(us):
                _first_use_row(first_use, v, sec, term, u, fig if i == 0 else None)
            if not us and fig and term in (vol_lists.get(v) or {}).get("entries", {}) and (other_uses(term, v) or 0) <= -1:
                continue                        # only in the volume's acronym list: the Acronym list tab reports it
            if not us and fig:                  # used only in figures, tables, captions or headers (2.23.0)
                first_use.append({"volume": v, "section": "", "term": term, "loc": "", "written": "",
                                  "list_definition": list_def(term), "status": status(term), "also_defined": "",
                                  "in_figures": fig_text(fig), "level": INFO,
                                  "state": "Only in figures or tables; not in the text, so no first use to check"})

    # ---------------- defined at first use but never used again (2.25.0)
    for r in first_use:
        if r["level"] == OK and r.get("loc") and r["state"].startswith("Defined at first use"):
            n = other_uses(r["term"], r["volume"])
            if n is not None and n <= 0:
                r["level"] = WARN
                r["state"] = "Defined at first use, but never used again: spell it out and leave out the (ACR)"
                r["once"] = True
    for r in first_use:
        r["check"] = CHECK_MARKS.get(r["level"], "")

    # ---------------- each volume's own acronym list (2.25.0)
    acr_list = []
    for v in volumes:
        info = vol_lists.get(v)
        if info is None:
            continue
        entries = info.get("entries", {})
        if not entries:
            acr_list.append({"volume": v, "term": "", "level": INFO, "listed_as": "", "list_definition": "", "uses": "",
                             "state": "No acronym list found in this volume (a table, or lines of acronym<tab>definition)"})
            continue
        listed_up = {t.upper(): t for t in entries}
        list_alts = {a: p for p, al in (info.get("alts") or {}).items() for a in al}
        for term in sorted(used[v], key=str.lower):
            d = results[term]
            k = kind(d.get("category", ""))
            if k not in ("Acronym", "Abbreviation") or term in rejected:
                continue
            if (term not in entries and term.upper() not in listed_up and term not in list_alts
                    and d.get("alt_of") not in entries):
                n = other_uses(term, v)
                acr_list.append({"volume": v, "term": term, "level": BAD, "listed_as": "", "list_definition": list_def(term),
                                 "uses": "" if n is None else n + 1,
                                 "state": "Used in this volume but not in its acronym list"})
        for t, written in entries.items():
            term = t if t in results else next((x for x in results if x.upper() == t.upper()), t)
            if term in results:
                n = other_uses(term, v)
                outside = None if n is None else n + 1          # uses outside the acronym list
                for a in (info.get("alts") or {}).get(t, []):   # its alt forms written in the list (2.26.0)
                    if outside is not None and a in results and not results[a].get("alt_of"):
                        outside += ((results[a].get("uses") or {}).get(v) or 0) - 1
            else:
                outside = 0
            want = list_def(term) if term in rows else ""
            row = {"volume": v, "term": t, "listed_as": written, "list_definition": want,
                   "uses": "" if outside is None else max(outside, 0)}
            if term in rejected:
                row.update(level=WARN, state="In the acronym list, but the effort list marks it Not a term")
            elif outside is not None and outside <= 0:
                row.update(level=WARN, state="In the acronym list, but not used anywhere else in this volume")
            elif want and not same_def(written, want):
                row.update(level=WARN, state="The list's definition differs from the effort list")
            elif not want:
                row.update(level=INFO, state="In the acronym list; not in the effort list yet")
            else:
                row.update(level=OK, state="Listed, used and matches the effort list")
            acr_list.append(row)
    for r in acr_list:
        r["check"] = CHECK_MARKS.get(r["level"], "")

    # ---------------- missing / not approved / differs
    missing, not_approved, differs = [], [], []
    all_terms = {}
    for v in volumes:
        for term, locs in used[v].items():
            all_terms.setdefault(term, []).append((v, locs))
    for term in sorted(all_terms, key=str.lower):
        d = results[term]
        vols = [v for v, _ in all_terms[term]]
        first_loc = all_terms[term][0][1][0]
        base = {"term": term, "type": kind(d.get("category", "")), "volumes": ", ".join(vols), "loc": first_loc,
                "count": sum(len(l) for _, l in all_terms[term])}
        r = rows.get(term)
        if r is None and d.get("alt_of") and d["alt_of"] in rows:
            continue                            # an alt form of a term in the effort list (2.26.0)
        if r is None:
            sd = d.get("selected_definition", "")
            sd = "" if _placeholder(sd) else sd
            missing.append(dict(base, suggestion=sd))
            continue
        if r["Status"] == "Rejected":
            continue
        if r["Status"] in ("Proposed", "Needs author"):
            not_approved.append(dict(base, status=r["Status"], definition=r.get("Definition", "")))
        want = r.get("Definition", "")
        for v in vols:
            for h in (vol_defs.get(v) or {}).get(term, []):
                if want and not same_def(h["definition"], want):
                    differs.append(dict(base, volume=v, written=h["definition"], list_definition=want,
                                        loc=h.get("where", first_loc), sentence=h.get("sentence", "")))

    # ---------------- unused list entries
    unused = []
    for term in sorted(rows, key=str.lower):
        r = rows[term]
        if r.get("Status") == "Rejected" or term in all_terms or any(a in all_terms for a in AF.parse_alts(r.get("Alt forms", ""))):
            continue
        unused.append({"term": term, "type": r.get("Type", ""), "status": r.get("Status", ""), "definition": r.get("Definition", "")})

    # ---------------- per-volume acronym lists
    lists = {}
    for v in volumes:
        out = []
        alt_used = {}
        for term in used[v]:
            p = results[term].get("alt_of")
            if p and term not in rejected:
                alt_used.setdefault(p, []).append(term)
        for term in sorted(set(used[v]) | set(alt_used), key=lambda t: (t.lower(), t)):
            if term in rejected or (results.get(term) or {}).get("alt_of"):
                continue                        # alt forms are listed with their primary: "TT&C (TTC)" (2.26.0)
            d = results.get(term) or (results[alt_used[term][0]] if term in alt_used else {})
            r = rows.get(term)
            if r:
                definition, st = r.get("Definition", ""), r["Status"]
                alts = AF.parse_alts(AF.merge_alts(r.get("Alt forms", ""), sorted(alt_used.get(term, []))))
            else:
                sd = d.get("selected_definition", "")
                definition, st = ("" if _placeholder(sd) else sd), "Not in effort list"
                alts = sorted(alt_used.get(term, []))
            out.append({"term": AF.join_term(term, alts), "key": term, "definition": definition, "status": st,
                        "type": kind(d.get("category", ""))})
        lists[v] = out

    summary = {"volumes": len(volumes), "rule": rule,
               "terms_used": len(all_terms),
               "first_use_problems": sum(1 for r in first_use if r["level"] in (BAD, WARN)),
               "missing": len(missing), "not_approved": len(not_approved), "differs": len(differs), "unused": len(unused),
               "rescan_for_sections": len(no_sections),
               "checks": {lv: sum(1 for r in first_use if r["level"] == lv) for lv in (OK, BAD, WARN, INFO)},
               "used_once": sum(1 for r in first_use if r.get("once")),
               "list_checks": {lv: sum(1 for r in acr_list if r["level"] == lv) for lv in (OK, BAD, WARN, INFO)}}
    return {"summary": summary, "first_use": first_use, "missing": missing, "not_approved": not_approved,
            "differs": differs, "unused": unused, "lists": lists, "volumes": list(volumes), "used": used,
            "acronym_list": acr_list}


def effort_updates(report):
    """{term: {"Used in volumes": ..., "First use": ...}} for terms already in the list."""
    vols = {}
    for v in report["volumes"]:
        for term in report["used"][v]:
            vols.setdefault(term, []).append(v)
    fu = {}
    for r in report["first_use"]:
        where = f"{r['volume']} section {r['section']}" if r.get("section") else r["volume"]
        fu.setdefault(r["term"], []).append(f"{where}: {r['state']}")
    return {t: {"Used in volumes": "; ".join(v), "First use": "; ".join(fu.get(t, []))[:400]} for t, v in vols.items()}


# ------------------------------------------------------------------ writers
def write_excel(report, path, effort_name=""):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    fills = {BAD: "FDE2E2", WARN: "FEF3E2", INFO: "E0F2FE", OK: "DCFCE7"}
    wb = Workbook()
    ws = wb.active; ws.title = "Summary"
    s = report["summary"]
    for r in (["Effort", effort_name], ["Volumes", ", ".join(report["volumes"])],
              ["First use rule", RULE_TEXT.get(s["rule"], s["rule"])],
              ["Terms used", s["terms_used"]], ["First-use problems", s["first_use_problems"]],
              ["Defined at first use (OK)", s.get("checks", {}).get(OK, 0)],
              ["Defined but never used again", s.get("used_once", 0)],
              ["Acronym list problems", sum(s.get("list_checks", {}).get(lv, 0) for lv in (BAD, WARN))],
              ["Checked on", "Word documents (.docx) only; PDF exports are not checked"],
              ["Not in effort list", s["missing"]], ["Not approved yet", s["not_approved"]],
              ["Definition differs", s["differs"]], ["In the list, not used", s["unused"]]):
        ws.append(r)
    ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 90

    def sheet(title, cols, rows, keys, widths, level_key=None):
        sh = wb.create_sheet(title[:31])
        sh.append(cols)
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1E3A8A")
        for row in rows:
            sh.append([row.get(k, "") for k in keys])
            f = fills.get(row.get(level_key)) if level_key else None
            if f:
                for c in sh[sh.max_row]:
                    c.fill = PatternFill("solid", fgColor=f)
        for i, w in enumerate(widths, 1):
            sh.column_dimensions[sh.cell(row=1, column=i).column_letter].width = w
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "B2"
        if sh.max_row > 1:
            sh.auto_filter.ref = f"A1:{sh.cell(row=1, column=len(cols)).column_letter}{sh.max_row}"

    if s["rule"] == "section":
        sheet("First use", ["Check", "Volume", "Section", "Acronym", "Result", "First use in text", "In figures or tables",
                            "Written as", "Effort list definition", "Also defined at", "List status"],
              report["first_use"], ["check", "volume", "section", "term", "state", "loc", "in_figures", "written",
                                    "list_definition", "also_defined", "status"], [9, 22, 12, 12, 46, 40, 40, 34, 34, 30, 14],
              level_key="level")
    else:
        sheet("First use", ["Check", "Volume", "Acronym", "Result", "First use in text", "In figures or tables", "Written as",
                            "Effort list definition", "Also defined at", "List status"],
              report["first_use"], ["check", "volume", "term", "state", "loc", "in_figures", "written", "list_definition",
                                    "also_defined", "status"], [9, 22, 12, 46, 40, 40, 34, 34, 30, 14], level_key="level")
    sheet("Not in effort list", ["Term", "Type", "Volumes", "Count", "Suggested definition", "First found at"],
          report["missing"], ["term", "type", "volumes", "count", "suggestion", "loc"], [14, 13, 30, 8, 40, 40])
    sheet("Not approved", ["Term", "Type", "Status", "Definition", "Volumes", "First found at"],
          report["not_approved"], ["term", "type", "status", "definition", "volumes", "loc"], [14, 13, 14, 40, 30, 40])
    sheet("Definition differs", ["Term", "Volume", "Volume says", "Effort list says", "Where", "Sentence"],
          report["differs"], ["term", "volume", "written", "list_definition", "loc", "sentence"], [14, 22, 36, 36, 36, 60])
    sheet("Unused in volumes", ["Term", "Type", "Status", "Definition"],
          report["unused"], ["term", "type", "status", "definition"], [14, 13, 14, 50])
    if report.get("acronym_list"):
        sheet("Acronym list check", ["Check", "Volume", "Acronym", "Result", "In the list as", "Effort list definition",
                                     "Uses outside the list"],
              report["acronym_list"], ["check", "volume", "term", "state", "listed_as", "list_definition", "uses"],
              [9, 22, 12, 50, 40, 40, 12], level_key="level")
    for i, v in enumerate(report["volumes"], 1):
        sheet(f"List {i} " + re.sub(r"[\[\]:*?/\\]", "", v)[:22], ["Term", "Definition", "Status", "Type"],
              report["lists"][v], ["term", "definition", "status", "type"], [14, 60, 18, 13])
    wb.save(path)


def _x(s):
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;"))


def write_docx(entries, path, title="Acronyms and Abbreviations"):
    """A Word document with a two-column table (Term, Definition), built from the Office XML directly
    (no add-in, no macros). entries: [(term, definition)]."""
    import zipfile
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

    def cell(text, width, bold=False, shade=None):
        rpr = "<w:rPr><w:b/></w:rPr>" if bold else ""
        sh = f'<w:shd w:val="clear" w:color="auto" w:fill="{shade}"/>' if shade else ""
        return (f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>{sh}</w:tcPr>'
                f'<w:p><w:pPr><w:spacing w:before="20" w:after="20"/></w:pPr><w:r>{rpr}<w:t xml:space="preserve">{_x(text)}</w:t></w:r></w:p></w:tc>')

    rows = ['<w:tr><w:trPr><w:tblHeader/></w:trPr>' + cell("Term", 2000, True, "D9E2F3") + cell("Definition", 7360, True, "D9E2F3") + "</w:tr>"]
    rows += ["<w:tr><w:trPr><w:cantSplit/></w:trPr>" + cell(t, 2000) + cell(d, 7360) + "</w:tr>" for t, d in entries]
    border = "".join(f'<w:{b} w:val="single" w:sz="4" w:space="0" w:color="808080"/>' for b in ("top", "left", "bottom", "right", "insideH", "insideV"))
    body = (f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>{_x(title)}</w:t></w:r></w:p>'
            f'<w:tbl><w:tblPr><w:tblW w:w="9360" w:type="dxa"/><w:tblBorders>{border}</w:tblBorders>'
            f'<w:tblLayout w:type="fixed"/></w:tblPr><w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="7360"/></w:tblGrid>'
            + "".join(rows) + '</w:tbl><w:p/><w:sectPr><w:pgSz w:w="12240" w:h="15840"/>'
            '<w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="720" w:footer="720" w:gutter="0"/></w:sectPr>')
    doc = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W}><w:body>{body}</w:body></w:document>'
    styles = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:styles {W}>'
              '<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Calibri"/>'
              '<w:sz w:val="22"/></w:rPr></w:rPrDefault></w:docDefaults>'
              '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>'
              '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/>'
              '<w:pPr><w:keepNext/><w:spacing w:before="240" w:after="120"/><w:outlineLvl w:val="0"/></w:pPr>'
              '<w:rPr><w:b/><w:sz w:val="28"/></w:rPr></w:style></w:styles>')
    ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
          '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    drels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
             '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>')
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)
        z.writestr("word/styles.xml", styles)
        z.writestr("word/_rels/document.xml.rels", drels)
