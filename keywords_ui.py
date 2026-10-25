"""Keyword Harvest (2.16.0): capture keywords while you work, get suggestions with reasons, export for SharePoint.

- Keywords view (third view next to Effort list and Scan review): every keyword of the effort, with category, status
  (Proposed, Keep, Reject), where it came from, why, who and when. Saved in the effort list workbook (Keywords sheet),
  with its History and recovery copy.
- Capture: Ctrl+K in the app (selected preview text, else the clipboard, fills it in); right-click a result, the
  preview text or a home row > Add as keyword; Ctrl+Alt+K anywhere in Windows (copy a phrase in Word, press the keys).
  The system-wide key uses Windows RegisterHotKey (one key combination, not a keyboard hook) and reads the clipboard
  only when you press it.
- Suggest keywords: phrases from the customer documents and our volumes, scored by rules and counts (No AI), with the
  reasons and sample sentences. Not a keyword puts it on the list as Reject, so it is not suggested again.
- Export: Excel or CSV laid out for the SharePoint columns. The app never writes to SharePoint (no network code).
"""
import csv, datetime, os, sys, threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import customtkinter as ctk
import diag
import effort as E

KW_FILTERS = ["All keywords", "Keep", "Proposed", "Reject"]
KW_COLS = (("kw", "Keyword", 220), ("cat", "Category", 130), ("status", "Status", 80), ("source", "Source", 220),
           ("why", "Why", 360), ("who", "Added by", 90), ("date", "Date", 90))
EXPORT_COLS = ["Keyword", "Category", "SharePoint field", "Why", "Source", "Status", "Added by", "Date"]


def _eng():
    import app_gui
    return app_gui


class KeywordsMixin:
    # ------------------------------------------------------------ build
    def kw_build(self, M):
        self.kw_view = K = ctk.CTkFrame(M, fg_color="transparent")
        self.lbl_kw = ctk.CTkLabel(K, text="", anchor="w", font=("Arial", 13, "bold")); self.lbl_kw.pack(fill="x", padx=4, pady=(0, 4))
        row = ctk.CTkFrame(K, fg_color="transparent"); row.pack(fill="x", pady=(0, 4))
        ctk.CTkLabel(row, text="Show").pack(side="left", padx=(4, 4))
        self.kw_filter = tk.StringVar(value=KW_FILTERS[0])
        ctk.CTkOptionMenu(row, variable=self.kw_filter, values=KW_FILTERS, width=150, command=lambda _v: self.kw_refresh()).pack(side="left")
        ctk.CTkLabel(row, text="Category").pack(side="left", padx=(12, 4))
        self.kw_cat = tk.StringVar(value="All")
        ctk.CTkOptionMenu(row, variable=self.kw_cat, values=["All"] + E.KW_CATEGORIES, width=170,
                          command=lambda _v: self.kw_refresh()).pack(side="left")
        tf = tk.Frame(K); tf.pack(fill="both", expand=True)
        tf.grid_rowconfigure(0, weight=1); tf.grid_columnconfigure(0, weight=1)
        self.kw_tree = tv = ttk.Treeview(tf, columns=[c[0] for c in KW_COLS], show="headings", selectmode="extended")
        for cid, head, w in KW_COLS:
            tv.heading(cid, text=head); tv.column(cid, width=w, minwidth=60, stretch=cid in ("kw", "why"))
        tv.tag_configure("Keep", background="#E7F6EC")
        tv.tag_configure("Reject", background="#EEEEEE", foreground="#6B7280")
        vsb = ttk.Scrollbar(tf, orient="vertical", command=tv.yview); tv.configure(yscrollcommand=vsb.set)
        tv.grid(row=0, column=0, sticky="nsew"); vsb.grid(row=0, column=1, sticky="ns")
        tv.bind("<Double-1>", lambda _e: self.kw_edit())
        tv.bind("<<TreeviewSelect>>", lambda _e: self._kw_buttons())
        b = ctk.CTkFrame(K, fg_color="transparent"); b.pack(fill="x", pady=(6, 0))
        k = dict(height=30)
        self.btn_kw_add = ctk.CTkButton(b, text="Add keyword... (Ctrl+K)", width=180, fg_color="#0F766E", command=self.kw_add_dialog, **k)
        self.btn_kw_keep = ctk.CTkButton(b, text="Keep", width=70, fg_color="#15803D", command=lambda: self.kw_set_status("Keep"), **k)
        self.btn_kw_rej = ctk.CTkButton(b, text="Reject", width=70, fg_color="#B91C1C", command=lambda: self.kw_set_status("Reject"), **k)
        self.btn_kw_edit = ctk.CTkButton(b, text="Edit...", width=70, fg_color="#1E3A8A", command=self.kw_edit, **k)
        self.btn_kw_del = ctk.CTkButton(b, text="Remove", width=80, fg_color="#475569", command=self.kw_remove, **k)
        self.btn_kw_cat = ctk.CTkButton(b, text="Add to catalog", width=110, fg_color="#475569", command=self.kw_to_catalog, **k)
        for w in (self.btn_kw_add, self.btn_kw_keep, self.btn_kw_rej, self.btn_kw_edit, self.btn_kw_del, self.btn_kw_cat):
            w.pack(side="left", padx=(0, 6))
        b2 = ctk.CTkFrame(K, fg_color="transparent"); b2.pack(fill="x", pady=(6, 0))   # second row: fits at 1280 px
        ctk.CTkButton(b2, text="Suggest keywords...", width=150, fg_color="#7C3AED", command=self.kw_suggest_dialog, **k).pack(side="left", padx=(0, 6))
        ctk.CTkButton(b2, text="Attach offerings...", width=140, fg_color="#B45309", command=self.attach_offerings_dialog, **k).pack(side="left", padx=(0, 6))
        ctk.CTkButton(b2, text="Offerings catalog...", width=150, fg_color="#475569", command=self.offerings_window, **k).pack(side="left", padx=(0, 6))
        ctk.CTkButton(b2, text="Past proposals...", width=130, fg_color="#475569", command=self.kw_baseline_dialog, **k).pack(side="left", padx=(0, 6))
        ctk.CTkButton(b2, text="Export...", width=90, fg_color="#0F766E", command=self.kw_export_dialog, **k).pack(side="right")
        self.bind_all("<Control-k>", lambda _e: (self.kw_quick(), "break")[1])
        self.bind_all("<Control-K>", lambda _e: (self.kw_quick(), "break")[1])
        self._kw_hotkey = None
        self.after(1200, self.kw_start_hotkey)
        self._kw_buttons()

    # ------------------------------------------------------------ list
    def kw_selected(self):
        return [self.kw_tree.set(i, "kw") for i in self.kw_tree.selection()] if hasattr(self, "kw_tree") else []

    def _kw_buttons(self):
        n = len(self.kw_selected())
        for w in (self.btn_kw_keep, self.btn_kw_rej, self.btn_kw_del, self.btn_kw_cat):
            w.configure(state="normal" if n else "disabled")
        self.btn_kw_edit.configure(state="normal" if n == 1 else "disabled")
        self.btn_kw_add.configure(state="normal" if self.effort is not None else "disabled")

    def kw_refresh(self):
        if not hasattr(self, "kw_tree"):
            return
        tv = self.kw_tree
        keep = set(self.kw_selected())
        tv.delete(*tv.get_children())
        L = self.effort
        if L is None:
            self.lbl_kw.configure(text=""); self._kw_buttons(); return
        f, c = self.kw_filter.get(), self.kw_cat.get()
        for i, k in enumerate(L.keywords):
            if f != KW_FILTERS[0] and k["Status"] != f:
                continue
            if c != "All" and k["Category"] != c:
                continue
            iid = f"kw{i}"
            tv.insert("", "end", iid=iid, tags=(k["Status"],),
                      values=(k["Keyword"], k["Category"], k["Status"], k["Source"], k["Why"], k["Added by"], k["Date"]))
            if k["Keyword"] in keep:
                tv.selection_add(iid)
        n = {s: sum(k["Status"] == s for k in L.keywords) for s in E.KW_STATUSES}
        self.lbl_kw.configure(text=f"{L.name}:  {len(L.keywords)} keywords  |  {n['Keep']} keep  |  {n['Proposed']} proposed  |  "
                                   f"{n['Reject']} rejected   (saved in the effort list's Keywords sheet)")
        self._kw_buttons()

    def _kw_saved(self):
        self._save_effort_list()
        self.kw_refresh()

    # ------------------------------------------------------------ actions (also used by tests)
    def kw_add(self, text, category="Other", source="", why="", status="Proposed"):
        if self.effort is None:
            return "no effort"
        r = self.effort.add_keyword(text, category, source, why, status)
        if r == "added":
            diag.activity("keyword_added")
            self._kw_saved()
        return r

    def kw_set_status(self, status, words=None):
        import keywords
        words = words or self.kw_selected()
        n = 0
        for w in (words if self.effort is not None else []):
            if self.effort.update_keyword(w, Status=status):
                n += 1
                keywords.learn(w, status)            # later suggestions in every effort (hashed, counts only)
        if n:
            diag.activity("keyword_status", n=n)
            self._kw_saved()
        return n

    def kw_remove(self, words=None):
        ask = words is None
        words = words or self.kw_selected()
        if words and (not ask or messagebox.askyesno("Remove", f"Remove {len(words)} keyword(s)?")):
            for w in words:
                self.effort.remove_keyword(w)
            diag.activity("keyword_removed", n=len(words))
            self._kw_saved()

    def kw_edit(self):
        sel = self.kw_selected()
        if len(sel) == 1:
            return self.kw_add_dialog(edit=self.effort.keyword(sel[0]))
        return None

    # ------------------------------------------------------------ capture
    def kw_grab_text(self):
        """Selected preview text, else a short clipboard text, else ''."""
        try:
            t = self.preview.selected_text() if getattr(self, "preview", None) is not None else ""
        except Exception:
            diag.swallowed(); t = ""
        if not t.strip():
            try:
                t = self.clipboard_get()
            except tk.TclError:
                t = ""
        t = " ".join((t or "").split())
        return t if len(t) <= 120 else ""

    def kw_quick(self, text=None, source=""):
        if self.effort is None:
            self.status.configure(text="Open an effort to collect keywords."); return None
        return self.kw_add_dialog(text=self.kw_grab_text() if text is None else text, source=source)

    def kw_add_dialog(self, text="", source="", edit=None):
        if self.effort is None:
            return None
        win = ctk.CTkToplevel(self); win.title("Edit keyword" if edit else "Add keyword"); win.geometry("560x330")
        win.transient(self); win.attributes("-topmost", True); win.after(150, win.lift)
        diag.register(win, "keyword_add")
        ctk.CTkLabel(win, text="Keyword (a phrase someone would search for)", anchor="w").pack(fill="x", padx=14, pady=(12, 0))
        e = ctk.CTkEntry(win, width=520); e.pack(padx=14, pady=2); e.insert(0, (edit or {}).get("Keyword", text))
        ctk.CTkLabel(win, text="Category", anchor="w").pack(fill="x", padx=14, pady=(8, 0))
        cat = ctk.CTkOptionMenu(win, values=E.KW_CATEGORIES, width=220); cat.set((edit or {}).get("Category", "Other")); cat.pack(anchor="w", padx=14)
        ctk.CTkLabel(win, text="Why (optional)", anchor="w").pack(fill="x", padx=14, pady=(8, 0))
        why = ctk.CTkEntry(win, width=520); why.pack(padx=14, pady=2); why.insert(0, (edit or {}).get("Why", ""))
        ctk.CTkLabel(win, text="SharePoint field (optional)", anchor="w").pack(fill="x", padx=14, pady=(8, 0))
        fld = ctk.CTkEntry(win, width=260); fld.pack(anchor="w", padx=14, pady=2); fld.insert(0, (edit or {}).get("SharePoint field", ""))
        msg = ctk.CTkLabel(win, text="", text_color="#B45309"); msg.pack(anchor="w", padx=14)

        def ok(_e=None):
            t = e.get().strip()
            if edit:
                if t and t != edit["Keyword"]:
                    if self.effort.keyword(t):
                        msg.configure(text="That keyword is already on the list."); return
                    edit["Keyword"] = t
                self.effort.update_keyword(edit["Keyword"], Category=cat.get(), Why=why.get().strip(), **{"SharePoint field": fld.get().strip()})
                self._kw_saved(); win.destroy(); return
            r = self.effort.add_keyword(t, cat.get(), source or "Typed in", why.get().strip(), field=fld.get().strip())
            if r == "exists":
                msg.configure(text="That keyword is already on the list."); return
            if r == "empty":
                msg.configure(text="Type a keyword."); return
            diag.activity("keyword_added")
            self._kw_saved(); win.destroy()
            self.status.configure(text=f"Keyword added: {t}")
        b = ctk.CTkFrame(win, fg_color="transparent"); b.pack(fill="x", padx=14, pady=10)
        ctk.CTkButton(b, text="Save" if edit else "Add", width=90, command=ok).pack(side="left")
        ctk.CTkButton(b, text="Cancel", width=90, fg_color="#475569", command=win.destroy).pack(side="right")
        e.bind("<Return>", ok); e.focus_set()
        win._api = {"entry": e, "cat": cat, "why": why, "field": fld, "ok": ok, "msg": msg}
        return win

    # ------------------------------------------------------------ Ctrl+Alt+K anywhere in Windows
    def kw_start_hotkey(self):
        if sys.platform != "win32" or os.environ.get("AH_TEST") or self._kw_hotkey is not None:
            return False
        import ctypes
        from ctypes import wintypes
        user32, kernel32 = ctypes.windll.user32, ctypes.windll.kernel32
        MOD_ALT, MOD_CONTROL, MOD_NOREPEAT, WM_HOTKEY, VK_K = 0x1, 0x2, 0x4000, 0x0312, 0x4B
        state = {"tid": None, "ok": False}

        def loop():
            state["tid"] = kernel32.GetCurrentThreadId()
            if not user32.RegisterHotKey(None, 1, MOD_CONTROL | MOD_ALT | MOD_NOREPEAT, VK_K):
                return                                  # another program owns Ctrl+Alt+K: the in-app Ctrl+K still works
            state["ok"] = True
            msg = wintypes.MSG()
            while user32.GetMessageW(ctypes.byref(msg), None, 0, 0) > 0:
                if msg.message == WM_HOTKEY:
                    try:
                        self.after(0, self._kw_hotkey_fired)
                    except Exception:
                        break
            user32.UnregisterHotKey(None, 1)
        t = threading.Thread(target=loop, daemon=True, name="keyword-hotkey")
        t.start()
        self._kw_hotkey = state
        return True

    def kw_stop_hotkey(self):
        st = getattr(self, "_kw_hotkey", None)
        if st and st.get("tid") and sys.platform == "win32":
            import ctypes
            ctypes.windll.user32.PostThreadMessageW(st["tid"], 0x0012, 0, 0)    # WM_QUIT ends the loop and unregisters
        self._kw_hotkey = None

    def _kw_hotkey_fired(self):
        diag.activity("keyword_hotkey")
        try:
            self.deiconify(); self.lift()
        except Exception:
            diag.swallowed()
        self.kw_quick(source="Copied (Ctrl+Alt+K)")

    # ------------------------------------------------------------ suggestions
    def kw_sources(self):
        """[(path, 'customer' or 'ours')]: the customer documents, and our volumes / scanned files."""
        L = self.effort
        info = (self.efforts.get(L.name) or {}) if L is not None else {}
        out, seen = [], set()
        cd = info.get("customer_dir") or ""
        if cd and os.path.isdir(cd):
            for f in sorted(os.listdir(cd)):
                p = os.path.join(cd, f)
                if os.path.isfile(p) and f.lower().endswith((".docx", ".pdf", ".pptx", ".xlsx")):
                    out.append((p, "customer")); seen.add(os.path.normcase(os.path.abspath(p)))
        cust_names = {os.path.basename(p) for p, _k in out}
        ours = []
        for v in (self.volumes() if hasattr(self, "volumes") else []):
            ours += [self.path_for(f) for f in v.get("files", [])]
        ours += list((getattr(self, "last_scan", None) or {}).get("paths", []))
        for p in ours:
            if p and os.path.isfile(p) and os.path.normcase(os.path.abspath(p)) not in seen and os.path.basename(p) not in cust_names:
                out.append((p, "ours")); seen.add(os.path.normcase(os.path.abspath(p)))
        return out

    def kw_compute(self, sources=None, progress=None):
        import keywords
        eng = _eng()
        chunks = []
        for i, (p, kind) in enumerate(sources if sources is not None else self.kw_sources(), 1):
            if progress:
                progress(f"Reading {i}: {os.path.basename(p)}")
            try:
                chunks += [(lbl, txt, kind) for lbl, txt in eng.typed_chunks(p)]
            except Exception:
                diag.swallowed()
        L = self.effort
        acr = [(t, r["Definition"]) for t, r in L.rows.items() if r["Status"] == "Approved" and r.get("Definition")] if L else []
        import offerings
        return keywords.suggest(chunks, acr, [k["Keyword"] for k in (L.keywords if L else [])], offerings=offerings.load(),
                                base=keywords.baseline(), learned=keywords.learning())

    def kw_suggest_dialog(self, sources=None, wait=False):
        if self.effort is None:
            return None
        win = ctk.CTkToplevel(self); win.title("Suggested keywords"); win.geometry("1000x620"); win.transient(self); win.after(150, win.lift)
        diag.register(win, "keyword_suggest")
        head = ctk.CTkLabel(win, text="Reading the customer documents and our volumes (typed text only)...", anchor="w")
        head.pack(fill="x", padx=12, pady=(10, 4))
        body = ctk.CTkFrame(win, fg_color="transparent"); body.pack(fill="both", expand=True, padx=12)
        tf = tk.Frame(body); tf.pack(side="left", fill="both", expand=True)
        tf.grid_rowconfigure(0, weight=1); tf.grid_columnconfigure(0, weight=1)
        tv = ttk.Treeview(tf, columns=("phrase", "bucket", "score", "why"), show="headings", selectmode="extended")
        for c, h, w in (("phrase", "Phrase", 230), ("bucket", "Kind", 130), ("score", "Score", 60), ("why", "Why", 420)):
            tv.heading(c, text=h); tv.column(c, width=w, stretch=c in ("phrase", "why"))
        vsb = ttk.Scrollbar(tf, orient="vertical", command=tv.yview); tv.configure(yscrollcommand=vsb.set)
        tv.grid(row=0, column=0, sticky="nsew"); vsb.grid(row=0, column=1, sticky="ns")
        side = ctk.CTkTextbox(body, width=320, wrap="word"); side.pack(side="left", fill="y", padx=(8, 0))
        st = {"items": []}

        def show(_e=None):
            sel = tv.selection()
            side.configure(state="normal"); side.delete("1.0", "end")
            if len(sel) == 1:
                it = st["items"][int(sel[0])]
                side.insert("1.0", it["phrase"] + "\n\n" + "\n".join("- " + w for w in it["why"]) + "\n\n" +
                            "\n\n".join(f"{lbl}:\n{s}" for lbl, s in it["samples"]))
            side.configure(state="disabled")
        tv.bind("<<TreeviewSelect>>", show)

        def fill(items):
            st["items"] = items
            tv.delete(*tv.get_children())
            for i, it in enumerate(items):
                tv.insert("", "end", iid=str(i), values=(it["phrase"], it["bucket"], it["score"], "; ".join(it["why"])))
            head.configure(text=f"{len(items)} suggestions from {len(srcs)} file(s). Select and add; Not a keyword keeps it from "
                                "coming back. Rules and counts only (No AI).")

        def picked():
            return [st["items"][int(i)] for i in tv.selection()]

        def add():
            got = 0
            for it in picked():
                cat = it["bucket"] if it["bucket"] in E.KW_CATEGORIES else "Other"
                got += self.effort.add_keyword(it["phrase"], cat, "Suggested", "; ".join(it["why"])) == "added"
            if got:
                diag.activity("keywords_suggested_added", n=got)
                self._kw_saved()
            for i in tv.selection():
                tv.delete(i)
            head.configure(text=f"Added {got} keyword(s) as Proposed.")
            return got

        def reject():
            got = 0
            import keywords
            for it in picked():
                if self.effort.add_keyword(it["phrase"], "Other", "Suggested", "Not a keyword", status="Reject") == "added":
                    got += 1
                    keywords.learn(it["phrase"], "Reject")
            if got:
                self._kw_saved()
            for i in tv.selection():
                tv.delete(i)
            return got

        b = ctk.CTkFrame(win, fg_color="transparent"); b.pack(fill="x", padx=12, pady=10)
        ctk.CTkButton(b, text="Add selected as keywords", width=200, fg_color="#0F766E", command=add).pack(side="left")
        ctk.CTkButton(b, text="Not a keyword", width=130, fg_color="#B91C1C", command=reject).pack(side="left", padx=6)
        ctk.CTkButton(b, text="Close", width=90, fg_color="#475569", command=win.destroy).pack(side="right")
        srcs = sources if sources is not None else self.kw_sources()
        win._api = {"tree": tv, "add": add, "reject": reject, "st": st, "head": head}
        if not srcs:
            head.configure(text="No documents yet: add customer documents to the effort or scan your volumes first.")
            return win
        diag.activity("keywords_suggest", files=len(srcs))
        if wait:                                   # tests
            fill(self.kw_compute(srcs)); return win

        def work():
            try:
                items = self.kw_compute(srcs, progress=lambda m: self.after(0, lambda: win.winfo_exists() and head.configure(text=m)))
            except Exception:
                diag.swallowed(); items = []
            self.after(0, lambda: win.winfo_exists() and fill(items))
        threading.Thread(target=work, daemon=True).start()
        return win

    # ------------------------------------------------------------ export
    def kw_export(self, path, keep_only=True):
        L = self.effort
        rows = [k for k in L.keywords if (k["Status"] == "Keep" if keep_only else k["Status"] != "Reject")]
        if path.lower().endswith(".csv"):
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f); w.writerow(EXPORT_COLS); w.writerows([[k.get(c, "") for c in EXPORT_COLS] for k in rows])
        else:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            wb = Workbook(); ws = wb.active; ws.title = "Keywords"
            ws.append(EXPORT_COLS)
            for c in ws[1]:
                c.font = Font(bold=True)
            for k in rows:
                ws.append([k.get(c, "") for c in EXPORT_COLS])
            for col, wd in zip("ABCDEFGH", (34, 18, 18, 60, 30, 10, 14, 12)):
                ws.column_dimensions[col].width = wd
            ab = wb.create_sheet("About")
            for r in (["Effort", L.name], ["Client", L.client], ["Solicitation", L.solicitation],
                      ["Exported", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")]):
                ab.append(r)
            wb.save(path)
        diag.activity("keywords_exported", rows=len(rows))
        return len(rows)

    def kw_export_dialog(self):
        L = self.effort
        if L is None:
            return None
        keep = [k for k in L.keywords if k["Status"] == "Keep"]
        keep_only = bool(keep) and messagebox.askyesno("Export keywords", f"Export only the {len(keep)} keyword(s) marked Keep?\n\n"
                                                       "No exports Keep and Proposed (never Reject).")
        name = f"{L.name.replace(' ', '_')}_Keywords_{datetime.date.today().isoformat()}.xlsx"
        path = filedialog.asksaveasfilename(title="Keywords for SharePoint", initialfile=name, defaultextension=".xlsx",
                                            filetypes=[("Excel", "*.xlsx"), ("CSV", "*.csv")])
        if not path:
            return None
        n = self.kw_export(path, keep_only)
        messagebox.showinfo("Export keywords", f"{n} keyword(s) saved to {os.path.basename(path)}. Copy them into the SharePoint "
                            "keyword fields; the app does not connect to SharePoint.")
        return path

    # ------------------------------------------------------------ offerings catalog (2.17.0)
    def kw_to_catalog(self, words=None):
        """Put selected keywords in the offerings catalog (and mark them Offering here)."""
        import offerings
        words = words or self.kw_selected()
        added = 0
        for w in words:
            k = self.effort.keyword(w) if self.effort is not None else None
            if offerings.add(w, description=(k or {}).get("Why", "") if (k or {}).get("Source") != "Suggested" else "") == "added":
                added += 1
            if k and k["Category"] != "Offering":
                self.effort.update_keyword(w, Category="Offering")
        if words:
            diag.activity("offerings_from_keywords", n=added)
            self._kw_saved()
            self.status.configure(text=f"Offerings catalog: {added} added, {len(words) - added} already there.")
        return added

    def attach_offerings(self, names):
        """Attach catalog offerings to this effort as keywords (category Offering). Returns how many were added."""
        import offerings
        if self.effort is None:
            return 0
        items = offerings.load()
        got = 0
        for n in names:
            it = offerings.find(items, n)
            if it:
                got += self.effort.add_keyword(it["Offering"], "Offering", "Offerings catalog", it.get("Description", "")) == "added"
        if got:
            diag.activity("offerings_attached", n=got)
            self._kw_saved()
        return got

    def attach_offerings_dialog(self):
        import offerings
        if self.effort is None:
            return None
        items = offerings.load()
        if not items:
            if messagebox.askyesno("Attach offerings", "The offerings catalog is empty. Open it to add your offerings?"):
                return self.offerings_window()
            return None
        win = ctk.CTkToplevel(self); win.title("Attach offerings"); win.geometry("760x520"); win.transient(self); win.after(150, win.lift)
        diag.register(win, "offerings_attach")
        ctk.CTkLabel(win, text=f"Pick the offerings this proposal uses. They join {self.effort.name}'s keywords as Offering.",
                     anchor="w").pack(fill="x", padx=12, pady=(12, 4))
        q = tk.StringVar(value="")
        ctk.CTkEntry(win, textvariable=q, width=300, placeholder_text="Search").pack(anchor="w", padx=12)
        tf = tk.Frame(win); tf.pack(fill="both", expand=True, padx=12, pady=6)
        tf.grid_rowconfigure(0, weight=1); tf.grid_columnconfigure(0, weight=1)
        tv = ttk.Treeview(tf, columns=("name", "group", "desc", "here"), show="headings", selectmode="extended")
        for c, h, w in (("name", "Offering", 220), ("group", "Group", 120), ("desc", "Description", 300), ("here", "In this effort", 100)):
            tv.heading(c, text=h); tv.column(c, width=w, stretch=c in ("name", "desc"))
        vsb = ttk.Scrollbar(tf, orient="vertical", command=tv.yview); tv.configure(yscrollcommand=vsb.set)
        tv.grid(row=0, column=0, sticky="nsew"); vsb.grid(row=0, column=1, sticky="ns")

        def fill(*_):
            tv.delete(*tv.get_children())
            t = q.get().strip().casefold()
            for i, it in enumerate(items):
                if t and t not in " ".join(it.values()).casefold():
                    continue
                tv.insert("", "end", iid=str(i), values=(it["Offering"], it["Group"], it["Description"],
                                                         "Yes" if self.effort.keyword(it["Offering"]) else ""))
        q.trace_add("write", fill)

        def attach():
            n = self.attach_offerings([items[int(i)]["Offering"] for i in tv.selection()])
            fill(); msg.configure(text=f"Attached {n} offering(s).")
            return n
        b = ctk.CTkFrame(win, fg_color="transparent"); b.pack(fill="x", padx=12, pady=10)
        ctk.CTkButton(b, text="Attach selected", width=140, fg_color="#B45309", command=attach).pack(side="left")
        msg = ctk.CTkLabel(b, text=""); msg.pack(side="left", padx=10)
        ctk.CTkButton(b, text="Close", width=90, fg_color="#475569", command=win.destroy).pack(side="right")
        fill()
        win._api = {"tree": tv, "attach": attach, "items": items, "search": q}
        return win

    def offerings_window(self):
        """The PC's offerings catalog: add, edit, remove, import a list, export to share."""
        import offerings
        if getattr(self, "_off_win", None) is not None and self._off_win.winfo_exists():
            self._off_win.lift(); return self._off_win
        win = ctk.CTkToplevel(self); win.title("Offerings catalog"); win.geometry("860x560"); win.transient(self); win.after(150, win.lift)
        self._off_win = win
        diag.register(win, "offerings")
        ctk.CTkLabel(win, text=("Sierra Space offerings, kept once on this PC for every effort. Attach them to an effort's keywords "
                                "(Keywords > Attach offerings); suggestions mark an offering when the documents name it."),
                     anchor="w", justify="left", wraplength=820).pack(fill="x", padx=12, pady=(12, 4))
        tf = tk.Frame(win); tf.pack(fill="both", expand=True, padx=12, pady=6)
        tf.grid_rowconfigure(0, weight=1); tf.grid_columnconfigure(0, weight=1)
        tv = ttk.Treeview(tf, columns=offerings.COLUMNS, show="headings", selectmode="extended")
        for c, w in zip(offerings.COLUMNS, (220, 200, 120, 320)):
            tv.heading(c, text=c); tv.column(c, width=w, stretch=c in ("Offering", "Description"))
        vsb = ttk.Scrollbar(tf, orient="vertical", command=tv.yview); tv.configure(yscrollcommand=vsb.set)
        tv.grid(row=0, column=0, sticky="nsew"); vsb.grid(row=0, column=1, sticky="ns")
        form = ctk.CTkFrame(win); form.pack(fill="x", padx=12)
        ents = {}
        for i, (c, w) in enumerate(zip(offerings.COLUMNS, (200, 180, 120, 260))):
            ctk.CTkLabel(form, text=c).grid(row=0, column=i, sticky="w", padx=4, pady=(6, 0))
            e = ctk.CTkEntry(form, width=w); e.grid(row=1, column=i, padx=4, pady=(0, 8)); ents[c] = e
        msg = ctk.CTkLabel(win, text="", text_color="#0F766E", anchor="w"); msg.pack(fill="x", padx=12)

        def fill():
            tv.delete(*tv.get_children())
            for i, it in enumerate(offerings.load()):
                tv.insert("", "end", iid=it["Offering"], values=[it[c] for c in offerings.COLUMNS])

        def pick(_e=None):
            sel = tv.selection()
            if len(sel) == 1:
                vals = tv.item(sel[0], "values")
                for c, v in zip(offerings.COLUMNS, vals):
                    ents[c].delete(0, "end"); ents[c].insert(0, v)

        def add_update():
            vals = {c: ents[c].get().strip() for c in offerings.COLUMNS}
            if not vals["Offering"]:
                msg.configure(text="Type the offering's name."); return None
            if offerings.find(offerings.load(), vals["Offering"]):
                offerings.update(vals["Offering"], **vals); msg.configure(text=f"Updated {vals['Offering']}.")
            else:
                offerings.add(vals["Offering"], vals["Other names"], vals["Group"], vals["Description"]); msg.configure(text=f"Added {vals['Offering']}.")
                diag.activity("offering_added")
            for e in ents.values():
                e.delete(0, "end")
            fill()
            return vals["Offering"]

        def remove():
            sel = list(tv.selection())
            if sel and messagebox.askyesno("Remove", f"Remove {len(sel)} offering(s) from the catalog? Efforts keep their keywords.", parent=win):
                offerings.remove(sel); fill()

        def imp(path=None):
            path = path or filedialog.askopenfilename(parent=win, title="Offerings list", filetypes=[
                ("Lists", "*.csv *.xlsx *.docx *.pdf *.txt"), ("All files", "*.*")])
            if not path:
                return None
            try:
                a, same = offerings.import_file(path)
            except Exception as ex:
                diag.swallowed(); msg.configure(text=f"Could not read {os.path.basename(path)} ({type(ex).__name__})."); return None
            diag.activity("offerings_imported", added=a)
            fill(); msg.configure(text=f"{a} offering(s) added, {same} already there.")
            return a, same

        def exp(path=None):
            path = path or filedialog.asksaveasfilename(parent=win, title="Share the offerings catalog", initialfile="Offerings_Catalog.csv",
                                                        defaultextension=".csv", filetypes=[("CSV", "*.csv")])
            if not path:
                return None
            import shutil
            offerings.save(offerings.load())
            shutil.copyfile(offerings.path(), path)
            msg.configure(text=f"Saved {os.path.basename(path)}. Others import it with Import list.")
            return path

        tv.bind("<<TreeviewSelect>>", pick)
        b = ctk.CTkFrame(win, fg_color="transparent"); b.pack(fill="x", padx=12, pady=10)
        ctk.CTkButton(b, text="Add / Update", width=120, fg_color="#0F766E", command=add_update).pack(side="left")
        ctk.CTkButton(b, text="Remove", width=90, fg_color="#B91C1C", command=remove).pack(side="left", padx=6)
        ctk.CTkButton(b, text="Import list...", width=110, fg_color="#475569", command=imp).pack(side="left")
        ctk.CTkButton(b, text="Export to share...", width=130, fg_color="#475569", command=exp).pack(side="left", padx=6)
        ctk.CTkButton(b, text="Close", width=90, fg_color="#475569", command=win.destroy).pack(side="right")
        fill()
        win._api = {"tree": tv, "ents": ents, "add": add_update, "import": imp, "export": exp, "remove": remove, "msg": msg}
        return win

    # ------------------------------------------------------------ capture from any text box (2.18.0)
    def kw_text_menu(self, event, text_widget, source):
        try:
            sel = text_widget.get("sel.first", "sel.last") if text_widget.tag_ranges("sel") else ""
        except tk.TclError:
            sel = ""
        sel = " ".join(sel.split())
        m = tk.Menu(self, tearoff=0)
        m.add_command(label=(f"Add as keyword: {sel[:50]}" if sel else "Add as keyword (select words first)"),
                      state="normal" if sel and self.effort is not None else "disabled",
                      command=lambda: self.kw_quick(sel, source=source))
        m.tk_popup(event.x_root, event.y_root)
        return "break"

    # ------------------------------------------------------------ past proposals baseline (2.18.0)
    def kw_build_baseline(self, folder, progress=None):
        """Read the typed text of every past proposal in a folder (one file = one proposal) and save phrase counts."""
        import keywords
        eng = _eng()
        files = []
        for root, _d, fs in os.walk(folder):
            files += [os.path.join(root, f) for f in fs if f.lower().endswith((".docx", ".pdf", ".pptx")) and not f.startswith("~$")]
        docs = []
        for i, p in enumerate(sorted(files), 1):
            if progress:
                progress(f"Reading {i} of {len(files)}")
            try:
                docs.append([t for _l, t in eng.typed_chunks(p)])
            except Exception:
                diag.swallowed()
        data = keywords.build_baseline(docs)
        diag.activity("keyword_baseline_built", files=len(docs))
        return data

    def kw_baseline_dialog(self):
        import keywords
        win = ctk.CTkToplevel(self); win.title("Past proposals"); win.geometry("620x320"); win.transient(self); win.after(150, win.lift)
        diag.register(win, "keyword_baseline")

        def state_text():
            b = keywords.baseline()
            return (f"Baseline: {b['docs']} past proposals, built {b.get('built', '')}." if b.get("docs") else "No baseline yet.")
        ctk.CTkLabel(win, text=("Pick a folder with past submissions (one file per proposal, Word, PDF or PowerPoint; 5 to 20 is a good "
                                "start). Suggestions then score phrases every proposal uses lower and phrases new to this one higher. "
                                "Only hashed phrases with counts are saved, never the text."), wraplength=580, justify="left",
                     anchor="w").pack(fill="x", padx=14, pady=(12, 6))
        lbl = ctk.CTkLabel(win, text=state_text(), anchor="w", font=("Segoe UI", 13, "bold")); lbl.pack(fill="x", padx=14, pady=6)

        def build(folder=None):
            folder = folder or filedialog.askdirectory(parent=win, title="Folder with past submissions")
            if not folder:
                return None
            lbl.configure(text="Reading...")
            def work():
                data = self.kw_build_baseline(folder, progress=lambda m: self.after(0, lambda: win.winfo_exists() and lbl.configure(text=m)))
                self.after(0, lambda: win.winfo_exists() and lbl.configure(text=state_text()))
                return data
            threading.Thread(target=work, daemon=True).start()
            return True

        def clear():
            keywords.clear_baseline(); lbl.configure(text=state_text())
        b = ctk.CTkFrame(win, fg_color="transparent"); b.pack(fill="x", padx=14, pady=12)
        ctk.CTkButton(b, text="Build from a folder...", width=170, fg_color="#0F766E", command=build).pack(side="left")
        ctk.CTkButton(b, text="Clear", width=80, fg_color="#475569", command=clear).pack(side="left", padx=6)
        ctk.CTkButton(b, text="Close", width=90, fg_color="#475569", command=win.destroy).pack(side="right")
        win._api = {"build": build, "clear": clear, "label": lbl}
        return win
