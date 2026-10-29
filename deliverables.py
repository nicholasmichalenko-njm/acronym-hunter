"""Deliverables (2.41.0): what the effort owes the customer, and where each stands.

- Suggested from the submission requirements (Attachment J-2, SF 33, Volume I...) found in the customer documents;
  the team adds each one or marks it not ours. Others are typed in.
- Each deliverable: name, kind, format, limit and rules, owner, status, notes, our working file, the customer
  form or template it is filled on (a document in the version library), and done marks per stage.
- What changed it: the requirement or page limit behind it changed in a newer customer version and is not
  reviewed yet; a newer version of its form arrived (fill the new one; the old one stays for the record).
- Page count of our Word file (as Word last saved it) against its limit.
Kept in the effort record (record.json), so the effort package carries it. Rules only (No AI); nothing leaves the PC.
"""
import datetime, os, re
import subreq as SR

KINDS = ("Volume", "Attachment", "Form", "Plan", "Presentation", "Other")
STATUSES = ("Not started", "Drafting", "In review", "Final", "Submitted")


def kind_of(name):
    n = (name or "").strip()
    if re.match(r"volume\b|.*\bvolume$", n, re.I):
        return "Volume"
    if re.match(r"(attachment|exhibit|annex|enclosure)\b", n, re.I):
        return "Attachment"
    if re.match(r"(sf|dd|standard form)\b", n, re.I) or re.search(r"representations|certifications", n, re.I):
        return "Form"
    if re.search(r"\bplan\b", n, re.I):
        return "Plan"
    if re.search(r"\b(slides|oral|presentation)\b", n, re.I):
        return "Presentation"
    return "Other"


def default_format(kind, rows=()):
    ff = next((r["value"] for r in rows if r["kind"] == "File" and r["subject"] == "File format"), "")
    if kind == "Volume":
        return "Word, sent as " + (ff or "PDF")
    if kind == "Form":
        return "Signed PDF"
    if kind == "Presentation":
        return "PowerPoint"
    return ff


def _now():
    return datetime.datetime.now().isoformat(timespec="minutes")


def _clean(s):
    return " ".join((s or "").split())


def items(rec):
    return rec.data.setdefault("deliverables", [])


def get(rec, did):
    return next((d for d in items(rec) if d["id"] == did), None)


def suggestions(rec, rows):
    """Deliverables the customer documents name that the team has not added or dismissed yet."""
    have = {d.get("req_key") for d in items(rec) if d.get("req_key")}
    names = {d["name"].casefold() for d in items(rec)}
    gone = set(rec.data.get("deliv_dismissed", []))
    return [r for r in rows if r["kind"] == "Deliverable" and not r.get("hidden") and r["key"] not in have
            and r["key"] not in gone and r["subject"].casefold() not in names]


def add(rec, name, kind="", req_key="", fmt="", limit="", owner="", doc_id="", form_version="", file="", notes="", who=None):
    name = _clean(name)
    if not name:
        return None
    import effortrec as ER
    d = {"id": ER._id(), "name": name, "kind": kind if kind in KINDS else kind_of(name), "req_key": req_key, "format": _clean(fmt),
         "limit": _clean(limit), "owner": _clean(owner), "status": "Not started", "doc_id": doc_id, "form_version": form_version,
         "file": file, "notes": (notes or "").strip(), "stages_done": {}, "created": _now(), "by": ER._who(who)}
    items(rec).append(d)
    rec.log("deliverable", f"Deliverable added: {name}" + (f" (owner {d['owner']})" if d["owner"] else "") +
            (" from the customer documents" if req_key else ""), ref={"deliverable": d["id"]}, who=who)
    rec.save()
    return d["id"]


def dismiss(rec, req_key, label="", who=None):
    gone = rec.data.setdefault("deliv_dismissed", [])
    if req_key not in gone:
        gone.append(req_key)
        rec.log("deliverable", f"Not ours: {label or req_key.partition(' | ')[2]}", who=who)
        rec.save()
    return True


def update(rec, did, who=None, **fields):
    d = get(rec, did)
    if d is None:
        return False
    changed = []
    for k, v in fields.items():
        if k not in ("name", "kind", "format", "limit", "owner", "status", "doc_id", "form_version", "file", "notes", "progress", "volume"):
            continue
        if k == "progress":                           # 2.48.0: percent done, 0 to 100
            try:
                v = max(0, min(100, int(round(float(str(v).strip().rstrip("%") or 0)))))
            except ValueError:
                continue
        v = v.strip() if k in ("notes", "file") and isinstance(v, str) else _clean(v) if isinstance(v, str) else v
        if k == "status" and v not in STATUSES:
            continue
        if d.get(k) != v:
            changed.append(k)
            d[k] = v
    if "status" in changed:
        rec.log("deliverable", f"{d['name']}: {d['status']}", ref={"deliverable": did}, who=who)
    if "progress" in changed:
        rec.log("deliverable", f"{d['name']}: {d['progress']}% done", ref={"deliverable": did}, who=who)
    if "owner" in changed:
        rec.log("assign", f"Deliverable {d['name']} assigned to {d['owner'] or 'nobody'}", ref={"deliverable": did}, who=who)
    if "form_version" in changed:
        rec.log("deliverable", f"{d['name']}: now filled on the newest form version", ref={"deliverable": did}, who=who)
    if changed:
        rec.save()
    return bool(changed)


def remove(rec, did, who=None):
    d = get(rec, did)
    if d is None:
        return False
    items(rec).remove(d)
    rec.log("deliverable", f"Deliverable removed: {d['name']}", who=who)
    rec.save()
    return True


def mark_stage(rec, did, done=True, who=None):
    d = get(rec, did)
    if d is None:
        return False
    import effortrec as ER
    sid = rec.data["stage"]
    if done:
        d["stages_done"][sid] = f"{datetime.date.today().isoformat()} {ER._who(who)}"
        rec.log("deliverable", f"{d['name']}: done for {rec.stage()['name']}", ref={"deliverable": did}, who=who)
    else:
        d["stages_done"].pop(sid, None)
    rec.save()
    return True


def limit_row(d, rows, volumes=()):
    """The page-limit requirement that applies to a deliverable (by volume number or name), or None."""
    names = [d["name"]] + [v["name"] for v in volumes if d.get("file") and os.path.basename(d["file"]) in v.get("files", [])]
    for r in rows:
        if r["kind"] == "Page limit" and r["subject"].lower() != "proposal" and any(SR.match_volume(r["subject"], n) for n in names):
            return r
    return None


def state(d, rows=(), lib=None, volumes=(), rec=None):
    """Everything the tracker shows for one deliverable: limit, pages, over, changed, form, set_by, done_here, flags."""
    out = {"limit": d.get("limit", ""), "limit_from": "", "pages": None, "over": False, "changed": [], "form": None,
           "set_by": "", "done_here": False, "flags": []}
    byk = {r["key"]: r for r in rows}
    src = byk.get(d.get("req_key", ""))
    if src is not None:
        out["set_by"] = f"{src['doc']}, {src['version']}" if src.get("doc") else ""
        if src.get("to_review"):
            out["changed"].append(f"changed in {src['doc']}, {src['version']}")
    lr = limit_row(d, rows, volumes)
    if lr is not None:
        if not out["limit"]:
            out["limit"] = lr["value"]
            out["limit_from"] = f"{lr['doc']}, {lr['version']}" if lr.get("doc") else "typed in"
        if lr.get("to_review"):
            prev = next((h["value"] for h in reversed(lr.get("history", [])[:-1]) if h["value"]), "")
            out["changed"].append(f"page limit {('was ' + prev + ', now ') if prev else ''}{lr['value']} ({lr['doc']}, {lr['version']})")
    if d.get("file") and d["file"].lower().endswith(".docx") and os.path.isfile(d["file"]):
        out["pages"] = SR.word_pages(d["file"])
        n = SR._limit_n(out["limit"])
        if out["pages"] is not None and n:
            out["over"] = out["pages"] > n
    if d.get("doc_id") and lib is not None:
        doc = lib.doc(d["doc_id"])
        if doc and doc["versions"]:
            latest = doc["versions"][-1]
            working = next((v for v in doc["versions"] if v["id"] == d.get("form_version")), None)
            out["form"] = {"doc": doc["name"], "latest": latest["label"], "latest_id": latest["id"],
                           "working": working["label"] if working else "", "newer": bool(working) and working["id"] != latest["id"],
                           "versions": [(v["id"], v["label"], v["added"][:10], v.get("sha256", "")) for v in doc["versions"]]}
            if out["form"]["newer"]:
                out["changed"].append(f"new form version: {latest['label']} (fill it; {working['label']} stays for the record)")
    if rec is not None:
        out["done_here"] = rec.data["stage"] in d.get("stages_done", {})
    if out["over"]:
        out["flags"].append("over")
    if out["changed"]:
        out["flags"].append("changed")
    if not d.get("owner"):
        out["flags"].append("noowner")
    return out


# ------------------------------------------------------------------ submission requirements per deliverable (2.48.0)
SPEC = ("Page limit", "Font", "Margins", "Page size", "Line spacing", "File format", "Copies", "Delivery", "Due")


def _applies(subject, d, volumes=()):
    """A requirement's subject applies to this deliverable: it names it (or its volume), or it is for the whole proposal."""
    s = (subject or "").strip().lower()
    names = [d["name"]] + [v["name"] for v in volumes if d.get("file") and os.path.basename(d["file"]) in v.get("files", [])]
    if any(SR.match_volume(subject, n) for n in names):
        return 2
    if s in ("", "proposal", "all volumes", "the proposal") or s in [x.lower() for x in SPEC]:
        return 1
    return 0


def spec_suggest(d, rows=(), volumes=()):
    """{field: requirement row} the customer documents set for this deliverable (its own volume first, then the whole
    proposal)."""
    out = {}
    for f in SPEC:
        best, score = None, 0
        for r in rows:
            if r.get("hidden"):
                continue
            hit = (r["kind"] == "Page limit" and f == "Page limit") or (r["kind"] == "Format" and r["subject"] == f) or \
                  (r["kind"] == "File" and f == "File format") or (r["kind"] == f and f in ("Copies", "Delivery", "Due"))
            if not hit:
                continue
            sc = _applies(r["subject"] if r["kind"] in ("Page limit", "Copies", "Delivery", "Due") else "", d, volumes)
            if r["kind"] == "Page limit" and sc == 1 and r["subject"].lower() != "proposal":
                sc = 0                                # another volume's limit
            if sc > score:
                best, score = r, sc
        if best is not None:
            out[f] = best
    return out


def spec(d, rows=(), volumes=()):
    """[{field, value, source, state, suggested}] for this deliverable. state: Confirmed, Set by the team, Suggested,
    Changed: review (the customer's value moved after it was confirmed), or Not set."""
    sugg = spec_suggest(d, rows, volumes)
    have = d.get("spec") or {}
    out = []
    for f in SPEC:
        s = sugg.get(f)
        sv = s["value"] if s else ""
        src = f"{s['doc']}, {s['version']}" if s and s.get("doc") else ("typed in" if s else "")
        h = have.get(f)
        if h is None and f == "Page limit" and d.get("limit"):       # a limit typed on the deliverable is the team's
            h = {"value": d["limit"], "from": "customer" if s and s["value"] == d["limit"] else "team",
                 "source": src if s and s["value"] == d["limit"] else "typed in", "confirmed": ""}
        if h:
            if h.get("from") == "customer" and sv and h.get("value") != sv:
                state = "Changed: review"
            else:
                state = "Confirmed" if h.get("from") == "customer" else "Set by the team"
            out.append({"field": f, "value": h.get("value", ""), "source": h.get("source", ""), "state": state, "suggested": sv,
                        "confirmed": h.get("confirmed", "")})
        elif sv:
            out.append({"field": f, "value": sv, "source": src, "state": "Suggested", "suggested": sv, "confirmed": ""})
        else:
            out.append({"field": f, "value": "", "source": "", "state": "Not set", "suggested": "", "confirmed": ""})
    return out


def spec_set(rec, did, field, value, who=None, from_customer=False, source=""):
    """Set (typed) or confirm (from the customer documents) one requirement of a deliverable. Empty clears it."""
    d = get(rec, did)
    if d is None or field not in SPEC:
        return False
    import effortrec as ER
    sp = d.setdefault("spec", {})
    value = _clean(value)
    if not value:
        if sp.pop(field, None) is not None:
            rec.log("deliverable", f"{d['name']}: {field} cleared", ref={"deliverable": did}, who=who)
            rec.save()
        return True
    sp[field] = {"value": value, "from": "customer" if from_customer else "team", "source": source if from_customer else "typed in",
                 "confirmed": f"{datetime.date.today().isoformat()} {ER._who(who)}"}
    if field == "Page limit":
        d["limit"] = value
    rec.log("deliverable", f"{d['name']}: {field} " + ("confirmed" if from_customer else "set") + f": {value}",
            ref={"deliverable": did}, who=who)
    rec.save()
    return True


def spec_confirm(rec, did, field, rows=(), volumes=(), who=None):
    """Confirm the value the customer documents give (also after it changed)."""
    d = get(rec, did)
    s = spec_suggest(d, rows, volumes).get(field) if d else None
    if s is None:
        return False
    return spec_set(rec, did, field, s["value"], who, from_customer=True,
                    source=f"{s['doc']}, {s['version']}" if s.get("doc") else "typed in")


def spec_confirm_all(rec, did, rows=(), volumes=(), who=None):
    d = get(rec, did)
    n = 0
    for x in spec(d, rows, volumes) if d else []:
        if x["state"] in ("Suggested", "Changed: review"):
            n += bool(spec_confirm(rec, did, x["field"], rows, volumes, who))
    return n


def spec_counts(d, rows=(), volumes=()):
    sp = spec(d, rows, volumes)
    return {"confirmed": sum(1 for x in sp if x["state"] in ("Confirmed", "Set by the team")),
            "to_confirm": sum(1 for x in sp if x["state"] in ("Suggested", "Changed: review")),
            "changed": sum(1 for x in sp if x["state"] == "Changed: review")}


def status_text(d, st):
    if st["over"]:
        return f"Over: {st['pages']} of {SR._limit_n(st['limit'])} pages"
    if st["changed"]:
        return "Changed: " + st["changed"][0]
    if st["pages"] is not None and SR._limit_n(st["limit"]):
        return f"{d['status']}  ({st['pages']} of {SR._limit_n(st['limit'])} pages)"
    return d["status"] + ("  (done for this stage)" if st["done_here"] else "")


def summary(rec, rows=(), lib=None, volumes=()):
    ds = items(rec)
    sts = [state(d, rows, lib, volumes, rec) for d in ds]
    return {"count": len(ds), "over": sum(1 for s in sts if s["over"]), "changed": sum(1 for s in sts if s["changed"]),
            "noowner": sum(1 for d in ds if not d.get("owner")), "suggested": len(suggestions(rec, rows)),
            "final": sum(1 for d in ds if d["status"] in ("Final", "Submitted"))}


def export_excel(path, rec, rows=(), lib=None, volumes=(), effort_name=""):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Deliverables"
    stage_names = [s["name"] for s in rec.stages]
    head = ["Deliverable", "Kind", "Set by", "Format", "Limit and rules", "Owner", "Status", "Progress", "Pages", "Changed", "Form version"] + \
           list(SPEC) + ["Requirements to confirm"] + [f"Done: {n}" for n in stage_names] + ["Notes"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1E3A8A")
    for d in items(rec):
        st = state(d, rows, lib, volumes, rec)
        form = st["form"]
        sp = spec(d, rows, volumes)
        ws.append([d["name"], d["kind"], st["set_by"], d["format"], st["limit"], d["owner"], d["status"],
                   f"{d.get('progress', 0)}%", "" if st["pages"] is None else st["pages"], "; ".join(st["changed"]),
                   (f"{form['working'] or '-'} (newest {form['latest']})" if form else "")] +
                  [(x["value"] + ("" if x["state"] in ("Confirmed", "Set by the team", "Not set") else f" ({x['state'].lower()})"))
                   for x in sp] + [sum(1 for x in sp if x["state"] in ("Suggested", "Changed: review"))] +
                  [d["stages_done"].get(s["id"], "") for s in rec.stages] + [d["notes"]])
        fill = "FDE2E2" if st["over"] else "FEF3C7" if st["changed"] else None
        if fill:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor=fill)
    for r in suggestions(rec, rows):
        ws.append([r["subject"], kind_of(r["subject"]), f"{r['doc']}, {r['version']}", "", "", "", "Suggested, not confirmed", "", "", "", ""]
                  + [""] * (len(SPEC) + 1) + [""] * len(stage_names) + [r["value"]])
        for c in ws[ws.max_row]:
            c.fill = PatternFill("solid", fgColor="E0F2FE")
    widths = [30, 12, 26, 24, 22, 16, 14, 10, 8, 40, 22] + [18] * len(SPEC) + [12] + [14] * len(stage_names) + [40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    if effort_name:
        wb.properties.title = f"{effort_name} deliverables"
    wb.save(path)
    return path
