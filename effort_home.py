"""Effort home, named volumes, export picker and Use as definition (2.14.0).

- Two views in the main window while an effort is open: Effort list (home) and Scan review. Home shows every
  decision in the effort list, where each term came from and which volumes use it; it reopens as it was left.
  Every scan lands in Scan review.
- Named volumes (Vol I Technical, Vol II Management...), each pointed at its files. Terms get their volumes from
  where they were found; a term can belong to several volumes; a manual assignment stays even when a rescan disagrees.
- Import a volume's acronym list (CSV, Excel, Word, PDF, text): its terms join the effort list with that volume.
- Export picker: columns (Term and Definition by default), the whole effort or chosen volumes, one file or one per
  volume, CSV, Excel or Word. Rejected terms and the omit list are left out.
- Use as definition: text selected in the Source preview (typed text or a picture's recognized text) becomes the
  term's definition, saved for later scans in this effort and written to the effort list.
"""
import csv, datetime, os, re
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk
import customtkinter as ctk
import diag
import effort as E

VIEW_HOME, VIEW_SCAN, VIEW_KW, VIEW_QA = "Effort list (home)", "Scan review", "Keywords", "QA/QC"
HOME_FILTERS = ["All terms", "Needs a definition", "Proposed", "Needs author", "Approved", "Rejected", "No longer found"]
NO_VOLUME = "No volume"
HOME_COLS = (("term", "Term", 110), ("definition", "Definition", 330), ("status", "Status", 110),
             ("volumes", "Volumes", 170), ("found", "Found in", 300), ("scan", "Scan", 170), ("source", "Source", 130))
EXPORT_COLS = ["Term", "Definition", "Type", "Status", "Volumes", "Found in", "Source", "SS Approver", "Approved on",
               "First found", "Last found", "Scan status", "Notes", "Effort", "Client", "Solicitation"]
EXPORT_DEFAULT = ["Term", "Definition"]
FORMATS = {"csv": "CSV", "xlsx": "Excel", "docx": "Word"}


def _eng():
    import app_gui
    return app_gui


def _split(s):
    return [x for x in re.split(r"\s*;\s*", s or "") if x]


def _safe(name):
    return re.sub(r"[^A-Za-z0-9 _.-]+", "", name).strip().replace(" ", "_") or "List"


def clean_definition(text, term=""):
    """Selected words -> a definition: one line, no surrounding quotes or punctuation, no '(TERM)' after it."""
    t = " ".join((text or "").split()).strip(" \t\"'\u201c\u201d\u2018\u2019.,;:")
    if term:
        t = re.sub(r"\s*[\(\[]\s*" + re.escape(term) + r"s?\s*[\)\]]\s*$", "", t, flags=re.I)
        t = re.sub(r"^\s*" + re.escape(term) + r"s?\s*[:=\-–—]\s*", "", t, flags=re.I)
    t = t.strip(" \t\"'“”‘’.,;:()[]")
    return t[:200]


class HomeMixin:
    # ------------------------------------------------------------ build
    def home_build(self, M, before):
        """The view switch and the home view, both above `before` (the scan view's frame)."""
        self.view_var = tk.StringVar(value=VIEW_SCAN)
        self.seg_view = ctk.CTkSegmentedButton(M, values=[VIEW_HOME, VIEW_SCAN, VIEW_KW, VIEW_QA], variable=self.view_var,
                                               command=lambda v: self.show_view(v))
        self.home_view = ctk.CTkFrame(M, fg_color="transparent")
        H = self.home_view
        self.lbl_home = ctk.CTkLabel(H, text="", anchor="w", font=("Arial", 13, "bold"))
        self.lbl_home.pack(fill="x", padx=4, pady=(0, 4))
        frow = ctk.CTkFrame(H, fg_color="transparent"); frow.pack(fill="x", pady=(0, 4))
        ctk.CTkLabel(frow, text="Show").pack(side="left", padx=(4, 4))
        self.home_filter = tk.StringVar(value=HOME_FILTERS[0])
        self.opt_home = ctk.CTkOptionMenu(frow, variable=self.home_filter, values=HOME_FILTERS, width=200,
                                          command=lambda _v: self._home_changed())
        self.opt_home.pack(side="left")
        ctk.CTkLabel(frow, text="Search").pack(side="left", padx=(12, 4))
        self.home_search = tk.StringVar(value="")
        self.ent_home = ctk.CTkEntry(frow, textvariable=self.home_search, width=220)
        self.ent_home.pack(side="left")
        self.home_search.trace_add("write", lambda *_: self._home_changed())
        self.lbl_home_n = ctk.CTkLabel(frow, text="", anchor="w")
        self.lbl_home_n.pack(side="left", padx=10)
        tf = tk.Frame(H); tf.pack(fill="both", expand=True)
        tf.grid_rowconfigure(0, weight=1); tf.grid_columnconfigure(0, weight=1)
        self.home_tree = tv = ttk.Treeview(tf, columns=[c[0] for c in HOME_COLS], show="headings", selectmode="extended")
        for cid, head, w in HOME_COLS:
            tv.heading(cid, text=head); tv.column(cid, width=w, minwidth=60, stretch=cid in ("definition", "found"))
        tv.tag_configure("Approved", background="#E7F6EC")
        tv.tag_configure("Needs author", background="#FEF3E2")
        tv.tag_configure("Rejected", background="#EEEEEE", foreground="#6B7280")
        tv.tag_configure("gone", background="#E5E7EB", foreground="#6B7280")
        tv.tag_configure("nodef", background="#FFF7ED")
        vsb = ttk.Scrollbar(tf, orient="vertical", command=tv.yview)
        hsb = ttk.Scrollbar(tf, orient="horizontal", command=tv.xview)
        tv.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        tv.grid(row=0, column=0, sticky="nsew"); vsb.grid(row=0, column=1, sticky="ns"); hsb.grid(row=1, column=0, sticky="ew")
        tv.bind("<<TreeviewSelect>>", lambda _e: (self._home_buttons(), self.resume_changed()))
        tv.bind("<Double-1>", lambda _e: self.home_edit_definition())
        tv.bind("<Button-3>", self._home_menu)
        tv.bind("<Control-a>", lambda _e: (tv.selection_set(tv.get_children()), "break")[1])
        b = ctk.CTkFrame(H, fg_color="transparent"); b.pack(fill="x", pady=(6, 0))
        K = dict(height=30)
        self.btn_h_def = ctk.CTkButton(b, text="Edit definition...", width=130, command=self.home_edit_definition, **K)
        self.btn_h_ok = ctk.CTkButton(b, text="Approve...", width=100, fg_color="#15803D",
                                      command=lambda: self.home_set_status("Approved"), **K)
        self.btn_h_auth = ctk.CTkButton(b, text="Needs author", width=110, fg_color="#B45309",
                                        command=lambda: self.home_set_status("Needs author"), **K)
        self.btn_h_rej = ctk.CTkButton(b, text="Not a term", width=100, fg_color="#B91C1C",
                                       command=lambda: self.home_set_status("Rejected"), **K)
        self.btn_h_vol = ctk.CTkButton(b, text="Volumes...", width=100, fg_color="#1E3A8A", command=self.assign_volumes_dialog, **K)
        self.btn_h_where = ctk.CTkButton(b, text="Where from...", width=110, fg_color="#475569", command=self.home_where_from, **K)
        self.btn_h_scan = ctk.CTkButton(b, text="Show in Scan review", width=150, fg_color="#475569", command=self.home_show_in_scan, **K)
        for w in (self.btn_h_def, self.btn_h_ok, self.btn_h_auth, self.btn_h_rej, self.btn_h_vol, self.btn_h_where, self.btn_h_scan):
            w.pack(side="left", padx=(0, 6))
        self.btn_h_export = ctk.CTkButton(b, text="Export...", width=100, fg_color="#0F766E", command=self.export_picker, **K)
        self.btn_h_export.pack(side="right")
        self.btn_h_import = ctk.CTkButton(b, text="Import a list...", width=120, fg_color="#475569", command=self.import_volume_list, **K)
        self.btn_h_import.pack(side="right", padx=6)
        self.btn_h_setvol = ctk.CTkButton(b, text="Name volumes...", width=130, fg_color="#475569", command=self.edit_volumes, **K)
        self.btn_h_setvol.pack(side="right")
        self._home_after = None
        self._home_buttons()
        if hasattr(self, "kw_build"):                  # Keyword Harvest (2.16.0)
            self.kw_build(M)
        if hasattr(self, "qa_build"):                  # QA/QC of our Word deliverables (2.24.0)
            self.qa_build(M)

    # ------------------------------------------------------------ views
    def _views(self):
        v = {VIEW_HOME: self.home_view, VIEW_SCAN: self.scan_view}
        if getattr(self, "kw_view", None) is not None:
            v[VIEW_KW] = self.kw_view
        if getattr(self, "qa_view", None) is not None:
            v[VIEW_QA] = self.qa_view
        return v

    def show_view(self, which=None, save=True):
        which = which or self.view_var.get()
        if self.effort is None or which not in self._views():
            which = VIEW_SCAN
        self.view_var.set(which)
        for name, frame in self._views().items():
            if name != which:
                frame.pack_forget()
        frame = self._views()[which]
        if not frame.winfo_manager():
            if which == VIEW_SCAN:
                frame.pack(fill="both", expand=True)
            else:
                frame.pack(fill="both", expand=True, padx=15, pady=(0, 10))
        if which == VIEW_HOME:
            self.home_refresh()
            diag.activity("view_home")
        elif which == VIEW_KW:
            self.kw_refresh()
            diag.activity("view_keywords")
        elif which == VIEW_QA:
            diag.activity("view_qa")
        else:
            diag.activity("view_scan")
        if save:
            self.resume_changed()

    def home_update(self):
        """Called whenever the effort or its list changes."""
        if not hasattr(self, "seg_view"):
            return
        if self.effort is None:
            self.seg_view.pack_forget()
            self.show_view(VIEW_SCAN, save=False)
            return
        if not self.seg_view.winfo_manager():
            anchor = next((f for f in self._views().values() if f.winfo_manager()), self.scan_view)
            self.seg_view.pack(pady=(0, 6), before=anchor)
        if self.view_var.get() == VIEW_HOME:
            self.home_refresh()
        elif self.view_var.get() == VIEW_KW:
            self.kw_refresh()

    def home_state(self):
        sel = self.home_tree.selection() if hasattr(self, "home_tree") else ()
        return {"view": self.view_var.get() if hasattr(self, "view_var") else VIEW_SCAN,
                "filter": self.home_filter.get() if hasattr(self, "home_filter") else HOME_FILTERS[0],
                "search": self.home_search.get() if hasattr(self, "home_search") else "",
                "selected": list(sel)[:50]}

    def home_restore(self, h):
        if not hasattr(self, "home_tree"):
            return
        h = h or {}
        self._home_restoring = True
        try:
            self.home_filter.set(h.get("filter") or HOME_FILTERS[0])
            self.home_search.set(h.get("search") or "")
        finally:
            self._home_restoring = False
        self._home_sel = [t for t in h.get("selected") or [] if isinstance(t, str)]
        self.show_view(h.get("view") or (VIEW_HOME if self.effort is not None else VIEW_SCAN), save=False)

    def _home_changed(self):
        if getattr(self, "_home_restoring", False):
            return
        if self._home_after:
            try:
                self.after_cancel(self._home_after)
            except Exception:
                diag.swallowed()
        self._home_after = self.after(200, lambda: (setattr(self, "_home_after", None), self.home_refresh(), self.resume_changed()))

    # ------------------------------------------------------------ home table
    def _home_match(self, t, r, f, q):
        if q and q not in t.casefold() and q not in r.get("Definition", "").casefold():
            return False
        if f == "Needs a definition":
            return not r.get("Definition") and r["Status"] != "Rejected"
        if f in E.STATUSES:
            return r["Status"] == f
        if f == "No longer found":
            return r.get("Scan status", "").startswith(E.SCAN_GONE)
        if f == NO_VOLUME:
            return not self.term_volumes(t, r)
        if f.startswith("Volume: "):
            return f[8:] in self.term_volumes(t, r)
        return True

    def home_filters(self):
        vols = [v["name"] for v in self.volumes()]
        return HOME_FILTERS + ([f"Volume: {v}" for v in vols] + [NO_VOLUME] if vols else [])

    def home_refresh(self):
        if not hasattr(self, "home_tree"):
            return
        tv = self.home_tree
        keep = list(tv.selection()) or getattr(self, "_home_sel", [])
        self._home_sel = []
        tv.delete(*tv.get_children())
        L = self.effort
        if L is None:
            self.lbl_home.configure(text=""); self.lbl_home_n.configure(text="")
            self._home_buttons(); return
        self.home_sync_volumes()
        opts = self.home_filters()
        self.opt_home.configure(values=opts)
        if self.home_filter.get() not in opts:
            self.home_filter.set(HOME_FILTERS[0])
        f, q = self.home_filter.get(), self.home_search.get().strip().casefold()
        shown = 0
        for t in L.order:
            r = L.rows[t]
            if not self._home_match(t, r, f, q):
                continue
            scan = r.get("Scan status", "")
            tag = ("gone" if scan.startswith(E.SCAN_GONE) else r["Status"] if r["Status"] != "Proposed"
                   else "nodef" if not r.get("Definition") else "")
            tv.insert("", "end", iid=t, tags=(tag,) if tag else (),
                      values=(t, r.get("Definition") or "(needs a definition)", r["Status"],
                              "; ".join(self.term_volumes(t, r)), r.get("Found in", ""), scan, r.get("Source", "")))
            shown += 1
        rows = list(L.rows.values())
        n_ok = sum(r["Status"] == "Approved" for r in rows)
        n_def = sum(not r.get("Definition") and r["Status"] != "Rejected" for r in rows)
        n_gone = sum(r.get("Scan status", "").startswith(E.SCAN_GONE) for r in rows)
        bits = [f"{len(rows)} terms", f"{n_ok} approved", f"{n_def} need a definition", f"{n_gone} no longer found"]
        vols = self.volumes()
        if vols:
            bits.append(f"{len(vols)} volumes")
        rec = getattr(self, "_last_effort_scan", None)
        if rec:
            bits.append(f"last scan: {rec['new']} new" + (" (Scan review)" if rec["new"] else ""))
        if L.pending:
            bits.append("NOT SAVED (close the list in Excel)")
        self.lbl_home.configure(text=f"{L.name}:  " + "  |  ".join(bits))
        self.lbl_home_n.configure(text=f"Showing {shown} of {len(rows)}")
        sel = [t for t in keep if tv.exists(t)]
        if sel:
            tv.selection_set(sel); tv.see(sel[0])
        self._home_buttons()

    def home_selected(self):
        if not hasattr(self, "home_tree") or self.effort is None:
            return []
        return [t for t in self.home_tree.selection() if t in self.effort.rows]

    def _home_buttons(self):
        sel = self.home_selected()
        st = "normal" if sel else "disabled"
        for w in (self.btn_h_ok, self.btn_h_auth, self.btn_h_rej, self.btn_h_vol):
            w.configure(state=st)
        one = "normal" if len(sel) == 1 else "disabled"
        for w in (self.btn_h_def, self.btn_h_where):
            w.configure(state=one)
        self.btn_h_scan.configure(state="normal" if len(sel) == 1 and sel[0] in self.results else "disabled")
        has = "normal" if self.effort is not None else "disabled"
        for w in (self.btn_h_export, self.btn_h_import, self.btn_h_setvol):
            w.configure(state=has)

    def _home_menu(self, event):
        iid = self.home_tree.identify_row(event.y)
        if not iid:
            return
        if iid not in self.home_tree.selection():
            self.home_tree.selection_set(iid)
        m = tk.Menu(self, tearoff=0)
        m.add_command(label="Edit definition...", command=self.home_edit_definition)
        m.add_command(label="Where from...", command=self.home_where_from)
        m.add_command(label="Volumes...", command=self.assign_volumes_dialog)
        r = self.effort.rows.get(iid, {})
        if hasattr(self, "kw_quick") and r.get("Definition"):
            m.add_command(label=f"Add as keyword: {r['Definition'][:40]}", command=lambda: self.kw_quick(r["Definition"], source=f"Effort list {iid}"))
        m.add_separator()
        for s in E.STATUSES:
            m.add_command(label=f"Set status: {s}", command=lambda s=s: self.home_set_status(s))
        m.add_separator()
        m.add_command(label="Remove from the effort list", command=self.home_remove)
        m.tk_popup(event.x_root, event.y_root)

    # ------------------------------------------------------------ home actions
    def home_set_definition(self, term, text, note=None):
        """Write a definition into the effort list (and into the scan results, if the term is there)."""
        text = clean_definition(text, term)
        L = self.effort
        if L is None or term not in L.rows or not text:
            return False
        if L.add_or_update(term, change_note=note or f"definition: {text[:80]}", Definition=text) != "unchanged":
            self._save_effort_list()
        self.ws_pick_definition(term, text)
        if term in self.results:
            self.apply_effort()
        self.home_refresh()
        return True

    def home_edit_definition(self):
        sel = self.home_selected()
        if len(sel) != 1:
            return None
        t = sel[0]
        cur = self.effort.rows[t].get("Definition", "")
        new = simpledialog.askstring("Definition", f"Definition for {t}:", initialvalue=cur, parent=self)
        if new is None or not new.strip() or new.strip() == cur:
            return None
        diag.activity("home_definition")
        return self.home_set_definition(t, new)

    def home_set_status(self, status):
        sel = self.home_selected()
        if not sel:
            return 0
        approver = None
        if status == "Approved":
            approver = simpledialog.askstring("Approve", "SS Approver (one name, or several separated by ;):",
                                              initialvalue=getattr(self, "_last_approver", ""), parent=self)
            if approver is None:
                return 0
            approver = approver.strip(); self._last_approver = approver
        n = self.effort.set_status(sel, status, approver=approver or None)
        diag.activity("home_status", n=n)
        self._save_effort_list()
        self.apply_effort()
        self.update_effort_ui()
        return n

    def home_remove(self):
        sel = self.home_selected()
        if sel and messagebox.askyesno("Remove", f"Remove {len(sel)} term(s) from the effort list? The History sheet keeps a line for each."):
            for t in sel:
                self.effort.remove(t)
            self._save_effort_list(); self.apply_effort(); self.update_effort_ui()

    def home_where_from(self):
        """Where a term came from: documents, customer wording, imports and every change, newest first."""
        sel = self.home_selected()
        if len(sel) != 1:
            return None
        t = sel[0]
        r = self.effort.rows[t]
        win = ctk.CTkToplevel(self); win.title(f"Where {t} came from"); win.geometry("820x480"); win.transient(self)
        win.after(150, win.lift)
        diag.register(win, "where_from")
        facts = [("Definition", r.get("Definition", "")), ("Source", r.get("Source", "")), ("Found in", r.get("Found in", "")),
                 ("First found", r.get("First found", "")), ("Last found", r.get("Last found", "")),
                 ("Scan status", r.get("Scan status", "")), ("Volumes", "; ".join(self.term_volumes(t, r))),
                 ("Customer document", r.get("Customer document", "")), ("Customer wording", r.get("Customer wording", "")),
                 ("Other definitions", r.get("Other definitions", ""))]
        txt = "\n".join(f"{k}: {v}" for k, v in facts if v)
        ctk.CTkLabel(win, text=txt or "Nothing recorded yet.", justify="left", anchor="w", wraplength=780).pack(fill="x", padx=12, pady=(12, 6))
        tv = ttk.Treeview(win, columns=("when", "who", "what"), show="headings")
        for c, h, w in (("when", "When", 130), ("who", "Who", 100), ("what", "What happened", 540)):
            tv.heading(c, text=h); tv.column(c, width=w, stretch=(c == "what"))
        tv.pack(fill="both", expand=True, padx=12, pady=6)
        for h in reversed([h for h in self.effort.history if len(h) > 3 and h[2] == t]):
            tv.insert("", "end", values=(h[0], h[1], h[3]))
        ctk.CTkButton(win, text="Close", width=90, command=win.destroy).pack(anchor="e", padx=12, pady=(0, 10))
        win._tree, win._text = tv, txt
        return win

    def home_show_in_scan(self):
        sel = self.home_selected()
        if len(sel) != 1 or sel[0] not in self.results:
            return
        t = sel[0]
        self.show_view(VIEW_SCAN)
        if not self.tree.exists(t):
            self.status_var.set(_eng().STATUS_FILTERS[0]); self.search_var.set("")
            self.populate_table()
        if self.tree.exists(t):
            self.tree.selection_set(t); self.tree.see(t); self.on_select(None)

    # ------------------------------------------------------------ volumes
    def volumes(self):
        ws = getattr(self, "ws", None)
        v = (ws.get("volumes", []) if ws is not None else []) or []
        return [x for x in v if isinstance(x, dict) and x.get("name")]

    def _set_volumes(self, vols):
        self.ws.set("volumes", vols)
        diag.activity("volumes_named", n=len(vols))
        self.home_sync_volumes(); self.home_refresh()

    def add_volume(self, name, files=()):
        name = " ".join((name or "").split())
        vols = self.volumes()
        if not name or any(v["name"].casefold() == name.casefold() for v in vols):
            return False
        vols.append({"name": name, "files": sorted(set(os.path.basename(f) for f in files), key=str.lower)})
        self._set_volumes(vols)
        return True

    def rename_volume(self, old, new):
        new = " ".join((new or "").split())
        vols = self.volumes()
        if not new or any(v["name"].casefold() == new.casefold() for v in vols if v["name"] != old):
            return False
        for v in vols:
            if v["name"] == old:
                v["name"] = new
        pins = self.ws.get("volume_pins", {}) or {}
        self.ws.set("volume_pins", {t: [new if x == old else x for x in vs] for t, vs in pins.items()})
        self._set_volumes(vols)
        return True

    def remove_volume(self, name):
        vols = [v for v in self.volumes() if v["name"] != name]
        pins = self.ws.get("volume_pins", {}) or {}
        self.ws.set("volume_pins", {t: [x for x in vs if x != name] for t, vs in pins.items()})
        self._set_volumes(vols)

    def set_volume_files(self, name, files):
        vols = self.volumes()
        for v in vols:
            if v["name"] == name:
                v["files"] = sorted(set(os.path.basename(f) for f in files), key=str.lower)
        self._set_volumes(vols)

    def term_volumes(self, term, row=None):
        """Volumes a term belongs to: the manual assignment when there is one, else the volumes whose files it was found in."""
        vols = self.volumes()
        if not vols:
            return _split((row or {}).get("Used in volumes", ""))
        names = [v["name"] for v in vols]
        pins = (self.ws.get("volume_pins", {}) or {}) if getattr(self, "ws", None) is not None else {}
        if term in pins:
            return [n for n in names if n in pins[term]]
        found = {f.casefold() for f in _split((row or {}).get("Found in", ""))}
        return [v["name"] for v in vols if found & {f.casefold() for f in v.get("files", [])}]

    def assign_volumes(self, terms, names):
        """names = a list of volumes (kept even when a rescan disagrees), or None for automatic (from where it was found)."""
        pins = dict(self.ws.get("volume_pins", {}) or {})
        valid = {v["name"] for v in self.volumes()}
        for t in terms:
            if names is None:
                pins.pop(t, None)
            else:
                pins[t] = [n for n in names if n in valid]
            if self.effort is not None and t in self.effort.rows:
                self.effort._hist(t, "volumes: " + ("automatic" if names is None else "; ".join(pins[t]) or "none"))
        self.ws.set("volume_pins", pins)
        diag.activity("volumes_assigned", n=len(terms))
        self.home_sync_volumes(force_save=True)
        self.home_refresh()

    def home_sync_volumes(self, force_save=False):
        """Write each term's volumes into the list's Used in volumes column (only when volumes are named)."""
        L = self.effort
        if L is None or not self.volumes():
            if force_save and L is not None and L.pending:
                self._save_effort_list()
            return False
        changed = False
        for t in L.order:
            r = L.rows[t]
            val = "; ".join(self.term_volumes(t, r))
            if r.get("Used in volumes", "") != val:
                r["Used in volumes"] = val; changed = True
        if changed:
            L.pending = True
        if changed or (force_save and L.pending):
            self._save_effort_list(backup=False)
        return changed

    def known_documents(self):
        """Document names this effort has seen: Found in, this session's scans, and the volumes folder."""
        L = self.effort
        names = set()
        if L is not None:
            for r in L.rows.values():
                names.update(_split(r.get("Found in", "")))
        names.update(getattr(self, "scanned_paths", {}).keys())
        info = (self.efforts.get(L.name) or {}) if L is not None else {}
        vd = info.get("volumes_dir") or ""
        if vd and os.path.isdir(vd):
            names.update(f for f in os.listdir(vd) if f.lower().endswith((".docx", ".pdf", ".pptx", ".xlsx")))
        cd = info.get("customer_dir") or ""
        cust = set(os.listdir(cd)) if cd and os.path.isdir(cd) else set()
        for v in self.volumes():
            names.update(v.get("files", []))
        return sorted((n for n in names if n not in cust), key=str.lower)

    def edit_volumes(self):
        """Name the volumes and tick each one's files."""
        if self.effort is None:
            return None
        win = ctk.CTkToplevel(self); win.title(f"Volumes: {self.effort.name}"); win.geometry("820x520"); win.transient(self)
        win.after(150, win.lift)
        diag.register(win, "volumes")
        ctk.CTkLabel(win, text="Name the volumes, then tick the files that make up each one. Terms get their volumes from the "
                               "files they were found in; Volumes... on the home view assigns them by hand.",
                     wraplength=780, justify="left").pack(anchor="w", padx=12, pady=(12, 6))
        body = ctk.CTkFrame(win, fg_color="transparent"); body.pack(fill="both", expand=True, padx=12)
        left = ctk.CTkFrame(body); left.pack(side="left", fill="y", padx=(0, 8))
        lb = tk.Listbox(left, height=14, exportselection=False, font=("Segoe UI", 11))
        lb.pack(fill="both", expand=True, padx=6, pady=6)
        right = ctk.CTkScrollableFrame(body, label_text="Files in the selected volume")
        right.pack(side="left", fill="both", expand=True)
        st = {"cur": None, "vars": {}}

        def fill():
            lb.delete(0, "end")
            for v in self.volumes():
                lb.insert("end", v["name"])

        def show(name):
            st["cur"] = name
            for w in right.winfo_children():
                w.destroy()
            st["vars"] = {}
            if not name:
                return
            have = set(next((v.get("files", []) for v in self.volumes() if v["name"] == name), []))
            for f in self.known_documents():
                var = tk.BooleanVar(value=f in have)
                ctk.CTkCheckBox(right, text=f, variable=var, command=save_files).pack(anchor="w", pady=1)
                st["vars"][f] = var

        def save_files():
            if st["cur"]:
                self.set_volume_files(st["cur"], [f for f, v in st["vars"].items() if v.get()])

        def on_pick(_e=None):
            s = lb.curselection()
            show(lb.get(s[0]) if s else None)

        def add():
            n = simpledialog.askstring("New volume", "Volume name (for example Vol I Technical):", parent=win)
            if n and self.add_volume(n):
                fill(); lb.selection_clear(0, "end"); lb.selection_set("end"); on_pick()

        def rename():
            if st["cur"]:
                n = simpledialog.askstring("Rename volume", "New name:", initialvalue=st["cur"], parent=win)
                if n and self.rename_volume(st["cur"], n):
                    fill(); show(None)

        def remove():
            if st["cur"] and messagebox.askyesno("Remove volume", f"Remove {st['cur']}? Terms keep their other volumes.", parent=win):
                self.remove_volume(st["cur"]); fill(); show(None)

        def add_files():
            if not st["cur"]:
                return
            ps = filedialog.askopenfilenames(parent=win, title="Files in this volume",
                                             filetypes=[("Documents", "*.docx *.pdf *.pptx *.xlsx"), ("All files", "*.*")])
            if ps:
                have = next((v.get("files", []) for v in self.volumes() if v["name"] == st["cur"]), [])
                self.set_volume_files(st["cur"], list(have) + list(ps)); show(st["cur"])

        lb.bind("<<ListboxSelect>>", on_pick)
        bb = ctk.CTkFrame(left, fg_color="transparent"); bb.pack(fill="x", padx=6, pady=(0, 6))
        for txt, fn in (("Add...", add), ("Rename...", rename), ("Remove", remove)):
            ctk.CTkButton(bb, text=txt, width=80, height=28, command=fn).pack(side="left", padx=(0, 4))
        foot = ctk.CTkFrame(win, fg_color="transparent"); foot.pack(fill="x", padx=12, pady=10)
        ctk.CTkButton(foot, text="Add files to this volume...", width=190, fg_color="#475569", command=add_files).pack(side="left")
        ctk.CTkButton(foot, text="Close", width=90, command=win.destroy).pack(side="right")
        fill()
        win._api = {"list": lb, "show": show, "vars": st["vars"], "st": st}
        return win

    def assign_volumes_dialog(self, terms=None):
        terms = terms or self.home_selected()
        if not terms:
            return None
        vols = self.volumes()
        if not vols:
            if messagebox.askyesno("Volumes", "No volumes are named for this effort yet. Name them now?"):
                return self.edit_volumes()
            return None
        win = ctk.CTkToplevel(self); win.title("Volumes"); win.geometry("460x420"); win.transient(self)
        win.after(150, win.lift)
        diag.register(win, "assign_volumes")
        L = self.effort
        ctk.CTkLabel(win, text=(f"{terms[0]} belongs to:" if len(terms) == 1 else f"{len(terms)} terms belong to:"),
                     font=("Segoe UI", 13, "bold"), anchor="w").pack(fill="x", padx=14, pady=(12, 4))
        cur = set.intersection(*[set(self.term_volumes(t, L.rows.get(t))) for t in terms])
        vars_ = {}
        for v in vols:
            var = tk.BooleanVar(value=v["name"] in cur)
            ctk.CTkCheckBox(win, text=v["name"], variable=var).pack(anchor="w", padx=18, pady=2)
            vars_[v["name"]] = var
        ctk.CTkLabel(win, text="Your ticks stay even when a rescan finds the term somewhere else. Automatic goes back to "
                               "the volumes whose files it was found in.", wraplength=420, justify="left").pack(anchor="w", padx=14, pady=8)
        b = ctk.CTkFrame(win, fg_color="transparent"); b.pack(fill="x", padx=14, pady=8)

        def ok():
            self.assign_volumes(terms, [n for n, var in vars_.items() if var.get()]); win.destroy()

        def auto():
            self.assign_volumes(terms, None); win.destroy()

        ctk.CTkButton(b, text="Save", width=90, command=ok).pack(side="left")
        ctk.CTkButton(b, text="Automatic", width=100, fg_color="#475569", command=auto).pack(side="left", padx=6)
        ctk.CTkButton(b, text="Cancel", width=90, fg_color="#475569", command=win.destroy).pack(side="right")
        win._api = {"vars": vars_, "ok": ok, "auto": auto}
        return win

    # ------------------------------------------------------------ import a volume's list
    def import_volume_list(self, path=None, volume=None, ask=True):
        """Bring in an acronym list exported earlier (or written by an author). Returns the counts, or None."""
        eng = _eng()
        L = self.effort
        if L is None:
            return None
        path = path or filedialog.askopenfilename(title="Acronym list to bring in", initialdir=os.path.dirname(L.path),
                                                  filetypes=[("Acronym lists", "*.csv *.xlsx *.docx *.pdf *.txt *.tsv"),
                                                             ("All files", "*.*")])
        if not path:
            return None
        try:
            pairs = eng.parse_client_list(path)
        except Exception as ex:
            diag.swallowed()
            messagebox.showerror("Import", f"Could not read {os.path.basename(path)}: {ex}")
            return None
        if not pairs:
            messagebox.showinfo("Import", f"No term and definition pairs found in {os.path.basename(path)}.")
            return None
        if volume is None and ask and self.volumes():
            names = [v["name"] for v in self.volumes()]
            volume = self._ask_choice("Which volume?", f"{len(pairs)} terms in {os.path.basename(path)}. Which volume is this list for?",
                                      names + ["(no volume)"])
            if volume is None:
                return None
            if volume == "(no volume)":
                volume = ""
        fname = os.path.basename(path)
        added = filled = differs = same = 0
        pins = dict(self.ws.get("volume_pins", {}) or {})
        for t, d in pairs:
            t, d = t.strip(), clean_definition(d, t)
            if not t:
                continue
            row = L.rows.get(t)
            if row is None:
                L.add_or_update(t, change_note=f"imported from {fname}" + (f" ({volume})" if volume else "") + (f": {d[:60]}" if d else ""),
                                Definition=d, Type="Acronym", Source=f"Imported list: {fname}")
                added += 1
            elif not row.get("Definition") and d:
                L.add_or_update(t, change_note=f"definition from imported {fname}: {d[:60]}", Definition=d); filled += 1
            elif d and eng.norm_def(d) != eng.norm_def(row.get("Definition", "")):
                others = [o for o in re.split(r"\s*\|\s*", row.get("Other definitions", "")) if o]
                if d not in others:
                    others.append(d)
                    row["Other definitions"] = " | ".join(others)[:500]
                    L._hist(t, f"imported {fname} has a different definition: {d[:60]}")
                differs += 1
            else:
                same += 1
            if volume:
                cur = self.term_volumes(t, L.rows[t])
                pins[t] = list(dict.fromkeys(cur + [volume]))
        if volume:
            self.ws.set("volume_pins", pins)
        diag.activity("volume_list_imported", added=added, n=len(pairs))
        self.home_sync_volumes()
        self._save_effort_list()
        self.apply_effort()
        self.update_effort_ui()
        res = {"added": added, "filled": filled, "differs": differs, "same": same, "volume": volume or ""}
        if ask:
            messagebox.showinfo("Import", f"{fname}: {added} added as Proposed, {filled} definitions filled in, "
                                f"{differs} with a different definition (kept under Other definitions), {same} already the same."
                                + (f"\n\nAll of them now belong to {volume}." if volume else ""))
        return res

    def _ask_choice(self, title, prompt, options):
        win = ctk.CTkToplevel(self); win.title(title); win.geometry("440x200"); win.transient(self)
        win.after(150, win.lift)
        ctk.CTkLabel(win, text=prompt, wraplength=400, justify="left").pack(anchor="w", padx=14, pady=(12, 6))
        var = ctk.StringVar(value=options[0])
        ctk.CTkOptionMenu(win, variable=var, values=options, width=300).pack(anchor="w", padx=14)
        out = {"v": None}
        b = ctk.CTkFrame(win, fg_color="transparent"); b.pack(fill="x", padx=14, pady=14)
        ctk.CTkButton(b, text="OK", width=90, command=lambda: (out.update(v=var.get()), win.destroy())).pack(side="left")
        ctk.CTkButton(b, text="Cancel", width=90, fg_color="#475569", command=win.destroy).pack(side="right")
        win.grab_set(); self.wait_window(win)
        return out["v"]

    # ------------------------------------------------------------ export picker
    def export_value(self, t, r, col):
        L = self.effort
        if col == "Volumes":
            return "; ".join(self.term_volumes(t, r))
        if col in ("Effort", "Client", "Solicitation"):
            return {"Effort": L.name, "Client": L.client, "Solicitation": L.solicitation}[col]
        return r.get(col, "")

    def export_terms(self, target, cols=None, scope="whole", volumes=None, fmt="csv", approved_only=False):
        """Write the effort's terms. scope 'whole' = one file (only terms in `volumes` when given); 'per_volume' = one
        file per volume. target = a file path (whole) or a folder (per volume). Returns the files written."""
        L = self.effort
        if L is None:
            return []
        cols = [c for c in (cols or EXPORT_DEFAULT) if c in EXPORT_COLS] or EXPORT_DEFAULT
        rows = [(r["Term"], r) for r in self._list_rows(approved_only)]
        stamp = datetime.date.today().isoformat()
        stem = _safe(L.name)
        made = []
        if scope == "per_volume":
            groups = []
            names = [v["name"] for v in self.volumes()]
            if names:
                for n in (volumes or names):
                    groups.append((n, [(t, r) for t, r in rows if n in self.term_volumes(t, r)]))
            else:                                       # no named volumes: one list per document it was found in
                info = self.efforts.get(L.name) or {}
                cd = info.get("customer_dir") or ""
                cust = set(os.listdir(cd)) if cd and os.path.isdir(cd) else set()
                docs = sorted({f for _t, r in rows for f in _split(r.get("Found in", "")) if f not in cust}, key=str.lower)
                for f in docs:
                    groups.append((os.path.splitext(f)[0], [(t, r) for t, r in rows if f in _split(r.get("Found in", ""))]))
            folder = target
            os.makedirs(folder, exist_ok=True)
            for name, rs in groups:
                p = os.path.join(folder, f"{stem}_Acronyms_{_safe(name)}_{stamp}.{fmt}")
                self._write_terms(p, cols, rs, fmt, f"Acronyms and Abbreviations: {name}")
                made.append(p)
        else:
            if volumes:
                rows = [(t, r) for t, r in rows if set(volumes) & set(self.term_volumes(t, r))]
            title = f"Acronyms and Abbreviations: {L.name}" + (f" ({', '.join(volumes)})" if volumes else "")
            self._write_terms(target, cols, rows, fmt, title)
            made.append(target)
        diag.activity("terms_exported", files=len(made), rows=len(rows))
        return made

    def _write_terms(self, path, cols, rows, fmt, title):
        vals = [[self.export_value(t, r, c) for c in cols] for t, r in rows]
        if fmt == "csv":
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f); w.writerow(cols); w.writerows(vals)
        elif fmt == "xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Font
            wb = Workbook(); ws = wb.active; ws.title = "Acronyms"
            ws.append(cols)
            for c in ws[1]:
                c.font = Font(bold=True)
            for v in vals:
                ws.append(v)
            for i, c in enumerate(cols, 1):
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 44 if c in ("Definition", "Found in", "Notes") else 16
            ws.freeze_panes = "A2"
            wb.save(path)
        else:
            import docxw
            L = self.effort
            d = docxw.Doc()
            d.props = {"AcronymHunterEffort": L.name, "Client": L.client, "Solicitation": L.solicitation}
            d.heading(title, 1)
            total = 9360
            first = 1800 if len(cols) > 1 else total
            rest = (total - first) // max(1, len(cols) - 1)
            d.table(cols, vals, [first] + [rest] * (len(cols) - 1))
            d.save(path)

    def export_picker(self):
        """Pick the columns and the volumes, then export (CSV, Excel or Word)."""
        L = self.effort
        if L is None:
            return None
        win = ctk.CTkToplevel(self); win.title(f"Export: {L.name}"); win.geometry("640x560"); win.transient(self)
        win.after(150, win.lift)
        diag.register(win, "export_picker")
        saved = self.ws.get("export_choice", {}) or {}
        ctk.CTkLabel(win, text="Columns", font=("Segoe UI", 13, "bold")).pack(anchor="w", padx=14, pady=(12, 2))
        cf = ctk.CTkFrame(win, fg_color="transparent"); cf.pack(fill="x", padx=14)
        colv = {}
        for i, c in enumerate(EXPORT_COLS):
            var = tk.BooleanVar(value=c in (saved.get("cols") or EXPORT_DEFAULT))
            ctk.CTkCheckBox(cf, text=c, variable=var, width=140).grid(row=i // 4, column=i % 4, sticky="w", pady=2)
            colv[c] = var
        ctk.CTkLabel(win, text="Which terms", font=("Segoe UI", 13, "bold")).pack(anchor="w", padx=14, pady=(10, 2))
        scope = tk.StringVar(value=saved.get("scope", "whole"))
        ctk.CTkRadioButton(win, text="One file", variable=scope, value="whole").pack(anchor="w", padx=18, pady=1)
        ctk.CTkRadioButton(win, text="One file per volume", variable=scope, value="per_volume").pack(anchor="w", padx=18, pady=1)
        vols = [v["name"] for v in self.volumes()]
        volv = {}
        if vols:
            vf = ctk.CTkFrame(win, fg_color="transparent"); vf.pack(fill="x", padx=18, pady=(4, 0))
            ctk.CTkLabel(vf, text="Volumes (none ticked = all terms):").grid(row=0, column=0, columnspan=3, sticky="w")
            for i, n in enumerate(vols):
                var = tk.BooleanVar(value=n in (saved.get("volumes") or []))
                ctk.CTkCheckBox(vf, text=n, variable=var).grid(row=1 + i // 3, column=i % 3, sticky="w", padx=(0, 10), pady=2)
                volv[n] = var
        else:
            ctk.CTkLabel(win, text="No volumes named: One file per volume makes one list per document the terms were found in.",
                         wraplength=600, justify="left", text_color="#64748B").pack(anchor="w", padx=18)
        appr = tk.BooleanVar(value=bool(saved.get("approved_only")))
        ctk.CTkCheckBox(win, text="Approved terms only", variable=appr).pack(anchor="w", padx=14, pady=(10, 2))
        ctk.CTkLabel(win, text="Format", font=("Segoe UI", 13, "bold")).pack(anchor="w", padx=14, pady=(10, 2))
        fmt = tk.StringVar(value=saved.get("fmt", "csv"))
        ff = ctk.CTkFrame(win, fg_color="transparent"); ff.pack(fill="x", padx=18)
        for k, lab in FORMATS.items():
            ctk.CTkRadioButton(ff, text=lab, variable=fmt, value=k).pack(side="left", padx=(0, 14))
        ctk.CTkLabel(win, text="Rejected terms and the omit list are always left out.", text_color="#64748B").pack(anchor="w", padx=14, pady=(10, 0))
        msg = ctk.CTkLabel(win, text="", text_color="#0F766E", wraplength=600, justify="left"); msg.pack(anchor="w", padx=14)

        def go(target=None):
            cols = [c for c in EXPORT_COLS if colv[c].get()]
            if not cols:
                msg.configure(text="Tick at least one column."); return None
            chosen = [n for n, v in volv.items() if v.get()]
            choice = {"cols": cols, "scope": scope.get(), "volumes": chosen, "fmt": fmt.get(), "approved_only": bool(appr.get())}
            self.ws.set("export_choice", choice)
            if target is None:
                if scope.get() == "per_volume":
                    target = filedialog.askdirectory(parent=win, title="Folder for the lists", initialdir=os.path.dirname(L.path))
                else:
                    target = filedialog.asksaveasfilename(parent=win, title="Save the list", initialdir=os.path.dirname(L.path),
                                                          defaultextension="." + fmt.get(),
                                                          initialfile=f"{_safe(L.name)}_Acronyms_{datetime.date.today().isoformat()}.{fmt.get()}",
                                                          filetypes=[(FORMATS[fmt.get()], "*." + fmt.get())])
            if not target:
                return None
            made = self.export_terms(target, cols, scope.get(), chosen or None, fmt.get(), bool(appr.get()))
            msg.configure(text="Saved: " + ", ".join(os.path.basename(p) for p in made))
            return made

        b = ctk.CTkFrame(win, fg_color="transparent"); b.pack(fill="x", padx=14, pady=12)
        ctk.CTkButton(b, text="Export", width=110, fg_color="#0F766E", command=go).pack(side="left")
        ctk.CTkButton(b, text="Close", width=90, fg_color="#475569", command=win.destroy).pack(side="right")
        win._api = {"cols": colv, "scope": scope, "vols": volv, "fmt": fmt, "approved": appr, "go": go, "msg": msg}
        return win

    # ------------------------------------------------------------ use as definition
    def ws_pick_definition(self, term, text):
        ws = getattr(self, "ws", None)
        if ws is None:
            return
        picks = dict(ws.get("picked_definitions", {}) or {})
        picks[term] = text
        ws.set("picked_definitions", picks)

    def apply_picked_definitions(self):
        """After every scan: definitions taken from the documents with Use as definition (this effort only)."""
        ws = getattr(self, "ws", None)
        picks = (ws.get("picked_definitions", {}) or {}) if ws is not None else {}
        for t, text in picks.items():
            d = self.results.get(t)
            if not d or d.get("excluded") or d.get("gone") or d.get("file_error") or d.get("known"):
                continue
            self._pick_into(d, text)

    @staticmethod
    def _pick_into(d, text):
        if text not in d.get("definitions", []):
            d["definitions"] = [text] + [x for x in d.get("definitions", []) if x]
        d["selected_definition"] = text
        d["status"] = "User Selected (from the document)"
        d["picked"] = text

    def use_as_definition(self, text, term=None):
        """Selected words become the definition of the selected term. Returns the definition used, or ''."""
        term = term or getattr(self, "selected_term", None)
        d = self.results.get(term or "")
        text = clean_definition(text, term or "")
        if not d or not text or d.get("file_error"):
            return ""
        self._pick_into(d, text)
        self.ws_pick_definition(term, text)
        diag.activity("use_as_definition")
        if self.effort is not None and term in self.effort.rows:
            if self.effort.add_or_update(term, change_note=f"definition from the document: {text[:80]}", Definition=text) != "unchanged":
                self._save_effort_list()
            self.apply_effort()
        if hasattr(self, "tree") and self.tree.exists(term):
            self.tree.item(term, values=self.row_values(term, d), tags=(self.row_tag(d),))
        try:                                         # Keep (Add to effort list) uses it straight away
            if str(self.entry_def.cget("state")) == "normal":
                self.entry_def.delete(0, "end"); self.entry_def.insert(0, text)
        except Exception:
            diag.swallowed()
        self.session_changed()
        if hasattr(self, "view_var") and self.view_var.get() == VIEW_HOME:
            self.home_refresh()
        return text

    def found_definition(self, term):
        """The ready line under the preview: the long form the documents, the customer's list, the database or the
        nearby words point to, or ''."""
        d = self.results.get(term or "", {})
        if not d or d.get("file_error"):
            return ""
        cands = list(d.get("customer_definitions", [])) + [(d.get("context") or {}).get("definition", "")]
        eng = _eng()
        for x in d.get("definitions", []):
            for tag in (eng.CLIENT_TAG, eng.DB_TAG):
                if x.startswith(tag):
                    x = x[len(tag):]
            cands.append(x)
        cur = (d.get("selected_definition") or "").strip()
        for c in cands:
            c = (c or "").strip()
            if c and not eng.status_text(c).startswith(("[", "MULTIPLE", "Rescan")) and c != cur:
                return c
        return ""

    def scan_landed(self):
        """Every scan lands in Scan review."""
        if hasattr(self, "view_var") and self.effort is not None:
            self.show_view(VIEW_SCAN, save=True)
