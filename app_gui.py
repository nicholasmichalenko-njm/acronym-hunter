import base64
import collections
import os
import sys
import shutil
import atexit
import time

# ==============================================================================
# 1. ZERO-FOOTPRINT FOLDER SANDBOX & OFFLINE NETWORK KILL-SWITCH
#    (Must execute BEFORE importing OCR, OpenCV, or Office libraries)
# ==============================================================================
def get_base_dir():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

BASE_DIR = get_base_dir()
if __name__ == "__main__":   # the Trainer window does "import app_gui": reuse this module, never load a second copy
    sys.modules.setdefault("app_gui", sys.modules[__name__])

SANDBOX_DIR = os.path.join(BASE_DIR, "_local_sandbox")
SANDBOX_TEMP = os.path.join(SANDBOX_DIR, "temp")
try:
    for d in (SANDBOX_DIR, SANDBOX_TEMP):
        os.makedirs(d, exist_ok=True)
except OSError as _e:
    import tkinter as _tk
    from tkinter import messagebox as _mb
    _r = _tk.Tk(); _r.withdraw()
    _mb.showerror("Acronym Hunter cannot start",
                  f"The app cannot write to its own folder:\n{BASE_DIR}\n\n"
                  "Unzip the whole AcronymHunter folder to a location you can write to "
                  "(for example Desktop or Documents), then run AcronymHunter.exe from there.\n\n"
                  f"Details: {_e}")
    sys.exit(1)

# Redirect only temp-file locations into the app folder.
# USERPROFILE / APPDATA / LOCALAPPDATA are deliberately NOT redirected: Windows expands its
# Desktop, Documents and Downloads paths from them, so redirecting breaks the file picker
# ("location is not available"). None of the app's libraries write to those folders.
os.environ["TEMP"] = SANDBOX_TEMP
os.environ["TMP"] = SANDBOX_TEMP
os.environ["TMPDIR"] = SANDBOX_TEMP
os.environ["PYTHONPYCACHEPREFIX"] = os.path.join(SANDBOX_DIR, "pycache")

def cleanup_sandbox():
    if os.path.exists(SANDBOX_TEMP):
        shutil.rmtree(SANDBOX_TEMP, ignore_errors=True)

cleanup_sandbox()                 # clear leftovers from a killed session
os.makedirs(SANDBOX_TEMP, exist_ok=True)   # recreate it: temp files must land here, never elsewhere on disk
atexit.register(cleanup_sandbox)

import traceback
import datetime
LOG_DIR = os.path.join(BASE_DIR, "logs")
import diag   # content-free diagnostics (logs\\events.jsonl): never file names, paths or document content

def log_error(context, with_traceback=True, area=None, **counts):
    """Record an error in logs\\events.jsonl. context must be a fixed message (tools\\check_diag.py enforces it):
    only the message, the error type, code locations and counts are kept."""
    exc = sys.exc_info() if with_traceback and sys.exc_info()[0] is not None else None
    if area is None:
        try:
            area = diag._clean_area(sys._getframe(1).f_code.co_name)
        except Exception:
            area = "app"
    diag.error(area, context, exc=exc, **counts)

def _excepthook(exc_type, exc, tb):
    diag.error("app", "Uncaught error", exc=(exc_type, exc, tb))

diag.setup(LOG_DIR, "?")
atexit.register(diag.shutdown)

import socket
_LOOPBACK = {"127.0.0.1", "::1", "localhost"}
_real_connect = socket.socket.connect
_real_getaddrinfo = socket.getaddrinfo

def _guarded_connect(self, address, *args, **kwargs):
    if isinstance(address, tuple) and address and address[0] not in _LOOPBACK:
        raise OSError("Offline Lockdown Active: outbound network connections are disabled.")
    return _real_connect(self, address, *args, **kwargs)

def _guarded_getaddrinfo(host, *args, **kwargs):
    if host not in (None, "", *_LOOPBACK):
        raise OSError("Offline Lockdown Active: DNS lookups are disabled.")
    return _real_getaddrinfo(host, *args, **kwargs)

socket.socket.connect = _guarded_connect
socket.getaddrinfo = _guarded_getaddrinfo

import urllib.request
def _blocked_network(*args, **kwargs):
    raise RuntimeError("Offline Lockdown Active: Outbound network connections are disabled.")
urllib.request.urlopen = _blocked_network
urllib.request.urlretrieve = _blocked_network

# ==============================================================================
# 2. IMPORTS & BUNDLED ONNX OCR INITIALIZATION
# ==============================================================================
import re
import csv
import json
import io
import getpass
import platform
import tkinter as tk
import codecs
import unicodedata
import tkinter.font as tkfont
import zipfile
import threading
import subprocess
import cv2
import pypdfium2 as pdfium        # Google PDFium (BSD-3 / Apache-2.0): PDF text, pictures, page images
import pypdfium2.raw as _pdfium_raw
import openpyxl
import numpy as np
import customtkinter as ctk
from tkinter import ttk, filedialog, messagebox, simpledialog
from rapidfuzz import process, fuzz
import importlib.util
import symbol_picker as SP
import help_guide as HG
import feedback as FB
from effort_ui import EffortMixin   # Effort List (2.4.0)
from volcheck_ui import VolumeCheckMixin   # volume check and first use (2.7.0)
from compliance_ui import ComplianceMixin   # compliance matrix (2.8.0)
from workspace_ui import WorkspaceMixin   # effort profiles, session, recovery, backups (2.11.0)
from ai_ui import AIMixin   # AI and ML labels and switch (2.12.0)
from setup_ui import SetupMixin   # Setup Builder in the compliance window (2.13.0)
from effort_wizard import WizardMixin   # effort start wizard, intake, review, record, exports (2.14.0)
from effort_home import HomeMixin       # effort home, Scan review, volumes, export picker, Use as definition (2.14.0)
from keywords_ui import KeywordsMixin   # Keyword Harvest (2.16.0)
from master_db import MasterMixin   # master database and Known in (2.14.0)
import equations as EQ             # equations and their symbols (equation-editor math, typed math, definitions)

DB_FILE = os.path.join(BASE_DIR, "acronym_database.csv")

# ---- Text encoding: read any common CSV encoding, always write UTF-8 with BOM (Excel-friendly) ----
CSV_WRITE_ENCODING = "utf-8-sig"   # holds every Unicode character; the BOM tells Excel it is UTF-8

def read_csv_rows(path):
    """DictReader rows from a CSV saved as UTF-8 (with or without BOM), UTF-16 (Excel "Unicode Text",
    tab-delimited), or legacy Windows-1252 ("CSV (Comma delimited)" from older Excel)."""
    text, delimiter = _decode_csv(path)
    return list(csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter))

def read_csv_header(path):
    text, delimiter = _decode_csv(path)
    return next(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter), [])

def _decode_csv(path):
    raw = open(path, "rb").read()
    if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        text = raw.decode("utf-16")
    elif raw.startswith(codecs.BOM_UTF8):
        text = raw[len(codecs.BOM_UTF8):].decode("utf-8")
    else:
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            text = raw.decode("cp1252", errors="replace")
            log_error("A CSV file is not UTF-8; read as Windows-1252", with_traceback=False, ext=".csv")
    header = text.split("\n", 1)[0]
    delimiter = "\t" if header.count("\t") > header.count(",") else ","
    return text, delimiter

# ---- Database columns: any order, common header names, extra columns kept ----
DB_FIELDS = ("Term", "Definition", "Category", "Action")   # order for new files
DB_ALIASES = {
    "term": "Term", "item": "Term", "acronym": "Term", "abbreviation": "Term", "acronym / abbreviation": "Term",
    "definition": "Definition", "meaning": "Definition", "description": "Definition", "expansion": "Definition",
    "category": "Category", "type": "Category",
    "action": "Action", "rule": "Action", "include/exclude": "Action",
}

def db_field(header):
    return DB_ALIASES.get((header or "").strip().lstrip("\ufeff").casefold())

def db_row(r):
    """A database row as {Term, Definition, Category, Action} whatever the column names or order."""
    out = {}
    for k, v in r.items():
        f = db_field(k)
        if f and f not in out:
            out[f] = (v or "").strip()
    return out

# Look-alike characters that PDFs and Office mix freely; folded so a database term matches either form.
_FOLD = str.maketrans({
    "\u00b5": "\u03bc",   # micro sign -> Greek mu
    "\u2126": "\u03a9",   # ohm sign -> Greek capital omega
    "\u212b": "\u00c5",   # angstrom sign -> A with ring
    "\u212a": "K",        # kelvin sign -> K
    "\u2212": "-", "\u2010": "-", "\u2011": "-", "\u2013": "-",   # minus, hyphens, en dash -> hyphen
    "\u00a0": " ", "\u202f": " ",                                  # non-breaking spaces -> space
    "\ufb01": "fi", "\ufb02": "fl",                                # typographic ligatures
})

def fold(text):
    """Unicode-normalize (NFC) and fold look-alike characters, applied to database terms and document text."""
    return unicodedata.normalize("NFC", text).translate(_FOLD)

# ---- User-tunable accuracy files (config\ folder beside the exe) ----
CONFIG_DIR = os.path.join(BASE_DIR, "config")
CHOICES_FILE = os.path.join(CONFIG_DIR, "choices.json")   # Apply Choice, kept for every later scan on this PC
OCR_FIXES_FILE = os.path.join(CONFIG_DIR, "ocr_corrections.csv")
OCR_IGNORE_FILE = os.path.join(CONFIG_DIR, "ocr_ignore.csv")   # OCR readings that are not text (diagram lines, specks)
OCR_SETTINGS_FILE = os.path.join(CONFIG_DIR, "ocr_settings.json")
MODELS_DIR = os.path.join(CONFIG_DIR, "models")
OCR_DEFAULTS = {
    "min_confidence": 0.5,        # drop OCR lines the engine is less sure of (0-1)
    "pdf_scan_dpi": 300,          # render resolution for scanned PDF pages (300 beat 200 in tools\\evaluate.py)
    "scanned_page_max_chars": 40, # PDF page with fewer digital characters is treated as scanned
    "upscale_below_px": 1200,     # enlarge images narrower than this before OCR
    "upscale_factor": 2.0,
    "contrast_clip": 2.0,         # CLAHE contrast boost; 0 disables
    "denoise_strength": 10,       # 0 disables denoising (faster, better for clean screenshots)
    "detect_side_len": 736,       # text-detection resolution; raise for tiny text
    "detect_box_thresh": 0.5,     # lower finds fainter text boxes (more noise)
    "detect_unclip_ratio": 1.6,   # how far each detected text box is expanded
    "fuzzy_min_score": 80,        # RapidFuzz score (0-100) to accept an OCR near-miss of a DB term
    "tile_large_images": 1,       # 1 = read large images in overlapping tiles so small text is not shrunk away
    "tile_size_px": 1800,         # tile edge; keep under 2000 (the OCR engine shrinks anything larger)
    "tile_overlap_px": 300,       # overlap so words on a tile edge are read whole in the next tile
    "tesseract_min_confidence": 0.30,  # Tesseract (US edition) line confidence floor; its scores run lower than RapidOCR's
    "doctr_enabled": 1,           # 1 = also read pictures with docTR (neural text finder + reader), add what Tesseract missed
    "doctr_min_confidence": 0.5,  # docTR word confidence floor (0-1)
    "doctr_extra_passes": 1,      # 1 = inverted pass for dark pictures and turned passes for vertical labels
    "doctr_override_confidence": 0.95,  # where Tesseract read text, a docTR word this sure that Tesseract lacks is added
}
OCR_ENGINE_MAX_SIDE = 2000        # RapidOCR downsizes images whose longest side exceeds this

# ---- OCR engine choice ----
# "tesseract": Tesseract 5 (Apache 2.0; HP origin, developed by Google) with Google-trained models. US edition.
# "rapidocr":  RapidOCR + PaddleOCR PP-OCRv4 models (Apache 2.0; Baidu / RapidAI). Legacy edition.
# The build decides which engine ships; AH_OCR_ENGINE overrides for testing.
TESS_MIN_WORD_CONF = 30           # Tesseract word confidence (0-100) below which a word is dropped as noise
TESS_PSM = 11                     # page segmentation: 11 = sparse text (diagram labels anywhere on the image)
CUSTOM_TESS_MODEL = "acronym"     # config\models\acronym.traineddata = model fine-tuned on your Trainer samples

def _find_tesseract():
    for d in (os.environ.get("AH_TESSERACT_DIR", ""), os.path.join(getattr(sys, "_MEIPASS", ""), "tesseract"),
              os.path.join(BASE_DIR, "_internal", "tesseract"), os.path.join(BASE_DIR, "vendor", "tesseract")):
        if d and os.path.isfile(os.path.join(d, "tesseract.exe")):
            return d
    return None

TESS_DIR = _find_tesseract()

DOCTR_FILES = ("doctr_det.onnx", "doctr_reco.onnx", "doctr_models.json")

def doctr_folder_ok(d):
    """A docTR model folder is used only when every ONNX file matches the SHA-256 its doctr_models.json records."""
    import hashlib
    try:
        info = json.load(open(os.path.join(d, "doctr_models.json"), encoding="utf-8"))
        for key in ("doctr_det", "doctr_reco"):
            m = info["models"][key]
            path = os.path.join(d, m["onnx"])
            h = hashlib.sha256()
            with open(path, "rb") as f:
                for block in iter(lambda: f.read(1 << 20), b""):
                    h.update(block)
            if h.hexdigest().upper() != str(m["onnx_sha256"]).upper():
                return False
        return True
    except Exception:
        diag.swallowed()
        return False

def _find_doctr():
    """Folder with docTR's ONNX models. Trained models installed in config\\models\\doctr (Training Kit or a
    Knowledge Pack) win when their SHA-256 checks pass; otherwise the bundled models (Mindee, Apache 2.0)."""
    trained = os.path.join(BASE_DIR, "config", "models", "doctr")
    if os.path.isfile(os.path.join(trained, "doctr_models.json")):
        if doctr_folder_ok(trained):
            return trained
        log_error("Trained docTR models failed their SHA-256 check; using the bundled models", with_traceback=False)
    for d in (os.environ.get("AH_DOCTR_DIR", ""), os.path.join(getattr(sys, "_MEIPASS", ""), "doctr"),
              os.path.join(BASE_DIR, "_internal", "doctr"), os.path.join(BASE_DIR, "vendor", "doctr")):
        if d and os.path.isfile(os.path.join(d, "doctr_models.json")):
            return d
    return None

DOCTR_DIR = _find_doctr()
EDITION_NAME = "US Classic (Tesseract only)" if not DOCTR_DIR else "US (Tesseract + docTR second reader)"

def _find_math():
    """Math reader add-on folder (math_addon folder next to the app, from AcronymHunter_MathAddon.zip). None = not installed."""
    for d in (os.environ.get("AH_MATH_DIR", ""), os.path.join(BASE_DIR, "math_addon"), os.path.join(BASE_DIR, "dist_math", "math_addon")):
        if d and os.path.isfile(os.path.join(d, "AcronymHunterMath.exe")) and os.path.isfile(os.path.join(d, "math_models.json")):
            return d
    return None

MATH_DIR = _find_math()
MATH_CACHE_FILE = os.path.join(BASE_DIR, "config", "math_results.json")   # picture fingerprint -> equations read
SP.set_storage(os.path.join(BASE_DIR, "config", "symbol_recent.json"))     # the symbol picker's "Recent" row
OCR_BACKEND = os.environ.get("AH_OCR_ENGINE") or ("tesseract" if TESS_DIR else "rapidocr")

def load_ocr_settings():
    cfg = dict(OCR_DEFAULTS)
    try:
        with open(OCR_SETTINGS_FILE, "r", encoding="utf-8-sig") as f:
            user = json.load(f)
        for k, v in user.items():
            if k in OCR_DEFAULTS and isinstance(v, (int, float)) and not isinstance(v, bool):
                cfg[k] = v
    except FileNotFoundError:
        pass
    except Exception:
        log_error("OCR settings file could not be read; using defaults")
    return cfg

OCR_CFG = load_ocr_settings()

APP_SETTINGS_FILE = os.path.join(CONFIG_DIR, "app_settings.json")
APP_DEFAULTS = {
    "advanced_features_default": False,   # Advanced tools (Fix OCR, OCR Fixes, Region Trainer) start hidden
    "allow_advanced_toggle": True,        # False = hide the Advanced switch; the default above is locked in
    "allow_training": True,               # False = no Trainer, no "Check in picture", no "Improve Reading" (nothing saved)
    "ocr_sandbox": True,                  # run Tesseract inside a Windows AppContainer with no network permission
    "ocr_sandbox_required": False,        # True = refuse to read pictures if the sandbox cannot start
    "usage_logging": True,                # False = do not record which controls are clicked (logs\\activity.jsonl)
    "allow_model_features": True,         # False = no docTR second reader and no math reader add-on (Tesseract only)
}

def load_app_settings():
    cfg = dict(APP_DEFAULTS)
    try:
        with open(APP_SETTINGS_FILE, "r", encoding="utf-8-sig") as f:
            for k, v in json.load(f).items():
                if k in APP_DEFAULTS and isinstance(v, bool):
                    cfg[k] = v
    except FileNotFoundError:
        pass
    except Exception:
        log_error("App settings file could not be read; using defaults")
    return cfg

APP_CFG = load_app_settings()

def load_ocr_fixes():
    """Misread -> Correct pairs from config/ocr_corrections.csv. Re-read on every scan."""
    fixes = {}
    try:
        for row in read_csv_rows(OCR_FIXES_FILE):
            wrong, right = fold((row.get("Misread") or "").strip()), fold((row.get("Correct") or "").strip())
            if wrong and right and wrong != right:
                fixes[wrong] = right
    except FileNotFoundError:
        pass
    except Exception:
        log_error("OCR fixes file could not be read")
    return fixes
# ---- OCR ignore list: exact readings that come from lines, arrows, or specks, never real text ----
_OCR_IGNORE = set()

def load_ocr_ignore():
    """Readings from config/ocr_ignore.csv. Applies only to text read by OCR, never to typed document text."""
    out = set()
    try:
        for row in read_csv_rows(OCR_IGNORE_FILE):
            r = fold((row.get("Reading") or "").strip())
            if r:
                out.add(r)
    except FileNotFoundError:
        pass
    except Exception:
        log_error("OCR ignore list could not be read")
    return out

def read_ocr_ignore_rows():
    try:
        return [[(r.get("Reading") or "").strip(), (r.get("Note") or "").strip()] for r in read_csv_rows(OCR_IGNORE_FILE)
                if (r.get("Reading") or "").strip()]
    except FileNotFoundError:
        return []

def write_ocr_ignore_rows(rows):
    os.makedirs(CONFIG_DIR, exist_ok=True)
    tmp = OCR_IGNORE_FILE + ".tmp"
    with open(tmp, "w", newline="", encoding=CSV_WRITE_ENCODING) as f:
        w = csv.writer(f)
        w.writerow(["Reading", "Note"])
        w.writerows(rows)
    os.replace(tmp, OCR_IGNORE_FILE)

def refresh_ocr_ignore():
    global _OCR_IGNORE
    _OCR_IGNORE = load_ocr_ignore()
    return _OCR_IGNORE

IMAGE_EXTS = {'.png', '.jpg', '.jpeg', '.tif', '.tiff', '.bmp', '.webp', '.gif'}
VECTOR_FILE_EXTS = {'.svg', '.svgz', '.emf', '.wmf', '.emz', '.wmz'}
SUPPORTED_EXTS = IMAGE_EXTS | VECTOR_FILE_EXTS | {'.pdf', '.pptx', '.docx', '.xlsx'}

# Create an EMPTY database (header only) if missing. Terms come from the user or a Knowledge Pack.
if not os.path.exists(DB_FILE):
    with open(DB_FILE, "w", newline="", encoding=CSV_WRITE_ENCODING) as f:
        csv.writer(f).writerow(list(DB_FIELDS))

# RapidOCR loads lazily on first image (uses pre-bundled offline ONNX weights)
_ocr_engine = None
_ocr_lock = threading.Lock()

OCR_MODEL_INFO = "not loaded yet"

def ocr_cutoff():
    """Confidence a reading needs to be kept by the scanner (engine-specific)."""
    return float(OCR_CFG["tesseract_min_confidence" if OCR_BACKEND == "tesseract" else "min_confidence"])

OCR_SANDBOX_STATUS = "not used yet"

def run_tesseract_tool(cmd, input_bytes=b"", timeout=180, read_dirs=(), write_dirs=()):
    """Run a Tesseract program (cmd[0] in TESS_DIR). Returns (exit code, stdout, stderr).
    Normally inside a Windows AppContainer with no network permission (sandbox.py). If the sandbox cannot start
    (for example the app folder is on a network share), it runs as a plain child process unless
    ocr_sandbox_required is true. The Tesseract build itself contains no network code either way."""
    global OCR_SANDBOX_STATUS
    env = dict(os.environ, OMP_THREAD_LIMIT="1")
    if APP_CFG["ocr_sandbox"] and os.name == "nt":
        try:
            import sandbox
            sandbox.grant(TESS_DIR)
            for d in read_dirs:
                sandbox.grant(d)
            for d in write_dirs:
                sandbox.grant(d, write=True)
            result = sandbox.run(list(cmd), input_bytes, timeout=timeout, env=env)
            OCR_SANDBOX_STATUS = "on: Windows AppContainer with no network permission"
            return result
        except TimeoutError:
            raise
        except Exception as e:
            log_error("OCR sandbox could not start")
            OCR_SANDBOX_STATUS = f"could not start ({type(e).__name__}: {e})"
            if APP_CFG["ocr_sandbox_required"]:
                raise RuntimeError("The OCR sandbox could not start and config\\app_settings.json requires it "
                                   "(ocr_sandbox_required). Text in pictures was not read.")
    elif not APP_CFG["ocr_sandbox"]:
        OCR_SANDBOX_STATUS = "off (config\\app_settings.json: ocr_sandbox = false)"
    r = subprocess.run(list(cmd), input=input_bytes, capture_output=True, timeout=timeout, env=env,
                       creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    return r.returncode, r.stdout, r.stderr

_TESS_LEAD_ONE = re.compile(r"^[1l|](?=[A-Z]{2})")

class TesseractEngine:
    """Tesseract 5 run as a child process (tesseract.exe from the app's own folder; images go in through stdin,
    results come back through stdout, nothing is written to disk). Called like RapidOCR so the scanner and the
    Trainer work with either engine:
        engine(img)                -> ([[box, text, confidence 0-1], ...], None)   one entry per text line
        engine(img, use_det=False) -> ([[text, confidence 0-1]], None)             the whole image read as one line
    Tesseract has no network code path in this use: the app passes only "stdin", never a file name or URL."""

    def __init__(self, exe_dir):
        global OCR_MODEL_INFO
        self.exe = os.path.join(exe_dir, "tesseract.exe")
        custom = os.path.join(MODELS_DIR, CUSTOM_TESS_MODEL + ".traineddata")
        if os.path.isfile(custom):
            self.tessdata, self.lang = MODELS_DIR, CUSTOM_TESS_MODEL
            OCR_MODEL_INFO = "Tesseract 5 with your trained model (acronym.traineddata)"
        else:
            self.tessdata, self.lang = os.path.join(exe_dir, "tessdata"), "eng"
            OCR_MODEL_INFO = "Tesseract 5 with Google's English model (tessdata_best)"

    def _run(self, img, psm):
        img = cv2.copyMakeBorder(img, 12, 12, 12, 12, cv2.BORDER_CONSTANT, value=255 if img.ndim == 2 else (255, 255, 255))
        ok, buf = cv2.imencode(".png", img)
        if not ok:
            return []
        cmd = [self.exe, "stdin", "stdout", "--tessdata-dir", self.tessdata, "-l", self.lang, "--psm", str(psm),
               "--dpi", "300", "-c", "tessedit_create_tsv=1", "-c", "tessedit_create_txt=0"]
        code, out, err = run_tesseract_tool(cmd, buf.tobytes(), timeout=180,
                                            read_dirs=[self.tessdata] if self.tessdata != os.path.join(TESS_DIR, "tessdata") else [])
        if code != 0:
            raise RuntimeError("Tesseract failed: " + err.decode("utf-8", "replace").strip()[-300:])
        lines = {}
        for row in out.decode("utf-8", "replace").splitlines()[1:]:
            c = row.split("\t")
            if len(c) < 12 or c[0] != "5":
                continue
            text, conf = c[11].strip(), float(c[10])
            if not text or conf < TESS_MIN_WORD_CONF:
                continue
            text = _TESS_LEAD_ONE.sub("I", text)   # "1CD-204" -> "ICD-204": a diagram line touching a capital I
            x, y, w, h = (int(v) - (12 if i < 2 else 0) for i, v in enumerate(c[6:10]))
            lines.setdefault((c[2], c[3], c[4]), []).append((x, y, x + w, y + h, text, conf))
        out = []
        for words in lines.values():
            words.sort(key=lambda t: t[0])
            x0, y0 = min(t[0] for t in words), min(t[1] for t in words)
            x1, y1 = max(t[2] for t in words), max(t[3] for t in words)
            conf = sum(t[5] for t in words) / len(words) / 100.0
            out.append(([[x0, y0], [x1, y0], [x1, y1], [x0, y1]], " ".join(t[4] for t in words), conf))
        return out

    def __call__(self, img, use_det=True, text_score=None, **_ignored):
        if not use_det:
            lines = self._run(img, 7)
            if not lines:
                return [], None
            return [[" ".join(t[1] for t in lines), sum(t[2] for t in lines) / len(lines)]], None
        floor = float(OCR_CFG["tesseract_min_confidence"]) if text_score is None else float(text_score)
        return [list(t) for t in self._run(img, TESS_PSM) if t[2] >= floor], None

def _build_ocr():
    """Tesseract (US edition) or RapidOCR (legacy edition), per OCR_BACKEND."""
    global OCR_MODEL_INFO
    if OCR_BACKEND == "tesseract":
        if not TESS_DIR:
            raise RuntimeError("Tesseract engine files are missing from this copy of the app.")
        return TesseractEngine(TESS_DIR)
    from rapidocr_onnxruntime import RapidOCR
    try:
        import onnxruntime as _ort   # ONNX Runtime telemetry (Windows TraceLogging) off
        _ort.disable_telemetry_events()
    except Exception:
        diag.swallowed()
        pass
    try:
        return _build_ocr_configured()
    except Exception:
        log_error("Custom OCR configuration failed; falling back to bundled models and default settings")
        OCR_MODEL_INFO = "bundled PP-OCRv4 (custom settings/models failed; see the logs folder)"
        return RapidOCR()

def _build_ocr_configured():
    global OCR_MODEL_INFO
    import yaml
    import rapidocr_onnxruntime as _ro
    from rapidocr_onnxruntime import RapidOCR
    with open(os.path.join(os.path.dirname(_ro.__file__), "config.yaml"), "r", encoding="utf-8") as f:
        cfg = yaml.safe_load(f)
    cfg["Global"]["text_score"] = float(OCR_CFG["min_confidence"])
    cfg["Det"]["limit_side_len"] = int(OCR_CFG["detect_side_len"])
    cfg["Det"]["box_thresh"] = float(OCR_CFG["detect_box_thresh"])
    cfg["Det"]["unclip_ratio"] = float(OCR_CFG["detect_unclip_ratio"])
    custom = []
    for part, section in (("det", "Det"), ("cls", "Cls"), ("rec", "Rec")):
        path = os.path.join(MODELS_DIR, f"{part}.onnx")
        if os.path.isfile(path):
            cfg[section]["model_path"] = os.path.abspath(path)
            custom.append(f"{part}.onnx")
    keys = os.path.join(MODELS_DIR, "rec_keys.txt")
    if os.path.isfile(keys):
        cfg["Rec"]["rec_keys_path"] = os.path.abspath(keys)
        custom.append("rec_keys.txt")
    os.makedirs(SANDBOX_TEMP, exist_ok=True)
    cfg_path = os.path.join(SANDBOX_TEMP, "rapidocr_config.yaml")
    with open(cfg_path, "w", encoding="utf-8") as f:
        yaml.safe_dump(cfg, f)
    try:
        engine = RapidOCR(config_path=cfg_path)
        OCR_MODEL_INFO = ("custom " + ", ".join(custom)) if custom else "bundled PP-OCRv4"
    except Exception:
        log_error("Custom OCR model failed to load; using bundled models")
        engine = RapidOCR()
        OCR_MODEL_INFO = "bundled PP-OCRv4 (custom model failed to load; see the logs folder)"
    return engine

# ---- AI and ML switch (2.12.0): per effort, on by default. Off = no trained model runs at all. ----
_AIML = {"on": True}

def ai_ml_on():
    return _AIML["on"]

def set_ai_ml(on):
    _AIML["on"] = bool(on)

def math_available():
    """The math reader add-on is installed and AI and ML are on."""
    return bool(MATH_DIR) and ai_ml_on()

def get_ocr():
    if not ai_ml_on():
        raise RuntimeError("AI and ML are off for this effort, so text inside pictures is not read.")
    global _ocr_engine
    with _ocr_lock:
        if _ocr_engine is None:
            _ocr_engine = _build_ocr()
    return _ocr_engine

# ==============================================================================
# 3. UNIFIED DATABASE LOADER (INCLUDES & EXCLUDES)
# ==============================================================================
def db_layout():
    """[(field or None, header text)] in the file's own column order; DB_FIELDS order for a new file."""
    try:
        header = [h.strip().lstrip("\ufeff") for h in read_csv_header(DB_FILE)]
    except (FileNotFoundError, StopIteration):
        diag.swallowed()
        header = []
    layout = [(db_field(h), h) for h in header if h]
    have = {f for f, _ in layout if f}
    layout += [(f, f) for f in DB_FIELDS if f not in have]   # add any missing core column at the end
    return layout

KNOWN_ACTION = "Omit"                 # database Action: the omit list (company names, locations...) (2.13.0)
KNOWN_BLANK = "(on the omit list)"
KNOWN_STATUS = "Found in omit list"
OMIT_REASON_COL = "Omit reason"
OMIT_REASONS = ["Company or organization", "Location", "Program or product", "Person or role", "Other"]

def load_database():
    include_db = {}
    exclude_db = {}
    exclude_upper = set()

    if not os.path.exists(DB_FILE):
        return include_db, exclude_db, exclude_upper

    for raw_row in read_csv_rows(DB_FILE):
        row = db_row(raw_row)
        term = fold(row.get("Term", ""))
        cat = row.get("Category") or "Acronym"
        defn = row.get("Definition", "")
        action = (row.get("Action") or "Include").capitalize()

        if not term:
            continue

        if action == "Exclude" or cat.lower() in {"exclude", "ignore", "noise"}:
            exclude_db[term] = {"category": cat, "definition": defn or "Excluded by user rule"}
            exclude_upper.add(term.upper())
        elif action.lower().startswith(("known", "omit")):   # 2.13.0: the omit list: recognized, kept, left off the lists
            if term not in include_db:
                include_db[term] = {"category": cat, "definitions": []}
            include_db[term]["known"] = True
            include_db[term]["omit_reason"] = (raw_row.get(OMIT_REASON_COL) or "").strip() or include_db[term].get("omit_reason", "")
            d = defn or KNOWN_BLANK
            if d not in include_db[term]["definitions"]:
                include_db[term]["definitions"].append(d)
        else:
            if not defn:
                continue
            if term not in include_db:
                include_db[term] = {"category": cat, "definitions": []}
            if defn not in include_db[term]["definitions"]:
                include_db[term]["definitions"].append(defn)

    return include_db, exclude_db, exclude_upper

# ==============================================================================
# 4. LAYER 2 & 3: OPENCV IMAGE PRE-PROCESSING & ONNX OCR READER
# ==============================================================================
def enhance_image_for_ocr(image_bytes):
    np_arr = np.frombuffer(image_bytes, np.uint8)
    img = cv2.imdecode(np_arr, cv2.IMREAD_COLOR)
    if img is None:
        return None
    height, width = img.shape[:2]
    # Ignore tiny 1x1 or icon-sized decorative graphics under 25px
    if width < 25 or height < 15:
        return None
    if width < OCR_CFG["upscale_below_px"] and OCR_CFG["upscale_factor"] > 1:
        f = float(OCR_CFG["upscale_factor"])
        img = cv2.resize(img, None, fx=f, fy=f, interpolation=cv2.INTER_CUBIC)
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    if OCR_CFG["contrast_clip"] > 0:
        gray = cv2.createCLAHE(clipLimit=float(OCR_CFG["contrast_clip"]), tileGridSize=(8, 8)).apply(gray)
    if OCR_CFG["denoise_strength"] > 0 and OCR_BACKEND != "tesseract":   # Tesseract binarizes on its own; denoise only cost time
        gray = cv2.fastNlMeansDenoising(gray, None, h=float(OCR_CFG["denoise_strength"]),
                                        templateWindowSize=7, searchWindowSize=21)
    return gray

# ---------------------------------------------------------------------------
# Shared helpers: OCR cache, skipped-content notes, Office XML parsing
# ---------------------------------------------------------------------------
import hashlib
import posixpath
from lxml import etree as _ET

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "v": "urn:schemas-microsoft-com:vml",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ss": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
}
def _q(prefix, tag):
    return "{%s}%s" % (NS[prefix], tag)

VECTOR_EXTS = {".emf", ".wmf", ".emz", ".wmz", ".svg", ".svgz"}
NS_ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"   # PowerPoint/Word SVG picture companion
VECTOR_TEXT_MIN_CHARS = 3   # vector graphic with less real text than this is drawn as a picture and OCR'd
_OCR_CACHE = {}   # sha1(image bytes) -> text; a logo repeated on 50 slides is OCR'd once
_SKIPPED = []     # (location label, reason) for content that could not be scanned

def note_skip(label, reason):
    if (label, reason) not in _SKIPPED:
        _SKIPPED.append((label, reason))

def _ocr_lines(img, x_off=0, y_off=0):
    """[(x0, y0, x1, y1, text)] for lines at or above min_confidence, in full-image coordinates."""
    result, _ = get_ocr()(img)
    out = []
    for box, text, conf in (result or []):
        if float(conf) < float(OCR_CFG["tesseract_min_confidence" if OCR_BACKEND == "tesseract" else "min_confidence"]) or not str(text).strip():
            continue
        if fold(str(text)).strip() in _OCR_IGNORE:   # marked "Not text" in the Trainer
            continue
        xs = [pt[0] for pt in box]; ys = [pt[1] for pt in box]
        out.append((min(xs) + x_off, min(ys) + y_off, max(xs) + x_off, max(ys) + y_off, str(text), float(conf)))
    return out

def _tile_starts(length, tile, step):
    if length <= tile:
        return [0]
    starts = list(range(0, length - tile + 1, step))
    if starts[-1] != length - tile:
        starts.append(length - tile)
    return starts

_doctr_engine = None
_doctr_failed = [False]
_allow_models = []

def _models_allowed():
    """Admin lock (config\\app_settings.json "allow_model_features": false): no docTR, no math reader."""
    if not _allow_models:
        try:
            _allow_models.append(bool(load_app_settings().get("allow_model_features", True)))
        except Exception:
            diag.swallowed()
            _allow_models.append(True)
    return _allow_models[0]

if not _models_allowed():
    MATH_DIR = None

def doctr_active():
    if not _models_allowed() or not ai_ml_on():
        return False
    return bool(DOCTR_DIR) and OCR_BACKEND == "tesseract" and bool(int(OCR_CFG.get("doctr_enabled", 1))) and not _doctr_failed[0]

def get_doctr():
    global _doctr_engine
    with _ocr_lock:
        if _doctr_engine is None:
            import doctr_engine
            _doctr_engine = doctr_engine.DocTREngine(DOCTR_DIR)
    return _doctr_engine

def doctr_info():
    if not DOCTR_DIR:
        return "not bundled"
    try:
        info = json.load(open(os.path.join(DOCTR_DIR, "doctr_models.json"), encoding="utf-8"))
        import onnxruntime as _ort
        m = info["models"]
        trained = (f"trained here by the Training Kit ({info.get('trained', '?')}), "
                   if info.get("trained_by") else "exported by this project from Mindee's weights, ")
        return (f"docTR (Mindee, Apache 2.0) {m['doctr_det']['arch']} + {m['doctr_reco']['arch']} on ONNX Runtime {_ort.__version__}, "
                + trained + ("on" if doctr_active() else "off"))
    except Exception as e:
        diag.swallowed()
        return f"unavailable ({type(e).__name__})"

def _decode_rgb(image_bytes):
    arr = cv2.imdecode(np.frombuffer(image_bytes, np.uint8), cv2.IMREAD_UNCHANGED)
    if arr is None:
        return None
    if arr.ndim == 2:
        arr = cv2.cvtColor(arr, cv2.COLOR_GRAY2BGR)
    elif arr.shape[2] == 4:                       # transparent picture: flatten onto white
        a = arr[:, :, 3:4].astype(np.float32) / 255.0
        arr = (arr[:, :, :3].astype(np.float32) * a + 255.0 * (1 - a)).astype(np.uint8)
    if arr.dtype != np.uint8:
        arr = cv2.convertScaleAbs(arr, alpha=255.0 / max(1, int(arr.max())))
    return cv2.cvtColor(arr, cv2.COLOR_BGR2RGB)

def _norm_tok(t):
    return re.sub(r"[^0-9a-z&]", "", t.lower())

def _doctr_extra(image_bytes, tess_lines, clean_shape):
    """Words docTR reads that Tesseract did not, as extra lines.
    - In a region Tesseract read, a docTR word is added only if Tesseract has no word there that matches it
      (capitals and punctuation ignored, or a close spelling). Tesseract stays the reader of record there.
    - In a region Tesseract found nothing, all of docTR's words are added: that is the text Tesseract missed."""
    if not doctr_active():
        return ""
    try:
        rgb = _decode_rgb(image_bytes)
        if rgb is None or min(rgb.shape[:2]) < 8:
            return ""
        check_cancel()
        lines = get_doctr().read(rgb, extra_passes=bool(int(OCR_CFG.get("doctr_extra_passes", 1))))
    except ScanCancelled:
        raise
    except Exception as e:
        diag.swallowed()
        _doctr_failed[0] = True
        log_error("docTR could not read a picture; it is switched off for this session")
        note_skip("docTR", f"docTR failed, so pictures were read by Tesseract only: {type(e).__name__}: {e}. Details are in the logs folder.")
        return ""
    from rapidfuzz import fuzz
    sx = clean_shape[1] / max(1, rgb.shape[1]); sy = clean_shape[0] / max(1, rgb.shape[0])
    floor = float(OCR_CFG.get("doctr_min_confidence", 0.5))
    out, seen = [], set()
    for ln in lines:
        box, text, conf = ln[0], ln[1], ln[2]
        words = ln[3] if len(ln) > 3 else [(t, conf) for t in text.split()]
        if fold(text).strip() in _OCR_IGNORE:
            continue
        xs = [p[0] * sx for p in box]; ys = [p[1] * sy for p in box]
        bx0, by0, bx1, by1 = min(xs), min(ys), max(xs), max(ys)
        area = max(1.0, (bx1 - bx0) * (by1 - by0))
        near = [t for t in tess_lines
                if max(0, min(bx1, t[2]) - max(bx0, t[0])) * max(0, min(by1, t[3]) - max(by0, t[1])) > 0.3 * area]
        tess_toks = [x for x in (_norm_tok(x) for t in near for x in t[4].split()) if x]
        new = []
        for tok, wconf in words:
            n = _norm_tok(tok)
            if not n or wconf < floor:
                continue
            if tess_toks:
                # Tesseract read this region: keep its reading unless docTR is sure and has more than Tesseract
                # and its word contains Tesseract's (TWTA where Tesseract read WTA). A docTR word inside a Tesseract
                # word ("TES" in NOTES) or merely different ("SOV" for SOW) never counts.
                if (wconf < float(OCR_CFG.get("doctr_override_confidence", 0.95)) or any(n in x for x in tess_toks)
                        or not any(len(x) >= 2 and x in n for x in tess_toks)):
                    continue
            new.append(tok)
        ln_txt = " ".join(new)
        if ln_txt and ln_txt not in seen:
            seen.add(ln_txt); out.append(ln_txt)
    return "\n".join(out)

# ---------------------------------------------------------------------------
# Math reader add-on (IBM Granite-Docling on PyTorch, in its own program, inside the no-network sandbox)
# ---------------------------------------------------------------------------
_math_cache = None

def math_cache():
    global _math_cache
    if _math_cache is None:
        try:
            _math_cache = json.load(open(MATH_CACHE_FILE, encoding="utf-8"))
        except Exception:
            diag.swallowed()
            _math_cache = {}
    return _math_cache

def save_math_cache():
    try:
        tmp = MATH_CACHE_FILE + ".tmp"
        json.dump(math_cache(), open(tmp, "w", encoding="utf-8"), ensure_ascii=False)
        os.replace(tmp, MATH_CACHE_FILE)
    except Exception:
        log_error("Could not save math_results.json")

def math_info():
    if not MATH_DIR:
        return "not installed (AcronymHunter_MathAddon.zip adds it)"
    try:
        man = json.load(open(os.path.join(MATH_DIR, "math_models.json"), encoding="utf-8"))
        return f"installed: {man.get('model', 'Granite-Docling')}; runs only in the no-network sandbox"
    except Exception as e:
        diag.swallowed()
        return f"damaged ({type(e).__name__})"

def run_math_reader(items, on_status, on_result, cancel):
    """items: [(id, png bytes)]. Calls on_result(id, formulas, cached, seconds) as each picture is read.
    Cached pictures return at once. The add-on program runs ONLY inside the AppContainer sandbox (no network);
    if the sandbox cannot start, nothing runs."""
    cache = math_cache()
    todo = []
    for pid, png in items:
        key = hashlib.sha1(png).hexdigest()
        if key in cache:
            on_result(pid, cache[key], True, 0.0)
        else:
            todo.append((pid, png, key))
    if not todo:
        return
    if not MATH_DIR:
        raise RuntimeError("The math reader add-on is not installed.")
    if not ai_ml_on():
        raise RuntimeError("AI and ML are off for this effort, so the math reader does not run.")
    import sandbox
    work = os.path.join(BASE_DIR, "training_data", "math_work")
    try:
        sandbox.grant(MATH_DIR, write=False)
        sandbox.grant(work, write=True)
    except Exception as e:
        diag.swallowed()
        raise RuntimeError(f"The math reader runs only inside the no-network sandbox, and the sandbox could not start: {e}")
    keys = {pid: key for pid, _png, key in todo}
    # Windows gives a sandboxed program its own temp folder (LOCALAPPDATA\Packages\<sandbox>\AC\Temp) based on these
    # variables, so they keep their real values. Model caches point at the one folder the sandbox may write.
    env = {k: os.environ[k] for k in ("SystemRoot", "SystemDrive", "NUMBER_OF_PROCESSORS", "PROCESSOR_ARCHITECTURE", "OS",
                                      "LOCALAPPDATA", "APPDATA", "USERPROFILE", "HOMEDRIVE", "HOMEPATH") if k in os.environ}
    env.update({"HF_HOME": work, "TORCH_HOME": work, "XDG_CACHE_HOME": work,
                "PATH": os.path.join(os.environ.get("SystemRoot", r"C:\Windows"), "System32")})
    payload = json.dumps({"items": [{"id": pid, "png": base64.b64encode(png).decode("ascii")} for pid, png, _k in todo],
                          "threads": os.cpu_count() or 4}).encode("utf-8")
    errors = []
    def line(b):
        try:
            obj = json.loads(b.decode("utf-8", "replace"))
        except Exception:
            diag.swallowed()
            return
        if obj.get("ready"):
            on_status(f"Math reader ready ({obj.get('load_sec')} s to load)")
        elif "error" in obj and "id" not in obj:
            errors.append(obj["error"])
        elif "id" in obj:
            forms = obj.get("formulas", []) if "error" not in obj else []
            if "error" not in obj:
                cache[keys.get(obj["id"], "")] = forms
            on_result(obj["id"], forms, False, float(obj.get("sec", 0)))
    exe = os.path.join(MATH_DIR, "AcronymHunterMath.exe")
    try:
        code, _out, err = sandbox.run([exe], payload, timeout=180 + 150 * len(todo), env=env, cwd=MATH_DIR,
                                      on_line=line, cancel=cancel)
    finally:
        save_math_cache()
    if errors:
        raise RuntimeError("Math reader: " + errors[0])
    if code != 0:
        raise RuntimeError("Math reader stopped: " + err.decode("utf-8", "replace").strip()[-300:])

def _ocr_raw(image_bytes):
    clean_img = enhance_image_for_ocr(image_bytes)
    if clean_img is None:
        return ""
    lines = _ocr_lines(clean_img)          # whole-image pass (large text, layout)
    h, w = clean_img.shape[:2]
    tile, over = int(OCR_CFG["tile_size_px"]), int(OCR_CFG["tile_overlap_px"])
    if OCR_BACKEND == "rapidocr" and OCR_CFG["tile_large_images"] and max(h, w) > OCR_ENGINE_MAX_SIDE and 0 <= over < tile:
        # The engine shrank this image; re-read it in full-resolution tiles so small text survives.
        tiled = []
        for y in _tile_starts(h, tile, tile - over):
            for x in _tile_starts(w, tile, tile - over):
                tiled += _ocr_lines(clean_img[y:y + tile, x:x + tile], x, y)
        kept = list(lines)
        for item in sorted(tiled, key=lambda t: -len(t[4])):   # longest reading of a line wins
            cx, cy, txt = (item[0] + item[2]) / 2, (item[1] + item[3]) / 2, item[4]
            duplicate = any(k[0] - 10 <= cx <= k[2] + 10 and k[1] - 10 <= cy <= k[3] + 10
                            and (txt in k[4] or k[4] in txt) for k in kept)
            if not duplicate:
                kept.append(item)
        lines = kept
    lines.sort(key=lambda t: (round(t[1] / 20), t[0]))
    text = "\n".join(t[4] for t in lines)
    extra = _doctr_extra(image_bytes, lines, clean_img.shape)
    return text + ("\n" + extra if extra else "")

_OCR_CALLS = [0]   # pictures sent to OCR while scanning the current file (a file with 0 does not depend on the reader)

_typed_only = threading.local()   # set by typed_chunks(): read typed text only, skip every picture

def typed_chunks(filepath):
    """Typed text of a file without reading any picture (fast; for the preview pane). [(location label, text)]"""
    _typed_only.on = True
    try:
        return extract_chunks_from_file(filepath, lambda _m: None, 1, 1)
    finally:
        _typed_only.on = False

def read_image_bytes(image_bytes):
    if getattr(_typed_only, "on", False):
        return ""
    _OCR_CALLS[0] += 1                 # the file has pictures: its saved results depend on the picture reader
    if not ai_ml_on():
        note_skip("AI and ML off", "AI and ML are off for this effort, so text inside pictures and scanned pages was not read.")
        return ""
    key = hashlib.sha1(image_bytes).hexdigest()
    if key not in _OCR_CACHE:
        try:
            _OCR_CACHE[key] = _ocr_raw(image_bytes)
        except Exception as e:
            log_error("OCR failed on one image")
            note_skip("OCR engine", f"OCR failed, so text inside images was not read: {type(e).__name__}: {e}. Details are in the logs folder.")
            _OCR_CACHE[key] = ""
    return _OCR_CACHE[key]

# Hardened XML parsing for untrusted Office files: no entity expansion (XXE / "billion laughs"),
# no DTD loading, no network access, no unbounded trees.
_SAFE_XML = _ET.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False)
MAX_EMBEDDED_IMAGE_BYTES = 100 * 1024 * 1024   # skip absurdly large embedded images (decompression-bomb guard)

def _xml(z, path):
    return _ET.fromstring(z.read(path), parser=_SAFE_XML)

def _in_fallback(el):
    return any(anc.tag == _q("mc", "Fallback") for anc in el.iterancestors())

def zip_rels(z, partpath):
    """rId -> (relationship kind, resolved target path, is_external) for one package part."""
    folder, fname = posixpath.split(partpath)
    relpath = posixpath.join(folder, "_rels", fname + ".rels")
    out = {}
    try:
        root = _xml(z, relpath)
    except KeyError:
        return out
    for rel in root.iter(_q("pr", "Relationship")):
        target, external = rel.get("Target", ""), rel.get("TargetMode") == "External"
        if not external:
            target = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(folder, target))
        out[rel.get("Id")] = (rel.get("Type", "").rsplit("/", 1)[-1], target, external)
    return out

def drawingml_lines(root):
    """Text of every DrawingML paragraph: PowerPoint shapes, tables, groups, SmartArt, charts, Excel shapes."""
    lines = []
    for para in root.iter(_q("a", "p")):
        if _in_fallback(para):
            continue
        text = "".join(t.text or "" for t in para.iter(_q("a", "t")))
        if text.strip():
            lines.append(text)
    return lines

def chart_lines(root):
    lines = drawingml_lines(root)  # titles, axis titles
    for cache in root.iter(_q("c", "strCache")):  # category labels, series names
        lines += [v.text for v in cache.iter(_q("c", "v")) if v.text and v.text.strip()]
    return lines

def word_paragraphs(root):
    """(text, is_heading) per Word paragraph; text-box paragraphs counted once, deleted text ignored."""
    W_P, W_T = _q("w", "p"), _q("w", "t")
    for para in root.iter(W_P):
        if _in_fallback(para):
            continue
        text = "".join(t.text or "" for t in para.iter(W_T) if next(t.iterancestors(W_P), None) is para)
        style = para.find("%s/%s" % (_q("w", "pPr"), _q("w", "pStyle")))
        sval = (style.get(_q("w", "val")) or "") if style is not None else ""
        yield text, sval.lower().startswith(("heading", "title"))

def shape_name(el):
    """Name the user sees in the Selection Pane / Alt Text for the shape holding this picture."""
    for anc in el.iterancestors():
        qn = _ET.QName(anc)
        if qn.localname in ("inline", "anchor") and qn.namespace == NS["wp"]:
            doc_pr = anc.find(_q("wp", "docPr"))
            if doc_pr is not None:
                return doc_pr.get("name") or ""
        if qn.localname == "bg":
            return "Slide background"
        if qn.namespace in (NS["p"], NS["xdr"]) and qn.localname in ("pic", "sp", "graphicFrame", "cxnSp", "grpSp"):
            c_nv = anc.find(".//{*}cNvPr")
            if c_nv is not None:
                return c_nv.get("name") or ""
    return ""

# ---------------------------------------------------------------------------
# Vector graphics: read real text first; draw and OCR only when there is none
# ---------------------------------------------------------------------------
import gzip
import struct

def _maybe_gunzip(data):
    if data[:2] == b"\x1f\x8b":   # .svgz / .emz / .wmz
        with gzip.GzipFile(fileobj=io.BytesIO(data)) as g:
            return g.read(200 * 1024 * 1024)
    return data

def svg_text(data):
    """Text from <text>/<tspan>/<textPath>, plus HTML labels in <foreignObject> (draw.io) when no <text> fallback exists."""
    root = _ET.fromstring(_maybe_gunzip(data), parser=_SAFE_XML)
    lines = []
    for el in root.iter():
        if not isinstance(el.tag, str):
            continue
        tag = _ET.QName(el).localname
        if tag == "text":
            t = " ".join(x.strip() for x in el.itertext() if x.strip())
            if t:
                lines.append(t)
        elif tag == "foreignObject":
            parent = el.getparent()
            has_fallback = parent is not None and any(isinstance(c.tag, str) and _ET.QName(c).localname == "text" for c in parent)
            if not has_fallback:
                t = " ".join(x.strip() for x in el.itertext() if x.strip())
                if t:
                    lines.append(t)
    return "\n".join(lines)

def _assemble_runs(runs):
    """Join text runs (x, y, text) into lines; single characters placed one by one are joined without spaces."""
    runs = sorted(runs, key=lambda r: (round(r[1] / 4), r[0]))
    lines, cur, cur_y, prev = [], [], None, ""
    for x, y, t in runs:
        if cur_y is None or abs(y - cur_y) > 4:
            if cur: lines.append("".join(cur))
            cur, cur_y, prev = [t], y, t
            continue
        cur.append(t if (len(t) == 1 and len(prev) == 1) else " " + t)
        prev = t
    if cur: lines.append("".join(cur))
    return "\n".join(l.strip() for l in lines if l.strip())

def emf_text(data):
    """Text from EMF records: EMR_EXTTEXTOUTW/A (real characters only) and EMF+ DrawString in EMR_COMMENT."""
    runs_gdi, runs_plus, pos, n = [], [], 0, len(data)
    while pos + 8 <= n:
        rtype, size = struct.unpack_from("<II", data, pos)
        if size < 8 or pos + size > n:
            break
        try:
            if rtype in (83, 84) and size >= 76:            # EMR_EXTTEXTOUTA / EMR_EXTTEXTOUTW
                x, y, nchars, off, opts = struct.unpack_from("<iiIII", data, pos + 36)
                if not (opts & 0x10) and nchars and off + nchars * (2 if rtype == 84 else 1) <= size:   # 0x10 = glyph indices
                    raw = data[pos + off: pos + off + nchars * (2 if rtype == 84 else 1)]
                    t = raw.decode("utf-16-le" if rtype == 84 else "cp1252", "ignore").strip("\x00")
                    if t.strip(): runs_gdi.append((x, y, t))
            elif rtype == 70 and size >= 16:                  # EMR_COMMENT, possibly EMF+
                dsize = struct.unpack_from("<I", data, pos + 8)[0]
                if data[pos + 12: pos + 16] == b"EMF+":
                    buf, p = data[pos + 16: pos + 12 + dsize], 0
                    while p + 12 <= len(buf):
                        t2, _flags, sz, _dsz = struct.unpack_from("<HHII", buf, p)
                        if sz < 12 or p + sz > len(buf): break
                        if t2 == 0x401C and sz >= 40:        # EmfPlusDrawString
                            length = struct.unpack_from("<I", buf, p + 20)[0]
                            lx, ly = struct.unpack_from("<ff", buf, p + 24)
                            t = buf[p + 40: p + 40 + length * 2].decode("utf-16-le", "ignore")
                            if t.strip(): runs_plus.append((lx, ly, t))
                        p += sz
            elif rtype == 14:                                 # EMR_EOF
                break
        except struct.error:
            pass
        pos += size
    return _assemble_runs(runs_plus or runs_gdi)              # EMF+ text is authoritative in dual-mode files

def wmf_text(data):
    """Text from WMF META_TEXTOUT / META_EXTTEXTOUT records (ANSI)."""
    pos = 22 if data[:4] == b"\xd7\xcd\xc6\x9a" else 0          # placeable header
    if pos + 18 > len(data): return ""
    pos += struct.unpack_from("<H", data, pos + 2)[0] * 2        # header size in 16-bit words
    runs = []
    while pos + 6 <= len(data):
        size_words, func = struct.unpack_from("<IH", data, pos)
        size = size_words * 2
        if size < 6 or pos + size > len(data): break
        try:
            if func == 0x0521:                                    # META_TEXTOUT
                ln = struct.unpack_from("<H", data, pos + 6)[0]
                t = data[pos + 8: pos + 8 + ln].decode("cp1252", "ignore")
                y, x = struct.unpack_from("<hh", data, pos + 8 + ln + (ln & 1))
                runs.append((x, y, t))
            elif func == 0x0A32:                                  # META_EXTTEXTOUT
                y, x, ln, opts = struct.unpack_from("<hhHH", data, pos + 6)
                start = pos + 14 + (8 if opts & 0x0006 else 0)
                runs.append((x, y, data[start: start + ln].decode("cp1252", "ignore")))
            elif func == 0x0000:
                break
        except struct.error:
            pass
        pos += size
    return _assemble_runs([r for r in runs if r[2].strip()])

def _pil_png(pil_img):
    buf = io.BytesIO()
    pil_img.save(buf, format="PNG")
    return buf.getvalue()

def render_svg(data):
    """Draw SVG bytes as PNG with Google Skia (BSD-3), from the bytes in memory. No resource loader is given, so
    images the SVG links to are never fetched; text is drawn with installed Windows fonts."""
    import skia
    dom = skia.SVGDOM.MakeFromStream(skia.MemoryStream(data))
    if dom is None:
        raise ValueError("SVG could not be parsed")
    size = dom.containerSize()
    w, h = size.width() or 800.0, size.height() or 600.0
    zoom = max(1.0, min(4.0, 2400.0 / max(w, h, 1)))
    surface = skia.Surface(max(1, int(w * zoom)), max(1, int(h * zoom)))
    with surface as canvas:
        canvas.clear(skia.ColorWHITE)
        canvas.scale(zoom, zoom)
        dom.setContainerSize(skia.Size(w, h))
        dom.render(canvas)
    return surface.makeImageSnapshot().encodeToData().bytes()

def render_pdf_page(path, index, dpi):
    """One PDF page as a PIL image (PDFium)."""
    doc = pdfium.PdfDocument(path)
    try:
        page = doc[index]
        try:
            return page.render(scale=dpi / 72.0).to_pil().convert("RGB")
        finally:
            page.close()
    finally:
        doc.close()

def pdf_page_count(path):
    doc = pdfium.PdfDocument(path)
    try:
        return len(doc)
    finally:
        doc.close()

def render_vector(data, ext):
    """Draw a vector graphic as PNG bytes for OCR: SVG via Google Skia, EMF/WMF via Windows GDI (Pillow)."""
    data = _maybe_gunzip(data)
    if ext in (".svg", ".svgz"):
        return render_svg(data)
    from PIL import Image
    im = Image.open(io.BytesIO(data))
    base_w, base_h = im.size
    dpi = max(150, min(600, int(72 * 2400 / max(base_w, base_h, 1))))   # aim for ~2400 px on the long side
    try:
        im.load(dpi=dpi)
    except TypeError:
        im.load()
    im = im.convert("RGB")
    if max(im.size) < 1600:   # GDI ignored the dpi request: scale up so small labels stay readable
        k = 2400 / max(im.size)
        im = im.resize((int(im.width * k), int(im.height * k)), Image.LANCZOS)
    out = io.BytesIO(); im.save(out, "PNG")
    return out.getvalue()

def read_vector(data, ext):
    """(text, method) for a vector graphic. method: 'vector text', 'OCR', or '' when nothing could be read."""
    ext = ext.lower()
    try:
        if ext in (".svg", ".svgz"):
            text = svg_text(data)
        else:
            raw = _maybe_gunzip(data)
            text = emf_text(raw) if raw[40:44] == b" EMF" else wmf_text(raw)
    except Exception:
        log_error("Vector text extraction failed")
        text = ""
    if len(text.strip()) >= VECTOR_TEXT_MIN_CHARS:
        return text, "vector text"
    try:
        return read_image_bytes(render_vector(data, ext)), "OCR"
    except Exception:
        log_error("Vector graphic could not be drawn for OCR")
        return "", ""

def part_images(z, partpath, root, label_base, chunks, seen, rels=None):
    """OCR every picture referenced by one part (slide, layout, Word body/header, Excel drawing)."""
    if getattr(_typed_only, "on", False):
        return
    rels = rels if rels is not None else zip_rels(z, partpath)
    refs = []
    for blip in root.iter(_q("a", "blip")):
        if not _in_fallback(blip):
            svg = next(blip.iter("{%s}svgBlip" % NS_ASVG), None)   # SVG original stored beside the PNG
            svg_rid = svg.get(_q("r", "embed")) if svg is not None else None
            if svg_rid and svg_rid in rels and not rels[svg_rid][2]:
                try:
                    stext = svg_text(z.read(rels[svg_rid][1]))
                except Exception:
                    diag.swallowed()
                    stext = ""
                if len(stext.strip()) >= VECTOR_TEXT_MIN_CHARS:
                    name = shape_name(blip); shown = f"'{name}' " if name else ""
                    target = rels[svg_rid][1]
                    if (label_base, target) not in seen:
                        seen.add((label_base, target))
                        chunks.append((f"{label_base}, SVG {shown}{posixpath.basename(target)})", stext))
                    continue   # exact SVG text beats OCR of the PNG copy
            refs.append((blip.get(_q("r", "embed")), blip.get(_q("r", "link")), shape_name(blip)))
    for img in root.iter(_q("v", "imagedata")):  # legacy VML pictures
        if not _in_fallback(img):
            refs.append((img.get(_q("r", "id")), None, img.getparent().get("alt") or ""))
    for rid_embed, rid_link, name in refs:
        shown = f"'{name}' " if name else ""
        if rid_embed and rid_embed in rels and not rels[rid_embed][2]:
            target = rels[rid_embed][1]
            if (label_base, target) in seen:
                continue
            seen.add((label_base, target))
            label = f"{label_base}, image {shown}{posixpath.basename(target)})"
            ext = posixpath.splitext(target)[1].lower()
            if ext in VECTOR_EXTS:
                text, method = read_vector(z.read(target), ext)
                if text.strip():
                    chunks.append((label if method == "OCR" else label[:-1] + ", vector text)", text))
                elif not method:
                    note_skip(label, "Vector graphic could not be read or drawn.")
                continue
            if ext not in IMAGE_EXTS:
                continue
            try:
                if z.getinfo(target).file_size > MAX_EMBEDDED_IMAGE_BYTES:
                    note_skip(label, "Embedded image is over 100 MB uncompressed; skipped as a safety limit.")
                    continue
                txt = read_image_bytes(z.read(target))
            except KeyError:
                note_skip(label, "Picture is referenced but missing from the file.")
                continue
            if txt.strip():
                chunks.append((label, txt))
        elif rid_link:
            link = rels.get(rid_link, ("", "unknown location", True))[1]
            note_skip(f"{label_base}, image {shown})".replace(" )", ")"),
                      f"Linked picture, not stored in the file (links to {link}). Insert it as an embedded picture to scan it.")

# ---------------------------------------------------------------------------
# Format readers
# ---------------------------------------------------------------------------
def _outer_maths(root):
    """Equation-editor equations (m:oMath) that are not inside another one and not in a Fallback copy."""
    M = EQ._m("oMath")
    for om in root.iter(M):
        if _in_fallback(om) or any(a.tag == M for a in om.iterancestors()):
            continue
        yield om

def _eq_chunk_text(om):
    linear = " ".join(EQ.omml_text(om).split())
    return linear, linear + "\t" + "|".join(EQ.omml_symbols(om))

def _is_equation(linear):
    return any(r in linear for r in EQ.RELATIONS) or bool(re.search(r"[∑∫∏√]|\)/\(", linear))

def omml_chunks(root, label, chunks):
    """Equation-editor equations in one part (a slide, a header): exact, read from the math markup."""
    for om in _outer_maths(root):
        linear, text = _eq_chunk_text(om)
        if linear and _is_equation(linear):
            chunks.append((label + EQ.EQ_LABEL, text))

def docx_math(body, filename, chunks):
    """Word equations with their paragraph location, and nomenclature tables whose symbol cells are equations."""
    W_P, W_T, W_TC, W_TR = _q("w", "p"), _q("w", "t"), _q("w", "tc"), _q("w", "tr")
    idx, heading, defs = 0, "", []
    for para in body.iter(W_P):
        if _in_fallback(para):
            continue
        idx += 1
        style = para.find("%s/%s" % (_q("w", "pPr"), _q("w", "pStyle")))
        if style is not None and (style.get(_q("w", "val")) or "").lower().startswith(("heading", "title")):
            heading = "".join(t.text or "" for t in para.iter(W_T)).strip() or heading
        for om in _outer_maths(para):
            if next(om.iterancestors(W_P), None) is not para:
                continue
            linear, text = _eq_chunk_text(om)
            if linear and _is_equation(linear):
                where = f"section '{heading[:60]}', " if heading else ""
                chunks.append((f"{filename} ({where}para {idx})" + EQ.EQ_LABEL, text))
    for tr in body.iter(W_TR):   # symbol in one cell (as an equation or text), meaning in another
        cells = []
        for tc in tr.iter(W_TC):
            math = " ".join(" ".join(EQ.omml_text(om).split()) for om in _outer_maths(tc)).strip()
            words = " ".join("".join(t.text or "" for t in tc.iter(W_T)).split())
            cells.append((math, words))
        if len(cells) >= 2:
            sym_cell = next((m for m, w in cells if m and not w and not _is_equation(m)), None)
            meaning = next((w for m, w in cells if w and not m and len(w.split()) >= 1 and w[:1].isalpha()), None)
            if sym_cell and meaning:
                defs.append((EQ.normalize_symbol(sym_cell), meaning))
    if defs:
        chunks.append((filename + EQ.DEF_LABEL, "\n".join(f"{a}\t{b}" for a, b in dict.fromkeys(defs))))

def extract_pptx(filepath, filename, status_callback, prefix, chunks):
    with zipfile.ZipFile(filepath) as z:
        names = set(z.namelist())
        pres = _xml(z, "ppt/presentation.xml")
        prels = zip_rels(z, "ppt/presentation.xml")
        slides = [prels[s.get(_q("r", "id"))][1] for s in pres.iter(_q("p", "sldId")) if s.get(_q("r", "id")) in prels]
        seen = set()
        for n, spath in enumerate(slides, start=1):
            status_callback(f"{prefix} - Slide {n}/{len(slides)}...")
            tag = f"Slide {n}"
            try:
                root = _xml(z, spath)
                if root.get("show") == "0":
                    tag += ", hidden"
                lines = drawingml_lines(root)
                if lines:
                    chunks.append((f"{filename} ({tag})", "\n".join(lines)))
                rels = zip_rels(z, spath)
                for kind, target, external in rels.values():
                    if external or target not in names:
                        continue
                    if kind == "notesSlide":
                        sub, extra = "notes", drawingml_lines(_xml(z, target))
                    elif kind == "diagramData":
                        sub, extra = "SmartArt", drawingml_lines(_xml(z, target))
                    elif kind == "chart":
                        sub, extra = "chart", chart_lines(_xml(z, target))
                    else:
                        continue
                    if extra:
                        chunks.append((f"{filename} ({tag} {sub})", "\n".join(extra)))
                part_images(z, spath, root, f"{filename} ({tag}", chunks, seen, rels)
                omml_chunks(root, f"{filename} ({tag})", chunks)
            except Exception as e:
                log_error("PowerPoint slide could not be fully read", ext=".pptx")
                note_skip(f"{filename} (Slide {n})", f"Slide could not be fully read: {type(e).__name__}: {e}")
        # Logos and graphics placed on layouts/masters appear on many slides; scan each once.
        for lpath in sorted(names):
            m = re.match(r"ppt/slide(Layout|Master)s/slide(?:Layout|Master)(\d+)\.xml$", lpath)
            if m:
                try:
                    part_images(z, lpath, _xml(z, lpath), f"{filename} (Slide {m.group(1).lower()} {m.group(2)}", chunks, seen)
                except Exception:
                    log_error("PowerPoint layout picture could not be read", ext=".pptx")

def extract_docx(filepath, filename, status_callback, prefix, chunks):
    with zipfile.ZipFile(filepath) as z:
        names = set(z.namelist())
        status_callback(f"{prefix} - Reading Word body...")
        body = _xml(z, "word/document.xml")
        block, start, heading, idx = [], 1, "", 0

        def flush(end):
            if block:
                where = f"section '{heading[:60]}', " if heading else ""
                chunks.append((f"{filename} ({where}para {start}-{end})", "\n".join(block)))

        for text, is_heading in word_paragraphs(body):
            idx += 1
            if is_heading and text.strip():
                flush(idx - 1)
                block, start, heading = [], idx, text.strip()
            if text.strip():
                block.append(text)
            if len(block) >= 25:
                flush(idx)
                block, start = [], idx + 1
        flush(idx)
        seen = set()
        part_images(z, "word/document.xml", body, f"{filename} (Body", chunks, seen)
        docx_math(body, filename, chunks)

        status_callback(f"{prefix} - Reading headers, footers, notes, comments...")
        extras = [(p, ("Header " if "header" in p else "Footer ") + re.sub(r"\D", "", p.rsplit("/", 1)[-1]))
                  for p in sorted(names) if re.match(r"word/(header|footer)\d*\.xml$", p)]
        extras += [(p, lbl) for p, lbl in (("word/footnotes.xml", "Footnotes"), ("word/endnotes.xml", "Endnotes"),
                                           ("word/comments.xml", "Comments")) if p in names]
        for path, label in extras:
            try:
                root = _xml(z, path)
                lines = [t for t, _h in word_paragraphs(root) if t.strip()]
                if lines:
                    chunks.append((f"{filename} ({label.strip()})", "\n".join(lines)))
                part_images(z, path, root, f"{filename} ({label.strip()}", chunks, seen)
            except Exception as e:
                log_error("Word part could not be read", ext=".docx")
                note_skip(f"{filename} ({label})", f"Could not be read: {type(e).__name__}: {e}")

def extract_xlsx(filepath, filename, status_callback, prefix, chunks):
    status_callback(f"{prefix} - Reading Excel cells...")
    wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    try:
        for sname in wb.sheetnames:
            cells = []
            for row in wb[sname].iter_rows(values_only=True):
                for val in row:
                    if val is not None and not isinstance(val, (int, float)):
                        s = str(val).strip()
                        if s:
                            cells.append(s)
            if cells:
                chunks.append((f"{filename} ({sname})", "\n".join(cells)))
    finally:
        wb.close()

    status_callback(f"{prefix} - Reading shapes, charts, images, comments...")
    with zipfile.ZipFile(filepath) as z:
        names = set(z.namelist())
        wbx, wrels = _xml(z, "xl/workbook.xml"), zip_rels(z, "xl/workbook.xml")
        sheets = [(wrels[s.get(_q("r", "id"))][1], s.get("name")) for s in wbx.iter(_q("ss", "sheet")) if s.get(_q("r", "id")) in wrels]
        seen = set()
        for spath, sname in sheets:
            for kind, target, external in zip_rels(z, spath).values():
                if external or target not in names:
                    continue
                try:
                    if kind == "drawing":
                        droot = _xml(z, target)
                        lines = drawingml_lines(droot)
                        if lines:
                            chunks.append((f"{filename} ({sname} shapes)", "\n".join(lines)))
                        drels = zip_rels(z, target)
                        for dkind, dtarget, dext in drels.values():
                            if dkind == "chart" and not dext and dtarget in names:
                                clines = chart_lines(_xml(z, dtarget))
                                if clines:
                                    chunks.append((f"{filename} ({sname} chart)", "\n".join(clines)))
                        part_images(z, target, droot, f"{filename} ({sname}", chunks, seen, drels)
                    elif kind == "comments":
                        croot = _xml(z, target)
                        lines = ["".join(t.text or "" for t in c.iter(_q("ss", "t"))) for c in croot.iter(_q("ss", "comment"))]
                        lines = [l for l in lines if l.strip()]
                        if lines:
                            chunks.append((f"{filename} ({sname} comments)", "\n".join(lines)))
                except Exception as e:
                    log_error("Excel part could not be read", ext=".xlsx")
                    note_skip(f"{filename} ({sname})", f"{kind} could not be read: {type(e).__name__}: {e}")

def _pdf_page_text(page):
    tp = page.get_textpage()
    try:
        return tp.get_text_range().replace("\r\n", "\n").replace("\r", "\n")
    finally:
        tp.close()

def extract_pdf(filepath, filename, status_callback, prefix, chunks):
    """PDF text and pictures with Google PDFium (pypdfium2). The file is only read."""
    doc = pdfium.PdfDocument(filepath)
    try:
        total_pages = len(doc)
        seen_images = set()
        for i in range(total_pages):
            check_cancel()
            pnum = i + 1
            if pnum == 1 or pnum % 5 == 0 or pnum == total_pages:
                status_callback(f"{prefix} - PDF page {pnum}/{total_pages}...")
            try:
                page = doc[i]
                try:
                    digital_text = _pdf_page_text(page)
                    parts = [digital_text]
                    if len(digital_text.strip()) < OCR_CFG["scanned_page_max_chars"]:  # scanned page: OCR it whole
                        pil = page.render(scale=int(OCR_CFG["pdf_scan_dpi"]) / 72.0).to_pil().convert("RGB")
                        parts.append(read_image_bytes(_pil_png(pil)))
                    else:  # digital page: OCR each embedded figure once per document
                        for obj in page.get_objects(filter=[_pdfium_raw.FPDF_PAGEOBJ_IMAGE], max_depth=5):
                            try:
                                pil = obj.get_bitmap(render=False).to_pil()
                                key = hashlib.sha1(pil.tobytes()).hexdigest()
                                if key in seen_images:
                                    continue
                                seen_images.add(key)
                                if pil.mode not in ("RGB", "L"):
                                    pil = pil.convert("RGB")
                                parts.append(read_image_bytes(_pil_png(pil)))
                            except ScanCancelled:
                                raise
                            except Exception:
                                diag.swallowed()
                                pass
                    chunks.append((f"{filename} (p.{pnum})", "\n".join(parts)))
                finally:
                    page.close()
            except ScanCancelled:
                raise
            except Exception as e:
                log_error("PDF page could not be read", ext=".pdf")
                note_skip(f"{filename} (p.{pnum})", f"Page could not be read: {type(e).__name__}: {e}")
    finally:
        doc.close()

# ==============================================================================
# 5. LAYER 1: ALL-FORMAT UNPACKER
# ==============================================================================
def extract_chunks_from_file(filepath, status_callback, file_idx, total_files):
    filename = os.path.basename(filepath)
    ext = os.path.splitext(filepath)[1].lower()
    chunks = []
    prefix = f"[{file_idx}/{total_files}] {filename}"
    if ext in VECTOR_FILE_EXTS:
        status_callback(f"{prefix} - Reading vector graphic...")
        with open(filepath, "rb") as f:
            text, method = read_vector(f.read(), ext)
        if text.strip():
            chunks.append((f"{filename} ({'Vector text' if method == 'vector text' else 'Image'})", text))
        else:
            note_skip(f"{filename} (Vector)", "Vector graphic has no readable text and could not be drawn for OCR.")
    elif ext in IMAGE_EXTS:
        status_callback(f"{prefix} - Running OCR on image...")
        with open(filepath, 'rb') as f:
            chunks.append((f"{filename} (Image)", read_image_bytes(f.read())))
    elif ext == '.pdf':
        extract_pdf(filepath, filename, status_callback, prefix, chunks)
    elif ext == '.pptx':
        extract_pptx(filepath, filename, status_callback, prefix, chunks)
    elif ext == '.docx':
        extract_docx(filepath, filename, status_callback, prefix, chunks)
    elif ext == '.xlsx':
        extract_xlsx(filepath, filename, status_callback, prefix, chunks)
    return chunks

# ==============================================================================
# 6. LAYER 4 & 5: 3-GATE ENGINE (EXCLUSIONS, DB DISAMBIGUATION, & DISCOVERY)
# ==============================================================================
NON_UNIT_WORDS = {
    "jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec",
    "days", "weeks", "months", "years", "hours", "minutes", "seconds", "items", "pages",
    "slides", "steps", "percent", "people", "users", "files", "lines", "rows", "times",
    "and", "or", "to", "of", "in", "on", "at", "for", "with", "by", "from", "is", "was"
}

OCR_CONFUSIONS = "0158|$"
OCR_MAP = str.maketrans({"0": "O", "1": "I", "5": "S", "8": "B", "|": "I", "$": "S"})

def record_file_error(findings, filepath, exc):
    """exc may be an exception or a message string."""
    key = f"[Error] {os.path.basename(filepath)}"
    findings[key] = {
        "category": "File Error", "count": 1, "locations": [filepath],
        "definitions": [], "in_db": True, "excluded": False, "file_error": True,
        "selected_definition": (exc if isinstance(exc, str) else f"{type(exc).__name__}: {exc}")[:300], "status": "Could not read file"
    }

def record_skip(findings, label, reason):
    key = f"[Not scanned] {label}"
    if key not in findings:
        findings[key] = {
            "category": "Not Scanned", "count": 1, "locations": [label],
            "definitions": [], "in_db": True, "excluded": False, "file_error": True,
            "selected_definition": reason, "status": "Not scanned"
        }

# ==============================================================================
# 5c. CLIENT ACRONYM LISTS (from RFPs): optional; when loaded, their definitions take precedence
# ==============================================================================
CLIENT_DB = {}   # fold(term) -> {"term": as written, "definitions": [...], "sources": [list names]}
CLIENT_LIST_EXTS = {".csv", ".tsv", ".txt", ".xlsx", ".docx", ".pdf"}
CLIENT_TAG, DB_TAG = "[Client] ", "[Database] "
_TERM_HDR = re.compile(r"^(acronyms?|abbreviations?|abbrev\.?|terms?|items?|symbols?|acronym\s*/\s*abbreviation)$", re.I)
_DEF_HDR = re.compile(r"^(definitions?|meanings?|descriptions?|expansions?|full\s+(form|name|term)|spelled\s+out|stands\s+for)$", re.I)
_LINE_SEPS = ("\t", " – ", " — ", " - ", " = ", ": ")

def set_client_lists(lists):
    """Rebuild CLIENT_DB from [{"name": ..., "rows": [[term, definition], ...]}]. Empty list = no client lists."""
    CLIENT_DB.clear()
    for lst in lists or []:
        for term, defn in lst.get("rows", []):
            key = fold(term).strip()
            if not key or not defn:
                continue
            e = CLIENT_DB.setdefault(key, {"term": key, "definitions": [], "sources": []})
            if defn.casefold() not in {x.casefold() for x in e["definitions"]}:
                e["definitions"].append(defn)
            if lst.get("name") not in e["sources"]:
                e["sources"].append(lst.get("name"))
    return CLIENT_DB

def _clean_pair(term, defn):
    t = re.sub(r"\s+", " ", fold(str(term or ""))).strip().strip(":–—-").strip()
    d = re.sub(r"\s+", " ", fold(str(defn or ""))).strip().strip(":–—-").strip()
    if not (1 <= len(t) <= 25) or not re.search(r"[^\W\d_]", t) or len(d) < 2 or t.casefold() == d.casefold():
        return None
    if _TERM_HDR.match(t) or _DEF_HDR.match(d) or len(t.split()) > 4:
        return None
    return t, d

def _pairs_from_table(rows):
    """Rows of cell strings -> (term, definition). Finds header columns by name, else uses the first two filled columns."""
    rows = [[(c or "").strip() if isinstance(c, str) else ("" if c is None else str(c)) for c in r] for r in rows]
    ti = di = None
    for r in rows[:3]:
        for i, c in enumerate(r):
            if ti is None and _TERM_HDR.match(c): ti = i
            elif di is None and _DEF_HDR.match(c): di = i
        if ti is not None and di is not None:
            break
    out = []
    for r in rows:
        if ti is not None and di is not None and max(ti, di) < len(r):
            pair = _clean_pair(r[ti], r[di])
        else:
            filled = [c for c in r if c]
            pair = _clean_pair(filled[0], " ".join(filled[1:])) if len(filled) >= 2 else None
        if pair:
            out.append(pair)
    return out

def _pair_from_line(line, strict=False):
    line = line.strip()
    for sep in _LINE_SEPS:
        if sep in line:
            t, d = line.split(sep, 1)
            break
    else:
        m = re.match(r"^(\S{1,25})\s{2,}(.+)$", line) or (None if strict else re.match(r"^([A-Z][A-Za-z0-9&/.\-]{1,15})\s+(.{3,})$", line))
        if not m:
            return None
        t, d = m.group(1), m.group(2)
    if strict and not re.search(r"[A-Z]", t):
        return None
    return _clean_pair(t, d)

def _pdf_table_rows(page):
    """Rows of cells from a PDF page. PDFium reports each run of text with its box; runs on the same line that sit
    far apart are separate cells. Recovers two-column tables (TERM | definition) without a table library."""
    tp = page.get_textpage()
    try:
        segs = []
        for k in range(tp.count_rects()):
            left, bottom, right, top = tp.get_rect(k)
            h = top - bottom
            txt = tp.get_text_bounded(left + 0.5, bottom + h * 0.25, right - 0.5, top - h * 0.25).strip()
            if txt:
                segs.append((left, bottom, right, top, " ".join(txt.split())))
    finally:
        tp.close()
    lines = []
    for seg in sorted(segs, key=lambda x: (-(x[1] + x[3]) / 2, x[0])):
        mid = (seg[1] + seg[3]) / 2
        for line in lines:
            if line[0][1] <= mid <= line[0][3]:
                line.append(seg); break
        else:
            lines.append([seg])
    rows = []
    for line in lines:
        line.sort(key=lambda x: x[0])
        height = max(1.0, max(x[3] - x[1] for x in line))
        cells, cur, prev_right = [], "", None
        for left, _b, right, _t, txt in line:
            if prev_right is not None and left - prev_right > 1.2 * height:
                cells.append(cur); cur = txt
            else:
                cur = (cur + " " + txt).strip()
            prev_right = right
        cells.append(cur)
        rows.append(cells)
    return rows

def parse_client_list(path):
    """(term, definition) pairs from a client acronym list: CSV/TSV/TXT, Excel, Word (tables or 'TERM - definition'
    lines), or PDF (tables, else lines). Files are only read; parsing uses the same hardened readers as scanning."""
    ext = os.path.splitext(path)[1].lower()
    pairs = []
    if ext in (".csv", ".tsv"):
        text, delim = _decode_csv(path)
        pairs = _pairs_from_table(list(csv.reader(io.StringIO(text, newline=""), delimiter="\t" if ext == ".tsv" else delim)))
    elif ext == ".txt":
        text, _ = _decode_csv(path)
        lines = [l for l in text.splitlines() if l.strip()]
        if sum("\t" in l for l in lines) > len(lines) / 2:
            pairs = _pairs_from_table([l.split("\t") for l in lines])
        else:
            pairs = [p for p in (_pair_from_line(l) for l in lines) if p]
    elif ext == ".xlsx":
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            for ws in wb.worksheets:
                pairs += _pairs_from_table([list(r) for r in ws.iter_rows(values_only=True)])
        finally:
            wb.close()
    elif ext == ".docx":
        W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with zipfile.ZipFile(path) as z:
            root = _xml(z, "word/document.xml")
        text_of = lambda el: "".join(t.text or "" for t in el.iter(W + "t"))
        for tbl in root.iter(W + "tbl"):
            pairs += _pairs_from_table([[text_of(tc) for tc in tr.iter(W + "tc")] for tr in tbl.iter(W + "tr")])
        body = root.find(W + "body")
        for para in (body if body is not None else []):
            if para.tag == W + "p":
                p = _pair_from_line(text_of(para), strict=True)
                if p:
                    pairs.append(p)
    elif ext == ".pdf":
        doc = pdfium.PdfDocument(path)
        try:
            rows, lines = [], []
            for i in range(len(doc)):
                page = doc[i]
                try:
                    rows += _pdf_table_rows(page)
                    lines += _pdf_page_text(page).splitlines()
                finally:
                    page.close()
            pairs = _pairs_from_table([r for r in rows if len(r) >= 2])
            if not pairs:
                pairs = [p for p in (_pair_from_line(l, strict=True) for l in lines) if p]
        finally:
            doc.close()
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    seen, out = set(), []
    for t, d in pairs:
        k = (t.casefold(), d.casefold())
        if k not in seen:
            seen.add(k); out.append([t, d])
    return out

def apply_client_overlay(findings, m):
    """Client definitions take precedence. Each matching row shows the client meaning(s) first, labeled, with any
    different database meanings after them, so the user can still pick."""
    if not CLIENT_DB:
        return findings
    db_include = m.get("db_include", {})
    for term, d in findings.items():
        c = CLIENT_DB.get(fold(term).strip())
        if not c or d.get("file_error"):
            continue
        cdefs = c["definitions"]
        cset = {x.casefold() for x in cdefs}
        dbdefs = [x for x in db_include.get(term, {}).get("definitions", []) if x.casefold() not in cset]
        d.update({"client": True, "client_definitions": cdefs, "db_definitions": dbdefs, "client_sources": c["sources"],
                  "excluded": False, "in_db": True,
                  "definitions": [CLIENT_TAG + x for x in cdefs] + [DB_TAG + x for x in dbdefs]})
        if d["category"].endswith("(Candidate)"):
            d["category"] = d["category"].replace(" (Candidate)", "")
        if len(cdefs) == 1:
            d["selected_definition"] = cdefs[0]
            d["status"] = "Client list (database differs)" if dbdefs else "Client list"
        else:
            d["selected_definition"] = f"MULTIPLE ({len(cdefs)} client meanings) - Choose: " + " OR ".join(cdefs)
            d["status"] = f"CHOOSE (client list: {len(cdefs)} meanings)"
    return findings

def build_matcher():
    """Load the database, client lists, corrections, ignore list, and patterns once per scan."""
    refresh_ocr_ignore()
    include_db, exclude_db, exclude_upper = load_database()
    db_include = {k: dict(v, definitions=list(v["definitions"])) for k, v in include_db.items()}
    for key, c in CLIENT_DB.items():   # client terms are always looked for, and win over database Exclude rules
        exclude_db.pop(key, None); exclude_upper.discard(key.upper())
        if key not in include_db:
            include_db[key] = {"category": "Acronym", "definitions": list(c["definitions"])}
    all_terms = set(include_db.keys()) | set(exclude_db.keys())
    direct_pattern = None
    if all_terms:
        sorted_terms = sorted(all_terms, key=len, reverse=True)
        # (?<![^\W\d_]) = not preceded by a letter; digits may touch a unit (25°C, 10µm, 150psi)
        direct_pattern = re.compile(r'(?<![^\W\d_])(' + '|'.join([re.escape(t) for t in sorted_terms]) + r")(?:'?s)?(?!\w)")
    return {
        "include_db": include_db, "exclude_db": exclude_db, "exclude_upper": exclude_upper, "db_include": db_include,
        "ocr_fixes": load_ocr_fixes(), "direct_pattern": direct_pattern,
        "fuzzy_targets": {t: t for t in include_db.keys() if len(t) >= 3 and t.isupper()},
        "acronym_re": re.compile(r'\b([A-Z][A-Z0-9/&]{1,7}s?)\b'),
        # Codes and designators kept whole: PCU-28V, MIL-STD-1553B, RS-422, PN 4521-003A, ST-200X, J1-P3, SN77X2
        "designator_re": re.compile(r'(?<![\w/-])((?:[A-Z0-9]+-)+[A-Z0-9]+|[A-Z]{1,4}[0-9][A-Z0-9]{2,8})(?![\w/-])'),
        # Mixed-case acronyms: DoD, DoE, IoT, MoU, PhD, DoDAF
        "mixed_re": re.compile(r'\b([A-Z][a-z]?[A-Z][A-Za-z]{0,2}[A-Z]|[A-Z][a-z][A-Z])\b'),
        "unit_re": re.compile(r'(?<!\w)\d+(?:\.\d+)?\s*([a-zA-Z\u00b0\u03bc\u03a9\u00c5]{1,5}(?:/[a-zA-Z0-9\^]{1,4})?)(?!\w)'),
        "common_words": common_words(),
        "symbol_defs": {},
        "abbrev_re": re.compile(r'(?<!\w)([a-zA-Z]{1,4}/[a-zA-Z]{1,4}|[a-zA-Z]{2,6}\.)(?!\w)'),
    }

_HITS = None   # terms recorded in the chunk being matched (for first use)

# Real abbreviations written with a period that are also (or look like) ordinary words: always reported.
PERIOD_ABBREVIATIONS = set("""
approx appx assy avg ave cal cap ch chap co col config corp ctrl dept dia doc docs dr dwg ea eng engr eq eqn est etc
fig figs ft ga gal govt hr hrs ht inc incl intl jr lb lbs ltd max mfg mgmt mgr min misc mo mr mrs ms mt natl no nos
oz para pg pgs pkg pp pt pts qty qual ref refs reqt rev rm rqmt sec sect spec specs sq sr st std sta tech tel temp
tot vol vols vs wk wt yr yrs
""".split())
_COMMON_WORDS = None

def common_words():
    """Common 2-6 letter English words (SCOWL, assets\\common_words.txt). Loaded once."""
    global _COMMON_WORDS
    if _COMMON_WORDS is None:
        words = set()
        try:
            with open(asset_path("common_words.txt"), encoding="utf-8") as f:
                words = {w.strip() for w in f if w.strip() and not w.startswith("#")}
        except (OSError, TypeError):
            log_error("Common words list missing; sentence-end words may be reported as abbreviations")
        _COMMON_WORDS = frozenset(words)
    return _COMMON_WORDS

def _file_of(label):
    """File name at the start of a location label ("Volume (Final).docx (Slide 3)" -> "Volume (Final).docx")."""
    exts = "|".join(re.escape(e) for e in sorted(SUPPORTED_EXTS, key=len, reverse=True))
    mt = re.match(rf"(.+?(?:{exts}))(?= \(| \[|$)", label, re.I)
    return mt.group(1) if mt else label

def note_first_use(findings, hits, label, text):
    """Per file, where each term first appears in reading order, and whether that first use is written the way
    proposals define an acronym: "Spelled Out Words (ACR)". Records a later "(ACR)" when the first use is bare."""
    f = _file_of(label)
    for term in dict.fromkeys(hits):
        d = findings.get(term)
        if d is None:
            continue
        pat = re.escape(term) + r"s?"
        mt = re.search(r"(?<![A-Za-z0-9])" + pat + r"(?![A-Za-z0-9])", text)
        paren = None if mt is None else text[max(0, mt.start() - 1):mt.start()] == "("
        defs = list(re.finditer(r"\(\s*" + pat + r"\s*\)", text))
        fu = d.setdefault("first_use", {})
        if f not in fu:
            fu[f] = {"loc": label, "defined": paren}
            if paren:                                   # the words written before "(ACR)" (2.7.0 volume check)
                fu[f]["lf"] = _long_form_before(term, text, mt.start() - 1)
                if len(defs) > 1:
                    fu[f]["redefs"] = [label]
            if paren is False and defs:
                fu[f]["later"] = label
        else:
            if fu[f]["defined"] is False and not fu[f].get("later") and defs:
                fu[f]["later"] = label
            elif fu[f]["defined"] and defs and label not in fu[f].get("redefs", []) and len(fu[f].get("redefs", [])) < 5:
                fu[f].setdefault("redefs", []).append(label)

def _long_form_before(term, text, paren_at):
    """Schwartz-Hearst long form ending right before the parenthesis, or "" when the words do not fit the acronym."""
    import custdefs
    before = re.split(r"[.;:!?\n(]\s", text[max(0, paren_at - 200):paren_at])[-1]
    words = before.split()[-(min(len(term) + 5, len(term) * 2) + 2):]
    try:
        return (custdefs._best_long_form(term, words) or "") if words else ""
    except Exception:
        diag.swallowed()
        return ""

def first_use_lines(d, with_file=True):
    """Readable first-use lines for acronyms only (units and abbreviations are never defined this way)."""
    if not d.get("category", "").startswith("Acronym") or d.get("excluded") or d.get("known"):
        return []
    out = []
    for f, u in d.get("first_use", {}).items():
        loc = u["loc"] if with_file else (u["loc"][len(f):].strip() or u["loc"])
        if u["defined"] is True:
            state = "defined here"
        elif u["defined"] is None:
            state = "check (OCR-corrected)"
        elif u.get("later") == u["loc"]:
            state = "NOT defined at first use; spelled out later in the same passage"
        elif u.get("later"):
            state = f"NOT defined at first use; spelled out later at {u['later']}"
        else:
            state = "NOT defined in this file"
        out.append(f"{loc}: {state}")
    return out

def ctx_text(d):
    """'Statement of Work (spelled out in the document; nearby: tasks)', 'differs by place: ...' or ''."""
    c = d.get("context") or {}
    if c.get("mixed"):
        means = list(dict.fromkeys(v[0] for v in c.get("by_loc", {}).values()))
        return "differs by place: " + " / ".join(means)
    if not c.get("definition"):
        return ""
    return c["definition"] + (f" ({'; '.join(c.get('cues', [])[:2])})" if c.get("cues") else "")

def first_use_flag(d):
    """True if any file's first use of an acronym is not written as "Words (ACR)"."""
    return any(u["defined"] is False for u in d.get("first_use", {}).values()) and bool(first_use_lines(d))

def record_symbol(findings, sym, loc, m):
    if sym in m['include_db']:
        record_included(findings, sym, loc, m['include_db'])
    elif sym in m['exclude_db']:
        record_excluded(findings, sym, loc, m['exclude_db']); return
    else:
        if sym not in findings:
            findings[sym] = {"category": "Symbol", "count": 0, "locations": [], "definitions": [], "in_db": False,
                             "excluded": False, "selected_definition": "[Not defined in the document or database]",
                             "status": "SYMBOL NOT DEFINED"}
        findings[sym]["count"] += 1
        if loc not in findings[sym]["locations"]:
            findings[sym]["locations"].append(loc)
    findings[sym]["symbol"] = True

def record_equation(findings, loc, text, m):
    linear, _sep, syms = text.partition("\t")
    linear = " ".join(linear.split())
    syms = [x for x in syms.split("|") if x] if _sep else EQ.symbols_in_text(linear, common=m['common_words'])
    key = linear if len(linear) <= 160 else linear[:157] + "..."
    if key not in findings:
        findings[key] = {"category": "Equation", "count": 0, "locations": [], "definitions": [], "in_db": False,
                         "excluded": False, "selected_definition": "", "status": "EQUATION", "equation": True,
                         "symbols": []}
    d = findings[key]
    d["count"] += 1
    if loc not in d["locations"]:
        d["locations"].append(loc)
    d["symbols"] = list(dict.fromkeys(d["symbols"] + syms))
    for sym in syms:
        record_symbol(findings, sym, loc, m)

def finalize_symbols(findings, m):
    """Give each symbol the meaning stated in the documents (where ṁ is ..., nomenclature lists and tables)."""
    stated = {}
    for _file, defs in m['symbol_defs'].items():
        for sym, desc in defs:
            lst = stated.setdefault(sym, [])
            if desc.casefold() not in {x.casefold() for x in lst}:
                lst.append(desc)
    for sym, d in findings.items():
        if not d.get("symbol") or d.get("in_db") or d.get("excluded"):
            continue
        descs = stated.get(sym, [])
        if len(descs) == 1:
            d.update({"definitions": descs, "selected_definition": descs[0], "status": "Defined in document"})
        elif len(descs) > 1:
            d.update({"definitions": descs, "status": f"CHOOSE (defined {len(descs)} ways)",
                      "selected_definition": f"MULTIPLE ({len(descs)} meanings) - Choose: " + " OR ".join(descs)})

def summarize_equations(findings):
    """Each equation row lists its symbols and which ones are still undefined."""
    for key, d in findings.items():
        if not d.get("equation"):
            continue
        undefined = [x for x in d["symbols"] if x in findings and status_kind(findings[x]["status"]) == "candidate"]
        d["selected_definition"] = "Symbols: " + ", ".join(d["symbols"]) + (f"  |  not defined: {', '.join(undefined)}" if undefined else "")
        d["status"] = f"EQUATION: {len(undefined)} symbol(s) not defined" if undefined else "EQUATION (all symbols defined)"

def match_chunks(findings, chunks, m):
    """Run the three gates over [(location label, text)] and record findings."""
    global _HITS
    for loc_label, text in chunks:
        if not text.strip():
            continue
        text = fold(text)
        if loc_label.endswith(EQ.EQ_LABEL):
            record_equation(findings, loc_label[:-len(EQ.EQ_LABEL)], text, m); continue
        if loc_label.endswith(EQ.DEF_LABEL):
            name = loc_label[:-len(EQ.DEF_LABEL)]
            m['symbol_defs'].setdefault(name, [])
            m['symbol_defs'][name] += [tuple(l.split("\t", 1)) for l in text.splitlines() if "\t" in l]
            continue
        if loc_label.endswith(EQ.PIC_LABEL):
            loc = loc_label[:-len(EQ.PIC_LABEL)]
            key = f"[equation picture] {_file_of(loc)}"
            if key not in findings:
                findings[key] = {"category": "Equation (picture)", "count": 0, "locations": [], "definitions": [],
                                 "in_db": False, "excluded": False, "status": "CHECK EQUATION PICTURE",
                                 "selected_definition": "Math inside a picture. OCR reads letters, not math layout: check it "
                                                        "in the preview. Read as: " + text[:120]}
            findings[key]["count"] += 1
            if loc not in findings[key]["locations"]:
                findings[key]["locations"].append(loc)
            continue
        _HITS = hits = []

        # Gate 1 & 2: Unified Database (Exclusions & Inclusions)
        if m['direct_pattern']:
            for match in m['direct_pattern'].finditer(text):
                term = match.group(1)
                if term in m['exclude_db'] and term not in m['include_db']:
                    record_excluded(findings, term, loc_label, m['exclude_db'])
                elif term in m['include_db']:
                    record_included(findings, term, loc_label, m['include_db'])

        # Gate 3a: codes and designators, reported whole. Their pieces are not reported as units or abbreviations.
        code_spans = []
        for match in m['designator_re'].finditer(text):
            tok = match.group(1)
            letters = sum(c.isalpha() for c in tok)
            if len(tok) < 4 or letters == 0 or (not any(c.isdigit() for c in tok) and "-" not in tok):
                continue
            if "-" in tok and not any(c.isdigit() for c in tok) and min(len(p) for p in tok.split("-")) < 2:
                continue   # "A-B" style fragments
            code_spans.append(match.span(1))
            if tok in m['include_db']:
                record_included(findings, tok, loc_label, m['include_db'])
            elif tok in m['exclude_db'] or tok.upper() in m['exclude_upper']:
                if tok in m['exclude_db']:
                    record_excluded(findings, tok, loc_label, m['exclude_db'])
            else:
                record_candidate(findings, tok, "Designator (Candidate)", loc_label)
        in_code = lambda span: any(a <= span[0] and span[1] <= b for a, b in code_spans)

        for match in m['mixed_re'].finditer(text):
            tok = match.group(1)
            if tok in m['include_db'] or tok in m['exclude_db'] or tok.upper() in m['exclude_upper'] or tok.isupper() or tok.lower() in m['common_words']:
                continue
            record_candidate(findings, tok, "Acronym (Candidate)", loc_label)

        # Gate 3: Heuristic Candidate Discovery (Items not in DB)
        for match in m['acronym_re'].finditer(text):
            raw = match.group(1)
            cand = raw[:-1] if raw.endswith('s') and len(raw) > 2 else raw
            if raw in m['ocr_fixes'] or cand in m['ocr_fixes']:
                wrong = raw if raw in m['ocr_fixes'] else cand
                apply_ocr_fix(findings, wrong, m['ocr_fixes'][wrong], "Acronym (Candidate)", loc_label, m['include_db'], m['exclude_db'])
                continue
            if cand in m['include_db'] or cand in m['exclude_db'] or cand.upper() in m['exclude_upper'] or cand.isdigit():
                continue
            if sum(1 for c in cand if c.isupper()) < 2:
                continue
            if any(ch in OCR_CONFUSIONS for ch in cand):
                fixed = cand.translate(OCR_MAP)
                if fixed in m['include_db']:
                    record_included(findings, fixed, f"{loc_label} [OCR fix: {cand}]", m['include_db'])
                    continue
            if any(ch.isdigit() for ch in cand) and m['fuzzy_targets']:
                best = process.extractOne(cand, m['fuzzy_targets'].keys(), scorer=fuzz.ratio)
                if best and best[1] >= OCR_CFG["fuzzy_min_score"]:
                    record_included(findings, best[0], f"{loc_label} [OCR fix: {cand}]", m['include_db'])
                    continue
            record_candidate(findings, cand, "Acronym (Candidate)", loc_label)

        for match in m['unit_re'].finditer(text):
            ucand = match.group(1)
            if in_code(match.span()):
                continue   # "28V" inside PCU-28V is part of a code, not a measurement
            if ucand in m['ocr_fixes']:
                apply_ocr_fix(findings, ucand, m['ocr_fixes'][ucand], "Unit (Candidate)", loc_label, m['include_db'], m['exclude_db'])
                continue
            if (ucand in m['include_db'] or ucand in m['exclude_db'] or
                    ucand.upper() in m['exclude_upper'] or ucand.lower() in NON_UNIT_WORDS or ucand.isdigit()):
                continue
            record_candidate(findings, ucand, "Unit (Candidate)", loc_label)

        for match in m['abbrev_re'].finditer(text):
            acand = match.group(1)
            if in_code(match.span(1)):
                continue
            if acand in m['ocr_fixes']:
                apply_ocr_fix(findings, acand, m['ocr_fixes'][acand], "Abbreviation (Candidate)", loc_label, m['include_db'], m['exclude_db'])
                continue
            base = acand.rstrip('.')
            if (acand in m['include_db'] or acand in m['exclude_db'] or base in m['include_db'] or base in m['exclude_db'] or
                    acand.upper() in m['exclude_upper'] or base.upper() in m['exclude_upper']):
                continue
            if acand.endswith('.') and '/' not in acand:
                if len(base) >= 2 and base.isupper():
                    continue   # "CCB." is an acronym ending a sentence: the acronym rule already reports CCB
                if base.lower() in m['common_words'] and base.lower() not in PERIOD_ABBREVIATIONS:
                    continue   # "award." is an ordinary word ending a sentence, not an abbreviation
            record_candidate(findings, acand, "Abbreviation (Candidate)", loc_label)
        _HITS = None
        note_first_use(findings, hits, loc_label, text)


CANCEL = threading.Event()   # set by the Cancel Scan button

class ScanCancelled(BaseException):
    """Raised to stop a scan. BaseException, so per-page error handlers (except Exception) do not swallow it."""

def check_cancel():
    if CANCEL.is_set():
        raise ScanCancelled()

def scan_file(filepath, status_callback, idx, total):
    """Read one file. Returns (chunks, skipped [(label, reason)], error message or None)."""
    _SKIPPED.clear()
    try:
        chunks = extract_chunks_from_file(filepath, status_callback, idx, total)
        chunks += EQ.text_math_chunks(chunks, common_words())
        err = None
    except Exception as e:
        log_error("File could not be read", ext=os.path.splitext(filepath)[1].lower())
        chunks, err = [], f"{type(e).__name__}: {e}"
    return chunks, list(_SKIPPED), err

# ---------------------------------------------------------------------------
# Full text mode: every word the scanner reads, typed and OCR, in reading order
# ---------------------------------------------------------------------------
_OCR_LOC = re.compile(r"\(Image\)|, image ")

def _para_span(loc):
    mt = re.search(r"para (\d+)(?:-(\d+))?\)", loc)
    return (int(mt.group(1)), int(mt.group(2) or mt.group(1))) if mt else None

def _reading_order(chunks):
    """Word equations sit in their paragraph's place: after the text block that holds or precedes that paragraph."""
    eqs = [c for c in chunks if c[0].endswith(EQ.EQ_LABEL) and _para_span(c[0])]
    rest = [c for c in chunks if c not in eqs]
    if not eqs:
        return chunks
    out, placed = [], set()
    for c in rest:
        out.append(c)
        span = _para_span(c[0])
        if span and not c[0].endswith(EQ.EQ_LABEL):
            for i, e in enumerate(eqs):
                if i not in placed and _para_span(e[0])[0] <= span[1]:
                    out.append(e); placed.add(i)
    out += [e for i, e in enumerate(eqs) if i not in placed]
    return out

def full_text_of(filepath, status_callback=lambda m: None, idx=1, total=1):
    """All text in one file, as the scanner reads it: typed text exactly, picture text through the sandboxed OCR
    reader, equations in linear form and in reading order, then every symbol definition found. Returns (text, error)."""
    chunks, skips, err = scan_file(filepath, status_callback, idx, total)
    name = os.path.basename(filepath)
    out = [f"===== {name} ====="]
    if err:
        out.append(f"[Could not read this file: {err}]")
    defs = {}
    for loc, text in _reading_order(chunks):
        if loc.endswith(EQ.DEF_LABEL):
            for line in text.splitlines():
                sym, _t, meaning = line.partition("\t")
                if sym and meaning:
                    defs.setdefault(sym.strip(), meaning.strip())
            continue
        if loc.endswith(EQ.PIC_LABEL):
            continue   # the picture's text is already shown where it was read
        where = loc[len(name):].strip() if loc.startswith(name) else loc
        if loc.endswith(EQ.EQ_LABEL):
            linear, _t, syms = text.partition("\t")
            out.append(f"--- {where[:-len(EQ.EQ_LABEL)].strip() or 'Equation'} [equation] ---")
            out.append(linear.strip() + (f"\n    symbols: {', '.join(x for x in syms.split('|') if x)}" if syms else ""))
            continue
        body = text.strip()
        if not body:
            continue
        tag = "  [read by OCR: check spelling]" if _OCR_LOC.search(loc) else ""
        out.append(f"--- {where or 'Text'}{tag} ---")
        out.append(body)
    if defs:
        out.append("--- Symbol definitions found in this file ---")
        out += [f"{k} = {v}" for k, v in defs.items()]
    for lbl, why in skips:
        out.append(f"[Not read: {lbl}: {why}]")
    return "\n".join(out) + "\n", err

def analyze_batch(filepaths, status_callback, capture=None):
    """Scan files once. capture (dict): filled with each file's terms-only cache entry, so the results can be saved
    as a project later without reading the files again."""
    m = build_matcher()
    findings = {}
    total = len(filepaths)
    _OCR_CACHE.clear()
    if capture is not None:
        weak = weak_terms(m); sigs = (engine_signature(weak), text_signature(weak), ocr_signature())
    for idx, filepath in enumerate(filepaths, start=1):
        check_cancel()
        _OCR_CALLS[0] = 0
        chunks, skips, err = scan_file(filepath, status_callback, idx, total)
        if capture is not None:
            try:
                capture[os.path.abspath(filepath)] = cache_entry(os.stat(filepath), weak, chunks, skips, err, _OCR_CALLS[0], sigs)
            except OSError:
                pass
        if err:
            record_file_error(findings, filepath, err)
        for label, reason in skips:
            record_skip(findings, label, reason)
        match_chunks(findings, chunks, m)
    finalize_symbols(findings, m)
    apply_client_overlay(findings, m)
    summarize_equations(findings)
    return dict(sorted(findings.items(), key=lambda x: x[1]["count"], reverse=True))

def refresh_batch(filepaths, status_callback, cache):
    """Match files again against the current database (2.14.0), reusing the text already read in this session when the
    file and everything its saved text depends on are unchanged; other files are read again. Updates cache."""
    m = build_matcher()
    findings = {}
    _OCR_CACHE.clear()
    weak = weak_terms(m); sigs = (engine_signature(weak), text_signature(weak), ocr_signature())
    total = len(filepaths)
    for idx, filepath in enumerate(filepaths, start=1):
        check_cancel()
        key = os.path.abspath(filepath)
        entry = cache.get(key)
        try:
            st = os.stat(filepath)
        except OSError:
            st = None
        fresh = entry and st and entry_current(entry, sigs) and entry["size"] == st.st_size and abs(entry["mtime"] - st.st_mtime) < 1.0
        if fresh:
            status_callback(f"[{idx}/{total}] {os.path.basename(filepath)} - unchanged, matching again")
        else:
            _OCR_CALLS[0] = 0
            chunks, skips, err = scan_file(filepath, status_callback, idx, total)
            if st is None:
                record_file_error(findings, filepath, err or "File not found")
                continue
            entry = cache_entry(st, weak, chunks, skips, err, _OCR_CALLS[0], sigs)
            cache[key] = entry
        if entry.get("error"):
            record_file_error(findings, filepath, entry["error"])
        for label, reason in entry.get("skips", []):
            record_skip(findings, label, reason)
        match_chunks(findings, entry["chunks"], m)
    finalize_symbols(findings, m)
    apply_client_overlay(findings, m)
    summarize_equations(findings)
    return dict(sorted(findings.items(), key=lambda x: x[1]["count"], reverse=True))

# ==============================================================================
# PROJECTS: remember a folder's findings, rescan only what changed, track changes
# ==============================================================================
APP_VERSION = "2.19.0"            # keep in step with version_info.txt
diag.VERSION = APP_VERSION
PROJECT_EXT = ".ahproj"
PROJECT_FORMAT = "AcronymHunterProject"
LOCK_STALE_HOURS = 12
SKIP_DIRS = {"_internal", "_local_sandbox", "training_data", "logs", "config", "__pycache__", ".git"}
CHANGE_LABELS = {"new": "New", "changed": "Changed", "unchanged": "Unchanged", "gone": "No longer found"}

_NUM_TOKEN = re.compile(r"[\d.,]+")
_EDGE_PUNCT = "()[]{}\"'\u201c\u201d\u2018\u2019,;:!?"

def _termlike(token):
    """True for tokens that can ever match: 2+ capitals, a digit, a symbol, or an abbreviation-style word ending in '.'
    Hyphenated words count if any part does (S-band. -> band.)."""
    core = token.strip(_EDGE_PUNCT)
    if not core:
        return False
    if "-" in core.strip("-"):
        return any(_termlike(part) for part in core.split("-") if part)
    if sum(c.isupper() for c in core) >= 2 or any(c.isdigit() for c in core):
        return True
    if any(not c.isalpha() and c not in ".-'" for c in core):
        return True
    return core.endswith(".") and core[:-1].isalpha() and 1 <= len(core) - 1 <= 6

def weak_terms(m):
    """Database / correction words that _termlike() would drop (e.g. lowercase units like 'psi'). They are kept
    explicitly, and changing them invalidates saved results so nothing is missed."""
    words = set()
    for term in list(m["include_db"]) + list(m["exclude_db"]) + list(m["ocr_fixes"]):
        words.update(w for w in term.split() if not _termlike(w))
    return words

def _sha12(obj):
    return hashlib.sha1(json.dumps(obj, sort_keys=True).encode()).hexdigest()[:12]

def text_signature(weak=()):
    """What typed-text results depend on: the app's extraction code (major.minor version) and the weak terms."""
    return _sha12([APP_VERSION.rsplit(".", 1)[0], sorted(weak)])

def ocr_signature():
    """What picture results depend on: the OCR engine, its settings, the installed reader, and the ignore list."""
    models = sorted((n, os.path.getsize(os.path.join(MODELS_DIR, n)), int(os.path.getmtime(os.path.join(MODELS_DIR, n))))
                    for n in os.listdir(MODELS_DIR)) if os.path.isdir(MODELS_DIR) else []
    dsig = ""
    if doctr_active():
        try:
            dsig = hashlib.sha1(open(os.path.join(DOCTR_DIR, "doctr_models.json"), "rb").read()).hexdigest()[:12]
        except Exception:
            diag.swallowed()
            dsig = "doctr"
    if not ai_ml_on():
        return _sha12(["ai_ml_off"])       # pictures were skipped: turning AI and ML back on re-reads them
    return _sha12([OCR_BACKEND, OCR_CFG, models, sorted(_OCR_IGNORE), dsig])

def engine_signature(weak=()):
    """Both signatures together (older project files store only this)."""
    return _sha12([text_signature(weak), ocr_signature()])

def cache_entry(st, weak, chunks, skips, err, ocr_used, sigs):
    sig, tsig, osig = sigs
    return {"size": st.st_size, "mtime": st.st_mtime, "engine": sig, "text_sig": tsig, "ocr_sig": osig,
            "ocr_used": bool(ocr_used), "scanned": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
            "chunks": [[label, text if label.endswith((EQ.EQ_LABEL, EQ.DEF_LABEL, EQ.PIC_LABEL)) else reduce_text(fold(text), weak)]
                       for label, text in chunks if text.strip()],
            "skips": skips, "error": err}

def entry_current(entry, sigs):
    """Saved results for a file stay valid when nothing they depend on changed. A file with no pictures is not
    re-read after Improve Reading, a new Knowledge Pack reader, or an OCR setting change."""
    sig, tsig, osig = sigs
    if "text_sig" in entry:
        return entry["text_sig"] == tsig and (not entry.get("ocr_used", True) or entry.get("ocr_sig") == osig)
    return entry.get("engine") == sig

def _numberish(token):
    t = token.strip(_EDGE_PUNCT)
    return bool(t) and t[-1].isdigit() and any(c.isdigit() for c in t)

def reduce_text(text, weak=()):
    """Terms-only text for the project file: keeps term-like tokens, the word after a number (units), and the
    database's own lowercase words. Each run of dropped words becomes '~' so neighbors never become adjacent
    (matching results stay identical to a full-text scan). Sentences cannot be reconstructed."""
    out, prev_num = [], False
    for line in text.splitlines():
        kept, gap = [], False
        for t in line.split():
            if _termlike(t) or prev_num or t.strip(_EDGE_PUNCT) in weak:
                if gap:
                    kept.append("~")
                kept.append(t); gap = False
            else:
                gap = True
            prev_num = _numberish(t)
        if gap and kept:
            kept.append("~")
        out.append(" ".join(kept) if kept else "~")
    return "\n".join(out)

class Project:
    def __init__(self, path):
        self.path = os.path.abspath(path)
        self.root = os.path.dirname(self.path)
        self.name = os.path.splitext(os.path.basename(self.path))[0]
        self.lock_path = self.path + ".lock"
        self.read_only = False
        self.locked_by = None
        self.data = None

    @staticmethod
    def new_data():
        now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
        return {"format": PROJECT_FORMAT, "version": 1, "app_version": APP_VERSION, "created": now,
                "last_scan": None, "last_stats": {}, "files": {}, "history": {}, "decisions": {}, "manual": {},
                "client_lists": [], "mode": "folder", "members": []}

    def load(self):
        with open(self.path, "r", encoding="utf-8-sig") as f:
            data = json.load(f)
        if data.get("format") != PROJECT_FORMAT:
            raise ValueError("This is not an Acronym Hunter project file.")
        base = self.new_data(); base.update(data); self.data = base

    def save(self):
        if self.read_only or self.data is None:
            return
        self.data["app_version"] = APP_VERSION
        tmp = self.path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self.data, f, ensure_ascii=False, indent=1)
        for attempt in range(6):   # OneDrive/SharePoint sync can hold the file for a moment while uploading
            try:
                os.replace(tmp, self.path); return
            except PermissionError:
                if attempt == 5:
                    raise
                time.sleep(0.5)

    # ---- single-editor lock ----
    @staticmethod
    def _me():
        return {"user": getpass.getuser(), "computer": platform.node()}

    def acquire_lock(self, force=False):
        """True if we hold the lock; False if someone else holds a fresh lock (self.locked_by says who)."""
        if os.path.exists(self.lock_path) and not force:
            try:
                info = json.load(open(self.lock_path, encoding="utf-8"))
                age_h = (datetime.datetime.now() - datetime.datetime.fromisoformat(info["time"])).total_seconds() / 3600
                me = self._me()
                if (info.get("user"), info.get("computer")) != (me["user"], me["computer"]) and age_h < LOCK_STALE_HOURS:
                    self.locked_by = info
                    return False
            except Exception:
                diag.swallowed()
                pass
        with open(self.lock_path, "w", encoding="utf-8") as f:
            json.dump(dict(self._me(), time=datetime.datetime.now().isoformat(timespec="seconds")), f)
        return True

    def release_lock(self):
        try:
            if not self.read_only and os.path.exists(self.lock_path):
                info = json.load(open(self.lock_path, encoding="utf-8"))
                if (info.get("user"), info.get("computer")) == (self._me()["user"], self._me()["computer"]):
                    os.remove(self.lock_path)
        except Exception:
            diag.swallowed()
            pass

    # ---- files ----
    def member_path(self, rel):
        """Absolute path of a project file entry (relative to the project folder, or absolute)."""
        return rel if os.path.isabs(rel) else os.path.join(self.root, rel.replace("/", os.sep))

    def to_member(self, path):
        path = os.path.abspath(path)
        try:
            rel = os.path.relpath(path, self.root)
        except ValueError:          # another drive
            return path
        return path if rel.startswith("..") else rel.replace(os.sep, "/")

    def list_files(self):
        if self.data.get("mode") == "files":   # a saved set of files, not a whole folder
            return sorted(m for m in self.data.get("members", []) if os.path.isfile(self.member_path(m)))
        out = []
        for dirpath, dirnames, files in os.walk(self.root):
            dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS]
            for f in files:
                if f.startswith("~$") or os.path.splitext(f)[1].lower() not in SUPPORTED_EXTS:
                    continue
                out.append(os.path.relpath(os.path.join(dirpath, f), self.root).replace(os.sep, "/"))
        return sorted(out)

def project_findings(project, status_callback, full=False, rescan=True):
    """Scan the project folder, reusing saved results for unchanged files. rescan=False rebuilds from saved results only."""
    m = build_matcher()
    findings, stats = {}, {"files": 0, "scanned": 0, "reused": 0, "removed": 0}
    _OCR_CACHE.clear()
    weak = weak_terms(m)
    cache, sig = project.data["files"], (engine_signature(weak), text_signature(weak), ocr_signature())
    files = project.list_files() if rescan else sorted(cache)
    stats["files"] = len(files)
    stats["removed"] = len(set(cache) - set(files)) if rescan else 0
    new_cache = {}
    try:
        _project_loop(project, files, cache, sig, weak, m, findings, stats, new_cache, status_callback, full, rescan)
    except ScanCancelled:
        if rescan:   # keep what was read so far; the next rescan continues from there
            project.data["files"].update(new_cache)
        raise
    if rescan:
        project.data["files"] = new_cache
    finalize_symbols(findings, m)
    apply_client_overlay(findings, m)
    summarize_equations(findings)
    return dict(sorted(findings.items(), key=lambda x: x[1]["count"], reverse=True)), stats

def _project_loop(project, files, cache, sig, weak, m, findings, stats, new_cache, status_callback, full, rescan):
    for idx, rel in enumerate(files, start=1):
        check_cancel()
        absf = project.member_path(rel)
        entry = cache.get(rel)
        st = os.stat(absf) if rescan and os.path.exists(absf) else None
        fresh = (entry and entry_current(entry, sig) and
                 (not rescan or (st and entry["size"] == st.st_size and abs(entry["mtime"] - st.st_mtime) < 1.0)))
        if (not full and fresh) or not rescan:
            if entry is None:
                continue
            stats["reused"] += 1
            if rescan:
                status_callback(f"[{idx}/{len(files)}] {rel} - unchanged, using saved results")
        else:
            _OCR_CALLS[0] = 0
            chunks, skips, err = scan_file(absf, status_callback, idx, len(files))
            entry = cache_entry(st, weak, chunks, skips, err, _OCR_CALLS[0], sig)
            stats["scanned"] += 1
        new_cache[rel] = entry
        if entry.get("error"):
            record_file_error(findings, rel, entry["error"])
        for label, reason in entry.get("skips", []):
            record_skip(findings, label, reason)
        match_chunks(findings, entry["chunks"], m)

def apply_project_history(project, findings, update=True):
    """Mark each finding New / Changed / Unchanged against the last scan; add 'No longer found' rows."""
    hist = project.data["history"]
    today = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    for term, d in findings.items():
        if d.get("file_error"):
            continue
        locs = sorted(d["locations"])[:200]
        h = hist.get(term)
        if update:
            if h is None or h.get("state") == "gone":
                state = "new"
            elif h.get("count") != d["count"] or h.get("locations") != locs:
                state = "changed"
            else:
                state = "unchanged"
            first = h["first_seen"] if h else today
            hist[term] = {"first_seen": first, "last_seen": today, "count": d["count"], "locations": locs,
                          "category": d["category"], "definition": d["selected_definition"], "state": state}
            h = hist[term]
        if h:
            d["change"] = CHANGE_LABELS.get(h.get("state"), "")
            d["first_seen"], d["last_seen"] = h.get("first_seen", ""), h.get("last_seen", "")
    for term, h in hist.items():
        if term in findings:
            continue
        if update:
            h["state"] = "gone"
        if h.get("state") == "gone":
            findings[term] = {"category": h.get("category", ""), "count": 0, "locations": h.get("locations", []),
                              "definitions": [], "in_db": True, "excluded": False, "gone": True,
                              "selected_definition": h.get("definition", ""), "status": "No longer found",
                              "change": CHANGE_LABELS["gone"], "first_seen": h.get("first_seen", ""),
                              "last_seen": h.get("last_seen", "")}
    if update:
        project.data["last_scan"] = today

def apply_ocr_fix(findings, wrong, right, cat, loc, include_db, exclude_db):
    tagged = f"{loc} [OCR fix: {wrong}]"
    if right in include_db:
        record_included(findings, right, tagged, include_db)
    elif right in exclude_db:
        record_excluded(findings, right, tagged, exclude_db)
    else:
        record_candidate(findings, right, cat, tagged)

def record_included(findings, term, loc, include_db):
    if _HITS is not None:
        _HITS.append(term)
    info = include_db[term]
    defs = info["definitions"]
    multi = len(defs) > 1
    if term not in findings:
        findings[term] = {
            "category": info["category"], "count": 0, "locations": [],
            "definitions": defs, "in_db": True, "excluded": False,
            "selected_definition": (f"MULTIPLE ({len(defs)} options) - Choose: " + " OR ".join(defs)) if multi else defs[0],
            "status": "CHOOSE DEFINITION" if multi else "Resolved (In DB)"
        }
        if info.get("known"):           # recognized, kept, left off the acronym lists; never asked to define at first use
            findings[term].update(known=True, status=KNOWN_STATUS + (f": {info['omit_reason']}" if info.get("omit_reason") else ""),
                                  selected_definition=defs[0], omit_reason=info.get("omit_reason", ""))
    findings[term]["count"] += 1
    if loc not in findings[term]["locations"]:
        findings[term]["locations"].append(loc)

def record_excluded(findings, term, loc, exclude_db):
    if _HITS is not None:
        _HITS.append(term)
    info = exclude_db[term]
    if term not in findings:
        findings[term] = {
            "category": info["category"], "count": 0, "locations": [],
            "definitions": [info["definition"]], "in_db": True, "excluded": True,
            "selected_definition": f"[EXCLUDED] {info['definition']}", "status": "Excluded by DB rule"
        }
    findings[term]["count"] += 1
    if loc not in findings[term]["locations"]:
        findings[term]["locations"].append(loc)

def record_candidate(findings, term, cat, loc):
    if _HITS is not None:
        _HITS.append(term)
    if term not in findings:
        findings[term] = {
            "category": cat, "count": 0, "locations": [],
            "definitions": [], "in_db": False, "excluded": False,
            "selected_definition": "[Not in Database — Definition omitted]", "status": "NEW CANDIDATE"
        }
    findings[term]["count"] += 1
    if loc not in findings[term]["locations"]:
        findings[term]["locations"].append(loc)

def write_db_rows(rows):
    """Atomic rewrite: write a temp file beside the DB, then swap it in (no half-written DB).
    rows: [term, category, definition, action(, {extra column: value})]. Keeps the file's own column order,
    header names, and extra columns, so a reordered or extended database stays the way the user set it up."""
    layout = db_layout()
    have = {h for _f, h in layout}
    for r in rows:                                   # a new extra column (Omit reason) is added at the end
        for h in (r[4] if len(r) > 4 and isinstance(r[4], dict) else {}):
            if h and h not in have and not db_field(h):
                layout.append((None, h)); have.add(h)
    tmp = DB_FILE + ".tmp"
    with open(tmp, "w", newline="", encoding=CSV_WRITE_ENCODING) as f:
        w = csv.writer(f)
        w.writerow([h for _, h in layout])
        for r in rows:
            vals = {"Term": r[0], "Category": r[1], "Definition": r[2], "Action": r[3]}
            extra = r[4] if len(r) > 4 else {}
            w.writerow([vals[f] if f else extra.get(h, "") for f, h in layout])
    os.replace(tmp, DB_FILE)

def read_db_rows():
    rows = []
    if os.path.exists(DB_FILE):
        for r in read_csv_rows(DB_FILE):
            d = db_row(r)
            extra = {k: v for k, v in r.items() if k and not db_field(k)}
            rows.append([d.get("Term", ""), d.get("Category") or "Acronym",
                         d.get("Definition", ""), d.get("Action") or "Include", extra])
    return rows

DB_LOCKED_MSG = "A file the app needs is open in another program (usually Excel). Close it and try again."

def read_ocr_fix_rows():
    rows = []
    try:
        for r in read_csv_rows(OCR_FIXES_FILE):
            wrong, right = (r.get("Misread") or "").strip(), (r.get("Correct") or "").strip()
            if wrong or right:
                rows.append([wrong, right, (r.get("Note") or "").strip()])
    except FileNotFoundError:
        pass
    return rows

def write_ocr_fix_rows(rows):
    os.makedirs(CONFIG_DIR, exist_ok=True)
    tmp = OCR_FIXES_FILE + ".tmp"
    with open(tmp, "w", newline="", encoding=CSV_WRITE_ENCODING) as f:
        w = csv.writer(f)
        w.writerow(["Misread", "Correct", "Note"])
        w.writerows(rows)
    os.replace(tmp, OCR_FIXES_FILE)

def update_csv_exclusion(term, cat, reason):
    update_csv_exclusions([(term, cat, reason)])

def norm_def(s):
    """Definitions compare without case, spacing or punctuation: 'Statement of Work' = 'statement-of-work'."""
    return re.sub(r"[^a-z0-9]+", " ", (s or "").lower()).strip()

def _action_kind(a):
    a = str(a or "").strip().lower()
    return "omit" if a.startswith(("known", "omit")) else "exclude" if a.startswith("exclude") else "include"

def db_add(term, cat, defn, action="Include", extra=None):
    """Add one database row. The same term may have several rows, one per different definition.
    Returns ('added' | 'updated' | 'duplicate', existing row or None): 'duplicate' = this term already has this
    definition with this action (nothing written); 'updated' = same definition, the action changed (no new row)."""
    rows = read_db_rows()
    ft, nd = fold(term), norm_def(defn)
    same = [r for r in rows if fold(r[0]) == ft and norm_def(r[2]) == nd]
    if same:
        r = same[0]
        if _action_kind(r[3]) == _action_kind(action):
            return "duplicate", r                     # nothing changes, not even its omit reason
        r[3] = action
        if extra:
            r[4].update(extra)
        write_db_rows(rows)
        return "updated", r
    rows.append([term, cat, defn, action, dict(extra or {})])
    write_db_rows(rows)
    return "added", None

def db_definitions(term):
    """[(definition, action)] the database holds for a term."""
    ft = fold(term)
    return [(r[2], r[3]) for r in read_db_rows() if fold(r[0]) == ft]

def update_csv_exclusions(items):
    """Mark several terms Exclude in one database write. items: [(term, category, reason)]."""
    rows = read_db_rows()
    for term, cat, reason in items:
        found = False
        for r in rows:
            if fold(r[0]) == fold(term):
                r[3] = "Exclude"
                found = True
        if not found:
            rows.append([term, cat.replace(" (Candidate)", ""), reason, "Exclude"])
    write_db_rows(rows)

# ==============================================================================
# 6b. KNOWLEDGE PACKS: trained models + database + OCR fixes + ignore list, shared via approved SharePoint
# ==============================================================================
KNOWLEDGE_FORMAT = "AcronymHunterKnowledge"
KNOWLEDGE_INFO_FILE = os.path.join(CONFIG_DIR, "knowledge_installed.json")
KNOWLEDGE_BACKUP_DIR = os.path.join(CONFIG_DIR, "knowledge_backups")
PACK_MODEL_FILES = ("rec.onnx", "det.onnx", "cls.onnx", "rec_keys.txt", CUSTOM_TESS_MODEL + ".traineddata") + \
    tuple("doctr/" + f for f in DOCTR_FILES)      # trained docTR models (Training Kit), checked by SHA-256 on load
PACK_DATA_FILES = {"acronym_database.csv": "DB", "ocr_corrections.csv": "FIXES", "ocr_ignore.csv": "IGNORE",
                   "compliance_learned.json": "LEARNED",   # compliance learning: hashed counts only, no RFP words (2.10.0)
                   "offerings.csv": "OFFERINGS",           # the offerings catalog (2.19.0)
                   "keyword_learning.json": "KWLEARN"}     # keyword Keep/Reject learning: hashed counts only (2.19.0)

def _team_files():
    """(name in the pack, path) for the offerings catalog and keyword learning (paths follow CONFIG_DIR)."""
    return [("offerings.csv", os.path.join(CONFIG_DIR, "offerings.csv")),
            ("keyword_learning.json", os.path.join(CONFIG_DIR, "keyword_learning.json"))]
COMPLIANCE_LEARNED_FILE = os.path.join(CONFIG_DIR, "compliance_learned.json")
PACK_MAX_MEMBER_BYTES = 400 * 1024 * 1024
PACK_ROOT = "knowledge/"

def _sha256_bytes(data):
    return hashlib.sha256(data).hexdigest()

def _sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()

def _knowledge_sources():
    """(name inside the pack, path on disk) for every knowledge file that exists in this copy."""
    out = []
    for name, path in ([("acronym_database.csv", DB_FILE), ("ocr_corrections.csv", OCR_FIXES_FILE),
                       ("ocr_ignore.csv", OCR_IGNORE_FILE), ("compliance_learned.json", COMPLIANCE_LEARNED_FILE)] + _team_files()):
        if os.path.isfile(path):
            out.append((name, path))
    for m in PACK_MODEL_FILES:
        path = os.path.join(MODELS_DIR, m)
        if os.path.isfile(path):
            out.append(("models/" + m, path))
    return out

def make_knowledge_pack(out_path, notes=""):
    """Write a Knowledge Pack zip of this copy's database, OCR fixes, ignore list, and models. Returns its info."""
    files = _knowledge_sources()
    now = datetime.datetime.now()
    info = {"format": KNOWLEDGE_FORMAT, "version": now.strftime("%Y.%m.%d-%H%M"), "created": now.strftime("%Y-%m-%d %H:%M"),
            "created_by": getpass.getuser(), "computer": platform.node(), "app_version": APP_VERSION,
            "notes": notes.strip(), "handling": "Contains knowledge learned from internal documents. Approved SharePoint only. Never GitHub or email.",
            "files": {}}
    for name, path in files:
        info["files"][name] = {"sha256": _sha256_file(path), "bytes": os.path.getsize(path)}
    try:
        info["counts"] = {"database_rows": len(read_db_rows()), "ocr_fixes": len(read_ocr_fix_rows()),
                          "ignored_readings": len(read_ocr_ignore_rows())}
    except Exception:
        diag.swallowed()
        info["counts"] = {}
    info["models"] = [n.split("/", 1)[1] for n, _ in files if n.startswith("models/")]
    tmp = out_path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(PACK_ROOT + "pack_info.json", json.dumps(info, indent=1, ensure_ascii=False))
        for name, path in files:
            z.write(path, PACK_ROOT + name)
    os.replace(tmp, out_path)
    return info

def read_knowledge_pack(path):
    """Validate a pack and return (info, {name: bytes}). Refuses unknown members, oversize members, and any
    file whose SHA-256 does not match pack_info.json (a damaged or altered pack)."""
    allowed = {PACK_ROOT + "pack_info.json"} | {PACK_ROOT + n for n in PACK_DATA_FILES} | \
              {PACK_ROOT + "models/" + m for m in PACK_MODEL_FILES}
    with zipfile.ZipFile(path) as z:
        entries = [e for e in z.infolist() if not e.is_dir()]
        for e in entries:
            if e.filename not in allowed:
                raise ValueError(f"Unexpected file in pack: {e.filename}. Not a valid Knowledge Pack.")
            if e.file_size > PACK_MAX_MEMBER_BYTES:
                raise ValueError(f"{e.filename} is too large.")
        if PACK_ROOT + "pack_info.json" not in {e.filename for e in entries}:
            raise ValueError("pack_info.json is missing. Not a Knowledge Pack.")
        info = json.loads(z.read(PACK_ROOT + "pack_info.json").decode("utf-8"))
        if info.get("format") != KNOWLEDGE_FORMAT:
            raise ValueError("Not an Acronym Hunter Knowledge Pack.")
        data = {}
        for name, meta in info.get("files", {}).items():
            member = PACK_ROOT + name
            if member not in allowed:
                raise ValueError(f"Pack lists an unexpected file: {name}")
            blob = z.read(member)
            if _sha256_bytes(blob) != meta.get("sha256"):
                raise ValueError(f"{name} does not match its fingerprint. The pack is damaged or was altered.")
            data[name] = blob
        if len(data) != len(entries) - 1:
            raise ValueError("The pack holds files that pack_info.json does not list.")
    return info, data

def _backup_targets():
    return [DB_FILE, OCR_FIXES_FILE, OCR_IGNORE_FILE, KNOWLEDGE_INFO_FILE, COMPLIANCE_LEARNED_FILE] + \
        [os.path.join(MODELS_DIR, m) for m in PACK_MODEL_FILES] + [p for _n, p in _team_files()]

def backup_knowledge(reason):
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = os.path.join(KNOWLEDGE_BACKUP_DIR, stamp)
    os.makedirs(dest, exist_ok=True)
    manifest = {"reason": reason, "time": stamp, "files": {}}
    for i, path in enumerate(_backup_targets()):
        rel = os.path.relpath(path, BASE_DIR)
        if os.path.isfile(path):
            shutil.copy2(path, os.path.join(dest, f"{i}.bak"))
            manifest["files"][rel] = f"{i}.bak"
        else:
            manifest["files"][rel] = None
    with open(os.path.join(dest, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=1)
    return dest

def list_knowledge_backups():
    if not os.path.isdir(KNOWLEDGE_BACKUP_DIR):
        return []
    return sorted(d for d in os.listdir(KNOWLEDGE_BACKUP_DIR) if os.path.isfile(os.path.join(KNOWLEDGE_BACKUP_DIR, d, "manifest.json")))

def restore_last_knowledge_backup():
    """Put back exactly what the last Install or Load replaced. Returns the reason recorded, or None."""
    backups = list_knowledge_backups()
    if not backups:
        return None
    src = os.path.join(KNOWLEDGE_BACKUP_DIR, backups[-1])
    manifest = json.load(open(os.path.join(src, "manifest.json"), encoding="utf-8"))
    for rel, bak in manifest["files"].items():
        path = os.path.join(BASE_DIR, rel)
        if bak is None:
            if os.path.isfile(path):
                os.remove(path)
        else:
            os.makedirs(os.path.dirname(path), exist_ok=True)
            shutil.copy2(os.path.join(src, bak), path)
    shutil.rmtree(src, ignore_errors=True)
    reset_ocr_engine()
    return manifest.get("reason", "last update")

def reset_ocr_engine():
    """Load models again on the next image (after a model change)."""
    global _ocr_engine, OCR_MODEL_INFO
    global _doctr_engine, DOCTR_DIR
    with _ocr_lock:
        _ocr_engine = None
        OCR_MODEL_INFO = "not loaded yet"
        _doctr_engine = None          # docTR reloads too: trained models may have been installed or removed
        DOCTR_DIR = _find_doctr()
        _doctr_failed[0] = False
    _OCR_CACHE.clear()

def _csv_bytes_rows(blob):
    os.makedirs(SANDBOX_TEMP, exist_ok=True)
    tmp = os.path.join(SANDBOX_TEMP, f"pack_{os.getpid()}_{time.time_ns()}.csv")
    with open(tmp, "wb") as f:
        f.write(blob)
    try:
        return read_csv_rows(tmp)
    finally:
        os.remove(tmp)

def merge_pack_database(blob):
    """Add-only: new terms and new definitions are added. The user's own rows, edits, and Exclude choices win."""
    rows = read_db_rows()
    have_pairs = {(fold(r[0]), r[2].strip().casefold()) for r in rows}
    have_terms = {fold(r[0]) for r in rows}
    excluded = {fold(r[0]) for r in rows if r[3].strip().capitalize() == "Exclude"}
    added = 0
    for raw in _csv_bytes_rows(blob):
        d = db_row(raw)
        term, defn = d.get("Term", ""), d.get("Definition", "")
        cat, action = d.get("Category") or "Acronym", (d.get("Action") or "Include").capitalize()
        key = fold(term)
        if not key:
            continue
        if action == "Exclude":
            if key in have_terms:
                continue
        elif key in excluded or (key, defn.casefold()) in have_pairs or not defn:
            continue
        rows.append([term, cat, defn, action]); have_terms.add(key); have_pairs.add((key, defn.casefold())); added += 1
    if added:
        write_db_rows(rows)
    return added

def merge_pack_fixes(blob):
    rows = read_ocr_fix_rows(); have = {fold(r[0]) for r in rows}; added = 0
    for r in _csv_bytes_rows(blob):
        wrong, right = (r.get("Misread") or "").strip(), (r.get("Correct") or "").strip()
        if wrong and right and fold(wrong) not in have:
            rows.append([wrong, right, (r.get("Note") or "").strip()]); have.add(fold(wrong)); added += 1
    if added:
        write_ocr_fix_rows(rows)
    return added

def merge_pack_ignore(blob):
    rows = read_ocr_ignore_rows(); have = {fold(r[0]).strip() for r in rows}; added = 0
    for r in _csv_bytes_rows(blob):
        reading = (r.get("Reading") or "").strip()
        if reading and fold(reading).strip() not in have:
            rows.append([reading, (r.get("Note") or "").strip()]); have.add(fold(reading).strip()); added += 1
    if added:
        write_ocr_ignore_rows(rows)
    return added

def _check_onnx(path):
    import onnxruntime as ort
    ort.InferenceSession(path, providers=["CPUExecutionProvider"])

def install_knowledge_pack(path):
    """Verify, back up, then install a pack. Models replace the current models; data files merge add-only."""
    info, data = read_knowledge_pack(path)
    models = {n.split("/", 1)[1]: b for n, b in data.items() if n.startswith("models/")}
    backup_knowledge(f"before installing Knowledge Pack {info.get('version', '?')}")
    report = [f"Installed Knowledge Pack {info.get('version', '?')} (made {info.get('created', '?')} by {info.get('created_by', '?')})"]
    if models:
        os.makedirs(MODELS_DIR, exist_ok=True)
        for m in PACK_MODEL_FILES:
            target = os.path.join(MODELS_DIR, m)
            if os.path.isfile(target):
                os.remove(target)
        for m, blob in models.items():
            target = os.path.join(MODELS_DIR, m)
            os.makedirs(os.path.dirname(target), exist_ok=True)
            with open(target, "wb") as f:
                f.write(blob)
        report.append("Models: " + ", ".join(sorted(models)))
    else:
        report.append("Models: none in this pack (current models kept)")
    if "acronym_database.csv" in data:
        report.append(f"Database: {merge_pack_database(data['acronym_database.csv'])} new rows added (your own rows unchanged)")
    if "ocr_corrections.csv" in data:
        report.append(f"OCR fixes: {merge_pack_fixes(data['ocr_corrections.csv'])} added")
    if "ocr_ignore.csv" in data:
        report.append(f"Ignored readings: {merge_pack_ignore(data['ocr_ignore.csv'])} added")
    if "compliance_learned.json" in data:
        import complearn
        L = complearn.Learner(COMPLIANCE_LEARNED_FILE)
        n = L.merge_pack(data["compliance_learned.json"])
        L.save()
        report.append(f"Compliance learning: {n} decisions added from the team (your own decisions unchanged)")
    if "offerings.csv" in data:                        # add-only: offerings already here are kept as they are
        import offerings
        os.makedirs(SANDBOX_TEMP, exist_ok=True)
        tmp = os.path.join(SANDBOX_TEMP, f"pack_offerings_{os.getpid()}.csv")
        with open(tmp, "wb") as f:
            f.write(data["offerings.csv"])
        try:
            a, _same = offerings.import_file(tmp)
        finally:
            os.remove(tmp)
        report.append(f"Offerings catalog: {a} offerings added (your own unchanged)")
    if "keyword_learning.json" in data:                # hashed counts: the larger count per phrase wins (no double counting)
        import keywords
        mine, theirs = keywords.learning(), json.loads(data["keyword_learning.json"].decode("utf-8"))
        n = 0
        for side in ("keep", "reject"):
            for h, c in (theirs.get(side) or {}).items():
                if isinstance(c, int) and c > mine[side].get(h, 0):
                    mine[side][h] = c; n += 1
        keywords._write("keyword_learning.json", mine)
        report.append(f"Keyword learning: {n} team decisions added")
    with open(KNOWLEDGE_INFO_FILE, "w", encoding="utf-8") as f:
        json.dump({"version": info.get("version"), "created": info.get("created"), "created_by": info.get("created_by"),
                   "notes": info.get("notes", ""), "models": sorted(models), "pack_sha256": _sha256_file(path),
                   "installed": datetime.datetime.now().strftime("%Y-%m-%d %H:%M")}, f, indent=1)
    reset_ocr_engine()
    return "\n".join(report)

def load_trained_models(folder):
    """Copy trained models from a Training Kit output folder into config\\models.
    docTR kit output (doctr_models.json + two ONNX files) goes to config\\models\\doctr after its SHA-256 check."""
    if os.path.isfile(os.path.join(folder, "doctr_models.json")):
        if not all(os.path.isfile(os.path.join(folder, f)) for f in DOCTR_FILES):
            raise ValueError("A docTR model folder needs doctr_det.onnx, doctr_reco.onnx and doctr_models.json.")
        if not doctr_folder_ok(folder):
            raise ValueError("These docTR models do not match the SHA-256 in their doctr_models.json. Not installed.")
        for f in DOCTR_FILES[:2]:
            _check_onnx(os.path.join(folder, f))
        backup_knowledge(f"before loading docTR models from {folder}")
        dest = os.path.join(MODELS_DIR, "doctr")
        os.makedirs(dest, exist_ok=True)
        for f in DOCTR_FILES:
            shutil.copy2(os.path.join(folder, f), os.path.join(dest, f))
        reset_ocr_engine()
        return ["doctr/" + f for f in DOCTR_FILES]
    found = [m for m in PACK_MODEL_FILES if os.path.isfile(os.path.join(folder, m))]
    if not any(m.endswith((".onnx", ".traineddata")) for m in found):
        raise ValueError("No acronym.traineddata, rec.onnx, det.onnx, or cls.onnx in that folder.")
    if "rec.onnx" in found and "rec_keys.txt" not in found:
        raise ValueError("rec.onnx needs its rec_keys.txt beside it.")
    for m in found:
        if m.endswith(".onnx"):
            _check_onnx(os.path.join(folder, m))
    backup_knowledge(f"before loading models from {folder}")
    os.makedirs(MODELS_DIR, exist_ok=True)
    for m in found:
        shutil.copy2(os.path.join(folder, m), os.path.join(MODELS_DIR, m))
    reset_ocr_engine()
    return found

def knowledge_status():
    try:
        k = json.load(open(KNOWLEDGE_INFO_FILE, encoding="utf-8"))
        return f"Knowledge Pack {k.get('version')} (installed {k.get('installed')})"
    except Exception:
        diag.swallowed()
        models = [m for m in PACK_MODEL_FILES if os.path.isfile(os.path.join(MODELS_DIR, m))]
        return "No Knowledge Pack installed" + (f"; local models: {', '.join(models)}" if models else "; built-in models")

# ---- App icons (assets\scanner.ico, assets\trainer.ico; bundled into _internal\assets by the build) ----
ACTIVE_ICON = "scanner.ico"   # the Trainer sets "trainer.ico"

def asset_path(name):
    for base in (getattr(sys, "_MEIPASS", ""), BASE_DIR):
        p = os.path.join(base, "assets", name)
        if base and os.path.isfile(p):
            return p
    return None

def apply_window_icon(win, name=None):
    path = asset_path(name or ACTIVE_ICON)
    if not path:
        return
    def _set():
        try:
            win.iconbitmap(path)
        except Exception:
            diag.swallowed()
            pass
    _set()
    win.after(300, _set)   # CustomTkinter swaps in its own icon about 200 ms after a window opens

_ctk_toplevel_init = ctk.CTkToplevel.__init__
def _toplevel_init(self, *a, **k):
    _ctk_toplevel_init(self, *a, **k)
    apply_window_icon(self)
ctk.CTkToplevel.__init__ = _toplevel_init   # every pop-up window gets the app icon too

# ==============================================================================
# 7. DESKTOP GUI APPLICATION
# ==============================================================================
SETTINGS_FILE = os.path.join(SANDBOX_DIR, "settings.json")
ALL_TYPES = "All types"
STATUS_FILTERS = ["All statuses", "Needs choice", "New candidates", "Resolved", "Client list", "Client and database differ",
                  "Not defined at first use", "Equations and symbols", "Undefined symbols",
                  "Added manually", "Customer documents", "Customer and database differ", "In effort list", "Not in effort list", "Effort: needs author", "Effort: approved",
                  "New since last scan", "Changed since last scan", "No longer found", "Errors / not scanned",
                  "Context suggests a meaning", "New in this scan (not in the effort list)", "Found in omit list",
                  "Known in: customer docs", "Known in: customer list", "Known in: master database only", "Known in: no source",
                  "Database update", "Excluded"]
MANUAL_MARK = ""   # statuses are plain words (2.7.0: no emoji in the interface or in exports)
# Status kinds by leading words. Older projects saved statuses with an emoji in front; status_text() drops it.
_KIND_WORDS = (("error", ("COULD NOT READ FILE", "NOT SCANNED")),
               ("choice", ("CHOOSE", "MULTIPLE", "CUSTOMER DOCS DIFFER")),
               ("client_differs", ("CLIENT LIST (DATABASE DIFFERS)",)),
               ("candidate", ("NEW CANDIDATE", "SYMBOL NOT DEFINED", "CHECK EQUATION PICTURE")),
               ("manual", ("ADDED MANUALLY",)),
               ("known", ("FOUND IN OMIT LIST", "KNOWN NAME")),
               ("resolved", ("RESOLVED", "DEFINED IN DOCUMENT", "USER SELECTED", "ADDED TO DB", "CLIENT LIST",
                             "CUSTOMER DOCUMENT", "EFFORT LIST:")),
               ("gone", ("NO LONGER FOUND",)),
               ("excluded", ("EXCLUDED", "NOT A TERM IN THIS EFFORT")))
RESOLVED_WORDS = ("Resolved", "Defined in document", "User Selected", "Added to DB", "Client list")   # the Resolved filter
STATUS_RANK = {"error": 0, "choice": 1, "client_differs": 2, "candidate": 3, "manual": 4, "resolved": 5, "known": 6, "gone": 7,
               "excluded": 8}

def status_text(st):
    """The status without any leading symbol (statuses saved by older versions began with an emoji)."""
    st = str(st or "")
    i = 0
    while i < len(st) and not (st[i].isascii() and (st[i].isalnum() or st[i] == "[")):
        i += 1
    return st[i:]

def status_kind(st):
    up = status_text(st).upper()
    for kind, words in _KIND_WORDS:
        if up.startswith(words):
            return kind
    return "other"
LOC_MAX_LINES = 3   # Locations wrap to at most this many lines in the table; double-click a row for all of them
MANUAL_CATEGORIES = ["Acronym", "Abbreviation", "Unit", "Other"]
COLUMNS = (  # id, heading, base width, anchor
    ("term", "Term / Unit", 120, "w"),
    ("category", "Type", 170, "w"),
    ("count", "Count", 70, "center"),
    ("change", "Change", 130, "w"),
    ("status", "Status", 210, "w"),
    ("definition", "Definition", 300, "w"),
    ("knownin", "Known in", 190, "w"),
    ("firstuse", "First Use (acronyms)", 260, "w"),
    ("locations", "Source Locations", 420, "w"),
)
FONT_MIN, FONT_MAX, FONT_DEFAULT = 9, 28, 13

def base_type(category):
    return category.replace(" (Candidate)", "")

def load_settings():
    try:
        with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        diag.swallowed()
        return {}

def save_settings(data):
    try:
        with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f)
    except Exception:
        diag.swallowed()
        pass

TRAIN_DIR = os.path.join(BASE_DIR, "training_data")   # the Trainer's saved samples (never shipped, never committed)

TRAINING_APP = False   # True only inside Acronym Hunter Training (training_app.py sets it) (2.15.0)

def trainer_available():
    """The Trainer, Improve Reading and Load Trained Models: only in Acronym Hunter Training (2.15.0)."""
    return TRAINING_APP and bool(APP_CFG["allow_training"]) and importlib.util.find_spec("trainer_gui") is not None

def corrections_on():
    """Report misreading and Not text save corrections for the training team, unless allow_training is false."""
    return bool(APP_CFG["allow_training"])

class AcronymAppGUI(KeywordsMixin, MasterMixin, HomeMixin, WizardMixin, AIMixin, WorkspaceMixin, EffortMixin, VolumeCheckMixin, SetupMixin, ComplianceMixin, ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title(f"Acronym Hunter {APP_VERSION}: acronyms, abbreviations and units (100% local)")
        apply_window_icon(self)
        self.geometry("1280x780")
        self.minsize(1000, 560)
        ctk.set_appearance_mode("System")
        ctk.set_default_color_theme("blue")

        self.report_callback_exception = self._on_tk_error
        self.queued_files = []
        self.results = {}
        self.selected_term = None
        self.sort_col, self.sort_desc = "count", True
        self.settings = load_settings()
        self.app_cfg = load_app_settings()
        diag.usage_on = bool(self.app_cfg.get("usage_logging", True))
        self.after(800, self._register_usage)
        self.manual_items = {}   # term -> finding dict; kept across rescans until removed
        self.project = None      # open Project, or None for one-off scans
        self.client_lists = []   # client acronym lists (RFPs); saved in the project when one is open
        set_client_lists(self.client_lists)
        self.project_stats = {}
        self.advanced_on = bool(self.settings.get("advanced", self.app_cfg["advanced_features_default"])) \
            if self.app_cfg["allow_advanced_toggle"] else self.app_cfg["advanced_features_default"]
        self.font_size = int(self.settings.get("font_size", FONT_DEFAULT))
        self.font_size = max(FONT_MIN, min(FONT_MAX, self.font_size))

        # ---- Layout (2.9.0): menu bar + a resizable, hideable sidebar on the left, the work area on the right ----
        self.body = tk.PanedWindow(self, orient="horizontal", sashwidth=5, bd=0, sashrelief="flat",
                                   bg="#CBD5E1" if ctk.get_appearance_mode() == "Light" else "#334155")
        self.body.pack(fill="both", expand=True)
        wsc = ctk.ScalingTracker.get_window_scaling(self)
        self.side = ctk.CTkFrame(self.body, corner_radius=0, fg_color=("#F1F5F9", "#1E293B"))
        self.main = ctk.CTkFrame(self.body, corner_radius=0, fg_color=("gray92", "gray14"))   # the window colour, light or dark
        self.side_w = int(self.settings.get("sidebar_w", int(250 * wsc)))
        self.body.add(self.side, minsize=int(190 * wsc), width=self.side_w, stretch="never")
        self.body.add(self.main, minsize=int(560 * wsc), stretch="always")
        self.body.bind("<ButtonRelease-1>", self._sidebar_resized, add="+")
        side_bottom = ctk.CTkFrame(self.side, fg_color="transparent")
        side_bottom.pack(side="bottom", fill="x", padx=10, pady=(4, 10))
        sb = ctk.CTkScrollableFrame(self.side, fg_color="transparent")
        sb.pack(fill="both", expand=True, padx=(4, 0), pady=(8, 0))
        self.sidebar = sb

        B = dict(height=32, font=("Segoe UI", 13), anchor="w", corner_radius=6)
        SEC = dict(fg_color=("#E2E8F0", "#334155"), hover_color=("#CBD5E1", "#475569"), text_color=("#0F172A", "#F1F5F9"),
                   text_color_disabled=("#94A3B8", "#64748B"))
        self._side_btn = lambda text, cmd, primary=False, **kw: ctk.CTkButton(sb, text=text, command=cmd, **B, **({} if primary else SEC), **kw)

        def section(title):
            self.ai_section_header(sb, title)      # title with its No AI / ML / Gen AI labels and a ? (2.12.0)
        self._side_section = section
        P = dict(fill="x", padx=8, pady=2)

        section("SCAN")
        self.btn_files = self._side_btn("Select Files...", self.select_files, primary=True); self.btn_files.pack(**P)
        self.btn_folder = self._side_btn("Scan Folder...", self.select_folder, primary=True); self.btn_folder.pack(**P)
        self.btn_again = self._side_btn("Scan Again", self.scan_again, state="disabled"); self.btn_again.pack(**P)
        self.btn_project = self._side_btn("Project", self.show_project_menu); self.btn_project.pack(**P)
        self.project_menu = tk.Menu(self, tearoff=0)

        section("REVIEW")
        self.type_var = ctk.StringVar(value=ALL_TYPES)
        ctk.CTkLabel(sb, text="Type", anchor="w").pack(fill="x", padx=8)
        self.opt_type = ctk.CTkOptionMenu(sb, variable=self.type_var, values=[ALL_TYPES], command=lambda _v: self.populate_table())
        self.opt_type.pack(**P)
        self.status_var = ctk.StringVar(value=STATUS_FILTERS[0])
        ctk.CTkLabel(sb, text="Status", anchor="w").pack(fill="x", padx=8)
        self.opt_status = ctk.CTkOptionMenu(sb, variable=self.status_var, values=STATUS_FILTERS, command=lambda _v: self.populate_table())
        self.opt_status.pack(**P)
        ctk.CTkLabel(sb, text="Search", anchor="w").pack(fill="x", padx=8)
        self.search_var = ctk.StringVar()
        self.entry_search = ctk.CTkEntry(sb, textvariable=self.search_var)
        self.entry_search.pack(**P)
        SP.attach(self.entry_search)
        self.search_var.trace_add("write", lambda *_a: self.populate_table())
        row = ctk.CTkFrame(sb, fg_color="transparent"); row.pack(**P)
        self.show_exc_var = ctk.BooleanVar(value=False)
        self.chk_exc = ctk.CTkCheckBox(row, text="Show excluded", variable=self.show_exc_var, command=self.populate_table)
        self.chk_exc.pack(side="left")
        ctk.CTkButton(row, text="Clear", width=56, height=26, command=self.clear_filters, **SEC).pack(side="right")
        self.lbl_showing = ctk.CTkLabel(sb, text="", text_color="gray", anchor="w")
        self.lbl_showing.pack(fill="x", padx=8)

        section("EFFORT")
        self.effort_init(sb, dict(B, **SEC))
        section("COMPLIANCE")
        self.compliance_init(sb, dict(B, **SEC))
        self.workspace_init()          # top bar with the effort switcher; per-effort saved work (2.11.0)
        self.ai_init()                 # AI and ML switch, state line (2.12.0)

        section("EXPORT")
        self.btn_export = self._side_btn("Export CSV...", self.export_csv, state="disabled"); self.btn_export.pack(**P)

        section("MORE")
        self.btn_add = self._side_btn("Add Item...", self.open_add_item); self.btn_add.pack(**P)
        self.btn_client = self._side_btn("Client Lists...", self.open_client_lists); self.btn_client.pack(**P)
        self.btn_tools = self._side_btn("Tools", self.show_tools_menu); self.btn_tools.pack(**P)
        self.btn_improve_top = self._side_btn("Improve Reading...", self.open_improve)
        self.tools_menu = tk.Menu(self, tearoff=0)
        # kept for code that refers to them; reached through the Tools menu
        self.btn_db = self.btn_know = self.btn_fixes = self.btn_trainer = None

        # pinned at the bottom of the sidebar: text size, Advanced, Guide, Feedback, Hide
        size_row = ctk.CTkFrame(side_bottom, fg_color="transparent"); size_row.pack(fill="x", pady=2)
        ctk.CTkLabel(size_row, text="Text size").pack(side="left")
        self.btn_bigger = ctk.CTkButton(size_row, text="A+", width=40, height=28, command=lambda: self.change_font_size(+1), **SEC)
        self.btn_bigger.pack(side="right")
        self.lbl_size = ctk.CTkLabel(size_row, text="", width=48)
        self.lbl_size.pack(side="right")
        self.btn_smaller = ctk.CTkButton(size_row, text="A-", width=40, height=28, command=lambda: self.change_font_size(-1), **SEC)
        self.btn_smaller.pack(side="right")
        self.adv_var = ctk.BooleanVar(value=self.advanced_on)
        self.sw_adv = ctk.CTkSwitch(side_bottom, text="Advanced", variable=self.adv_var, command=self.toggle_advanced)
        if self.app_cfg["allow_advanced_toggle"]:
            self.sw_adv.pack(fill="x", pady=(4, 4))
        hb = ctk.CTkFrame(side_bottom, fg_color="transparent"); hb.pack(fill="x", pady=2)
        self.btn_help = ctk.CTkButton(hb, text="Guide", width=10, height=28, command=lambda: (diag.activity("guide_opened"), HG.open_guide(self)), **SEC)
        self.btn_help.pack(side="left", fill="x", expand=True, padx=(0, 3))
        self.btn_feedback = ctk.CTkButton(hb, text="Feedback", width=10, height=28, command=lambda: FB.open_feedback(self), **SEC)
        self.btn_feedback.pack(side="left", fill="x", expand=True, padx=(3, 0))
        self.btn_hide_side = ctk.CTkButton(side_bottom, text="Hide sidebar  (Ctrl+B)", height=26, command=self.toggle_sidebar,
                                           fg_color="transparent", text_color=("#475569", "#94A3B8"), hover_color=("#E2E8F0", "#334155"))
        self.btn_hide_side.pack(fill="x", pady=(4, 0))
        self.bind_all("<Control-Shift-F>", lambda e: FB.open_feedback(self))
        self.bind_all("<Control-Shift-f>", lambda e: FB.open_feedback(self))
        self.bind_all("<F1>", lambda e: (diag.activity("guide_opened"), HG.open_guide(self)))

        # ---- work area ----
        M = self.main
        self.hdr = ctk.CTkFrame(M, fg_color="transparent", height=30)   # packed only while the sidebar is hidden
        self.btn_show_side = ctk.CTkButton(self.hdr, text="Show sidebar", width=110, height=26, command=self.toggle_sidebar, **SEC)
        self.lbl_project = ctk.CTkLabel(M, text="", font=("Arial", 13, "bold"), text_color="#1E3A8A")
        self.status = ctk.CTkLabel(M, text="Ready. Running offline: nothing leaves this PC.", font=("Arial", 13))
        self.status.pack(pady=(8, 4))
        self.protocol("WM_DELETE_WINDOW", self.on_close)
        atexit.register(self._release_project_lock)

        self.banner = ctk.CTkFrame(M, fg_color=("#ECFDF5", "#064E3B"))   # "Improve Reading" (shown when there are new corrections)
        self.lbl_banner = ctk.CTkLabel(self.banner, text="", text_color=("#065F46", "#D1FAE5"))
        self.lbl_banner.pack(side="left", padx=12, pady=6)
        self.btn_improve = ctk.CTkButton(self.banner, text="Improve Reading...", width=150, height=28, fg_color="#0F766E", command=self.open_improve)
        self.btn_improve.pack(side="left", padx=6, pady=6)
        self._training = False
        self.scanned_paths = {}   # file name -> full path of files scanned this session (for "Check in picture")

        self.prog = ctk.CTkProgressBar(M)
        self.prog.pack(fill="x", padx=15, pady=(0, 8))
        self.btn_cancel = ctk.CTkButton(M, text="Cancel Scan", fg_color="#DC2626", width=130, height=30,
                                        command=self.cancel_scan)   # shown only while a scan runs
        self.prog.set(0)

        # ---- Scan review (the results table) and the effort home (2.14.0) ----
        self.scan_view = ctk.CTkFrame(M, fg_color="transparent")
        self.scan_view.pack(fill="both", expand=True)
        self.home_build(M, self.scan_view)
        t_frame = ctk.CTkFrame(self.scan_view)
        t_frame.pack(fill="both", expand=True, padx=15, pady=(0, 10))
        t_frame.grid_rowconfigure(0, weight=1)
        t_frame.grid_columnconfigure(0, weight=1)
        self.split = tk.PanedWindow(t_frame, orient="horizontal", sashwidth=6, bd=0, sashrelief="flat", bg="#94A3B8")
        self.split.grid(row=0, column=0, sticky="nsew")
        tf = tk.Frame(self.split)
        tf.grid_rowconfigure(0, weight=1)
        tf.grid_columnconfigure(0, weight=1)
        self.split.add(tf, stretch="always", minsize=480)

        self.style = ttk.Style()
        self.style.theme_use("default")
        # Raw Tk text is not scaled by CustomTkinter: scale the table for the display (200% = twice the pixels)
        self.dpi = max(1.0, float(ctk.ScalingTracker.get_widget_scaling(self)))
        self.tree_font = tkfont.Font(family="Segoe UI", size=int(round(self.font_size * self.dpi)))
        self.head_font = tkfont.Font(family="Segoe UI", size=int(round(self.font_size * self.dpi)), weight="bold")

        self.tree = ttk.Treeview(tf, columns=[c[0] for c in COLUMNS], show="headings", selectmode="extended")  # Ctrl/Shift+click: several rows
        for cid, label, width, anchor in COLUMNS:
            self.tree.column(cid, width=width, minwidth=50, anchor=anchor, stretch=(cid == "locations"))
        self.update_headings()

        self.tree.tag_configure("ambiguous", background="#FEF3C7", foreground="#92400E")
        self.tree.tag_configure("candidate", background="#E0F2FE", foreground="#075985")
        self.tree.tag_configure("resolved", background="#FFFFFF", foreground="#1E293B")
        self.tree.tag_configure("excluded", background="#F1F5F9", foreground="#94A3B8")
        self.tree.tag_configure("fileerror", background="#FEE2E2", foreground="#991B1B")
        self.tree.tag_configure("manual", background="#FEF9C3", foreground="#713F12")
        self.tree.tag_configure("gone", background="#E5E7EB", foreground="#6B7280")
        self.tree.tag_configure("known", background="#EEF2FF", foreground="#475569")   # known names: kept, not listed
        self.tree.tag_configure("client", background="#EDE9FE", foreground="#4C1D95")
        self.tree.tag_configure("clientdiff", background="#F5D0FE", foreground="#701A75")
        self.tree.tag_configure("equation", background="#E0E7FF", foreground="#312E81")
        self.tree.tag_configure("effort", background="#CCFBF1", foreground="#115E59")
        self.tree.tag_configure("customer", background="#DCFCE7", foreground="#166534")

        vsb = ttk.Scrollbar(tf, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(tf, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        self.tree.bind("<<TreeviewSelect>>", self.on_select)
        self.tree.bind("<Control-a>", lambda _e: (self.tree.selection_set(self.tree.get_children()), "break")[1])
        self.tree.bind("<Double-1>", self.show_details)
        self.tree.bind("<Button-3>", self.on_right_click)
        self.tree.bind("<ButtonRelease-1>", self._maybe_rewrap, add="+")   # column resized -> rewrap Locations
        self._loc_width = None
        import preview_pane
        self.preview = preview_pane.PreviewPane(self.split, self)
        self.preview_var = tk.BooleanVar(value=bool(self.settings.get("preview", True)))
        if self.preview_var.get():
            self.split.add(self.preview, minsize=int(380 * ctk.ScalingTracker.get_window_scaling(self)) + 40, width=int(400 * ctk.ScalingTracker.get_window_scaling(self)) + 40, stretch="never")

        # Text size shortcuts: Ctrl + / Ctrl - / Ctrl 0 / Ctrl + mouse wheel
        for seq in ("<Control-plus>", "<Control-equal>", "<Control-KP_Add>"):
            self.bind_all(seq, lambda _e: self.change_font_size(+1))
        for seq in ("<Control-minus>", "<Control-KP_Subtract>"):
            self.bind_all(seq, lambda _e: self.change_font_size(-1))
        self.bind_all("<Control-0>", lambda _e: self.change_font_size(0, reset=True))
        self.tree.bind("<Control-MouseWheel>", lambda e: self.change_font_size(+1 if e.delta > 0 else -1))
        self.apply_font_size()

        act = ctk.CTkFrame(self.scan_view)
        act.pack(side="bottom", fill="x", padx=15, pady=(0, 15), before=t_frame)  # never pushed off-screen

        self.lbl_act = ctk.CTkLabel(act, text="Select a row above (Ctrl+click or Shift+click for several):", font=("Arial", 12, "bold"), anchor="w")
        self.lbl_act.pack(fill="x", padx=10, pady=(6, 0))
        act = ctk.CTkFrame(act, fg_color="transparent")   # the controls, on their own row
        act.pack(fill="x")

        self.combo = ctk.CTkComboBox(act, values=["Select a row..."], width=240, state="disabled")
        self.combo.pack(side="left", padx=5, pady=(4, 8))

        self.btn_apply = ctk.CTkButton(act, text="Apply Choice", command=self.apply_choice, state="disabled", width=110)
        self.btn_apply.pack(side="left", padx=5, pady=(4, 8))

        self.entry_def = ctk.CTkEntry(act, placeholder_text="Definition, or correct term for Fix OCR", width=210, state="disabled")
        self.entry_def.pack(side="left", padx=(5, 2), pady=(4, 8))
        SP.button(act, self.entry_def).pack(side="left", padx=(0, 5), pady=(4, 8))

        self.btn_save = ctk.CTkButton(act, text="Include in DB", command=self.save_candidate, state="disabled", fg_color="#0284C7", width=105)
        self.btn_save.pack(side="left", padx=5, pady=(4, 8))
        self.effort_actionbar(act)

        self.btn_ocr = ctk.CTkButton(act, text="Fix OCR", command=self.save_ocr_fix, state="disabled", fg_color="#7C3AED", width=90)
        self.btn_ocr.pack(side="left", padx=5, pady=(4, 8))

        self.btn_exc = ctk.CTkButton(act, text="Exclude", command=self.exclude_term, state="disabled", fg_color="#DC2626", width=110)
        self.btn_exc.pack(side="right", padx=10, pady=(4, 8))

        self.apply_advanced_visibility()
        self._build_menubar()
        if not self.settings.get("sidebar", True):
            self.after(10, lambda: self.toggle_sidebar(save=False))
        self.after(1500, self._banner_tick)
        self.after(300, self.workspace_startup)   # resume the last effort (2.11.0)

    # ---------- training inside the scanner ----------
    def training_state(self):
        """(saved text lines, lines saved since the reader was last improved)."""
        import tess_train
        n = len(tess_train.read_samples(TRAIN_DIR))
        try:
            last = json.load(open(os.path.join(TRAIN_DIR, "model", "training_report.json"), encoding="utf-8")).get("samples", 0)
        except (OSError, ValueError):
            diag.swallowed()
            last = 0
        return n, max(0, n - last)

    def _banner_tick(self):
        try:
            self.update_banner()
        finally:
            self.after(5000, self._banner_tick)

    def update_banner(self):
        if not TRAINING_APP:                      # Acronym Hunter: corrections go to the training team (2.15.0)
            return self._send_banner()
        import tess_train
        show = trainer_available() and OCR_BACKEND == "tesseract" and ai_ml_on()   # ML: hidden while AI and ML are off
        n, new = self.training_state() if show else (0, 0)
        if self._training:
            self.lbl_banner.configure(text="Improving the reader in the background (see the Improve Reading window)...")
            self.btn_improve.configure(state="normal")
        elif not show or new == 0:
            self.banner.pack_forget(); return
        elif n < tess_train.MIN_SAMPLES:
            self.lbl_banner.configure(text=f"{n} correction(s) saved in the Trainer. Improve Reading unlocks at {tess_train.MIN_SAMPLES}.")
            self.btn_improve.configure(state="disabled")
        else:
            steps = min(3000, max(400, 4 * n))
            mins = max(1, round(steps * 0.08 / 60 + n * 0.1 / 60))
            self.lbl_banner.configure(text=f"{new} new corrected line(s) since the reader was last improved ({n} in all). "
                                           f"Improve Reading teaches and applies them (about {mins} min, offline).")
            self.btn_improve.configure(state="normal")
        if not self.banner.winfo_ismapped():
            self.banner.pack(fill="x", padx=15, pady=(0, 6), before=self.prog)

    def improve_reading(self):
        self.open_improve()

    # ---------- corrections for the training team (2.15.0) ----------
    def _send_banner(self):
        try:
            import corrections
            lines, regions = corrections.unsent() if corrections_on() else ([], [])
        except Exception:
            diag.swallowed(); lines, regions = [], []
        n = len(lines) + len(regions)
        if not n:
            self.banner.pack_forget(); return
        self.lbl_banner.configure(text=f"{n} corrected reading(s) saved and not sent yet. Send them to the training team "
                                       "so the next reader learns them.")
        self.btn_improve.configure(text="Send corrections...", command=self.send_corrections, state="normal")
        if not self.banner.winfo_ismapped():
            self.banner.pack(fill="x", padx=15, pady=(0, 6), before=self.prog)

    def send_corrections(self, path=None, everything=False, ask=True):
        """Save a corrections pack for the training team (approved SharePoint folder). Returns the result, or None."""
        import corrections, getpass
        diag.activity("corrections_send_opened")
        lines, regions = corrections.unsent()
        if not (lines or regions or everything):
            if ask:
                messagebox.showinfo("Send corrections", "Nothing new to send. Report misreading and Not text in the picture "
                                    "preview save corrections; they collect here until you send them.")
            return None
        if ask and not messagebox.askyesno(
                "Send corrections", f"{len(lines)} corrected line(s) and {len(regions)} not-text region(s) are ready.\n\n"
                "They are small pictures cut from your documents, with the right text. Save the pack ONLY to the approved "
                "SharePoint folder for the training team. Never email, GitHub or personal storage.\n\nSave the pack now?"):
            return None
        if path is None:
            name = f"AcronymHunter_Corrections_{re.sub(r'[^A-Za-z0-9]+', '', getpass.getuser()) or 'user'}_{datetime.datetime.now():%Y.%m.%d-%H%M}.zip"
            path = filedialog.asksaveasfilename(title="Save corrections to the APPROVED SHAREPOINT folder", initialfile=name,
                                                defaultextension=".zip", filetypes=[("Corrections pack", "*.zip")])
            if not path:
                return None
        r = corrections.export_pack(path, everything=everything)
        diag.activity("corrections_sent", n=r["lines"], items=r["regions"])
        self.update_banner()
        if ask:
            messagebox.showinfo("Send corrections", f"Saved {os.path.basename(path)}: {r['lines']} line(s), {r['regions']} "
                                "region(s). The training team imports it in Acronym Hunter Training; the better reader comes "
                                "back to you in a Knowledge Pack.")
        return r

    def _register_usage(self):
        """Usage logging: name the controls after their code attributes, then log clicks (no text, no data)."""
        diag.register_attrs(self)
        if getattr(self, "preview", None) is not None:
            diag.register_attrs(self.preview, "preview.")
            diag.register_attrs(self.preview.pv, "preview.")
        diag.register(self, "main")
        diag.install_click_logger(self)

    # ---------- menu bar and sidebar (2.9.0) ----------
    def _menu_cmd(self, name, func):
        return lambda: (diag.menu(name), func())

    def _build_menubar(self):
        """Every command lives in the menu bar; the sidebar keeps the everyday ones. Menus that depend on state
        (Project, Effort, Tools) are rebuilt each time they open."""
        mb = tk.Menu(self, tearoff=0)
        c = self._menu_cmd
        f = tk.Menu(mb, tearoff=0)
        f.add_command(label="Select Files...", accelerator="Ctrl+O", command=c("file_select_files", self.select_files))
        f.add_command(label="Scan Folder...", accelerator="Ctrl+Shift+O", command=c("file_scan_folder", self.select_folder))
        f.add_command(label="Scan Again", accelerator="F5", command=c("file_scan_again", self.scan_again))
        f.add_separator()
        proj = tk.Menu(f, tearoff=0, postcommand=lambda: self.show_project_menu(proj))
        f.add_cascade(label="Project", menu=proj)
        f.add_separator()
        f.add_command(label="Load master database...", command=c("file_master_load", self.master_load))
        f.add_command(label="Update results against the master database", command=c("file_master_refresh", self.refresh_from_database))
        f.add_separator()
        f.add_command(label="Export CSV...", accelerator="Ctrl+E", command=c("file_export_csv", self.export_csv))
        f.add_command(label="Restore an earlier version...", command=c("file_restore", self.restore_dialog))
        f.add_separator()
        f.add_command(label="Exit", command=c("file_exit", self.on_close))
        mb.add_cascade(label="File", menu=f)
        e = tk.Menu(mb, tearoff=0)
        e.add_command(label="Add Item...", command=c("edit_add_item", self.open_add_item))
        e.add_command(label="Client Lists...", command=c("edit_client_lists", self.open_client_lists))
        e.add_separator()
        e.add_command(label="Find (search box)", accelerator="Ctrl+F", command=c("edit_find", lambda: self.entry_search.focus_set()))
        e.add_command(label="Clear filters", command=c("edit_clear_filters", self.clear_filters))
        e.add_separator()
        e.add_command(label="Add selected to the omit list...", command=c("edit_omit_add", self.mark_known))
        e.add_command(label="Omit list...", command=c("edit_omit_list", self.open_omit_list))
        e.add_command(label="Accept context suggestions (rows that have one)", command=c("edit_accept_context", lambda: self.accept_context(list(self.results))))
        mb.add_cascade(label="Edit", menu=e)
        ef = tk.Menu(mb, tearoff=0, postcommand=lambda: self.show_effort_menu(ef))
        mb.add_cascade(label="Effort", menu=ef)
        cm = tk.Menu(mb, tearoff=0)
        cm.add_command(label="Compliance Matrix...", accelerator="Ctrl+M", command=c("compliance_open", self.open_compliance))
        mb.add_cascade(label="Compliance", menu=cm)
        t = tk.Menu(mb, tearoff=0, postcommand=lambda: self.show_tools_menu(t))
        mb.add_cascade(label="Tools", menu=t)
        v = tk.Menu(mb, tearoff=0)
        self.side_var = tk.BooleanVar(value=bool(self.settings.get("sidebar", True)))
        v.add_checkbutton(label="Sidebar", accelerator="Ctrl+B", variable=self.side_var, command=c("view_sidebar", lambda: self.toggle_sidebar(var=True)))
        v.add_checkbutton(label="Preview pane", variable=self.preview_var, command=c("view_preview", self.toggle_preview))
        if self.app_cfg["allow_advanced_toggle"]:
            v.add_checkbutton(label="Advanced features", variable=self.adv_var, command=c("view_advanced", self.toggle_advanced))
        v.add_separator()
        v.add_command(label="Larger text", accelerator="Ctrl++", command=c("view_text_larger", lambda: self.change_font_size(+1)))
        v.add_command(label="Smaller text", accelerator="Ctrl+-", command=c("view_text_smaller", lambda: self.change_font_size(-1)))
        v.add_command(label="Reset text size", accelerator="Ctrl+0", command=c("view_text_reset", lambda: self.change_font_size(0, reset=True)))
        mb.add_cascade(label="View", menu=v)
        h = tk.Menu(mb, tearoff=0)
        h.add_command(label="How it works (Guide)", accelerator="F1", command=c("help_guide", lambda: HG.open_guide(self)))
        h.add_command(label="Send feedback...", accelerator="Ctrl+Shift+F", command=c("help_feedback", lambda: FB.open_feedback(self)))
        h.add_separator()
        h.add_command(label="What runs where (AI and ML)...", command=c("help_what_runs", self.show_what_runs))
        h.add_command(label="About Acronym Hunter", command=c("help_about", self.show_about))
        mb.add_cascade(label="Help", menu=h)
        self.configure(menu=mb)
        self.menubar = mb
        keys = (("<Control-o>", self.select_files), ("<Control-O>", self.select_folder), ("<F5>", self.scan_again),
                ("<Control-e>", self.export_csv), ("<Control-b>", self.toggle_sidebar), ("<Control-m>", self.open_compliance),
                ("<Control-f>", lambda: self.entry_search.focus_set()))
        for seq, fn in keys:
            self.bind_all(seq, lambda _e, fn=fn: (fn(), "break")[1])

    def toggle_sidebar(self, save=True, var=False):
        """Hide or show the sidebar (View > Sidebar, Ctrl+B, or the Hide / Show buttons)."""
        shown = str(self.side) in [str(p) for p in self.body.panes()]
        want = bool(self.side_var.get()) if var and hasattr(self, "side_var") else not shown
        if want and not shown:
            self.body.add(self.side, before=self.main, width=self.side_w, minsize=int(190 * ctk.ScalingTracker.get_window_scaling(self)), stretch="never")
            self.btn_show_side.pack_forget(); self.hdr.pack_forget()
        elif not want and shown:
            self.body.forget(self.side)
            self.hdr.pack(fill="x", padx=15, pady=(8, 0), before=self.status)
            self.btn_show_side.pack(side="left")
        if hasattr(self, "side_var"):
            self.side_var.set(want)
        if save:
            if want:
                diag.activity("sidebar_shown")
            else:
                diag.activity("sidebar_hidden")
            self.settings["sidebar"] = want
            save_settings(self.settings)

    def _sidebar_resized(self, _e=None):
        try:
            if str(self.side) in [str(p) for p in self.body.panes()]:
                w = self.side.winfo_width()
                if w > 50 and abs(w - self.side_w) > 4:
                    self.side_w = w
                    self.settings["sidebar_w"] = w
                    save_settings(self.settings)
        except tk.TclError:
            diag.swallowed()

    def show_about(self):
        edition = "US Classic (Tesseract only)" if not DOCTR_DIR else "US (Tesseract + docTR second reader)"
        messagebox.showinfo("About Acronym Hunter",
                            f"Acronym Hunter {APP_VERSION}\nEdition: {edition}\n\nRuns fully offline on this PC. It makes no network "
                            "connections; picture reading runs in a no-network sandbox.\n\nFinds acronyms, abbreviations, units, symbols "
                            "and equations; keeps an effort list per proposal; checks volumes; builds compliance matrices.", parent=self)

    def toggle_preview(self):
        on = bool(self.preview_var.get())
        self.settings["preview"] = on
        save_settings(self.settings)
        if on:
            self.split.add(self.preview, minsize=int(380 * ctk.ScalingTracker.get_window_scaling(self)) + 40, width=int(400 * ctk.ScalingTracker.get_window_scaling(self)) + 40, stretch="never")
            if self.selected_term:
                self.preview.show(self.selected_term)
        else:
            self.split.forget(self.preview)

    def open_improve(self):
        """One window, one action: teach the reader from the saved corrections, test it, and apply it if it reads better."""
        diag.activity("improve_reading_opened")
        if not ai_ml_on():
            messagebox.showinfo("Improve Reading", "AI and ML are off for this effort. Improve Reading trains the picture reader "
                                "(ML), so it is off too. Turn AI and ML on in the top bar to use it.")
            return
        import tess_train
        if getattr(self, "improve_win", None) is not None and self.improve_win.winfo_exists():
            self.improve_win.lift(); self.improve_win.focus(); return
        win = ctk.CTkToplevel(self); self.improve_win = win
        win.title("Improve Reading"); win.geometry("720x640"); win.transient(self); win.after(100, win.lift)
        n, new = self.training_state()
        ctk.CTkLabel(win, text="Improve Reading: teach the reader your corrections", font=("Arial", 15, "bold")).pack(anchor="w", padx=16, pady=(14, 4))
        ctk.CTkLabel(win, justify="left", wraplength=680, anchor="w", text=(
            "1. You correct readings (preview pane: Correct + Teach, or the Trainer). Each one is saved as a corrected line.\n"
            "2. Teach and Apply trains a new reader from all saved lines, on this PC's processor, offline.\n"
            "3. It tests the new reader on lines it did not learn from, against the reader you use now.\n"
            "4. Only if it reads better, it becomes your reader right away. Otherwise nothing changes.")).pack(anchor="w", padx=16)
        st = ctk.CTkLabel(win, text=f"Saved corrected lines: {n} ({new} new since the last time)."
                          + ("" if n >= tess_train.MIN_SAMPLES else f" At least {tess_train.MIN_SAMPLES} are needed."),
                          font=("Arial", 13, "bold"), anchor="w")
        st.pack(anchor="w", padx=16, pady=(10, 4))
        btn = ctk.CTkButton(win, text="Teach and Apply", height=38, fg_color="#0F766E", font=("Segoe UI", 14, "bold"),
                            state="normal" if n >= tess_train.MIN_SAMPLES else "disabled")
        btn.pack(fill="x", padx=16, pady=(4, 6))
        bar = ctk.CTkProgressBar(win); bar.set(0); bar.pack(fill="x", padx=16, pady=(2, 2))
        stage = ctk.CTkLabel(win, text="", anchor="w", wraplength=680, justify="left"); stage.pack(anchor="w", padx=16)
        result = ctk.CTkTextbox(win, height=230, wrap="word", font=("Segoe UI", 13)); result.pack(fill="both", expand=True, padx=16, pady=(8, 6))
        try:
            last = json.load(open(os.path.join(TRAIN_DIR, "model", "training_report.json"), encoding="utf-8"))
            result.insert("1.0", f"Last time ({last.get('date', '?')}):\n" + tess_train.describe(last))
        except (OSError, ValueError):
            diag.swallowed()
            result.insert("1.0", "No reader trained yet. The scanner uses Google's English reader.")
        result.configure(state="disabled")
        foot = ctk.CTkFrame(win, fg_color="transparent"); foot.pack(fill="x", padx=16, pady=(0, 12))
        ctk.CTkButton(foot, text="Put back the previous reader", fg_color="#475569", width=210,
                      command=lambda: self._undo_reader(result)).pack(side="left")
        ctk.CTkButton(foot, text="Share with team...", fg_color="#B45309", width=150, command=self.open_knowledge).pack(side="left", padx=8)
        ctk.CTkButton(foot, text="Close", width=90, command=win.destroy).pack(side="right")
        ctk.CTkLabel(win, text="Share with team saves a Knowledge Pack: your trained reader plus the database, OCR fixes, and ignore list, "
                     "for colleagues to install. Approved SharePoint folder only.", text_color="gray", wraplength=680,
                     justify="left").pack(anchor="w", padx=16, pady=(0, 10))

        def say(msg, frac=None):
            def ui():
                if not win.winfo_exists(): return
                stage.configure(text=msg)
                if frac is None or (0.3 <= frac < 0.95):
                    if bar.cget("mode") != "indeterminate":
                        bar.configure(mode="indeterminate"); bar.start()
                else:
                    bar.stop(); bar.configure(mode="determinate"); bar.set(frac)
            self.after(0, ui)
        def run():
            out_dir = os.path.join(TRAIN_DIR, "model")
            current = os.path.join(MODELS_DIR, CUSTOM_TESS_MODEL + ".traineddata")
            btn.configure(state="disabled", text="Working... (you can keep using the scanner)")
            self._training = True; self.update_banner()
            def work():
                try:
                    r = tess_train.train(TESS_DIR, TRAIN_DIR, out_dir, progress=say, current_model=current,
                                         runner=lambda cmd, timeout=600: run_tesseract_tool(cmd, b"", timeout, write_dirs=[out_dir]))
                    self.after(0, lambda: done(r, None))
                except Exception as e:
                    log_error("Improve Reading failed")
                    msg = f"{e}"
                    self.after(0, lambda: done(None, msg))
            threading.Thread(target=work, daemon=True).start()
        def done(r, err):
            self._training = False
            self.update_banner()
            if not win.winfo_exists():
                return
            bar.stop(); bar.configure(mode="determinate"); bar.set(1.0 if r else 0)
            btn.configure(state="normal", text="Teach and Apply")
            result.configure(state="normal"); result.delete("1.0", "end")
            if err:
                stage.configure(text="Stopped. Your reader was not changed.")
                result.insert("1.0", err)
            elif r.get("better"):
                load_trained_models(os.path.dirname(r["model"]))
                stage.configure(text="Applied. The scanner now reads pictures with your improved reader.")
                lines = ["APPLIED: the new reader read your held-out lines better.", "", tess_train.describe(r)]
                if r.get("fixed_examples"):
                    lines += ["", "Now read correctly (was wrong before):"] + [f"  {t}    (was: {b})" for t, b in r["fixed_examples"]]
                if r.get("new_mistakes"):
                    lines += ["", "New mistakes to watch:"] + [f"  {t}    (now: {a})" for t, a in r["new_mistakes"]]
                lines += ["", "Next: rescan to see it in your results (only files with pictures are read again)."]
                result.insert("1.0", "\n".join(lines))
                if self.project is not None and not self.project.read_only:
                    if messagebox.askyesno("Rescan?", "Rescan the project now? Only files with pictures are read again.", parent=win):
                        self.rescan_project()
                elif self.queued_files and messagebox.askyesno("Rescan?", "Scan the same files again with the improved reader?", parent=win):
                    self.start_batch(self.queued_files)
            else:
                stage.configure(text="Not applied. Your current reader still reads these lines as well or better.")
                result.insert("1.0", "NOT APPLIED: the new reader did not beat the current one on held-out lines.\n\n"
                              + tess_train.describe(r) + "\n\nSave more corrected lines (varied pictures and fonts), then try again.")
            result.configure(state="disabled")
        btn.configure(command=run)

    def _undo_reader(self, box=None):
        what = restore_last_knowledge_backup()
        msg = f"Put back the reader and knowledge from {what}." if what else "Nothing to undo."
        if box is not None:
            box.configure(state="normal"); box.insert("end", "\n\n" + msg); box.configure(state="disabled")
        else:
            messagebox.showinfo("Undo", msg)

    def path_for(self, filename):
        """Full path of a scanned file, from this session's scans or the open project."""
        if filename in self.scanned_paths and os.path.exists(self.scanned_paths[filename]):
            return self.scanned_paths[filename]
        if self.project is not None:
            for rel in self.project.data.get("files", {}):
                if os.path.basename(rel.replace("/", os.sep)) == filename:
                    path = self.project.member_path(rel)
                    if os.path.exists(path):
                        return path
        return None

    def picture_locations(self, term):
        d = self.results.get(term, {})
        if d.get("manual") or d.get("file_error"):
            return []
        pat = re.compile(r"\((Image|Vector text|Vector graphic)\)|, (image|vector) |\(p\.\d+\)")
        return [loc for loc in d.get("locations", []) if pat.search(loc)]

    def on_right_click(self, event):
        iid = self.tree.identify_row(event.y)
        if not iid or iid not in self.results:
            return
        if iid not in self.tree.selection():
            self.tree.selection_set(iid)
        self.tree.focus(iid)
        menu = tk.Menu(self, tearoff=0)
        menu.add_command(label="Details...", command=self.show_details)
        pics_sel = [k for k in self.tree.selection() if self.results.get(k, {}).get("category") == "Equation (picture)"]
        if pics_sel:
            if math_available():
                menu.add_command(label=f"Read equation picture{'s' if len(pics_sel) > 1 else ''} with the math reader ({len(pics_sel)})",
                                 command=lambda: self.read_equations(pics_sel))
            else:
                menu.add_command(label="Read equation pictures (math add-on not installed)" if not MATH_DIR else
                                 "Read equation pictures (Gen AI: off for this effort)", state="disabled")
        if trainer_available():
            pics = self.picture_locations(iid)
            if len(pics) == 1:
                menu.add_command(label="Check in picture...", command=lambda: self.check_in_picture(iid, pics[0]))
            elif pics:
                sub = tk.Menu(menu, tearoff=0)
                for loc in pics[:25]:
                    sub.add_command(label=loc if len(loc) < 110 else loc[:107] + "...", command=lambda l=loc: self.check_in_picture(iid, l))
                menu.add_cascade(label="Check in picture", menu=sub)
            else:
                menu.add_command(label="Check in picture (found only in typed text)", state="disabled")
        menu.add_separator()
        dd = self.results.get(iid, {})
        nsel = len([t for t in self.tree.selection() if t in self.results]) or 1
        if self.effort is not None:            # Keyword Harvest (2.16.0)
            kd = dd.get("selected_definition", "")
            kw = kd if kd and not status_text(kd).startswith(("[", "MULTIPLE", "Rescan")) else iid
            menu.add_command(label=f"Add as keyword: {kw[:50]}...", command=lambda: self.kw_quick(kw, source=f"Result {iid}"))
        if dd.get("known"):
            menu.add_command(label=f"Take off the omit list ({nsel})", command=self.unmark_known)
        else:
            menu.add_command(label=f"Add to omit list: found, kept, left off the final list ({nsel})...", command=self.mark_known)
        ctx = (dd.get("context") or {})
        if ctx.get("definition") and status_kind(dd.get("status", "")) == "choice":
            menu.add_command(label=f"Accept context suggestion: {ctx['definition'][:60]}", command=lambda: self.accept_context([iid]))
        if dd.get("excluded_instances"):
            menu.add_command(label=f"Restore excluded instances ({len(dd['excluded_instances'])})",
                             command=lambda: self.restore_instances(iid))
        self.effort_context_menu(menu, iid)
        menu.tk_popup(event.x_root, event.y_root)

    def check_in_picture(self, term, location):
        name = _file_of(location)
        path = self.path_for(name)
        if not path:
            messagebox.showinfo("Check in picture", f"{name} is not where it was scanned. Open it from the Trainer instead."); return
        self.open_trainer(path, focus=(location, term))

    # ---------- advanced tools ----------
    def toggle_advanced(self):
        self.advanced_on = bool(self.adv_var.get())
        self.settings["advanced"] = self.advanced_on
        save_settings(self.settings)
        self.apply_advanced_visibility()

    def apply_advanced_visibility(self):
        """Advanced shows Fix OCR in the action bar and the OCR Fixes / Trainer items in the Tools menu."""
        self.btn_ocr.pack_forget()
        self.btn_improve_top.pack_forget()
        if self.advanced_on:
            self.btn_ocr.pack(side="left", padx=5, pady=(4, 8), after=self.btn_save)
        if self.advanced_on and trainer_available() and OCR_BACKEND == "tesseract":   # always in Tools and the banner
            self.btn_improve_top.pack(fill="x", padx=8, pady=2, after=self.btn_tools)

    def show_tools_menu(self, menu=None):
        m = menu or self.tools_menu
        m.delete(0, "end")
        ml = "normal" if ai_ml_on() else "disabled"
        if trainer_available() and OCR_BACKEND == "tesseract":
            m.add_command(label="Improve Reading (teach the reader your corrections) [ML]...", command=self.open_improve, state=ml)
        if trainer_available():
            m.add_command(label="Trainer (check a picture, box missed text) [ML]...", command=self.open_trainer, state=ml)
        if not TRAINING_APP:
            m.add_command(label="Send corrections to the training team...", command=self.send_corrections,
                          state="normal" if corrections_on() else "disabled")
        m.add_command(label=("Knowledge (share with your team, install, undo)..." if TRAINING_APP else
                             "Knowledge Packs (install the training team's reader, undo)..."), command=self.open_knowledge)
        m.add_command(label="Full Text (all typed and picture text, save as .txt)...", command=self.open_fulltext)
        m.add_command(label="Read all equation pictures (math add-on) [Gen AI]..." if MATH_DIR else "Read all equation pictures (math add-on not installed)",
                      command=lambda: self.read_equations([k for k, d in self.results.items() if d.get("category") == "Equation (picture)"]),
                      state="normal" if math_available() else "disabled")
        m.add_command(label="What runs where (AI and ML)...", command=self.show_what_runs)
        m.add_command(label="How it works (guide, F1)...", command=lambda: (diag.activity("guide_opened"), HG.open_guide(self)))
        m.add_command(label="Feedback: bug, improvement or idea (Ctrl+Shift+F)...", command=lambda: FB.open_feedback(self))
        m.add_separator()
        m.add_command(label="Edit Database CSV", command=lambda: self.open_file(DB_FILE))
        m.add_command(label="OCR Fixes...", command=self.open_fix_manager,
                      state="normal" if self.advanced_on else "disabled")
        m.add_separator()
        m.add_checkbutton(label="Show preview pane", onvalue=True, offvalue=False, variable=self.preview_var,
                          command=self.toggle_preview)
        if menu is None:
            m.tk_popup(self.btn_tools.winfo_rootx() + self.btn_tools.winfo_width(), self.btn_tools.winfo_rooty())

    # ---------- math reader (add-on) ----------
    def _picture_png(self, loc):
        """PNG bytes of the picture at a location label (same loaders as the preview pane)."""
        import pictures as tg
        path = self.path_for(_file_of(loc))
        if not path:
            return None
        srcs = tg.list_sources(path)
        k = tg.match_source([lbl for lbl, _ in srcs], loc)
        if k is None:
            return None
        img = srcs[k][1]()
        b = io.BytesIO(); img.save(b, "PNG")
        return b.getvalue()

    def read_equations(self, keys):
        """Read the pictures behind the selected equation-picture rows with the math reader, in the background."""
        diag.activity("math_reader_started")
        targets = [(k, loc) for k in keys for loc in self.results.get(k, {}).get("locations", [])]
        if not targets:
            messagebox.showinfo("Read equations", "No equation pictures selected."); return
        if getattr(self, "math_win", None) is not None and self.math_win.winfo_exists():
            self.math_win.focus(); return
        win = ctk.CTkToplevel(self); self.math_win = win
        win.title("Math reader"); win.geometry("520x210"); win.resizable(False, False); win.transient(self)
        ctk.CTkLabel(win, text=f"Reading {len(targets)} equation picture(s) with the math reader.", anchor="w").pack(fill="x", padx=14, pady=(14, 2))
        ctk.CTkLabel(win, text="Runs on this PC inside the no-network sandbox. About 5 s per small equation; whole figures take longer.",
                     anchor="w", text_color=("#475569", "#94A3B8"), wraplength=490, justify="left").pack(fill="x", padx=14)
        prog = ctk.CTkProgressBar(win, width=490); prog.set(0); prog.pack(padx=14, pady=10)
        stat = ctk.CTkLabel(win, text="Starting...", anchor="w"); stat.pack(fill="x", padx=14)
        cancel = threading.Event()
        btn = ctk.CTkButton(win, text="Stop", width=90, fg_color="#DC2626", command=cancel.set); btn.pack(pady=10)
        win.protocol("WM_DELETE_WINDOW", lambda: (cancel.set(), win.destroy()))
        m = build_matcher()
        done = {"n": 0, "eq": 0}
        ids = {f"{i}": t for i, t in enumerate(targets)}
        def ui(f):
            self.after(0, lambda: win.winfo_exists() and f())
        def on_result(pid, formulas, cached, sec):
            key, loc = ids[pid]
            done["n"] += 1; done["eq"] += len(formulas)
            self.after(0, lambda: self.apply_math_result(key, loc, formulas, m))
            ui(lambda: (prog.set(done["n"] / len(targets)),
                        stat.configure(text=f"{done['n']} of {len(targets)} read{' (saved result)' if cached else f' in {sec:.0f} s'}: "
                                            f"{len(formulas)} equation(s)")))
        def work():
            try:
                items = []
                for pid, (key, loc) in ids.items():
                    if cancel.is_set():
                        return
                    png = self._picture_png(loc)
                    if png:
                        items.append((pid, png))
                    else:
                        done["n"] += 1
                ui(lambda: stat.configure(text="Loading the math reader..."))
                run_math_reader(items, lambda t: ui(lambda: stat.configure(text=t)), on_result, cancel)
                ui(lambda: (prog.set(1), stat.configure(text=f"Done: {done['eq']} equation(s) from {done['n']} picture(s)."),
                            btn.configure(text="Close", fg_color="#475569", command=win.destroy)))
            except InterruptedError:
                ui(lambda: (stat.configure(text=f"Stopped. {done['n']} picture(s) read; results kept."),
                            btn.configure(text="Close", fg_color="#475569", command=win.destroy)))
            except Exception as e:
                log_error("Math reader failed")
                msg = f"{type(e).__name__}: {e}"
                ui(lambda: (stat.configure(text=msg[:160]), btn.configure(text="Close", fg_color="#475569", command=win.destroy)))
        threading.Thread(target=work, daemon=True).start()

    def apply_math_result(self, pic_key, loc, formulas, m):
        """Replace OCR's pieces of an equation picture with what the math reader read."""
        ocr_cats = ("Equation", "Symbol", "Unit (Candidate)", "Abbreviation (Candidate)", "Acronym (Candidate)")
        for term in list(self.results):
            d = self.results[term]
            if d.get("category") in ocr_cats and not d.get("in_db") and not d.get("manual") and loc in d.get("locations", []) \
                    and not d.get("math_reader"):
                d["locations"].remove(loc)
                d["count"] = max(0, d.get("count", 1) - 1)
                if not d["locations"]:
                    del self.results[term]
        read_loc = loc + " (math reader)"
        lines = []
        for f in formulas:
            lin, syms = EQ.latex_symbols(f)
            if not lin:
                continue
            record_equation(self.results, read_loc, lin + "\t" + "|".join(syms), m)
            key = lin if len(lin) <= 160 else lin[:157] + "..."
            if key in self.results:
                self.results[key]["math_reader"] = True
                self.results[key]["latex"] = f
            lines.append(lin)
        summarize_equations(self.results)
        pic = self.results.get(pic_key)
        if pic is not None:
            pic.setdefault("math_read", {})[loc] = lines
            n = sum(len(v) for v in pic["math_read"].values())
            left = [l for l in pic["locations"] if l not in pic["math_read"]]
            pic["status"] = (f"READ BY MATH READER: {n} equation(s)" if not left
                             else f"READ {len(pic['math_read'])} of {len(pic['locations'])} PICTURES")
            pic["selected_definition"] = "Math reader: " + ("; ".join(l for v in pic["math_read"].values() for l in v)[:300] or "no equation found")
        self.populate_table()

    # ---------- full text mode ----------
    def open_fulltext(self):
        diag.activity("full_text_opened")
        """Tools > Full Text: every word in the chosen files, typed and OCR, saved as .txt. Same reader and sandbox
        as the scan. Nothing leaves the computer."""
        if getattr(self, "ft_win", None) is not None and self.ft_win.winfo_exists():
            self.ft_win.focus(); return
        win = ctk.CTkToplevel(self); self.ft_win = win
        win.title("Full Text - all typed and picture text"); win.geometry("900x680")
        win.transient(self)
        st = {"files": [], "texts": {}, "busy": False}
        top = ctk.CTkFrame(win, fg_color="transparent"); top.pack(fill="x", padx=10, pady=(10, 4))
        ctk.CTkLabel(top, text="Pull out all the text: typed text exactly, picture text through the OCR reader, "
                     "equations in linear form.", anchor="w").pack(fill="x")
        ctk.CTkLabel(top, text="The text carries the same classification and markings as the source files. "
                     "Save and handle it the same way.", anchor="w", text_color="#B45309").pack(fill="x", pady=(2, 6))
        row = ctk.CTkFrame(top, fg_color="transparent"); row.pack(fill="x")
        lbl = ctk.CTkLabel(row, text="No files chosen.", anchor="w")
        def set_files(fps):
            st["files"] = [f for f in fps if os.path.splitext(f)[1].lower() in SUPPORTED_EXTS and not os.path.basename(f).startswith("~$")]
            lbl.configure(text=f"{len(st['files'])} file(s) chosen." if st["files"] else "No supported files chosen.")
            btn_go.configure(state="normal" if st["files"] else "disabled")
        def pick_files():
            fps = filedialog.askopenfilenames(parent=win, title="Choose files")
            if fps: set_files(list(fps))
        def pick_folder():
            d = filedialog.askdirectory(parent=win, title="Choose a folder")
            if d:
                set_files([os.path.join(r, f) for r, _d, fs in os.walk(d) for f in fs])
        def use_last():
            set_files(list((getattr(self, "last_scan", None) or {}).get("paths", [])))
        ctk.CTkButton(row, text="Choose Files", width=110, height=28, command=pick_files).pack(side="left", padx=(0, 4))
        ctk.CTkButton(row, text="Choose Folder", width=110, height=28, command=pick_folder).pack(side="left", padx=4)
        b_last = ctk.CTkButton(row, text="Files from last scan", width=140, height=28, command=use_last,
                               state="normal" if (getattr(self, "last_scan", None) or {}).get("paths") else "disabled")
        b_last.pack(side="left", padx=4)
        lbl.pack(side="left", padx=10)
        bar = ctk.CTkFrame(win, fg_color="transparent"); bar.pack(fill="x", padx=10, pady=4)
        btn_go = ctk.CTkButton(bar, text="Read All Text", width=130, height=30, fg_color="#0F766E", state="disabled")
        btn_go.pack(side="left")
        btn_stop = ctk.CTkButton(bar, text="Stop", width=70, height=30, fg_color="#DC2626", state="disabled",
                                 command=lambda: CANCEL.set())
        btn_stop.pack(side="left", padx=6)
        prog = ctk.CTkProgressBar(bar, width=260); prog.set(0); prog.pack(side="left", padx=10)
        stat = ctk.CTkLabel(bar, text="", anchor="w"); stat.pack(side="left", fill="x", expand=True)
        box = ctk.CTkTextbox(win, wrap="word", font=("Consolas", 12)); box.pack(fill="both", expand=True, padx=10, pady=4)
        bot = ctk.CTkFrame(win, fg_color="transparent"); bot.pack(fill="x", padx=10, pady=(4, 10))
        def save_one():
            path = filedialog.asksaveasfilename(parent=win, defaultextension=".txt", filetypes=[("Text", "*.txt")],
                                                initialfile="Full_Text.txt", title="Save all text as one file")
            if path:
                with open(path, "w", encoding="utf-8") as fh:
                    fh.write("\n".join(st["texts"][f] for f in st["files"] if f in st["texts"]))
                stat.configure(text=f"Saved {os.path.basename(path)}")
        def save_each():
            d = filedialog.askdirectory(parent=win, title="Folder for one .txt per file")
            if not d: return
            n = 0
            for f, t in st["texts"].items():
                out = os.path.join(d, os.path.splitext(os.path.basename(f))[0] + "_fulltext.txt")
                with open(out, "w", encoding="utf-8") as fh:
                    fh.write(t)
                n += 1
            stat.configure(text=f"Saved {n} file(s) to {d}")
        def copy_all():
            win.clipboard_clear(); win.clipboard_append(box.get("1.0", "end")); stat.configure(text="Copied.")
        b1 = ctk.CTkButton(bot, text="Save as one .txt", width=140, height=30, command=save_one, state="disabled")
        b2 = ctk.CTkButton(bot, text="Save one .txt per file", width=170, height=30, command=save_each, state="disabled")
        b3 = ctk.CTkButton(bot, text="Copy", width=80, height=30, command=copy_all, state="disabled")
        for b in (b1, b2, b3): b.pack(side="left", padx=(0, 6))
        ctk.CTkButton(bot, text="Close", width=80, height=30, fg_color="#475569",
                      command=lambda: (CANCEL.set() if st["busy"] else None, win.destroy())).pack(side="right")
        def done(msg):
            st["busy"] = False
            btn_go.configure(state="normal"); btn_stop.configure(state="disabled")
            for b in (b1, b2, b3): b.configure(state="normal" if st["texts"] else "disabled")
            stat.configure(text=msg)
        def add_text(t):
            if win.winfo_exists():
                box.insert("end", t + "\n"); box.see("end")
        def work(files):
            n = len(files)
            def cb(m):
                check_cancel()
                self.after(0, lambda: win.winfo_exists() and stat.configure(text=str(m)[:90]))
            try:
                for i, f in enumerate(files, 1):
                    check_cancel()
                    self.after(0, lambda i=i: win.winfo_exists() and prog.set((i - 1) / n))
                    t, _err = full_text_of(f, cb, i, n)
                    st["texts"][f] = t
                    self.after(0, lambda t=t: add_text(t))
                self.after(0, lambda: (prog.set(1), done(f"Done: {n} file(s).")))
            except ScanCancelled:
                self.after(0, lambda: done(f"Stopped. {len(st['texts'])} file(s) read."))
            except Exception as e:
                log_error("Full text failed")
                self.after(0, lambda: done(f"Error: {type(e).__name__}: {e}"))
        def go():
            if self.btn_cancel.winfo_ismapped():
                messagebox.showinfo("Busy", "Wait for the current scan to finish.", parent=win); return
            st["texts"].clear(); st["busy"] = True; box.delete("1.0", "end"); prog.set(0)
            CANCEL.clear()
            btn_go.configure(state="disabled"); btn_stop.configure(state="normal")
            for b in (b1, b2, b3): b.configure(state="disabled")
            threading.Thread(target=work, args=(list(st["files"]),), daemon=True).start()
        btn_go.configure(command=go)
        if (getattr(self, "last_scan", None) or {}).get("paths"):
            use_last()

    # ---------- knowledge packs ----------
    def open_knowledge(self):
        diag.activity("knowledge_opened")
        if getattr(self, "know_win", None) is not None and self.know_win.winfo_exists():
            self.know_win.focus(); return
        win = ctk.CTkToplevel(self); self.know_win = win
        win.title("Knowledge"); win.geometry("600x560"); win.resizable(False, False)
        win.transient(self); win.after(100, win.lift)
        ctk.CTkLabel(win, text="Knowledge Pack = your trained reader + database + OCR fixes + ignore list",
                     font=("Arial", 13, "bold")).pack(anchor="w", padx=16, pady=(14, 2))
        ctk.CTkLabel(win, text=("Improve Reading makes the reader better on this PC. A Knowledge Pack SHARES everything learned "
                     "with colleagues. Export it after Improve Reading." if TRAINING_APP else
                     "The training team builds a better reader from everyone's corrections and ships it as a Knowledge "
                     "Pack. Install it here; Undo puts the previous one back."), wraplength=560,
                     justify="left", text_color="gray").pack(anchor="w", padx=16, pady=(0, 4))
        lbl = ctk.CTkLabel(win, text="", justify="left", wraplength=560); lbl.pack(anchor="w", padx=16, pady=(0, 8))
        msg = ctk.CTkLabel(win, text="", justify="left", wraplength=560, text_color="#0F766E")
        def refresh(text=None):
            lbl.configure(text=knowledge_status() + f"\nBackups available for Undo: {len(list_knowledge_backups())}")
            if text is not None: msg.configure(text=text)
        def after_change(text):
            refresh(text)
            if self.project is not None and not self.project.read_only:
                if messagebox.askyesno("Rescan project?", "Knowledge changed. Rescan the open project now to apply it?", parent=win):
                    self.rescan_project()
        def install():
            path = filedialog.askopenfilename(parent=win, title="Install Knowledge Pack (from the approved SharePoint folder)",
                                              filetypes=[("Knowledge Pack", "*.zip")])
            if not path: return
            try:
                after_change(install_knowledge_pack(path))
            except PermissionError:
                messagebox.showerror("File in use", DB_LOCKED_MSG, parent=win)
            except Exception as e:
                log_error("Knowledge Pack install failed")
                messagebox.showerror("Not installed", f"{e}\n\nNothing was changed.", parent=win)
        def undo():
            what = restore_last_knowledge_backup()
            refresh(f"Restored the state from {what}." if what else "Nothing to undo.")
        ctk.CTkButton(win, text="Install Knowledge Pack...", command=install).pack(fill="x", padx=16, pady=4)
        ctk.CTkButton(win, text="Undo Last Knowledge Change", fg_color="#475569", command=undo).pack(fill="x", padx=16, pady=4)
        if TRAINING_APP:   # loading trained models and exporting packs belong to Acronym Hunter Training (2.15.0)
            ctk.CTkLabel(win, text="Share and load", font=("Arial", 12, "bold")).pack(anchor="w", padx=16, pady=(14, 2))
            def load_models():
                folder = filedialog.askdirectory(parent=win, title="Training Kit output folder (docTR: output\\models; or acronym.traineddata)")
                if not folder: return
                try:
                    found = load_trained_models(folder)
                    after_change("Loaded " + ", ".join(found) + ". Score with tools\\evaluate.py; Undo puts the old models back.")
                except Exception as e:
                    log_error("Load trained models failed")
                    messagebox.showerror("Not loaded", f"{e}\n\nNothing was changed.", parent=win)
            def export():
                notes = notes_box.get().strip()
                name = f"AcronymHunter_Knowledge_{datetime.datetime.now():%Y.%m.%d-%H%M}.zip"
                path = filedialog.asksaveasfilename(parent=win, title="Save Knowledge Pack to the APPROVED SHAREPOINT folder",
                                                    initialfile=name, defaultextension=".zip", filetypes=[("Knowledge Pack", "*.zip")])
                if not path: return
                info = make_knowledge_pack(path, notes)
                refresh(f"Saved {os.path.basename(path)} (version {info['version']}, models: {', '.join(info['models']) or 'none'}).")
            ctk.CTkButton(win, text="Load Trained Models...", fg_color="#7C3AED", command=load_models).pack(fill="x", padx=16, pady=4)
            notes_box = ctk.CTkEntry(win, placeholder_text="Notes for this pack (what changed, score)")
            notes_box.pack(fill="x", padx=16, pady=(8, 2))
            ctk.CTkButton(win, text="Export Knowledge Pack...", fg_color="#B45309", command=export).pack(fill="x", padx=16, pady=4)
            ctk.CTkLabel(win, text="Knowledge Packs contain what was learned from internal documents. Save them ONLY to the "
                         "approved SharePoint folder. Never GitHub, email, or personal storage.",
                         text_color="#B45309", wraplength=560, justify="left").pack(anchor="w", padx=16, pady=(4, 0))
        msg.pack(anchor="w", padx=16, pady=(10, 0))
        refresh()

    # ---------- manual items ----------
    def open_add_item(self):
        """Add a term, or another meaning of a term already found. One acronym may have several meanings; the same
        meaning twice is flagged instead of added (2.13.0)."""
        diag.activity("add_item_opened")
        if getattr(self, "add_win", None) is not None and self.add_win.winfo_exists():
            self.add_win.focus(); return self.add_win
        win = ctk.CTkToplevel(self); self.add_win = win
        win.title("Add an item"); win.geometry("580x560"); win.resizable(False, False)
        win.transient(self); win.after(100, win.lift)
        pad = {"padx": 16, "pady": (8, 0), "anchor": "w"}
        ctk.CTkLabel(win, text="Term (required). A term already listed gets the definition as another meaning.").pack(**pad)
        e_term = ctk.CTkEntry(win, width=540); e_term.pack(padx=16, pady=2); SP.attach(e_term)
        known = ctk.CTkLabel(win, text="", anchor="w", justify="left", wraplength=540, text_color=("#475569", "#94A3B8"))
        known.pack(padx=16, anchor="w")
        ctk.CTkLabel(win, text="Type").pack(**pad)
        c_cat = ctk.CTkComboBox(win, values=MANUAL_CATEGORIES, width=540); c_cat.set("Acronym"); c_cat.pack(padx=16, pady=2)
        ctk.CTkLabel(win, text="Definition").pack(**pad)
        e_def = ctk.CTkEntry(win, width=540); e_def.pack(padx=16, pady=2); SP.attach(e_def)
        SP.button(win, e_def, text="Ω  Symbols (for the field you clicked last)", width=300).pack(padx=16, pady=(4, 0), anchor="w")
        ctk.CTkLabel(win, text="Where found (optional, e.g. 'Volume II, p.14')").pack(**pad)
        e_loc = ctk.CTkEntry(win, width=540); e_loc.pack(padx=16, pady=2)
        ctk.CTkLabel(win, text="Save it").pack(**pad)
        dest = tk.StringVar(value="results")
        df = ctk.CTkFrame(win, fg_color="transparent"); df.pack(padx=16, anchor="w")
        for val, lab in (("results", "In these results only"), ("db", "Database: future scans find it"),
                         ("omit", "Omit list: found, left off the final lists")):
            ctk.CTkRadioButton(df, text=lab, variable=dest, value=val).pack(anchor="w", pady=1)
        why = ctk.CTkComboBox(win, values=OMIT_REASONS, width=260); why.set(OMIT_REASONS[0])
        ctk.CTkLabel(win, text="Why omitted (omit list)").pack(**pad)
        why.pack(padx=16, pady=2, anchor="w")
        save_var = dest                                       # older tests read _add_api["save"]
        msg = ctk.CTkLabel(win, text="", text_color="#DC2626", wraplength=540, justify="left"); msg.pack(padx=16, pady=(6, 0), anchor="w")

        def show_known(*_):
            t = fold(e_term.get().strip())
            have = []
            if t in self.results:
                have += [x for x in self.results[t].get("definitions", []) if x]
            have += [dd for dd, a in db_definitions(t) if dd and dd not in have and _action_kind(a) != "exclude"]
            known.configure(text=("Already known: " + " | ".join(have[:6])) if have else "")
        e_term.bind("<KeyRelease>", show_known)

        def add():
            term, cat = fold(e_term.get().strip()), (c_cat.get().strip() or "Acronym")
            defn, loc, where = e_def.get().strip(), e_loc.get().strip(), dest.get()
            if not term:
                msg.configure(text="Enter a term."); return
            if where == "db" and not defn:
                msg.configure(text="A definition is needed to save it to the database."); return
            d = self.results.get(term)
            have = [x for x in (d.get("definitions", []) if d else []) if x]
            if d is not None and not defn and where != "omit":
                msg.configure(text=f"'{term}' is already listed. Type a different definition to add another meaning."); return
            if defn and any(norm_def(defn) == norm_def(x) for x in have):
                msg.configure(text=f"'{term}' already has the meaning \"{defn}\". Nothing added. A different meaning is fine."); return
            res = None
            if where == "db":
                res, _r = db_add(term, cat, defn, "Include")
            elif where == "omit":
                res, _r = db_add(term, cat, defn, KNOWN_ACTION, {OMIT_REASON_COL: why.get()})
            if res == "duplicate" and d is None:
                msg.configure(text=f"The database already has '{term}' = \"{defn}\"" + (" on the omit list" if where == "omit" else "")
                                   + ". Nothing added. A different meaning is fine."); return
            if d is not None:                                  # another meaning for a term already listed
                if defn:
                    d.setdefault("definitions", []).append(defn)
                if where == "omit":
                    d.update(known=True, in_db=True, excluded=False, omit_reason=why.get(), selected_definition=defn or d.get("selected_definition"),
                             status=KNOWN_STATUS + f": {why.get()}")
                elif len([x for x in d["definitions"] if x]) > 1 and status_kind(d.get("status", "")) != "choice":
                    defs = [x for x in d["definitions"] if x]
                    d["in_db"] = d.get("in_db") or where == "db"
                    d["selected_definition"] = f"MULTIPLE ({len(defs)} options) - Choose: " + " OR ".join(defs)
                    d["status"] = "CHOOSE DEFINITION"
                diag.activity("add_item_meaning")
                self.selected_term = term
                self.populate_table()
                msg.configure(text=f"Added another meaning to {term}" + (" and saved it to the database." if where == "db" else "."),
                              text_color="#047857")
                e_def.delete(0, "end"); show_known()
                return
            item = {"category": cat, "count": 1, "locations": [loc or "Added manually"],
                    "definitions": [defn] if defn else [], "in_db": True, "excluded": False, "manual": True,
                    "selected_definition": defn or "[No definition given]",
                    "status": "Added manually" + (" (saved to DB)" if where == "db" else "")}
            if where == "omit":
                item.update(known=True, omit_reason=why.get(), status=KNOWN_STATUS + f": {why.get()}",
                            selected_definition=defn or KNOWN_BLANK)
            self.manual_items[term] = item
            self.results[term] = item
            self.save_project()
            self.selected_term = term
            self.refresh_type_options(); self.populate_table()
            self.btn_export.configure(state="normal")
            win.destroy()

        btns = ctk.CTkFrame(win, fg_color="transparent"); btns.pack(fill="x", padx=16, pady=14)
        ctk.CTkButton(btns, text="Add", command=add).pack(side="left")
        ctk.CTkButton(btns, text="Close", fg_color="#475569", command=win.destroy).pack(side="right")
        e_term.focus_set()
        win.bind("<Return>", lambda _e: add())
        win._add_api = {"term": e_term, "cat": c_cat, "defn": e_def, "loc": e_loc, "save": save_var, "dest": dest, "why": why,
                        "add": add, "msg": msg, "known": known}
        return win

    def remove_manual_item(self):
        term = self.selected_term
        if not term or not self.results.get(term, {}).get("manual"): return
        self.manual_items.pop(term, None); self.results.pop(term, None)
        self.save_project()
        self.selected_term = None
        self.btn_exc.configure(text="Exclude", command=self.exclude_term, state="disabled")
        self.lbl_act.configure(text="Select a row above (Ctrl+click or Shift+click for several):")
        self.populate_table()

    def _on_tk_error(self, exc_type, exc, tb):
        _excepthook(exc_type, exc, tb)
        if isinstance(exc, PermissionError):
            messagebox.showerror("Database locked", DB_LOCKED_MSG)
        else:
            messagebox.showerror("Error", f"{exc_type.__name__}: {exc}\n\nDetails were saved to the logs folder (no file names or document content).")

    def open_trainer(self, path=None, focus=None):
        """Open the Trainer as a window of this program (same engine, same sandbox, same settings)."""
        diag.activity("trainer_opened")
        if not ai_ml_on():
            messagebox.showinfo("Trainer", "AI and ML are off for this effort. The Trainer reads pictures (ML), so it is off "
                                "too. Turn AI and ML on in the top bar to use it.")
            return
        if not trainer_available():
            messagebox.showinfo("Trainer", "Training features are turned off (config\\app_settings.json: allow_training)."); return
        import trainer_gui
        w = getattr(self, "trainer_win", None)
        if w is not None and w.winfo_exists():
            w.deiconify(); w.lift(); w.focus()
            if path and focus:
                w.focus_on(path, *focus)
            elif path:
                w.load_file(path)
            return
        self.trainer_win = trainer_gui.TrainerWindow(self, start_file=path, focus=focus)

    def open_file(self, path):
        if sys.platform == "win32": os.startfile(path)
        elif sys.platform == "darwin": subprocess.call(["open", path])
        else: subprocess.call(["xdg-open", path])

    def start_dir(self):
        last = self.settings.get("last_dir", "")
        if last and os.path.isdir(last):
            return last
        docs = os.path.join(os.path.expanduser("~"), "Documents")
        return docs if os.path.isdir(docs) else os.path.expanduser("~")

    def remember_dir(self, path):
        folder = path if os.path.isdir(path) else os.path.dirname(path)
        if folder:
            self.settings["last_dir"] = folder
            save_settings(self.settings)

    def select_files(self):
        fps = filedialog.askopenfilenames(title="Select Files", initialdir=self.start_dir(), filetypes=[("Supported", "*.pdf *.docx *.xlsx *.pptx *.png *.jpg *.jpeg *.bmp *.tif *.tiff *.webp *.gif *.svg *.svgz *.emf *.wmf *.emz *.wmz")])
        if fps:
            self.remember_dir(fps[0])
            self.start_batch(list(fps))

    def select_folder(self):
        folder = filedialog.askdirectory(title="Select Folder", initialdir=self.start_dir())
        if not folder: return
        self.remember_dir(folder)
        collected = []
        for root, _, files in os.walk(folder):
            for f in files:
                if not f.startswith("~$") and os.path.splitext(f)[1].lower() in SUPPORTED_EXTS:
                    collected.append(os.path.join(root, f))
        if not collected:
            messagebox.showwarning("Empty", "No supported files found.")
            return
        self.start_batch(collected, folder=folder)

    def scan_again(self):
        """Scan the same files again (picks up new choices, definitions, OCR fixes). A folder is re-listed, so new files count."""
        diag.activity("scan_again")
        if self.project is not None:
            self.rescan_project(); return
        last = getattr(self, "last_scan", None)
        if not last:
            return
        folder = last.get("folder")
        if folder and os.path.isdir(folder):
            fps = [os.path.join(r, f) for r, _, fs in os.walk(folder) for f in fs
                   if not f.startswith("~$") and os.path.splitext(f)[1].lower() in SUPPORTED_EXTS]
        else:
            fps = [f for f in last["paths"] if os.path.isfile(f)]
        if not fps:
            messagebox.showwarning("Scan Again", "The files from the last scan are no longer there.")
            return
        self.start_batch(fps, folder=folder)

    def start_batch(self, fps, folder=None, after=None):
        """after(status_callback): extra work on the worker thread once the files are read (volume check)."""
        self.queued_files = fps
        self.scanned_paths.update({os.path.basename(f): f for f in fps})
        capture = {}
        self.last_scan = {"paths": list(fps), "folder": folder, "cache": capture}
        diag.scan_started(len(fps), dict(collections.Counter(os.path.splitext(f)[1].lower() for f in fps)))
        def job(cb):
            res = analyze_batch(fps, cb, capture=capture)
            if after:
                after(cb)
            return res
        self.run_job(job)

    def run_job(self, job):
        """Run job(status_callback) -> results on a worker thread, then show them."""
        for b in (self.btn_files, self.btn_folder, self.btn_export, self.btn_project, self.btn_again):
            b.configure(state="disabled")
        self.prog.configure(mode="indeterminate")
        self.prog.start()
        self._prev_results = self.results
        CANCEL.clear()
        self.btn_cancel.configure(state="normal", text="Cancel Scan")
        self.btn_cancel.pack(pady=(0, 8), before=self.prog)
        for i in self.tree.get_children(): self.tree.delete(i)
        threading.Thread(target=self.worker, args=(job,), daemon=True).start()

    def cancel_scan(self):
        CANCEL.set()
        self.btn_cancel.configure(state="disabled", text="Stopping...")
        self.status.configure(text="Stopping after the current page...")

    def worker(self, job):
        def cb(m):
            check_cancel()   # stops mid-file at the next page, slide, or picture
            self.after(0, lambda: self.status.configure(text=m))
        try:
            results = job(cb)
            self.after(0, lambda: self.scan_done(results))
        except ScanCancelled:
            self.after(0, self.scan_cancelled)
        except Exception as e:
            log_error("Scan failed")
            diag.scan_finished(errors=1)
            msg = f"{type(e).__name__}: {e}\n\nDetails were saved to the logs folder (no file names or document content)."
            self.after(0, lambda: messagebox.showerror("Scan failed", msg))
            self.after(0, self.reset_ui)

    # ---------- projects ----------
    def show_project_menu(self, menu=None):
        m, has, rw = menu or self.project_menu, self.project is not None, self.project is not None and not self.project.read_only
        m.delete(0, "end")
        m.add_command(label="New Project from a Folder...", command=self.new_project)
        m.add_command(label="Save Current Results as Project...", command=self.save_results_as_project,
                      state="normal" if (not has and getattr(self, "last_scan", None) and self.results) else "disabled")
        m.add_command(label="Open Project...", command=self.open_project)
        m.add_separator()
        m.add_command(label="Scan / Rescan Project (changed files only)", command=self.rescan_project, state="normal" if rw else "disabled")
        m.add_command(label="Add Files to Project...", command=self.add_files_to_project,
                      state="normal" if rw and self.project.data.get("mode") == "files" else "disabled")
        m.add_command(label="Full Rescan (read every file again)", command=lambda: self.rescan_project(full=True), state="normal" if rw else "disabled")
        m.add_command(label='Clear "No longer found" items', command=self.clear_gone, state="normal" if rw else "disabled")
        m.add_separator()
        m.add_command(label="Close Project", command=self.close_project, state="normal" if has else "disabled")
        if menu is None:
            b = self.btn_project
            m.tk_popup(b.winfo_rootx() + b.winfo_width(), b.winfo_rooty())

    def new_project(self):
        folder = filedialog.askdirectory(title="Choose the folder that holds the project's documents", initialdir=self.start_dir())
        if not folder: return
        self.remember_dir(folder)
        name = os.path.basename(os.path.normpath(folder)) or "Project"
        path = os.path.join(folder, name + PROJECT_EXT)
        if os.path.exists(path):
            if messagebox.askyesno("Project exists", f"This folder already has a project:\n{path}\n\nOpen it?"):
                self.open_project(path)
            return
        if not self.close_project(): return
        proj = Project(path); proj.data = Project.new_data()
        try:
            proj.acquire_lock(); proj.save()
        except OSError as e:
            messagebox.showerror("Cannot create project", f"Could not write to this folder:\n{e}"); return
        self.project = proj
        self.manual_items = proj.data["manual"]
        self.set_client_lists(proj.data.setdefault("client_lists", []))
        self.results = {}; self.populate_table(); self.lbl_showing.configure(text="")
        self.update_project_ui()
        self.status.configure(text=f"Project '{proj.name}' created. Click Scan Project when you are ready.")
        messagebox.showinfo("Project created", f"Project file:\n{path}\n\nNothing has been scanned yet. Click Scan Project when you "
                            "are ready. The first scan reads every file; later scans read only new or changed files.\n\n"
                            "It stores the terms found, your choices and the scan history, never sentences or document text.")

    def save_results_as_project(self):
        """Turn the current one-off scan into a project, reusing what was just read: nothing is scanned again."""
        ls = getattr(self, "last_scan", None)
        if not ls or not self.results or self.project is not None:
            return
        paths = ls["paths"]
        folder = ls["folder"]
        start = folder or os.path.dirname(os.path.abspath(paths[0]))
        stem = os.path.basename(os.path.normpath(folder)) if folder else os.path.splitext(os.path.basename(paths[0]))[0]
        path = filedialog.asksaveasfilename(title="Save results as a project (the project file goes with the documents)",
                                            initialdir=start, initialfile=stem + PROJECT_EXT, defaultextension=PROJECT_EXT,
                                            filetypes=[("Acronym Hunter project", "*" + PROJECT_EXT)])
        if not path:
            return
        proj = Project(path); proj.data = Project.new_data()
        whole_folder = bool(folder) and os.path.normcase(os.path.abspath(folder)) == os.path.normcase(proj.root)
        proj.data["mode"] = "folder" if whole_folder else "files"
        members = [proj.to_member(p) for p in paths]
        if not whole_folder:
            proj.data["members"] = members
        proj.data["files"] = {proj.to_member(p): e for p, e in ls["cache"].items()}
        proj.data["manual"] = self.manual_items
        proj.data["client_lists"] = self.client_lists
        proj.data["decisions"] = dict(getattr(self, "session_decisions", {}))
        apply_project_history(proj, {k: dict(v) for k, v in self.results.items() if not v.get("manual")}, update=True)
        proj.data["last_stats"] = {"files": len(paths), "scanned": len(paths), "reused": 0, "removed": 0}
        try:
            proj.acquire_lock(); proj.save()
        except OSError as e:
            messagebox.showerror("Cannot save project", f"{e}"); return
        self.project, self.project_stats = proj, proj.data["last_stats"]
        apply_project_history(proj, self.results, update=False)
        self.update_project_ui(); self.populate_table()
        kind = "the whole folder" if whole_folder else f"{len(members)} file(s)"
        messagebox.showinfo("Project saved", f"Saved as {os.path.basename(path)} ({kind}).\n\nNothing was scanned again. Future "
                            "rescans reuse these results and read only files that are new or changed."
                            + ("" if whole_folder else "\nProject > Add Files to Project adds more files later."))

    def add_files_to_project(self):
        proj = self.project
        if proj is None or proj.read_only or proj.data.get("mode") != "files":
            return
        fps = filedialog.askopenfilenames(title="Add files to the project", initialdir=proj.root,
                                          filetypes=[("Supported", " ".join("*" + e for e in sorted(SUPPORTED_EXTS)))])
        if not fps:
            return
        have = set(proj.data["members"])
        new = [m for m in (proj.to_member(f) for f in fps) if m not in have]
        proj.data["members"] += new
        self.save_project()
        if new:
            self.rescan_project()   # reads only the new files

    def open_project(self, path=None):
        path = path or filedialog.askopenfilename(title="Open Project", initialdir=self.start_dir(),
                                                  filetypes=[("Acronym Hunter project", "*" + PROJECT_EXT)])
        if not path: return
        if not self.close_project(): return
        proj = Project(path)
        try:
            proj.load()
        except Exception as e:
            diag.swallowed()
            messagebox.showerror("Cannot open project", f"{type(e).__name__}: {e}"); return
        if not proj.acquire_lock():
            who = proj.locked_by or {}
            ans = messagebox.askyesnocancel("Project in use",
                f"{who.get('user', '?')} on {who.get('computer', '?')} has had this project open since {who.get('time', '?')}.\n\n"
                "Yes: open read-only (view and export only).\nNo: take over editing (only if they have finished).\nCancel: do not open.")
            if ans is None: return
            if ans: proj.read_only = True
            else: proj.acquire_lock(force=True)
        self.project = proj
        self.remember_dir(proj.root)
        self.manual_items = proj.data["manual"]
        self.set_client_lists(proj.data.setdefault("client_lists", []))
        self.project_stats = proj.data.get("last_stats", {})
        def job(cb):
            cb(f"Opening project '{proj.name}' from saved results...")
            res, _ = project_findings(proj, cb, rescan=False)
            apply_project_history(proj, res, update=False)
            return res
        self.run_job(job)

    def rescan_project(self, full=False):
        proj = self.project
        if proj is None or proj.read_only: return
        def job(cb):
            res, stats = project_findings(proj, cb, full=full)
            apply_project_history(proj, res, update=True)
            proj.data["last_stats"] = stats
            self.project_stats = stats
            return res
        diag.scan_started(len(proj.data.get("members", []) or []), {})
        self.run_job(job)

    def clear_gone(self):
        proj = self.project
        if proj is None or proj.read_only: return
        gone = [t for t, h in proj.data["history"].items() if h.get("state") == "gone"]
        if not gone:
            messagebox.showinfo("Nothing to clear", 'No "No longer found" items.'); return
        if not messagebox.askyesno("Clear items", f'Remove {len(gone)} "No longer found" item(s) from this project?'): return
        for t in gone:
            proj.data["history"].pop(t, None)
            if self.results.get(t, {}).get("gone"): self.results.pop(t, None)
        self.save_project(); self.populate_table()

    def close_project(self):
        """Close the open project (saving it). Returns False if the user cancels."""
        if self.project is None: return True
        if str(self.btn_project.cget("state")) == "disabled":
            messagebox.showwarning("Busy", "Wait for the scan to finish first."); return False
        self.save_project()
        self._release_project_lock()
        self.project, self.project_stats, self.manual_items, self.results = None, {}, {}, {}
        self.set_client_lists([])
        self.selected_term = None
        self.populate_table(); self.lbl_showing.configure(text="")
        self.status.configure(text="Project closed. Ready.")
        self.update_project_ui()
        return True

    def save_project(self):
        proj = self.project
        if proj is None or proj.read_only: return
        try:
            proj.save(); proj.acquire_lock(force=True)   # refresh lock time
        except OSError as e:
            messagebox.showerror("Project not saved", f"Could not save the project file:\n{e}")

    def _release_project_lock(self):
        if self.project is not None: self.project.release_lock()

    def on_close(self):
        try:
            self.kw_stop_hotkey()              # Ctrl+Alt+K is released with the window (2.16.0)
        except Exception:
            diag.swallowed()
        if self.project is not None and str(self.btn_project.cget("state")) != "disabled":
            self.save_project()
        self._release_project_lock()
        try:
            if self.effort is not None and self.effort.pending:
                self._save_effort_list()
            self.leave_workspace()          # session, where you were, compliance matrix; releases the effort
        except Exception:
            diag.swallowed()
        self.destroy()

    def update_project_ui(self):
        proj = self.project
        if proj is None:
            self.lbl_project.pack_forget()
            self.btn_folder.configure(text="Scan Folder...", command=self.select_folder)
            self.btn_files.configure(state="normal")
            self.btn_project.configure(text="Project")
            return
        st = self.project_stats or {}
        bits = [f"Project: {proj.name}", proj.root]
        if proj.data.get("last_scan"): bits.append(f"last scan {proj.data['last_scan']}")
        if st: bits.append(f"{st.get('files', 0)} files ({st.get('scanned', 0)} read, {st.get('reused', 0)} unchanged, {st.get('removed', 0)} removed)")
        if proj.read_only:
            who = proj.locked_by or {}
            bits.append(f"READ-ONLY (in use by {who.get('user', '?')})")
        self.lbl_project.configure(text="  |  ".join(bits))
        self.lbl_project.pack(pady=(0, 2), before=self.status)
        self.btn_project.configure(text=f"Project: {proj.name[:18]}")
        self.btn_folder.configure(text="Rescan Project" if proj.data.get("last_scan") else "Scan Project", command=self.rescan_project,
                                  state="disabled" if proj.read_only else "normal")
        self.btn_files.configure(state="disabled")

    def apply_saved_decisions(self):
        """Choices made with Apply Choice: this PC's saved choices, then the open project's (the project wins)."""
        decisions = dict(self.load_choices())
        if self.project is not None:
            decisions.update(self.project.data["decisions"])
        for term, choice in decisions.items():
            d = self.results.get(term)
            if d and (d.get("in_db") or d.get("customer_definitions")) and not d.get("excluded") and not d.get("gone") and \
                    (choice == "Keep All Possibilities" or choice in d.get("definitions", [])):
                self._set_choice(d, choice)

    def load_choices(self):
        """Apply Choice decisions for the open effort only (2.11.0), or for scans with no effort."""
        try:
            return self.ws.load_choices()
        except Exception:
            diag.swallowed()
            return {}

    def save_choice(self, term, choice):
        """Remember a definition choice for every later scan in this effort (written atomically)."""
        self.ws.save_choice(term, choice)

    @staticmethod
    def _set_choice(data, choice):
        if choice == "Keep All Possibilities":
            data["selected_definition"] = "MULTIPLE: " + " | ".join(data["definitions"])
            data["status"] = "MULTIPLE (Kept All)"
            return "ambiguous"
        for tag in (CLIENT_TAG, DB_TAG):
            if choice.startswith(tag):
                choice = choice[len(tag):]
        data["selected_definition"] = choice
        data["status"] = "User Selected" + (" (client list)" if data.get("client") and choice in data.get("client_definitions", []) else "")
        return "resolved"

    # ---------- view: text size ----------
    def apply_font_size(self):
        size = self.font_size
        dpi = getattr(self, "dpi", 1.0)
        self.tree_font.configure(size=int(round(size * dpi)))
        self.head_font.configure(size=int(round(size * dpi)))
        row_h = self.tree_font.metrics("linespace") + int(max(8, size // 2) * dpi)
        self.style.configure("Treeview", font=self.tree_font, rowheight=row_h)
        self.style.configure("Treeview.Heading", font=self.head_font, padding=(6, int(max(4, size // 3) * dpi)))
        scale = size / 11.0 * dpi
        for cid, _label, width, _anchor in COLUMNS:
            self.tree.column(cid, width=int(width * scale))
        self.lbl_size.configure(text=f"{size} pt")
        if self.results:
            self.populate_table()   # rewrap Locations for the new text size

    def change_font_size(self, step, reset=False):
        new = FONT_DEFAULT if reset else max(FONT_MIN, min(FONT_MAX, self.font_size + step))
        if new != self.font_size or reset:
            self.font_size = new
            self.apply_font_size()
            self.settings["font_size"] = new
            save_settings(self.settings)
        return "break"

    # ---------- view: sorting ----------
    def update_headings(self):
        for cid, label, _w, _a in COLUMNS:
            arrow = (" \u25BC" if self.sort_desc else " \u25B2") if cid == self.sort_col else ""
            self.tree.heading(cid, text=label + arrow, command=lambda c=cid: self.on_heading_click(c))

    def on_heading_click(self, col):
        if self.sort_col == col:
            self.sort_desc = not self.sort_desc
        else:
            self.sort_col = col
            self.sort_desc = col in ("count", "locations")  # numbers start high-to-low, text A-Z
        self.update_headings()
        self.populate_table()

    def sort_key(self, term, data):
        c = self.sort_col
        if c == "count":
            return (data["count"], term.casefold())
        if c == "locations":
            return (len(data["locations"]), term.casefold())
        if c == "category":
            return (data["category"].casefold(), term.casefold())
        if c == "status":
            rank = STATUS_RANK.get(status_kind(data["status"]), 9)
            return (rank, term.casefold())
        if c == "definition":
            return (data["selected_definition"].casefold(), term.casefold())
        if c == "firstuse":
            return (0 if first_use_flag(data) else 1 if first_use_lines(data) else 2, term.casefold())
        if c == "knownin":
            return (self.known_cell(term, data), term.casefold())
        if c == "change":
            order = {"New": 0, "Changed": 1, "Database update": 1, "No longer found": 2, "Unchanged": 3}
            return (order.get(data.get("change", ""), 4), term.casefold())
        return (term.casefold(),)

    # ---------- view: filtering ----------
    def clear_filters(self):
        self.type_var.set(ALL_TYPES)
        self.status_var.set(STATUS_FILTERS[0])
        self.search_var.set("")   # triggers populate_table
        self.populate_table()

    def refresh_type_options(self):
        types = sorted({base_type(d["category"]) for d in self.results.values()})
        values = [ALL_TYPES] + types
        self.opt_type.configure(values=values)
        if self.type_var.get() not in values:
            self.type_var.set(ALL_TYPES)

    @staticmethod
    def status_matches(data, wanted):
        st = data["status"]
        kind = status_kind(st)
        return {
            "Needs choice": kind == "choice",
            "New candidates": kind == "candidate",
            "Resolved": kind == "resolved" and status_text(st).startswith(RESOLVED_WORDS),
            "Client list": data.get("client", False),
            "Not defined at first use": first_use_flag(data),
            "Equations and symbols": data.get("category", "").startswith(("Equation", "Symbol")) or bool(data.get("symbol")),
            "Undefined symbols": bool(data.get("symbol")) and kind == "candidate" or (data.get("equation") and "not defined" in st),
            "Client and database differ": bool(data.get("client") and data.get("db_definitions")),
            "Added manually": data.get("manual", False),
            "In effort list": bool(data.get("effort")),
            "Customer documents": bool(data.get("customer_definitions")),
            "Customer and database differ": bool(data.get("cust_db_differ")),
            "Not in effort list": not data.get("effort") and not data.get("excluded", False) and not data.get("known"),
            "Found in omit list": bool(data.get("known")),
            "New in this scan (not in the effort list)": bool(data.get("new_in_effort")) and not data.get("effort"),
            "Context suggests a meaning": bool((data.get("context") or {}).get("definition")) and kind == "choice",
            "Effort: needs author": data.get("effort") in ("Proposed", "Needs author"),
            "Effort: approved": data.get("effort") == "Approved",
            "New since last scan": data.get("change") == "New",
            "Changed since last scan": data.get("change") == "Changed",
            "No longer found": data.get("gone", False),
            "Errors / not scanned": kind == "error",
            "Excluded": data.get("excluded", False),
            "Database update": data.get("change") == "Database update",
            "Known in: customer docs": "Customer docs" in (data.get("known_in") or []),
            "Known in: customer list": "Customer list" in (data.get("known_in") or []),
            "Known in: master database only": (data.get("known_in") or []) == ["Master database"],
            "Known in: no source": not data.get("known_in") and not data.get("file_error"),
        }.get(wanted, True)

    def row_values(self, term, data):
        defn = data["selected_definition"]
        if status_kind(data.get("status", "")) == "choice" and ctx_text(data):
            defn = f"Context: {ctx_text(data)}\n{defn}"
        return (term, data["category"], data["count"], data.get("change", ""), data["status"],
                defn, self.known_cell(term, data), self.first_use_cell(data), self.wrap_locations(data["locations"]))

    def first_use_cell(self, data):
        """Compact table text: where the first use is and whether it is defined there. Details and export say more."""
        if not first_use_lines(data):
            return ""
        items = list(data["first_use"].items())
        one_file = len(set(_file_of(l) for l in data["locations"])) <= 1
        f, u = items[0]
        loc = (u["loc"][len(f):].strip() if one_file else u["loc"]) or u["loc"]
        mark = {True: "defined", None: "check"}.get(u["defined"], "defined later" if u.get("later") else "NOT defined")
        text = self._wrap_px(f"{mark}: {loc}", "firstuse", 2)
        if len(items) > 1:
            bad = sum(1 for _f, x in items if x["defined"] is False)
            text += f"\n+{len(items) - 1} more file(s)" + (f", {bad} not defined" if bad else "")
        return text

    def _wrap_px(self, text, col, max_lines):
        """Word-wrap text to a column's pixel width, at most max_lines lines (last line ends in ... if cut)."""
        if not hasattr(self, "tree"):
            return text
        width = max(60, int(self.tree.column(col, "width")) - 16)
        lines, cur = [], ""
        for word in text.split(" "):
            trial = (cur + " " + word).strip()
            if cur and self.tree_font.measure(trial) > width:
                lines.append(cur); cur = word
            else:
                cur = trial
        lines.append(cur)
        if len(lines) > max_lines:
            lines = lines[:max_lines]; lines[-1] = lines[-1][:max(1, len(lines[-1]) - 3)] + "..."
        return "\n".join(lines)

    def wrap_locations(self, locations):
        """Locations wrapped to the column width, at most LOC_MAX_LINES lines (the rest: double-click the row)."""
        text = "; ".join(locations)
        if not hasattr(self, "tree"):
            return text
        width = max(80, int(self.tree.column("locations", "width")) - 16)
        measure = self.tree_font.measure
        lines, cur = [], ""
        for i, piece in enumerate(locations):
            add = piece + ("; " if i < len(locations) - 1 else "")
            if cur and measure(cur + add) > width:
                lines.append(cur.rstrip()); cur = ""
            while measure(add) > width and len(add) > 10:      # one very long location: break it
                cut = len(add)
                while cut > 10 and measure(add[:cut]) > width:
                    cut -= 5
                lines.append(add[:cut]); add = add[cut:]
            cur += add
        if cur:
            lines.append(cur.rstrip())
        if len(lines) > LOC_MAX_LINES:
            lines = lines[:LOC_MAX_LINES]
            lines[-1] = lines[-1][:max(0, len(lines[-1]) - 28)] + "...  (double-click for all)"
        return "\n".join(lines)

    def _fit_row_height(self):
        lines = 1
        for iid in self.tree.get_children():
            lines = max(lines, str(self.tree.set(iid, "locations")).count("\n") + 1,
                        str(self.tree.set(iid, "firstuse")).count("\n") + 1)
            if lines >= LOC_MAX_LINES:
                break
        line_h = self.tree_font.metrics("linespace")
        self.style.configure("Treeview", rowheight=line_h * lines + int(max(8, self.font_size // 2) * getattr(self, "dpi", 1.0)))

    def _maybe_rewrap(self, _e=None):
        w = (int(self.tree.column("locations", "width")), int(self.tree.column("firstuse", "width")))
        if self._loc_width is not None and w != self._loc_width:
            self.populate_table()
        self._loc_width = w

    @staticmethod
    def row_tag(data):
        if data.get("excluded", False): return "excluded"
        if data.get("manual"): return "manual"
        if data.get("gone"): return "gone"
        if data.get("known"): return "known"
        if data.get("effort") == "Approved": return "resolved"
        if data.get("effort"): return "effort"
        st = data["status"]
        kind = status_kind(st)
        if data.get("customer_definitions") and kind != "choice": return "customer"
        if kind == "error": return "fileerror"
        if kind == "choice": return "ambiguous"
        if data.get("client"): return "clientdiff" if data.get("db_definitions") else "client"
        if data.get("category", "").startswith("Equation"): return "equation"
        if kind == "candidate": return "candidate"
        return "resolved"

    def scan_cancelled(self):
        diag.scan_finished(cancelled=True)
        self._volcheck = None
        self.results = getattr(self, "_prev_results", {}) or {}
        if self.project is not None:
            self.save_project()   # keeps the files already read, so the next rescan continues from there
        self.reset_ui()
        self.refresh_type_options(); self.populate_table()
        self.status.configure(text="Scan cancelled. Showing the results from before this scan."
                              + (" Files already read are saved in the project." if self.project is not None else ""))

    def scan_done(self, results):
        diag.scan_finished(terms=len(results))
        self.results = results
        for term, item in self.manual_items.items():   # manual items survive rescans
            if term not in self.results:
                self.results[term] = item
        self.apply_instance_excludes()     # instances excluded one by one in this effort (2.13.0)
        self.apply_customer_defs()
        self.apply_saved_decisions()
        self.apply_picked_definitions()    # Use as definition, this effort only (2.14.0)
        self.apply_effort(reload=True)
        for d in self.results.values():   # known names keep their status unless the effort list holds them (2.13.0)
            if d.get("known") and not d.get("effort") and not d.get("excluded"):
                d["status"] = KNOWN_STATUS + (f": {d['omit_reason']}" if d.get("omit_reason") else "")
        self.mark_known_in()               # customer docs, customer list, master database, or no source (2.14.0)
        rec = self.effort_after_scan(getattr(self, "queued_files", []) or []) if self.effort is not None else None   # 2.14.0
        dbr = self.after_db_refresh()      # a database update keeps every row and marks what changed (2.14.0)
        self.save_project()
        self.reset_ui()
        self.refresh_type_options()
        self.populate_table()
        self.volume_check_after_scan()
        self.start_context_suggestions()   # meaning from nearby words, in the background (2.13.0)
        if rec:
            self._last_effort_scan = rec
            self.update_effort_ui()
            self.status.configure(text=f"Effort list updated: {rec['found']} found again, {rec['new']} new (not in the list"
                                       f"{': Effort > Review new terms' if rec['new'] else ''}), {rec['gone']} no longer found.")
            self.scan_landed()             # every scan lands in Scan review (2.14.0)
        if dbr:
            self.status.configure(text=f"Results updated against the master database: {dbr['changed']} changed (Status > Database "
                                       f"update), {dbr['new']} new, {dbr['kept']} kept from before. Nothing was removed.")

    def populate_table(self):
        if not hasattr(self, "tree"):
            return
        for i in self.tree.get_children(): self.tree.delete(i)
        show_exc = self.show_exc_var.get()
        ftype, fstat = self.type_var.get(), self.status_var.get()
        query = self.search_var.get().strip().casefold()

        for t, d in self.results.items():     # Known in for results saved before 2.14.0
            if "known_in" not in d and not d.get("file_error"):
                d["known_in"] = self.known_in(t, d)
        v_count = e_count = a_count = c_count = 0
        rows = []
        for term, data in self.results.items():
            excluded = data.get("excluded", False)
            if excluded:
                e_count += 1
            else:
                v_count += 1
                k = status_kind(data["status"])
                if k == "choice": a_count += 1
                elif k == "candidate": c_count += 1

            if excluded and not (show_exc or fstat == "Excluded"): continue
            if ftype != ALL_TYPES and base_type(data["category"]) != ftype: continue
            if fstat != STATUS_FILTERS[0] and not self.status_matches(data, fstat): continue
            if query and query not in term.casefold() and query not in data["selected_definition"].casefold(): continue
            rows.append((term, data))

        rows.sort(key=lambda r: self.sort_key(*r), reverse=self.sort_desc)
        for term, data in rows:
            self.tree.insert("", "end", iid=term, tags=(self.row_tag(data),), values=self.row_values(term, data))
        self._fit_row_height()
        self._loc_width = (int(self.tree.column("locations", "width")), int(self.tree.column("firstuse", "width")))

        self.lbl_showing.configure(text=f"Showing {len(rows)} of {len(self.results)}" if self.results else "")
        if self.results and time.time() >= getattr(self, "_status_hold", 0):   # a fresh "Resumed ..." stays readable
            self.status.configure(text=f"Active Findings: {v_count} ({a_count} need choice, {c_count} new candidates) | Suppressed Exclusions: {e_count}")
        if v_count > 0: self.btn_export.configure(state="normal")
        if self.selected_term and self.tree.exists(self.selected_term):
            self.tree.selection_set(self.selected_term)
            self.tree.see(self.selected_term)
        self.session_changed()      # saved in a moment, so a crash or a close loses nothing (2.11.0)
        self.resume_changed()

    def excludable(self, term):
        d = self.results.get(term)
        return bool(d) and not (d.get("excluded") or d.get("manual") or d.get("gone") or d.get("file_error"))

    def on_select(self, event):
        sel = self.tree.selection()
        if not sel: return
        self.resume_changed()
        self.update_effort_buttons()
        if len(sel) > 1:   # several rows: only Exclude applies to a group
            n = sum(1 for t in sel if self.excludable(t))
            self.selected_term = sel[0]
            for w in (self.btn_apply, self.btn_save, self.btn_ocr, self.entry_def, self.combo):
                w.configure(state="disabled")
            self.btn_exc.configure(text=f"Exclude {n}", command=self.exclude_term, state="normal" if n else "disabled")
            self.lbl_act.configure(text=f"{len(sel)} rows selected ({n} can be excluded):")
            return
        term = sel[0]
        self.selected_term = term
        data = self.results.get(term)
        if not data: return
        if self.preview_var.get():
            self.preview.show(term)

        if data.get("file_error"):
            self.btn_exc.configure(text="Exclude", command=self.exclude_term)
            for w in (self.btn_exc, self.btn_apply, self.btn_save, self.btn_ocr, self.entry_def, self.combo):
                w.configure(state="disabled")
            self.lbl_act.configure(text="File could not be read:")
            return

        if data.get("gone"):
            for w in (self.btn_exc, self.btn_apply, self.btn_save, self.btn_ocr, self.entry_def, self.combo):
                w.configure(state="disabled")
            self.lbl_act.configure(text=f"'{term}' no longer found (last seen {data.get('last_seen', '?')}):")
            return

        if data.get("manual"):
            for w in (self.btn_apply, self.btn_save, self.btn_ocr, self.entry_def, self.combo):
                w.configure(state="disabled")
            self.btn_exc.configure(text="Remove Item", command=self.remove_manual_item, state="normal")
            self.lbl_act.configure(text=f"Added manually '{term}':")
            return
        self.btn_exc.configure(text="Exclude", command=self.exclude_term)
        self.btn_exc.configure(state="disabled" if data.get("excluded", False) else "normal")

        if (data["in_db"] or data.get("customer_definitions")) and not data.get("excluded", False) and len(data["definitions"]) > 1:
            opts = list(data["definitions"]) + ["Keep All Possibilities"]
            ctx = (data.get("context") or {}).get("definition", "")
            pick = next((o for o in opts if ctx and (o == ctx or o.endswith(ctx))), opts[0])
            self.combo.configure(state="normal", values=opts); self.combo.set(pick)
            self.btn_apply.configure(state="normal")
            self.entry_def.configure(state="disabled"); self.btn_save.configure(state="disabled"); self.btn_ocr.configure(state="disabled")
            self.lbl_act.configure(text=f"Client list term '{term}' (client meaning pre-selected):" if data.get("client")
                                   else f"Disambiguate '{term}' (context suggests: {ctx}):" if ctx
                                   else f"Disambiguate '{term}':")
        elif not data["in_db"]:
            shown = (data.get("customer_definitions") or ["Not in DB"])[:1]
            self.combo.configure(state="normal", values=shown); self.combo.set(shown[0]); self.combo.configure(state="disabled")
            self.btn_apply.configure(state="disabled")
            self.entry_def.configure(state="normal"); self.entry_def.delete(0, "end")
            self.btn_save.configure(state="normal"); self.btn_ocr.configure(state="normal")
            self.lbl_act.configure(text=f"Customer document '{term}':" if data.get("customer_definitions") else f"New Candidate '{term}':")
        else:
            self.combo.configure(state="normal", values=data["definitions"][:1] or [""])
            self.combo.set(data["definitions"][0] if data["definitions"] else "")
            self.combo.configure(state="disabled")
            self.btn_apply.configure(state="disabled")
            self.entry_def.configure(state="disabled"); self.btn_save.configure(state="disabled"); self.btn_ocr.configure(state="disabled")
            self.lbl_act.configure(text=f"Selected '{term}':")

    def apply_choice(self):
        if not self.selected_term: return
        choice = self.combo.get()
        data = self.results[self.selected_term]
        tag = self._set_choice(data, choice)
        self.session_decisions = getattr(self, "session_decisions", {}); self.session_decisions[self.selected_term] = choice
        try:
            self.save_choice(self.selected_term, choice)
        except OSError:
            diag.swallowed()
            messagebox.showwarning("Not saved", "The choice applies now, but could not be saved for later scans "
                                   "(the config folder is not writable).")
        diag.activity("apply_choice")
        if self.project is not None:
            self.project.data["decisions"][self.selected_term] = choice
            self.save_project()
        if self.effort is not None and self.selected_term in self.effort.rows and choice != "Keep All Possibilities":
            chosen = data.get("selected_definition", "")     # 2.14.0: the choice goes into the effort list too
            if self.effort.add_or_update(self.selected_term, change_note=f"definition chosen: {chosen[:80]}", Definition=chosen) != "unchanged":
                self._save_effort_list()
        self.tree.item(self.selected_term, values=self.row_values(self.selected_term, data), tags=(tag,))
        self.session_changed()

    def open_fix_manager(self):
        """View, add, edit, delete, and undo entries in config\\ocr_corrections.csv."""
        diag.activity("ocr_fixes_opened")
        if getattr(self, "fix_win", None) is not None and self.fix_win.winfo_exists():
            self.fix_win.focus(); return
        win = ctk.CTkToplevel(self); self.fix_win = win
        win.title("OCR Fixes - config\\ocr_corrections.csv")
        win.geometry("760x520"); win.minsize(620, 400)
        win.transient(self); win.after(100, win.lift)

        ctk.CTkLabel(win, text="Each line turns an OCR misread into the correct term on the next scan.",
                     anchor="w").pack(fill="x", padx=14, pady=(12, 4))
        frame = ctk.CTkFrame(win); frame.pack(fill="both", expand=True, padx=14, pady=4)
        frame.grid_rowconfigure(0, weight=1); frame.grid_columnconfigure(0, weight=1)
        tree = ttk.Treeview(frame, columns=("misread", "correct", "note"), show="headings", selectmode="browse")
        for cid, label, width in (("misread", "Misread", 160), ("correct", "Correct term", 160), ("note", "Note", 340)):
            tree.heading(cid, text=label); tree.column(cid, width=int(width * self.font_size / 11.0 * getattr(self, "dpi", 1.0)), anchor="w")
        vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview); tree.configure(yscrollcommand=vsb.set)
        tree.grid(row=0, column=0, sticky="nsew"); vsb.grid(row=0, column=1, sticky="ns")

        form = ctk.CTkFrame(win); form.pack(fill="x", padx=14, pady=6)
        e_wrong = ctk.CTkEntry(form, placeholder_text="Misread (e.g. RFF)", width=150); e_wrong.pack(side="left", padx=4, pady=8)
        e_right = ctk.CTkEntry(form, placeholder_text="Correct (e.g. RFP)", width=150); e_right.pack(side="left", padx=4, pady=8)
        SP.attach(e_wrong); SP.button(form, e_right).pack(side="left", padx=(0, 4), pady=8)
        e_note = ctk.CTkEntry(form, placeholder_text="Note (optional)", width=220); e_note.pack(side="left", padx=4, pady=8)
        status = ctk.CTkLabel(win, text="", anchor="w", text_color="gray"); status.pack(fill="x", padx=14)

        def load():
            tree.delete(*tree.get_children())
            for i, row in enumerate(read_ocr_fix_rows()):
                tree.insert("", "end", iid=str(i), values=row)
        def rows_now():
            return [list(tree.item(i, "values")) for i in tree.get_children()]
        def save(rows, msg):
            write_ocr_fix_rows(rows); load(); status.configure(text=msg + "  Rescan to apply.")
        def on_pick(_e=None):
            sel = tree.selection()
            if not sel: return
            wrong, right, note = (list(tree.item(sel[0], "values")) + ["", "", ""])[:3]
            for ent, val in ((e_wrong, wrong), (e_right, right), (e_note, note)):
                ent.delete(0, "end"); ent.insert(0, val)
        def add_update():
            wrong, right, note = e_wrong.get().strip(), e_right.get().strip(), e_note.get().strip()
            if not wrong or not right or wrong == right:
                status.configure(text="Enter a misread and a different correct term."); return
            rows = rows_now()
            for r in rows:
                if r[0] == wrong:
                    r[1], r[2] = right, note or r[2]
                    save(rows, f"Updated: {wrong} > {right}."); return
            rows.append([wrong, right, note or "Added in OCR Fixes window"])
            save(rows, f"Added: {wrong} > {right}.")
        def delete():
            sel = tree.selection()
            if not sel:
                status.configure(text="Select a line to delete."); return
            rows = rows_now(); gone = rows.pop(int(tree.index(sel[0])))
            save(rows, f"Deleted: {gone[0]} > {gone[1]}.")
        def undo_last():
            rows = rows_now()
            if not rows:
                status.configure(text="Nothing to undo."); return
            gone = rows.pop()
            save(rows, f"Removed last entry: {gone[0]} > {gone[1]}.")

        btns = ctk.CTkFrame(win); btns.pack(fill="x", padx=14, pady=(4, 12))
        ctk.CTkButton(btns, text="Add / Update", width=120, command=add_update).pack(side="left", padx=4, pady=8)
        ctk.CTkButton(btns, text="Delete Selected", width=130, fg_color="#DC2626", command=delete).pack(side="left", padx=4, pady=8)
        ctk.CTkButton(btns, text="Undo Last Entry", width=130, fg_color="#64748B", command=undo_last).pack(side="left", padx=4, pady=8)
        ctk.CTkButton(btns, text="Close", width=90, fg_color="#475569", command=win.destroy).pack(side="right", padx=4, pady=8)
        tree.bind("<<TreeviewSelect>>", on_pick)
        load()
        win._fix_api = {"tree": tree, "wrong": e_wrong, "right": e_right, "note": e_note, "status": status,
                        "add": add_update, "delete": delete, "undo": undo_last}   # used by automated tests

    def save_ocr_fix(self):
        """Record that the selected term is an OCR misread of the term typed in the box."""
        diag.activity("fix_ocr_saved")
        if not self.selected_term: return
        right = self.entry_def.get().strip()
        if not right:
            messagebox.showwarning("Fix OCR", "Type the correct term in the box first (for example RFP), then click Fix OCR.")
            return
        wrong = self.selected_term
        rows = [r for r in read_ocr_fix_rows() if r[0] != wrong]
        rows.append([wrong, right, "Added with Fix OCR"])
        write_ocr_fix_rows(rows)
        data = self.results[wrong]
        data["in_db"] = True
        data["status"] = f"OCR fix saved: {wrong} > {right}"
        data["selected_definition"] = "Rescan to apply. View, edit, or undo with the OCR Fixes button."
        self.tree.item(wrong, values=self.row_values(wrong, data), tags=("resolved",))
        self.session_changed()
        self.entry_def.delete(0, "end")
        self.entry_def.configure(state="disabled")
        self.btn_save.configure(state="disabled"); self.btn_ocr.configure(state="disabled")

    def save_candidate(self):
        diag.activity("include_in_db")
        if not self.selected_term: return
        nd = self.entry_def.get().strip()
        if not nd:
            messagebox.showwarning("Error", "Type a definition first.")
            return
        data = self.results[self.selected_term]
        cat = data["category"].replace(" (Candidate)", "")
        res, _row = db_add(self.selected_term, cat, nd, "Include")   # one row per different meaning; the same one is flagged
        if res == "duplicate":
            messagebox.showinfo("Already in the database", f"The database already has {self.selected_term} = \"{nd}\". "
                                "Nothing was added; the result uses that definition.")
        data["in_db"] = True
        data["excluded"] = False
        data["category"] = cat
        data["definitions"] = [nd]
        data["selected_definition"] = nd
        data["status"] = "Added to DB"
        self.tree.item(self.selected_term, values=self.row_values(self.selected_term, data), tags=("resolved",))
        self.session_changed()
        self.entry_def.delete(0, "end")
        self.entry_def.configure(state="disabled")
        self.btn_save.configure(state="disabled"); self.btn_ocr.configure(state="disabled")

    # ---------- exclude one instance (2.13.0): the term stays wherever else it appears ----------
    @staticmethod
    def _drop_instance(d, loc):
        if loc not in d.get("locations", []):
            return False
        d["locations"].remove(loc)
        d["count"] = max(len(d["locations"]), int(d.get("count", 1)) - 1)
        lst = d.setdefault("excluded_instances", [])
        if loc not in lst:
            lst.append(loc)
        if not d["locations"] and not d.get("excluded"):
            d["_pre_instances"] = [d.get("status", ""), d.get("selected_definition", "")]
            d["excluded"] = True
            d["status"] = "All instances excluded"
        return True

    def apply_instance_excludes(self):
        """After every scan: take out the instances excluded in this effort."""
        try:
            ex = self.ws.instance_excludes()
        except Exception:
            diag.swallowed()
            return 0
        n = 0
        for term, locs in ex.items():
            d = self.results.get(term)
            if d:
                for loc in locs:
                    n += self._drop_instance(d, loc)
        return n

    def exclude_instance(self, term, loc):
        """Exclude one place a term was found (a misreading, a heading word). Saved for this effort only."""
        d = self.results.get(term)
        if not d or not self._drop_instance(d, loc):
            return False
        ex = self.ws.instance_excludes()
        if loc not in ex.setdefault(term, []):
            ex[term].append(loc)
        self.ws.save_instance_excludes(ex)
        diag.activity("exclude_instance", n=1)
        self.selected_term = term if not d.get("excluded") else None
        self.populate_table()
        return True

    def restore_instances(self, term):
        """Bring back every excluded instance of a term (they show again now and on later scans)."""
        d = self.results.get(term)
        ex = self.ws.instance_excludes()
        locs = ex.pop(term, [])
        self.ws.save_instance_excludes(ex)
        if d:
            back = [l for l in d.get("excluded_instances", []) if l not in d["locations"]]
            d["locations"].extend(back)
            d["count"] = int(d.get("count", 0)) + len(back)
            d.pop("excluded_instances", None)
            pre = d.pop("_pre_instances", None)
            if pre and d.get("status") == "All instances excluded":
                d["excluded"] = False
                d["status"], d["selected_definition"] = pre[0], pre[1]
        diag.activity("restore_instances", n=len(locs))
        self.selected_term = term
        self.populate_table()
        return len(locs)

    # ---------- meaning from context (2.13.0): rules only, the person decides ----------
    def _context_targets(self):
        out = {}
        for term, d in self.results.items():
            st = status_text(d.get("status", ""))
            if d.get("excluded") or d.get("known") or d.get("gone") or not st.startswith(("CHOOSE DEFINITION", "CUSTOMER DOCS DIFFER")):
                continue
            defs = d.get("customer_definitions") if st.startswith("CUSTOMER DOCS DIFFER") else d.get("definitions", [])
            defs = [x for x in (defs or []) if x and not x.startswith("MULTIPLE")]
            if len(defs) >= 2:
                lfs = [u.get("lf", "") for u in d.get("first_use", {}).values() if u.get("lf")]
                out[term] = (defs, list(d.get("locations", [])), lfs)
        return out

    def start_context_suggestions(self):
        targets = self._context_targets()
        if not targets:
            return
        files = {}
        for _t, (_defs, locs, _lfs) in targets.items():
            for loc in locs:
                name = _file_of(loc)
                if name not in files:
                    files[name] = self.path_for(name)
        def work():
            import context as X
            texts = {}
            for name, path in files.items():
                if path and os.path.isfile(path) and os.path.splitext(path)[1].lower() in (".docx", ".pptx", ".xlsx", ".pdf"):
                    try:
                        texts[name] = typed_chunks(path)
                    except Exception:
                        diag.swallowed()
            res = {}
            for term, (defs, locs, lfs) in targets.items():
                places = {}
                for loc in locs:
                    chunks = texts.get(_file_of(loc), [])
                    want = loc.split(" [OCR fix", 1)[0]
                    hit = [t for lbl, t in chunks if lbl == want]
                    wins = [w for t in hit for w in X.windows(t, term)]
                    if wins:
                        places[loc] = wins
                sug = X.suggest(defs, places, lfs)
                if sug:
                    res[term] = sug
            self.after(0, lambda: self._apply_context(res))
        threading.Thread(target=work, daemon=True).start()

    def _apply_context(self, res):
        n = 0
        for term, sug in res.items():
            d = self.results.get(term)
            if not d:
                continue
            d["context"] = sug
            if status_text(d.get("status", "")).startswith("CHOOSE DEFINITION"):
                if sug.get("definition"):
                    d["status"] = "CHOOSE DEFINITION (context suggests one)"
                    n += 1
                elif sug.get("mixed"):
                    d["status"] = "CHOOSE DEFINITION (meaning differs by place)"
        diag.activity("context_suggested", terms=n)
        self.populate_table()

    def accept_context(self, terms=None):
        """Apply the context suggestion to the selected rows (or every row that has one)."""
        terms = [t for t in (terms or self.tree.selection() or list(self.results)) if (self.results.get(t, {}).get("context") or {}).get("definition")
                 and status_kind(self.results[t].get("status", "")) == "choice"]
        for t in terms:
            d = self.results[t]
            choice = d["context"]["definition"]
            opt = next((o for o in d.get("definitions", []) if o == choice or o.endswith(choice)), choice)
            self._set_choice(d, opt)
            try:
                self.save_choice(t, opt)
            except OSError:
                diag.swallowed()
            if self.project is not None:
                self.project.data["decisions"][t] = opt
        if self.project is not None and terms:
            self.save_project()
        diag.activity("context_accepted", n=len(terms))
        self.status.configure(text=f"Context suggestions accepted for {len(terms)} term(s). Each one is saved like Apply Choice.")
        self.populate_table()
        return len(terms)

    # ---------- known names (2.13.0): recognized and kept in the database, left off the acronym lists ----------
    def mark_known(self, terms=None):
        """Add the selected terms to the omit list: recognized and kept in the database, left off the final lists."""
        terms = [t for t in (terms or self.tree.selection()) if t in self.results and not self.results[t].get("file_error")
                 and not self.results[t].get("manual")]
        if not terms:
            return None
        win = ctk.CTkToplevel(self); win.title("Add to the omit list"); win.geometry("560x330"); win.transient(self)
        win.after(150, win.lift)
        diag.register(win, "omit_add")
        one = len(terms) == 1
        ctk.CTkLabel(win, text=("Omit list: found and recognized on every scan, saved in the database, left off the final "
                                "lists (CSV export, volume acronym lists)."), wraplength=520, justify="left").pack(anchor="w", padx=14, pady=(12, 6))
        ctk.CTkLabel(win, text=(f"Term: {terms[0]}" if one else f"{len(terms)} terms: " + ", ".join(terms[:8]) + ("..." if len(terms) > 8 else "")),
                     anchor="w", font=("Segoe UI", 13, "bold")).pack(anchor="w", padx=14)
        e_def = None
        if one:
            ctk.CTkLabel(win, text="What it stands for (optional)").pack(anchor="w", padx=14, pady=(8, 0))
            e_def = ctk.CTkEntry(win, width=520); e_def.pack(padx=14, pady=2)
            cur = self.results[terms[0]].get("selected_definition", "")
            if cur and not status_text(cur).startswith(("[", "MULTIPLE", "Rescan")) and cur != KNOWN_BLANK:
                e_def.insert(0, cur)
        ctk.CTkLabel(win, text="Why it is omitted").pack(anchor="w", padx=14, pady=(8, 0))
        reason = ctk.CTkComboBox(win, values=OMIT_REASONS, width=260); reason.set(OMIT_REASONS[0]); reason.pack(anchor="w", padx=14, pady=2)
        msg = ctk.CTkLabel(win, text="", text_color="#B45309", wraplength=520, justify="left"); msg.pack(anchor="w", padx=14, pady=4)

        def ok():
            why = reason.get().strip()
            out = []
            for t in terms:
                d = self.results[t]
                cur = d.get("selected_definition", "")
                defn = (e_def.get().strip() if e_def is not None else
                        ("" if status_text(cur).startswith(("[", "MULTIPLE", "Rescan")) or cur == KNOWN_BLANK else cur))
                res, _row = db_add(t, base_type(d.get("category", "Acronym")), defn, KNOWN_ACTION, {OMIT_REASON_COL: why})
                out.append(res)
                d.update(known=True, in_db=True, excluded=False, omit_reason=why, selected_definition=defn or KNOWN_BLANK,
                         status=KNOWN_STATUS + (f": {why}" if why else ""))
            diag.activity("omit_added", n=len(terms))
            dup = out.count("duplicate")
            self.status.configure(text=f"{len(terms)} term(s) on the omit list" + (f" ({dup} were already there)" if dup else "") +
                                  ". Found on every scan, left off the final lists.")
            win.destroy()
            self.populate_table()
        b = ctk.CTkFrame(win, fg_color="transparent"); b.pack(fill="x", padx=14, pady=10)
        ctk.CTkButton(b, text="Add to omit list", command=ok).pack(side="left")
        ctk.CTkButton(b, text="Cancel", fg_color="#475569", command=win.destroy).pack(side="right")
        win._api = {"defn": e_def, "reason": reason, "ok": ok, "msg": msg}
        return win

    def open_omit_list(self):
        """Edit > Omit list: everything on it, add (several meanings of one acronym are fine), paste many, remove."""
        diag.activity("omit_list_opened")
        win = ctk.CTkToplevel(self); win.title("Omit list"); win.geometry("900x600"); win.transient(self); win.after(150, win.lift)
        diag.register(win, "omit_list")
        ctk.CTkLabel(win, text="Found and recognized on every scan, saved in the company database (and Knowledge Packs), left off "
                               "the final lists. One acronym can have several entries when the meanings differ.",
                     wraplength=860, justify="left").pack(anchor="w", padx=12, pady=(10, 6))
        tv = ttk.Treeview(win, columns=("term", "def", "why"), show="headings", selectmode="extended")
        for c, h, w in (("term", "Term", 140), ("def", "Stands for", 420), ("why", "Why omitted", 220)):
            tv.heading(c, text=h); tv.column(c, width=w, stretch=(c == "def"))
        tv.pack(fill="both", expand=True, padx=12)
        form = ctk.CTkFrame(win, fg_color="transparent"); form.pack(fill="x", padx=12, pady=(8, 0))
        e_term = ctk.CTkEntry(form, width=120, placeholder_text="Term"); e_term.pack(side="left")
        e_def = ctk.CTkEntry(form, width=360, placeholder_text="Stands for (optional)"); e_def.pack(side="left", padx=6)
        why = ctk.CTkComboBox(form, values=OMIT_REASONS, width=200); why.set(OMIT_REASONS[0]); why.pack(side="left")
        msg = ctk.CTkLabel(win, text="", anchor="w", wraplength=860, justify="left", text_color="#B45309")
        msg.pack(fill="x", padx=12, pady=4)

        def fill():
            tv.delete(*tv.get_children())
            for i, r in enumerate(read_db_rows()):
                if _action_kind(r[3]) == "omit":
                    tv.insert("", "end", iid=str(i), values=(r[0], r[2], r[4].get(OMIT_REASON_COL, "")))
            win.title(f"Omit list ({len(tv.get_children())})")

        def add_one(term, defn, reason):
            term = term.strip()
            if not term:
                return "empty"
            res, row = db_add(term, "Acronym", defn.strip(), KNOWN_ACTION, {OMIT_REASON_COL: reason})
            d = self.results.get(fold(term))
            if d is not None and res != "duplicate":
                d.update(known=True, in_db=True, excluded=False, omit_reason=reason,
                         selected_definition=defn.strip() or d.get("selected_definition") or KNOWN_BLANK,
                         status=KNOWN_STATUS + (f": {reason}" if reason else ""))
            return res

        def add():
            res = add_one(e_term.get(), e_def.get(), why.get())
            t = e_term.get().strip()
            if res == "duplicate":
                msg.configure(text=f"{t} already has this meaning on the list. Nothing added. A different meaning is fine.")
            elif res == "empty":
                msg.configure(text="Type a term.")
            else:
                msg.configure(text=f"{t} added." + (" (It was in the database with this meaning; it is now on the omit list.)"
                                                    if res == "updated" else ""))
                e_term.delete(0, "end"); e_def.delete(0, "end")
                diag.activity("omit_added", n=1)
            fill(); self.populate_table()

        def paste():
            box = ctk.CTkToplevel(win); box.title("Paste several"); box.geometry("620x420"); box.transient(win)
            ctk.CTkLabel(box, text="One per line:  TERM, stands for, why   (the last two are optional; a tab works too)",
                         anchor="w").pack(anchor="w", padx=12, pady=(10, 4))
            t = tk.Text(box, height=14, wrap="none"); t.pack(fill="both", expand=True, padx=12)
            def go():
                counts = {"added": 0, "updated": 0, "duplicate": 0, "empty": 0}
                for line in t.get("1.0", "end").splitlines():
                    if not line.strip():
                        continue
                    parts = [p.strip() for p in re.split(r"\t|,", line, maxsplit=2)]
                    parts += ["", ""]
                    reason = parts[2] if parts[2] else why.get()
                    counts[add_one(parts[0], parts[1], reason)] += 1
                diag.activity("omit_pasted", added=counts["added"], n=counts["duplicate"])
                msg.configure(text=f"{counts['added']} added, {counts['updated']} moved from the database to the omit list, "
                                   f"{counts['duplicate']} already on the list with the same meaning (not added again).")
                box.destroy(); fill(); self.populate_table()
            ctk.CTkButton(box, text="Add all", command=go).pack(anchor="w", padx=12, pady=10)
            box._go = go; box._text = t
            win._paste = box
            return box

        def remove():
            sel = set(tv.selection())
            if not sel:
                return
            rows = read_db_rows()
            gone = [rows[int(i)][0] for i in sel]
            keep = [r for i, r in enumerate(rows) if str(i) not in sel]
            write_db_rows(keep)
            for term in gone:
                d = self.results.get(fold(term))
                if d and d.get("known") and not any(fold(r[0]) == fold(term) and _action_kind(r[3]) == "omit" for r in keep):
                    d.pop("known", None); d.pop("omit_reason", None)
                    d["status"], d["in_db"] = "NEW CANDIDATE", False
                    d["selected_definition"] = "[Removed from the omit list: Scan Again to refresh]"
            diag.activity("omit_removed", n=len(sel))
            msg.configure(text=f"{len(sel)} removed from the omit list.")
            fill(); self.populate_table()
        b = ctk.CTkFrame(form, fg_color="transparent"); b.pack(side="left", padx=6)
        ctk.CTkButton(b, text="Add", width=60, command=add).pack(side="left")
        ctk.CTkButton(win, text="Paste several...", width=120, command=paste).pack(side="left", padx=12, pady=(0, 12))
        ctk.CTkButton(win, text="Remove selected", width=130, fg_color="#B91C1C", command=remove).pack(side="left", pady=(0, 12))
        ctk.CTkButton(win, text="Close", width=80, fg_color="#475569", command=win.destroy).pack(side="right", padx=12, pady=(0, 12))
        e_term.bind("<Return>", lambda _e: add()); e_def.bind("<Return>", lambda _e: add())
        fill()
        win._api = {"tree": tv, "term": e_term, "defn": e_def, "why": why, "add": add, "paste": paste, "remove": remove, "msg": msg}
        return win

    def unmark_known(self, terms=None):
        terms = [t for t in (terms or self.tree.selection()) if self.results.get(t, {}).get("known")]
        if not terms:
            return
        rows = read_db_rows()
        keep = []
        for r in rows:
            if any(fold(r[0]) == fold(t) for t in terms) and str(r[3]).lower().startswith(("known", "omit")):
                if not r[2] or r[2] == KNOWN_BLANK:
                    continue                                   # no definition: the row goes; the term is a candidate again
                r[3] = "Include"
            keep.append(r)
        write_db_rows(keep)
        for t in terms:
            d = self.results[t]
            d.pop("known", None)
            has = d.get("selected_definition") not in ("", KNOWN_BLANK)
            d["status"] = "Resolved (In DB)" if has else "NEW CANDIDATE"
            d["in_db"] = has
        diag.activity("known_name_unmarked", n=len(terms))
        self.populate_table()

    def exclude_term(self):
        diag.activity("exclude_term")
        terms = [t for t in self.tree.selection() if self.excludable(t)]
        if not terms and self.selected_term and self.excludable(self.selected_term):
            terms = [self.selected_term]
        if not terms: return
        if len(terms) > 1:
            shown = ", ".join(terms[:12]) + (f" and {len(terms) - 12} more" if len(terms) > 12 else "")
            if not messagebox.askyesno("Exclude terms", f"Exclude {len(terms)} terms?\n\n{shown}\n\n"
                                       "They are saved to the database as Exclude rules, so future scans skip them. "
                                       "Show Excluded (or Status > Excluded) lists them again."):
                return
        items = []
        for term in terms:
            data = self.results[term]
            items.append((term, data["category"], data["definitions"][0] if data["definitions"] else "Excluded via GUI"))
        update_csv_exclusions(items)   # one database write for the whole group
        for term, _cat, reason in items:
            data = self.results[term]
            data["excluded"] = True
            data["in_db"] = True
            data["status"] = "Excluded by rule"
            data["selected_definition"] = f"[EXCLUDED] {reason}"
        self.btn_exc.configure(text="Exclude", state="disabled")
        self.lbl_act.configure(text=f"Excluded {len(items)} term(s). Select a row above (Ctrl+click or Shift+click for several):")
        self.selected_term = None
        self.populate_table()

    def reset_ui(self):
        self.prog.stop()
        self.prog.configure(mode="determinate")
        self.prog.set(1)
        self.btn_files.configure(state="normal")
        self.btn_folder.configure(state="normal")
        self.btn_project.configure(state="normal")
        self.btn_again.configure(state="normal" if (self.project is not None or getattr(self, "last_scan", None)) else "disabled")
        self.btn_cancel.pack_forget()
        self.update_project_ui()

    def export_name(self):
        slug = lambda t: re.sub(r"[^A-Za-z0-9]+", "-", t).strip("-")
        parts = []
        if self.project is not None:
            parts.append(slug(self.project.name))
        parts.append("Acronym_Report")
        if self.status_var.get() != STATUS_FILTERS[0]:
            parts.append(slug(self.status_var.get()))
        if self.type_var.get() != ALL_TYPES:
            parts.append(slug(self.type_var.get()))
        q = self.search_var.get().strip()
        if q:
            parts.append("search-" + slug(q)[:30])
        parts.append(datetime.datetime.now().strftime("%Y-%m-%d"))
        return "_".join(parts) + ".csv"

    def export_csv(self):
        """Export exactly the rows shown (current Type / Status / Search filters, in the current sort order)."""
        diag.activity("export_csv")
        terms = [t for t in self.tree.get_children() if t in self.results]
        left_out = 0
        if self.status_var.get() != "Found in omit list":      # the omit list stays off the list (2.13.0)
            left_out = sum(1 for t in terms if self.results[t].get("known"))
            terms = [t for t in terms if not self.results[t].get("known")]
        if not terms:
            messagebox.showinfo("Nothing to export", "No rows are showing. Change the filters, then export."); return
        path = filedialog.asksaveasfilename(initialdir=self.start_dir(), defaultextension=".csv",
                                            filetypes=[("CSV Report", "*.csv")], initialfile=self.export_name())
        if not path:
            return
        strip = lambda x: x[len(CLIENT_TAG):] if x.startswith(CLIENT_TAG) else x[len(DB_TAG):] if x.startswith(DB_TAG) else x
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["Term", "Type", "Count", "Change", "Status", "Selected definition", "Possible definitions",
                        "Client list definition", "Database definition", "Author decision",
                        "First use (acronyms: define here)", "Locations",
                        "First found by a scan (date)", "Last found by a scan (date)", "Known in", "Context suggests"])
            for term in terms:
                d = self.results[term]
                defs = d.get("definitions", [])
                client = d.get("client_definitions", [])
                dbdefs = d.get("db_definitions", []) if d.get("client") else ([] if d.get("excluded") else
                                                                              [x for x in defs if not d.get("manual")])
                w.writerow([term, d["category"], d["count"], d.get("change", ""), d["status"], d["selected_definition"],
                            "\n".join(defs) if len(defs) > 1 or d.get("client") else "",
                            "\n".join(client), "\n".join(strip(x) for x in dbdefs), "",
                            "\n".join(first_use_lines(d)), "\n".join(d["locations"]), d.get("first_seen", ""), d.get("last_seen", ""),
                            self.known_cell(term, d), ctx_text(d)])
        view = self.status_var.get() if self.status_var.get() != STATUS_FILTERS[0] else "all statuses"
        messagebox.showinfo("Exported", f"{len(terms)} rows exported ({view}) to:\n{path}\n\n"
                            + (f"{left_out} term(s) on the omit list left out (Status > Found in omit list shows them).\n\n" if left_out else "")
                            + "Possible definitions lists every meaning (client first). Authors can fill in Author decision.")

    # ---------- details ----------
    def show_details(self, event=None):
        iid = self.tree.identify_row(event.y) if event is not None else (self.tree.selection() or [None])[0]
        d = self.results.get(iid) if iid else None
        if not d:
            return
        win = ctk.CTkToplevel(self); win.title(f"Details: {iid}"); win.geometry("760x560"); win.transient(self)
        win.after(100, win.lift)
        box = ctk.CTkTextbox(win, wrap="word", font=("Segoe UI", 13)); box.pack(fill="both", expand=True, padx=12, pady=12)
        lines = [f"{iid}   ({d['category']}, found {d['count']} times)", f"Status: {d['status']}", "",
                 f"Selected definition: {d['selected_definition']}"]
        if d.get("client"):
            lines += ["", "Client list definition(s):"] + [f"  - {x}" for x in d.get("client_definitions", [])]
            lines += [f"  (from: {', '.join(d.get('client_sources', []))})"]
            if d.get("db_definitions"):
                lines += ["", "Database definition(s), different from the client:"] + [f"  - {x}" for x in d["db_definitions"]]
        elif len(d.get("definitions", [])) > 1:
            lines += ["", "Possible definitions:"] + [f"  - {x}" for x in d["definitions"]]
        fu = first_use_lines(d)
        if fu:
            lines += ["", "First use in each file (define it here as: Spelled Out Words (ACR)):"] + [f"  - {x}" for x in fu]
        lines += ["", f"Locations ({len(d['locations'])}):"] + [f"  - {x}" for x in d["locations"]]
        if d.get("first_seen"):
            lines += ["", f"Project history: first found by a scan on {d['first_seen']}, last found {d.get('last_seen', '')}"]
        box.insert("1.0", "\n".join(lines)); box.configure(state="disabled")

    # ---------- client acronym lists ----------
    def set_client_lists(self, lists):
        self.client_lists = lists
        set_client_lists(lists)
        n = len(lists)
        if hasattr(self, "btn_client"):
            self.btn_client.configure(text=f"Client Lists ({n})" if n else "Client Lists")

    def _preview_client_list(self, name, rows):
        """Show what was read from a client list; True if the user accepts it."""
        win = ctk.CTkToplevel(self); win.title(f"Client list: {name}"); win.geometry("720x520"); win.transient(self)
        win.grab_set(); win.after(100, win.lift)
        ctk.CTkLabel(win, text=f"Read {len(rows)} acronyms from {name}. Check a few, then add the list.",
                     font=("Arial", 13, "bold")).pack(anchor="w", padx=12, pady=(12, 4))
        box = ctk.CTkTextbox(win, wrap="none", font=("Consolas", 12)); box.pack(fill="both", expand=True, padx=12, pady=4)
        box.insert("1.0", "\n".join(f"{t:<16} {dfn}" for t, dfn in rows[:400]) + ("\n..." if len(rows) > 400 else ""))
        box.configure(state="disabled")
        result = {"ok": False}
        btns = ctk.CTkFrame(win, fg_color="transparent"); btns.pack(fill="x", padx=12, pady=10)
        ctk.CTkButton(btns, text="Add This List", command=lambda: (result.update(ok=True), win.destroy())).pack(side="left")
        ctk.CTkButton(btns, text="Cancel", fg_color="#475569", command=win.destroy).pack(side="right")
        self.wait_window(win)
        return result["ok"]

    def open_client_lists(self):
        diag.activity("client_lists_opened")
        if getattr(self, "client_win", None) is not None and self.client_win.winfo_exists():
            self.client_win.focus(); return
        win = ctk.CTkToplevel(self); self.client_win = win
        win.title("Client Acronym Lists"); win.geometry("680x500"); win.transient(self); win.after(100, win.lift)
        where = (f"Saved in the project file ({self.project.name})." if self.project is not None
                 else f"Saved with the effort ({self.effort.name}) and used on every scan." if self.effort is not None
                 else "Kept for scans with no effort open.")
        ctk.CTkLabel(win, justify="left", wraplength=640, text=(
            "Acronym lists from the client (for example an RFP's acronym list). When lists are added:\n"
            "  - their terms are found in your documents even if they are not in the database\n"
            "  - the client definition is selected first; database meanings that differ stay available to pick\n"
            "  - rows are marked 'Client list' (purple) or 'Client list (database differs)' (pink)\n"
            "With no lists added, scans work exactly as before. " + where)).pack(anchor="w", padx=12, pady=(12, 6))
        lb = tk.Listbox(win, height=10, font=("Segoe UI", 11), activestyle="none", exportselection=False)
        lb.pack(fill="both", expand=True, padx=12, pady=4)
        msg = ctk.CTkLabel(win, text="", text_color="#6D28D9", wraplength=640, justify="left")
        def refresh():
            lb.delete(0, "end")
            for lst in self.client_lists:
                lb.insert("end", f"{lst['name']}   ({len(lst['rows'])} acronyms, added {lst.get('added', '')})")
        def changed(text):
            set_client_lists(self.client_lists); self.set_client_lists(self.client_lists)
            self.save_project(); refresh(); msg.configure(text=text)
            if self.project is None:
                try:
                    self.ws.set("client_lists", self.client_lists)   # remembered by the effort (2.14.0)
                except Exception:
                    diag.swallowed()
            if self.results and messagebox.askyesno("Rescan?", "Client lists changed. Rescan now to apply them?", parent=win):
                if self.project is not None and not self.project.read_only:
                    self.rescan_project()
                elif self.queued_files:
                    self.start_batch(self.queued_files)
        def add():
            paths = filedialog.askopenfilenames(parent=win, title="Add client acronym list(s)", initialdir=self.start_dir(),
                                                filetypes=[("Acronym lists", "*.csv *.tsv *.txt *.xlsx *.docx *.pdf")])
            added = []
            for path in paths:
                try:
                    rows = parse_client_list(path)
                except Exception as e:
                    log_error("Client list could not be read", ext=os.path.splitext(path)[1].lower())
                    messagebox.showerror("Not read", f"{os.path.basename(path)}: {type(e).__name__}: {e}", parent=win); continue
                if not rows:
                    messagebox.showwarning("No acronyms found", f"No acronym / definition pairs were found in {os.path.basename(path)}.\n"
                                           "Lists work best as a two-column table (acronym, definition).", parent=win); continue
                if self._preview_client_list(os.path.basename(path), rows):
                    self.client_lists.append({"name": os.path.basename(path), "rows": rows,
                                              "added": datetime.datetime.now().strftime("%Y-%m-%d %H:%M")})
                    added.append(f"{os.path.basename(path)} ({len(rows)})")
            if added:
                changed("Added: " + ", ".join(added))
        def remove():
            sel = lb.curselection()
            if not sel: return
            gone = self.client_lists.pop(sel[0])
            changed(f"Removed {gone['name']}.")
        def view():
            sel = lb.curselection()
            if not sel: return
            lst = self.client_lists[sel[0]]
            v = ctk.CTkToplevel(win); v.title(lst["name"]); v.geometry("700x520"); v.transient(win); v.after(100, v.lift)
            box = ctk.CTkTextbox(v, wrap="none", font=("Consolas", 12)); box.pack(fill="both", expand=True, padx=12, pady=12)
            box.insert("1.0", "\n".join(f"{t:<16} {dfn}" for t, dfn in lst["rows"])); box.configure(state="disabled")
        btns = ctk.CTkFrame(win, fg_color="transparent"); btns.pack(fill="x", padx=12, pady=8)
        ctk.CTkButton(btns, text="Add List...", fg_color="#6D28D9", command=add).pack(side="left", padx=(0, 6))
        ctk.CTkButton(btns, text="View Terms", fg_color="#475569", command=view).pack(side="left", padx=6)
        ctk.CTkButton(btns, text="Remove", fg_color="#DC2626", command=remove).pack(side="left", padx=6)
        msg.pack(anchor="w", padx=12, pady=(0, 10))
        refresh()

def run_selftest(paths):
    os.makedirs(LOG_DIR, exist_ok=True)
    out = os.path.join(LOG_DIR, "selftest.txt")
    with open(out, "w", encoding="utf-8") as f:
        try:
            res = analyze_batch(paths, lambda m: f.write(m + "\n"))
            f.write(f"\nOCR models: {OCR_MODEL_INFO}\nText finder: {doctr_info()}\nMath reader: {math_info()}\nKnowledge: {knowledge_status()}\nOCR settings: {json.dumps(OCR_CFG)}\n")
            import openpyxl.xml as _oxml
            f.write(f"Security: network guard={socket.socket.connect is _guarded_connect}, "
                    f"office XML entities disabled={_SAFE_XML is not None}, openpyxl defusedxml={_oxml.DEFUSEDXML}\n")
            if trainer_available():
                import trainer_gui, tess_train   # the Trainer window and Improve Reading load inside this program
                f.write(f"Training features: on (Trainer window and Improve Reading loaded; "
                        f"{len(tess_train.read_samples(TRAIN_DIR))} saved lines)\n")
            else:
                f.write("Training features: in Acronym Hunter Training; this program saves corrections only"
                        + ("" if APP_CFG["allow_training"] else " (off: allow_training = false)") + "\n")
            if OCR_BACKEND == "tesseract":
                f.write(f"OCR sandbox: {OCR_SANDBOX_STATUS}\n")
                if APP_CFG["ocr_sandbox"] and os.name == "nt":
                    import sandbox
                    f.write(f"Sandbox network probe (Windows curl.exe run in the same sandbox): {sandbox.network_probe()}\n")
            f.write(f"\nRESULTS: {len(res)} terms\n")
            for term, d in res.items():
                f.write(f"{term}\t{d['category']}\t{d['count']}\t{d['status']}\t{d['selected_definition']}\n")
            f.write("SELFTEST OK\n")
        except Exception:
            f.write(traceback.format_exc())
            f.write("SELFTEST FAILED\n")

if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--selftest":
        run_selftest(sys.argv[2:])
        sys.exit(0)
    if len(sys.argv) > 2 and sys.argv[1] == "--install-knowledge":   # used by build.ps1 -KnowledgePack
        os.makedirs(LOG_DIR, exist_ok=True)
        with open(os.path.join(LOG_DIR, "knowledge_install.txt"), "w", encoding="utf-8") as f:
            try:
                f.write(install_knowledge_pack(sys.argv[2]) + "\nINSTALL OK\n")
            except Exception:
                f.write(traceback.format_exc() + "INSTALL FAILED\n")
        sys.exit(0)
    try:
        app = AcronymAppGUI()
        app.mainloop()
    except Exception:
        log_error("Application failed to start")
        raise