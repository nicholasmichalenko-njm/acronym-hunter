"""Submission requirements (2.39.0): what to deliver and how, read from the customer documents and tracked across their
versions. Rules and counts only (No AI); nothing leaves the PC.

From every version of every customer document (the text copies kept by docversions), sentences that set:
- Due date and time (with the time zone when written), and the deadline for questions;
- Page limits, per volume when the sentence names one ("Volume I", "Technical Volume");
- Format: font and point size, margins, page size, line spacing;
- Files: file format (searchable PDF, Excel, Word), file naming, copies, delivery method;
- Deliverables: volumes, attachments (Attachment J-3), forms (SF 33), representations and certifications, plans.

Each requirement has a key (kind and subject, such as "Page limit | Volume I") so it can be followed from version to
version: its history lists each version's value and sentence, and a change shows until someone marks it reviewed.
The person can confirm, edit (the edit wins over what the documents say), hide (not a requirement) or add requirements.
The register's state is saved in the effort's folder: documents\\register.json.
"""
import datetime, getpass, re

KINDS = ("Due", "Questions due", "Page limit", "Format", "File", "Copies", "Delivery", "Contact", "Deliverable")   # Contact: 2.49.0
ROMAN = {"i": 1, "ii": 2, "iii": 3, "iv": 4, "v": 5, "vi": 6, "vii": 7, "viii": 8}
MONTHS = r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
DATE = re.compile(r"\b" + MONTHS + r"\.?\s+\d{1,2}(?:st|nd|rd|th)?,?\s+\d{4}\b|\b\d{1,2}/\d{1,2}/\d{2,4}\b|\b\d{4}-\d{2}-\d{2}\b|"
                  r"\b\d{1,2}\s+" + MONTHS + r"\.?\s+\d{4}\b", re.I)
ZONE = r"(?:\s*\(?(?:(?:Mountain|Eastern|Central|Pacific|Alaska|Hawaii)(?:\s+(?:Standard|Daylight))?\s+Time|local time|(?-i:[ECMP][SD]?T|UTC|GMT|Zulu|Z))\)?)?"
TIME = re.compile(r"\b\d{1,2}(?::\d{2})?\s*(?:a\.?m\.?|p\.?m\.?)" + ZONE + r"|\b\d{3,4}\s*(?:hours|hrs)\b" + ZONE, re.I)
DUE = re.compile(r"\b(?:due|no later than|nlt|deadline|must be received|shall be received|submitted by|closing date)\b", re.I)
QUESTIONS = re.compile(r"\b(?:questions?|clarifications?|inquir(?:y|ies))\b", re.I)
PAGES = re.compile(r"\b(?:not\s+(?:to\s+)?exceed|limited\s+to|maximum\s+of|no\s+more\s+than|up\s+to|shall\s+not\s+be\s+more\s+than)\s+"
                   r"(?:\w+\s+)?\(?(\d{1,3})\)?\s*(?:-\s*)?(?:single[- ]sided\s+)?pages?\b|\b(\d{1,3})[- ]page\s+(?:limit|maximum)\b", re.I)
VOLUME = re.compile(r"\bvolume\s+(viii|vii|vi|iv|v|iii|ii|i|[1-8])\b(?:\s*[-–—:,(]?\s*((?-i:(?!Volume\b)[A-Z][a-z]+)(?:\s(?-i:(?!Volume\b)[A-Z][a-z]+)){0,2}))?|"
                    r"\b(technical|management|cost|price|past performance|small business|executive summary|oral)\s+(?:volume|proposal)\b", re.I)
FONT = re.compile(r"\b(\d{1,2})[- ]?(?:point|pt)\b[^.]{0,40}?\b(Times New Roman|Arial|Calibri|Helvetica|Garamond|Cambria|Courier(?: New)?|Verdana|Tahoma)?|"
                  r"\b(Times New Roman|Arial|Calibri|Helvetica|Garamond|Cambria|Courier(?: New)?)\b", re.I)
FACE = re.compile(r"\b(?:Times New Roman|Arial|Calibri|Helvetica|Garamond|Cambria|Courier(?: New)?|Verdana|Tahoma)\b", re.I)
MARGIN = re.compile(r"\b(\d(?:\.\d+)?|one|one-half)[- ]?inch(?:es)?\s+margins?|\bmargins?\s+(?:of|shall be)\s+(?:at least\s+)?(\d(?:\.\d+)?)[- ]?inch", re.I)
PAGESIZE = re.compile(r"\b8[.,]5\s*(?:x|by|×)\s*11\b|\bletter[- ]size\b|\bA4\b", re.I)
SPACING = re.compile(r"\b(single|double|1\.5)[- ]spaced\b|\b(single|double)\s+spacing\b", re.I)
FILEFMT = re.compile(r"\bsearchable\s+PDF\b|\b(?:in|as)\s+(?:a\s+)?(?:PDF|Adobe PDF|Microsoft Excel|Excel|Microsoft Word|Word)\b(?:\s+(?:format|files?))?|"
                     r"\.(?:pdf|xlsx|docx)\b", re.I)
NAMING = re.compile(r"\b(?:file\s+names?|naming\s+convention|shall\s+be\s+named|name\s+(?:each|the)\s+files?)\b", re.I)
COPIES = re.compile(r"\b(\w+)\s+\(?(\d+)?\)?\s*(?:hard|paper|printed|electronic)\s+cop(?:y|ies)\b|\b(\d+)\s+cop(?:y|ies)\b", re.I)
DELIVERY = re.compile(r"\b(?:submitted|delivered|sent|uploaded|e-?mailed)\s+(?:electronically\s+)?(?:via|to|through|by|using)\s+([^.;]{3,80})", re.I)
ATTACH = re.compile(r"\b(?:Attachment|Exhibit|Enclosure|Annex)\s+[A-Z]?-?\d{1,3}[a-z]?\b(?:\s*,?\s*(?:the\s+)?((?:[A-Z][a-z]+\s?){1,4}))?")
FORMS = re.compile(r"\b(?:SF|DD|Standard Form)\s?-?\s?\d{2,4}\b|\bRepresentations and Certifications\b|\bSmall Business Subcontracting Plan\b|"
                   r"\bCover (?:letter|page)\b", re.I)
EMAIL = re.compile(r"\b[\w.+-]+@[\w-]+(?:\.[\w-]+)+\b")
CONTACT_ROLE = re.compile(r"\b(Contracting Officer|Contract Specialist|Contracting Specialist|Point of Contact|Buyer)\b(?:\s*\([A-Z]{2,4}\))?"
                          r"(?:\s*,\s*|\s+)((?:(?:Mr|Ms|Mrs|Dr)\.?\s+)?[A-Z][a-z]+(?:\s+[A-Z]\.)?\s+[A-Z][a-zA-Z'-]+)?")
SUBMIT = re.compile(r"\b(?:shall|must|will)\s+(?:submit|provide|include|deliver|complete|sign|furnish|consist|contain)\b|\bsubmit(?:ted)?\b|\binclude[sd]?\b|\brequired\b", re.I)


def volume_key(text):
    """'Volume II - Management' or 'Management Volume' -> 'Volume II' / 'Management volume'."""
    m = VOLUME.search(text or "")
    if not m:
        return ""
    if m.group(1):
        n = m.group(1).lower()
        num = ROMAN.get(n, int(n) if n.isdigit() else 0)
        roman = {v: k for k, v in ROMAN.items()}.get(num, n).upper()
        return f"Volume {roman}"
    return m.group(3).strip().title() + " volume"


def _clean(s):
    return " ".join((s or "").split())


def find(sentence):
    """[(kind, subject, value)] the submission requirements one sentence sets."""
    s = _clean(sentence)
    out = []
    if not s:
        return out
    low = s.lower()
    if DUE.search(s) and DATE.search(s):
        when = DATE.search(s).group(0)
        t = TIME.search(s)
        val = when + (f", {t.group(0)}" if t else "")
        out.append(("Questions due" if QUESTIONS.search(s) else "Due", "Questions" if QUESTIONS.search(s) else "Proposal", val))
    for m in PAGES.finditer(s):
        n = m.group(1) or m.group(2)
        # the volume named closest before the limit, else anywhere in the sentence
        vk = volume_key(s[:m.start()][::-1][:120][::-1]) or volume_key(s) or "Proposal"
        out.append(("Page limit", vk, f"{n} pages"))
    f = FONT.search(s)
    if f and ("font" in low or "point" in low or "pt" in low or f.group(3)):
        size = f.group(1)
        fm = FACE.search(s)
        face = fm.group(0) if fm else ""
        tables = re.search(r"\b(tables?|figures?|graphics?|captions?|footnotes?|charts?)\b", s, re.I)
        subj = "Font in tables and figures" if tables and size and re.search(r"\b(?:may|can|allowed)\b", s, re.I) else "Font"
        val = " ".join(x for x in ((size + "-point") if size else "", face) if x)
        if val:
            out.append(("Format", subj, val))
    m = MARGIN.search(s)
    if m:
        out.append(("Format", "Margins", (m.group(1) or m.group(2)).replace("one-half", "0.5").replace("one", "1") + "-inch margins"))
    m = PAGESIZE.search(s)
    if m:
        out.append(("Format", "Page size", "8.5 x 11" if "8" in m.group(0) else m.group(0)))
    m = SPACING.search(s)
    if m:
        out.append(("Format", "Line spacing", (m.group(1) or m.group(2)).lower() + "-spaced"))
    if FILEFMT.search(s) and SUBMIT.search(s) or re.search(r"\bsearchable\s+PDF\b", s, re.I):
        names = {"pdf": "PDF", "searchable pdf": "Searchable PDF", "excel": "Excel", "word": "Word", "xlsx": "Excel", "docx": "Word"}
        fmts = sorted({names.get(" ".join(x.strip(".").lower().split()), x) for x in
                       re.findall(r"searchable\s+PDF|\bPDF\b|\bExcel\b|\bWord\b|\.pdf\b|\.xlsx\b|\.docx\b", s, re.I)})
        if "Searchable PDF" in fmts and "PDF" in fmts:
            fmts.remove("PDF")
        if fmts:
            out.append(("File", "File format", ", ".join(fmts)))
    if NAMING.search(s):
        m = re.search(r"\b[\w\[\]<>{}-]+(?:_[\w\[\]<>{}-]+)+\.(?:pdf|docx|xlsx)\b", s, re.I)
        out.append(("File", "File naming", m.group(0) if m else s[:120]))
    m = COPIES.search(s)
    if m and SUBMIT.search(s):
        out.append(("Copies", "Copies", m.group(0)))
    em = EMAIL.search(s)                               # 2.49.0: who receives the proposal or the questions
    if em and (SUBMIT.search(s) or QUESTIONS.search(s) or CONTACT_ROLE.search(s)):
        role = CONTACT_ROLE.search(s)
        who = " ".join(x for x in ((role.group(2) or "") if role else "", f"({role.group(1)})" if role else "") if x)
        out.append(("Contact", "Questions contact" if QUESTIONS.search(s) and not re.search(r"\bproposals?\b", s, re.I)
                    else "Submission contact", (who + " " + em.group(0)).strip()))
    m = DELIVERY.search(s)
    if m and re.search(r"\bproposals?|\bvolumes?|\bofferors?|\bresponses?|\bquotes?", s, re.I):
        how = re.split(r"\s+(?:to the|to|at|no later than|not later than|by|on or before)\s+", _clean(m.group(1)), maxsplit=1)[0]
        out.append(("Delivery", "Delivery", (how or _clean(m.group(1)))[:80]))
    if SUBMIT.search(s):
        for m in ATTACH.finditer(s):
            name = re.match(r"\w+\s+[A-Z]?-?\d{1,3}[a-z]?", m.group(0)).group(0)
            what = re.search(r"\b(?:provide|submit|include|complete|deliver|furnish|sign)\s+(?:an?\s+|the\s+)?(.{3,90}?)\s+(?:as|in|using)\s+"
                             + re.escape(name), s, re.I)
            desc = what.group(1) if what else (m.group(1) or "").strip()
            out.append(("Deliverable", name, f"{name}: {desc}" if desc else name))
        for m in FORMS.finditer(s):
            out.append(("Deliverable", m.group(0).strip().title() if not m.group(0)[:2].isupper() else m.group(0).strip(),
                        m.group(0).strip()))
        vk = volume_key(s)
        if vk and re.search(r"\b(?:shall|must)\s+(?:submit|include|consist|contain)|\bproposal shall (?:consist|be organized)", s, re.I):
            for m in VOLUME.finditer(s):
                k = volume_key(m.group(0))
                if k:
                    out.append(("Deliverable", k, _clean(m.group(0))[:60]))
    seen, uniq = set(), []
    for k in out:
        if (k[0], k[1]) not in seen:
            seen.add((k[0], k[1])); uniq.append(k)
    return uniq


def key_of(kind, subject):
    return f"{kind} | {subject}"


# ------------------------------------------------------------------ the register
RULES = 1                      # 2.82.0: bump when find() changes, so saved results are worked out again
_FOUND = {}                    # (library folder, version id, fingerprint) -> what that version's text sets
_BUILT = {}                    # (library folder, library signature) -> build(lib)


def _found_for(lib, v):
    """What one version sets: {key: (kind, subject, value, sentence, loc)}. A version's text never changes, so this is
    worked out once, kept in memory and in documents\\found\\<version>.json, and read back after that (2.82.0)."""
    import os
    import workspace as W
    key = (lib.folder, v["id"], v.get("sha256", ""))
    if key in _FOUND:
        return _FOUND[key]
    p = os.path.join(lib.folder, "found", v["id"] + ".json")
    d = W.read_json(p, None)
    if isinstance(d, dict) and d.get("rules") == RULES and d.get("sha") == v.get("sha256", "") and isinstance(d.get("found"), dict):
        found = {k: tuple(x) for k, x in d["found"].items()}
    else:
        found = {}
        for s in lib.text(v):
            for kind, subj, val in find(s["x"]):
                k = key_of(kind, subj)
                if k not in found:
                    found[k] = (kind, subj, val, s["x"], s.get("l", ""))
        try:
            os.makedirs(os.path.dirname(p), exist_ok=True)
            W.write_json(p, {"rules": RULES, "sha": v.get("sha256", ""), "found": {k: list(x) for k, x in found.items()}})
        except OSError:
            pass
    _FOUND[key] = found
    return found


def build(lib):
    """Read every version of every document in the library. Returns {key: requirement} where requirement is
    {"key", "kind", "subject", "value", "sentence", "doc", "doc_id", "version", "version_id", "loc", "history": [...],
     "changed": bool, "removed": bool}. History entries: {"doc", "version", "version_id", "added", "value", "sentence",
     "loc", "change": "first"|"changed"|"same"|"removed"}."""
    reqs = {}
    for d in lib.docs:
        prev = {}
        for v in d["versions"]:
            found = _found_for(lib, v)
            for k, (kind, subj, val, sent, loc) in found.items():
                r = reqs.setdefault(k, {"key": k, "kind": kind, "subject": subj, "history": []})
                before = prev.get(k)
                if before is None:
                    change = "first" if not r["history"] else ("same" if r["history"][-1]["value"] == val else "changed")
                else:
                    change = "same" if before == val else "changed"
                r["history"].append({"doc": d["name"], "version": v["label"], "version_id": v["id"], "added": v.get("added", ""),
                                     "value": val, "sentence": sent, "loc": loc, "change": change, "file": v.get("file", ""),
                                     "sha256": v.get("sha256", ""), "who": v.get("who", "")})
            for k in prev:
                if k not in found and k in reqs:
                    reqs[k]["history"].append({"doc": d["name"], "version": v["label"], "version_id": v["id"],
                                               "added": v.get("added", ""), "value": "", "sentence": "", "loc": "",
                                               "change": "removed", "file": v.get("file", ""), "sha256": v.get("sha256", ""),
                                               "who": v.get("who", "")})
            prev = {k: x[2] for k, x in found.items()}
    for r in reqs.values():
        r["history"].sort(key=lambda h: h["added"])
        last = next((h for h in reversed(r["history"]) if h["change"] != "removed"), r["history"][-1])
        r.update(value=last["value"], sentence=last["sentence"], doc=last["doc"], version=last["version"],
                 version_id=last["version_id"], loc=last["loc"])
        r["removed"] = r["history"][-1]["change"] == "removed"
        r["changed"] = any(h["change"] in ("changed", "removed") for h in r["history"])
    return reqs


class Register:
    """The requirements of the open effort with the person's decisions (saved in documents\\register.json)."""

    def __init__(self, lib):
        self.lib = lib
        self.state = {"items": {}, "manual": []}
        self._built = (None, {})
        self.load()

    def found(self):
        """build(lib), kept until the library gains, loses or renames a version."""
        sig = tuple((d["name"], v["id"], v["label"], v.get("added", "")) for d in self.lib.docs for v in d["versions"])
        if self._built[0] != sig:
            key = (self.lib.folder, sig)              # 2.82.0: kept across pages, so moving between them is instant
            if key not in _BUILT:
                if len(_BUILT) > 8:
                    _BUILT.clear()
                _BUILT[key] = build(self.lib)
            self._built = (sig, _BUILT[key])
        return self._built[1]

    @property
    def path(self):
        import os
        return os.path.join(self.lib.folder, "register.json")

    def load(self):
        import workspace as W
        d = W.read_json(self.path, None) or {}
        self.state = {"items": d.get("items", {}), "manual": d.get("manual", [])}
        return self

    def save(self):
        import workspace as W
        W.write_json(self.path, {"v": 1, **self.state})

    def _st(self, key):
        return self.state["items"].setdefault(key, {})

    def rows(self, include_hidden=False):
        """Every requirement, in kind order: found ones with the person's edits, then the ones the person added."""
        out = []
        for k, r in self.found().items():
            st = self.state["items"].get(k, {})
            if st.get("hidden") and not include_hidden:
                continue
            row = dict(r)
            row["found"] = r["value"]
            if st.get("edit"):
                row["value"] = st["edit"]
            row["confirmed"] = st.get("confirmed", "")
            row["hidden"] = bool(st.get("hidden"))
            last_change = next((h for h in reversed(r["history"]) if h["change"] in ("changed", "removed")), None)
            row["to_review"] = bool(last_change) and st.get("reviewed_version") != last_change["version_id"]
            row["reviewed"] = st.get("reviewed", "")
            row["manual"] = False
            out.append(row)
        for m in self.state["manual"]:
            out.append({"key": m["key"], "kind": m["kind"], "subject": m["subject"], "value": m["value"], "found": "",
                        "sentence": "", "doc": "Typed in", "version": "", "version_id": "", "loc": "", "history": [],
                        "changed": False, "removed": False, "confirmed": m.get("who", ""), "hidden": False,
                        "to_review": False, "reviewed": "", "manual": True})
        order = {k: i for i, k in enumerate(KINDS)}
        out.sort(key=lambda r: (order.get(r["kind"], 99), r["subject"]))
        return out

    def stamp(self, who=None):
        return f"{datetime.date.today().isoformat()} {who or getpass.getuser()}"

    def review(self, key, who=None):
        r = self.found().get(key)
        if r is None:
            return False
        last_change = next((h for h in reversed(r["history"]) if h["change"] in ("changed", "removed")), None)
        st = self._st(key)
        st["reviewed_version"] = last_change["version_id"] if last_change else ""
        st["reviewed"] = self.stamp(who)
        self.save()
        return True

    def confirm(self, key, who=None):
        self._st(key)["confirmed"] = self.stamp(who)
        self.save()
        return True

    def edit(self, key, value, who=None):
        st = self._st(key)
        value = _clean(value)
        if value:
            st["edit"] = value
            st["confirmed"] = self.stamp(who)
        else:
            st.pop("edit", None)
        self.save()
        return True

    def hide(self, key, hidden=True):
        st = self._st(key)
        if hidden:
            st["hidden"] = True
        else:
            st.pop("hidden", None)
        self.save()
        return True

    def add(self, kind, subject, value, who=None):
        subject, value = _clean(subject), _clean(value)
        if not value:
            return None
        k = key_of(kind, subject or kind)
        self.state["manual"] = [m for m in self.state["manual"] if m["key"] != k]
        self.state["manual"].append({"key": k, "kind": kind, "subject": subject or kind, "value": value, "who": self.stamp(who)})
        self.save()
        return k

    def remove_manual(self, key):
        n = len(self.state["manual"])
        self.state["manual"] = [m for m in self.state["manual"] if m["key"] != key]
        self.save()
        return len(self.state["manual"]) < n


def summary(rows):
    """What the card and the bar show: due date, page limits, format line, file line, deliverables, changes."""
    by = {}
    for r in rows:
        by.setdefault(r["kind"], []).append(r)
    due = next((r for r in by.get("Due", [])), None)
    pages = [(r["subject"], r["value"], r) for r in by.get("Page limit", [])]
    fmt = [r for r in by.get("Format", [])]
    files = [r for r in by.get("File", [])]
    return {"due": due, "pages": pages, "format": fmt, "files": files, "deliverables": by.get("Deliverable", []),
            "contact": by.get("Contact", []), "delivery": by.get("Delivery", []), "copies": by.get("Copies", []),
            "questions_due": next((r for r in by.get("Questions due", [])), None),
            "to_review": [r for r in rows if r.get("to_review")], "count": len(rows)}


def days_until(text, today=None):
    """Days from today to a written date ('December 15, 2026', '12/15/2026', '2026-12-15'), None when unreadable."""
    today = today or datetime.date.today()
    t = (text or "").split(",")
    s = ",".join(t[:2]) if len(t) > 1 and re.search(r"\d{4}", t[1]) else t[0]
    for fmt in ("%B %d, %Y", "%b %d, %Y", "%B %d %Y", "%b %d %Y", "%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%d %B %Y", "%d %b %Y"):
        try:
            return (datetime.datetime.strptime(re.sub(r"(\d)(st|nd|rd|th)", r"\1", s.strip().replace(".", "")), fmt).date() - today).days
        except ValueError:
            continue
    return None


def match_volume(subject, name):
    """Does a page-limit subject ('Volume I', 'Technical volume') belong to an effort volume or file name?"""
    a, b = (subject or "").lower(), (name or "").lower()
    if not a or a == "proposal":
        return False
    m = re.match(r"volume\s+(\w+)", a)
    if m:
        num = ROMAN.get(m.group(1), int(m.group(1)) if m.group(1).isdigit() else 0)
        for tok in re.findall(r"(?:vol(?:ume)?[\s_.-]*)(viii|vii|vi|iv|v|iii|ii|i|[1-8])(?![a-z0-9])", b):
            if ROMAN.get(tok, int(tok) if tok.isdigit() else -1) == num:
                return True
        return False
    word = a.replace(" volume", "")
    return bool(word) and word in b


# ------------------------------------------------------------------ QA/QC: our volumes against the page limits
def word_pages(path):
    """The page count Word saved in the file (docProps/app.xml), None when Word has not counted pages."""
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            x = z.read("docProps/app.xml").decode("utf-8", "replace")
    except (KeyError, OSError, zipfile.BadZipFile):
        return None
    m = re.search(r"<(?:\w+:)?Pages>(\d+)</(?:\w+:)?Pages>", x)
    return int(m.group(1)) if m else None


def _limit_n(value):
    m = re.search(r"\d+", value or "")
    return int(m.group(0)) if m else None


def page_check(paths, rows, volumes=()):
    """Rows for the QA/QC Page limits tab: each Word volume against the limit for its volume, and all volumes against a
    limit for the whole proposal. paths {file name: path}; rows from Register.rows(); volumes [{"name", "files"}]."""
    limits = [r for r in rows if r["kind"] == "Page limit" and _limit_n(r["value"])]
    if not limits:
        return []
    out, total, counted = [], 0, True
    for name, p in paths.items():
        names = [name] + [v["name"] for v in volumes if name in v.get("files", [])]
        n = word_pages(p)
        if n is None:
            counted = False
        else:
            total += n
        for r in limits:
            if r["subject"].lower() == "proposal" or not any(match_volume(r["subject"], x) for x in names):
                continue
            lim = _limit_n(r["value"])
            if n is None:
                lv, res = "info", "Word has not counted the pages: open the file in Word and save it, then check again"
            elif n > lim:
                lv, res = "bad", f"Over by {n - lim} page{'s' if n - lim != 1 else ''}"
            else:
                lv, res = "ok", f"Within the limit ({lim - n} to spare)"
            out.append({"level": lv, "volume": name, "limit_for": r["subject"], "limit": r["value"],
                        "pages": "" if n is None else n, "result": res, "source": f"{r['doc']}, {r['version']}".strip(", ")})
    for r in limits:
        if r["subject"].lower() != "proposal":
            continue
        lim = _limit_n(r["value"])
        if not counted:
            lv, res = "info", "Not every file has a page count from Word: open each in Word and save it, then check again"
        elif total > lim:
            lv, res = "bad", f"Over by {total - lim} page{'s' if total - lim != 1 else ''} (all volumes together)"
        else:
            lv, res = "ok", f"Within the limit ({lim - total} to spare, all volumes together)"
        out.append({"level": lv, "volume": "All volumes", "limit_for": "Proposal", "limit": r["value"], "pages": total if counted else "",
                    "result": res, "source": f"{r['doc']}, {r['version']}".strip(", ")})
    return out


def export_excel(path, rows, effort_name=""):
    """Submission requirements (one row each, with status) and their history across versions."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()

    def sheet(sh, cols, data, widths):
        sh.append(cols)
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1E3A8A")
        for r, fill in data:
            sh.append(r)
            if fill:
                for c in sh[sh.max_row]:
                    c.fill = PatternFill("solid", fgColor=fill)
        for i, w in enumerate(widths, 1):
            sh.column_dimensions[sh.cell(row=1, column=i).column_letter].width = w
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"
        if sh.max_row > 1:
            sh.auto_filter.ref = f"A1:{sh.cell(row=1, column=len(cols)).column_letter}{sh.max_row}"

    def status(r):
        if r["manual"]:
            return "Typed in " + r["confirmed"]
        if r["to_review"]:
            return "Changed: not reviewed yet"
        if r["hidden"]:
            return "Hidden (not a requirement)"
        if r["removed"]:
            return "No longer in the documents"
        return ("Confirmed " + r["confirmed"]) if r["confirmed"] else "Found"

    ws = wb.active; ws.title = "Submission requirements"
    sheet(ws, ["Kind", "Requirement", "Says", "Status", "Document", "Version", "Where", "Sentence", "Documents say (when edited)", "Reviewed"],
          [([r["kind"], r["subject"], r["value"], status(r), r["doc"], r["version"], r["loc"], r["sentence"],
             r["found"] if r["found"] and r["found"] != r["value"] else "", r["reviewed"]],
            "FEF3C7" if r["to_review"] else "E0F2FE" if r["manual"] else None) for r in rows],
          [14, 24, 40, 24, 24, 16, 30, 70, 30, 22])
    hs = wb.create_sheet("History")
    word = {"first": "First found", "same": "No change", "changed": "Changed", "removed": "Removed"}
    sheet(hs, ["Kind", "Requirement", "Document", "Version", "Added", "Added by", "File", "Fingerprint (SHA-256)", "Change", "Says", "Sentence"],
          [([r["kind"], r["subject"], h["doc"], h["version"], h["added"].replace("T", " "), h["who"], h["file"], h["sha256"],
             word[h["change"]], h["value"], h["sentence"]], "FEF3C7" if h["change"] in ("changed", "removed") else None)
           for r in rows for h in r["history"]],
          [14, 24, 24, 16, 17, 14, 30, 30, 12, 34, 70])
    if effort_name:
        wb.properties.title = f"{effort_name} submission requirements"
    wb.save(path)
    return path


def summary_text(rows):
    """One line: due date (days left), page limits, format, files, changes to review (2.63.0, from the classic bar)."""
    if not rows:
        return ""
    s = summary(rows)
    bits = []
    if s["due"]:
        n = days_until(s["due"]["value"])
        bits.append(f"Due {s['due']['value']}" + ("" if n is None else " (due date passed)" if n < 0 else
                                                  " (today)" if n == 0 else f" ({n} day{'s' if n != 1 else ''})"))
    if s["pages"]:
        bits.append("Pages: " + ", ".join(f"{subj} {_limit_n(v)}" for subj, v, _r in s["pages"]))
    f = [r["value"] for r in s["format"] if r["subject"] in ("Font", "Margins")]
    if f:
        bits.append(", ".join(f))
    ff = [r["value"] for r in s["files"] if r["subject"] == "File format"]
    if ff:
        bits.append(ff[0])
    if s["to_review"]:
        n = len(s["to_review"])
        bits.append(f"{n} change{'s' if n != 1 else ''} to review")
    return "  |  ".join(bits)
