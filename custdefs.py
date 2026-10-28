"""Definitions written in customer documents (2.5.0): "Spelled Out Words (ACR)", "ACR (Spelled Out Words)", and
glossary or acronym-list tables. The customer's own definition outranks the company database.

Matching follows the Schwartz & Hearst method (2003): the long form ends right before the parenthesis, and the short
form's characters appear in it in order, the first one starting a word. Glossary lines are accepted inside a section
titled Acronyms / Abbreviations / Glossary / Definitions, or anywhere when the definition's word initials fit the
term. Pure functions: no files are written, nothing leaves the PC.
"""
import re

SF_RE = r"[A-Za-z0-9][A-Za-z0-9&/\-.]{0,11}"
_PAREN = re.compile(r"\(\s*(" + SF_RE + r")s?\s*\)")
_PAREN_LONG = re.compile(r"\b(" + SF_RE + r")\s*\(\s*([^()]{4,120}?)\s*\)")
_GLOSS_HEAD = re.compile(r"^\s*(?:\d+(?:\.\d+)*\.?\s+)?(?:list of\s+)?(acronyms?|abbreviations?|glossary|definitions|acronyms and abbreviations|"
                         r"abbreviations and acronyms|terms and definitions)\b", re.I)
_GLOSS_LINE = re.compile(r"^\s*(" + SF_RE + r")\s*(?:\t+|\s{2,}|\s*[-–—:=|]\s*)\s*([A-Za-z][^\n]{2,140}?)\s*$")
_STOP = {"a", "an", "and", "of", "the", "for", "to", "in", "on", "by", "with", "at", "or", "as", "from"}
_GLOSS_LINE1 = re.compile(r"^\s*(" + SF_RE + r")\s+([A-Za-z][^\n]{2,140}?)\s*$")   # single space: only inside a glossary


def _is_sf(sf):
    """A plausible short form: has a capital or digit, not a plain lowercase word, at most 12 characters."""
    s = sf.strip(".")
    if not (2 <= len(s) <= 12) or not s[0].isalnum():
        return False
    caps = sum(c.isupper() for c in s)
    if caps == 0 and not any(c.isdigit() for c in s):
        return False
    if s.isdigit() or re.fullmatch(r"\d+(\.\d+)?%?", s) or re.fullmatch(r"[A-Z][a-z]+", s):   # "2.1", "Figure"
        return False
    return True


def _best_long_form(sf, words):
    """Schwartz-Hearst: walk the short form backwards through the candidate words."""
    s = [c.lower() for c in sf if c.isalnum()]
    if not s:
        return None
    text = " ".join(words)
    si, li = len(s) - 1, len(text) - 1
    while si >= 0:
        c = s[si]
        while li >= 0 and (text[li].lower() != c or (si == 0 and li > 0 and text[li - 1].isalnum())):
            li -= 1
        if li < 0:
            return None
        si -= 1; li -= 1
    start = text.rfind(" ", 0, li + 1) + 1
    lf = text[start:].strip(" ,;:-")
    if not lf or len(lf.split()) > len(s) + 5 or len(lf) <= len(sf):
        return None
    if lf.lower() == sf.lower() or lf.split()[0].lower() in _STOP and len(lf.split()) > 1 and lf.split()[0][0].lower() != s[0]:
        return None
    return lf


def _initials_fit(sf, definition):
    """Share of the short form's letters that match word initials of the definition, in order."""
    letters = [c.lower() for c in sf if c.isalpha()]
    words = [w for w in re.findall(r"[A-Za-z][A-Za-z0-9\-]*", definition) if w.lower() not in _STOP]
    if not letters or not words:
        return 0.0
    initials = "".join(w[0].lower() for w in words) + "".join(p[0].lower() for w in words for p in w.split("-")[1:] if p)
    hit, j = 0, 0
    for ch in letters:
        k = initials.find(ch, j)
        if k >= 0:
            hit += 1; j = k + 1
    return hit / len(letters)


def _sentence(text, a, b):
    s = max(text.rfind(".", 0, a), text.rfind("\n", 0, a)) + 1
    e_dot, e_nl = text.find(".", b), text.find("\n", b)
    ends = [x for x in (e_dot, e_nl) if x >= 0]
    e = min(ends) + 1 if ends else len(text)
    return " ".join(text[s:e].split())[:400]


def extract(chunks):
    """[(location label, text)] -> {term: [{"definition", "where", "sentence", "how"}]} (first occurrence first)."""
    found = {}

    def add(sf, lf, where, sentence, how):
        sf = sf.strip().rstrip(".")
        lf = " ".join(lf.split()).strip(" ,;:.-")
        if not _is_sf(sf) or len(lf) < 3:
            return
        lst = found.setdefault(sf, [])
        if any(x["definition"].lower() == lf.lower() for x in lst):
            return
        lst.append({"definition": lf, "where": where, "sentence": sentence, "how": how})

    for where, text in chunks:
        if not text:
            continue
        # 1. "Long Form (SF)"
        for m in _PAREN.finditer(text):
            sf = m.group(1)
            if not _is_sf(sf):
                continue
            before = text[max(0, m.start() - 200):m.start()]
            before = re.split(r"[.;:!?\n(]\s", before)[-1]
            words = before.split()[-(min(len(sf) + 5, len(sf) * 2) + 2):]
            lf = _best_long_form(sf, words) if words else None
            if lf:
                add(sf, lf, where, _sentence(text, m.start(), m.end()), "spelled out")
        # 2. "SF (Long Form)"
        for m in _PAREN_LONG.finditer(text):
            sf, inner = m.group(1), m.group(2)
            if _is_sf(sf) and inner[:1].isalpha() and _initials_fit(sf, inner) >= 0.75 and len(inner.split()) <= len(sf) + 5:
                lf = _best_long_form(sf, inner.split())
                if lf:
                    add(sf, lf, where, _sentence(text, m.start(), m.end()), "spelled out")
        # 3. glossary / acronym list lines
        label = where.split("(", 1)[-1] if "(" in where else ""
        in_gloss = bool(re.search(r"acronym|abbreviation|glossary|definition", label, re.I))   # e.g. a sheet named Glossary
        for line in text.splitlines():
            if _GLOSS_HEAD.match(line) and len(line.split()) <= 8:
                in_gloss = True; continue
            m = _GLOSS_LINE.match(line)
            if not m and in_gloss:
                m1 = _GLOSS_LINE1.match(line)
                if m1 and _is_sf(m1.group(1)) and _initials_fit(m1.group(1), m1.group(2)) >= 0.5:
                    m = m1
            if not m:
                if in_gloss and line.strip() and len(line.split()) > 14:
                    in_gloss = False          # back in running text
                continue
            sf, lf = m.group(1), m.group(2)
            if not _is_sf(sf):
                continue
            fit = _initials_fit(sf, lf)
            if (in_gloss and fit >= 0.3) or fit >= 0.8:
                add(sf, lf, where, " ".join(line.split())[:400], "glossary")
    return found


def merge_into(results, defs, folded=None):
    """Attach customer definitions to scan results. Returns the number of terms that got one.
    folded: optional function used to match terms the scanner folded (look-alike characters)."""
    n = 0
    index = {k: v for k, v in defs.items()}
    if folded:
        for k, v in defs.items():
            index.setdefault(folded(k), v)
    for term, d in results.items():
        hits = index.get(term) or (index.get(folded(term)) if folded else None)
        if not hits:
            d.pop("customer_definitions", None); d.pop("custdef_source", None); d.pop("custdef_all", None)
            continue
        d["customer_definitions"] = [h["definition"] for h in hits]
        d["custdef_source"] = {"where": hits[0]["where"], "sentence": hits[0]["sentence"]}
        d["custdef_all"] = hits
        n += 1
    return n


def table_chunks(path):
    """Rows of Excel sheets and Word tables as "cell<TAB>cell" lines (the scanner flattens cells, which loses rows).
    A table where most rows look like "ACR | definition" is labelled an acronym table, so all its rows count."""
    import os, zipfile
    out = []
    name = os.path.basename(path)
    ext = os.path.splitext(path)[1].lower()
    tables = []
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                rows = []
                for r in ws.iter_rows(values_only=True):
                    cells = [str(v).strip() for v in r if v is not None and str(v).strip()]
                    if cells:
                        rows.append(cells)
                tables.append((ws.title, rows))
        finally:
            wb.close()
    elif ext == ".docx":
        import xml.etree.ElementTree as ET
        W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with zipfile.ZipFile(path) as z:
            root = ET.fromstring(z.read("word/document.xml"))
        for i, tbl in enumerate(root.iter(W + "tbl"), 1):
            rows = []
            for tr in tbl.iter(W + "tr"):
                cells = [__import__("textcheck").w_text(tc).strip() for tc in tr.iter(W + "tc")]
                cells = [c for c in cells if c]
                if cells:
                    rows.append(cells)
            tables.append((f"table {i}", rows))
    for label, rows in tables:
        pairs = [r for r in rows if len(r) >= 2]
        good = sum(1 for r in pairs if _is_sf(r[0]) and _initials_fit(r[0], r[1]) >= 0.5)
        is_list = len(pairs) >= 2 and good >= 0.5 * len(pairs)
        text = "\n".join("\t".join(r[:2]) for r in pairs)
        if text:
            out.append((f"{name} ({'acronym table, ' if is_list else ''}{label})", text))
    return out


def acronym_list(path):
    """The acronym list inside one of our Word volumes (2.25.0 QA/QC): {"where": [...], "entries": {term: definition}}.
    Found as a table whose rows mostly read "ACR | definition", or as lines "ACR<tab>definition" (at least 3 in a row,
    or any number under an Acronyms / Abbreviations / Glossary heading). Empty entries when the volume has none."""
    import os, zipfile
    import xml.etree.ElementTree as ET
    import alt_forms as AF
    out = {"where": [], "entries": {}, "alts": {}}
    if os.path.splitext(path)[1].lower() != ".docx":
        return out

    def put(t, d):
        primary, alts = AF.split_term(t)             # "TT&C (TTC)": primary with an alt form (2.26.0)
        out["entries"].setdefault(primary, d)
        if alts:
            out["alts"][primary] = AF.parse_alts(AF.merge_alts(out["alts"].get(primary, []), alts))
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    for i, tbl in enumerate(root.iter(W + "tbl"), 1):
        pairs = []
        for tr in tbl.iter(W + "tr"):
            cells = [__import__("textcheck").w_text(tc).strip() for tc in tr.iter(W + "tc")]
            cells = [c for c in cells if c]
            if len(cells) >= 2:
                pairs.append(cells[:2])
        good = [p for p in pairs if (_is_sf(p[0]) or AF.split_term(p[0])[1])
                and _initials_fit(AF.split_term(p[0])[0], p[1]) >= 0.5]
        if len(pairs) >= 2 and len(good) >= 0.5 * len(pairs):
            out["where"].append(f"table {i}")
            for t, d in pairs:
                if _is_sf(t) or AF.split_term(t)[1]:
                    put(t, d)
    # tab-separated lines outside tables
    in_table = {id(p) for tbl in root.iter(W + "tbl") for p in tbl.iter(W + "p")}
    run, under_head = [], False

    def close():
        nonlocal run
        good = [r for r in run if _initials_fit(r[0], r[1]) >= 0.5]
        if run and (len(run) >= 3 and len(good) >= 0.5 * len(run) or under_head and good):
            out["where"].append("lines under a heading" if under_head else "tab-separated lines")
            for t, d in run:
                put(t, d)
        run = []

    for p in root.iter(W + "p"):
        if id(p) in in_table:
            continue
        text = "".join(("\t" if el.tag == W + "tab" else (el.text or "")) for el in p.iter() if el.tag in (W + "t", W + "tab"))
        m = re.match(r"^\s*(" + SF_RE + r")\t+\s*([A-Za-z][^\t]{2,140}?)\s*$", text)
        if m and _is_sf(m.group(1)):
            run.append((m.group(1), m.group(2)))
            continue
        if text.strip():
            close()
            under_head = bool(_GLOSS_HEAD.match(text)) and len(text.split()) <= 6
    close()
    return out
