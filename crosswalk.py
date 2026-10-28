"""The Section L/M crosswalk (2.61.0): each instruction (Section L) beside the evaluation factor that scores it
(Section M) and the work it describes (Section C, SOW or PWS), with the gaps flagged:
- an instruction no factor evaluates ("not evaluated by any Section M factor found"),
- a factor no instruction asks for ("evaluated, but no instruction asks for it"),
- an instruction tied to no work section (shown, not a gap: many instructions are about format).
Links come from rules only (No AI), each with how it was found, in this order:
1. the person's own links (added or taken out on the crosswalk page; kept with the matrix),
2. cross-references written in the RFP ("Section M.2", "C.2.1", "SOW 3.2", "PWS paragraph 4"),
3. the same subject words in the section titles ("Technical Approach" and "Technical Factor"),
4. a subsection without links of its own takes its parent's ("from L.4").
Nothing leaves the PC.
"""
import re

STOP = {"the", "and", "of", "for", "to", "a", "an", "in", "on", "or", "with", "by", "its", "their", "section", "volume",
        "factor", "subfactor", "approach", "general", "instructions", "instruction", "requirements", "requirement",
        "proposal", "offeror", "offerors", "evaluation", "evaluate", "criteria", "basis", "award", "information",
        "description", "plan", "content", "contents", "format", "submission", "part", "other", "specific", "overview"}
_NUM = re.compile(r"\b(?:(?:Section|Sec\.?|Para(?:graph)?\.?|SOW|PWS|SOO)\s+)?([LMC]\.\d+(?:\.\d+)*|[LMC])\b|"
                  r"\b(?:SOW|PWS|SOO)\s+(?:para(?:graph)?\.?\s*)?(\d+(?:\.\d+)*)\b", re.I)
_WORK_FILE = re.compile(r"\b(?:sow|pws|soo|statement[\s_-]*of[\s_-]*(?:work|objectives)|performance[\s_-]*work)\b", re.I)


def _words(title):
    return {w for w in re.findall(r"[a-z]+", (title or "").lower()) if len(w) > 2 and w not in STOP}


def kind_of(num, file=""):
    """'L', 'M', 'C' (work: Section C, SOW, PWS) or '' for a section number."""
    n = (num or "").strip()
    m = re.match(r"^(?:Section\s+)?([LMC])(?:\.|$)", n, re.I)
    if m:
        return m.group(1).upper()
    if re.match(r"^\d+(\.\d+)*$", n) and _WORK_FILE.search(file or ""):
        return "C"
    return ""


def sections(M):
    """{num: {"num", "title", "file", "kind", "items": [ids]}} from the RFP headings and the sentences' sections."""
    out = {}
    for h in M.headings:
        num = (h.get("num") or "").strip()
        if num and num not in out:
            out[num] = {"num": num, "title": (h.get("title") or "").strip(), "file": h.get("file", ""),
                        "kind": kind_of(num, h.get("file", "")), "items": []}
    for it in M.items:
        num = (it.get("section") or "").strip()
        if not num:
            continue
        s = out.setdefault(num, {"num": num, "title": (it.get("section_title") or "").strip(), "file": it.get("file", ""),
                                 "kind": kind_of(num, it.get("file", "")), "items": []})
        if not s["title"] and it.get("section_title"):
            s["title"] = it["section_title"].strip()
        s["items"].append(it["id"])
    return out


def _sec_key(num):
    parts = re.split(r"[.\s]", num)
    return [(0, int(p)) if p.isdigit() else (1, p) for p in parts if p]


def _parent(num, known):
    p = num
    while "." in p:
        p = p.rsplit(".", 1)[0]
        if p in known:
            return p
    return None


def _refs(text, secs):
    """Section numbers of kind M or C a text points to ("Section M.2", "C.2.1", "SOW 3.2", "Section C")."""
    out = []
    for m in _NUM.finditer(text or ""):
        if m.group(1):
            n = m.group(1).upper() if len(m.group(1)) == 1 else m.group(1)[0].upper() + m.group(1)[1:]
            if n in ("M", "C"):                     # "Section M" / "Section C": the top sections of that kind
                out += [k for k, s in secs.items() if s["kind"] == n and "." not in k[2:] and k.count(".") == 1]
            elif n in secs:
                out.append(n)
            else:                                   # C.2.1 not a heading: its nearest known parent
                p = _parent(n, secs)
                if p:
                    out.append(p)
        elif m.group(2):
            n = m.group(2)
            hit = [k for k, s in secs.items() if s["kind"] == "C" and (k == n or k.endswith("." + n) or k == "C." + n)]
            out += hit
    return [r for r in dict.fromkeys(out) if secs.get(r, {}).get("kind") in ("M", "C") and len(r) > 1]


def build(M):
    """[row] in RFP order: {"num", "title", "m": [(num, title, how)], "c": [(num, title, how)], "gaps": [text]} for each
    Section L section, then a row for each Section M factor no instruction asks for ("kind": "M")."""
    import setupdoc as SD
    secs = sections(M)
    by_id = {it["id"]: it for it in M.items}
    own = SD.setup_of(M).get("crosswalk") or {}
    # a numbered paragraph under a heading with the same title (L.2.1 under L.2 "Volume Organization") is part of it
    fold = {}
    for k in sorted(secs, key=_sec_key):
        p = _parent(k, secs)
        if p and secs[k]["title"].lower() == secs[p]["title"].lower():
            fold[k] = fold.get(p, p)
    for k, p in fold.items():
        secs[p]["items"] += secs[k]["items"]
    whole = {k for k in secs if len(k) == 1}         # "L", "M", "C": a whole RFP section, not one instruction
    Ls = sorted((k for k, s in secs.items() if s["kind"] == "L" and k not in fold and k not in whole), key=_sec_key)
    Ms = [k for k, s in secs.items() if s["kind"] == "M" and k not in fold and k not in whole]
    Cs = [k for k, s in secs.items() if s["kind"] == "C" and k not in fold and k not in whole]
    links = {}
    for k in Ls:
        s = secs[k]
        found = {"M": {}, "C": {}}
        text = " ".join(by_id[i].get("text", "") + " " + by_id[i].get("xrefs", "") for i in s["items"] if i in by_id)
        for r in _refs(text, secs):
            r = fold.get(r, r)
            if r in whole:
                continue
            found[secs[r]["kind"]].setdefault(r, "cross-reference in the RFP")
        mine = _words(s["title"])
        for kind, pool in (("M", Ms), ("C", Cs)):
            if found[kind]:
                continue
            best, score = [], 0.0
            for r in pool:
                w = _words(secs[r]["title"])
                if not mine or not w:
                    continue
                sc = len(mine & w) / min(len(mine), len(w))
                if sc > score:
                    best, score = [r], sc
                elif sc == score and sc > 0:
                    best.append(r)
            if score >= 0.5:
                for r in best[:3]:
                    found[kind].setdefault(r, "the same words in the titles")
        links[k] = found
    for k in Ls:                                    # a subsection without links of its own takes its parent's,
        p = _parent(k, {x: 1 for x in Ls})          # unless its title names a subject of its own
        mine, theirs = _words(secs[k]["title"]), _words(secs[p]["title"]) if p else set()
        if p is None or (mine and not (mine & theirs)):
            continue
        for kind in ("M", "C"):
            if not links[k][kind] and links.get(p, {}).get(kind):
                links[k][kind] = {r: f"from {p}" for r in links[p][kind]}
    for k in Ls:                                    # the person's own links last: they win
        o = own.get(k) or {}
        for r in o.get("drop", []):
            for kind in ("M", "C"):
                links[k][kind].pop(r, None)
        for r in o.get("add", []):
            if r in secs and secs[r]["kind"] in ("M", "C"):
                links[k][secs[r]["kind"]][r] = "added by you"
    rows = []
    used_m = set()
    for k in Ls:
        s = secs[k]
        m = [(r, secs[r]["title"], how) for r, how in sorted(links[k]["M"].items(), key=lambda x: _sec_key(x[0]))]
        c = [(r, secs[r]["title"], how) for r, how in sorted(links[k]["C"].items(), key=lambda x: _sec_key(x[0]))]
        used_m |= {r for r, _t, _h in m}
        gaps = [] if m else ["not evaluated by any Section M factor found"]
        rows.append({"kind": "L", "num": k, "title": s["title"], "m": m, "c": c, "gaps": gaps, "items": len(s["items"])})
    for r in sorted(Ms, key=_sec_key):
        if r in used_m or any(x.startswith(r + ".") for x in used_m) or any(r.startswith(x + ".") for x in used_m):
            continue
        rows.append({"kind": "M", "num": r, "title": secs[r]["title"], "m": [], "c": [],
                     "gaps": ["evaluated, but no instruction asks for it"], "items": len(secs[r]["items"])})
    return rows


def link(M, l_num, other, on=True):
    """The person ties an instruction to a factor or work section (on) or takes a link out (off)."""
    import setupdoc as SD
    cw = SD.setup_of(M).setdefault("crosswalk", {})
    o = cw.setdefault(l_num, {"add": [], "drop": []})
    if on:
        o["drop"] = [x for x in o["drop"] if x != other]
        if other not in o["add"]:
            o["add"].append(other)
    else:
        o["add"] = [x for x in o["add"] if x != other]
        if other not in o["drop"]:
            o["drop"].append(other)
    M.pending = True


def counts(rows):
    return {"instructions": sum(1 for r in rows if r["kind"] == "L"),
            "not_evaluated": sum(1 for r in rows if r["kind"] == "L" and not r["m"]),
            "no_instruction": sum(1 for r in rows if r["kind"] == "M"),
            "no_work": sum(1 for r in rows if r["kind"] == "L" and not r["c"])}


def export(M, path, effort_name=""):
    """The crosswalk as an Excel workbook (one row per instruction, then factors with no instruction)."""
    import os
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    rows = build(M)
    wb = Workbook()
    ws = wb.active
    ws.title = "L-M Crosswalk"
    ws.append([f"{effort_name or M.name}: Section L/M crosswalk"])
    ws["A1"].font = Font(bold=True, size=14, color="1E3A8A")
    ws.append(["Built from the RFP by rules (cross-references, section titles) and your own links. No AI."])
    ws["A2"].font = Font(italic=True, color="475569")
    ws.append([])
    heads = ["Section L", "Instruction", "Evaluated by (Section M)", "How linked", "Work (Section C, SOW, PWS)", "How linked", "Gaps"]
    ws.append(heads)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A8A")
    gap_fill = PatternFill("solid", fgColor="FCEFD9")
    for r in rows:
        ws.append([r["num"] if r["kind"] == "L" else "", r["title"] if r["kind"] == "L" else "",
                   "\n".join(f"{n} {t}".strip() for n, t, _h in r["m"]) or (f"{r['num']} {r['title']}" if r["kind"] == "M" else ""),
                   "\n".join(h for _n, _t, h in r["m"]),
                   "\n".join(f"{n} {t}".strip() for n, t, _h in r["c"]), "\n".join(h for _n, _t, h in r["c"]),
                   "; ".join(r["gaps"])])
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if r["gaps"]:
                cell.fill = gap_fill
    for col, wd in zip("ABCDEFG", (12, 36, 36, 24, 36, 24, 34)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A5"
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    tmp = path + ".tmp"
    wb.save(tmp)
    os.replace(tmp, path)
    return {"path": path, "rows": len(rows)}
