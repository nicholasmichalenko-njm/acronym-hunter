"""Read and assign any document (2.84.0): what a person takes from a library document that does not build the
compliance matrix (an attachment, an ICD, Q&A, a standard, a past proposal), and where it goes.

The rule is the same as for the matrix: select text, then choose where it goes. Each sentence of a library document
can be attached, any number of times, as:
- section:     context for a section of the response structure (target = the section's label);
- requirement: context for a compliance requirement (target = the matrix row's ID);
- note, question, task: made in the effort record (target = the note or task id), quoting the sentence;
- promoted:    made a compliance requirement in the matrix (target = the new row's ID), keeping its source.
An attachment keeps the quoted text, the document, the version and the sentence's place, so it reads the same if the
file moves or is left out of a hand-off, and it stays tied to the version it was read from (a newer version is a
change to review, not a silent move). Topics: any words (the folder tags, the sections and what people type).

The sentences themselves come from the library's text copy of the version (docversions), numbered in reading order:
the ID of a sentence is the document's tag and its number (J1.42), stable because a version's text never changes.
Stored in <effort data>\\context\\attach.json. Rules only (No AI); nothing leaves the PC.
"""
import datetime, getpass, os, re, uuid
import docversions as DV

KINDS = (("section", "Context for section"), ("requirement", "Context for requirement"), ("note", "Note"),
         ("question", "Question for the customer"), ("task", "Task"), ("promoted", "Made a requirement"))
KIND_NAMES = dict(KINDS)
TAG = {"section": "CTX", "requirement": "CTX", "note": "NOTE", "question": "Q", "task": "TASK", "promoted": "REQ"}


def section_context(data_dir):
    """{section label: [(text, document, where)]} for the Response Setup documents, or {} (2.84.0)."""
    try:
        st = Store(data_dir).load()
        if not st.att:
            return {}
        return st.by_section(DV.Library(os.path.join(data_dir, "documents")).load())
    except Exception:
        return {}


def doc_tag(d, lib=None):
    """The short tag a library document's sentence IDs start with: its initial, else letters from its name."""
    if d.get("initial"):
        return d["initial"]
    words = re.findall(r"[A-Za-z0-9]+", d.get("name", "")) or ["D"]
    tag = "".join(w[0] for w in words[:4]).upper()
    if lib is not None:
        taken = {x.get("initial") for x in lib.all if x.get("initial")} | {doc_tag(x) for x in lib.all if x is not d and not x.get("initial")
                                                                              and x["id"] < d["id"]}
        base, n = tag, 2
        while tag in taken:
            tag = f"{base}{n}"
            n += 1
    return tag


def items_of(lib, d, v=None):
    """The sentences of a document's version as rows the page view can box: [{"id", "text", "file", "loc", "n"}]."""
    v = v or lib.latest(d)
    if v is None:
        return []
    tag = doc_tag(d, lib)
    return [{"id": f"{tag}.{n + 1}", "text": s.get("x", ""), "file": v.get("file", ""), "loc": s.get("l", ""), "n": n,
             "section": s.get("s", ""), "section_title": s.get("t", ""), "doc": d["id"], "ver": v["id"]}
            for n, s in enumerate(lib.text(v)) if (s.get("x") or "").strip()]


class Store:
    def __init__(self, data_dir):
        self.path = os.path.join(data_dir, "context", "attach.json")
        self.att = []

    def load(self):
        import workspace as W
        d = W.read_json(self.path, None) or {}
        self.att = [a for a in d.get("att", []) if isinstance(a, dict)]
        return self

    def save(self):
        import workspace as W
        os.makedirs(os.path.dirname(self.path), exist_ok=True)
        W.write_json(self.path, {"v": 1, "att": self.att})

    def attach(self, row, kind, target="", note="", topics=(), who=None, save=True):
        """Attach a sentence (a row from items_of). The same sentence and target twice is kept once. Returns it."""
        if kind not in KIND_NAMES:
            return None
        for a in self.att:
            if a["doc"] == row["doc"] and a["iid"] == row["id"] and a["kind"] == kind and a.get("target", "") == (target or ""):
                return a
        a = {"id": uuid.uuid4().hex[:10], "doc": row["doc"], "ver": row["ver"], "iid": row["id"], "n": row.get("n", 0),
             "text": " ".join((row.get("text") or "").split()), "file": row.get("file", ""), "loc": row.get("loc", ""),
             "kind": kind, "target": target or "", "note": " ".join((note or "").split()), "topics": DV.clean_tags(topics),
             "who": who or getpass.getuser(), "when": datetime.datetime.now().isoformat(timespec="minutes")}
        self.att.append(a)
        if save:
            self.save()
        return a

    def remove(self, att_id=None, doc=None, iid=None, save=True):
        """Take off one attachment, or every attachment of a sentence. Returns how many."""
        keep = [a for a in self.att if not ((att_id and a["id"] == att_id) or (not att_id and a["doc"] == doc and a["iid"] == iid))]
        n = len(self.att) - len(keep)
        self.att = keep
        if n and save:
            self.save()
        return n

    def set_topics(self, att_id, topics):
        a = self.get(att_id)
        if a is None:
            return False
        a["topics"] = DV.clean_tags(topics)
        self.save()
        return True

    def get(self, att_id):
        return next((a for a in self.att if a["id"] == att_id), None)

    def of_row(self, doc, iid):
        return [a for a in self.att if a["doc"] == doc and a["iid"] == iid]

    def of_doc(self, doc, ver=None):
        return [a for a in self.att if a["doc"] == doc and (ver is None or a["ver"] == ver)]

    def for_section(self, label):
        return [a for a in self.att if a["kind"] == "section" and a.get("target") == label]

    def for_requirement(self, iid):
        return [a for a in self.att if a["kind"] == "requirement" and a.get("target") == iid]

    def section_counts(self):
        out = {}
        for a in self.att:
            if a["kind"] == "section":
                out[a["target"]] = out.get(a["target"], 0) + 1
        return out

    def by_section(self, lib=None):
        """{section label: [(quoted text, document name, where)]} for the Response Setup documents."""
        out = {}
        for a in self.att:
            if a["kind"] != "section":
                continue
            d = lib.doc(a["doc"]) if lib is not None else None
            out.setdefault(a["target"], []).append((a["text"], d["name"] if d else a.get("file", ""), a.get("loc", "")))
        return out

    def rename_section(self, old, new):
        n = 0
        for a in self.att:
            if a["kind"] == "section" and a.get("target") == old:
                a["target"] = new
                n += 1
        if n:
            self.save()
        return n

    def old_version(self, lib):
        """Attachments read from a version that is no longer the document's newest (to look at again)."""
        out = []
        for a in self.att:
            d = lib.doc(a["doc"])
            last = lib.latest(d) if d else None
            if last is not None and last["id"] != a["ver"]:
                out.append(a)
        return out

    def rows(self, lib, kind="", doc="", target="", topic="", text=""):
        """Every attachment for the list and the export, filtered: [dict] with the document's name added."""
        out = []
        q = (text or "").lower()
        for a in self.att:
            if (kind and a["kind"] != kind) or (doc and a["doc"] != doc) or (target and a.get("target") != target):
                continue
            if topic and topic.lower() not in [t.lower() for t in a.get("topics", [])]:
                continue
            if q and q not in (a["text"] + " " + a.get("note", "")).lower():
                continue
            d = lib.doc(a["doc"])
            out.append(dict(a, doc_name=d["name"] if d else a.get("file", ""), kind_name=KIND_NAMES[a["kind"]]))
        return out

    def export_xlsx(self, path, lib):
        """Every attachment in one sheet for writers: what, from which document and where, and what it is attached to."""
        import openpyxl
        from openpyxl.styles import Alignment, Font
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "From the library"
        ws.append(["Attached as", "To", "Text", "Document", "Where", "Sentence", "Topics", "Note", "Who", "When"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in sorted(self.rows(lib), key=lambda r: (r["kind"], r.get("target", ""), r["doc_name"], r["n"])):
            ws.append([r["kind_name"], r.get("target", ""), r["text"], r["doc_name"], r.get("loc", ""), r["iid"],
                       ", ".join(r.get("topics", [])), r.get("note", ""), r.get("who", ""), r.get("when", "")])
        for col, wdt in zip("ABCDEFGHIJ", (22, 22, 80, 30, 24, 10, 20, 30, 14, 17)):
            ws.column_dimensions[col].width = wdt
        for row in ws.iter_rows(min_row=2):
            row[2].alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        wb.save(path)
        return path
