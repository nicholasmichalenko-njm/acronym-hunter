"""The formal customer Q&A list (2.98.0): the questions we submit to the customer, each with the customer's exact
wording quoted, where it is (document, section, page, paragraph), our question kept separate from the quote, a
category, a status (Draft, Approved, Submitted, Answered) and, once it comes back, the customer's answer with the
document it came in (usually an amendment).

Kept in the effort record (record.json, "customer_qa"), numbered in the order they will be submitted. Exports:
- Word: a polished document with a cover block (solicitation, offeror, date, point of contact), a numbered table in
  landscape, and the effort's sensitivity marking in every header and footer;
- Excel: the layout most solicitations ask for (No., Document, Section, Page, Paragraph, RFP Text, Question, and an empty
  Government Response column), one row per question, ready to send;
- the customer's own template (Excel or Word): its header row is found by its column names and the questions are
  written under it, leaving the rest of the customer's file as it was.
Rules only (No AI); files are only read and written here; nothing leaves the PC.
"""
import datetime, getpass, os, re, uuid
import zipfile

STATUSES = ("Draft", "Approved", "Submitted", "Answered")
CATEGORIES = ("Technical", "Management", "Contractual", "Pricing", "Format and submission", "Schedule", "Other")
FIELDS = ("question", "quote", "doc", "section", "page", "para", "category", "status", "answer", "answer_doc", "note", "owner")
KEY = "customer_qa"

# the columns of the customer's template, found by their names (lower case, words only)
COLUMN_WORDS = {
    "n": ("no", "number", "item", "item no", "question no", "question number", "q", "q no", "#", "seq", "serial"),
    "doc": ("document", "doc", "rfp document", "reference document", "attachment", "volume", "source", "rfp section document"),
    "section": ("section", "rfp section", "paragraph reference", "reference", "ref", "section paragraph", "rfp reference",
                "section no", "clause", "solicitation section", "paragraph no"),
    "page": ("page", "page no", "page number", "pg", "rfp page"),
    "para": ("paragraph", "para", "para no", "line", "paragraph number"),
    "quote": ("rfp text", "rfp language", "solicitation text", "text", "excerpt", "current language", "rfp verbiage",
              "solicitation language", "statement", "requirement", "original text", "quoted text"),
    "question": ("question", "questions", "comment", "comments", "question comment", "offeror question", "industry question",
                 "inquiry", "query", "vendor question"),
    "category": ("category", "type", "topic", "area", "subject"),
    "answer": ("answer", "response", "government response", "government answer", "agency response"),
}


def _now():
    return datetime.datetime.now().isoformat(timespec="minutes")


def _clean(v, n=4000):
    return " ".join(str(v or "").split())[:n]


def items(rec):
    """The questions, in their order (No. 1 first)."""
    return sorted((q for q in rec.data.get(KEY, []) if isinstance(q, dict)), key=lambda q: q.get("n", 0))


def get(rec, qid):
    return next((q for q in rec.data.get(KEY, []) if q.get("id") == qid), None)


def _renumber(rec):
    for i, q in enumerate(items(rec), start=1):
        q["n"] = i


def add(rec, question, quote="", doc="", section="", page="", para="", category="Technical", owner="", sid="", who=None,
        save=True):
    """A new Draft question at the end of the list. Returns its id ('' without a question or a quote)."""
    question, quote = _clean(question), _clean(quote)
    if not question and not quote:
        return ""
    q = {"id": uuid.uuid4().hex[:10], "n": len(items(rec)) + 1, "question": question, "quote": quote, "doc": _clean(doc, 200),
         "section": _clean(section, 120), "page": _clean(page, 20), "para": _clean(para, 20),
         "category": category if category in CATEGORIES else "Other", "status": "Draft", "answer": "", "answer_doc": "",
         "note": "", "owner": _clean(owner, 80), "sid": sid, "who": who or getpass.getuser(), "at": _now(), "history": []}
    rec.data.setdefault(KEY, []).append(q)
    rec.log("qa", f"Customer question {q['n']} drafted" + (f" ({q['doc']})" if q["doc"] else ""), who=who)
    if save:
        rec.save()
    return q["id"]


def update(rec, qid, who=None, **f):
    """Change a question's fields (FIELDS). A status change is kept in its history with who and when."""
    q = get(rec, qid)
    if q is None:
        return False
    for k in FIELDS:
        if k not in f:
            continue
        v = f[k]
        if k == "status":
            if v not in STATUSES or v == q.get("status"):
                continue
            q.setdefault("history", []).append({"status": v, "was": q.get("status", ""), "who": who or getpass.getuser(), "at": _now()})
            if v == "Submitted" and not q.get("submitted"):
                q["submitted"] = _now()[:10]
            rec.log("qa", f"Customer question {q['n']}: {v.lower()}", who=who)
        elif k == "category":
            v = v if v in CATEGORIES else "Other"
        else:
            v = _clean(v, 200 if k in ("doc", "answer_doc") else 8000)
        q[k] = v
    if f.get("answer") and q.get("status") != "Answered":
        update(rec, qid, who=who, status="Answered")
    rec.save()
    return True


def remove(rec, qid, who=None):
    q = get(rec, qid)
    if q is None:
        return False
    rec.data[KEY] = [x for x in rec.data.get(KEY, []) if x.get("id") != qid]
    _renumber(rec)
    rec.log("qa", f"Customer question {q.get('n')} removed", who=who)
    rec.save()
    return True


def move(rec, qid, delta):
    """Move a question up (-1) or down (+1) the list; the numbers follow."""
    lst = items(rec)
    i = next((k for k, q in enumerate(lst) if q["id"] == qid), None)
    if i is None or not 0 <= i + delta < len(lst):
        return False
    lst[i], lst[i + delta] = lst[i + delta], lst[i]
    for k, q in enumerate(lst, start=1):
        q["n"] = k
    rec.save()
    return True


def counts(rec):
    out = {s: 0 for s in STATUSES}
    for q in items(rec):
        out[q.get("status", "Draft")] = out.get(q.get("status", "Draft"), 0) + 1
    return out


def summary(rec):
    c = counts(rec)
    total = sum(c.values())
    if not total:
        return "No customer questions yet."
    return f"{total} question(s): " + ", ".join(f"{n} {s.lower()}" for s, n in c.items() if n)


def chosen(rec, statuses=None, ids=None):
    """The questions to export: those with the statuses given (all when None), or the ids given."""
    out = items(rec)
    if ids:
        out = [q for q in out if q["id"] in ids]
    elif statuses:
        out = [q for q in out if q.get("status") in statuses]
    return out


def where(q):
    """'Section L.3.1, p.12, para 2' (the parts there are)."""
    return ", ".join(x for x in ((f"Section {q['section']}" if q.get("section") and not str(q["section"]).lower().startswith("section")
                                  else q.get("section", "")), (f"p.{q['page']}" if q.get("page") else ""),
                                 (f"para {q['para']}" if q.get("para") else "")) if x)


def file_name(effort, kind, when=None):
    when = when or datetime.datetime.now()
    return f"{re.sub(r'[^A-Za-z0-9]+', '_', effort).strip('_') or 'Effort'}_Customer_Questions_{when:%Y-%m-%d}.{kind}"


# ------------------------------------------------------------------ Word
def export_docx(path, rows, effort="", solicitation="", offeror="", contact="", marking="", when=None, numbers="list",
                answers=False):
    """The polished Word document. rows: questions (chosen()). numbers: 'list' keeps each question's number in the list;
    'new' numbers the exported ones 1, 2, 3. answers: fill Government Response with the answers we have. Only what the
    customer needs goes out: never the owner, status, category or internal notes. Returns the path."""
    import docxw
    when = when or datetime.date.today()
    d = docxw.Doc()
    d.landscape = True
    if marking:
        d.marking = (marking, "C00000")
    d.props = {"Stationkeeper": "Customer questions", "Effort": effort}
    d.rich([("Questions for the Government", {"bold": True, "size": 36, "color": "1F3864"})])
    cover = [("Solicitation", solicitation or effort), ("Offeror", offeror), ("Date submitted", when.strftime("%B %d, %Y").replace(" 0", " ")),
             ("Point of contact", contact), ("Questions", str(len(rows)))]
    d.table(["", ""], [[[(k, {"bold": True})], v] for k, v in cover if v], [2600, 10360])
    # 3.0.1: one Reference column (the document, then where in it) so the customer's words and the question get the room,
    # and a Government Response column for the customer to answer in (filled only with answers=True)
    head = ["No.", "Reference", "RFP Text", "Question", "Government Response"]
    widths = [620, 2140, 3700, 3700, 2800]
    body = []
    for k, q in enumerate(rows, start=1):
        n = q.get("n", k) if numbers == "list" else k
        ref = [(q.get("doc", ""), {"bold": True})] + ([("\n" + where(q), None)] if where(q) else []) if q.get("doc") else where(q)
        quote = [(f"“{q['quote']}”", {"italic": True})] if q.get("quote") else ""
        body.append([str(n), ref, quote, q.get("question", ""), q.get("answer", "") if answers else ""])
    d.table(head, body, widths)
    d.save(path)
    return path


# ------------------------------------------------------------------ Excel
def export_xlsx(path, rows, effort="", solicitation="", offeror="", marking="", when=None, numbers="list", answers=False):
    """The government layout, one row per question. answers: include the Government Response we have."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    when = when or datetime.date.today()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Questions"
    top = 1
    if marking:
        ws.cell(1, 1, marking).font = Font(bold=True, color="C00000")
        top = 2
    ws.cell(top, 1, f"Questions for the Government: {solicitation or effort}").font = Font(bold=True, size=14, color="1F3864")
    ws.cell(top + 1, 1, " · ".join(x for x in (offeror, when.strftime("%B %d, %Y").replace(" 0", " ")) if x)).font = Font(italic=True, color="595959")
    head = ["No.", "Document", "Section", "Page", "Paragraph", "RFP Text", "Question", "Government Response"]
    hr = top + 3
    thin = Side(style="thin", color="A6A6A6")
    for c, h in enumerate(head, start=1):
        cell = ws.cell(hr, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for k, q in enumerate(rows, start=1):
        n = q.get("n", k) if numbers == "list" else k
        vals = [n, q.get("doc", ""), q.get("section", ""), q.get("page", ""), q.get("para", ""), q.get("quote", ""),
                q.get("question", ""), q.get("answer", "") if answers else ""]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(hr + k, c, v)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for col, wd in zip("ABCDEFGH", (6, 22, 14, 7, 10, 50, 50, 40)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = ws.cell(hr + 1, 1)
    ws.print_title_rows = f"{hr}:{hr}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if marking:
        ws.oddHeader.center.text = marking
        ws.oddFooter.center.text = marking
    wb.save(path)
    return path


# ------------------------------------------------------------------ the customer's template
def _norm(s):
    return " ".join(re.findall(r"[a-z0-9#]+", str(s or "").lower()))


def column_of(header):
    """Which field a template column holds ('' when none)."""
    h = _norm(header)
    if not h:
        return ""
    for field, words in COLUMN_WORDS.items():
        if h in words:
            return field
    for field, words in COLUMN_WORDS.items():          # "Offeror Question / Comment", "RFP Section/Para"
        if any(re.search(r"(^| )" + re.escape(w) + r"( |$)", h) for w in words if len(w) > 2):
            return field
    return ""


def _values(q, k, numbers, full_ref=False, doc_in_section=False):
    """full_ref: the template has no page column, so its section column carries section, page and paragraph.
    doc_in_section (3.0.1): the template has no document column, so the section column names the document first: every
    question always says which customer file it is about."""
    section = where(q) if (full_ref or not q.get("section")) else q.get("section", "")
    if doc_in_section and q.get("doc"):
        section = f"{q['doc']}: {section}" if section else q["doc"]
    return {"n": q.get("n", k) if numbers == "list" else k, "doc": q.get("doc", ""),
            "section": section, "page": q.get("page", ""),
            "para": q.get("para", ""),
            "quote": q.get("quote", ""), "question": q.get("question", ""), "category": q.get("category", ""), "answer": ""}


def template_columns(path):
    """(where, {column: field}) of the customer's template: Excel (sheet name, header row) or Word (table index, header
    row index). None when no header row with a Question column is found."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(path)
        for ws in wb.worksheets:
            for r in range(1, min(ws.max_row, 40) + 1):
                cols = {c: column_of(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 40) + 1)}
                cols = {c: f for c, f in cols.items() if f}
                if "question" in cols.values() and len(cols) >= 2:
                    return (ws.title, r), cols
        return None
    if path.lower().endswith(".docx"):
        tables = _docx_tables(path)
        for ti, rows in enumerate(tables):
            for ri, cells in enumerate(rows[:4]):
                cols = {c: column_of(t) for c, t in enumerate(cells)}
                cols = {c: f for c, f in cols.items() if f}
                if "question" in cols.values() and len(cols) >= 2:
                    return (ti, ri), cols
        return None
    return None


def fill_template(template, out, rows, numbers="list"):
    """Write the questions into a copy of the customer's template (Excel or Word): under its header row, in the columns
    it names. Returns (rows written, [the fields we filled]). Raises ValueError when no Question column is found."""
    found = template_columns(template)
    if not found:
        raise ValueError("No header row with a Question column was found in the template.")
    (where_, hrow), cols = found
    if template.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        from copy import copy
        wb = openpyxl.load_workbook(template, keep_vba=template.lower().endswith(".xlsm"))
        ws = wb[where_]
        r = hrow + 1
        while any(ws.cell(r, c).value not in (None, "") for c in cols) and r < hrow + 2000:   # rows already filled
            r += 1
        style_row = hrow + 1
        full = "page" not in cols.values()
        dis = "doc" not in cols.values()
        for k, q in enumerate(rows, start=1):
            vals = _values(q, k, numbers, full, dis)
            for c, field in cols.items():
                cell = ws.cell(r, c, vals.get(field, ""))
                src = ws.cell(style_row, c)
                if src.has_style:
                    cell.font, cell.border, cell.alignment, cell.fill = copy(src.font), copy(src.border), copy(src.alignment), copy(src.fill)
                    cell.number_format = src.number_format
            r += 1
        wb.save(out)
        return len(rows), sorted(set(cols.values()))
    full, dis = "page" not in cols.values(), "doc" not in cols.values()
    _docx_fill(template, out, where_, hrow, cols, [_values(q, k, numbers, full, dis) for k, q in enumerate(rows, start=1)])
    return len(rows), sorted(set(cols.values()))


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_tables(path):
    """[[cell texts of each row]] of each table in a Word file's body."""
    from lxml import etree
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"), etree.XMLParser(resolve_entities=False, no_network=True))
    out = []
    for tbl in root.iter(_W + "tbl"):
        rows = []
        for tr in tbl.findall(_W + "tr"):
            rows.append(["".join(t.text or "" for t in tc.iter(_W + "t")) for tc in tr.findall(_W + "tc")])
        out.append(rows)
    return out


def _docx_fill(template, out, ti, hrow, cols, values):
    """Copy the row under the header (or the header's look) once per question, set each cell's text, keep everything
    else of the customer's file."""
    import copy
    from lxml import etree
    with zipfile.ZipFile(template) as z:
        names = z.namelist()
        parts = {n: z.read(n) for n in names}
    root = etree.fromstring(parts["word/document.xml"], etree.XMLParser(resolve_entities=False, no_network=True))
    tbl = list(root.iter(_W + "tbl"))[ti]
    trs = tbl.findall(_W + "tr")
    proto = trs[hrow + 1] if len(trs) > hrow + 1 else trs[hrow]
    blank = all(not "".join(t.text or "" for t in tc.iter(_W + "t")).strip() for tc in proto.findall(_W + "tc"))
    anchor = proto if blank else trs[-1]
    for vals in values:
        tr = copy.deepcopy(proto)
        for th in tr.iter(_W + "tblHeader"):
            th.getparent().remove(th)
        for c, tc in enumerate(tr.findall(_W + "tc")):
            text = str(vals.get(cols.get(c, ""), "")) if c in cols else ""
            ps = tc.findall(_W + "p")
            for p in ps[1:]:
                tc.remove(p)
            p = ps[0] if ps else etree.SubElement(tc, _W + "p")
            runs = p.findall(_W + "r")
            rpr = None
            for r in runs:
                if rpr is None and r.find(_W + "rPr") is not None:
                    rpr = copy.deepcopy(r.find(_W + "rPr"))
                p.remove(r)
            r = etree.SubElement(p, _W + "r")
            if rpr is not None:
                for b in rpr.findall(_W + "b"):
                    rpr.remove(b)
                r.append(rpr)
            t = etree.SubElement(r, _W + "t")
            t.text = text
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        anchor.addnext(tr)
        anchor = tr
    if blank and proto.getparent() is not None:
        proto.getparent().remove(proto)
    parts["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n in names:
            z.writestr(n, parts[n])
