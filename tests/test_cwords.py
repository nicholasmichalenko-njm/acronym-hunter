"""2.31.0 compliance words: found in context, styled in every output (Excel, coverage report, Word outline, setup)."""
import os, sys, tempfile, zipfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import cwords as W, compliance as C, setupdoc as SD


def run():
    d = tempfile.mkdtemp()
    cfg = W.load(d)
    assert cfg["style"]["bold"] and not cfg["style"]["italic"], cfg["style"]
    hits = lambda t: [x for x, w in W.segments(t, cfg) if w]
    assert hits("The Offeror shall describe and provide the thermal design.") == ["shall", "describe", "provide"]
    assert hits("Describe the approach in the document.") == ["Describe"]          # a verb at the start; "document" is a noun here
    assert hits("The state of the art document will show margins.") == ["will", "show"]
    assert hits("Proposals shall not exceed 25 pages.") == ["shall not", "exceed"]
    assert hits("(a) Provide resumes; state the contract number.") == ["Provide", "state"]
    assert "".join(t for t, _w in W.segments("A shall B.", cfg)) == "A shall B."
    # the user's style is saved and read back; bad colors are dropped
    W.save({"style": {"bold": True, "italic": True, "underline": True, "color": "C00000", "highlight": "zzz"},
            "words": ["shall", "Shall", "  must "]}, d)
    cfg2 = W.load(d)
    assert cfg2["style"] == {"bold": True, "italic": True, "underline": True, "color": "C00000", "highlight": ""}, cfg2
    assert cfg2["words"] == ["shall", "must"], cfg2["words"]
    parts = W.docx_parts("They shall comply.", {"color": "7030A0"}, cfg2)
    assert parts[1] == ("shall", {"color": "C00000", "bold": True, "italic": True, "underline": True}), parts
    assert parts[0][1] == {"color": "7030A0"}
    # outputs: a small matrix with an outline
    M = C.Matrix(os.path.join(d, "M.xlsx"), "Test")
    base = {"lead": "", "file": "f.docx", "page": "", "section": "L.4", "section_title": "Technical", "actor": "Offeror",
            "type": "Instruction", "strength": "Mandatory", "suggest": "Keep", "modal": "shall", "xrefs": "", "compound": False,
            "owner": "", "status": "", "response_page": "", "notes": "", "links": "", "learned": "", "new": "", "context": "", "loc": ""}
    M.items = [dict(base, id="L.4-01", text="The Offeror shall describe the thermal design.", decision="Keep", map="1.1 Thermal",
                    key=C.text_key("a")),
               dict(base, id="L.4-02", text="Resumes do not count.", decision="Context", map="1.1 Thermal", type="Context",
                    key=C.text_key("b"))]
    M.outline = C.parse_outline("Volume I Technical\n1.1 Thermal")
    import types
    sys.modules["app_gui"] = types.SimpleNamespace(CONFIG_DIR=d)                     # outputs read the style from config
    M.save()
    from openpyxl import load_workbook
    wb = load_workbook(M.path, rich_text=True)
    ws = wb["Compliance Matrix"]
    col = C.COLUMNS.index("Requirement") + 1
    v = ws.cell(row=2, column=col).value
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    assert isinstance(v, CellRichText), type(v)
    blk = [b for b in v if isinstance(b, TextBlock)]
    assert blk and blk[0].text == "shall" and blk[0].font.b and blk[0].font.i, blk
    assert ws.cell(row=3, column=col).value == "Resumes do not count."               # nothing to style: plain text
    N = C.Matrix(M.path).load()
    assert N.items[0]["text"] == "The Offeror shall describe the thermal design.", N.items[0]["text"]
    W.save({"style": {"bold": True, "highlight": "FFFF00"}, "words": W.DEFAULT_WORDS}, d)
    try:
        cp = os.path.join(d, "cov.docx"); C.write_coverage_docx(N, cp)
        x = zipfile.ZipFile(cp).read("word/document.xml").decode()
        assert '<w:b/><w:shd w:val="clear" w:color="auto" w:fill="FFFF00"/></w:rPr><w:t xml:space="preserve">shall' in x, x[:400]
        op = os.path.join(d, "out.docx"); C.write_outline_docx(N, op)
        x = zipfile.ZipFile(op).read("word/document.xml").decode()
        assert '<w:b/><w:shd w:val="clear" w:color="auto" w:fill="FFFF00"/></w:rPr><w:t xml:space="preserve">shall' in x
        paths = SD.export(N, d)
        x = zipfile.ZipFile(paths[0]).read("word/document.xml").decode()
        assert '<w:t xml:space="preserve">shall</w:t>' in x and "<w:b/>" in x
        assert "describe" in x
    finally:
        del sys.modules["app_gui"]
    print("compliance words ok")


if __name__ == "__main__":
    run()
