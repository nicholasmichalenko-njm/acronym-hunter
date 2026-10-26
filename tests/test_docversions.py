"""2.37.0 customer document versions and compare: name cleanup and matching, a library that records each version
(fingerprint, who, when, text copy) and survives a reload, moving and renaming versions, compare with paired changes
and word diffs in reading order, submission-requirement flags, and the Word redline and Excel exports."""
import os, sys, tempfile, zipfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import docversions as DV


def amend(src, dst, pairs):
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == "word/document.xml":
                x = data.decode("utf-8")
                for a, b in pairs:
                    assert a in x, a
                    x = x.replace(a, b)
                data = x.encode("utf-8")
            zo.writestr(item, data)
    return dst


def run():
    assert DV.stem("Kestrel_RFP_CLM_Amd2 (1).docx") == "kestrel rfp clm"
    assert DV.stem("RFP Rev B 2026-10-28.pdf") == "rfp"
    assert DV.stem("Attachment_J-2_Pricing_Template.xlsx") == "attachment j 2 pricing template"
    assert DV.name_similarity("Kestrel_RFP_Amd2.docx", "Kestrel_RFP.docx") == 1.0
    assert DV.is_submission("Volume I Technical shall not exceed 25 pages.") and not DV.is_submission("The Offeror shall identify risks.")
    # compare: paired changes, removed and added in reading order, submission flags
    old = [{"s": "L.2", "t": "Volumes", "x": "Volume I Technical shall not exceed 25 pages."},
           {"s": "L.4", "t": "Approach", "x": "The Offeror shall identify the five highest risks."},
           {"s": "L.4", "t": "Approach", "x": "The Offeror shall provide a risk burn-down chart."},
           {"s": "L.4", "t": "Approach", "x": "Same sentence here."}]
    new = [{"s": "L.2", "t": "Volumes", "x": "Volume I Technical shall not exceed 30 pages."},
           {"s": "L.4", "t": "Approach", "x": "The Offeror shall identify the ten highest risks."},
           {"s": "L.4", "t": "Approach", "x": "Same sentence here."},
           {"s": "L.4", "t": "Approach", "x": "The Offeror shall describe handling for each high risk."}]
    rows = DV.compare(old, new)
    assert [r["kind"] for r in rows] == ["changed", "changed", "removed", "same", "added"], [r["kind"] for r in rows]
    assert rows[0]["diff"] == [("=", "Volume I Technical shall not exceed"), ("-", "25"), ("+", "30"), ("=", "pages.")]
    assert rows[0]["submission"] and not rows[1]["submission"] and rows[0]["section"] == "L.2 Volumes"
    c = DV.summary(rows)
    assert DV.summary_text(c) == "2 changed, 1 added, 1 removed; 1 touches submission requirements"
    # the library, with real files
    import app_gui as A
    d = tempfile.mkdtemp()
    src = os.path.join(A.BASE_DIR, "test_samples", "rfp", "Kestrel_RFP_Sections_C_L_M.docx")
    amd = amend(src, os.path.join(d, "Kestrel_RFP_Sections_C_L_M_Amd1.docx"),
                [("exceed 50 pages", "exceed 40 pages"), ("the five highest", "the ten highest"), ("December 15, 2026", "December 22, 2026")])
    other = os.path.join(A.BASE_DIR, "test_samples", "rfp", "Kestrel_RFP_Section_L_Part2.pdf")
    s1 = DV.snapshot(os.path.basename(src), A.typed_chunks(src))
    s2 = DV.snapshot(os.path.basename(amd), A.typed_chunks(amd))
    s3 = DV.snapshot(os.path.basename(other), A.typed_chunks(other))
    assert len(s1) > 10 and all(set(x) == {"s", "t", "x", "l"} for x in s1)
    lib = DV.Library(os.path.join(d, "documents"))
    doc, v1 = lib.add(src, s1, doc_name="RFP Sections C, L, M", who="tester", when="2026-10-02T09:00")
    assert v1["label"] == "Original" and v1["sha256"] == DV.sha256(src) and len(v1["sha256"]) == 64
    sug = lib.suggest(os.path.basename(amd), [x["x"] for x in s2])
    assert sug[0][0] is doc and DV.Library.band(sug[0][1]) == "strong" and "of the text is the same" in sug[0][2], sug[0][1:]
    sug3 = lib.suggest(os.path.basename(other), [x["x"] for x in s3])
    assert DV.Library.band(sug3[0][1]) != "strong", sug3[0][1:]
    _d, v2 = lib.add(amd, s2, doc_id=doc["id"], who="tester", when="2026-10-14T09:00")
    assert v2["label"] == "Amendment 1" and lib.find_sha(DV.sha256(amd))[1] is v2
    d3, v3 = lib.add(other, s3, doc_name="Section L Part 2")
    # survives a reload; the text copy is kept apart from the file
    lib2 = DV.Library(lib.folder).load()
    assert [x["name"] for x in lib2.docs] == ["RFP Sections C, L, M", "Section L Part 2"]
    assert lib2.counts() == {"docs": 2, "versions": 3, "unreviewed": 1}
    rows = DV.compare(lib2.text(lib2.docs[0]["versions"][0]), lib2.text(lib2.docs[0]["versions"][1]))
    ch = [r for r in rows if r["kind"] != "same"]
    assert len(ch) == 3 and all(r["kind"] == "changed" for r in ch), [(r["kind"], r["old"][:40]) for r in ch]
    assert any(("-", "50") in r["diff"] and ("+", "40") in r["diff"] and r["submission"] for r in ch)
    assert any(r["submission"] for r in ch if "December" in r["new"]), "a due-date change is a submission change"
    # review, rename, move
    assert lib2.mark_reviewed(v2["id"], who="J. Smith") and lib2.counts()["unreviewed"] == 0
    assert lib2.rename(vid=v2["id"], name="Amendment 1 (conformed)") and lib2.docs[0]["versions"][1]["label"] == "Amendment 1 (conformed)"
    assert lib2.move_version(v3["id"], lib2.docs[0]["id"], label="Part 2") and len(lib2.docs) == 1 and len(lib2.docs[0]["versions"]) == 3
    # exports
    rp = DV.export_redline(os.path.join(d, "r.docx"), "RFP", lib2.docs[0]["versions"][0], lib2.docs[0]["versions"][1], rows)
    x = zipfile.ZipFile(rp).read("word/document.xml").decode("utf-8")
    assert "<w:strike/>" in x and ">40 <" in x and "SHA-256 " + v1["sha256"] in x and "submission requirement" in x
    from xml.dom import minidom
    minidom.parseString(x.encode("utf-8"))
    xp = DV.export_list(os.path.join(d, "h.xlsx"), lib2, compares=[(lib2.docs[0], lib2.docs[0]["versions"][0], lib2.docs[0]["versions"][1], rows)])
    from openpyxl import load_workbook
    wb = load_workbook(xp)
    hist = list(wb["Version history"].iter_rows(values_only=True))
    assert hist[0][5] == "SHA-256" and hist[1][5] == v1["sha256"] and len(hist) == 4
    ch_sheet = wb.worksheets[1]
    assert ch_sheet.cell(row=2, column=1).value == "Change" and ch_sheet.max_row == 2 + 3
    print("doc versions ok")


if __name__ == "__main__":
    run()
