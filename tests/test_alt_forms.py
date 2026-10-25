"""2.26.0: alternate forms (TT&C written as TTC) and the master database's use history.
Pure parts: parsing and writing "TT&C (TTC)", the effort list's Alt forms column (set, fold, read from Excel), the
master merge with Used in efforts / Times used / Last used, the scanner marking a found alt form, and the volume lists."""
import os, sys, tempfile, zipfile
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.environ["AH_TEST"] = "1"
import alt_forms as AF


def run():
    # ---------------- parsing and writing
    assert AF.split_term("TT&C (TTC)") == ("TT&C", ["TTC"])
    assert AF.split_term("SATCOM (SatCom, SATCOMM)") == ("SATCOM", ["SatCom", "SATCOMM"])
    assert AF.split_term("Statement of Work (SOW)") == ("Statement of Work (SOW)", [])      # a long form, not alt forms
    assert AF.split_term("C(3)") == ("C(3)", []) and AF.split_term("SOW") == ("SOW", [])
    assert AF.join_term("TT&C", ["TTC"]) == "TT&C (TTC)" and AF.join_term("TT&C", "TTC | T&C") == "TT&C (TTC, T&C)"
    assert AF.join_term("SOW", "") == "SOW"
    assert AF.parse_alts("TTC | T&C, TTC; TC") == ["TTC", "T&C", "TC"]
    assert AF.merge_alts("TTC", ["T&C", "TTC"]) == "TTC | T&C"

    import app_gui as A
    import effort as E
    import master_db as M
    import volcheck as V
    from openpyxl import load_workbook
    d = tempfile.mkdtemp()
    A.CONFIG_DIR = os.path.join(d, "config"); os.makedirs(A.CONFIG_DIR)
    A.DB_FILE = os.path.join(d, "acronym_database.csv")
    with open(A.DB_FILE, "w", encoding="utf-8") as f:
        f.write("Term,Definition,Category,Action\n")

    # ---------------- effort list: Alt forms column, folding an alt form's own row, Excel cell "TT&C (TTC)"
    L = E.EffortList(os.path.join(d, "Kestrel_Effort_List.xlsx"), "Kestrel", solicitation="FA8810-27-R-0042")
    L.add_or_update("TT&C", Type="Acronym", Definition="Telemetry, Tracking and Command")
    L.add_or_update("TTC", Type="Acronym", Definition="Telemetry, Tracking and Command", **{"Used in volumes": "Vol II"})
    L.add_or_update("SOW", Type="Acronym", Definition="Statement of Work")
    folded = L.set_alts("TT&C", ["TTC"])
    assert folded == ["TTC"] and "TTC" not in L.rows and L.alts("TT&C") == ["TTC"], (folded, L.rows.keys())
    assert "Vol II" in L.rows["TT&C"]["Used in volumes"], "the folded row's volumes stay on the primary"
    assert L.alt_map() == {"TTC": "TT&C"}
    L.set_status(["TT&C", "SOW"], "Approved", approver="J. Smith")
    L.save()
    wb = load_workbook(L.path)
    ws = wb[E.SHEET]
    head = [c.value for c in ws[1]]
    assert head[:4] == ["Term", "Type", "Definition", "Alt forms"], head
    ws.append(["SATCOM (SatCom)", "Acronym", "Satellite Communications"])   # typed in Excel by an author
    wb.save(L.path)
    L2 = E.EffortList(L.path); L2.load()
    assert "SATCOM" in L2.rows and L2.alts("SATCOM") == ["SatCom"] and "SATCOM (SatCom)" not in L2.rows

    # ---------------- master database: combine an effort, one row per term and meaning, no repeats
    master = [["TT&C", "Acronym", "Telemetry, Tracking, and Command", "Include", {}],
              ["CDR", "Acronym", "Critical Design Review", "Include", {}],
              ["CDR", "Acronym", "Call Detail Record", "Include", {}]]
    rows, rep = M.combine(master, [L.rows[t] for t in L.order], M.effort_label(L), when="2026-10-25")
    by = {(r[0], r[2]): r for r in rows}
    tc = by[("TT&C", "Telemetry, Tracking, and Command")]
    assert tc[4][M.USED_COL] == "Kestrel (FA8810-27-R-0042)" and tc[4][M.TIMES_COL] == "1" and tc[4][AF.ALT_COL] == "TTC", tc
    assert ("SOW", "Statement of Work") in by and rep["added"] == ["SOW"] and rep["updated"] == ["TT&C (TTC)"], rep
    assert len([r for r in rows if r[0] == "CDR"]) == 2, "same term, different meanings: two rows"
    rows2, rep2 = M.combine(rows, [L.rows[t] for t in L.order], M.effort_label(L), when="2026-10-26")
    assert rep2["added"] == [] and rep2["updated"] == [] and rep2["same"] == 2 and len(rows2) == len(rows), rep2
    O = E.EffortList(os.path.join(d, "Osprey.xlsx"), "Osprey")
    O.add_or_update("TTC", Type="Acronym", Definition="Telemetry Tracking and Command")   # written with the alt form
    O.add_or_update("SOW", Type="Acronym", Definition="Statement of Work")
    O.set_status(["TTC", "SOW"], "Approved")
    rows3, rep3 = M.combine(rows2, [O.rows[t] for t in O.order], M.effort_label(O), when="2026-11-02")
    tc3 = [r for r in rows3 if r[0] == "TT&C"][0]
    assert tc3[4][M.USED_COL] == "Kestrel (FA8810-27-R-0042) | Osprey" and tc3[4][M.TIMES_COL] == "2", tc3
    assert tc3[4][M.LAST_COL] == "2026-11-02" and len([r for r in rows3 if r[0] in ("TT&C", "TTC")]) == 1
    sow3 = [r for r in rows3 if r[0] == "SOW"][0]
    assert sow3[4][M.TIMES_COL] == "2" and len(rows3) == len(rows2)

    # write, read back, load into this PC's database: alt forms and history survive
    mp = os.path.join(d, "Master_Acronym_Database.xlsx")
    M.write_master(mp, rows3)
    back = M.read_master_file(mp)
    tcb = [r for r in back if r[0] == "TT&C"][0]
    assert tcb[4][AF.ALT_COL] == "TTC" and tcb[4][M.USED_COL].endswith("Osprey"), tcb
    M.write_master(os.path.join(d, "empty.xlsx"), [])
    assert M.read_master_file(os.path.join(d, "empty.xlsx")) == [], "an empty master is not read as a plain list"
    info = M.load_master(mp)
    assert info["path"] == os.path.abspath(mp) and info["rows"] == len(rows3)
    db = {r[0]: r for r in A.read_db_rows()}
    assert db["TT&C"][4].get(AF.ALT_COL) == "TTC" and db["TT&C"][4].get(M.TIMES_COL) == "2"
    bk = M.backup_master(mp)
    assert os.path.isfile(bk) and os.path.basename(os.path.dirname(bk)) == "Master database backups"

    # ---------------- the scanner looks for alt forms and labels them with the primary
    inc = A.load_database()[0]
    assert inc["TTC"]["alt_of"] == "TT&C" and inc["TTC"]["definitions"] == ["Telemetry, Tracking, and Command"], inc.get("TTC")
    from tests.test_volcheck import make_docx
    v1 = os.path.join(d, "Vol 1 Technical.docx")
    make_docx(v1, ["The Telemetry, Tracking, and Command (TT&C) radios link the bus. The TTC link closes at 2 Mbps.",
                   "The TT&C subsystem reports health.", ("Heading1", "Acronyms"),
                   ("TABLE", [["TT&C (TTC)", "Telemetry, Tracking, and Command"], ["SOW", "Statement of Work"]])])
    res = A.analyze_batch([v1], lambda m: None)
    assert res["TTC"]["alt_of"] == "TT&C" and res["TTC"]["status"].startswith(A.ALT_STATUS), res["TTC"]["status"]
    assert res["TTC"]["selected_definition"] == "Telemetry, Tracking, and Command"
    assert A.status_kind(res["TTC"]["status"]) == "resolved"

    # ---------------- QA/QC and the volume lists: alt forms in parentheses, never a row of their own
    import custdefs
    al = custdefs.acronym_list(v1)
    assert al["entries"].get("TT&C") and al["alts"] == {"TT&C": ["TTC"]}, al
    r = V.build(res, L.rows, ["Vol 1 Technical.docx"], "volume", file_of=A._file_of, vol_lists={"Vol 1 Technical.docx": al})
    lst = {x["key"]: x for x in r["lists"]["Vol 1 Technical.docx"]}
    assert lst["TT&C"]["term"] == "TT&C (TTC)" and "TTC" not in lst, lst
    assert not any(m["term"] == "TTC" for m in r["missing"]), "an alt form is not missing from the effort list"
    assert not any(x["term"] == "TTC" for x in r["first_use"]), "an alt form has no first use of its own"
    li = {x["term"]: x for x in r["acronym_list"]}
    assert "TTC" not in li or li["TTC"]["level"] != V.BAD, li.get("TTC")
    assert li["TT&C"]["level"] == V.OK, li["TT&C"]
    print("alt forms ok")


if __name__ == "__main__":
    run()
