"""2.91.0 the post-submission checklist, offscreen: a new effort's list starts filled in (owners and days after
submission); tick items done (who and when kept), assign an owner, add, reword, remove and reorder items; due dates count
from the due date until Mark submitted, then from the submission; late items show; Save as default fills the next
effort's list; Export writes a formatted Excel workbook in the effort's Exports folder."""
import datetime, os, sys, tempfile
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE)); sys.path.insert(0, HERE)
os.environ["QT_QPA_PLATFORM"] = "offscreen"
os.environ["AH_TEST"] = "1"
os.environ["SK_SANDBOX_TEMP"] = "temp_qt"
import app_gui as A
tmp = tempfile.mkdtemp()
A.CONFIG_DIR = os.path.join(tmp, "appconfig"); os.makedirs(A.CONFIG_DIR)
A.SETTINGS_FILE = os.path.join(tmp, "settings.json"); A.CHOICES_FILE = os.path.join(tmp, "choices.json")
import effortrec as RR, postsub as PS
from PySide6.QtWidgets import QApplication, QMessageBox
QMessageBox.information = staticmethod(lambda *a, **k: None)
from skqt import design as D
from skqt.mainwin import MainWindow

checks = []


def check(name, ok, info=""):
    checks.append(bool(ok)); print("PASS" if ok else "FAIL", name, info)


cfg = os.path.join(tmp, "config"); os.makedirs(cfg)
app = QApplication([]); D.apply(app)
w = MainWindow(app, cfg, A.APP_VERSION); w.resize(1500, 950); w.show()
L = os.path.join(tmp, "Efforts", "Heron", "Heron_Effort_List.xlsx")
w.settings_page.new_effort({"name": "Heron", "path": L})
check("Post-Submission sits in Deliver", "postsub" in w.step_btns and w.step_btns["postsub"].text() == "Post-Submission")
due = (datetime.date.today() - datetime.timedelta(days=10)).isoformat() + "T14:00"
RR.Record(w.ws.root).set_due(due)
w.go("postsub"); app.processEvents()
pp = w.postsub_page
items = PS.data(RR.Record(w.ws.root))["items"]
check("a new effort's list starts filled in, with owners", len(items) == len(PS.BUILT_IN) and pp.table.rowCount() == len(items)
      and items[0]["owner"] == "Proposal Manager")
r = RR.Record(w.ws.root)
check("due dates count from the due date until it is marked submitted", PS.due_of(r, items[2]) == datetime.date.today() - datetime.timedelta(days=9))
check("late items show", PS.progress(r)[2] > 0 and "late" in pp.state.text())
iid = items[0]["id"]
pp.set_status(iid, "Done")
it = PS.data(RR.Record(w.ws.root))["items"][0]
check("tick an item done: who and when are kept", it["status"] == "Done" and it["done_by"] and it["done_at"])
pp.update(iid, owner="Dana Ruiz")
check("assign an owner", PS.data(RR.Record(w.ws.root))["items"][0]["owner"] == "Dana Ruiz")
new = pp.add_item("Return the loaner laptops", "A. Lee", 2)
items = PS.data(RR.Record(w.ws.root))["items"]
check("add an item", items[-1]["id"] == new and items[-1]["owner"] == "A. Lee" and pp.table.rowCount() == len(PS.BUILT_IN) + 1)
pp.table.item(len(items) - 1, 1).setText("Return the loaner laptops and badges")
check("reword an item in place", PS.data(RR.Record(w.ws.root))["items"][-1]["text"] == "Return the loaner laptops and badges")
pp.move(-1, new)
check("reorder", PS.data(RR.Record(w.ws.root))["items"][-2]["id"] == new)
pp.remove_item(items[1]["id"])
check("remove", len(PS.data(RR.Record(w.ws.root))["items"]) == len(PS.BUILT_IN))
sub = (datetime.datetime.now()).replace(second=0, microsecond=0).isoformat(timespec="minutes")
pp.mark_submitted(sub)
r = RR.Record(w.ws.root)
x = next(i for i in PS.data(r)["items"] if i["days"] == 7)
check("Mark submitted: due dates count from the submission", PS.due_of(r, x) == datetime.date.today() + datetime.timedelta(days=7)
      and "Submitted" in pp.state.text() and any("submitted" in e["text"] for e in r.events(limit=5)))
out = pp.export(os.path.join(os.path.dirname(L), "Exports", "Heron_Post-Submission_Checklist.xlsx"))
import openpyxl
wb = openpyxl.load_workbook(out); ws = wb.active
heads = [c.value for c in ws[7]]
check("Export: a formatted workbook for the effort's folder", heads[:5] == ["#", "Task", "Owner", "Due", "Status"]
      and ws["A1"].value.startswith("Heron") and ws.freeze_panes == "A8" and ws.data_validations.dataValidation
      and len(ws.conditional_formatting) >= 1 and ws.page_setup.orientation == "landscape", heads)
check("one row per item, done ones marked done", ws.max_row == 7 + len(PS.BUILT_IN) and any(ws.cell(r, 5).value == "Done" for r in range(8, ws.max_row + 1)))
pp.save_default()
w.settings_page.new_effort({"name": "Osprey", "path": os.path.join(tmp, "Efforts", "Osprey", "Osprey_Effort_List.xlsx")})
w.go("postsub"); app.processEvents()
texts = [i["text"] for i in PS.data(RR.Record(w.ws.root))["items"]]
check("Save as default: the next effort starts from this list", "Return the loaner laptops and badges" in texts and len(texts) == len(PS.BUILT_IN))
w.close()
n = sum(checks)
print(f"QT POSTSUB PASS ({n}/{len(checks)})" if n == len(checks) else f"QT POSTSUB FAIL ({n}/{len(checks)})")
sys.exit(0 if n == len(checks) else 1)
