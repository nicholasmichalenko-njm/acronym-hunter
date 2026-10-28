"""2.48.0 the compliance matrix in Excel, designed in the app: a design keeps only known columns (renamed, in order) and
roles; rows follow the response structure; the workbook has the title, the headings as designed, one row per compliance
requirement (and context when chosen), a Customer's sheet, About; the team default saves and loads; the matrix and the
Response Setup export together into one dated folder."""
import os, sys, tempfile
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE)); sys.path.insert(0, HERE)
import compliance as C, setupdoc as SD, matrixdesign as MD
from test_assign import row


def build():
    a = row(1, "L.5", "The Offeror shall provide the schedule.")
    b = row(2, "L.4", "The Offeror shall describe the thermal design.")
    c = row(3, "L.4", "Thermal margins matter because the orbit has long eclipses.")
    e = row(4, "L.4", "The Government will evaluate the thermal design.")
    items = [a, b, c, e]
    M = C.Matrix(os.path.join(tempfile.mkdtemp(), "Kestrel_Matrix.xlsx"), "Kestrel")
    M.items = items
    M.outline = C.parse_outline("Volume I\n1.1 Thermal\n1.2 Schedule")
    C.assign_requirement(items, [a], "1.2 Schedule")
    C.assign_requirement(items, [b], "1.1 Thermal")
    C.assign_context(items, [c], ["L.4-02"])
    C.assign_customer(items, [e])
    SD.set_people(M, "1.1 Thermal", lead="J. Smith")
    a["response_page"] = "12"; a["status"] = "Compliant"
    return M


def run():
    M = build()
    d = MD.clean({"columns": [{"key": "text", "heading": "  Requirement   text "}, {"key": "bogus"}, {"key": "ref"},
                              {"key": "text"}, {"key": "owner", "heading": "Author"}], "roles": ["Keep", "Nope"],
                  "order": "structure"})
    assert [c["key"] for c in d["columns"]] == ["text", "ref", "owner"] and d["columns"][0]["heading"] == "Requirement text"
    assert d["roles"] == ["Keep"] and d["customer_sheet"] is True
    old_cols = [{"key": k, "heading": MD.COLUMNS[k][0]} for k in ("ref", "text", "section", "owner", "page", "status", "notes")]
    keep_only = dict(MD.default(), columns=old_cols, roles=["Keep"], order="structure", context="rows")
    items, data = MD.rows(M, keep_only)
    assert [it["id"] for it in items] == ["L.4-02", "L.5-01"], "in the structure's order: 1.1 Thermal before 1.2 Schedule"
    assert data[0][3] == "J. Smith" and data[1][4] == "12" and data[1][5] == "Compliant", data
    both = dict(keep_only, roles=["Keep", "Context"])
    assert [it["id"] for it in MD.rows(M, both)[0]] == ["L.4-02", "L.4-03", "L.5-01"]
    # 2.58.0: support in its requirement's cell (the default), with where it is in the RFP
    cell_items, cell_data = MD.rows(M, dict(both, context="cell"))
    assert [it["id"] for it in cell_items] == ["L.4-02", "L.5-01"] and [c["id"] for c in cell_items[0]["_ctx"]] == ["L.4-03"]
    tcol = [c["key"] for c in both["columns"]].index("text")
    assert "\n\nSupporting context:\n(" in cell_data[0][tcol] and "Thermal margins" in cell_data[0][tcol], cell_data[0][tcol]
    assert MD.default()["context"] == "cell" and MD.clean({"context": "rows"})["context"] == "rows" and MD.clean({})["context"] == "cell"
    # 2.51.0: Order, Context under its requirement, Volume, deliverable facts, empty sections, own columns
    olddef = [{"key": k} for k in MD._OLD_DEFAULT_KEYS]
    full = MD.clean(dict(MD.default(), v=2, columns=olddef + [{"key": "deliverable"}, {"key": "format"}, {"key": "far"},
                                                                       {"key": "custom:Reviewer"}], order="structure", context="rows"))
    ctx = {"deliverables": [{"name": "Volume I", "limit": "25 pages", "format": "Font: 12 point Times"}]}
    M.outline = C.parse_outline("Volume I\n1.1 Thermal\n1.2 Schedule\n1.3 Risk")
    its, dat = MD.rows(M, full, ctx)
    col = {c["key"]: k for k, c in enumerate(full["columns"])}
    assert [it.get("_kind") for it in its] == ["req", "ctx", "req", "empty"], [it.get("_kind") for it in its]
    assert [r[col["order"]] for r in dat] == [1, 2, 3, 4]
    assert dat[1][col["text"]].startswith("Thermal margins") and its[1]["_host"] == "L.4-02", "context under its requirement"
    assert dat[0][col["volume"]] == "Volume I" and dat[0][col["page_limit"]] == "25 pages" and dat[0][col["deliverable"]] == "Volume I"
    assert dat[0][col["format"]] == "Font: 12 point Times" and dat[0][col["custom:Reviewer"]] == ""
    assert dat[3][col["section"]] == "1.3 Risk" and dat[3][col["text"]] == "(nothing assigned yet)", "a section with nothing assigned"
    assert full["columns"][-1] == {"key": "custom:Reviewer", "heading": "Reviewer"}
    assert len(MD.rows(M, dict(full, empty_sections=False))[0]) == 3
    assert "id" not in [k for k in its[0] if k.startswith("_")] and "_kind" not in M.items[0], "the matrix rows are not touched"
    rfp = MD.rows(M, dict(full, order="rfp"), ctx)[0]                 # 2.54.0: the RFP's order, support under its requirement
    assert [(x.get("id"), x["_kind"]) for x in rfp] == [("L.5-01", "req"), ("L.4-02", "req"), ("L.4-03", "ctx"), ("", "empty")], \
        [(x.get("id"), x["_kind"]) for x in rfp]
    assert MD.clean({})["order"] == "rfp" and MD.clean({"order": "nonsense"})["order"] == "rfp"
    # 2.56.0: the default columns; an old default never changed takes the new one; a changed one is kept
    assert [c["heading"] for c in MD.default()["columns"]] == ["Volume", "Response section no.", "Response section name", "Reference",
                                                               "Doc", "L/M", "Section", "RFP content", "Author(s)", "Notes"]
    assert [c["key"] for c in MD.clean({"columns": olddef})["columns"]] == [c["key"] for c in MD.default()["columns"]]
    assert [c["key"] for c in MD.clean({"columns": olddef[:3]})["columns"]] == MD._OLD_DEFAULT_KEYS[:3]
    nd = MD.default()
    ncol = {c["key"]: k for k, c in enumerate(nd["columns"])}
    its2, dat2 = MD.rows(M, dict(nd, order="structure"), ctx)
    r0 = dat2[0]
    assert r0[ncol["section_no"]] == "1.1" and r0[ncol["section_name"]] == "Thermal" and r0[ncol["volume"]] == "Volume I", r0
    assert r0[ncol["rfp_ref"]].startswith("L.4") and r0[ncol["authors"]] == "J. Smith" and r0[ncol["text"]].startswith("The Offeror"), r0
    assert MD._value(M, its[1], "role", {}, {}) == "Support for L.4-02" and MD._value(M, its[0], "role", {}, {}) == "Compliance requirement"
    xo = os.path.join(tempfile.mkdtemp(), "styled.xlsx")
    MD.export(M, xo, full, "Kestrel", ctx=ctx)
    from openpyxl import load_workbook as _lw
    cell = _lw(xo, rich_text=True)["Compliance Matrix"].cell(row=5, column=col["text"] + 1).value
    assert type(cell).__name__ == "CellRichText" and any(getattr(b, "text", "") == "shall" for b in cell), "compliance words styled in the export"
    M.outline = C.parse_outline("Volume I\n1.1 Thermal\n1.2 Schedule")
    out = os.path.join(tempfile.mkdtemp(), "x", "Kestrel_Compliance_Matrix.xlsx")
    res = MD.export(M, out, d, "Kestrel")
    assert res == {"path": out, "rows": 2, "customer": 1}, res
    from openpyxl import load_workbook
    wb = load_workbook(out)
    ws = wb["Compliance Matrix"]
    assert ws["A1"].value == "Kestrel: Compliance Matrix" and [c.value for c in ws[4]] == ["Requirement text", "Section L No.", "Author"]
    assert ws.cell(row=5, column=1).value == "The Offeror shall describe the thermal design." and ws.cell(row=5, column=3).value == "J. Smith"
    assert wb.sheetnames == ["Compliance Matrix", "Customer's", "About"]
    assert "Government will evaluate" in wb["Customer's"].cell(row=3, column=2).value
    cfg = tempfile.mkdtemp()
    assert MD.load_team(cfg) is None
    MD.save_team(cfg, d)
    assert MD.load_team(cfg) == MD.clean(d) and MD.design_for(M, cfg) == MD.clean(d)
    MD.set_matrix(M, both)
    assert MD.design_for(M, cfg)["roles"] == ["Keep", "Context"], "the matrix's own design wins"
    M.save()
    assert MD.of_matrix(C.Matrix(M.path).load())["roles"] == ["Keep", "Context"], "kept with the matrix"
    t = MD.export_together(M, tempfile.mkdtemp(), MD.design_for(M, cfg), "Kestrel")
    names = sorted(os.path.basename(p) for p in t["files"])
    assert os.path.basename(t["folder"]).startswith("Kestrel compliance export ") and any(n.endswith(".xlsx") for n in names) \
        and any(n.endswith(".docx") for n in names), names
    print("matrix design ok")


if __name__ == "__main__":
    run()
