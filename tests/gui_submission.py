"""2.37.0 Submission, phase 1 (and 2.39.0 phase 2, requirements): the customer documents folder shows files not in the history yet; Add reads them,
suggests matches and records versions; a new amendment is matched to its document (strong match, preselected);
the same file twice is skipped; the history shows each version with its fingerprint; Compare lines up the two
versions with removed and added words marked; filters, jumping between changes, exports; Mark reviewed; the rail
status; Export all includes the version history."""
import os, sys, tempfile, shutil, zipfile
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE)); sys.path.insert(0, HERE)
os.environ["AH_OCR_ENGINE"] = "tesseract"
os.environ["AH_TEST"] = "1"
import app_gui as A, effort as E, docversions as DV
from tkinter import messagebox
from test_docversions import amend
tmp = tempfile.mkdtemp()
A.CONFIG_DIR = os.path.join(tmp, "config"); A.CHOICES_FILE = os.path.join(tmp, "config", "choices.json")
A.SETTINGS_FILE = os.path.join(tmp, "settings.json")
messagebox.showinfo = lambda *a, **k: None
messagebox.askyesno = lambda *a, **k: True
LIST = os.path.join(tmp, "Kestrel", "Kestrel_Effort_List.xlsx")
CDIR = os.path.join(tmp, "Kestrel", "Customer documents"); os.makedirs(CDIR)
RFP = os.path.join(A.BASE_DIR, "test_samples", "rfp")
for f in os.listdir(RFP):
    shutil.copy2(os.path.join(RFP, f), CDIR)
INBOX = os.path.join(tmp, "inbox"); os.makedirs(INBOX)
AMD = amend(os.path.join(RFP, "Kestrel_RFP_Sections_C_L_M.docx"), os.path.join(INBOX, "Kestrel_RFP_Sections_C_L_M_Amd1.docx"),
            [("exceed 50 pages", "exceed 40 pages"), ("the five highest", "the ten highest")])
checks = []


def check(name, ok, info=""):
    checks.append(bool(ok)); print("PASS" if ok else "FAIL", name, info)


app = A.AcronymAppGUI()
app.geometry("1500x940+0+0")


def go():
    L = E.EffortList(LIST, "Kestrel"); L.pending = True; L.save()
    app.efforts.put("Kestrel", path=LIST, customer_dir=CDIR)
    app.open_effort("Kestrel"); app.update()
    rail = [app._rail[k][0].cget("text") for k in app._rail]
    check("Customer Documents sits second in the feature list", rail[1].endswith("Customer Documents"), rail[:3])
    app.shell_go("submission"); app.update()
    pills = [w.cget("text").split("  ")[0] for w in app.strip_pills.winfo_children()]
    check("its steps: Requirements, Deliverables, Documents, Impact, Compare, Export; with nothing read yet it opens on Documents",
          pills == ["1 Requirements", "2 Deliverables", "3 Documents", "4 Impact", "5 Compare", "6 Export"] and app.shell_step == "documents", (pills, app.shell_step))
    check("no bar until requirements are found", not app.subbar.winfo_manager())
    check("files in the customer documents folder not in the history yet show in a banner",
          app.sub_banner.winfo_manager() and "2 customer document" in app.sub_banner_lbl.cget("text"), app.sub_banner_lbl.cget("text"))
    w = app.sub_add_dialog(app.sub_unadded()); w.update()
    plan = w._api["plan"]
    check("each file is read and offered as a new document", len(plan) == 2 and all(e["choice"] == "new" and e["sents"] for e in plan),
          [(os.path.basename(e["path"]), e["choice"]) for e in plan])
    made = w._api["add"](); app.update()
    lib = app.doclib()
    check("added: two documents, each with its Original version", len(made) == 2 and [len(d["versions"]) for d in lib.docs] == [1, 1]
          and all(v["label"] == "Original" for _d, v in made))
    check("the banner goes away", not app.sub_banner.winfo_manager())
    # an amendment arrives
    w = app.sub_add_dialog([AMD]); w.update()
    e = w._api["plan"][0]
    rfp = next(d for d in lib.docs if d["name"] == "Kestrel RFP Sections C L M")
    check("the amendment is matched to its document and preselected", e["choice"] == "version" and e["doc"] == rfp["id"]
          and DV.Library.band(e["best"][1]) == "strong", (e["choice"], e["best"] and e["best"][1:]))
    var, om, ids, names, lab, nm = w._api["rows"][0]
    check("named as the next version", lab.get() == "Amendment 1", lab.get())
    w._api["add"](); app.update()
    check("the amendment is a new version, copied into the customer documents folder",
          len(rfp["versions"]) == 2 and os.path.isfile(os.path.join(CDIR, os.path.basename(AMD))))
    check("a new version opens its impact review (2.42.0)", app.shell_step == "impact" and app.imp_r is not None
          and app.imp_r["version"] == "Amendment 1", app.shell_step)
    w = app.sub_add_dialog([AMD]); w.update()
    check("the same file twice is skipped (same fingerprint)", w._api["plan"][0].get("dup") is not None)
    w.destroy()
    # the history
    app.sub_show(rfp["id"]); app.update()
    texts = []
    def walk(wd):
        for c in wd.winfo_children():
            try:
                texts.append(c.cget("text"))
            except Exception:
                pass
            walk(c)
    walk(app.sub_detail)
    joined = " | ".join(t for t in texts if isinstance(t, str))
    check("the history shows each version newest first, with its fingerprint and what changed",
          joined.index("Amendment 1") < joined.index("Original") and "Fingerprint (SHA-256)" in joined and "2 changed" in joined
          and "Not reviewed yet" in joined, joined[:300])
    check("the rail says a change waits for review", app.shell_status("submission") == "1 change to review", app.shell_status("submission"))
    # compare
    rows = app.sub_compare(rfp["id"]); app.update()
    check("Compare opens on the last two versions", app.shell_step == "compare" and app.cmp_from.get().startswith("Original")
          and app.cmp_to.get().startswith("Amendment 1"), (app.cmp_from.get(), app.cmp_to.get()))
    check("only the changes are listed, lined up", len(app.cmp_rows) == 2 and all(r["kind"] == "changed" for r in app.cmp_rows), len(app.cmp_rows))
    L_, R_ = app.cmp_panes
    dels = [L_.get(a, b) for a, b in zip(L_.tag_ranges("del")[0::2], L_.tag_ranges("del")[1::2])]
    ins = [R_.get(a, b) for a, b in zip(R_.tag_ranges("ins")[0::2], R_.tag_ranges("ins")[1::2])]
    check("removed words struck through on the left, added words on the right", "50 " in dels and "40 " in ins and "ten " in ins, (dels, ins))
    a0, b0 = L_.tag_ranges("row:1"), R_.tag_ranges("row:1")
    check("both sides line up row by row", L_.index(a0[0]).split(".")[0] == R_.index(b0[0]).split(".")[0])
    check("the page-limit change is marked as a submission requirement", any(r["submission"] for r in app.cmp_rows))
    app.cmp_subonly.set(True); app.cmp_refresh(); app.update()
    check("Only submission requirements narrows the list", len(app.cmp_rows) == 1 and "pages" in app.cmp_rows[0]["new"], len(app.cmp_rows))
    app.cmp_subonly.set(False); app.cmp_only.set(False); app.cmp_refresh(); app.update()
    check("all sentences when Only changes is off", len(app.cmp_rows) > 10)
    i1 = app.cmp_jump(1); i2 = app.cmp_jump(1)
    check("Down jumps from change to change", app.cmp_rows[i1]["kind"] == "changed" and app.cmp_rows[i2]["kind"] == "changed" and i2 > i1, (i1, i2))
    app.cmp_only.set(True); app.cmp_refresh()
    rp = app.cmp_export_redline(os.path.join(tmp, "redline.docx"))
    xp = app.cmp_export_list(os.path.join(tmp, "changes.xlsx"))
    x = zipfile.ZipFile(rp).read("word/document.xml").decode("utf-8")
    check("redline and change list written", "<w:strike/>" in x and os.path.isfile(xp))
    # review and export
    v2 = rfp["versions"][1]
    app.shell_go("submission", "documents"); app.update()
    check("Mark reviewed records who and when; the rail then says a requirement changed", app.sub_mark_reviewed(v2["id"]) and v2["reviewed"]
          and app.shell_status("submission") == "1 requirement changed", app.shell_status("submission"))
    # 2.39.0: submission requirements, followed across versions
    app.shell_go("submission", "requirements"); app.update()
    check("the Requirements step opens", app.shell_step == "requirements", app.shell_step)
    rows = {app.req_rows[i]["key"]: (i, app.req_rows[i]) for i in app.req_items()}
    groups = [app.req_tree.item(g, "text") for g in app.req_tree.get_children()]
    check("2.49.0: requirements are grouped: submittal items, page limits, format, files, how and to whom, dates",
          [g.split(" (")[0] for g in groups] == ["Submittal items", "Page limits", "Format", "Files", "How, where and to whom", "Dates"], groups)
    check("Requirements lists what the documents set", {"Due | Proposal", "Page limit | Proposal", "Format | Font", "File | File format",
                                                         "File | File naming", "Deliverable | Attachment J-2"} <= set(rows), sorted(rows))
    key = "Page limit | Proposal"
    i, r = rows[key]
    check("the changed page limit is marked to review, with its new value", r["to_review"] and r["value"] == "40 pages"
          and app.req_tree.set(i, "state") == "Changed: review" and "review" in app.req_tree.item(i, "tags"), app.req_tree.item(i))
    check("the card shows the due date and changes to review", app.req_cells["due"][0].cget("text").startswith("December 15, 2026")
          and app.req_cells["review"][0].cget("text") == "1", [v.cget("text") for v, _s in app.req_cells.values()])
    app.req_fact_edit("method", "DoD SAFE"); app.req_fact_edit("contact", "Jane Doe, Contracting Officer, jane.doe@us.af.mil"); app.update()
    check("2.49.0: how it is submitted and the submission contact are on the screen and editable",
          app.req_cells["method"][0].cget("text") == "DoD SAFE" and "jane.doe@us.af.mil" in app.req_cells["contact"][0].cget("text"))
    bar = app.sub_bar_refresh(); app.update()
    check("the bar shows due date, page limits, format and changes, across every feature", app.subbar.winfo_manager()
          and "Due December 15, 2026" in bar and "Pages: Proposal 40" in bar and "12-point Times New Roman" in bar and "1 change to review" in bar, bar)
    app.shell_go("compliance"); app.update()
    check("still there in another feature", app.subbar.winfo_manager())
    app.subbar_lbl._label.event_generate("<Button-1>"); app.update()
    check("clicking the bar opens Requirements", (app.shell_feature, app.shell_step) == ("submission", "requirements"))
    i = next(x for x in app.req_items() if app.req_rows[x]["key"] == key)
    app.req_tree.selection_set(i); app.update()
    texts.clear(); walk(app.req_detail)
    joined = " | ".join(t for t in texts if isinstance(t, str))
    check("the history opens on the latest version: was, now, the sentence, the file and its fingerprint",
          "Version 2 of 2" in joined and "Was: 50 pages" in joined and "Now: 40 pages" in joined and "exceed 40 pages" in joined
          and "Fingerprint (SHA-256)" in joined and "Amendment 1" in joined, joined[:400])
    app.req_history(key, 0); app.update()
    texts.clear(); walk(app.req_detail)
    joined = " | ".join(t for t in texts if isinstance(t, str))
    check("Previous goes back to the Original version", "Version 1 of 2" in joined and "First found" in joined and "exceed 50 pages" in joined)
    hit = app.req_show_change(key); app.update()
    check("Show the change opens Compare on the changed sentence", app.shell_step == "compare" and hit is not None
          and "40 pages" in app.cmp_rows[hit]["new"] and app.cmp_panes[1].tag_ranges("cur"), hit)
    app.shell_go("submission", "requirements"); app.update()
    check("Mark reviewed clears it; the rail and the bar follow", app.req_review(key) and not app.req_selected()["to_review"]
          and app.shell_status("submission") == "2 documents, 3 versions" and "to review" not in app.sub_bar_refresh(), app.shell_status("submission"))
    app.req_edit("Format | Margins", "1-inch margins on all four sides")
    check("Edit keeps what the documents say beside it", app.req_selected()["value"] == "1-inch margins on all four sides"
          and app.req_selected()["found"] == "1-inch margins")
    k = app.req_add_dialog("Page limit", "Volume I", "20 pages")
    check("a typed-in requirement", k and app.req_selected()["manual"] and app.req_tree.set(app.req_tree.selection()[0], "state") == "Typed in")
    app.req_hide("Format | Page size")
    check("Not a requirement hides it", "Format | Page size" not in {r["key"] for r in app.req_rows.values()})
    app.req_hidden_var.set(True); app.req_refresh()
    check("Show items marked Not a requirement lists it again", "Format | Page size" in {r["key"] for r in app.req_rows.values()})
    app.req_hidden_var.set(False); app.req_refresh()
    app.shell_go("home"); app.update()
    card = app.ov_rows["submission"].cget("text")
    check("the Submission card on Effort home shows the due date and page limits", "Due December 15, 2026" in card and "Pages: Proposal 40, Volume I 20" in card, card)
    xp = app.req_export(os.path.join(tmp, "reqs.xlsx"))
    from openpyxl import load_workbook as _lw
    wbr = _lw(xp)
    check("the requirements export has the list and the history", wbr.sheetnames == ["Submission requirements", "History"]
          and wbr["History"].max_row > wbr["Submission requirements"].max_row)
    hp = app.sub_export_history(os.path.join(tmp, "history.xlsx"))
    from openpyxl import load_workbook
    wb = load_workbook(hp)
    check("the version history export has the trail and a sheet of changes", wb["Version history"].max_row == 4 and len(wb.worksheets) == 2)
    sub, files, probs = app.export_all(os.path.join(tmp, "out"), ["versions", "requirements"])
    check("Export all includes the version history and the requirements", any(f.endswith("_Customer_Document_Versions.xlsx") for f in files)
          and any(f.endswith("_Submission_Requirements.xlsx") for f in files) and not probs, files)
    check("saved in the effort folder as it happens", os.path.isfile(os.path.join(app.ws.root, "documents", "index.json"))
          and DV.Library(os.path.join(app.ws.root, "documents")).load().counts()["versions"] == 3)
    # 2.38.0: Acronyms with an effort open: straight to Review, one button scans the effort's documents, Found in column
    VOL = os.path.join(A.BASE_DIR, "test_samples", "Kestrel_Technical_Volume.docx")
    app.add_volume("Vol I Technical", [VOL])
    import subreq as SR
    pg = SR.page_check({os.path.basename(VOL): VOL}, app.subreg().rows(), app.volumes())
    check("QA/QC page limits: our volume against its limit and all volumes against the proposal limit",
          [(p["volume"], p["level"], p["pages"]) for p in pg] == [("Kestrel_Technical_Volume.docx", "ok", 1), ("All volumes", "ok", 1)], pg)
    app.results = {}
    app.shell_go("acronyms"); app.update()
    check("Acronyms opens on Review, offering to scan the effort's documents", app.shell_step == "review"
          and app.rv_scanbar.winfo_manager() and "customer" in app.rv_scanlbl.cget("text"), app.rv_scanlbl.cget("text"))
    n = app.acronyms_effort_scan(force=True)
    wait(lambda: not app.busy() and app.results, lambda: found(n))


def found(n):
    app.update()
    vals = {app.tree.set(i, "foundin") for i in app.tree.get_children()}
    check("the scan covers customer documents and ours", n >= 3, n)
    check("Found in says where each term was found", {"Customer", "Ours"} & vals and vals <= {"Customer", "Ours", "Both", ""}, vals)
    app.found_var.set("Customer documents"); app.populate_table(); app.update()
    only = {app.tree.set(i, "foundin") for i in app.tree.get_children()}
    check("filter by where it was found", only == {"Customer"}, only)
    app.found_var.set("All documents"); app.populate_table()
    check("the scan offer goes away once something is found", not app.rv_scanbar.winfo_manager())
    print(f"GUI SUBMISSION {'PASS' if all(checks) else 'FAIL'} ({sum(checks)}/{len(checks)})")
    app.after(200, app.on_close)


def wait(cond, then, n=0):
    if not cond() and n < 240:
        app.after(500, lambda: wait(cond, then, n + 1)); return
    then()


A.set_ai_ml(True)
app.after(900, go)
app.mainloop()
