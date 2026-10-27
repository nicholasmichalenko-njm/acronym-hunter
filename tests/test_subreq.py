"""2.39.0 submission requirements: the rules find due dates, page limits per volume, format, files and deliverables;
each requirement is followed across versions (first, same, changed, removed), a change shows until marked reviewed;
confirm, edit, hide and typed-in requirements are saved in register.json; the summary, days until due, and matching
a page limit to our volume files."""
import os, sys, tempfile, datetime
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE)); sys.path.insert(0, HERE)
os.environ.setdefault("AH_OCR_ENGINE", "tesseract")
import subreq as S, docversions as DV
from test_docversions import amend


def run():
    f = S.find
    assert f("Proposals are due no later than 2:00 p.m. Mountain Time on December 15, 2026.") == \
        [("Due", "Proposal", "December 15, 2026, 2:00 p.m. Mountain Time")]
    assert f("Questions must be submitted by November 2, 2026 at 1600 hours MT.")[0][:2] == ("Questions due", "Questions")
    assert f("Volume I, Technical, shall not exceed 25 pages.") == [("Page limit", "Volume I", "25 pages")]
    assert f("The Technical Volume shall not exceed thirty (30) pages.") == [("Page limit", "Technical volume", "30 pages")]
    assert ("Format", "Font", "12-point Times New Roman") in f("Text shall be 12-point Times New Roman with 1-inch margins on 8.5 x 11 inch pages.")
    assert f("Tables and figures may use 10-point Arial font.") == [("Format", "Font in tables and figures", "10-point Arial")]
    assert f("Text shall be single-spaced.") == [("Format", "Line spacing", "single-spaced")]
    assert f("Electronic copies shall be submitted as searchable PDF files.") == [("File", "File format", "Searchable PDF")]
    assert f("File names shall follow the convention Offeror_Volume_Number.pdf.") == [("File", "File naming", "Offeror_Volume_Number.pdf")]
    assert f("The Offeror shall submit three (3) hard copies and one electronic copy.") == [("Copies", "Copies", "three (3) hard copies")]
    assert f("Proposals shall be submitted electronically via the PIEE Solicitation Module.") == [("Delivery", "Delivery", "the PIEE Solicitation Module")]
    assert f("The Offeror must provide an Integrated Master Schedule (IMS) as Attachment J-2.") == \
        [("Deliverable", "Attachment J-2", "Attachment J-2: Integrated Master Schedule (IMS)")]
    assert [x[1] for x in f("The proposal shall consist of Volume I Technical, Volume II Cost and Volume III Past Performance.")] == \
        ["Volume I", "Volume II", "Volume III"]
    assert f("Offerors shall complete and sign the SF 33.") == [("Deliverable", "SF 33", "SF 33")]
    assert f("The Offeror shall identify the five highest risks.") == []
    assert S.days_until("December 15, 2026, 2:00 p.m.", datetime.date(2026, 12, 1)) == 14
    assert S.days_until("12/15/2026", datetime.date(2026, 12, 1)) == 14 and S.days_until("soon") is None
    assert S.match_volume("Volume II", "Kestrel_Vol_2_Mgmt.docx") and not S.match_volume("Volume I", "Kestrel_Vol_2_Mgmt.docx")
    assert S.match_volume("Volume I", "Vol I Technical") and S.match_volume("Technical volume", "Vol I Technical")
    assert not S.match_volume("Proposal", "Vol I Technical")
    # across versions, with real files
    import app_gui as A
    d = tempfile.mkdtemp()
    src = os.path.join(A.BASE_DIR, "test_samples", "rfp", "Kestrel_RFP_Sections_C_L_M.docx")
    amd = amend(src, os.path.join(d, "Kestrel_RFP_Sections_C_L_M_Amd1.docx"),
                [("exceed 50 pages", "exceed 40 pages"), ("December 15, 2026", "December 22, 2026")])
    other = os.path.join(A.BASE_DIR, "test_samples", "rfp", "Kestrel_RFP_Section_L_Part2.pdf")
    lib = DV.Library(os.path.join(d, "documents"))
    doc, v1 = lib.add(src, DV.snapshot(os.path.basename(src), A.typed_chunks(src)), doc_name="RFP", who="tester", when="2026-10-02T09:00")
    lib.add(other, DV.snapshot(os.path.basename(other), A.typed_chunks(other)), doc_name="Section L Part 2", who="tester", when="2026-10-02T09:05")
    R = S.Register(lib)
    rows = {r["key"]: r for r in R.rows()}
    assert {"Due | Proposal", "Page limit | Proposal", "Format | Font", "Format | Margins", "Format | Page size", "File | File format",
            "File | File naming", "Deliverable | Attachment J-2"} <= set(rows), sorted(rows)
    assert not any(r["to_review"] for r in rows.values())
    pl = rows["Page limit | Proposal"]
    assert pl["value"] == "50 pages" and pl["doc"] == "RFP" and pl["version"] == "Original" and "50 pages" in pl["sentence"]
    _d, v2 = lib.add(amd, DV.snapshot(os.path.basename(amd), A.typed_chunks(amd)), doc_id=doc["id"], who="tester", when="2026-10-14T09:00")
    rows = {r["key"]: r for r in R.rows()}
    pl = rows["Page limit | Proposal"]
    assert pl["value"] == "40 pages" and pl["version"] == "Amendment 1" and pl["to_review"]
    assert [(h["version"], h["value"], h["change"]) for h in pl["history"]] == [("Original", "50 pages", "first"), ("Amendment 1", "40 pages", "changed")]
    assert rows["Due | Proposal"]["value"].startswith("December 22, 2026") and rows["Due | Proposal"]["to_review"]
    assert rows["Format | Font"]["history"][-1]["change"] == "same" and not rows["Format | Font"]["to_review"]
    s = S.summary(R.rows())
    assert s["due"]["value"].startswith("December 22") and [p[:2] for p in s["pages"]] == [("Proposal", "40 pages")] and len(s["to_review"]) == 2
    # reviewed, confirmed, edited, hidden, typed in: saved and read back
    assert R.review("Page limit | Proposal", who="tester")
    R.confirm("Format | Font", who="tester")
    R.edit("Format | Margins", "1-inch margins on all sides", who="tester")
    R.hide("Format | Page size")
    k = R.add("Page limit", "Volume II", "15 pages", who="tester")
    R2 = S.Register(DV.Library(lib.folder).load())
    rows = {r["key"]: r for r in R2.rows()}
    assert not rows["Page limit | Proposal"]["to_review"] and rows["Page limit | Proposal"]["reviewed"].endswith("tester")
    assert rows["Format | Font"]["confirmed"].endswith("tester")
    assert rows["Format | Margins"]["value"] == "1-inch margins on all sides" and rows["Format | Margins"]["found"] == "1-inch margins"
    assert "Format | Page size" not in rows and "Format | Page size" in {r["key"] for r in R2.rows(include_hidden=True)}
    assert rows[k]["manual"] and rows[k]["value"] == "15 pages" and k == "Page limit | Volume II"
    assert os.path.isfile(os.path.join(lib.folder, "register.json"))
    # a later version changing the limit again brings the review back
    amd2 = amend(src, os.path.join(d, "Kestrel_RFP_Sections_C_L_M_Amd2.docx"), [("exceed 50 pages", "exceed 45 pages")])
    lib2 = R2.lib
    lib2.add(amd2, DV.snapshot(os.path.basename(amd2), A.typed_chunks(amd2)), doc_id=doc["id"], who="tester", when="2026-10-20T09:00")
    rows = {r["key"]: r for r in R2.rows()}
    assert rows["Page limit | Proposal"]["to_review"] and rows["Page limit | Proposal"]["value"] == "45 pages"
    assert R2.remove_manual(k) and k not in {r["key"] for r in R2.rows()}
    R2.hide("Format | Page size", False)
    assert "Format | Page size" in {r["key"] for r in R2.rows()}
    # QA/QC: our Word volumes against the page limits (Word's saved page count)
    import zipfile

    def vol(name, pages):
        fp = os.path.join(d, name)
        with zipfile.ZipFile(fp, "w") as z:
            z.writestr("word/document.xml", "<w:document/>")
            if pages is not None:
                z.writestr("docProps/app.xml", f"<Properties><Pages>{pages}</Pages></Properties>")
        return fp
    v1, v2, v3 = vol("Kestrel_Vol_1_Technical.docx", 27), vol("Kestrel_Vol_2_Mgmt.docx", 12), vol("Kestrel_Vol_3.docx", None)
    assert S.word_pages(v1) == 27 and S.word_pages(v3) is None
    lim = [{"kind": "Page limit", "subject": "Volume I", "value": "25 pages", "doc": "RFP", "version": "Amendment 1"},
           {"kind": "Page limit", "subject": "Management volume", "value": "15 pages", "doc": "RFP", "version": "Original"},
           {"kind": "Page limit", "subject": "Proposal", "value": "50 pages", "doc": "RFP", "version": "Original"}]
    out = S.page_check({os.path.basename(v1): v1, os.path.basename(v2): v2}, lim, [{"name": "Management", "files": [os.path.basename(v2)]}])
    assert [(o["volume"], o["level"], o["pages"]) for o in out] == [("Kestrel_Vol_1_Technical.docx", "bad", 27),
                                                                     ("Kestrel_Vol_2_Mgmt.docx", "ok", 12), ("All volumes", "ok", 39)], out
    assert out[0]["result"] == "Over by 2 pages" and out[0]["source"] == "RFP, Amendment 1"
    out = S.page_check({os.path.basename(v3): v3}, lim, [])
    assert [(o["volume"], o["level"]) for o in out] == [("All volumes", "info")] and "open each in Word" in out[0]["result"]
    assert S.page_check({os.path.basename(v1): v1}, [], []) == []
    S.export_excel(os.path.join(d, "reqs.xlsx"), R2.rows(include_hidden=True), "Kestrel")
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(d, "reqs.xlsx"))
    assert wb.sheetnames == ["Submission requirements", "History"] and wb["History"].max_row > 5
    # 2.49.0: the submission contact and how the proposal is submitted
    got = S.find("Proposals shall be submitted electronically via DoD SAFE to the Contracting Officer, Jane Doe, at "
                  "jane.doe@us.af.mil no later than December 15, 2026.")
    assert ("Contact", "Submission contact", "Jane Doe (Contracting Officer) jane.doe@us.af.mil") in got and ("Delivery", "Delivery", "DoD SAFE") in got, got
    got = S.find("Questions shall be sent to the Contract Specialist, Mr. Bob Lee, at robert.lee.3@us.af.mil by November 2, 2026.")
    assert got == [("Contact", "Questions contact", "Mr. Bob Lee (Contract Specialist) robert.lee.3@us.af.mil")], got
    print("subreq ok")


if __name__ == "__main__":
    run()
