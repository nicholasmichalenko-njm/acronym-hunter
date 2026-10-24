"""Compliance C3 and complements (2.10.0): learn Keep/Drop from decisions (hashed, no RFP words stored), re-learning a
matrix replaces its contribution, cross-check, mapping from similar mapped rows, merge, amendment changes with word
differences, Excel export of changes, PDF source page with the sentence highlighted, sources keep paths."""
import json, os, sys, tempfile
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import compliance as C, complearn as CL

RFP = os.path.join(ROOT, "test_samples", "rfp")


def synth(n_keep=30, n_drop=15):
    """A matrix the way a team decides: offeror duties kept, Government statements and background dropped."""
    keep = ["The Offeror shall describe its approach to {x}.", "The Contractor shall deliver the {x} report monthly.",
            "Provide a plan for {x} with milestones.", "The proposal shall include a {x} matrix."]
    drop = ["The Government will provide {x} after award.", "This section explains the history of {x}.",
            "The Government may review {x} at any time."]
    topics = ["thermal control", "risk management", "configuration", "safety", "software assurance", "quality", "logistics",
              "training", "security", "test"]
    items = []
    for i in range(n_keep):
        t = keep[i % len(keep)].format(x=topics[i % len(topics)])
        it = C.classify(t, "", ("L.4", "Approach"), "") or {"text": t, "lead": "", "section": "L.4", "section_title": "",
                                                          "actor": "Offeror", "type": "Instruction", "strength": "Mandatory",
                                                          "suggest": "Keep", "modal": "shall", "xrefs": "", "compound": False}
        it.update(id=f"K{i}", decision="Keep", file="x.docx", loc="x", key=C.text_key(t), map="")
        items.append(it)
    for i in range(n_drop):
        t = drop[i % len(drop)].format(x=topics[i % len(topics)])
        it = C.classify(t, "", ("C.1", "Background"), "") or {"text": t, "lead": "", "section": "C.1", "section_title": "",
                                                            "actor": "Unclear", "type": "Requirement", "strength": "Check",
                                                            "suggest": "Review", "modal": "", "xrefs": "", "compound": False}
        it.update(id=f"D{i}", decision="Drop", file="x.docx", loc="x", key=C.text_key(t), map="")
        items.append(it)
    return items


def run():
    d = tempfile.mkdtemp()
    model = os.path.join(d, "config", "compliance_learned.json")
    L = CL.Learner(model)
    assert L.p_keep({"text": "x"}) is None and not L.ready()
    items = synth()
    assert L.learn(os.path.join(d, "Kestrel_Compliance_Matrix.xlsx"), items) == 45
    L.save()
    raw = open(model, encoding="utf-8").read()
    for word in ("thermal", "Offeror", "Government", "risk", "Kestrel"):
        assert word.lower() not in raw.lower(), f"model file must not hold RFP words ({word})"
    M = CL.Learner(model)
    assert M.summary() == {"decisions": 45, "keep": 30, "drop": 15, "matrices": 1}
    # learning the same matrix again replaces its contribution
    M.learn(os.path.join(d, "Kestrel_Compliance_Matrix.xlsx"), items[:20] + items[30:40])
    assert M.summary()["decisions"] == 30
    M.learn(os.path.join(d, "Kestrel_Compliance_Matrix.xlsx"), items)
    new_keep = {"text": "The Offeror shall describe its approach to cable routing.", "lead": "", "type": "Instruction",
                "actor": "Offeror", "strength": "Mandatory", "modal": "shall", "section": "L.5", "suggest": "Keep"}
    new_drop = {"text": "The Government will provide cable routing drawings after award.", "lead": "", "type": "Information",
                "actor": "Government", "strength": "Information", "modal": "will", "section": "C.3", "suggest": "Drop"}
    pk, pd = M.p_keep(new_keep), M.p_keep(new_drop)
    assert pk > 0.8 and pd < 0.2, (pk, pd)
    assert CL.Learner.label(pk) == "Keep" and CL.Learner.label(pd) == "Drop" and CL.Learner.label(0.5) == "Review"
    agree, total, rule = CL.Learner.cross_check(items)
    assert total == 45 and agree >= 40, (agree, total, rule)
    # mapping from similar rows the person already mapped
    rows = [dict(text="The Offeror shall describe the thermal control design and margins.", map="1.1 Thermal Design", decision="Keep", lead=""),
            dict(text="The Offeror shall describe the risk register and mitigation.", map="1.2 Risk Management", decision="Keep", lead="")]
    q = dict(text="Describe thermal control margins for each mode.", decision="Keep", lead="", map="")
    lab, sim = CL.map_from_neighbors(q, rows + [q])
    assert lab == "1.1 Thermal Design" and sim >= 0.35, (lab, sim)
    assert CL.map_from_neighbors(dict(text="Submit two hard copies.", lead=""), rows)[0] == ""
    # merge
    import app_gui as A
    files = sorted(os.listdir(RFP))
    heads = []
    real = C.extract([(f, A.typed_chunks(os.path.join(RFP, f))) for f in files], heads)
    a = [it for it in real if it["text"].startswith("(a) identify")][0]
    b = [it for it in real if it["text"].startswith("(b) describe")][0]
    n0 = len(real)
    m = C.merge_items(real, [a["id"], b["id"]])
    assert len(real) == n0 - 1 and m["text"].startswith("(a) identify") and "(b) describe" in m["text"] and b["id"] in m["notes"]
    # amendment changes: one changed, one added, one removed
    old = C.extract([(f, A.typed_chunks(os.path.join(RFP, f))) for f in files])
    new = C.extract([(f, A.typed_chunks(os.path.join(RFP, f))) for f in files])
    ch = [x for x in new if x["text"].startswith("Proposals shall not exceed 50 pages")][0]
    ch["text"] = ch["text"].replace("50 pages", "40 pages"); ch["key"] = C.text_key(ch["text"])
    new.append(dict(new[0], text="The Offeror shall provide a cybersecurity plan per NIST SP 800-171.", key="added1", id=""))
    gone = [x for x in new if x["text"].startswith("The Offeror may include")][0]
    new.remove(gone)
    for x in old:
        if x["text"].startswith("Proposals shall not exceed"):
            x["decision"] = "Keep"
    merged = C.merge_reread(old, new)
    changes = C.amendment_changes(merged)
    kinds = [c["kind"] for c in changes]
    assert kinds.count("Changed") == 1 and kinds.count("Added") == 1 and kinds.count("Removed") == 1, kinds
    c0 = changes[0]
    assert c0["kind"] == "Changed" and "[-50-] {+40+}" in c0["diff"] and c0["decision"] == "Keep", c0
    xl = os.path.join(d, "changes.xlsx"); C.write_changes_excel(changes, xl)
    from openpyxl import load_workbook
    ws = load_workbook(xl)["Amendment changes"]
    assert ws.max_row == 4 and ws["A2"].value == "Changed"
    # PDF source page with highlight
    img = C.render_page(os.path.join(RFP, "Kestrel_RFP_Section_L_Part2.pdf"), 2, "identify a point of contact.")
    px = img.convert("RGB").getpixel((int(img.width * 0.25), int(img.height * 0.142)))
    yellow = sum(1 for x in range(0, img.width, 4) for y in range(0, int(img.height * 0.3), 4)
                 if (lambda p: p[0] > 230 and p[1] > 200 and p[2] < 200)(img.getpixel((x, y))))
    assert yellow > 50, yellow
    # sources keep the path so Show page can find the file
    W = C.Matrix(os.path.join(d, "m.xlsx")); W.items = real; W.sources = [["Kestrel_RFP_Section_L_Part2.pdf", "AB", "2026-10-23", os.path.join(RFP, "Kestrel_RFP_Section_L_Part2.pdf")]]
    W.save()
    R = C.Matrix(W.path).load()
    assert R.path_of("Kestrel_RFP_Section_L_Part2.pdf").endswith("Kestrel_RFP_Section_L_Part2.pdf") and R.path_of("nope.pdf") == ""
    assert "learned" in R.items[0]
    print("complearn ok")


if __name__ == "__main__":
    run()


def run_pack():
    """Knowledge Pack carries the learning (no words) and installs add-only."""
    import app_gui as A, zipfile
    d = tempfile.mkdtemp()
    A.COMPLIANCE_LEARNED_FILE = os.path.join(d, "mine.json")
    L = CL.Learner(A.COMPLIANCE_LEARNED_FILE); L.learn(os.path.join(d, "Team_A.xlsx"), synth()); L.save()
    A.KNOWLEDGE_BACKUP_DIR = os.path.join(d, "bk"); A.KNOWLEDGE_INFO_FILE = os.path.join(d, "ki.json")
    A.DB_FILE = os.path.join(d, "db.csv"); open(A.DB_FILE, "w").write("Term,Definition,Category\n")
    A.OCR_FIXES_FILE = os.path.join(d, "fx.csv"); A.OCR_IGNORE_FILE = os.path.join(d, "ig.csv"); A.MODELS_DIR = os.path.join(d, "models")
    pack = os.path.join(d, "pack.zip")
    A.make_knowledge_pack(pack, "test")
    names = zipfile.ZipFile(pack).namelist()
    assert "knowledge/compliance_learned.json" in names, names
    other = os.path.join(d, "other.json")
    O = CL.Learner(other); O.learn(os.path.join(d, "Team_B.xlsx"), synth(10, 5)); O.save()
    A.COMPLIANCE_LEARNED_FILE = other
    rep = A.install_knowledge_pack(pack)
    assert "Compliance learning: 45 decisions added" in rep, rep
    assert CL.Learner(other).summary() == {"decisions": 60, "keep": 40, "drop": 20, "matrices": 2}
    assert "Compliance learning: 0 decisions added" in A.install_knowledge_pack(pack), "add-only: installing twice adds nothing"
    print("complearn pack ok")


if __name__ == "__main__":
    run_pack()
