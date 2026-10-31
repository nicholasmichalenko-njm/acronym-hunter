"""3.0.4 graphic needs, offscreen: in Review, Graphic need (P) tags the selected sentence, or just the words dragged across
in it, with what the graphic shows, its kind, section and owner; the sentence's card lists it; Effort Home's Graphic needs
card lists every one with where it was found, opens the sentence in Review, changes it (in work, done) or removes it;
+ Graphic need adds one by hand and Graphic on a note lists the note; Export makes a workbook; it is all in the effort
record, and the card is in the Proposal Manager, Production and Author layouts."""
import os, sys, tempfile, time
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
os.environ["QT_QPA_PLATFORM"] = "offscreen"
os.environ["AH_TEST"] = "1"
os.environ["SK_SANDBOX_TEMP"] = "temp_qt"
import app_gui as A
import effort as E, workspace as W, office_pdf as OP, effortrec as RR, graphics as GX, branding as BR
OP.available = lambda: False
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QMessageBox
QMessageBox.information = staticmethod(lambda *a, **k: None)
from skqt import theme
from skqt.mainwin import MainWindow

checks = []


def check(name, ok, info=""):
    checks.append(bool(ok)); print("PASS" if ok else "FAIL", name, info)


tmp = tempfile.mkdtemp()
cfg = os.path.join(tmp, "config"); os.makedirs(cfg)
EDIR = os.path.join(tmp, "Kestrel"); os.makedirs(EDIR)
LIST = os.path.join(EDIR, "Kestrel_Effort_List.xlsx")
L = E.EffortList(LIST, "Kestrel"); L.pending = True; L.save()
reg = E.Registry(os.path.join(cfg, "efforts.json")); reg.put("Kestrel", path=LIST)
ws0 = W.Workspace(W.data_dir_for(LIST), "Kestrel")
ws0.set("outline_text", "Volume I Technical\n1.1 Technical Approach\n1.2 Risk")
PDFS = [os.path.join(A.BASE_DIR, "test_samples", "rfp_walk", "Section_L_print.pdf")]
app = QApplication([]); theme.apply(app)
w = MainWindow(app, cfg, A.APP_VERSION); w.resize(1500, 950); w.show()


def wait(cond, secs=60):
    t = time.time()
    while not cond() and time.time() - t < secs:
        app.processEvents(); time.sleep(0.02)
    return cond()


w.open_effort("Kestrel")
w.read_page.read(PDFS, os.path.join(EDIR, "Kestrel_Compliance_Matrix.xlsx"))
check("the matrix starts", wait(lambda: w.review is not None and w.stack.currentWidget() is w.assign_page))
pg = w.assign_page; app.processEvents()
pg.files.setCurrentText("Section_L_print.pdf"); app.processEvents()
it = next(x for x in w.review.M.items if x["text"].startswith("The Offeror shall describe its technical approach"))
pg.set_target("1.1 Technical Approach")
pg.select(it["id"]); app.processEvents()
check("Review has Graphic need beside More captures", "Graphic need" in pg.cap_keys["graphic"].text())
from skqt import captureview as CV
d = CV.graphic_dialog(pg, it["text"], w.review.labels(), ["A. Lee"], pg._graphic_section())
check("... its form starts from the sentence, in the target section", d.w["concept"].text().startswith("The Offeror shall")
      and d.w["section"].currentText() == "1.1 Technical Approach")
gid = pg.do_capture("graphic", {"concept": "Technical approach overview", "kind": "Diagram", "section": "1.1 Technical Approach", "owner": "A. Lee"})
rec = RR.Record(w.ws.root)
g = GX.get(rec, gid)
check("Graphic need tags the sentence: what, kind, section, owner, where it was found", g and g["concept"] == "Technical approach overview"
      and g["kind"] == "Diagram" and g["owner"] == "A. Lee" and g["sid"] == it["id"] and g["doc"] == "Section_L_print.pdf"
      and g["quote"].startswith("The Offeror shall") and g["status"] == "Needed", g)
check("... the sentence's card lists it", any("Technical approach overview" in x.text() for x in pg.gfx_labels))
pg.show_grab(it["id"], 1, 3, final=True); app.processEvents()
gid2 = pg.do_capture("graphic", {"concept": "Approach words", "kind": "", "section": "", "owner": ""})
g2 = GX.get(RR.Record(w.ws.root), gid2)
check("... and the words dragged across are what it quotes", g2 and g2["quote"] == pg.grab["text"] and len(g2["quote"]) < len(it["text"]), g2 and g2["quote"])
check("... into the timeline", any(e["kind"] == "graphic" and "Technical approach overview" in e["text"] for e in RR.Record(w.ws.root).events()))
w.go("home"); app.processEvents()
hp = w.home_page; hp.refresh(); app.processEvents()
check("Effort Home lists them under Graphic needs", hp.gfx_card.isVisibleTo(w) and gid in hp.gfx_now and gid2 in hp.gfx_now
      and "2 open" in hp.gfx_card.title.text(), hp.gfx_card.title.text())
check("... Open goes to the sentence in Review", hp.open_sentence(it["id"]) and w.stack.currentWidget() is pg and pg.cur == it["id"])
w.go("home"); app.processEvents()
hp.graphic_dialog(gid, {"concept": "Technical approach overview", "kind": "Diagram", "section": "1.1 Technical Approach",
                        "owner": "A. Lee", "status": "In work"})
check("changing one: in work", GX.get(RR.Record(w.ws.root), gid)["status"] == "In work")
hp.graphic_dialog(gid, {"concept": "Technical approach overview", "kind": "Diagram", "section": "1.1 Technical Approach",
                        "owner": "A. Lee", "status": "Done"})
check("... done, and the count says so", GX.get(RR.Record(w.ws.root), gid)["done"] and "1 open" in hp.gfx_card.title.text(), hp.gfx_card.title.text())
gid3 = hp.graphic_dialog(None, {"concept": "Org chart of the program office", "kind": "Org chart", "section": "1.2 Risk", "owner": "",
                                "status": "Needed"})
check("+ Graphic need adds one by hand", GX.get(RR.Record(w.ws.root), gid3)["kind"] == "Org chart" and gid3 in hp.gfx_now)
nid = hp.note_dialog(text="Show the thermal margin across the mission phases.")
gid4 = hp.graphic_from_note(nid)
g4 = GX.get(RR.Record(w.ws.root), gid4)
check("Graphic on a note lists the note", g4 and g4["note"] == nid and g4["concept"].startswith("Show the thermal margin")
      and GX.where_text(g4) == "From a note")
check("... and the note no longer offers it", not any(b.text() == "Graphic" for b in hp.notes.findChildren(type(pg.cap_keys["graphic"]))))
hp.graphic_dialog(gid2, {"remove": True})
check("Remove takes one off", GX.get(RR.Record(w.ws.root), gid2) is None)
x = hp.export_graphics(os.path.join(tmp, "g.xlsx"))
wb = load_workbook(x); sh = wb.active
rows = [[c.value for c in r] for r in sh.iter_rows(values_only=False)]
check("Export: a workbook of every graphic need, open first, with where it was found", rows[0][1] == "Graphic needed" and len(rows) == 4
      and rows[-1][5] == "Done" and any(r[6] and "Section_L_print.pdf" in r[6] for r in rows[1:]), [r[1] for r in rows])
check("the export has a name in Folders and file names", "Graphic_Needs" in w.out_path("graphic_needs"))
check("the card is in the Proposal Manager, Production and Author layouts", all("graphics" in BR.ROLES[r] for r in
      ("Proposal Manager", "Production", "Author")))
w.people = lambda: ["A. Lee"]
hp.who.addItem("A. Lee"); hp.who.setCurrentText("A. Lee"); hp.draw_graphics(RR.Record(w.ws.root))
check("one person picked: only their graphic needs", hp.gfx_now == [gid], hp.gfx_now)
w.close()
t = sum(checks)
print(f"QT GRAPHICS {'PASS' if t == len(checks) else 'FAIL'} ({t}/{len(checks)})")
sys.exit(0 if t == len(checks) else 1)
