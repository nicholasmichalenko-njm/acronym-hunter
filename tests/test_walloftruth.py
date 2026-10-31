"""3.3.0 Wall of Truth, without screens: truths (add, change keeps the old value, confirm, stale, remove), the team's
wall (save, add to an effort without duplicates), numbers read the way writers write them, and the draft check: a
number beside a truth's words that differs, an old value still in use, a wrong form, the right value left alone, a
finding marked right as written left out; the Word page, the workbook and the readable copy."""
import datetime, os, sys, tempfile, zipfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import walloftruth as WT, effortrec as R, docxw, readable as RD
from openpyxl import load_workbook

checks = []


def check(name, ok, info=""):
    checks.append(bool(ok)); print("PASS" if ok else "FAIL", name, info)


tmp = tempfile.mkdtemp()
cfg = os.path.join(tmp, "cfg")
rec = R.Record(os.path.join(tmp, "e")); os.makedirs(rec.folder)
check("numbers read as written: $, commas, percent, K/M, million", [WT.norm_number(x) for x in ("$1,200", "15 percent", "15%", "2.5M", "1.5 million", "142.0")]
      == ["1200", "15%", "15%", "2.5m", "1.5m", "142"])
check("a truth needs what it is and its value", WT.add(rec, "Cleared staff", "") == "" and WT.add(rec, "", "142") == "")
a = WT.add(rec, "Cleared staff", "142", "Number", ["cleared staff", "cleared employees"], owner="HR", source="HR report")
b = WT.add(rec, "Years supporting the customer", "12", "Number", ["years"])
c = WT.add(rec, "Customer", "the Government", "Term", [], ["the client", "the customer agency"])
d = WT.add(rec, "Program name", "Kestrel Thermal Subsystem", "Name", [], ["Kestrel Thermal System"])
check("truths in order, with today's confirmation", [t["label"] for t in WT.items(rec)][:2] == ["Cleared staff", "Years supporting the customer"]
      and WT.get(rec, a)["confirmed"] == datetime.date.today().isoformat())
check("a truth with no words given looks beside its own name", WT.get(R.Record(rec.folder), c)["context"] == ["customer"])
check("a new value keeps the old one", WT.update(rec, a, value="145") and WT.get(rec, a)["old"][0]["value"] == "142" and WT.get(rec, a)["value"] == "145")
check("... and is logged", any("Cleared staff now 145 (was 142)" in e["text"] for e in rec.events()))
WT.get(rec, b)["confirmed"] = "2020-01-01"; rec.save()
check("not confirmed for 30 days is stale; Confirm dates it today", WT.stale(WT.get(rec, b)) and WT.confirm(rec, b) and not WT.stale(WT.get(rec, b)))
# the draft
doc = docxw.Doc()
doc.heading("1.1 Staffing", 2)
doc.para("We bring 142 cleared staff to the program. Our 145 cleared employees are ready on day one. For 10 years we have "
         "supported the client. That is over 12 years of support. The Kestrel Thermal System flies in 2027.")
doc.heading("1.2 Facilities", 2)
doc.para("Our facility has 30 years of history and 145 cleared staff.")
p = os.path.join(tmp, "Volume I.docx"); doc.save(p)
found = WT.check(rec, [("Volume I", p)], ["1.1 Staffing", "1.2 Facilities"])
probs = sorted((f["label"], f["section"], f["problem"]) for f in found)
check("the check: an old value still in use", ("Cleared staff", "1.1 Staffing", "old value 142; now 145") in probs, probs)
check("... a number beside the words that differs", ("Years supporting the customer", "1.1 Staffing", "says 10; the wall says 12") in probs
      and ("Years supporting the customer", "1.2 Facilities", "says 30; the wall says 12") in probs)
check("... wrong forms of terms and names", ("Customer", "1.1 Staffing", "says “the client”; the wall says “the Government”") in probs
      and any(p_[0] == "Program name" for p_ in probs))
check("... the right value is left alone", not any("145 cleared" in f["sentence"] and f["label"] == "Cleared staff" for f in found)
      and not any("over 12 years" in f["sentence"] for f in found))
check("... five findings in all", len(found) == 5, probs)
k = next(f["key"] for f in found if "says 30" in f["problem"])
check("right as written leaves it out next time", WT.ignore(rec, k) and len(WT.check(rec, [("Volume I", p)], ["1.1 Staffing", "1.2 Facilities"])) == 4)
# the team's wall
check("save to the team's wall; saving the same again changes nothing", WT.to_team(cfg, WT.get(rec, a)) and not WT.to_team(cfg, WT.get(rec, a)))
WT.to_team(cfg, {"label": "Employees", "value": "1,800", "kind": "Number", "context": ["employees"], "wrong": []})
rec2 = R.Record(os.path.join(tmp, "e2")); os.makedirs(rec2.folder)
WT.add(rec2, "Employees", "1,750")
check("add the team's facts: only those the effort lacks", WT.from_team(rec2, cfg) == 1 and sorted(t["label"] for t in WT.items(rec2)) == ["Cleared staff", "Employees"]
      and next(t for t in WT.items(rec2) if t["label"] == "Employees")["value"] == "1,750")
# outputs
out = WT.export_word(os.path.join(tmp, "wall.docx"), rec, "Kestrel", ("CUI", "502B85"))
x = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
check("the Word page: by kind, with what never to write and earlier values", "Wall of Truth: Kestrel" in x and "Numbers" in x and "Terms" in x
      and "the client" in x and ">142<" in x and "145" in x)
xl = WT.export_findings(os.path.join(tmp, "f.xlsx"), found, "Kestrel")
check("the findings workbook", load_workbook(xl).active.max_row == len(found) + 1)
tbl = RD.more(rec, None).get("Wall of Truth.csv")
check("the readable copy has Wall of Truth.csv", tbl and len(tbl[1]) == 4 and tbl[0][0] == "What")
check("removing a truth", WT.remove(rec, d) and WT.get(R.Record(rec.folder), d) is None)
n = sum(checks)
print(f"WALL OF TRUTH {'PASS' if n == len(checks) else 'FAIL'} ({n}/{len(checks)})")
sys.exit(0 if n == len(checks) else 1)
