"""2.68.0 in Qt, offscreen: Read RFP will not read a document without an initial and a type (a taken initial is
refused); Assign works with no structure yet (R keeps the sentence with no section); L and M on requirements (the E key
and the buttons), the card shows the document, L or M and the reference (R-L-3.1); the structure is built beside the
page: + Volume, + Heading, + Subheading down eight levels, From sentence, rename, move, indent, outdent, delete, and
the requirements follow their sections; the volume switcher shows one volume and new sections go into it; a sentence in
two volumes shows both; Customer Documents shows and sets each document's initial and type; the matrix has Doc, L/M,
Section and Reference, and a tab per volume."""
import os, sys, tempfile, time
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
os.environ["QT_QPA_PLATFORM"] = "offscreen"
os.environ["AH_TEST"] = "1"
os.environ["SK_SANDBOX_TEMP"] = "temp_qt"
import app_gui as A
import compliance as C, effort as E, setupdoc as SD, workspace as W, office_pdf as OP, outline_ops as O, docversions as DV
OP.available = lambda: False
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox
QMessageBox.information = staticmethod(lambda *a, **k: None)
from skqt import theme
from skqt.mainwin import MainWindow

checks = []


def check(name, ok, info=""):
    checks.append(bool(ok)); print("PASS" if ok else "FAIL", name, info)


tmp = tempfile.mkdtemp()
cfg = os.path.join(tmp, "config"); os.makedirs(cfg)
EDIR = os.path.join(tmp, "Kestrel"); os.makedirs(EDIR)
LIST = os.path.join(EDIR, "Kestrel_Effort_List.xlsx")
L = E.EffortList(LIST, "Kestrel"); L.pending = True; L.save()
E.Registry(os.path.join(cfg, "efforts.json")).put("Kestrel", path=LIST)
PDFS = [os.path.join(A.BASE_DIR, "test_samples", "rfp_walk", n) for n in ("Section_L_print.pdf", "SOW.pdf")]
app = QApplication([])
theme.apply(app)
w = MainWindow(app, cfg, A.APP_VERSION)
w.resize(1500, 950); w.show()


def wait(cond, secs=60):
    t = time.time()
    while not cond() and time.time() - t < secs:
        app.processEvents(); time.sleep(0.02)
    return cond()


w.open_effort("Kestrel")
rp = w.read_page
out = os.path.join(EDIR, "Kestrel_Compliance_Matrix.xlsx")
check("both files need an initial and a type", len(rp.unidentified(PDFS)) == 2)
check("a taken initial is refused: nothing is read", rp.read(PDFS, out, identities={PDFS[0]: ("R", "RFP"), PDFS[1]: ("R", "SOW")}) is None
      and w.review is None)
check("no initial: refused", rp.read(PDFS, out, identities={PDFS[0]: ("", "RFP"), PDFS[1]: ("S", "SOW")}) is None)
rp.read(PDFS, out, identities={PDFS[0]: ("R", "Section L"), PDFS[1]: ("S", "SOW")})
check("with an initial and type each, Read starts the matrix", wait(lambda: w.review is not None and w.stack.currentWidget() is w.assign_page))
M = w.review.M
check("the matrix knows each file's initial and type", SD.docs_of(M)["Section_L_print.pdf"]["initial"] == "R"
      and SD.docs_of(M)["SOW.pdf"]["type"] == "SOW", SD.docs_of(M))
check("the RFP's Section L headings are the starting structure", bool(M.outline), len(M.outline))
pg = w.assign_page
M.outline[:] = []                                  # this team builds its own as it reads
SD.set_volumes(M, "", "", [])
pg.after()
app.processEvents()
check("no structure: the tree is empty and Assign still works", pg.tree.topLevelItemCount() == 0)
pg.files.setCurrentText("Section_L_print.pdf"); app.processEvents()
by = lambda s: next(it for it in M.items if it["text"].startswith(s))
a, b = by("The Offeror shall describe its technical approach"), by("Each volume shall be written")


def click(iid, ctrl=False):
    box = pg.view.boxes[iid]
    pg.view.ensureVisible(box)
    app.processEvents()
    r = box.path().boundingRect()
    pt = pg.view.mapFromScene(r.left() + 6, r.top() + 4)
    QTest.mouseClick(pg.view.viewport(), Qt.LeftButton, Qt.ControlModifier if ctrl else Qt.NoModifier, pt)
    app.processEvents()


click(a["id"])
QTest.keyClick(pg.view, Qt.Key_R); app.processEvents()
check("R with no structure: a compliance requirement with no section yet", a["decision"] == "Keep" and not a.get("map"))
check("L suggested from the section (L.3.1)", SD.lm_of(M, a) == "L" and not a.get("lm"))
QTest.keyClick(pg.view, Qt.Key_L); app.processEvents()
check("L sets it", a.get("lm") == "L")
check("the card: the document, L, and the reference", "Document R (Section L)" in pg.ident_label.text()
      and pg.ref_label.text().endswith("R-L-3.1"), (pg.ident_label.text(), pg.ref_label.text()))
click(b["id"])
QTest.keyClick(pg.view, Qt.Key_E); app.processEvents()
check("E (M) makes a sentence a requirement designated M in one step", b["decision"] == "Keep" and b.get("lm") == "M")
check("the margin shows L or M beside the ID", any(l.text().split()[1:3] == ["M", b["id"]] for l in pg.labels),
      [l.text() for l in pg.labels][:6])
# build the structure as you go
v1 = pg.new_volume("Technical")
check("+ Volume: Volume I Technical, the volume you are working in", v1 == "Volume I Technical" and pg.active_volume == v1
      and pg.volume.currentData() == v1, (v1, pg.volume.currentData()))
h1 = pg.new_section(title="Technical Approach")
check("+ Heading with the volume as target: its first section, 1", h1 == "1 Technical Approach", h1)
deep = h1
for k in range(7):
    deep = pg.new_section(sub=True, title=f"Level {k + 3}")
check("+ Subheading seven times: eight levels, numbered 1.1.1.1.1.1.1.1", deep.startswith("1.1.1.1.1.1.1.1 ")
      and M.outline[O.index_of(M, deep)]["level"] == 9, deep)
click(a["id"]); QTest.keyClick(pg.view, Qt.Key_R); app.processEvents()
check("R puts the sentence into the new deepest section", a["map"] == deep)
pg.set_target(h1)
click(b["id"])
fs = pg.section_from_sentence()
check("From sentence: a section named with the sentence's words", fs.startswith("1.2 Each volume shall be written"), fs)
QTest.keyClick(pg.view, Qt.Key_R); app.processEvents()
check("and R puts it there", b["map"] == fs)
mv = pg.outline_op(O.move, -1)
check("Up: it moves above its sibling, its requirement follows", mv.startswith("1.1 Each volume") and b["map"] == mv
      and a["map"].startswith("1.2.1"), (mv, a["map"]))
ind = pg.outline_op(O.indent)
check("Indent: none above at its level, so it stays", ind is None)
pg.set_target(a["map"].split(" ")[0] and next(l for l in O.labels(M) if l.startswith("1.2 ")))
ind = pg.outline_op(O.indent)
check("Indent: under the section above", ind.startswith("1.1.1 ") and b["map"] == mv, ind)
od = pg.outline_op(O.outdent)
check("Outdent: back up a level", od.startswith("1.2 "), od)
rn = pg.rename_section(mv, "Volume format")
check("Rename keeps the requirement in it", rn == "1.1 Volume format" and b["map"] == rn, rn)
v2 = pg.new_volume("Management")
check("a second volume: the switcher has both", pg.volume.count() == 3 and pg.active_volume == v2)
pm = pg.new_section(title="Program Management")
check("new sections go into the volume you are working in, numbered from 1", pm == "1 Program Management"
      and O.volumes_of_item(M, {"map": pm, "id": "x"}) == [v2], pm)
check("the structure shows only that volume", pg.tree.topLevelItemCount() == 1 and pg.tree.topLevelItem(0).text(0).startswith(v2))
SD.add_section(M, b, pm); pg.after()
click(b["id"])
check("a sentence in two volumes shows both", [l.text() for l in pg.vol_badges.findChildren(type(pg.ident_label))][1:] == [v1, v2]
      or all(v in " ".join(x.text() for x in pg.vol_badges.findChildren(type(pg.ident_label))) for v in (v1, v2)))
pg.pick_volume("")
check("All volumes shows everything", pg.tree.topLevelItemCount() == 2)
gone = pg.delete_section(rn)
check("Delete: its sentences go up to the section above", gone == v1 and b["map"] in (v1, pm) or SD.item_sections(M, b), gone)
w.save_now()
# Customer Documents: the initial and type of each document
w.go("docs"); app.processEvents()
dp = w.docs_page
docs = os.path.join(EDIR, "docs"); os.makedirs(docs)
import shutil
p1 = shutil.copy2(PDFS[0], os.path.join(docs, "Kestrel_Amendment_0001.pdf"))
made = dp.add([p1], choices=[{"choice": "new"}])
d = made[0][0]
check("a new document gets its type and initial (an amendment: A1)", d["type"] == "Amendment" and d["initial"] == "A1", (d.get("type"), d.get("initial")))
check("a taken initial is refused", dp.add([PDFS[1]], choices=[{"choice": "new", "initial": "A1"}]) is None)
check("the documents list shows the initial and type", dp.tree.topLevelItem(0).text(1) == "A1" and dp.tree.topLevelItem(0).text(2) == "Amendment")
check("set the initial and type", dp.identity("K1", "Amendment", doc_id=d["id"]) == ""
      and DV.Library(w.ws.p("documents")).load().doc(d["id"])["initial"] == "K1")
d2 = dp.add([PDFS[1]], choices=[{"choice": "new", "initial": "S2", "type": "SOW"}])[0][0]
check("a taken one is refused there too", dp.identity("K1", "SOW", doc_id=d2["id"]) != ""
      and DV.Library(w.ws.p("documents")).load().doc(d2["id"])["initial"] == "S2")
# the matrix: Doc, L/M, Section, Reference; a tab per volume
w.go("matrix"); app.processEvents()
mp = w.matrix_page
res = mp.export(os.path.join(tmp, "m.xlsx"))
from openpyxl import load_workbook
wb = load_workbook(res["path"])
ws = wb["Compliance Matrix"]
heads = [c.value for c in ws[4]]
check("the matrix columns: Reference, Doc, L/M, Section", heads[3:7] == ["Reference", "Doc", "L/M", "Section"], heads)
rows = [[c.value for c in r] for r in ws.iter_rows(min_row=5)]
ra = next(r for r in rows if str(r[heads.index("RFP content")] or "").startswith("The Offeror shall describe its technical approach"))
check("a row: R, L, and R-L-3.1", ra[heads.index("Doc")] == "R" and ra[heads.index("L/M")] == "L" and ra[heads.index("Reference")] == "R-L-3.1", ra[:8])
check("a tab per volume", v1 in wb.sheetnames and v2 in wb.sheetnames, wb.sheetnames)
w.close()
n = sum(checks)
print(f"QT STRUCTURE LM PASS ({n}/{len(checks)})" if n == len(checks) else f"QT STRUCTURE LM FAIL ({n}/{len(checks)})")
sys.exit(0 if n == len(checks) else 1)
