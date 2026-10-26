"""2.38.0 references from the document itself and people per section: unnumbered files are listed until a rule is
chosen; name + paragraph, page + paragraph and prefix references; a section's lead and contributors, inherited by
subsections and by its requirements unless they have their own owner; the Owner column shows the inherited lead and a
reload keeps it inherited; the Response Setup names the people and can include the requirement tables."""
import os, sys, tempfile, zipfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import compliance as C, setupdoc as SD


def row(i, f, text, sec="", dec="Keep", mp="", loc=""):
    return {"id": f"X-{i:02d}", "text": text, "lead": "", "file": f, "page": None, "section": sec, "section_title": "",
            "actor": "", "type": "Instruction", "strength": "", "suggest": "", "modal": "", "xrefs": "", "compound": False,
            "decision": dec, "map": mp, "owner": "", "status": "", "response_page": "", "notes": "", "links": "",
            "learned": "", "new": "", "context": "", "loc": loc, "key": C.text_key(text), "context_for": "", "draft": "", "draft_cf": ""}


def run():
    d = tempfile.mkdtemp()
    M = C.Matrix(os.path.join(d, "m.xlsx"), "K")
    M.items = [row(1, "Section_L.docx", "The Offeror shall describe the thermal design.", sec="L.4", mp="1.1 Thermal",
                   loc="Section_L.docx (section 'L.4 Tech', para 3)"),
               row(2, "Kestrel_Design_Note_Amd1.docx", "The radiator shall reject 400 W.", mp="1.1 Thermal",
                   loc="Kestrel_Design_Note_Amd1.docx (para 7-8)"),
               row(3, "SOW.pdf", "The Contractor shall deliver the schedule.", mp="1.2 Schedule", loc="SOW.pdf (p.4, para 2)"),
               row(4, "SOW.pdf", "The Contractor shall hold monthly reviews.", mp="1.2 Schedule", loc="SOW.pdf (p.5, para 1)")]
    M.items[2]["page"] = 4
    M.outline = C.parse_outline("Volume I Technical\n1.1 Thermal\n1.2 Schedule")
    # files with no numbering need a rule
    assert SD.files_needing_reference(M) == ["Kestrel_Design_Note_Amd1.docx", "SOW.pdf"]
    assert SD.ref_of(M, M.items[0])[0] == "L.4"
    SD.set_refmode(M, "Kestrel_Design_Note_Amd1.docx", "para")
    assert SD.ref_of(M, M.items[1]) == ("Kestrel Design Note, para 7", "The document itself"), SD.ref_of(M, M.items[1])
    SD.set_refmode(M, "SOW.pdf", "page", name="SOW")
    assert SD.ref_of(M, M.items[2])[0] == "SOW, p. 4, para 2", SD.ref_of(M, M.items[2])
    SD.set_refmode(M, "SOW.pdf", "prefix", prefix="SOW")
    assert [SD.ref_of(M, it)[0] for it in M.items[2:]] == ["SOW-1", "SOW-2"]
    assert SD.files_needing_reference(M) == []
    # people: a volume's lead flows down to its sections and their requirements
    SD.set_people(M, "Volume I Technical", "J. Smith", "A. Lee,  B. Chen")
    SD.set_people(M, "1.2 Schedule", "R. Diaz")
    assert SD.section_people(M, "1.1 Thermal") == {"lead": "J. Smith", "with": "A. Lee, B. Chen", "from": "Volume I Technical"}
    assert SD.owner_of(M, M.items[0]) == "J. Smith" and SD.owner_of(M, M.items[2]) == "R. Diaz"
    M.items[1]["owner"] = "K. Park"
    assert SD.owner_of(M, M.items[1]) == "K. Park"
    assert SD.people(M) == ["A. Lee", "B. Chen", "J. Smith", "K. Park", "R. Diaz"]
    assert SD.sections_of(M, "R. Diaz") == ["1.2 Schedule"] and "1.1 Thermal" in SD.sections_of(M, "A. Lee")
    SD.set_people(M, "1.2 Schedule", "", "")
    assert SD.owner_of(M, M.items[2]) == "J. Smith"                   # back to the volume's lead
    SD.set_people(M, "1.2 Schedule", "R. Diaz")
    # the Owner column shows the inherited lead; a reload keeps it inherited, and keeps the real owner
    M.save()
    from openpyxl import load_workbook
    ws = load_workbook(M.path)["Compliance Matrix"]
    hdr = [c.value for c in ws[1]]
    owners = [ws.cell(row=r, column=hdr.index("Owner") + 1).value for r in range(2, 6)]
    assert owners == ["J. Smith", "K. Park", "R. Diaz", "R. Diaz"], owners
    N = C.Matrix(M.path).load()
    assert [it["owner"] for it in N.items] == ["", "K. Park", "", ""], [it["owner"] for it in N.items]
    assert SD.setup_of(N)["refmode"]["SOW.pdf"]["prefix"] == "SOW" and SD.section_people(N, "1.1 Thermal")["lead"] == "J. Smith"
    # Response Setup: people under each heading; tables only when asked
    out = os.path.join(d, "out"); os.makedirs(out)
    paths = SD.export(N, out)
    x = zipfile.ZipFile(paths[0]).read("word/document.xml").decode("utf-8")
    assert "_Response_Setup_" in os.path.basename(paths[0])
    assert "Lead: J. Smith" in x and "Contributors: A. Lee, B. Chen" in x and "Lead: R. Diaz" in x
    assert "<w:tbl>" not in x
    SD.setup_of(N)["tables"] = True
    x = zipfile.ZipFile(SD.export(N, out, when=__import__("datetime").datetime(2026, 1, 1))[0]).read("word/document.xml").decode("utf-8")
    assert x.count("<w:tbl>") == 2 and "SOW-1" in x and "Kestrel Design Note, para 7" in x and ">K. Park<" in x
    print("refs owners ok")


if __name__ == "__main__":
    run()
