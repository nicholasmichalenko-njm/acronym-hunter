"""2.41.0 deliverables: suggested from the submission requirements (add or not ours); kind and format guessed; a
volume's page limit found by its number, our Word file's page count against it, over the limit flagged; a changed
page limit and a newer form version show as changes; status, owner and stage marks logged; the checklist export;
kept in the effort record across a reload."""
import os, sys, tempfile, zipfile
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE)); sys.path.insert(0, HERE)
os.environ.setdefault("AH_OCR_ENGINE", "tesseract")
import deliverables as D, effortrec as ER, subreq as SR, docversions as DV
from test_docversions import amend


def run():
    assert [D.kind_of(n) for n in ("Volume I", "Technical volume", "Attachment J-2", "SF 33", "Representations and Certifications",
                                   "Small Business Subcontracting Plan", "Oral presentation slides", "Cover letter")] == \
        ["Volume", "Volume", "Attachment", "Form", "Form", "Plan", "Presentation", "Other"]
    import app_gui as A
    d = tempfile.mkdtemp()
    src = os.path.join(A.BASE_DIR, "test_samples", "rfp", "Kestrel_RFP_Sections_C_L_M.docx")
    lib = DV.Library(os.path.join(d, "documents"))
    doc, v1 = lib.add(src, DV.snapshot(os.path.basename(src), A.typed_chunks(src)), doc_name="RFP", who="t", when="2026-10-02T09:00")
    reg = SR.Register(lib)
    reg.add("Page limit", "Volume I", "30 pages", who="t")
    rec = ER.Record(d)
    rows = reg.rows()
    sug = D.suggestions(rec, rows)
    assert [r["subject"] for r in sug] == ["Attachment J-2"], [r["subject"] for r in sug]
    assert D.default_format("Volume", rows) == "Word, sent as PDF" and D.default_format("Volume", [{"kind": "File", "subject": "File format", "value": "Searchable PDF"}]) == "Word, sent as Searchable PDF" and D.default_format("Form") == "Signed PDF"
    j2 = D.add(rec, sug[0]["subject"], req_key=sug[0]["key"], owner="A. Lee", who="pm")
    assert get_kind(rec, j2) == "Attachment" and not D.suggestions(rec, rows)
    # a volume: limit by its number, our file's page count
    vol = os.path.join(d, "Kestrel_Vol_1_Technical.docx")
    with zipfile.ZipFile(vol, "w") as z:
        z.writestr("word/document.xml", "<w:document/>")
        z.writestr("docProps/app.xml", "<Properties><Pages>27</Pages></Properties>")
    v = D.add(rec, "Volume I Technical", fmt=D.default_format("Volume", rows), owner="J. Smith", file=vol, who="pm")
    st = D.state(D.get(rec, v), rows, lib, (), rec)
    assert st["limit"] == "30 pages" and st["pages"] == 27 and not st["over"] and D.status_text(D.get(rec, v), st) == "Not started  (27 of 30 pages)"
    reg.add("Page limit", "Volume I", "25 pages", who="t")             # the typed-in limit changed
    st = D.state(D.get(rec, v), reg.rows(), lib, (), rec)
    assert st["over"] and D.status_text(D.get(rec, v), st) == "Over: 27 of 25 pages" and "over" in st["flags"]
    # a changed requirement behind a deliverable: the RFP's own page limit changes in an amendment
    amd = amend(src, os.path.join(d, "RFP_Amd1.docx"), [("exceed 50 pages", "exceed 40 pages"), ("an Integrated Master Schedule (IMS) as", "an Integrated Master Schedule (IMS) and a risk register as")])
    lib.add(amd, DV.snapshot(os.path.basename(amd), A.typed_chunks(amd)), doc_id=doc["id"], who="t", when="2026-10-14T09:00")
    rows = reg.rows()
    stj = D.state(D.get(rec, j2), rows, lib, (), rec)
    assert any("changed in RFP, Amendment 1" in c for c in stj["changed"]), stj["changed"]
    # a customer form with two versions: fill the newest; the old one stays
    f1 = os.path.join(d, "J-3_Pricing_Template.docx"); f2 = os.path.join(d, "J-3_Pricing_Template_v2.docx")
    for p, t in ((f1, "CLIN 0001 price"), (f2, "CLIN 0001 price and CLIN 0002 travel")):
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("word/document.xml", f'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>{t}</w:t></w:r></w:p></w:body></w:document>')
    fd, fv1 = lib.add(f1, [{"s": "", "t": "", "x": "CLIN 0001 price", "l": ""}], doc_name="Attachment J-3 Pricing Template", who="t", when="2026-10-02T10:00")
    cost = D.add(rec, "Volume IV Cost", kind="Volume", doc_id=fd["id"], form_version=fv1["id"], owner="B. Chen", who="pm")
    assert D.state(D.get(rec, cost), rows, lib, (), rec)["form"]["newer"] is False
    _d, fv2 = lib.add(f2, [{"s": "", "t": "", "x": "CLIN 0001 price and CLIN 0002 travel", "l": ""}], doc_id=fd["id"], who="t", when="2026-10-20T10:00")
    st = D.state(D.get(rec, cost), rows, lib, (), rec)
    assert st["form"]["newer"] and st["form"]["latest"] == "Amendment 1" and st["form"]["working"] == "Original"
    assert any(c.startswith("new form version: Amendment 1") for c in st["changed"])
    assert D.update(rec, cost, form_version=fv2["id"], who="bchen") and not D.state(D.get(rec, cost), rows, lib, (), rec)["form"]["newer"]
    # status, owner, stage marks, dismiss, remove
    assert D.update(rec, v, status="Drafting", owner="R. Diaz", who="pm") and not D.update(rec, v, status="Done")
    assert D.mark_stage(rec, v, who="pm") and D.state(D.get(rec, v), rows, lib, (), rec)["done_here"]
    extra = D.add(rec, "Cover letter", who="pm")
    assert D.remove(rec, extra) and D.get(rec, extra) is None
    sm = D.summary(rec, reg.rows(), lib)
    assert sm["count"] == 3 and sm["over"] == 1 and sm["noowner"] == 0, sm
    texts = " | ".join(e["text"] for e in rec.events())
    assert "Deliverable added: Attachment J-2 (owner A. Lee) from the customer documents" in texts and "Volume I Technical: Drafting" in texts
    assert "Deliverable Volume I Technical assigned to R. Diaz" in texts and "Volume I Technical: done for Intake" in texts
    assert "Volume IV Cost: now filled on the newest form version" in texts
    # a new suggestion can be dismissed
    reg2rows = reg.rows() + [{"kind": "Deliverable", "subject": "SF 33", "key": "Deliverable | SF 33", "value": "SF 33", "hidden": False,
                              "doc": "RFP", "version": "Original"}]
    assert [r["subject"] for r in D.suggestions(rec, reg2rows)] == ["SF 33"]
    D.dismiss(rec, "Deliverable | SF 33", "SF 33", who="pm")
    assert not D.suggestions(rec, reg2rows) and "Not ours: SF 33" in rec.events()[0]["text"]
    # the checklist, and a reload keeps it all
    out = D.export_excel(os.path.join(d, "deliv.xlsx"), rec, reg.rows(), lib, (), "Kestrel")
    from openpyxl import load_workbook
    ws = load_workbook(out)["Deliverables"]
    head = [c.value for c in ws[1]]
    assert ws.max_row == 4 and head[0] == "Deliverable" and "Done: Intake" in head and "Progress" in head and "Font" in head, head
    # 2.48.0: submission requirements per deliverable: suggested from the customer documents, confirmed or set by the team
    vol = next(x for x in D.items(rec) if x["kind"] == "Volume")
    rows2 = reg.rows()
    sp = {x["field"]: x for x in D.spec(vol, rows2)}
    assert sp["Font"]["state"] == "Suggested" and sp["Font"]["value"] and sp["Page limit"]["value"], sp
    n = D.spec_confirm_all(rec, vol["id"], rows2, who="pm")
    sp = {x["field"]: x for x in D.spec(D.get(rec, vol["id"]), rows2)}
    assert n >= 2 and sp["Font"]["state"] == "Confirmed" and "pm" in sp["Font"]["confirmed"], sp["Font"]
    D.spec_set(rec, vol["id"], "Copies", "3 printed copies", who="pm")
    sp = {x["field"]: x for x in D.spec(D.get(rec, vol["id"]), rows2)}
    assert sp["Copies"]["state"] == "Set by the team" and sp["Copies"]["value"] == "3 printed copies"
    changed = [dict(r, value="11-point Arial") if r["kind"] == "Format" and r["subject"] == "Font" else r for r in rows2]
    assert {x["field"]: x for x in D.spec(D.get(rec, vol["id"]), changed)}["Font"]["state"] == "Changed: review"
    D.update(rec, vol["id"], progress="45%")
    assert D.get(rec, vol["id"])["progress"] == 45 and D.spec_counts(D.get(rec, vol["id"]), changed)["changed"] == 1
    rec2 = ER.Record(d)
    assert len(D.items(rec2)) == 3 and rec2.data["deliv_dismissed"] == ["Deliverable | SF 33"]
    print("deliverables ok")


def get_kind(rec, did):
    return D.get(rec, did)["kind"]


if __name__ == "__main__":
    run()
