"""3.0.4 graphic needs (graphics.py): add from a sentence or a note, change, mark done, remove, the order (open first),
where each was found, the timeline, the workbook, and the readable copy's Graphic needs.csv."""
import os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import effortrec as RR, graphics as GX
from openpyxl import load_workbook

checks = []


def check(name, ok, info=""):
    checks.append(bool(ok)); print("PASS" if ok else "FAIL", name, info)


tmp = tempfile.mkdtemp()
rec = RR.Record(tmp)
check("nothing to say what it shows: nothing added", GX.add(rec, "   ") == "")
a = GX.add(rec, "Thermal margin chart", kind="Chart", section="1.1 Thermal Design", owner="J. Garcia",
           source={"doc": "Section L.pdf", "loc": "p.3, para 2", "text": "thermal margin", "sid": "L.4-02"}, who="jsmith")
b = GX.add(rec, "Program office org chart", kind="Org chart", note_id="n1", who="jsmith")
rec = RR.Record(tmp)
g = GX.get(rec, a)
check("kept in the effort record", g["concept"] == "Thermal margin chart" and g["status"] == "Needed" and g["sid"] == "L.4-02"
      and len(GX.items(rec)) == 2)
check("where it was found, in words", GX.where_text(g) == "Section L.pdf p.3, para 2 (L.4-02)" and GX.where_text(GX.get(rec, b)) == "From a note")
check("found by sentence and by note", [x["id"] for x in GX.for_sentence(rec, "L.4-02")] == [a] and [x["id"] for x in GX.for_note(rec, "n1")] == [b])
check("in the timeline", any(e["kind"] == "graphic" and "Thermal margin chart" in e["text"] for e in rec.events()))
check("a status that is not one is refused", not GX.update(rec, a, status="Lost"))
check("in work, then done (dated)", GX.update(rec, a, status="In work") and GX.update(rec, a, status="Done") and GX.get(rec, a)["done"])
check("open ones first", [x["id"] for x in GX.items(rec)] == [b, a] and GX.open_count(rec) == 1)
check("changing what it shows; an empty one is refused", GX.update(rec, b, concept="Org chart, program office")
      and not GX.update(rec, b, concept=" ") and GX.get(rec, b)["concept"] == "Org chart, program office")
p = GX.export_excel(os.path.join(tmp, "g.xlsx"), rec, "Kestrel")
rows = list(load_workbook(p).active.iter_rows(values_only=True))
check("the workbook: a row each, open first", len(rows) == 3 and rows[1][1] == "Org chart, program office" and rows[2][5] == "Done"
      and rows[2][7] == "thermal margin", rows)
import readable as RD
out = RD.more(rec, None)
tbl = next((v for k, v in (out.items() if isinstance(out, dict) else []) if k == "Graphic needs.csv"), None)
check("the readable copy has Graphic needs.csv", tbl is not None and len(tbl[1]) == 2 and tbl[0][0] == "Graphic needed", type(out))
check("removing one", GX.remove(rec, b) and GX.get(RR.Record(tmp), b) is None and not GX.remove(rec, b))
n = sum(checks)
print(f"GRAPHICS {'PASS' if n == len(checks) else 'FAIL'} ({n}/{len(checks)})")
sys.exit(0 if n == len(checks) else 1)
