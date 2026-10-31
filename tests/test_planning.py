"""3.1.0 planning, without screens: the schedule from the due date (backward, fitted, weekends and holidays, the team's
defaults, applied to the stages), what changed since a person last opened an effort, the page budget (limits from the
customer's rows, splits, totals that count a section once), and My work across efforts (soonest first, late first)."""
import datetime, os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import effortsched as S, sincelast as SL, pagebudget as PB, mywork as MW
import effortrec as R, compliance as C, setupdoc as SD, effort as E, workspace as W, graphics as GX

checks = []


def check(name, ok, info=""):
    checks.append(bool(ok)); print("PASS" if ok else "FAIL", name, info)


D = datetime.date
tmp = tempfile.mkdtemp()
# ------------------------------------------------ schedule
cal = S.Calendar()
d = S.load_defaults(os.path.join(tmp, "cfg"))
check("built-in working days by stage name, any case; others take 3", S.days_for("author DRAFTS", d) == 10 and S.days_for("Orals prep", d) == 3)
plan = S.plan_back([("a", 3), ("b", 2), ("c", 1)], "2026-12-15", cal)
check("back from the due date: the last stage ends on it, each ends the working day before the next starts",
      plan == [("a", D(2026, 12, 8), D(2026, 12, 10)), ("b", D(2026, 12, 11), D(2026, 12, 14)), ("c", D(2026, 12, 15), D(2026, 12, 15))], plan)
plan = S.plan_back([("a", 2)], "2026-12-14", cal, last_ends_on_due=False)
check("... or the working day before it (a Monday due date: Friday)", plan == [("a", D(2026, 12, 10), D(2026, 12, 11))], plan)
cal_h = S.Calendar(holidays=["2026-11-26", "2026-11-27"])
plan = S.plan_back([("a", 2), ("b", 1)], "2026-11-30", cal_h)
check("holidays and weekends are skipped", plan == [("a", D(2026, 11, 24), D(2026, 11, 25)), ("b", D(2026, 11, 30), D(2026, 11, 30))], plan)
plan = S.plan_back([("a", 2), ("b", 1)], "2026-11-29", S.Calendar(weekends=True))
check("weekends can count", plan[0][1] == D(2026, 11, 27) and plan[-1][2] == D(2026, 11, 29), plan)
fp, avail = S.fit([("a", 10), ("b", 10)], "2026-12-01", "2026-12-14", cal)
check("fit between a start and the due date: every day used, the proportions kept", avail == 10 and fp[0][1] == D(2026, 12, 1)
      and fp[-1][2] == D(2026, 12, 14) and cal.count(fp[0][1], fp[0][2]) == 5, fp)
check("... not enough days for one each: nothing", S.fit([("a", 1), ("b", 1), ("c", 1)], "2026-12-14", "2026-12-15", cal)[0] == [])
S.save_defaults(os.path.join(tmp, "cfg"), {"Red team": 4}, True, ["2026-12-25", "bad"])
d2 = S.load_defaults(os.path.join(tmp, "cfg"))
check("the team's defaults: days, weekends, holidays", S.days_for("Red team", d2) == 4 and d2["weekends"] and d2["holidays"] == ["2026-12-25"])
eroot = os.path.join(tmp, "eff"); os.makedirs(eroot)
rec = R.Record(eroot); rec.save()
plan = S.plan_back([(s["id"], S.days_for(s["name"], d)) for s in rec.stages], "2026-12-15", cal)
n = S.apply(rec, plan, who="jsmith")
rec = R.Record(eroot)
check("applied: every stage gets its planned start and end; names and goals stay", n == len(rec.stages) == 7
      and rec.stages[-1]["planned_end"] == "2026-12-15" and rec.stages[0]["planned"] == plan[0][1].isoformat()
      and [s["name"] for s in rec.stages][0] == "Intake", [(s["name"], s["planned"], s["planned_end"]) for s in rec.stages])
check("... and logged", any("Schedule set from the due date" in e["text"] for e in rec.events()))
# ------------------------------------------------ since last opened
check("first visit: nothing to compare", SL.on_open(eroot, rec, None, user="jsmith", now="2026-10-30T09:00") is None)
rec.log("task", "Task added: Draft 1.1", who="alee", at="2026-10-30T10:00")
rec.log("task", "Task added: my own", who="jsmith", at="2026-10-30T10:05")
rec.log("note", "Old one", who="alee", at="2026-10-29T10:00")
rec.save()
M = C.Matrix(os.path.join(tmp, "m.xlsx"))
M.history = [["2026-10-30 11:00", "kpark", "L.1-01", "Keep"], ["2026-10-30 11:01", "kpark", "L.1-02", "Keep"],
             ["2026-10-30 08:00", "kpark", "L.1-03", "Keep"], ["2026-10-30 11:02", "jsmith", "L.1-04", "Keep"]]
c = SL.on_open(eroot, R.Record(eroot), M, user="jsmith", now="2026-10-30T12:00")
check("since the last visit: others' changes only, newest first, with decisions by person", c and c["count"] == 3
      and [e["who"] for e in c["events"]] == ["alee"] and c["decisions"] == {"kpark": 2} and c["people"] == ["alee", "kpark"], c)
check("... in one line", SL.summary(c) == "3 changes by alee and kpark since Oct 30, 9:00 a.m.", SL.summary(c))
check("... and this visit is noted", SL.last_opened(eroot, "jsmith") == "2026-10-30T12:00")
c2 = SL.on_open(eroot, R.Record(eroot), M, user="jsmith", now="2026-10-30T13:00")
check("nothing new since: no changes", c2 is not None and c2["count"] == 0 and SL.summary(c2) == "")
check("each person has their own last visit", SL.on_open(eroot, R.Record(eroot), M, user="alee") is None)
# ------------------------------------------------ page budget
M2 = C.Matrix(os.path.join(tmp, "m2.xlsx"))
M2.outline = [{"num": "Volume I", "title": "Technical", "level": 1}, {"num": "1.1", "title": "Approach", "level": 2},
              {"num": "1.1.1", "title": "Design", "level": 3}, {"num": "1.2", "title": "Risk", "level": 2},
              {"num": "Volume II", "title": "Management", "level": 1}, {"num": "2.1", "title": "Schedule", "level": 2}]
labs = SD.labels(M2)
M2.items = [{"id": f"r{k}", "text": "x", "decision": "Keep", "map": labs[1 if k < 3 else 3], "type": "Sentence"} for k in range(4)]
vols = PB.volumes(M2)
check("volumes from the structure's Volume headings", [v for v, _l in vols] == [labs[0], labs[4]] and vols[0][1] == labs[:4], vols)
rows = [{"kind": "Page limit", "subject": "Volume I", "value": "20 pages", "doc": "Section L", "version": "v1"}]
check("the limit from the customer's page limit for the volume", PB.limit_for(M2, labs[0], rows) == (20, "Section L, v1"))
PB.set_limit(M2, labs[0], 25)
check("... a typed limit wins", PB.limit_for(M2, labs[0], rows) == (25, "set here"))
check("the volume's limit splits across the sections under its heading", PB.parts(M2, vols[0][1]) == [labs[1], labs[3]])
got = PB.split(M2, [labs[1], labs[3]], 25, by="reqs")
check("split by requirements: 3 and 1 of 4 requirements, in half pages, every page used", got == {labs[1]: 19.0, labs[3]: 6.0}, got)
got = PB.split(M2, [labs[1], labs[3]], 25)
check("split evenly", got == {labs[1]: 12.5, labs[3]: 12.5}, got)
PB.set_pages(M2, labs[2], 4)
check("a section counts once: its subsection's pages are inside it", PB.total(M2, vols[0][1]) == 25 and PB.children_total(M2, labs[1]) == 4)
PB.set_pages(M2, labs[3], 0)
check("0 takes a budget off", PB.pages(M2, labs[3]) is None and PB.total(M2, vols[0][1]) == 12.5)
p = os.path.join(tmp, "m2s.xlsx"); M2.path = p
try:
    M2.save()
    M3 = C.Matrix(p).load()
    check("saved with the matrix", PB.pages(M3, labs[1]) == 12.5 and PB.limit_for(M3, labs[0]) == (25, "set here"))
except Exception as e:
    check("saved with the matrix", False, repr(e))
# ------------------------------------------------ my work
cfg = os.path.join(tmp, "cfg2"); os.makedirs(cfg)
regs = E.Registry(os.path.join(cfg, "efforts.json"))
for nm, due in (("Kestrel", "2026-12-15T14:00"), ("Osprey", "2026-11-20T12:00"), ("Old", "")):
    ed = os.path.join(tmp, nm); os.makedirs(ed)
    lp = os.path.join(ed, f"{nm}_Effort_List.xlsx")
    regs.put(nm, path=lp, archived=(nm == "Old"))
    root = W.data_dir_for(lp); os.makedirs(root)
    r = R.Record(root); r.data["due"] = due
    r.add_task(f"Draft {nm} 1.1", owner="A. Lee, K. Park", due="2026-11-01" if nm == "Kestrel" else "2026-11-10", who="x")
    r.add_task("Someone else's", owner="K. Park", due="2026-11-02", who="x")
    nid = r.add_note("What page limit?", question=True, who="x"); r.update_note(nid, assignee="A. Lee")
    GX.add(r, f"{nm} chart", owner="A. Lee")
    r.save()
work = MW.collect(regs, "A. Lee", load_matrix=False)
check("My work: every effort with work for the person, archived ones left out", sorted(w["effort"] for w in work) == ["Kestrel", "Osprey"])
rows_ = MW.soonest(work, today=datetime.datetime(2026, 11, 5))
check("soonest first: late first, then by date, then by the effort's due date", [(r_["kind"], r_["effort"], r_["late"]) for r_ in rows_][:2]
      == [("task", "Kestrel", True), ("task", "Osprey", False)] and rows_[2]["effort"] == "Osprey", [(r_["kind"], r_["effort"]) for r_ in rows_])
check("... a name among several owners counts; others' tasks do not", all("Someone" not in r_["text"] for r_ in rows_)
      and sum(1 for r_ in rows_ if r_["kind"] == "task") == 2)
check("... questions for the person and graphic needs they own", sum(1 for r_ in rows_ if r_["kind"] == "question") == 2
      and sum(1 for r_ in rows_ if r_["kind"] == "graphic") == 2)
x = MW.export_excel(os.path.join(tmp, "mw.xlsx"), "A. Lee", work)
from openpyxl import load_workbook
wb = load_workbook(x)
check("export: the list and the sections", wb.sheetnames == ["My work", "Sections"] and wb["My work"].max_row == len(rows_) + 1)
t = sum(checks)
print(f"PLANNING {'PASS' if t == len(checks) else 'FAIL'} ({t}/{len(checks)})")
sys.exit(0 if t == len(checks) else 1)
