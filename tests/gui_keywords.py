"""2.16.0 Keyword Harvest: keywords live in the effort list (Keywords sheet) and survive a reopen; Ctrl+K and the
right-click capture fill in the text; the same keyword twice is one; Keep, Reject, edit and remove; suggestions come
from the customer documents and our volumes with reasons, and Not a keyword keeps a phrase from coming back; export
writes the SharePoint columns."""
import os, sys, tempfile, csv, shutil
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ["AH_OCR_ENGINE"] = "tesseract"
os.environ["AH_TEST"] = "1"
import app_gui as A, effort as E, effort_home as H
from tkinter import messagebox
from openpyxl import load_workbook
tmp = tempfile.mkdtemp()
A.CONFIG_DIR = os.path.join(tmp, "config"); A.CHOICES_FILE = os.path.join(tmp, "config", "choices.json")
A.SETTINGS_FILE = os.path.join(tmp, "settings.json")
messagebox.showinfo = lambda *a, **k: None
messagebox.showerror = lambda *a, **k: print("ERROR BOX:", a[:2])
messagebox.askyesno = lambda *a, **k: True
LIST = os.path.join(tmp, "Kestrel", "Kestrel_Effort_List.xlsx")
CDIR = os.path.join(tmp, "Kestrel", "Customer documents"); os.makedirs(CDIR)
RFP = os.path.join(A.BASE_DIR, "test_samples", "rfp")
for f in os.listdir(RFP):
    shutil.copy2(os.path.join(RFP, f), CDIR)
VOL = os.path.join(A.BASE_DIR, "test_samples", "Kestrel_Technical_Volume.docx")
checks = {}


def check(name, ok, info=""):
    checks[name] = bool(ok)
    print(("PASS " if ok else "FAIL ") + name, info)


app = A.AcronymAppGUI()


def step1():
    L = E.EffortList(LIST, "Kestrel"); L.pending = True; L.save()
    app.efforts.put("Kestrel", path=LIST, customer_dir=CDIR)
    assert app.open_effort("Kestrel")
    app.show_view(H.VIEW_KW); app.update()
    check("Keywords view opens", app.view_var.get() == H.VIEW_KW and app.kw_view.winfo_manager())
    check("add a keyword", app.kw_add("thermal control system", "Our solution", "Vol 1") == "added")
    check("the same words again are one keyword", app.kw_add("Thermal  Control System") == "exists")
    app.clipboard_clear(); app.clipboard_append("radiator panel sizing")
    w = app.kw_quick(); w.update()
    check("Ctrl+K fills in the clipboard text", w._api["entry"].get() == "radiator panel sizing", w._api["entry"].get())
    w._api["cat"].set("Differentiator"); w._api["ok"]()
    k = app.effort.keyword("radiator panel sizing")
    check("dialog adds with category", k and k["Category"] == "Differentiator" and k["Added by"], k)
    app.kw_set_status("Keep", ["radiator panel sizing"])
    app.kw_set_status("Reject", ["thermal control system"])
    app.kw_filter.set("Keep"); app.kw_refresh()
    shown = [app.kw_tree.set(i, "kw") for i in app.kw_tree.get_children()]
    check("filter: keep", shown == ["radiator panel sizing"], shown)
    app.kw_filter.set(app.kw_filter.get()); app.kw_filter.set("All keywords"); app.kw_refresh()
    check("saved in the Keywords sheet", "radiator panel sizing" in [x["Keyword"] for x in E.EffortList(LIST).load().keywords])
    # suggestions
    app.last_scan = {"paths": [VOL], "folder": None, "cache": {}}
    srcs = app.kw_sources()
    check("sources: customer documents and our volume", {k for _p, k in srcs} == {"customer", "ours"} and len(srcs) >= 3, srcs)
    sw = app.kw_suggest_dialog(wait=True); sw.update()
    items = sw._api["st"]["items"]
    check("suggestions with reasons", len(items) >= 5 and all(it["why"] for it in items), [i["phrase"] for i in items[:6]])
    check("a heading phrase from the RFP is suggested", any(i["phrase"].lower() == "thermal design" for i in items), [i["phrase"] for i in items[:8]])
    sides = {i.get("side") for i in items}
    tied = [i for i in items if i.get("answers")]
    check("2.34.0 suggestions come from both sides", {"customer", "both", "ours"} <= sides, sides)
    check("2.34.0 our wording is tied to the customer ask it answers", tied and all(i["side"] == "ours" for i in tied)
          and all(i["bucket"] in ("Our response", "Differentiator") for i in tied), [(i["phrase"], i["answers"]) for i in tied[:4]])
    sw._api["side"]("In both"); sw.update()
    shown = sw._api["st"]["items"]
    check("2.34.0 Show In both lists only shared phrases", shown and all(i["side"] == "both" for i in shown), len(shown))
    sw._api["side"]("All"); sw.update(); items = sw._api["st"]["items"]
    tv = sw._api["tree"]
    first = items[0]["phrase"]
    tv.selection_set(tv.get_children()[:1]); srcw = sw._api["source"]()
    check("Show in document highlights the phrase", srcw is not None and srcw._hits >= 1, getattr(srcw, "_hits", None))
    if srcw is not None:              # let the window finish opening first: closing it at once crashed Tk on Windows
        import time
        for _ in range(6):
            app.update(); time.sleep(0.1)
        srcw.destroy(); app.update()
    tv.selection_set(tv.get_children()[:2]); got = sw._api["add"]()
    check("add selected suggestions", got == 2 and app.effort.keyword(first)["Source"] == "Suggested")
    tv.selection_set(tv.get_children()[:1]); rejected = sw._api["st"]["items"][int(tv.selection()[0])]["phrase"]
    sw._api["reject"]()
    sw.destroy()
    again = [i["phrase"] for i in app.kw_compute(srcs)]
    check("added and rejected phrases are not suggested again", first not in again and rejected not in again)
    # offerings catalog (2.17.0)
    import offerings
    ow = app.offerings_window(); ow.update(); api = ow._api
    api["ents"]["Offering"].insert(0, "Orbital Thermal Suite"); api["ents"]["Other names"].insert(0, "Northwind Orbital")
    api["ents"]["Group"].insert(0, "Thermal"); api["add"]()
    lst = os.path.join(tmp, "offer.csv")
    with open(lst, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows([["Name", "Description"], ["Dream Chaser", "Spaceplane"], ["LIFE Habitat", "Inflatable habitat"],
                                 ["Orbital Thermal Suite", "dup"]])
    a, same = api["import"](lst)
    check("catalog: add and import a list", a == 2 and same == 1 and len(offerings.load()) == 3, (a, same))
    exp = api["export"](os.path.join(tmp, "Offerings_Catalog.csv"))
    check("catalog: export to share keeps every column", open(exp, encoding="utf-8-sig").readline().startswith("Offering,Other names,Group"))
    ow.destroy()
    items = app.kw_compute(srcs)
    off = [i for i in items if i["bucket"] == "Offering"]
    check("suggestions mark an offering named in the documents", off and off[0]["phrase"] == "Orbital Thermal Suite"
          and "offerings catalog" in off[0]["why"][0] and not any(i["phrase"] == "Northwind Orbital" for i in items),
          [(i["phrase"], i["why"]) for i in off])
    aw = app.attach_offerings_dialog(); aw.update()
    tv2 = aw._api["tree"]
    pick = [i for i in tv2.get_children() if tv2.set(i, "name") in ("Dream Chaser", "LIFE Habitat")]
    tv2.selection_set(pick); n = aw._api["attach"](); aw.destroy()
    k = app.effort.keyword("Dream Chaser")
    check("attach offerings to the effort", n == 2 and k and k["Category"] == "Offering" and k["Source"] == "Offerings catalog", k)
    app.kw_add("ThermaLoop radiator", "Our solution")
    check("a keyword goes into the catalog", app.kw_to_catalog(["ThermaLoop radiator"]) == 1
          and offerings.find(offerings.load(), "ThermaLoop radiator") and app.effort.keyword("ThermaLoop radiator")["Category"] == "Offering")
    # past proposals and team history (2.18.0)
    import keywords, tkinter as tk
    past = os.path.join(tmp, "past"); os.makedirs(past)
    for f in ("Kestrel_Technical_Volume.docx", "Kestrel_Design_Note.docx", "Kestrel_Program_Manual.pdf", "Kestrel_PDR_Briefing.pptx"):
        shutil.copy2(os.path.join(A.BASE_DIR, "test_samples", f), past)
    data = app.kw_build_baseline(past)
    raw = open(os.path.join(A.CONFIG_DIR, "keyword_baseline.json"), encoding="utf-8").read()
    check("baseline from past proposals: counts only, no text", data["docs"] == 4 and "thermal" not in raw.lower() and len(data["df"]) > 50)
    items = app.kw_compute(srcs)
    check("suggestions compare with past proposals", any("past proposals" in w for i in items for w in i["why"]),
          [i["why"] for i in items[:3]])
    lr = open(os.path.join(A.CONFIG_DIR, "keyword_learning.json"), encoding="utf-8").read()
    check("Keep and Reject are learned as hashed counts", keywords.learning()["keep"].get(keywords.phrase_hash("radiator panel sizing")) == 1
          and "radiator" not in lr)
    t = tk.Text(app); t.insert("1.0", "The cryogenic fluid management system is new."); t.tag_add("sel", "1.4", "1.37")
    shown = {}
    real_popup = tk.Menu.tk_popup
    tk.Menu.tk_popup = lambda self, *a: shown.setdefault("label", self.entrycget(self.index("end"), "label"))   # 2.50.0: last entry
    class Ev: x_root = y_root = 10
    app.kw_text_menu(Ev, t, "Compliance requirement")
    tk.Menu.tk_popup = real_popup
    check("right-click selected text: Add as keyword", shown.get("label", "").startswith("Add as keyword: cryogenic fluid management"), shown)
    # export
    out = os.path.join(tmp, "kw.xlsx")
    n = app.kw_export(out, keep_only=True)
    ws = load_workbook(out)["Keywords"]
    check("export keep only, SharePoint columns", n == 1 and [c.value for c in ws[1]][:3] == ["Keyword", "Category", "SharePoint field"])
    out2 = os.path.join(tmp, "kw.csv")
    n2 = app.kw_export(out2, keep_only=False)
    rows = list(csv.reader(open(out2, encoding="utf-8-sig")))
    check("export keep and proposed, never reject", n2 == len(rows) - 1 and all(r[5] != "Reject" for r in rows[1:]) and n2 >= 3, n2)
    app.kw_remove(["radiator panel sizing"])
    check("remove", app.effort.keyword("radiator panel sizing") is None)
    check("history records keyword changes", any(h[2] == "(keyword)" for h in app.effort.history))
    app.flush_session(); app.save_resume_now()
    app.on_close()
    app2 = A.AcronymAppGUI()

    def after():
        app2.update()
        check("reopen: back on the Keywords view with the list", app2.view_var.get() == H.VIEW_KW and len(app2.kw_tree.get_children()) == len(app2.effort.keywords) >= 3,
              (app2.view_var.get(), len(app2.kw_tree.get_children())))
        ok = all(checks.values())
        print("GUI KEYWORDS", "PASS" if ok else "FAIL", f"({sum(checks.values())}/{len(checks)})")
        app2.on_close()
    app2.after(1500, after)
    app2.mainloop()


app.after(1500, step1)
app.mainloop()
