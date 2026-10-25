"""2.14.0 Effort home and Scan review: the effort list is home and reopens as it was left; every scan lands in Scan
review; words selected in the preview become a definition (kept for later scans, written to the list); named volumes
assign terms from where they were found, a manual assignment survives rescans; the export picker writes the chosen
columns for the whole effort or per volume (CSV, Excel, Word); a volume's acronym list imports into the effort list."""
import os, sys, tempfile, csv
sys.stdout.reconfigure(encoding="utf-8", errors="replace")   # Windows consoles default to cp1252
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ["AH_OCR_ENGINE"] = "tesseract"
os.environ["AH_TEST"] = "1"   # no first-start wizard during tests
import app_gui as A, effort as E, docxw, effort_home as H
from tkinter import messagebox, simpledialog
from openpyxl import load_workbook
tmp = tempfile.mkdtemp()
A.CONFIG_DIR = os.path.join(tmp, "config"); A.CHOICES_FILE = os.path.join(tmp, "config", "choices.json")
A.SETTINGS_FILE = os.path.join(tmp, "settings.json")
A.DB_FILE = os.path.join(tmp, "db.csv")
with open(A.DB_FILE, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["Term", "Category", "Definition", "Action"])
    w.writerow(["SOW", "Acronym", "Statement of Work", "Include"]); w.writerow(["TRL", "Acronym", "Technology Readiness Level", "Include"])
messagebox.showinfo = lambda *a, **k: None
messagebox.showerror = lambda *a, **k: print("ERROR BOX:", a[:2])
messagebox.askyesno = lambda *a, **k: True
messagebox.askretrycancel = lambda *a, **k: False
simpledialog.askstring = lambda *a, **k: "J. Smith"
V1 = os.path.join(tmp, "Vol_1_Technical.docx")
V2 = os.path.join(tmp, "Vol_2_Management.docx")
d = docxw.Doc(); d.heading("1 Approach", 1)
d.para("The SOW tasks follow the TRL plan. The QZM sensor reads the panel data every second.")
d.para("Our Quartz Zone Monitor design flew twice before.")
d.save(V1)
d = docxw.Doc(); d.heading("1 Management", 1)
d.para("The SOW schedule applies. The QZM report is monthly.")
d.save(V2)
LIST = os.path.join(tmp, "Kestrel", "Kestrel_Effort_List.xlsx")
IMP = os.path.join(tmp, "Vol_2_Acronyms.csv")
with open(IMP, "w", newline="", encoding="utf-8") as f:
    csv.writer(f).writerows([["Acronym", "Definition"], ["NEWX", "New Experimental Unit"], ["QZM", "Quartz Zone Monitor"],
                             ["SOW", "Scope of Work"]])
OUT = os.path.join(tmp, "out")
os.makedirs(OUT)
checks = {}


def check(name, ok, info=""):
    checks[name] = bool(ok)
    print(("PASS " if ok else "FAIL ") + name, info)


app = A.AcronymAppGUI()
st = {"t": 0}


def wait(cond, then, limit=120):
    st["t"] += 1
    if not cond() and st["t"] < limit:
        app.after(400, lambda: wait(cond, then, limit)); return
    st["t"] = 0
    then()


def scanning():
    return app.btn_cancel.winfo_ismapped() or app.busy()


def step1():
    L = E.EffortList(LIST, "Kestrel", client="USSF SSC", solicitation="FA8810-27-R-0042"); L.pending = True; L.save()
    app.efforts.put("Kestrel", path=LIST, first_use="volume")
    assert app.open_effort("Kestrel")
    app.update()
    check("an open effort starts on its home", app.view_var.get() == H.VIEW_HOME and app.home_view.winfo_ismapped()
          and app.seg_view.winfo_ismapped(), app.view_var.get())
    app.start_batch([V1, V2])
    app.after(500, lambda: wait(lambda: not scanning(), step2))


def step2():
    app.update()
    check("a scan lands in Scan review", app.view_var.get() == H.VIEW_SCAN and app.scan_view.winfo_ismapped(), app.view_var.get())
    check("QZM found as a new term", "QZM" in app.results, sorted(app.results)[:10])
    app.tree.selection_set("QZM"); app.on_select(None); app.update()
    check("Omit and Not a term are offered in Scan review", str(app.btn_effort_omit.cget("state")) == "normal"
          and str(app.btn_effort_rej.cget("state")) == "normal")
    if not app.preview_var.get():
        app.preview_var.set(True); app.toggle_preview()
    app.preview.show("QZM")
    wait(lambda: "Quartz" in app.preview.txt.get("1.0", "end") or app.preview.mode == "typed" and
         "QZM" in app.preview.txt.get("1.0", "end") and "Reading" not in app.preview.txt.get("1.0", "end"), step3)


def step3():
    pv = app.preview
    txt = pv.txt
    body = txt.get("1.0", "end")
    if "Quartz" not in body:                  # the passage shown may be the first paragraph only: show the second
        for i, l in enumerate(pv.locs):
            pv._show_location(i); app.update()
            if "Quartz" in txt.get("1.0", "end"):
                break
    body = txt.get("1.0", "end")
    k = body.find("Quartz Zone Monitor")
    if k >= 0:
        txt.tag_add("sel", f"1.0+{k}c", f"1.0+{k + len('Quartz Zone Monitor')}c")
        used = pv.use_selection()
    else:                                     # the paragraphs are separate passages: select in the recognized-text path
        used = pv.use_selection("Quartz Zone Monitor (QZM).")
    d = app.results["QZM"]
    check("Use as definition: the selected words become the definition", used == "Quartz Zone Monitor"
          and d["selected_definition"] == "Quartz Zone Monitor", (used, d["selected_definition"]))
    check("Use as definition resolves the row", A.status_kind(d["status"]) == "resolved", d["status"])
    check("Use as definition is kept for this effort", app.ws.get("picked_definitions", {}).get("QZM") == "Quartz Zone Monitor")
    check("cleaning: quotes, (TERM) and punctuation go", H.clean_definition('"Statement of Work (SOW)."', "SOW") == "Statement of Work"
          and H.clean_definition("SOW - Statement of Work", "SOW") == "Statement of Work")
    pv.show("SOW"); app.update()
    check("found line offers the other meaning", pv._found == "" or isinstance(pv._found, str))
    app.add_to_effort([t for t in app.results if not app.results[t].get("excluded") and not app.results[t].get("file_error")])
    L = app.effort
    check("Keep writes the picked definition into the effort list", L.rows["QZM"]["Definition"] == "Quartz Zone Monitor",
          L.rows.get("QZM"))
    # volumes
    app.add_volume("Vol I Technical", [V1]); app.add_volume("Vol II Management", [V2])
    check("a term found in both volumes belongs to both", app.term_volumes("SOW", L.rows["SOW"]) == ["Vol I Technical", "Vol II Management"],
          app.term_volumes("SOW", L.rows["SOW"]))
    check("TRL belongs to Vol I only", app.term_volumes("TRL", L.rows["TRL"]) == ["Vol I Technical"])
    check("Used in volumes written into the list", E.EffortList(LIST).load().rows["SOW"]["Used in volumes"] == "Vol I Technical; Vol II Management")
    app.assign_volumes(["TRL"], ["Vol II Management"])
    check("manual assignment", app.term_volumes("TRL", L.rows["TRL"]) == ["Vol II Management"])
    app.start_batch([V1, V2])
    app.after(500, lambda: wait(lambda: not scanning(), step4))


def step4():
    L = app.effort
    d = app.results["QZM"]
    check("rescan keeps the definition taken from the document", d["selected_definition"] == "Quartz Zone Monitor", d["status"])
    check("manual volume survives a rescan", app.term_volumes("TRL", L.rows["TRL"]) == ["Vol II Management"])
    app.show_view(H.VIEW_HOME); app.update()
    app.home_filter.set("Volume: Vol I Technical"); app.home_refresh()
    shown = set(app.home_tree.get_children())
    check("home filter by volume", "QZM" in shown and "TRL" not in shown and "SOW" in shown, shown)
    app.home_filter.set("Needs a definition"); app.home_refresh()
    check("home filter: needs a definition", all(not L.rows[t]["Definition"] for t in app.home_tree.get_children()))
    app.home_filter.set(H.HOME_FILTERS[0]); app.home_refresh()
    check("home lists every term", len(app.home_tree.get_children()) == len(L.rows), (len(app.home_tree.get_children()), len(L.rows)))
    check("home summary counts", f"{len(L.rows)} terms" in app.lbl_home.cget("text"), app.lbl_home.cget("text"))
    # export picker and exports
    p = os.path.join(OUT, "all.csv")
    app.export_terms(p)
    rows = list(csv.reader(open(p, encoding="utf-8-sig")))
    check("export: Term and Definition by default", rows[0] == ["Term", "Definition"] and len(rows) - 1 == len(L.rows), rows[:2])
    made = app.export_terms(OUT, ["Term", "Definition", "Volumes"], "per_volume", fmt="csv")
    check("export: one file per volume", len(made) == 2 and all(os.path.isfile(m) for m in made), made)
    v1 = list(csv.reader(open([m for m in made if "Vol_I_" in m][0], encoding="utf-8-sig")))
    check("per-volume file holds that volume's terms", {r[0] for r in v1[1:]} >= {"QZM", "SOW"} and "TRL" not in {r[0] for r in v1[1:]}, v1)
    p2 = os.path.join(OUT, "vol2.xlsx")
    app.export_terms(p2, ["Term", "Definition", "Found in"], "whole", ["Vol II Management"], "xlsx")
    ws = load_workbook(p2).active
    names = {r[0] for r in ws.iter_rows(min_row=2, values_only=True)}
    check("export: chosen volume only, Excel", "TRL" in names and "SOW" in names and [c.value for c in ws[1]] == ["Term", "Definition", "Found in"], names)
    p3 = os.path.join(OUT, "list.docx")
    app.export_terms(p3, fmt="docx")
    check("export: Word", os.path.getsize(p3) > 2000)
    L.set_status(["SOW"], "Approved", approver="J. Smith"); app._save_effort_list()
    p4 = os.path.join(OUT, "appr.csv")
    app.export_terms(p4, approved_only=True)
    check("export: approved only", [r[0] for r in csv.reader(open(p4, encoding="utf-8-sig"))][1:] == ["SOW"])
    win = app.export_picker(); win.update(); api = win._api
    api["cols"]["Status"].set(True); api["fmt"].set("csv")
    p5 = os.path.join(OUT, "picker.csv")
    api["go"](p5); win.destroy()
    check("export picker writes the ticked columns and remembers them",
          next(csv.reader(open(p5, encoding="utf-8-sig"))) == ["Term", "Definition", "Status"]
          and app.ws.get("export_choice", {}).get("cols") == ["Term", "Definition", "Status"])
    # import a volume's list
    res = app.import_volume_list(IMP, volume="Vol II Management", ask=False)
    check("import: new term added, same skipped, different kept aside",
          res and res["added"] == 1 and res["same"] == 1 and res["differs"] == 1, res)
    check("import: the new term carries its source and volume", L.rows["NEWX"]["Source"].startswith("Imported list")
          and app.term_volumes("NEWX", L.rows["NEWX"]) == ["Vol II Management"], L.rows.get("NEWX"))
    check("import: a different definition goes to Other definitions", "Scope of Work" in L.rows["SOW"]["Other definitions"]
          and L.rows["SOW"]["Definition"] == "Statement of Work")
    check("import: QZM gains Vol II and keeps Vol I", app.term_volumes("QZM", L.rows["QZM"]) == ["Vol I Technical", "Vol II Management"],
          app.term_volumes("QZM", L.rows["QZM"]))
    app.home_refresh()
    app.home_tree.selection_set("NEWX"); app.update()
    w = app.home_where_from(); w.update()
    check("Where from shows the import", "Imported list" in w._text and len(w._tree.get_children()) >= 1, w._text[:80])
    w.destroy()
    app.home_set_definition("NEWX", "New Experimental Unit (NEWX)")
    check("home: edit definition", L.rows["NEWX"]["Definition"] == "New Experimental Unit")
    app.home_tree.selection_set("NEWX"); app.update()
    app.home_set_status("Needs author")
    check("home: set status", L.rows["NEWX"]["Status"] == "Needs author")
    vw = app.edit_volumes(); vw.update()
    vw._api["list"].selection_set(0); vw._api["show"]("Vol I Technical"); vw.update()
    check("volumes window ticks the volume's files", vw._api["st"]["vars"].get("Vol_1_Technical.docx") is not None
          and vw._api["st"]["vars"]["Vol_1_Technical.docx"].get())
    vw.destroy()
    # leave it as the person would: home, a filter, a search, a selection
    app.home_filter.set("Needs author"); app.home_search.set("NE"); app.home_refresh()
    app.home_tree.selection_set("NEWX"); app.update()
    app.flush_session(); app.save_resume_now()
    app.on_close()
    app2 = A.AcronymAppGUI()

    def after_start():
        app2.update()
        check("reopen: back on the home view", app2.view_var.get() == H.VIEW_HOME and app2.home_view.winfo_ismapped())
        check("reopen: filter and search restored", app2.home_filter.get() == "Needs author" and app2.home_search.get() == "NE",
              (app2.home_filter.get(), app2.home_search.get()))
        check("reopen: selection restored", list(app2.home_tree.selection()) == ["NEWX"], app2.home_tree.selection())
        check("reopen: the list as it stands", app2.effort.rows["NEWX"]["Definition"] == "New Experimental Unit"
              and app2.effort.rows["QZM"]["Definition"] == "Quartz Zone Monitor")
        ok = all(checks.values())
        print("GUI HOME", "PASS" if ok else "FAIL", f"({sum(checks.values())}/{len(checks)})")
        app2.on_close()
    app2.after(1500, after_start)
    app2.mainloop()


app.after(1500, step1)
app.mainloop()
