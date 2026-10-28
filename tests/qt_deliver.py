"""2.61.0 Deliver in Qt and the L/M crosswalk: the crosswalk links instructions to factors and work by the RFP's
cross-references and titles, flags the gaps, and keeps the person's own links; the Matrix page changes the design
(columns, headings, a column of the team's own, context in the cell or on rows), previews it and exports it; the
compliance package writes the matrix, the setup, the coverage report and the crosswalk into one folder; Response
Setup orders blocks, takes one out and exports the Word documents."""
import os, sys, tempfile, time
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
os.environ["QT_QPA_PLATFORM"] = "offscreen"
os.environ["AH_TEST"] = "1"
os.environ["SK_SANDBOX_TEMP"] = "temp_qt"
os.environ["SK_THEME"] = "light"
import app_gui as A
import compliance as C, effort as E, setupdoc as SD, workspace as W, office_pdf as OP, crosswalk as X, matrixdesign as MD
OP.available = lambda: False
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QMessageBox
QMessageBox.information = staticmethod(lambda *a, **k: None)
from skqt import design as D
from skqt.mainwin import MainWindow
from openpyxl import load_workbook

checks = []


def check(name, ok, info=""):
    checks.append(bool(ok)); print("PASS" if ok else "FAIL", name, info)


tmp = tempfile.mkdtemp()
cfg = os.path.join(tmp, "config"); os.makedirs(cfg)
EDIR = os.path.join(tmp, "Kestrel"); os.makedirs(EDIR)
LIST = os.path.join(EDIR, "Kestrel_Effort_List.xlsx")
L = E.EffortList(LIST, "Kestrel"); L.pending = True; L.save()
E.Registry(os.path.join(cfg, "efforts.json")).put("Kestrel", path=LIST)
W.Workspace(W.data_dir_for(LIST), "Kestrel").set("outline_text", "Volume I Technical\n1.1 Technical Approach\n1.2 Management\nVolume II Past Performance")
app = QApplication([]); D.apply(app)
w = MainWindow(app, cfg, A.APP_VERSION); w.resize(1500, 950); w.show()
w.open_effort("Kestrel")
B = os.path.join(A.BASE_DIR, "test_samples", "rfp")
w.read_page.read([os.path.join(B, "Kestrel_RFP_Sections_C_L_M.docx"), os.path.join(B, "Kestrel_RFP_Section_L_Part2.pdf")],
                 os.path.join(EDIR, "Kestrel_Compliance_Matrix.xlsx"))
t = time.time()
while w.review is None and time.time() - t < 60:
    app.processEvents(); time.sleep(0.02)
M = w.review.M
# the crosswalk
rows = {r["num"]: r for r in X.build(M)}
check("L.4.1 is evaluated by M.2 (same words) and points at C.2 (its cross-reference)",
      [n for n, _t, _h in rows["L.4.1"]["m"]] == ["M.2"] and "C.2" in [n for n, _t, _h in rows["L.4.1"]["c"]]
      and any(h == "cross-reference in the RFP" for _n, _t, h in rows["L.4.1"]["c"]), rows["L.4.1"])
check("Management Approach: no factor evaluates it (a gap)", not rows["L.4.2"]["m"] and rows["L.4.2"]["gaps"])
check("a numbered paragraph under the same heading is not a row of its own", "L.2.1" not in rows and "L.2" in rows)
w.go("crosswalk"); app.processEvents()
cp = w.crosswalk_page
check("the crosswalk page: one row per instruction, gaps counted", cp.table.rowCount() == len(X.build(M)) and "not evaluated" in cp.p_noeval.text())
cp.select("L.4.2"); app.processEvents()
cp.link("L.4.2", "M.1", True); app.processEvents()
r = {x["num"]: x for x in X.build(M)}["L.4.2"]
check("+ factor: the person's own link", [(n, h) for n, _t, h in r["m"]] == [("M.1", "added by you")] and not r["gaps"])
cp.link("L.4.1", "M.2", False)
r = {x["num"]: x for x in X.build(M)}["L.4.1"]
check("x on a found link takes it out and it stays out", not r["m"] and r["gaps"])
w.save_now()
M2 = C.Matrix(M.path).load()
check("the person's links are kept with the matrix", {x["num"]: x for x in X.build(M2)}["L.4.2"]["m"][0][0] == "M.1")
res = cp.export(os.path.join(tmp, "x.xlsx"))
ws = load_workbook(res["path"]).active
check("the crosswalk exports to Excel", ws["A4"].value == "Section L" and any(c.value == "L.4.2" for c in ws["A"]))
# assign a few so the matrix and the setup have content
by = lambda s: next(it for it in M.items if s in it["text"])
R = w.review                                                     # (the Word file has no page copy here: no Office)
a = by("thermal design"); b = by("Integrated Master Schedule")
R.req([a], "1.1 Technical Approach"); R.req([b], "1.2 Management")
ctx = next(it for it in M.items if it["decision"] == "" and it["file"] == a["file"] and it is not a and it is not b)
R.ctx([ctx], [a["id"]]); w.changed()
# the matrix page
w.go("matrix"); app.processEvents()
mp = w.matrix_page
check("the matrix page shows the design and a preview", mp.cols.count() >= len(MD.COLUMNS) and mp.model.rowCount() >= 2
      and "Context in its requirement's cell" in mp.info.text(), mp.info.text())
notes = next(i for i in range(mp.cols.count()) if mp.cols.item(i).data(Qt.UserRole) == "notes")
mp.cols.item(notes).setCheckState(Qt.Unchecked); app.processEvents()
check("unticking a column takes it out of the matrix", "notes" not in [c["key"] for c in MD.of_matrix(M)["columns"]])
vol = next(i for i in range(mp.cols.count()) if mp.cols.item(i).data(Qt.UserRole) == "volume")
mp.cols.setCurrentRow(vol); mp.heading.setText("Vol."); mp.rename(); app.processEvents()
check("renaming a column", MD.of_matrix(M)["columns"][0]["heading"] == "Vol.")
k = mp.add_column("Reviewer")
check("a column of the team's own", k == "custom:Reviewer" and MD.of_matrix(M)["columns"][-1]["key"] == "custom:Reviewer")
mp.context.setCurrentIndex(1); app.processEvents(); mp.preview()
check("context on its own rows", MD.of_matrix(M)["context"] == "rows" and "own rows" in mp.info.text())
mp.context.setCurrentIndex(0); app.processEvents(); mp.preview()
res = mp.export(os.path.join(tmp, "m.xlsx"))
wb = load_workbook(res["path"])["Compliance Matrix"]
heads = [c.value for c in wb[4]]
check("Export the matrix: the designed columns", heads[0] == "Vol." and "Notes" not in heads and heads[-1] == "Reviewer", heads)
check("context in its requirement's cell in Excel", any("Supporting context:" in str(c.value) for row in wb.iter_rows(min_row=5) for c in row))
mp.team_save()
check("save as the team default", MD.load_team(cfg)["columns"][0]["heading"] == "Vol.")
# response setup
w.go("setup"); app.processEvents()
sp = w.setup_page
check("Response Setup: each section with what is placed in it", sp.tree.topLevelItemCount() == len(M.outline)
      and "blocks placed" in sp.p_blocks.text())
sec = next(sp.tree.topLevelItem(i) for i in range(sp.tree.topLevelItemCount()) if sp.tree.topLevelItem(i).data(0, Qt.UserRole)[1] == "1.1 Technical Approach")
check("the requirement and its context under 1.1", sec.childCount() >= 2, sec.childCount())
first = sec.child(0).data(0, Qt.UserRole)
sp.tree.setCurrentItem(sec.child(0)); sp.move(1)
check("Down moves a block later in its section", SD.placed_in(M)["1.1 Technical Approach"][1] == first[2])
paths = sp.export(os.path.join(tmp, "setup"))
check("Export the setup: Word documents", paths and all(p.endswith(".docx") and os.path.isfile(p) for p in paths), paths)
w.go("matrix")
pkg = mp.export_package(os.path.join(tmp, "pkg"))
names = sorted(os.path.basename(p) for p in pkg["files"])
check("the compliance package: matrix, setup, coverage and crosswalk in one folder", any(n.endswith("_Compliance_Matrix_" + __import__("datetime").date.today().isoformat() + ".xlsx") for n in names)
      and any("RFP_Coverage" in n for n in names) and any("LM_Crosswalk" in n for n in names) and any(n.endswith(".docx") and "Volume" in n for n in names), names)
w.close()
n = sum(checks)
print(f"QT DELIVER PASS ({n}/{len(checks)})" if n == len(checks) else f"QT DELIVER FAIL ({n}/{len(checks)})")
sys.exit(0 if n == len(checks) else 1)
