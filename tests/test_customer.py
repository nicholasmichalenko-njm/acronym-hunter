"""2.48.0 Customer's sentences: a sentence that reads like a requirement but is the customer's to act on ("The
Government will evaluate...") is identified and counted, never answered: no section, no links (context that pointed at
it loses the link), counted in the coverage (sheet and report), kept through save and load, named in the readable copy,
and left out of the setup documents."""
import os, sys, tempfile, zipfile, csv
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE)); sys.path.insert(0, HERE)
import compliance as C, setupdoc as SD, readable as RD
from test_assign import row


def run():
    a = row(1, "L.4", "The Offeror shall describe the thermal design.")
    b = row(2, "L.4", "The Government will evaluate the thermal design for risk.")
    c = row(3, "L.4", "Thermal margins matter because the orbit has long eclipses.")
    items = [a, b, c]
    C.assign_requirement(items, [a], "1.1 Thermal")
    C.assign_requirement(items, [b], "1.1 Thermal")
    C.assign_context(items, [c], ["L.4-01", "L.4-02"])
    C.assign_customer(items, [b])
    assert b["decision"] == "Customer" and b["map"] == "" and b["context_for"] == ""
    assert c["context_for"] == "L.4-01", "context that pointed at it loses the link"
    d = tempfile.mkdtemp()
    M = C.Matrix(os.path.join(d, "Kestrel_Matrix.xlsx"), "Kestrel"); M.items = items
    M.outline = C.parse_outline("Volume I\n1.1 Thermal")
    cov = M.coverage()
    assert cov[-1]["Customer"] == 1 and cov[-1][""] == 0 and M.counts()["Customer"] == 1
    assert SD.placed_in(M)["1.1 Thermal"] == ["L.4-01", "L.4-03"], "never in our response"
    M.save()
    N = C.Matrix(M.path).load()
    assert [x["decision"] for x in N.items] == ["Keep", "Customer", "Context"]
    from openpyxl import load_workbook
    wb = load_workbook(M.path, read_only=True)
    head = [c.value for c in next(wb["Coverage"].iter_rows(max_row=1))]
    assert "Customer" in head, head
    out = os.path.join(d, "cov.docx"); C.write_coverage_docx(M, out)
    with zipfile.ZipFile(out) as z:
        x = z.read("word/document.xml").decode("utf-8")
    assert "1 the customer's" in x and "Customer's (not ours to answer)" in x
    parts = RD.build(None, None, None, [], M, {"effort": "Kestrel"}, "2.48.0")
    folder = os.path.join(d, "readable")
    RD.write(folder, parts, None)
    with open(os.path.join(folder, "Compliance assignments.csv"), encoding="utf-8-sig", newline="") as f:
        roles = {r["ID"]: r["Role"] for r in csv.DictReader(f)}
    assert roles["L.4-02"] == "Customer's (not ours to answer)", roles
    print("customer ok")


if __name__ == "__main__":
    run()
