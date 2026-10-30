"""2.96.0 every kind of customer file, offscreen: a picture (and a scanned PDF) is read into the review from a copy with
its words laid over the page, and Assign boxes its sentences on that copy like a typed PDF; an Excel workbook opens in
Sheet view (its cells as a grid, each sentence boxed in its cell, clickable and splittable), with Page view a click
away; pictures and old Office files are accepted as customer documents; the copy is made once per file contents."""
import os, sys, tempfile, time
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE)); sys.path.insert(0, HERE)
os.environ["QT_QPA_PLATFORM"] = "offscreen"
os.environ["AH_TEST"] = "1"
os.environ["SK_SANDBOX_TEMP"] = "temp_qt"
import app_gui as A
import effort as E, workspace as W, docintake as DI, ocrpdf as OP2, office_pdf as OP, sheetview as SV, watch as WA
OP.available = lambda: False
from PIL import Image, ImageDraw, ImageFont
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QMessageBox
from PySide6.QtTest import QTest
QMessageBox.information = staticmethod(lambda *a, **k: None)
from skqt import theme
from skqt.mainwin import MainWindow

checks = []


def check(name, ok, info=""):
    checks.append(bool(ok)); print("PASS" if ok else "FAIL", name, info)


tmp = tempfile.mkdtemp()
cfg = os.path.join(tmp, "config"); os.makedirs(cfg)
EDIR = os.path.join(tmp, "Heron"); CDIR = os.path.join(EDIR, "Customer Documents"); os.makedirs(CDIR)
# a picture of a page, and the reader's lines for it (this PC's OCR engine is not in the test machine)
TEXT = [(150, 300, "L.5 MANAGEMENT VOLUME"), (150, 400, "The Offeror shall provide a staffing plan for the program."),
        (150, 470, "The Offeror shall describe its risk management process."), (150, 540, "Each volume shall be bound separately.")]
img = Image.new("RGB", (1700, 2200), "white"); dr = ImageDraw.Draw(img)
try:
    font = ImageFont.truetype("DejaVuSans.ttf", 40)
except OSError:
    font = ImageFont.load_default()
LINES = []
for x, y, t in TEXT:
    dr.text((x, y), t, fill="black", font=font)
    l, tp, r, b = dr.textbbox((x, y), t, font=font)
    LINES.append((l, tp, r, b, t))
PIC = os.path.join(CDIR, "Section L page 5 (scan).png"); img.save(PIC, dpi=(200, 200))
calls = []
OP2.ocr_lines = lambda im: (calls.append(1), [tuple(v * (im.width / 1700) if i < 4 else v for i, v in enumerate(ln)) for ln in LINES])[1]
# a workbook
import openpyxl
XL = os.path.join(CDIR, "Attachment J-2 Requirements.xlsx")
wb = openpyxl.Workbook(); sh = wb.active; sh.title = "Reqs"
sh.append(["ID", "Requirement", "Notes"])
sh.append(["R-1", "The contractor shall deliver a monthly status report to the program office.", "CDRL A001"])
sh.append(["R-2", "The contractor shall maintain a risk register and review it every two weeks with the government.", ""])
sh.append(["R-3", "The contractor shall provide a transition plan within 30 days of award.", "See L.4"])
wb.save(XL)
LIST = os.path.join(EDIR, "Heron_Effort_List.xlsx")
L = E.EffortList(LIST, "Heron"); L.pending = True; L.save()
E.Registry(os.path.join(cfg, "efforts.json")).put("Heron", path=LIST, customer_dir=CDIR)
W.Workspace(W.data_dir_for(LIST), "Heron").set("outline_text", "Volume I Technical\n1.1 Technical Approach\nVolume II Management\n2.1 Program Management")
app = QApplication([]); theme.apply(app)
w = MainWindow(app, cfg, A.APP_VERSION); w.resize(1500, 950); w.show()


def wait(cond, secs=90):
    t = time.time()
    while not cond() and time.time() - t < secs:
        app.processEvents(); time.sleep(0.02)
    return cond()


check("pictures and old Office files are customer documents", all(x in DI.FILE_TYPES for x in (".png", ".jpg", ".tif", ".doc", ".xls", ".ppt"))
      and ".png" in WA.KINDS)
w.open_effort("Heron")
out = os.path.join(EDIR, "Heron_Compliance_Matrix.xlsx")
w.read_page.read([PIC, XL], out)
check("Read RFP reads a picture and a workbook", wait(lambda: w.review is not None))
M = w.review.M
pic_items = [it for it in M.items if it.get("file") == os.path.basename(PIC)]
check("the picture's sentences come in (from its copy with words)", any("staffing plan" in it["text"] for it in pic_items)
      and any("risk management process" in it["text"] for it in pic_items), [it["text"][:40] for it in pic_items])
copies = w.ws.p("page copies")
cp = OP2.copy_path(PIC, copies)
check("the copy with words is kept in the effort's page copies", os.path.isfile(cp) and len(calls) == 1)
check("... and made once per file contents", OP2.make(PIC, copies) == cp and len(calls) == 1)
import pagesent as PS
doc = PS.read(cp)
check("the copy's words sit where they are on the picture", [l["text"] for l in doc["pages"][0]["lines"]][:2] == [TEXT[0][2], TEXT[1][2]]
      and abs(doc["pages"][0]["lines"][1]["box"][0] - 150 * 72 / 200) < 4, [(l["text"][:20], [round(v) for v in l["box"]]) for l in doc["pages"][0]["lines"][:2]])
pg = w.assign_page
w.go("assign"); app.processEvents()
pg.files.setCurrentText(os.path.basename(PIC)); app.processEvents()
if os.environ.get("SHOT_DIR"):
    app.processEvents(); w.grab().save(os.path.join(os.environ["SHOT_DIR"], "picture_view.png"))
check("Assign shows the picture with its sentences boxed", pg.pdf == cp and sum(1 for i in pic_items if i["id"] in pg.view.boxes) >= 2,
      (pg.pdf, len(pg.view.boxes)))
# Sheet view
xl_items = [it for it in M.items if it.get("file") == os.path.basename(XL)]
pg.files.setCurrentText(os.path.basename(XL)); app.processEvents()
check("an Excel workbook opens in Sheet view", pg.sheet is not None and pg.sheet_btn.isVisible() and pg.sheet_btn.text() == "Page view"
      and not pg.pdf and len(pg.sheet["cells"]) == 12)
if os.environ.get("SHOT_DIR"):
    for _ in range(5):
        app.processEvents(); time.sleep(0.05)
    w.grab().save(os.path.join(os.environ["SHOT_DIR"], "sheet_view.png"))
reqs = [it for it in xl_items if it["text"].startswith("The contractor shall")]
check("each sentence is boxed in its cell", len(reqs) == 3 and all(it["id"] in pg.view.boxes for it in reqs), [it["text"][:30] for it in xl_items])
r1 = next(it for it in reqs if "monthly" in it["text"])
cell = next(c for c in pg.sheet["cells"] if c["text"] == r1["text"])
br = pg.view.boxes[r1["id"]].path().boundingRect()
check("... inside its cell (row 2, column B)", cell["r"] == 1 and cell["c"] == 1 and br.left() >= cell["rect"][0] + 8 - 2
      and br.top() >= cell["rect"][1] - 1 and br.bottom() <= cell["rect"][1] + cell["rect"][3] + 1)
check("... with one rectangle per word, for splitting", len(pg.view.words[r1["id"]]) == len(SV.WORD.findall(r1["text"])))
pg.select(r1["id"]); app.processEvents()
k = [m.group(0) for m in SV.WORD.finditer(r1["text"])].index("to")
pg.on_word(r1["id"], k); app.processEvents()
check("a click on a word in a cell sets the split there, shown on the sheet", pg.split_at is not None and pg._split_items and pg.keys["Split"].isEnabled())
pg.do_split(); app.processEvents()
parts = [it for it in M.items if it["id"].startswith(r1["id"])]
check("Split works in Sheet view", len(parts) == 2 and parts[1]["text"].startswith("to the program office"), [p["text"] for p in parts])
check("the queue: two documents, Next document goes to the one not reviewed", "of 2" in pg.queue_lbl.text())
pg.mark_reviewed(); cur = pg.file
nx = pg.next_file(); app.processEvents()
check("... Next document", nx == os.path.basename(PIC) and pg.file == nx and not pg.next_btn.isVisible())
pg.files.setCurrentText(os.path.basename(XL)); app.processEvents()
pg.toggle_sheet(); app.processEvents()
check("Page view: the workbook's PDF copy (Office makes it on the PC)", pg.sheet is None and pg.sheet_btn.text() == "Sheet view")
pg.toggle_sheet(); app.processEvents()
check("... and back to Sheet view", pg.sheet is not None)
# a scanned PDF
SCAN = os.path.join(tmp, "scan.pdf"); img.save(SCAN, "PDF", resolution=200)
check("a scanned PDF needs a copy with words; a typed one does not", OP2.needs_copy(SCAN)
      and not OP2.needs_copy(os.path.join(A.BASE_DIR, "test_samples", "rfp_walk", "Section_L_print.pdf")))
sc = OP2.make(SCAN, copies)
check("the scanned PDF's copy has its words", [l["text"] for l in PS.read(sc)["pages"][0]["lines"]][1] == TEXT[1][2])
mixed = os.path.join(tmp, "mixed.pdf")
import pypdfium2 as pdfium
a = pdfium.PdfDocument(os.path.join(A.BASE_DIR, "test_samples", "rfp_walk", "Section_L_print.pdf"))
a.import_pages(pdfium.PdfDocument(SCAN)); a.save(mixed); a.close()
mc = OP2.make(mixed, copies)
md = PS.read(mc)
check("a mixed PDF: typed pages keep their own words, the scanned page gets the reader's", OP2.scanned_pages(mixed) == [len(md["pages"])]
      and any("SECTION L" in l["text"] for l in md["pages"][0]["lines"]) and md["pages"][-1]["lines"][1]["text"] == TEXT[1][2])
n = sum(checks)
print(f"QT FILETYPES {'PASS' if all(checks) else 'FAIL'} ({n}/{len(checks)})")
sys.exit(0 if all(checks) else 1)
