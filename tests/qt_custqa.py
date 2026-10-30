"""2.98.0 customer Q&A, key dates and what changed, exactly as you see them, offscreen: Capture > Question in Assign
quotes the sentence with its document, section, page and paragraph into Customer Q&A; Capture > Key date puts questions
due on the calendar (and a proposal due date on the effort); the Customer Q&A page approves, edits, orders, exports to
Word and Excel, fills the customer's template and marks the exported questions Submitted; Effort Home shows the counts
and when questions are due; adding an amendment says at once what it changed; Assign outlines the changed and new
sentences of the amendment, the card says what each was, and Show the old page finds it on the old version."""
import os, sys, tempfile, time, zipfile
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE)); sys.path.insert(0, HERE)
os.environ["QT_QPA_PLATFORM"] = "offscreen"
os.environ["AH_TEST"] = "1"
os.environ["SK_SANDBOX_TEMP"] = "temp_qt"
import app_gui as A
import effort as E, workspace as W, effortrec as RR, office_pdf as OP, pdfmini, custqa as Q, compliance as C
OP.available = lambda: False
from PySide6.QtWidgets import QApplication, QMessageBox
QMessageBox.information = staticmethod(lambda *a, **k: None)
from skqt import theme
from skqt.mainwin import MainWindow
from skqt import captureview as CV

checks = []


def check(name, ok, info=""):
    checks.append(bool(ok)); print("PASS" if ok else "FAIL", name, info)


tmp = tempfile.mkdtemp()
cfg = os.path.join(tmp, "config"); os.makedirs(cfg)
EDIR = os.path.join(tmp, "Osprey"); CUST = os.path.join(EDIR, "Customer documents"); os.makedirs(CUST)
IN = os.path.join(tmp, "incoming"); os.makedirs(IN)
BASE = ["L.4 PROPOSAL PREPARATION",
        "Questions shall be submitted no later than 2:00 p.m. Eastern Time on November 20, 2026.",
        "Volume II shall not exceed 20 pages.",
        "The Offeror shall describe its approach to testing the spacecraft bus.",
        "The Offeror shall describe its risk management process."]
AMD = ["L.4 PROPOSAL PREPARATION",
       "Questions shall be submitted no later than 2:00 p.m. Eastern Time on November 20, 2026.",
       "Volume II shall not exceed 25 pages.",
       "The Offeror shall describe its approach to testing the spacecraft bus.",
       "The Offeror shall describe its risk management process and its mitigation plans.",
       "The Offeror shall provide a software quality plan."]
P1 = os.path.join(IN, "Section_L.pdf"); pdfmini.write(P1, [BASE], size=9)
P2 = os.path.join(IN, "Section_L_Amendment_1.pdf"); pdfmini.write(P2, [AMD], size=9)
LIST = os.path.join(EDIR, "Osprey_Effort_List.xlsx")
L = E.EffortList(LIST, "Osprey"); L.pending = True; L.save()
E.Registry(os.path.join(cfg, "efforts.json")).put("Osprey", path=LIST, customer_dir=CUST, solicitation="FA8810-26-R-0001")
W.Workspace(W.data_dir_for(LIST), "Osprey").set("outline_text", "Volume I Technical\n1.1 Technical Approach")
app = QApplication([]); theme.apply(app)
w = MainWindow(app, cfg, A.APP_VERSION); w.resize(1500, 950); w.show()


def wait(cond, secs=60):
    t = time.time()
    while not cond() and time.time() - t < secs:
        app.processEvents(); time.sleep(0.02)
    return cond()


w.open_effort("Osprey")
w.go("docs"); app.processEvents()
w.docs_page.add([P1], choices=[{}])
C1 = os.path.join(CUST, "Section_L.pdf")
w.read_page.read([C1], os.path.join(EDIR, "Osprey_Compliance_Matrix.xlsx"), identities="suggested")
check("read", wait(lambda: w.review is not None))
M = w.review.M
pg = w.assign_page
w.go("assign"); app.processEvents()
by = lambda s: next(it for it in M.items if s in it["text"])
# Capture > Question
pg.open_file("Section_L.pdf"); app.processEvents()
vol = by("20 pages")
pg.select(vol["id"]); app.processEvents()
check("six captures, in two rows", list(pg.cap_keys) == ["contact", "submit", "format", "acronym", "question", "keydate"])
text, src = pg.capture_source()
d = CV.QuestionDialog(pg, text, src, w.people())
check("Question: their wording, the document and the page come from the sentence", d.quote.toPlainText() == "Volume II shall not exceed 20 pages."
      and d.doc.text() == "Section_L.pdf" and d.page.text() == "1", (d.doc.text(), d.page.text(), d.section.text()))
d.question.setPlainText("Will the Government consider a 25-page limit for Volume II, given the added content?")
d.category.setCurrentText("Format and submission")
qid = pg.do_capture("question", answer=d.values()); d.deleteLater()
q = Q.get(RR.Record(w.ws.root), qid)
check("... filed in Customer Q&A as Draft 1, tied to the sentence", q and q["n"] == 1 and q["status"] == "Draft" and q["sid"] == vol["id"]
      and q["category"] == "Format and submission", q)
pg.select(by("risk management")["id"]); app.processEvents()
qid2 = pg.do_capture("question", answer={"question": "Should the risk process cover supplier risk?", "quote": by("risk management")["text"],
                                         "doc": "Section_L.pdf", "category": "Technical"})
# Capture > Key date
pg.select(by("Questions shall be submitted")["id"]); app.processEvents()
kd = CV.KeyDateDialog(pg, pg.capture_source()[0])
vals = kd.values(); kd.deleteLater()
check("Key date: questions due, the date, time and zone are read", vals == {"kind": "Questions due", "title": "Questions due", "date": "2026-11-20",
                                                                         "start": "14:00", "zone": "ET"}, vals)
eid = pg.do_capture("keydate", answer=vals)
ev = next((e for e in RR.Record(w.ws.root).data.get("events", []) if e["id"] == eid), None)
check("... on the calendar, with where it came from", ev and ev["kind"] == "Questions due" and ev["zone"] == "ET" and "Section_L.pdf" in ev["details"], ev)
got = pg.do_capture("keydate", answer={"kind": "Proposal due", "title": "Proposal due", "date": "2026-12-15", "start": "2:00 PM", "zone": "ET"})
r = RR.Record(w.ws.root)
check("a proposal due date becomes the effort's", got == "due" and r.data["due"] == "2026-12-15T14:00" and r.zone_key() == "ET")
# the Customer Q&A page
w.go("custqa"); app.processEvents()
cp = w.custqa_page
check("the page lists them, with when questions are due", cp.table.rowCount() == 2 and cp.summary.text() == "2 question(s): 2 draft"
      and "Nov 20, 2026" in cp.due.text(), (cp.summary.text(), cp.due.text()))
cp.table.selectRow(0); cp.set_status("Approved")
check("Approve", Q.get(RR.Record(w.ws.root), qid)["status"] == "Approved")
cp.select_id(qid2); cp.move(-1)
check("Up: the numbers follow", [x["id"] for x in Q.items(RR.Record(w.ws.root))] == [qid2, qid])
cp.select_id(qid2)
cp.edit(values={"question": "Should the risk process include supplier risk?", "status": "Approved", "owner": "J. Garcia"})
q2 = Q.get(RR.Record(w.ws.root), qid2)
check("Edit", q2["question"].startswith("Should the risk process include") and q2["owner"] == "J. Garcia" and q2["status"] == "Approved")
wd = cp.export("docx", path=os.path.join(tmp, "Q.docx"), choice={"statuses": ["Approved"], "numbers": "list", "mark": False})
x = zipfile.ZipFile(wd).read("word/document.xml").decode()
check("Export to Word: the solicitation on the cover, both questions", "FA8810-26-R-0001" in x and "25-page limit" in x and "supplier risk" in x)
import openpyxl
t = os.path.join(tmp, "Template.xlsx"); wb = openpyxl.Workbook(); ws = wb.active
for k, h in enumerate(["No.", "Reference", "Question", "Answer"], start=1):
    ws.cell(2, k, h)
wb.save(t)
out = cp.export("template", path=os.path.join(tmp, "Filled.xlsx"), choice={"statuses": ["Approved"], "numbers": "new", "mark": True}, template=t)
ws2 = openpyxl.load_workbook(out).active
check("Fill the customer's template: renumbered 1, 2 in their columns", ws2["A3"].value == 1 and ws2["A4"].value == 2
      and ws2["C3"].value.startswith("Should the risk process"), [ws2.cell(3, k).value for k in range(1, 5)])
check("... and marked Submitted", all(q["status"] == "Submitted" for q in Q.items(RR.Record(w.ws.root))))
xl = cp.export("xlsx", path=os.path.join(tmp, "Q.xlsx"), choice={"statuses": ["Submitted"], "numbers": "list", "mark": False, "answers": False})
check("Export to Excel", os.path.isfile(xl))
if os.environ.get("SHOT_DIR"):
    w.go("custqa"); app.processEvents()
    w.grab().save(os.path.join(os.environ["SHOT_DIR"], "custqa_page.png"))
w.go("home"); app.processEvents()
hp = w.home_page
check("Effort Home: the Customer Q&A card", "custqa" in hp.shown and hp.qa_now.get("Submitted") == 2)
# an amendment: what changed, at once
w.go("docs"); app.processEvents()
made = w.docs_page.add([P2], choices=[{}])
check("the amendment is added as Amendment 1", made and made[0][1]["label"] == "Amendment 1")
s = (w.last_changes or [None])[0]
check("what changed is said at once", s and s["title"].startswith("What changed in Amendment 1") and s["changed"] == 2 and s["added"] == 1
      and any("25 pages" in x for x in s["rules"]), s and s["text"])
from skqt.whatview import WhatChangedDialog
wc = WhatChangedDialog(w, s)
check("... its card: Review the changes opens Changes on the pair", wc.review() is not None and w.stack.currentWidget() is w.compare_page)
cpg = w.compare_page
ci = next(i for i, r_ in enumerate(cpg.rows) if r_["kind"] == "changed" and "mitigation" in r_["new"])
cpg.select_index(ci); app.processEvents()
n0 = len(Q.items(RR.Record(w.ws.root)))
cpg.decide_panel.decide("question")
qn = Q.items(RR.Record(w.ws.root))
check("Changes: a Question for the customer is drafted in Customer Q&A, quoting the change", len(qn) == n0 + 1
      and "mitigation plans" in qn[-1]["quote"] and qn[-1]["status"] == "Draft" and "Amendment 1" in qn[-1]["doc"], qn[-1])
dq = next(x for x in __import__("changereview").Review(w.ws.root).load().decisions(cpg.pk, cpg.rows[ci]["key"]) if x["kind"] == "question")
cpg.select_index(ci); app.processEvents()
cpg.decide_panel.undo(dq["id"])
check("... and taking the decision off takes the draft away", len(Q.items(RR.Record(w.ws.root))) == n0)
import readable as RD
files = RD.more(RR.Record(w.ws.root), None, w.ws.root, {})
check("the readable copy has the customer questions", "Customer questions.csv" in files and len(files["Customer questions.csv"][1]) == n0)
# the amendment on the page
w.read_page.read([os.path.join(CUST, "Section_L_Amendment_1.pdf")], None, identities="suggested")
wait(lambda: not w.read_page.running())
w.go("assign"); app.processEvents()
pg.open_file("Section_L_Amendment_1.pdf"); app.processEvents()
M = w.review.M
new25 = next(it for it in M.items if "25 pages" in it["text"] and it.get("file") == "Section_L_Amendment_1.pdf")
newq = next(it for it in M.items if "software quality plan" in it["text"])
check("Assign outlines what the amendment changed and added", pg.change_of.get(new25["id"], {}).get("kind") == "changed"
      and pg.change_of.get(newq["id"], {}).get("kind") == "added" and "changed or new since" in pg.note.text(), pg.note.text())
check("... the unchanged sentences are not outlined", not any("approach to testing" in it["text"] for it in M.items if it["id"] in pg.change_of))
pg.select(new25["id"]); app.processEvents()
check("the card says what it was", getattr(pg, "change_box", None) is not None and "Changed in Amendment 1" in
      " ".join(l.text() for l in pg.change_box.findChildren(type(pg.note))))
dlg = pg.old_page(new25["id"], run=False)
check("Show the old page: the old sentence is found there", dlg is not None and dlg.found is not None and dlg.found[0] == 0)
if os.environ.get("SHOT_DIR"):
    w.grab().save(os.path.join(os.environ["SHOT_DIR"], "changes_on_page.png"))
pg.toggle_changes(); app.processEvents()
check("Hide changes", not pg.change_of and pg.changes_btn.text() == "Show changes")
pg.toggle_changes(); app.processEvents()
w.close_effort()
n = sum(checks)
print(f"QT CUSTQA {'PASS' if n == len(checks) else 'FAIL'} ({n}/{len(checks)})")
sys.exit(0 if n == len(checks) else 1)
