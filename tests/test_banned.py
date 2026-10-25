"""2.27.0 QA/QC banned words: matching rules, the team list (starter, save, import, Knowledge Pack replace), finding
candidates in a Word volume with where and the sentence, and the QA/QC report rows and Excel sheet."""
import os, sys, tempfile
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.environ["AH_TEST"] = "1"


def run():
    import app_gui as A, banned as B, volcheck as V
    from openpyxl import load_workbook
    from tests.test_volcheck import make_docx
    d = tempfile.mkdtemp()
    A.CONFIG_DIR = os.path.join(d, "config"); os.makedirs(A.CONFIG_DIR)

    def hits(phrase, text):
        return [m.group(0) for m in B.pattern(phrase).finditer(text)]
    assert hits("world-class", "a world class team, a world-class team") == ["world class", "world-class"]
    assert hits("state of the art", "our state-of-the-art lab") == ["state-of-the-art"]
    assert hits("guarantee", "We guarantee it; it is guaranteed; guarantees") == ["guarantee", "guaranteed", "guarantees"]
    assert hits("warrant", "the warranty covers it") == [], "a longer different word is not a hit"
    assert hits("utiliz*", "we utilize and utilization") == ["utilize", "utilization"]
    assert hits("TBD", "TBD and Tbd and tbd") == ["TBD"], "short capitals match capitals only"
    assert hits("etc.", "parts, tools, etc. and more") == ["etc."]
    assert hits("[insert*", "[Insert value here]") == ["[Insert"]
    assert hits("very", "every everyday very") == ["very"]

    # the list: starter until saved, save, import, a team pack replaces the starter
    assert not os.path.exists(B.path()) and len(B.load()) == len(B.STARTER)
    B.save(B.starter()[:3] + [{"Phrase": "game changer", "Level": "avoid", "Use instead": "the measured gain"}])
    L = B.load()
    assert len(L) == 4 and next(x for x in L if x["Phrase"] == "game changer")["Level"] == "Avoid"
    csvp = os.path.join(d, "team.csv")
    with open(csvp, "w", encoding="utf-8") as f:
        f.write("Banned word,Severity,Replacement\ngame changer,Banned,\npivot,Avoid,change\n")
    assert B.import_file(csvp) == (1, 1) and len(B.load()) == 5
    os.remove(B.path())
    B.import_file(csvp, keep_starter=False)
    assert {x["Phrase"] for x in B.load()} == {"game changer", "pivot"}, "a team pack replaces the starter list"
    txt = os.path.join(d, "list.txt")
    open(txt, "w").write("# one per line\nsynergy\n")
    assert B.read_file(txt)[0]["Phrase"] == "synergy"

    # find in a Word volume: where it is (text, table, caption) and the sentence around it
    words = B.starter()
    v = os.path.join(d, "Vol 1 Technical.docx")
    make_docx(v, [("Heading1", "1. Approach"),
                  "Our world-class team will guarantee delivery. The value is TBD until CDR.",
                  ("TABLE", [["Item", "Status"], ["Mass margin", "TBD"]]),
                  "Figure 2. The state of the art test stand.",
                  "Plain sentence with nothing to flag."])
    found = B.find(A.typed_chunks(v), words, A.loc_kind)
    by = {(h["phrase"], h["kind"]) for h in found}
    assert ("world-class", "text") in by and ("guarantee", "text") in by and ("TBD", "text") in by, by
    assert ("TBD", "table") in by and ("state of the art", "caption") in by, by
    g = next(h for h in found if h["phrase"] == "guarantee")
    assert "guarantee delivery" in g["sentence"] and g["level"] == "Banned" and g["use"], g
    assert all(h["kind"] != "notes" for h in found)

    # the QA/QC report
    res = A.analyze_batch([v], lambda m: None)
    r = V.build(res, {}, ["Vol 1 Technical.docx"], "volume", file_of=A._file_of, vol_banned={"Vol 1 Technical.docx": found})
    rows = r["banned"]
    assert rows and r["summary"]["banned"][V.BAD] >= 3 and r["summary"]["banned"][V.WARN] >= 1, r["summary"]["banned"]
    g = next(x for x in rows if x["phrase"] == "guarantee")
    assert g["check"] == V.CHECK_MARKS[V.BAD] and g["where"].startswith("Text: ") and g["state"] == "Banned"
    t = next(x for x in rows if x["phrase"] == "TBD" and x["where"].startswith("Table"))
    assert t["level"] == V.BAD
    xl = os.path.join(d, "qa.xlsx"); V.write_excel(r, xl, "Kestrel")
    wb = load_workbook(xl)
    assert "Banned words" in wb.sheetnames and [c.value for c in wb["Banned words"][1]][:4] == ["Check", "Volume", "Found", "Banned phrase"]
    r0 = V.build(res, {}, ["Vol 1 Technical.docx"], "volume", file_of=A._file_of)
    xl0 = os.path.join(d, "qa0.xlsx"); V.write_excel(r0, xl0, "Kestrel")
    assert "Banned words" not in load_workbook(xl0).sheetnames, "no banned check run: no sheet"
    assert "banned_words.csv" in A.PACK_DATA_FILES and any(n == "banned_words.csv" for n, _p in A._team_files())
    print("banned ok")


if __name__ == "__main__":
    run()
