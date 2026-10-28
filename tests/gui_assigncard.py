"""2.58.0 the sentence card in Assign: click a sentence and the card shows its source (file, page and paragraph, RFP
reference), its decision (sections as chips with x and + Section; the requirements it supports as chips), its
author(s) (several names), its relationships with where each is in the RFP, and its notes. The compliance matrix puts
the support in its requirement's cell, each piece with its RFP place, repeated under every requirement it supports."""
import os, sys, tempfile
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE)); sys.path.insert(0, HERE)
os.environ["AH_OCR_ENGINE"] = "tesseract"; os.environ["AH_TEST"] = "1"
import app_gui as A, compliance as C, effort as E, setupdoc as SD, office_pdf as OP, pagesent as PS, matrixdesign as MD
OP.available = lambda: False
from tkinter import messagebox
tmp = tempfile.mkdtemp()
A.CONFIG_DIR = os.path.join(tmp, "config"); A.CHOICES_FILE = os.path.join(tmp, "choices.json"); A.SETTINGS_FILE = os.path.join(tmp, "settings.json")
os.makedirs(A.CONFIG_DIR, exist_ok=True)
for n in ("showinfo", "showwarning"):
    setattr(messagebox, n, lambda *a, **k: None)
messagebox.askyesno = lambda *a, **k: True
checks = []


def check(name, ok, info=""):
    checks.append(bool(ok)); print("PASS" if ok else "FAIL", name, info)


EDIR = os.path.join(tmp, "Kestrel"); os.makedirs(EDIR)
LIST = os.path.join(EDIR, "Kestrel_Effort_List.xlsx")
PDF = os.path.join(A.BASE_DIR, "test_samples", "rfp_walk", "Section_L_print.pdf")
M = C.Matrix(os.path.join(EDIR, "Kestrel_Matrix.xlsx"), "Kestrel")
items, heads, _info = PS.read_review(PDF, "Section_L_print.pdf", C.extract)
M.items, M.headings = items, heads
C.number_items(M.items) if hasattr(C, "number_items") else None
for k, it in enumerate(M.items):
    it.setdefault("id", f"L-{k + 1:02d}")
    it.setdefault("decision", "")
M.sources = [["Section_L_print.pdf", "", "", PDF]]
M.outline = C.parse_outline("1 Technical\n1.1 Approach\n2 Management")
SD.setup_of(M)["complete_v"] = 1
M.save()
app = A.AcronymAppGUI()
app.geometry("1500x1000+0+0")


def go():
    L = E.EffortList(LIST, "Kestrel"); L.pending = True; L.save()
    app.efforts.put("Kestrel", path=LIST)
    app.open_effort("Kestrel")
    app.ws.set("compliance_path", M.path)
    app.open_compliance(); app._c_load(M.path, quiet=True); app.update()
    app.shell_go("compliance", "decide"); app.update()
    win = app._cwin
    app.cp_open_file("Section_L_print.pdf"); app.update()
    Mx = app.matrix
    by = lambda w: next(it for it in Mx.items if it["text"].startswith(w))
    req = by("The Offeror shall describe its technical approach")
    req2 = by("Each volume shall be written")
    sup = by("Information provided in one volume")
    app._cp_select(req["id"]); app.update()
    where = win.cp_where.cget("text")
    check("the card shows the file, the page and paragraph, and the RFP section",
          "File: Section_L_print.pdf" in where and "Where: p.1, para" in where and "RFP: L.3.1 Technical Approach" in where, where)
    check("an unassigned sentence says so", "Not used" in win.ac_kind.cget("text"), win.ac_kind.cget("text"))
    app.cp_section_click("1.1 Approach"); app.cp_req(); app.update()
    app._cp_select(req["id"]); app.update()
    check("a requirement: its kind and its section as a chip", "Compliance requirement" in win.ac_kind.cget("text")
          and any("1.1 Approach" in w.cget("text") for ch in win.ac_dec.winfo_children()
                  for w in ch.winfo_children() if isinstance(w, A.ctk.CTkLabel)), win.ac_kind.cget("text"))
    app.ac_add_section(req["id"], "2 Management"); app.update()
    check("+ Section puts it in a second section; the card stays on it", SD.item_sections(Mx, req) == ["1.1 Approach", "2 Management"]
          and app._c_selected()[0]["id"] == req["id"], SD.item_sections(Mx, req))
    app.cp_remove_section(req["id"], "2 Management"); app.update()
    check("the chip's x takes it out of that section", SD.item_sections(Mx, req) == ["1.1 Approach"])
    win.ac_auth.delete(0, "end"); win.ac_auth.insert(0, "A. Lee;  K. Park")
    app.ac_set_authors(win.ac_auth.get()); app.update()
    check("several authors", req["owner"] == "A. Lee, K. Park" and SD.owners_list(Mx, req) == ["A. Lee", "K. Park"]
          and "K. Park" in SD.people(Mx), req.get("owner"))
    app.ac_set_notes("Use the thermal model figure."); app.update()
    check("notes saved on the sentence", req["notes"] == "Use the thermal model figure.")
    # the second requirement, and a support sentence for both
    app._cp_select(req2["id"]); app.cp_section_click("2 Management"); app.cp_req(); app.update()
    app._cp_select(sup["id"]); app.update()
    C.assign_context(Mx.items, [sup], [req["id"], req2["id"]])
    app._c_changed(); app._cp_select(sup["id"]); app.update()
    labels = [w.cget("text") for ch in win.ac_dec.winfo_children() for w in ch.winfo_children() if hasattr(w, "cget")
              and isinstance(w, A.ctk.CTkLabel)]
    check("a support sentence: the requirements it supports as chips, with their sections",
          "Supporting context" in win.ac_kind.cget("text") and any(req["id"] in x and "1.1 Approach" in x for x in labels)
          and any(req2["id"] in x and "2 Management" in x for x in labels), labels)
    rel = [w.cget("text") for line in win.cp_rel.winfo_children() for w in line.winfo_children() if isinstance(w, A.ctk.CTkLabel)]
    check("relationships say where each one is in the RFP", any(x.startswith("(L.3.1") for x in rel), rel)
    # the matrix
    d = MD.default()
    check("the default matrix puts support in its requirement's cell", d["context"] == "cell")
    heads_, rows, n = MD.preview(Mx, d)
    ti = heads_.index("RFP content"); ai = heads_.index("Author(s)")
    r1 = next(r for r in rows if r[ti].startswith("The Offeror shall describe its technical approach"))
    r2 = next(r for r in rows if r[ti].startswith("Each volume shall be written"))
    check("the support is in both requirements' cells, with its RFP place",
          "Supporting context:\n(L.1, p.1) Information provided in one volume" in r1[ti] and "(L.1, p.1) Information provided" in r2[ti], r1[ti][-160:])
    check("no row of its own for it", not any(r[ti].startswith("Information provided in one volume") for r in rows))
    check("the authors column lists both", r1[ai] == "A. Lee, K. Park", r1[ai])
    d["context"] = "rows"
    _h, rows2, _n = MD.preview(Mx, d)
    check("On its own rows: a row under each requirement it supports", sum(1 for r in rows2 if r[ti].startswith("Information provided")) == 2)
    out = os.path.join(tmp, "m.xlsx")
    MD.export(Mx, out, MD.default())
    from openpyxl import load_workbook
    ws = load_workbook(out)["Compliance Matrix"]
    vals = [str(c.value) for row in ws.iter_rows(min_row=5) for c in row if c.value]
    check("the Excel cell holds the requirement and its support", any("Supporting context:" in v and "technical approach" in v for v in vals))
    n = sum(checks)
    print(f"GUI ASSIGNCARD PASS ({n}/{len(checks)})" if n == len(checks) else f"GUI ASSIGNCARD FAIL ({n}/{len(checks)})")
    app.after(300, app.destroy)


def run():
    try:
        go()
    except Exception:
        import traceback
        traceback.print_exc()
        print("GUI ASSIGNCARD FAIL (exception)")
        app.after(300, app.destroy)


app.after(400, run)
app.mainloop()
