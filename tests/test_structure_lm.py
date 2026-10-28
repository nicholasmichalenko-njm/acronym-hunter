"""2.68.0 every customer document has an initial and a type (RFP, SOW, SOO, Amendment ...); requirements are L or M
(Shipley), suggested from the section and the document type, and support may serve any number of L and M items; the
matrix carries Doc, L/M, Section and the reference R-L-4.2.1, and a tab for each volume; the response structure is built
while assigning, as deep as needed (7 levels and more), and everything assigned follows its section when numbers change."""
import os, sys, tempfile
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE)); sys.path.insert(0, HERE)
import compliance as C, setupdoc as SD, matrixdesign as MD, outline_ops as O, docversions as DV, docintake as DI
from test_assign import row


def library():
    tmp = tempfile.mkdtemp()
    lib = DV.Library(tmp)
    files = []
    for name in ("Kestrel_RFP.docx", "Kestrel_SOW.docx"):
        p = os.path.join(tmp, name)
        open(p, "w").write(name)
        files.append(p)
    return lib, files


def run():
    # document identity
    assert DV.guess_type("Kestrel_SOW_rev2.docx") == "SOW" and DV.guess_type("FA8810_Amendment_0003.pdf") == "Amendment"
    assert DV.guess_type("Statement_of_Objectives.docx") == "SOO" and DV.guess_type("Kestrel.docx") == "RFP"
    assert DV.clean_initial(" r-2! ") == "R2"
    lib, files = library()
    plan = DI.plan(lib, files, reader=lambda p: [{"x": f"Text of {os.path.basename(p)} number {k}.", "p": 1} for k in range(3)])
    for e in plan:
        e["choice"] = "new"
    assert [(e["type"], e["initial"]) for e in plan] == [("RFP", "R"), ("SOW", "S")], [(e["type"], e["initial"]) for e in plan]
    plan[1]["initial"] = "R"
    assert any("already used" in x for x in DI.check(lib, plan)), "two documents cannot share an initial"
    plan[1]["initial"] = ""
    assert any("give it an initial" in x for x in DI.check(lib, plan)), "an initial is required"
    plan[1]["initial"], plan[1]["type"] = "S", "Nonsense"
    assert any("pick its type" in x for x in DI.check(lib, plan)), "a type is required"
    plan[1]["type"] = "SOW"
    assert DI.check(lib, plan) == []
    DI.apply(lib, plan)
    assert lib.identity_map()["Kestrel_SOW.docx"]["initial"] == "S" and not lib.missing_identity()
    rfp = lib.doc_of_file("Kestrel_RFP.docx")
    assert lib.set_identity(rfp["id"], "S", "RFP"), "a taken initial is refused"
    assert lib.set_identity(rfp["id"], "K", "RFP") == "" and lib.identity_map()["Kestrel_RFP.docx"]["initial"] == "K"
    assert lib.suggest_initial("Amendment") == "A1" and lib.suggest_initial("RFP") == "R"

    # L and M, and the reference
    a = row(1, "L.4.2", "The Offeror shall describe the thermal design.")
    b = row(2, "M.2.1", "The Government will evaluate the thermal approach.")
    c = row(3, "C.3.1", "The contractor shall deliver monthly reports.")
    s = row(4, "L.4.2", "Thermal margins matter because the orbit has long eclipses.")
    for it in (a, b, s):
        it["file"] = "Kestrel_RFP.docx"
    c["file"] = "Kestrel_SOW.docx"
    M = C.Matrix(os.path.join(tempfile.mkdtemp(), "Kestrel_Matrix.xlsx"), "Kestrel")
    M.items = [a, b, c, s]
    M.outline = []
    assert SD.sync_docs(M, lib) == 2 and SD.initial_of(M, a) == "K" and SD.initial_of(M, c) == "S"
    assert SD.lm_suggest(M, a) == "L" and SD.lm_suggest(M, b) == "M" and SD.lm_suggest(M, c) == ""
    # no structure yet: requirements are kept with no section, and L or M is set on them
    C.assign_requirement(M.items, [a, b, c], "")
    assert all(it["decision"] == "Keep" and not it.get("map") for it in (a, b, c)), "assign before any structure"
    SD.set_lm(M, [c], "L")
    assert SD.full_ref(M, a) == "K-L-4.2" and SD.full_ref(M, b) == "K-M-2.1" and SD.full_ref(M, c) == "S-L-C.3.1"
    C.assign_context(M.items, [s], [a["id"], b["id"]])
    assert set(C.cf_ids(s["context_for"])) == {a["id"], b["id"]}, "support serves an L and an M"

    # build the structure while assigning: volumes, headings, 8 levels down
    v1 = O.add_volume(M, "Technical")
    t = O.add(M, "Technical Approach", parent=v1)
    deep = t
    for k in range(7):
        deep = O.add(M, f"Level {k + 3}", parent=deep)
    assert deep.split(" ")[0] == "1.1.1.1.1.1.1.1" and M.outline[O.index_of(M, deep)]["level"] == 9, deep
    C.assign_requirement(M.items, [a], deep)
    th = O.add(M, "Thermal", parent=t)
    C.assign_requirement(M.items, [b], th)
    mg = O.add(M, "Management", after=t)
    C.assign_requirement(M.items, [c], mg)
    assert mg.startswith("2 ") and th.startswith("1.2 "), (mg, th)
    up = O.move(M, mg, -1)
    assert up.startswith("1 ") and c["map"] == up, "a requirement follows its section when it moves"
    assert b["map"].startswith("2.2 ") and a["map"].startswith("2.1.1.1"), (a["map"], b["map"])
    v2 = O.add_volume(M, "Cost")
    cost = O.add(M, "Basis of Estimate", parent=v2)
    assert cost.startswith("1 "), "numbers restart in each volume"
    SD.add_section(M, b, cost)                                 # one more section: b is in both volumes
    assert O.volumes_of_item(M, b) == [v1, v2] and O.volumes_of_item(M, c) == [v1]
    assert set(O.in_volume(M, v2)) == {v2, cost}
    gone = O.delete(M, th)
    assert b["map"] == gone or b["map"] == cost or SD.item_sections(M, b), "a deleted section's sentences go up"
    ren = O.rename(M, up, "Program Management")
    assert c["map"] == ren

    # the matrix: Doc, L/M, Section, Reference, and a tab per volume
    d = MD.default()
    keys = [x["key"] for x in d["columns"]]
    assert keys[3:7] == ["reference", "doc", "lm", "rfp_ref"], keys
    its, dat = MD.rows(M, dict(d, context="rows"))
    ra = dat[[x.get("id") for x in its].index(a["id"])]
    assert ra[keys.index("doc")] == "K" and ra[keys.index("lm")] == "L" and ra[keys.index("reference")] == "K-L-4.2", ra
    rs = dat[[x.get("id") for x in its].index(s["id"])]
    assert rs[keys.index("lm")] in ("L/M", "L"), rs
    out = os.path.join(tempfile.mkdtemp(), "Kestrel_Compliance_Matrix.xlsx")
    MD.export(M, out, d, "Kestrel")
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb.sheetnames[0] == "Compliance Matrix" and v1[:31] in wb.sheetnames and v2[:31] in wb.sheetnames, wb.sheetnames
    vtext = " ".join(str(c.value or "") for r in wb[v2[:31]].iter_rows() for c in r)
    assert "thermal approach" in vtext and "monthly reports" not in vtext, "a volume's tab has only its rows"
    paths = SD.export(M, tempfile.mkdtemp())
    assert len(paths) == 2 and all(os.path.isfile(p) for p in paths), "the Response Setup: one file per volume, eight levels deep"
    print("structure, L and M, document identity ok")


if __name__ == "__main__":
    run()
