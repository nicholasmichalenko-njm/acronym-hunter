"""2.14.0 Effort start wizard and the effort record: the wizard creates the effort with client, solicitation and
internal number, remembers the customer's lists, copies and reads the customer documents, and opens the review of new
terms; one click adds them as Proposed with Found in; scans update Found in / Last found / Scan status, mark new terms,
and flag terms no longer found; Apply Choice writes the definition into the list; Activity shows the history; exports
for the whole effort, per volume, and all efforts carry effort, client and Found in."""
import os, sys, tempfile, csv, time
sys.stdout.reconfigure(encoding="utf-8", errors="replace")   # Windows consoles default to cp1252
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ["AH_OCR_ENGINE"] = "tesseract"
os.environ["AH_TEST"] = "1"   # no first-start wizard during tests
import app_gui as A, effort as E, docxw
from tkinter import messagebox
from openpyxl import load_workbook
tmp = tempfile.mkdtemp()
A.CONFIG_DIR = os.path.join(tmp, "config"); A.CHOICES_FILE = os.path.join(tmp, "config", "choices.json")
A.SETTINGS_FILE = os.path.join(tmp, "settings.json")
A.DB_FILE = os.path.join(tmp, "db.csv")
with open(A.DB_FILE, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["Term", "Category", "Definition", "Action"])
    w.writerow(["SOW", "Acronym", "Statement of Work", "Include"]); w.writerow(["SOW", "Acronym", "Scope of Work", "Include"])
messagebox.showinfo = lambda *a, **k: None
messagebox.showerror = lambda *a, **k: print("ERROR BOX:", a[:2])
messagebox.askyesno = lambda *a, **k: False
CL = os.path.join(tmp, "Customer_Acronyms.csv")
with open(CL, "w", newline="", encoding="utf-8") as f:
    csv.writer(f).writerows([["Acronym", "Definition"], ["TRL", "Technology Readiness Level"], ["CDR", "Critical Design Review"]])
RFP = os.path.join(A.BASE_DIR, "test_samples", "rfp")
DOCS = [os.path.join(RFP, f) for f in sorted(os.listdir(RFP))]
VOL = os.path.join(tmp, "Vol_1_Technical.docx")
def write_vol(with_zeta):
    d = docxw.Doc(); d.heading("1 Approach", 1)
    d.para("The SOW tasks and deliverables follow the TRL plan." + (" The ZQX interface carries data." if with_zeta else ""))
    d.save(VOL)
write_vol(True)
FOLDER = os.path.join(tmp, "Efforts", "Kestrel Thermal")
checks = {}


def check(name, ok, info=""):
    checks[name] = bool(ok)
    print(("PASS " if ok else "FAIL ") + name, info)


app = A.AcronymAppGUI()
st = {"t": 0}


def wait(cond, then, limit=120):
    st["t"] += 1
    if not cond() and st["t"] < limit:
        app.after(500, lambda: wait(cond, then, limit)); return
    st["t"] = 0
    then()


def step1():
    w = app.open_wizard(); w.update(); api = w._api; v = api["v"]
    v["name"].set("Kestrel Thermal"); v["client"].set("USSF SSC"); v["sol"].set("FA8810-27-R-0042"); v["internal"].set("SS-2027-114")
    st_ = api["st"]; st_["folder_set"] = True; v["folder"].set(FOLDER)
    api["next"](); w.update(); w._add_files([CL])
    api["next"](); w.update(); w._add_files(DOCS)
    api["next"](); w.update(); v["format"].set("both")
    api["next"](); w.update()
    api["next"]()                                   # Create effort
    wait(lambda: getattr(app, "_review_win", None) is not None, step2)


def step2():
    info = app.efforts.get("Kestrel Thermal") or {}
    check("effort created with client fields", app.effort is not None and info.get("client") == "USSF SSC"
          and info.get("solicitation") == "FA8810-27-R-0042" and info.get("internal_no") == "SS-2027-114", info)
    L = E.EffortList(info["path"]).load()
    check("list's About sheet names the client", L.client == "USSF SSC" and L.solicitation == "FA8810-27-R-0042")
    check("customer documents copied into the effort folder",
          sorted(os.listdir(os.path.join(FOLDER, "Customer documents"))) == sorted(os.path.basename(p) for p in DOCS))
    check("customer lists remembered by the effort", "TRL" in A.CLIENT_DB and app.ws.get("client_lists"))
    rw = app._review_win; tv = rw._api["tree"]
    terms = list(tv.get_children())
    check("review opens with the new terms", len(terms) > 3, len(terms))
    pick = terms[:3]
    tv.selection_set(pick); rw._api["add"]()
    row = app.effort.rows.get(pick[0], {})
    check("one click adds them as Proposed with Found in", all(t in app.effort.rows for t in pick) and row.get("Status") == "Proposed"
          and row.get("Found in", "").endswith((".docx", ".pdf")) and row.get("Scan status") == E.SCAN_NEW, row.get("Found in"))
    check("history says where it came from", any("added (Proposed) from" in h[3] for h in app.effort.history), app.effort.history[-1:])
    rw.destroy()
    st["pick"] = pick
    app.effort.add_or_update("SOW", Definition="", Type="Acronym"); app._save_effort_list()
    app.start_batch([VOL]); wait(lambda: not app.busy() and not app.btn_cancel.winfo_ismapped(), step3)


def step3():
    L = app.effort
    check("scan updates the list: Found in and Last found", "Vol_1_Technical.docx" in L.rows["SOW"].get("Found in", "")
          and L.rows["SOW"].get("Last found") == E.today(), L.rows["SOW"].get("Found in"))
    check("terms not in the list marked new", app.results.get("ZQX", {}).get("new_in_effort"), app.results.get("ZQX", {}).get("status"))
    app.status_var.set("New in this scan (not in the effort list)"); app.populate_table()
    check("filter shows new terms", "ZQX" in app.tree.get_children() and "SOW" not in app.tree.get_children())
    app.status_var.set(A.STATUS_FILTERS[0]); app.populate_table()
    check("effort line reports the last scan", "last scan: " in app.lbl_effort.cget("text") and "USSF SSC" in app.lbl_effort.cget("text"),
          app.lbl_effort.cget("text")[-80:])
    app.tree.selection_set("ZQX"); app.add_to_effort(["ZQX"])
    # Apply Choice writes into the list
    app.selected_term = "SOW"; app.combo.configure(state="normal", values=app.results["SOW"]["definitions"]); app.combo.set("Scope of Work")
    app.apply_choice()
    check("Apply Choice goes into the effort list", L.rows["SOW"]["Definition"] == "Scope of Work", L.rows["SOW"]["Definition"])
    write_vol(False); time.sleep(1.1)
    app.scan_again(); wait(lambda: not app.busy() and not app.btn_cancel.winfo_ismapped(), step4)


def step4():
    L = app.effort
    check("term gone from a rescanned file is flagged", L.rows["ZQX"].get("Scan status", "").startswith("No longer found in Vol_1_Technical.docx"),
          L.rows["ZQX"].get("Scan status"))
    d = app.results.get("ZQX", {})
    check("and shows in the results as no longer found", d.get("gone") and d.get("status", "").startswith("No longer found"), d.get("status"))
    w = app.show_activity(); w.update()
    check("activity lists what happened", len(w._tree.get_children()) >= 5, len(w._tree.get_children())); w.destroy()
    out = os.path.join(tmp, "lists"); os.makedirs(out)
    made = app.export_effort_lists(folder=out, scope="both")
    names = [os.path.basename(p) for p in made]
    check("whole-effort and per-volume lists", any("Vol_1_Technical" in n for n in names) and any(n.endswith(".xlsx") for n in names)
          and not any("Kestrel_RFP" in n for n in names), names)
    ws = load_workbook([p for p in made if p.endswith(".xlsx")][0])["Acronym list"]
    head = [c.value for c in ws[1]]
    rows = [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    check("Excel list has Found in, effort and client", rows and all(r["Effort"] == "Kestrel Thermal" and r["Client"] == "USSF SSC" for r in rows)
          and any(r["Found in"] for r in rows))
    got = app.export_all_efforts(os.path.join(tmp, "all.xlsx"))
    ws = load_workbook(got[0])["All efforts"]
    head = [c.value for c in ws[1]]
    check("all efforts workbook carries effort, client, solicitation, Found in", got[1] >= 4 and
          {"Effort", "Client", "Solicitation", "Found in", "Internal number"} <= set(head), (got[1], head))
    ok = all(checks.values())
    print("GUI WIZARD", "PASS" if ok else "FAIL", f"({sum(checks.values())}/{len(checks)})")
    app.on_close()


app.after(1500, step1)
app.mainloop()
