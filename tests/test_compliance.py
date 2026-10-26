"""Compliance matrix rules (2.8.0) on the fictional RFP in test_samples/rfp: sections, sentences, list items, actor,
type, strength, suggestion, cross-references, links, split, amendment re-read, outline, mapping, workbook, Word outline."""
import os, sys, tempfile, zipfile
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import compliance as C

RFP = os.path.join(ROOT, "test_samples", "rfp")


def by_text(items, start):
    hit = [it for it in items if it["text"].startswith(start)]
    assert hit, start
    return hit[0]


def run():
    import app_gui as A
    files = sorted(os.listdir(RFP))
    heads = []
    items = C.extract([(f, A.typed_chunks(os.path.join(RFP, f))) for f in files], heads)
    req = [it for it in items if it["type"] != "Context"]
    assert len(req) == 27, len(req)
    ctx = [it for it in items if it["type"] == "Context"]                  # 2.29.0: every sentence is a row
    assert ctx and all(it["suggest"] in ("Context", "Drop") for it in ctx), [(it["text"], it["suggest"]) for it in ctx]
    assert any("Resumes do not count toward the page limit" in it["text"] for it in ctx)
    assert len(C.extract([(f, A.typed_chunks(os.path.join(RFP, f))) for f in files], every=False)) == 27
    ids = [it["id"] for it in items]
    assert len(set(ids)) == len(ids)
    # Section C (SOW)
    c21 = by_text(items, "The Contractor shall provide a thermal design")
    assert (c21["id"], c21["type"], c21["actor"], c21["strength"], c21["suggest"]) == ("C.2.1-01", "SOW", "Contractor", "Mandatory", "Keep")
    assert c21["section_title"] == "TECHNICAL REQUIREMENTS" and c21["file"].endswith(".docx") and "C.2" in c21["loc"]
    icd = by_text(items, "The Contractor shall deliver the Interface Control Document")
    assert icd["compound"] and icd["xrefs"] == "CDRL A003"
    gov = by_text(items, "The Government will provide the launch vehicle")
    assert (gov["type"], gov["strength"], gov["suggest"]) == ("Information", "Information", "Drop")
    assert by_text(items, "The Contractor should use")["strength"] == "Desired"
    assert by_text(items, "This acquisition is for")["type"] == "Context", "plain statements are Context rows, not requirements"
    # Section L (instructions, format, lists)
    assert by_text(items, "Proposals shall not exceed 50 pages")["type"] == "Format"
    assert by_text(items, "Proposals are due no later than 2:00 p.m. Mountain Time")["type"] == "Format"
    th = by_text(items, "The Offeror shall describe its approach to the thermal design")
    assert th["type"] == "Instruction" and th["xrefs"] == "C.2.1" and th["links"] == "C.2.1-01"
    a, b, c = (by_text(items, x) for x in ("(a) identify", "(b) describe", "(c) provide a risk register"))
    assert a["lead"] == "The offeror shall:" and c["xrefs"] == "Attachment J-3" and not a["compound"]
    assert by_text(items, "Describe the verification approach")["actor"] == "Offeror"
    may = by_text(items, "The Offeror may include")
    assert (may["strength"], may["suggest"]) == ("Optional", "Review")
    # PDF: wrapped lines rejoined, list items with their lead-in across a page break
    vol = by_text(items, "The Offeror shall organize the proposal into three volumes")
    assert "Volume III Cost" in vol["text"] and vol["page"] == 1
    poc = by_text(items, "identify a point of contact")
    assert poc["page"] == 2 and poc["lead"].startswith("For each reference, the Offeror shall") and poc["section"] == "L.3"
    assert not any(it["text"].endswith(":") for it in items), "lead-ins are not items"
    assert by_text(items, "The Government will not evaluate")["suggest"] == "Drop"
    # Section M (evaluation) linked back to the instructions
    m21 = by_text(items, "The Government will evaluate the offeror's thermal design")
    assert m21["type"] == "Evaluation" and m21["suggest"] == "Keep" and m21["links"] == "L.4.1-01"
    # split a compound requirement
    parts = C.split_text(icd["text"])
    assert parts == ["The Contractor shall deliver the Interface Control Document (CDRL A003).",
                     "The Contractor shall maintain it through launch."], parts
    # decisions, outline, mapping
    for it in items:
        it["decision"] = it["suggest"] if it["suggest"] != "Review" else ""
    outline = C.outline_from_headings(heads, "L")
    labels = [(C.outline_label(o), o["level"]) for o in outline]
    assert labels == [("L.1 GENERAL INSTRUCTIONS", 1), ("L.2 VOLUME ORGANIZATION", 1), ("L.3 PAST PERFORMANCE", 1),
                      ("L.4 TECHNICAL VOLUME", 1), ("L.4.1 Technical Approach", 2), ("L.4.2 Management Approach", 2)], labels
    stats = C.word_df(items)
    assert C.suggest_map(th, outline, stats)[0] == "L.4.1 Technical Approach"
    typed = C.parse_outline("Volume I - Technical\n1.0 Technical Approach\n    1.1 Thermal Design\n1.2 Risk Management\nVolume II Management")
    assert [(o["num"], o["title"], o["level"]) for o in typed] == [("Volume I", "Technical", 1), ("1.0", "Technical Approach", 2),
                                                                   ("1.1", "Thermal Design", 2), ("1.2", "Risk Management", 2), ("Volume II", "Management", 1)], typed
    assert C.suggest_map(th, typed, stats)[0] == "1.1 Thermal Design", C.suggest_map(th, typed, stats)
    assert C.suggest_map(c, typed, stats)[0] in ("", "1.2 Risk Management"), "never a wrong section"
    scope = by_text(items, "The Contractor shall design, build, and test")
    assert C.suggest_map(scope, typed, stats)[0] == "", "one shared common word is not enough"
    assert C.suggest_map(vol, typed, stats)[0] != "Volume I Technical", "volumes with subsections take no items"
    # workbook round trip (Excel edits count)
    d = tempfile.mkdtemp()
    M = C.Matrix(os.path.join(d, "Kestrel_Compliance.xlsx"), "Kestrel")
    M.items, M.outline, M.headings = items, outline, heads
    th["map"], th["owner"] = "L.4.1 Technical Approach", "J. Smith"
    M.log(th["id"], "mapped"); M.save()
    N = C.Matrix(M.path).load()
    assert len(N.items) == len(items) and N.name == "Kestrel" and len(N.outline) == len(outline) and len(N.headings) == len(heads)
    t2 = [it for it in N.items if it["id"] == th["id"]][0]
    assert t2["map"] == "L.4.1 Technical Approach" and t2["owner"] == "J. Smith" and t2["decision"] == "Keep"
    assert [it for it in N.items if it["id"] == icd["id"]][0]["compound"] is True
    from openpyxl import load_workbook
    wb = load_workbook(M.path)
    assert wb.sheetnames == ["Compliance Matrix", "Outline", "Sources", "Coverage", "RFP headings", "History", "About"], wb.sheetnames
    ws = wb["Compliance Matrix"]
    for r in ws.iter_rows(min_row=2):
        if r[0].value == may["id"]:
            r[1].value = "keep"                       # someone types the decision in Excel
    wb.save(M.path)
    assert [it for it in C.Matrix(M.path).load().items if it["id"] == may["id"]][0]["decision"] == "Keep"
    # amendment: re-read keeps decisions; changed sentences show up as new / no longer present
    new = C.extract([(f, A.typed_chunks(os.path.join(RFP, f))) for f in files])
    changed = [x for x in new if x["text"].startswith("Proposals shall not exceed 50 pages")][0]
    changed["text"] = changed["text"].replace("50 pages", "40 pages"); changed["key"] = C.text_key(changed["text"])
    merged = C.merge_reread(N.items, new)
    assert [x for x in merged if x["id"] == th["id"]][0]["map"] == "L.4.1 Technical Approach"
    assert any(x["new"] == "New in latest read" and "40 pages" in x["text"] for x in merged)
    assert any(x["new"] == "Not in the latest read" and "50 pages" in x["text"] for x in merged)
    assert len({x["id"] for x in merged}) == len(merged)
    # Word outline: headings + mapped table + unmapped list, readable by the app
    out = os.path.join(d, "outline.docx")
    info = C.write_outline_docx(N, out)
    assert info["mapped"] == 1 and info["unmapped"] >= 15, info
    xml = zipfile.ZipFile(out).read("word/document.xml").decode()
    assert 'w:val="Heading3"' in xml and "Kept requirements not mapped" in xml and th["id"] in xml
    from xml.dom import minidom
    minidom.parseString(xml.encode())
    heads = C.outline_from_docx_xml(xml.encode())
    assert any(h["title"].startswith("Technical Approach") for h in heads), heads[:4]
    assert A.typed_chunks(out)
    # 2.29.0: every sentence gets a decision; the coverage report says plainly whether any line was missed
    cov = N.coverage()
    tot = cov[-1]
    assert tot["file"] == "Total" and tot["total"] == len(N.items), tot
    cp = os.path.join(d, "coverage.docx")
    ci = C.write_coverage_docx(N, cp)
    cx = zipfile.ZipFile(cp).read("word/document.xml").decode()
    minidom.parseString(cx.encode())
    assert ci["sentences"] == len(N.items) and ci["undecided"] > 0 and "Every sentence was read:" in cx and "Not used" in cx, ci
    for it in N.items:
        it["decision"] = it["decision"] or "Context"
    ci = C.write_coverage_docx(N, cp)
    cx = zipfile.ZipFile(cp).read("word/document.xml").decode()
    assert ci["undecided"] == 0 and "Every sentence was read:" in cx and C.COVER_FILL["Context"] in cx
    print("compliance ok")


if __name__ == "__main__":
    run()
