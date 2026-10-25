"""Volume check (2.7.0): two Word volumes through the app's own scanner, checked against an Effort List.
Covers first use per volume and across the submission, re-definitions, wording that differs, terms missing from the
list, not approved, unused list entries, the per-volume acronym lists, the Excel report and the Word acronym list."""
import os, sys, tempfile, zipfile
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import volcheck as V


def make_docx(path, paras):
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    def para(p):
        style = ""
        if isinstance(p, tuple):                      # ("Heading1", "2. Technical Approach")
            style, p = f'<w:pPr><w:pStyle w:val="{p[0]}"/></w:pPr>', p[1]
        return f'<w:p>{style}<w:r><w:t xml:space="preserve">{V._x(p)}</w:t></w:r></w:p>'
    body = "".join(para(p) for p in paras)
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
                   '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr("word/document.xml", f'<?xml version="1.0" encoding="UTF-8"?><w:document {W}><w:body>{body}</w:body></w:document>')


def run():
    import app_gui as A
    import custdefs
    import effort as E
    d = tempfile.mkdtemp()
    v1, v2 = os.path.join(d, "Vol 1 Technical.docx"), os.path.join(d, "Vol 2 Management.docx")
    make_docx(v1, ["This Statement of Work (SOW) covers the whole program for the customer.",
                   "The QZR follows the design phase. The Quarterly Zonal Review (QZR) happens in month nine.",
                   "Telemetry Tracking and Control (TTK) radios link the vehicle to the ground.",
                   "Thermal blankets called XBL wrap the bus, and the GNX team owns pointing.",
                   "As noted, the Statement of Work (SOW) applies to every subcontract."])
    make_docx(v2, ["The SOW applies here as well, and the QZR schedule holds.",
                   "Ground stations use TTK links during every pass."])
    L = E.EffortList(os.path.join(d, "Effort.xlsx"), "Kestrel", "volume")
    L.add_or_update("SOW", Type="Acronym", Definition="Statement of Work"); L.set_status(["SOW"], "Approved", "J. Smith")
    L.add_or_update("QZR", Type="Acronym", Definition="Quarterly Zonal Review"); L.set_status(["QZR"], "Approved", "J. Smith")
    L.add_or_update("TTK", Type="Acronym", Definition="Telemetry, Tracking, and Command"); L.set_status(["TTK"], "Approved", "J. Smith")
    L.add_or_update("XBL", Type="Acronym", Definition="Extra Blanket Layer")
    L.add_or_update("PDQ", Type="Acronym", Definition="Pretty Darn Quick"); L.set_status(["PDQ"], "Approved", "J. Smith")
    L.add_or_update("ZZZ", Type="Acronym", Definition="Not a term"); L.set_status(["ZZZ"], "Rejected")

    res = A.analyze_batch([v1, v2], lambda m: None)
    for t in ("SOW", "QZR", "TTK", "XBL", "GNX"):
        assert t in res, (t, sorted(res)[:40])
    fu = res["SOW"]["first_use"]["Vol 1 Technical.docx"]
    assert fu["defined"] is True and fu["lf"] == "Statement of Work", fu
    assert fu.get("redefs"), fu
    names = ["Vol 1 Technical.docx", "Vol 2 Management.docx"]
    defs = {os.path.basename(p): custdefs.extract(A.typed_chunks(p)) for p in (v1, v2)}
    pic = lambda loc: bool(A._OCR_LOC.search(loc))

    r = V.build(res, L.rows, names, "volume", defs, file_of=A._file_of, in_picture=pic)
    by = {(x["volume"], x["term"]): x for x in r["first_use"]}
    sow = by[("Vol 1 Technical.docx", "SOW")]
    assert sow["level"] == V.INFO and "defined again" in sow["state"] and sow["also_defined"], sow
    assert by[("Vol 1 Technical.docx", "QZR")]["level"] == V.BAD and "later in the same passage" in by[("Vol 1 Technical.docx", "QZR")]["state"]
    ttk = by[("Vol 1 Technical.docx", "TTK")]
    assert ttk["level"] == V.WARN and "differ" in ttk["state"] and ttk["written"].startswith("Telemetry"), ttk
    assert by[("Vol 1 Technical.docx", "XBL")]["level"] == V.BAD and "Never defined" in by[("Vol 1 Technical.docx", "XBL")]["state"]
    # per-volume rule: volume 2 must define SOW again
    assert by[("Vol 2 Management.docx", "SOW")]["level"] == V.BAD
    assert any(m["term"] == "GNX" for m in r["missing"]) and not any(m["term"] in ("SOW", "ZZZ") for m in r["missing"])
    assert [x["term"] for x in r["not_approved"]] == ["XBL"], r["not_approved"]
    assert any(x["term"] == "TTK" and x["volume"] == "Vol 1 Technical.docx" for x in r["differs"]), r["differs"]
    assert [x["term"] for x in r["unused"]] == ["PDQ"], r["unused"]
    l1 = {x["term"]: x for x in r["lists"]["Vol 1 Technical.docx"]}
    assert l1["SOW"]["definition"] == "Statement of Work" and l1["GNX"]["status"] == "Not in effort list" and "ZZZ" not in l1
    assert set(x["term"] for x in r["lists"]["Vol 2 Management.docx"]) >= {"SOW", "QZR", "TTK"}

    # across the submission: volume 2 may use SOW bare (defined in volume 1); QZR was spelled out later in volume 1
    r2 = V.build(res, L.rows, names, "submission", defs, file_of=A._file_of, in_picture=pic)
    by2 = {(x["volume"], x["term"]): x for x in r2["first_use"]}
    assert by2[("Vol 2 Management.docx", "SOW")]["level"] == V.OK and "earlier" in by2[("Vol 2 Management.docx", "SOW")]["state"]
    assert by2[("Vol 2 Management.docx", "QZR")]["level"] == V.OK
    assert by2[("Vol 1 Technical.docx", "XBL")]["level"] == V.BAD

    up = V.effort_updates(r)
    assert up["SOW"]["Used in volumes"] == "Vol 1 Technical.docx; Vol 2 Management.docx"

    xl = os.path.join(d, "check.xlsx"); V.write_excel(r, xl, "Kestrel")
    from openpyxl import load_workbook
    wb = load_workbook(xl)
    assert wb.sheetnames[:6] == ["Summary", "First use", "Not in effort list", "Not approved", "Definition differs", "Unused in volumes"], wb.sheetnames
    assert len(wb.sheetnames) == 8
    dx = os.path.join(d, "list.docx")
    V.write_docx([(x["term"], x["definition"]) for x in r["lists"]["Vol 1 Technical.docx"]], dx)
    with zipfile.ZipFile(dx) as z:
        docxml = z.read("word/document.xml").decode()
    assert "Statement of Work" in docxml and "<w:tblHeader/>" in docxml
    from xml.dom import minidom
    minidom.parseString(docxml.encode())
    assert A.typed_chunks(dx), "the app reads its own acronym list back"

    # 2.22.0: first use restarts in each major section of each volume
    sec = A.section_of
    assert sec("V.docx (section '2. KEY PERFORMANCE', para 4-9)") == "2"
    assert sec("V.docx (section '2.1 Thermal Design', para 10-12)") == "2"
    assert sec("V.docx (section 'C.3 Work Breakdown', para 1-3)") == "C"
    assert sec("V.docx (section 'Section M - Evaluation', para 1-3)") == "M"
    assert sec("V.docx (section 'Executive Summary', para 1-3)") == "Executive Summary"
    assert sec("V.docx (section 'A Novel Approach', para 1-3)") == "A Novel Approach"
    assert sec("V.docx (para 1-3)") == "(start)"
    assert sec("V.docx (Body, image 'Figure 1') [OCR]") is None and sec("V.docx (Footnotes)") is None
    assert sec("V.pdf (Page 3)") == "" and sec("V.pptx (Slide 2)") == ""
    v3 = os.path.join(d, "Vol 3 Past Performance.docx")
    make_docx(v3, ["Kestrel past performance volume.",
                   ("Heading1", "1. Introduction"),
                   "The Statement of Work (SOW) sets the tasks. The PDQ goal holds.",
                   ("Heading2", "1.1 Scope"),
                   "The SOW lists every task.",
                   ("Heading1", "2. Contracts"),
                   "Each SOW task maps to a past contract.",
                   ("Heading1", "3. Lessons"),
                   "Pretty Darn Quick (PDQ) delivery and the Statement of Work (SOW) review."])
    res3 = A.analyze_batch([v3], lambda m: None)
    fs = res3["SOW"]["first_use_sec"]["Vol 3 Past Performance.docx"]
    assert list(fs) == ["1", "2", "3"], fs
    assert fs["1"]["defined"] is True and fs["2"]["defined"] is False and fs["3"]["defined"] is True, fs
    r3 = V.build(res3, L.rows, ["Vol 3 Past Performance.docx"], "section", file_of=A._file_of, in_picture=pic)
    by3 = {(x["section"], x["term"]): x for x in r3["first_use"]}
    assert by3[("1", "SOW")]["level"] == V.OK and by3[("3", "SOW")]["level"] == V.OK, by3
    assert by3[("2", "SOW")]["level"] == V.BAD and "Never defined in this section" in by3[("2", "SOW")]["state"]
    assert by3[("1", "PDQ")]["level"] == V.BAD and "later" not in by3[("1", "PDQ")]["state"]
    assert by3[("3", "PDQ")]["level"] == V.OK and r3["summary"]["rescan_for_sections"] == 0
    # the per-volume rule on the same volume sees one first use
    rv = V.build(res3, L.rows, ["Vol 3 Past Performance.docx"], "volume", file_of=A._file_of, in_picture=pic)
    assert [x["term"] for x in rv["first_use"]].count("SOW") == 1 and not any(x["section"] for x in rv["first_use"])
    # results saved before 2.22.0 (no section data) fall back to the volume and say so
    old = {t: {k: v for k, v in dd.items() if k != "first_use_sec"} for t, dd in res3.items()}
    ro = V.build(old, L.rows, ["Vol 3 Past Performance.docx"], "section", file_of=A._file_of, in_picture=pic)
    assert ro["summary"]["rescan_for_sections"] >= 1 and [x["term"] for x in ro["first_use"]].count("SOW") == 1
    assert "section 2" in V.effort_updates(r3)["SOW"]["First use"]
    xl3 = os.path.join(d, "check3.xlsx"); V.write_excel(r3, xl3, "Kestrel")
    ws = load_workbook(xl3)["First use"]
    assert [c.value for c in ws[1]][:3] == ["Volume", "Section", "Acronym"]
    assert load_workbook(xl3)["Summary"]["B3"].value == "Restarts in each section of each volume"
    L3 = E.EffortList(os.path.join(d, "Effort3.xlsx"), "Kestrel", "section"); L3.save()
    L4 = E.EffortList(L3.path); L4.load(); assert L4.first_use == "section"
    print("volcheck ok")


if __name__ == "__main__":
    run()
