"""2.13.0 Setup Builder logic: references (RFP numbering with list markers, inherited, file rules, typed), placement
shared with the Outline column, copies and order, management direction, plain export (purple RFP text with the
reference on the same line, green direction, every heading, one file per volume, matrix name in the document
properties), export check, export changes only, compare a styled document, save and reload in the workbook."""
import os, sys, tempfile, zipfile, re
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.environ["AH_OCR_ENGINE"] = "tesseract"
import compliance as C, setupdoc as SD
import app_gui as A

RFP = os.path.join(ROOT, "test_samples", "rfp")


def doc_paras(path):
    return SD.docx_paragraphs(path)


def colors(path):
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    return set(re.findall(r'w:color w:val="([0-9A-F]{6})"', xml))


def run():
    d = tempfile.mkdtemp()
    files = sorted(os.listdir(RFP))
    heads = []
    M = C.Matrix(os.path.join(d, "Kestrel_Matrix.xlsx"), "Kestrel")
    M.items = C.extract([(f, A.typed_chunks(os.path.join(RFP, f))) for f in files], heads)
    M.headings = heads
    M.outline = C.parse_outline("Volume I Technical\n1.1 Thermal Design\n1.2 Risk Management\nVolume II Management\n2.1 Schedule\n2.2 Staffing")
    for it in M.items:
        it["decision"] = "Keep" if it["suggest"] != "Drop" else "Drop"
    kept = [it for it in M.items if it["decision"] == "Keep"]
    assert len(kept) > 10
    labs = SD.labels(M)
    assert labs[1] == "1.1 Thermal Design" and len(SD.volume_groups(M)) == 2
    # references
    S = SD.setup_of(M)
    a = next(it for it in M.items if it["text"].startswith("(a) identify"))
    ref, src = SD.ref_of(M, a)
    assert ref.endswith("(a)") and src == "RFP numbering", (ref, src)
    num = next(it for it in kept if it["section"] and it["section"][0].isalpha())
    r2, _ = SD.ref_of(M, num)
    assert r2.startswith(num["section"]), r2
    fake = dict(num, id="Z-1", section="3.2", file="Attachment_3_Instructions.pdf", text="Provide a plan.")
    M.items.append(fake)
    assert SD.suggest_prefix("Attachment_3_Instructions.pdf") == "L" and SD.suggest_prefix("Kestrel_SOW.docx") == "SOW"
    S["rules"]["Attachment_3_Instructions.pdf"] = "L"
    assert SD.ref_of(M, fake) == ("L-3.2", "File rule: L"), SD.ref_of(M, fake)
    nosec = dict(fake, id="Z-2", section="", text="Show how risks are retired.")
    M.items.append(nosec)
    assert SD.ref_of(M, nosec) == ("L-3.2", "File rule: L") or SD.ref_of(M, nosec)[1] in ("Inherited", "File rule: L"), SD.ref_of(M, nosec)
    lone = dict(fake, id="Z-3", section="", file="other.pdf", text="Nothing to inherit here at all.")
    M.items.append(lone)
    assert SD.ref_of(M, lone) == ("", "No reference found")
    S["refs"]["Z-3"] = "L.2 (para 4)"
    assert SD.ref_of(M, lone) == ("L.2 (para 4)", "Typed")
    assert SD.block_text(M, lone).startswith("L.2 (para 4): Nothing")
    S["ref_format"] = "brackets"
    assert SD.block_text(M, lone).startswith("[L.2 (para 4)] Nothing")
    S["ref_format"] = "colon"
    del S["refs"]["Z-3"]
    assert SD.block_text(M, lone).startswith(SD.MISSING)
    M.items.remove(lone)
    # placement shares the Outline column
    ids = [it["id"] for it in kept[:6]]
    SD.place(M, ids[0], "1.1 Thermal Design"); SD.place(M, ids[1], "1.1 Thermal Design"); SD.place(M, ids[2], "1.1 Thermal Design", before=ids[0])
    assert SD.placed_in(M)["1.1 Thermal Design"] == [ids[2], ids[0], ids[1]]
    assert SD.by_id(M)[ids[0]]["map"] == "1.1 Thermal Design"
    SD.place(M, ids[0], "2.1 Schedule", copy=True)
    P = SD.placed_in(M)
    assert ids[0] in P["2.1 Schedule"] and ids[0] in P["1.1 Thermal Design"], "a copy shows in both sections"
    SD.move(M, ids[1], "1.1 Thermal Design", -1)
    assert SD.placed_in(M)["1.1 Thermal Design"][1] == ids[1]
    SD.place(M, ids[3], "1.2 Risk Management")
    SD.unplace(M, ids[3], "1.2 Risk Management")
    assert SD.by_id(M)[ids[3]]["map"] == "" and ids[3] not in SD.placed_in(M)["1.2 Risk Management"]
    S["mgmt"]["1.1 Thermal Design"] = "Lead with the radiator. Owner: J. Ortiz."
    S["notes"][ids[1]] = "We exceed this: cite the 2024 test."
    # context paragraphs
    n = SD.context_from_chunks(M, files[0], A.typed_chunks(os.path.join(RFP, files[0])))
    assert n > 0 and S["context"][0]["id"] == "X-001" and SD.kind_of(S["context"][0]) == "Context"
    SD.place(M, "X-001", "1.2 Risk Management")
    # export check and export
    chk = SD.export_check(M)
    assert chk["unplaced"] and chk["blocks"] >= 4, chk["blocks"]
    out = os.path.join(d, "out"); os.makedirs(out)
    paths = SD.export(M, out)
    assert len(paths) == 2 and all(os.path.isfile(p) for p in paths), paths
    p1 = doc_paras(paths[0])
    texts = [t for t, _s, _l in p1]
    assert texts[0].startswith("Volume I") and "1.2 Risk Management" in texts, texts[:6]   # empty headings export too
    styles = {t: s for t, s, _l in p1}
    assert styles["1.1 Thermal Design"] == "Heading2" and styles[texts[0]] == "Heading1"
    i = texts.index("1.1 Thermal Design")
    assert re.match(r"^[A-Z][\w.\-()]*: ", texts[i + 1]), texts[i + 1]                    # reference on the same line
    assert "Lead with the radiator. Owner: J. Ortiz." in texts and "We exceed this: cite the 2024 test." in texts
    assert texts.index("We exceed this: cite the 2024 test.") < texts.index("Lead with the radiator. Owner: J. Ortiz.")
    assert colors(paths[0]) == {SD.PURPLE, SD.GREEN}, colors(paths[0])
    custom = zipfile.ZipFile(paths[0]).read("docProps/custom.xml").decode("utf-8")
    assert "AcronymHunterMatrix" in custom and "Kestrel" in custom
    assert "Volume II" in doc_paras(paths[1])[0][0]
    # export changes only
    assert SD.changes_since_last(M) == []
    SD.place(M, ids[4], "1.2 Risk Management")
    S["mgmt"]["2.1 Schedule"] = "Use the IMS from the capture plan."
    SD.unplace(M, ids[1], "1.1 Thermal Design")
    ch = SD.changes_since_last(M)
    kinds = sorted(k for _l, k, _i, _t in ch)
    assert kinds == ["Added", "Direction changed", "Removed"], kinds
    cp = os.path.join(d, "changes.docx")
    SD.export_changes(M, cp)
    ct = [t for t, _s, _l in doc_paras(cp)]
    assert "1.2 Risk Management" in ct and "Use the IMS from the capture plan." in ct
    assert SD.changes_since_last(M) == [], "handed-over changes are not listed again"
    # compare a styled document: the first export (before the changes) is missing the new block
    rows = SD.compare(M, paths[0])
    probs = {(r[1].split(":")[0], r[2]) for r in rows}
    assert ("Missing", ids[4]) in probs, rows
    assert ("Not in the setup", ids[1]) in probs, rows
    # save and reload
    M.save()
    R = C.Matrix(M.path).load()
    RS = SD.setup_of(R)
    assert RS["mgmt"]["2.1 Schedule"] == "Use the IMS from the capture plan." and RS["rules"]["Attachment_3_Instructions.pdf"] == "L"
    assert SD.placed_in(R)["1.1 Thermal Design"] == SD.placed_in(M)["1.1 Thermal Design"]
    assert len(RS["context"]) == len(S["context"]) and len(RS["exports"]) == 1
    from openpyxl import load_workbook
    wb = load_workbook(M.path, read_only=True)
    assert {"Setup", "Setup direction", "Setup data"} <= set(wb.sheetnames)
    wb.close()
    print("setup ok")


if __name__ == "__main__":
    run()
