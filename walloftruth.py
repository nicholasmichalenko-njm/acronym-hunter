"""Wall of Truth (3.3.0): the one approved value for every fact the proposal states more than once (cleared staff,
years supporting the customer, facility square footage, contract values, team size) and the one way to write names
and terms. Each truth has the words it appears beside in a draft; Check the drafts reads the Word drafts sentence by
sentence and lists every number beside those words that does not match, every old value still in use after the
truth changed, and every wrong form written, with the volume and section. The team's wall (company facts for every
effort) lives with the app; each effort's wall lives in its record. Rules only: no AI and no ML."""
import datetime, getpass, hashlib, json, os, re, uuid

KINDS = ("Number", "Fact", "Name", "Term")
TEAM_FILE = "wall_of_truth.json"
NEAR = 6                                     # words between a truth's context words and a number beside them
_NUM = re.compile(r"(?<![\w.$])\$?\d[\d,]*(?:\.\d+)?(?:\s?(?:%|percent\b|[KMB]\b|million\b|billion\b))?", re.I)
_SENT = re.compile(r"(?<=[.!?;])\s+(?=[A-Z0-9(\"'])")


def _now():
    return datetime.datetime.now().isoformat(timespec="minutes")


def _clean(v, n=600):
    return " ".join(str(v or "").split())[:n]


def norm_number(text):
    """'$1,200' -> '1200', '15 percent' -> '15%', '2.5M' -> '2.5m', '1.5 million' -> '1.5m'; '' when no number."""
    m = _NUM.search(text or "")
    if not m:
        return ""
    t = m.group(0).lower().replace("$", "").replace(",", "").replace(" ", "")
    t = t.replace("percent", "%").replace("million", "m").replace("billion", "b").replace("k", "k")
    if re.fullmatch(r"\d+\.0+", t):
        t = t.split(".")[0]
    return t


def _list(v):
    if isinstance(v, str):
        v = re.split(r"[,;\n]+", v)
    return [x for x in (_clean(y, 200) for y in (v or [])) if x]


# ------------------------------------------------------------------ an effort's wall, or the team's
def items(rec):
    return sorted((t for t in rec.data.get("truths", []) if isinstance(t, dict)), key=lambda t: (t.get("order", 0)))


def get(rec, tid):
    return next((t for t in rec.data.get("truths", []) if t.get("id") == tid), None)


def _new(label, value, kind="Number", context=(), wrong=(), owner="", source="", notes="", who=None):
    return {"id": uuid.uuid4().hex[:10], "label": _clean(label, 160), "value": _clean(value, 300),
            "kind": kind if kind in KINDS else "Fact", "context": _list(context), "wrong": _list(wrong), "old": [],
            "owner": _clean(owner, 120), "source": _clean(source, 300), "notes": _clean(notes, 600),
            "confirmed": datetime.date.today().isoformat(), "who": who or getpass.getuser(), "at": _now()}


def add(rec, label, value, kind="Number", context=(), wrong=(), owner="", source="", notes="", who=None):
    """A truth on the effort's wall; '' without both a label and a value."""
    if not _clean(label) or not _clean(value):
        return ""
    ts = rec.data.setdefault("truths", [])
    t = _new(label, value, kind, context or [_clean(label).lower()], wrong, owner, source, notes, who)
    t["order"] = 1 + max([x.get("order", 0) for x in ts] or [0])
    ts.append(t)
    rec.log("effort", f"Wall of Truth: {t['label']} = {t['value']}", ref={"truth": t["id"]}, who=who)
    rec.save()
    return t["id"]


def update(rec, tid, who=None, **fields):
    """Changes a truth. A new value keeps the one before it in "old", dated, so the check finds it still in use."""
    t = get(rec, tid)
    if t is None:
        return False
    changed = []
    for k, v in fields.items():
        if k in ("label", "value", "owner", "source", "notes"):
            v = _clean(v, 600)
            if k in ("label", "value") and not v:
                continue
        elif k == "kind":
            if v not in KINDS:
                continue
        elif k in ("context", "wrong"):
            v = _list(v)
        else:
            continue
        if t.get(k) != v:
            if k == "value":
                t.setdefault("old", []).append({"value": t["value"], "until": _now(), "who": who or getpass.getuser()})
            t[k] = v
            changed.append(k)
    if changed:
        if "value" in changed:
            t["confirmed"] = datetime.date.today().isoformat()
            rec.log("effort", f"Wall of Truth: {t['label']} now {t['value']} (was {t['old'][-1]['value']})", ref={"truth": tid}, who=who)
        else:
            rec.log("effort", f"Wall of Truth changed: {t['label']}", ref={"truth": tid}, who=who)
        rec.save()
    return bool(changed)


def confirm(rec, tid, who=None):
    """Marks a truth checked again today (its value still right)."""
    t = get(rec, tid)
    if t is None:
        return False
    t["confirmed"] = datetime.date.today().isoformat()
    rec.log("effort", f"Wall of Truth confirmed: {t['label']} = {t['value']}", ref={"truth": tid}, who=who)
    rec.save()
    return True


def remove(rec, tid, who=None):
    t = get(rec, tid)
    if t is None:
        return False
    rec.data["truths"] = [x for x in rec.data.get("truths", []) if x.get("id") != tid]
    rec.log("effort", f"Wall of Truth removed: {t['label']}", who=who)
    rec.save()
    return True


def stale(t, days=30, today=None):
    """True when the truth has not been confirmed for `days` days."""
    try:
        d = datetime.date.fromisoformat(t.get("confirmed", "")[:10])
    except ValueError:
        return True
    return ((today or datetime.date.today()) - d).days > days


# ------------------------------------------------------------------ the team's wall
def team_load(config_dir):
    try:
        with open(os.path.join(config_dir, TEAM_FILE), encoding="utf-8") as f:
            return [t for t in (json.load(f) or {}).get("truths", []) if isinstance(t, dict) and t.get("label")]
    except (OSError, ValueError):
        return []


def team_save(config_dir, truths):
    os.makedirs(config_dir, exist_ok=True)
    p = os.path.join(config_dir, TEAM_FILE)
    with open(p + ".tmp", "w", encoding="utf-8") as f:
        json.dump({"truths": truths}, f, indent=1, ensure_ascii=False)
    os.replace(p + ".tmp", p)


def to_team(config_dir, t):
    """Puts a truth on the team's wall (or updates the one with the same label). Returns True when it changed."""
    team = team_load(config_dir)
    keep = {k: t.get(k) for k in ("label", "value", "kind", "context", "wrong", "source", "notes")}
    keep["confirmed"] = t.get("confirmed", datetime.date.today().isoformat())
    hit = next((x for x in team if x["label"].casefold() == t["label"].casefold()), None)
    if hit is not None and all(hit.get(k) == v for k, v in keep.items()):
        return False
    if hit is not None:
        hit.update(keep)
    else:
        team.append(keep)
    team_save(config_dir, team)
    return True


def from_team(rec, config_dir, who=None):
    """Adds the team's truths this effort does not have yet (by label). Returns how many."""
    have = {t["label"].casefold() for t in items(rec)}
    n = 0
    for x in team_load(config_dir):
        if x["label"].casefold() in have:
            continue
        if add(rec, x["label"], x.get("value", ""), x.get("kind", "Fact"), x.get("context", []), x.get("wrong", []),
               source=x.get("source", "") or "team wall", notes=x.get("notes", ""), who=who):
            n += 1
    return n


# ------------------------------------------------------------------ the drafts
def sentences(text):
    return [s.strip() for s in _SENT.split(text or "") if s.strip()]


def _words(s):
    return [(m.start(), m.end()) for m in re.finditer(r"\S+", s)]


def _near_numbers(sentence, phrase):
    """Numbers within NEAR words of the phrase in the sentence: [(normalized, as written)], nearest first."""
    low = sentence.lower()
    words = _words(sentence)
    out = []
    for m in re.finditer(r"\b" + re.escape(phrase.lower()) + r"\b", low):
        wi = next((k for k, (a, b) in enumerate(words) if a <= m.start() < b), None)
        if wi is None:
            continue
        we = next((k for k, (a, b) in enumerate(words) if a < m.end() <= b), wi)
        for n in _NUM.finditer(sentence):
            ni = next((k for k, (a, b) in enumerate(words) if a <= n.start() < b), None)
            if ni is None:
                continue
            dist = (wi - ni) if ni < wi else (ni - we) if ni > we else 0
            if 0 <= dist <= NEAR:
                out.append((dist, norm_number(n.group(0)), n.group(0).strip()))
    out.sort()
    return [(n, w) for _d, n, w in out]


def _key(tid, vol, sentence):
    return hashlib.sha1(f"{tid}|{vol}|{sentence}".encode("utf-8")).hexdigest()[:12]


def check(rec, drafts, labels=()):
    """Every place a draft disagrees with the wall. drafts: [(volume, .docx path)]. Returns [{"id", "label", "value",
    "volume", "section", "sentence", "problem", "found", "key"}], the ignored ones left out."""
    import themes as TH
    ignored = set(rec.data.get("truth_ignored", []))
    parts = []
    for vol, path in drafts:
        try:
            parts.append((vol, TH.draft_sections(path, list(labels))))
        except Exception:
            continue
    out = []
    for t in items(rec):
        want = norm_number(t["value"]) if t.get("kind") in ("Number", "Fact") else ""
        olds = {norm_number(o["value"]) for o in t.get("old", []) if norm_number(o["value"])} - {want}
        old_text = [o["value"] for o in t.get("old", []) if not norm_number(o["value"]) and o["value"] != t["value"]]
        for vol, secs in parts:
            for sec, text in secs.items():
                for s in sentences(text):
                    probs = []
                    low = s.lower()
                    for w in t.get("wrong", []):
                        if re.search(r"(?<!\w)" + re.escape(w.lower()) + r"(?!\w)", low):
                            probs.append((f"says “{w}”; the wall says “{t['value']}”", w))
                    for w in old_text:
                        if re.search(r"(?<!\w)" + re.escape(w.lower()) + r"(?!\w)", low):
                            probs.append((f"old wording “{w}”; now “{t['value']}”", w))
                    if want:
                        for ph in t.get("context", []):
                            near = _near_numbers(s, ph)
                            if not near or any(n == want for n, _w in near):
                                continue
                            n, w = near[0]
                            probs.append((f"old value {w}; now {t['value']}" if n in olds else f"says {w}; the wall says {t['value']}", w))
                            break
                    for p, w in dict.fromkeys(probs):
                        k = _key(t["id"], vol, s)
                        if k in ignored:
                            continue
                        out.append({"id": t["id"], "label": t["label"], "value": t["value"], "volume": vol, "section": sec,
                                    "sentence": s[:600], "problem": p, "found": w, "key": k})
    return out


def ignore(rec, key, on=True):
    """Marks one finding as right as written (or not), so the next check leaves it out."""
    ig = rec.data.setdefault("truth_ignored", [])
    if on and key not in ig:
        ig.append(key)
    elif not on and key in ig:
        ig.remove(key)
    else:
        return False
    rec.save()
    return True


# ------------------------------------------------------------------ outputs
def export_word(path, rec, effort="", marking=None):
    """The wall on a page or two: what, the approved value, words to avoid, owner and when it was last confirmed."""
    import docxw
    d = docxw.Doc()
    d.marking = marking
    d.props = {"StationkeeperWallOfTruth": effort}
    d.heading(f"Wall of Truth{': ' + effort if effort else ''}", 1)
    d.para(f"The one approved way to state each of these. Updated {datetime.date.today().strftime('%B %d, %Y').replace(' 0', ' ')}. "
           "Ask the owner before writing anything different.")
    for kind in KINDS:
        ts = [t for t in items(rec) if t.get("kind") == kind]
        if not ts:
            continue
        d.heading({"Number": "Numbers", "Fact": "Facts", "Name": "Names", "Term": "Terms"}[kind], 2)
        d.table(["What", "Write it as", "Not", "Owner", "Confirmed"],
                [[t["label"], t["value"], "; ".join(t.get("wrong", []) + [o["value"] for o in t.get("old", [])]),
                  t.get("owner", ""), t.get("confirmed", "")] for t in ts], [2300, 2700, 2160, 1200, 1000])
    d.save(path)
    return path


def export_findings(path, findings, effort=""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Wall of Truth check"
    ws.append(["Truth", "Approved", "Volume", "Section", "Problem", "Sentence"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1E3A8A")
    for f in findings:
        ws.append([f["label"], f["value"], f["volume"], f["section"], f["problem"], f["sentence"]])
    for i, w in enumerate([24, 18, 22, 24, 40, 90], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    if effort:
        wb.properties.title = f"{effort} Wall of Truth check"
    wb.save(path)
    return path
