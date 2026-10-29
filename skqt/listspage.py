"""Lists to submit (2.86.0): the acronym lists that go in the submission, built from the submission check. Each list
holds only the terms its volume actually uses, with the effort list's approved definition; a term the customer's
documents use but no volume uses is left out, and the page says which ones and why. One list per volume, one master
list for all volumes, or both: set once per effort. Rules only (No AI); nothing leaves the PC."""
import os
from PySide6.QtCore import Qt
from PySide6.QtWidgets import (QButtonGroup, QFileDialog, QHBoxLayout, QHeaderView, QRadioButton, QTableWidget, QTableWidgetItem,
                               QTabWidget, QVBoxLayout, QWidget)
import diag
import volcheck as V
from . import design as D

MODES = (("per", "One list per volume"), ("master", "One master list for all volumes"), ("both", "Both"))
MASTER = "All volumes"


def lists_of(report, mode):
    """[(name, rows)] in submission order: each volume's list, the master list, or both."""
    if not report:
        return []
    sets = [(v, report["lists"].get(v, [])) for v in report["volumes"]]
    combo = {}
    for _v, rows in sets:
        for x in rows:
            combo.setdefault(x.get("key", x["term"]), x)
    master = (MASTER, sorted(combo.values(), key=lambda x: (x["term"].lower(), x["term"])))
    if mode == "per":
        return sets
    if mode == "master":
        return [master]
    return sets + ([master] if len(sets) > 1 else [])


def _sheet(name, wb):
    """A worksheet name Excel accepts (no []:*?/\\, at most 31 characters, not used yet)."""
    base = "".join(ch for ch in os.path.splitext(name)[0] if ch not in "[]:*?/\\")[:28] or "List"
    out, n = base, 2
    while out in wb.sheetnames:
        out = f"{base[:26]} {n}"; n += 1
    return out


def open_items(rows):
    return [x for x in rows if x["status"] != "Approved" or not x["definition"]]


class ListsPage(QWidget):
    def __init__(self, host):
        super().__init__()
        self.setObjectName("page")
        self.host = host
        v = QVBoxLayout(self)
        v.setContentsMargins(D.SPACE["xl"], D.SPACE["l"], D.SPACE["xl"], D.SPACE["l"]); v.setSpacing(D.SPACE["m"])
        self.where = D.label("<b>END</b> &nbsp;The acronym lists that go in your submission: only the terms each volume uses, "
                             "with the effort list's approved definitions. Built from the latest Submission check.", "muted")
        self.where.setTextFormat(Qt.RichText); v.addWidget(self.where)
        mrow = QHBoxLayout(); v.addLayout(mrow)
        mrow.addWidget(D.label("Lists", "label", wrap=False))
        self.mode = QButtonGroup(self)
        for k, text in MODES:
            rb = QRadioButton(text); rb.setProperty("mode", k); self.mode.addButton(rb); mrow.addWidget(rb)
        self.mode.buttonClicked.connect(lambda b: self.set_mode(b.property("mode")))
        mrow.addStretch(1)
        self.b_word = D.button("Export the lists (Word)...", "primary", lambda: self.export_word())
        self.b_xlsx = D.button("Export (Excel)...", None, lambda: self.export_excel())
        self.b_record = D.button("Record in the effort list", None, lambda: self.record(), "Used in volumes and first use, for each term")
        for b in (self.b_word, self.b_xlsx, self.b_record):
            mrow.addWidget(b)
        prow = QHBoxLayout(); v.addLayout(prow)
        self.p_check = D.Pill("", "accent"); self.p_open = D.Pill("", "ok"); self.p_out = D.Pill("", "accent")
        for p in (self.p_check, self.p_open, self.p_out):
            prow.addWidget(p)
        prow.addStretch(1)
        self.tabs = QTabWidget(); v.addWidget(self.tabs, 1)
        self.empty = D.EmptyState("No submission check yet", "The lists are built from what your volumes use. Mark the volumes "
                                  "you will submit and check them first.", "Go to Submission check", lambda: self.host.go("qa"))
        v.addWidget(self.empty)

    @property
    def report(self):
        s = getattr(self.host.acronym_page, "session", None)
        return s.volume_report if s is not None else None

    def current_mode(self):
        m = self.host.ws.get("acronym_list_mode", "") if self.host.ws is not None else ""
        if m not in dict(MODES):
            r = self.report
            m = "both" if r and len(r.get("volumes", [])) > 1 else "per"
        return m

    def set_mode(self, m):
        if self.host.ws is None or m not in dict(MODES):
            return False
        self.host.ws.set("acronym_list_mode", m)
        diag.activity("acronym_list_mode_set")
        self.refresh()
        return True

    def refresh(self):
        h = self.host
        self.setEnabled(h.ws is not None)
        if h.ws is not None and getattr(h.acronym_page, "session", None) is None:
            h.acronym_page.load()
        r = self.report
        m = self.current_mode() if h.ws is not None else "per"
        for b in self.mode.buttons():
            b.setChecked(b.property("mode") == m)
        self.tabs.clear()
        on = bool(r)
        self.tabs.setVisible(on); self.empty.setVisible(not on)
        for b in (self.b_word, self.b_xlsx, self.b_record):
            b.setEnabled(on)
        for p in (self.p_check, self.p_open, self.p_out):
            p.setVisible(on)
        if not on:
            return []
        sets = lists_of(r, m)
        opened = sum(len(open_items(rows)) for name, rows in sets if name != MASTER)
        self.p_check.setText(f"From the check of {len(r['volumes'])} volume(s)")
        self.p_open.setText(f"{opened} need a definition or approval" if opened else "every term approved and defined")
        self.p_open.set_tone("warn" if opened else "ok")
        unused = r.get("unused", [])
        self.p_out.setText(f"{len(unused)} effort-list term(s) left out: no volume uses them")
        for name, rows in sets:
            self.tabs.addTab(self._table(rows), f"{os.path.splitext(name)[0]} ({len(rows)})")
        if unused:
            t = QTableWidget(len(unused), 4)
            t.setHorizontalHeaderLabels(["Left out", "Definition", "Status", "Why"])
            t.verticalHeader().setVisible(False); t.setEditTriggers(QTableWidget.NoEditTriggers)
            t.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
            for i, x in enumerate(unused):
                for c, val in enumerate((x["term"], x.get("definition", ""), x.get("status", ""), "Not used in any submission volume")):
                    it = QTableWidgetItem(val)
                    it.setForeground(D.color("muted"))
                    t.setItem(i, c, it)
            self.tabs.addTab(t, f"Left out ({len(unused)})")
        return sets

    def _table(self, rows):
        t = QTableWidget(len(rows), 3)
        t.setHorizontalHeaderLabels(["Term", "Definition", "Status"])
        t.verticalHeader().setVisible(False); t.setEditTriggers(QTableWidget.NoEditTriggers)
        t.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        t.setColumnWidth(0, 160); t.setColumnWidth(2, 160)
        for i, x in enumerate(rows):
            bad = x["status"] != "Approved" or not x["definition"]
            for c, val in enumerate((x["term"], x["definition"] or "(no definition yet)", x["status"])):
                it = QTableWidgetItem(val)
                if bad and c > 0:
                    it.setForeground(D.color("warn"))
                t.setItem(i, c, it)
        return t

    # ------------------------------------------------------------ export
    def export_word(self, folder=None):
        r = self.report
        if not r:
            return None
        if folder is None:
            folder = QFileDialog.getExistingDirectory(self, "Folder for the acronym lists", os.path.dirname(self.host.effort_path))
            if not folder:
                return None
        import re
        written = []
        try:
            os.makedirs(folder, exist_ok=True)
            for name, rows in lists_of(r, self.current_mode()):
                p = os.path.join(folder, re.sub(r"[^A-Za-z0-9 _.-]+", "_", os.path.splitext(name)[0]) + "_Acronym_List.docx")
                V.write_docx([(x["term"], x["definition"]) for x in rows], p)
                written.append(p)
        except PermissionError:
            self.host.status("A list with that name is open in Word. Close it, then export again.")
            return None
        diag.activity("volume_lists_exported", files=len(written))
        n_open = sum(len(open_items(rows)) for name, rows in lists_of(r, self.current_mode()) if name != MASTER)
        self.host.toast(f"{len(written)} acronym list(s) written" + (f"; {n_open} entries need a definition or approval" if n_open else ""))
        return written

    def export_excel(self, path=None):
        r = self.report
        if not r:
            return None
        if path is None:
            path, _ = QFileDialog.getSaveFileName(self, "Acronym lists", "Acronym_Lists.xlsx", "Excel (*.xlsx)")
            if not path:
                return None
        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        for name, rows in lists_of(r, self.current_mode()):
            ws = wb.create_sheet(_sheet(name, wb))
            ws.append(["Term", "Definition", "Status"])
            for c in ws[1]:
                c.font = Font(bold=True)
            for x in rows:
                ws.append([x["term"], x["definition"], x["status"]])
            ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 70; ws.column_dimensions["C"].width = 18
        ws = wb.create_sheet("Left out")
        ws.append(["Term", "Definition", "Status", "Why"])
        for x in r.get("unused", []):
            ws.append([x["term"], x.get("definition", ""), x.get("status", ""), "Not used in any submission volume"])
        try:
            wb.save(path)
        except PermissionError:
            self.host.status("That file is open in Excel. Close it, then export again.")
            return None
        self.host.toast("Acronym lists written")
        return path

    def record(self):
        s = getattr(self.host.acronym_page, "session", None)
        if s is None or not self.report:
            return 0
        n = s.qa_record_use()
        self.host.toast(f"Used in volumes and First use updated for {n} terms")
        return n
