"""Compliance matrix (2.8.0), C1 extract and triage + C2 outline and mapping. Rules only: no language model.

Pipeline: the app's readers give [(location label, text)] per RFP file -> lines rejoined (PDF line wraps) -> section
numbers tracked (L.4.2.1, C.3.4, 3.2.1, "SECTION L") -> sentences and list items -> candidates by trigger words
(shall / must / will / should / may / is required to, imperatives, evaluation and format language) -> rule-based
actor, type, strength, cross-references and a Keep / Drop / Review suggestion. The person decides.

Items keep file, location, section and paragraph text, so every row cites its source. The matrix lives in one Excel
workbook (atomic save, re-read on open, edits made in Excel count). Nothing leaves the PC.
"""
import datetime, hashlib, json, os, re, tempfile

# 2.50.0: characters a PDF can hand over that Excel and Word refuse (U+FFFE is PDFium's hyphen at a line break)
_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def clean_chars(s):
    """Text safe for Excel and Word: PDFium's line-break hyphen becomes a hyphen, control characters go."""
    if not isinstance(s, str):
        return s
    return _BAD_XML.sub(lambda m: "-" if m.group(0) == "\ufffe" else "", s) if _BAD_XML.search(s) else s


# ------------------------------------------------------------------ vocabulary
MODALS = r"shall not|must not|will not|may not|shall|must|will|should|may|is required to|are required to|is to|are to|is expected to|are expected to"
_MODAL = re.compile(r"\b(" + MODALS + r")\b", re.I)
IMPERATIVES = ("describe", "provide", "submit", "identify", "demonstrate", "discuss", "explain", "include", "address",
               "propose", "deliver", "furnish", "specify", "list", "present", "show", "ensure", "maintain", "comply",
               "state", "indicate", "define", "detail", "outline", "summarize", "prepare", "document", "certify",
               "complete", "sign", "attach", "use", "limit", "organize", "label", "number", "format", "do not")
_IMPER = re.compile(r"^(?:\(?[a-z0-9ivx]{1,4}[.)]\s*)?(" + "|".join(IMPERATIVES) + r")\b", re.I)
_EVAL = re.compile(r"\b(will|shall)\s+(?:be\s+)?(evaluat\w*|assess\w*|consider\w*|rat(?:e|ed)|review\w*|scor\w*)\b|"
                   r"\bevaluation (?:factor|criteri\w+|subfactor)|\bbasis (?:for|of) (?:award|evaluation)\b|"
                   r"\b(?:strength|weakness|deficienc\w+|acceptab\w+|unacceptab\w+)\b.*\b(?:assign|assess|rat)", re.I)
_FORMAT = re.compile(r"\b(page limit|pages?\b(?=.*\b(?:limit|exceed|maximum|not to exceed|no more than))|font|point size|"
                     r"\d+[- ]point|\d+(?:\.\d+)?[- ]?inch(?:es)? margins?|margins? of \d|single[- ]spaced|double[- ]spaced|8[.,]5\s*(?:x|by)\s*11|letter[- ]size|"
                     r"due (?:date|no later)|no later than|\bNLT\b|electronic copy|hard ?cop(?:y|ies)|\bcopies\b|"
                     r"file format|\.pdf\b|\bPDF\b|searchable|file size|naming convention|cover (?:page|letter)|"
                     r"table of contents|tab(?:bed)? divider|binder|page numbers?)", re.I)
OFFEROR = r"offerors?|proposers?|bidders?|vendors?|quoters?|respondents?|you|your proposal|the proposal|proposals|each volume|the volume|the response|responses|the submission|submissions|the quote|the bid"
CONTRACTOR = r"contractors?|the contractor|subcontractors?|the seller|the supplier|the developer|the provider"
GOVERNMENT = r"government|contracting officer|\bco\b|\bpco\b|\baco\b|agency|nasa|usaf|ussf|air force|space force|army|navy|dod|department|evaluators?|evaluation (?:team|board|panel)|source selection|\bssa\b|\bssEB\b|the panel|the board"
_OFF, _CON, _GOV = (re.compile(r"\b(" + p + r")\b", re.I) for p in (OFFEROR, CONTRACTOR, GOVERNMENT))
_XREF = re.compile(r"\b(?:Section|Sec\.|Paragraph|Para\.?|¶)\s+[A-Z]?(?:\.?\d+)+(?:\.\d+)*\b|"
                   r"\b[A-M]\.\d+(?:\.\d+)*\b|"
                   r"\bSection\s+[A-M]\b|"
                   r"\bAttachment\s+(?:J[- ]?)?[A-Z0-9]+(?:[- ][A-Z0-9]+)?\b|\bExhibit\s+[A-Z0-9]+\b|\bAppendix\s+[A-Z0-9]+\b|"
                   r"\bCDRL\s+[A-Z]?\d{2,4}\b|\bDID\s+[A-Z]+-[A-Z]+-\d+\b|\bDI-[A-Z]+-\d+[A-Z]?\b|"
                   r"\b(?:FAR|DFARS|NFS)\s+\d{1,3}\.\d{2,4}(?:-\d+)?\b|\b(?:MIL|DOD|NASA)-(?:STD|HDBK|SPEC)-\d+[A-Z]?\b|"
                   r"\b(?:SOW|PWS|SOO)\s+(?:paragraph\s+|para\.?\s+|section\s+)?\d+(?:\.\d+)*\b", re.I)
_ABBR = ("e.g.", "i.e.", "etc.", "U.S.", "No.", "Inc.", "vs.", "para.", "Para.", "Fig.", "fig.", "approx.", "Sec.", "sec.",
         "Corp.", "Co.", "Ltd.", "a.m.", "p.m.", "A.M.", "P.M.", "St.", "Mr.", "Ms.", "Dr.", "Jr.", "Sr.", "ea.", "Vol.", "vol.", "Rev.", "rev.", "al.", "Ref.")

# Section headings: "L.4.2.1 Technical Volume", "SECTION L - INSTRUCTIONS", "3.2.1 Thermal", "C.3 SCOPE"
_SEC_LETTER = re.compile(r"^\s*(?:SECTION\s+)?([A-M])\.((?:\d+\.)*\d+)\.?(?:\s+|$)(.*)$")
_SEC_WORD = re.compile(r"^\s*SECTION\s+([A-M])\b\s*[-–—:.]?\s*(.*)$", re.I)
_SEC_NUM = re.compile(r"^\s*((?:\d+\.)+\d+|\d+\.)\s+(.*)$")
_LIST = re.compile(r"^\s*(?:[•●▪◦·\-–*]|\(?[a-z]\)|\(?[ivx]{1,4}\)|\(?\d{1,2}\)|[a-z]\.(?=\s)|\(\d{1,2}\))\s+")
_INLINE_ITEM = re.compile(r"(?:(?<=[:;])|(?<=;\s)|(?<=\s))\(([a-z]|[ivx]{1,4}|\d{1,2})\)\s+")
_STOPW = set("""a an and or of the to for in on by with at as from is are be been shall must will should may that this these those
its it their his her our your any all each such which who whom than then also not no into within upon under over per
offeror offerors contractor government proposal provide describe include""".split())

DECISIONS = ("Keep", "Context", "Customer", "Drop", "")   # 2.29.0: Context gets a home too; 2.48.0: Customer = the
#   customer's own obligation ("The Government will evaluate..."): identified, never answered, never in our response
TYPES = ("Instruction", "Evaluation", "SOW", "Format", "Information", "Requirement", "Context")
_BOILER = re.compile(r"^(?:page\s+\d+(?:\s+of\s+\d+)?|\d{1,4}|-\s*\d+\s*-|(?:un)?classified|cui|controlled unclassified information|"
                     r"for official use only|fouo|proprietary(?: information)?|source selection information|"
                     r"(?:solicitation|rfp)\s*(?:no\.?|number)?\s*[:#]?\s*[A-Z0-9-]+|draft|this page intentionally left blank)\.?$", re.I)


def context_item(text, lead, cur, label=""):
    """A sentence that is not a requirement (2.29.0): every sentence of the RFP becomes a row, so every one gets a
    decision. Repeated headers, footers, page numbers and markings are suggested Drop; the rest Context."""
    t = " ".join(text.split())
    if not re.search(r"[A-Za-z0-9]", t):           # 2.51.0: a phone number, a table cell "25" or "I" is text too
        return None
    boiler = bool(_BOILER.match(t)) or bool(re.search(r"\((?:Header|Footer)\b", label or ""))
    return {"text": t, "lead": lead if lead and lead != t else "", "section": cur[0], "section_title": cur[1],
            "actor": "", "type": "Context", "strength": "", "suggest": "Drop" if boiler else "Context", "modal": "",
            "xrefs": "; ".join(dict.fromkeys(m.group(0).strip() for m in _XREF.finditer(t) if m.group(0).strip() != cur[0])),
            "compound": False}


# ------------------------------------------------------------------ reading
def _join_lines(text, paragraphs=False):
    """PDF text wraps mid-sentence: join a line to the next unless it ends a sentence, is a heading or a list item.
    2.51.0: paragraphs=True (Word, PowerPoint, Excel): every line is a paragraph of its own and is never joined to the
    next ("Font: 12 point" and "Margins: 1 inch" stay two rows)."""
    out = []
    lens = sorted(len(x.strip()) for x in text.splitlines() if x.strip())
    full = lens[int(len(lens) * 0.8)] if len(lens) >= 5 else 0      # a full line's length on this page
    prev_raw = ""
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            out.append("")
            prev_raw = ""
            continue
        if paragraphs:
            out.append(line)
            continue
        # 2.51.0: a short line followed by one starting with a capital ends its paragraph ("Font: 12 point Times New
        # Roman" then "Margins: 1 inch"), even without a period
        short_end = bool(full) and bool(prev_raw) and len(prev_raw) < 0.6 * full and line[:1].isupper()
        prev_raw = line
        if short_end:
            out.append(line)
            continue
        starts_block = bool(_LIST.match(line) or _SEC_LETTER.match(line) or _SEC_WORD.match(line) or
                            (_SEC_NUM.match(line) and len(line) < 140))
        if out and out[-1] and not starts_block and not re.search(r"[.:;?!]\s*$", out[-1]) and not _is_heading(out[-1]):
            out[-1] = out[-1] + " " + line
        else:
            out.append(line)
    return out


def _is_heading(line):
    s = line.strip()
    if len(s) > 120 or s.endswith((".", ";", ",")) and not re.match(r"^[A-M]\.\d", s):
        return False
    if _SEC_WORD.match(s):
        return True
    m = _SEC_LETTER.match(s) or _SEC_NUM.match(s)
    if m:
        rest = m.groups()[-1]
        return not _MODAL.search(rest) and len(rest.split()) <= 14 and not rest.endswith(".")
    letters = [c for c in s if c.isalpha()]
    return bool(letters) and len(s.split()) <= 10 and sum(c.isupper() for c in letters) / len(letters) > 0.8


def _label_section(label):
    """Word passages carry their heading in the label: "File.docx (section 'L.4 Technical', para 3-9)"."""
    m = re.search(r"section '([^']*)'", label)
    return m.group(1) if m else ""


def _page(label):
    m = re.search(r"\(p\.(\d+)", label)
    return int(m.group(1)) if m else None


def _sentences(par):
    """Split a paragraph into sentences without breaking at abbreviations, decimals or section numbers."""
    protected = par
    for a in _ABBR:
        protected = protected.replace(a, a.replace(".", "․"))
    parts = re.split(r"(?<=[.!?])\s+(?=[\"'(\[]?[A-Z0-9])", protected)
    return [p.replace("․", ".").strip() for p in parts if p.strip()]


def _split_inline_list(sentence):
    """"The Offeror shall: (a) provide X; (b) describe Y; and (c) list Z." -> lead + items."""
    if ":" not in sentence:
        return sentence, []
    lead, rest = sentence.split(":", 1)
    marks = list(_INLINE_ITEM.finditer(" " + rest))
    if len(marks) < 2:
        return sentence, []
    items = []
    s = " " + rest
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(s)
        t = s[m.start():end].strip().rstrip(";,").strip()
        t = re.sub(r"[;,]?\s+and$|[;,]?\s+or$", "", t).strip()
        items.append(t)
    return lead.strip() + ":", items


def _section_from_line(line, cur):
    """(section, title, remaining text) when the line starts a numbered section or paragraph, else None."""
    m = _SEC_LETTER.match(line)
    if m:
        sec, rest = f"{m.group(1)}.{m.group(2)}", m.group(3).strip()
        if _is_heading(line):
            return sec, rest, ""
        parent_title = cur[1] if cur[0] and sec.startswith(cur[0]) else ""
        return sec, parent_title, rest                 # numbered paragraph: "L.4.2.1 The Offeror shall ..."
    m = _SEC_WORD.match(line)
    if m and _is_heading(line):
        return m.group(1).upper(), m.group(2).strip(), ""
    m = _SEC_NUM.match(line)
    if m:
        num, rest = m.group(1).rstrip("."), m.group(2).strip()
        if _is_heading(line):
            return num, rest, ""
        return num, (cur[1] if cur[0] and num.startswith(cur[0]) else ""), rest
    return None


_FURNITURE = re.compile(r"\((?:Header|Footer)\b[^()]*\)$")


def read_items(chunks, file_name, file_kind="", heads_out=None, every=True):
    """chunks: [(label, text)] from one file. file_kind: '' or a hint from the file name (Instruction / Evaluation / SOW).
    Returns items (dicts): the requirement candidates and, with every=True (2.29.0), every other sentence as Context."""
    items = []
    cur = ("", "")                                   # (section number, title)
    heads = {}                                       # heading number -> title, for numbered paragraphs under it

    def titled(sec, title):
        if title:
            heads[sec] = title
            if heads_out is not None and "." in sec:
                heads_out.append({"num": sec, "title": title, "file": file_name})
            return sec, title
        best = max((k for k in heads if sec == k or sec.startswith(k + ".")), key=len, default="")
        return sec, heads.get(best, "")
    for label, text in chunks:
        if _FURNITURE.search(label or ""):
            continue                    # 2.57.0: headers and footers are page furniture, not text of the review
        if not text and not _label_section(label):
            continue
        # 2.40.0: a table cell, header, footer or footnote is never a heading (a short "CLIN 0001" cell or a page header
        # must stay text to assign, and must not rename the section of what follows)
        plain = bool(re.search(r"[(,]\s*(?:table|Header|Footer|Footnotes|Endnotes|text box)\b", label or "", re.I))
        lab_sec = _label_section(label)
        if lab_sec:
            got = _section_from_line(lab_sec, cur)
            bare = re.match(r"^\s*(\d{1,2})\.?\s+([A-Za-z].{0,90})$", lab_sec)    # 2.57.0: "3 Program Management"
            if got:
                cur = titled(got[0], got[1] or lab_sec)
            elif bare:
                cur = titled(bare.group(1), bare.group(2).strip())
            elif _is_heading(lab_sec) or label.endswith(", heading)"):
                cur = (cur[0], lab_sec)
        if not text:
            continue
        lead_line = ""
        all_lines = _join_lines(text, paragraphs=file_name.lower().endswith((".docx", ".docm", ".pptx", ".pptm", ".xlsx", ".xlsm")))
        for li, line in enumerate(all_lines):
            # 2.51.0: a line ending in ':' is a lead-in only when list items follow it ("Submit questions to:" followed
            # by a name and an e-mail address is text to keep)
            nxt = next((x for x in all_lines[li + 1:] if x), "")
            lead_in_ok = bool(nxt) and bool(_LIST.match(nxt))
            if not line:
                lead_line = ""
                continue
            got = None if plain else _section_from_line(line, cur)
            if got:
                cur = titled(got[0], got[1] if not got[2] else "")
                line = got[2]
                if not line:
                    continue
            elif not plain and _is_heading(line):
                cur = (cur[0], line.strip())
                if every and not (_SEC_WORD.match(line.strip()) or _SEC_LETTER.match(line.strip()) or _SEC_NUM.match(line.strip())):
                    # 2.40.0: an unnumbered capitalized line ("NOTE TO OFFERORS: SUBMIT QUESTIONS EARLY") is also a
                    # sentence to assign, so no text is only a title
                    it = context_item(line.strip(), "", cur, label)
                    if it:
                        it.update({"type": "Heading", "file": file_name, "loc": label, "page": _page(label), "context": line.strip()[:700]})
                        items.append(it)
                continue
            is_item = bool(_LIST.match(line))
            body = _LIST.sub("", line, count=1) if is_item else line
            if is_item:
                body = re.sub(r"\s*[;,]\s*(?:and|or)?\s*$", "", body).strip()
            sents = [body] if is_item else _sentences(body)
            for s in sents:
                lead, sub = _split_inline_list(s)
                group = [(x, lead) for x in sub] if sub else [(s, lead_line if is_item else "")]
                if sub and lead.strip(" :"):                 # 2.70.0: an inline list's lead-in is a sentence too
                    li_it = classify(lead, "", cur, file_kind) or (context_item(lead, "", cur, label) if every else None)
                    if li_it:
                        li_it.update({"file": file_name, "loc": label, "page": _page(label), "context": body[:700], "lead_in": True})
                        items.append(li_it)
                for txt, ld in group:
                    lead_in = txt.rstrip().endswith(":") and not sub and lead_in_ok and s is sents[-1]
                    # 2.70.0: a lead-in ("the Offeror shall address the following:") is a sentence of its own, to
                    # assign like any other (it used to live only in its items, so the page had no box for it); its
                    # items still carry it in the matrix
                    it = classify(txt, ld, cur, file_kind)
                    if it is None and every:
                        it = context_item(txt, ld, cur, label)
                    if it:
                        it.update({"file": file_name, "loc": label, "page": _page(label), "context": body[:700]})
                        if lead_in:
                            it["lead_in"] = True
                        items.append(it)
            if line.rstrip().endswith(":") and (lead_in_ok or is_item):
                lead_line = (sents[-1] if sents else body) if not is_item else lead_line
            elif not is_item:
                lead_line = ""
    return items


def file_kind(name):
    n = name.lower().replace("_", " ")
    if re.search(r"\bsection l\b|instruction|\bl\b(?= )|proposal prep", n):
        return "Instruction"
    if re.search(r"\bsection m\b|evaluation|basis (?:for|of) award", n):
        return "Evaluation"
    if re.search(r"\bsow\b|\bpws\b|\bsoo\b|statement of (?:work|objectives)|performance work|\bsection c\b", n):
        return "SOW"
    return ""


# ------------------------------------------------------------------ classifying one sentence
# 2.50.0: compliance words the team adds (Compliance words...) also make a sentence a requirement candidate
_EXTRA = {"words": (), "re": None}


def set_extra_words(words):
    """The added compliance words (the ones not in the built-in list); [] turns them off."""
    ws = tuple(sorted({" ".join(str(w).split()).lower() for w in (words or ()) if str(w).strip()}, key=len, reverse=True))
    if ws != _EXTRA["words"]:
        _EXTRA["words"] = ws
        _EXTRA["re"] = re.compile(r"(?<![\w-])(" + "|".join(re.escape(w).replace(r"\ ", r"\s+") for w in ws) + r")(?![\w-])",
                                  re.I) if ws else None
    return list(ws)


def sync_extra_words(cfg=None):
    """Read the added words from the compliance words list on this PC (cwords)."""
    try:
        import cwords
        if cfg is None and not cwords.path():
            return list(_EXTRA["words"])          # no word list on this PC (tests, tools): keep what was set
        cfg = cfg or cwords.load()
        builtin = set(cwords.DEFAULT_WORDS) | set(MODALS.split("|"))
        return set_extra_words([w for w in cfg["words"] if w not in builtin])
    except Exception:
        return set_extra_words([])


def classify(text, lead, cur, kind_hint=""):
    t = " ".join(text.split())
    if len(t) < 12 or len(t) > 1500:
        return None
    probe = (lead + " " + t) if lead and not _MODAL.search(t) and not _IMPER.match(t) else t
    modal = _MODAL.search(probe)
    imper = _IMPER.match(t)
    ev = _EVAL.search(probe)
    fm = _FORMAT.search(t)
    xw = _EXTRA["re"].search(probe) if _EXTRA["re"] is not None else None
    if not (modal or imper or ev or xw or (fm and re.search(r"\d", t))):
        return None
    word = modal.group(1).lower() if modal else ""
    # actor: the subject right before the modal; imperatives speak to the offeror
    actor = "Unclear"
    if modal:
        before = probe[:modal.start()][-90:]
        hits = [(m.start(), "Offeror") for m in _OFF.finditer(before)] + \
               [(m.start(), "Contractor") for m in _CON.finditer(before)] + \
               [(m.start(), "Government") for m in _GOV.finditer(before)]
        if hits:
            actor = max(hits)[1]
        elif imper:
            actor = "Offeror"
    elif imper:
        actor = "Offeror"
    sec = cur[0]
    letter = sec.split(".")[0] if sec and sec[0].isalpha() else ""
    kind = {"L": "Instruction", "M": "Evaluation", "C": "SOW"}.get(letter, kind_hint)
    if ev or (kind == "Evaluation" and actor == "Government"):
        typ = "Evaluation"
    elif fm and actor != "Government":
        typ = "Format"
    elif actor == "Government":
        typ = "Information"
    elif kind:
        typ = kind
    elif actor == "Contractor":
        typ = "SOW"
    elif actor == "Offeror" or imper:
        typ = "Instruction"
    else:
        typ = "Requirement"
    if typ == "Evaluation":
        strength = "Evaluation"
    elif actor == "Government":
        strength = "Information"
    elif word in ("shall", "must", "shall not", "must not", "is required to", "are required to", "is to", "are to", "may not", "will not") or \
            (imper and not word) or (word == "will" and actor != "Government"):
        strength = "Mandatory"
    elif word in ("should", "is expected to", "are expected to"):
        strength = "Desired"
    elif word == "may":
        strength = "Optional"
    elif actor == "Government":
        strength = "Information"
    else:
        strength = "Mandatory" if (fm or xw) else "Check"
    if typ == "Information" or (actor == "Government" and typ != "Evaluation"):
        suggest = "Drop"
    elif strength == "Optional" or actor == "Unclear" and not fm and not xw and typ == "Requirement":
        suggest = "Review"
    else:
        suggest = "Keep"
    side = [m for m in _MODAL.finditer(t) if m.group(1).lower() in ("shall", "must", "will", "should")]
    compound = len(side) >= 2 or bool(re.search(r"\band\s+(?:shall|must|will)\b", t, re.I)) or \
        (bool(imper) and len(re.findall(r"[;]", t)) >= 1)
    return {"text": t, "lead": lead if lead and lead != t else "", "section": sec, "section_title": cur[1],
            "actor": actor, "type": typ, "strength": strength, "suggest": suggest, "modal": word or (imper.group(1).lower() if imper else "") or (" ".join(xw.group(1).lower().split()) if xw else ""),
            "xrefs": "; ".join(dict.fromkeys(m.group(0).strip() for m in _XREF.finditer(t) if m.group(0).strip() != sec)),
            "compound": compound}


def split_text(text):
    """Split a compound requirement at '; ' or at 'and shall / must / will' into its parts."""
    parts = re.split(r";\s+(?:and\s+)?|\s*,?\s+and\s+(?=(?:shall|must|will|should)\b)", text)
    parts = [p.strip().rstrip(".;,") for p in parts if p and len(p.strip()) > 3]
    if len(parts) < 2:
        return [text]
    subj = re.match(r"^(.*?\b(?:shall|must|will|should))\b", parts[0], re.I)
    head = subj.group(1).rsplit(" ", 1)[0] if subj else ""
    out = [parts[0] + "."]
    for p in parts[1:]:
        if head and re.match(r"^(shall|must|will|should)\b", p, re.I):
            p = head + " " + p
        out.append(p[0].upper() + p[1:] + ".")
    return out


# ------------------------------------------------------------------ the matrix
def text_key(text):
    return hashlib.sha1(re.sub(r"\W+", " ", text.lower()).strip().encode()).hexdigest()[:12]


def assign_ids(items):
    """Stable, readable IDs: section-based (L.4.2-03) or file-based (SOW-07)."""
    counts = {}
    used = {it["id"] for it in items if it.get("id")}
    for it in items:
        if it.get("id"):
            continue
        base = it["section"] or re.sub(r"[^A-Za-z0-9]+", "", os.path.splitext(it["file"])[0])[:10] or "REQ"
        n = counts.get(base, 0)
        while True:
            n += 1
            cand = f"{base}-{n:02d}"
            if cand not in used:
                break
        counts[base] = n
        it["id"] = cand
        used.add(cand)
    return items


_RUNNING_NO = re.compile(r"\b(shall|must|will|should|required|offeror|contractor)\b", re.I)


def extract(files_chunks, headings=None, every=True):
    """files_chunks: [(file name, chunks)]. Returns items with IDs, links and suggestions.
    headings (list): filled with the RFP's numbered headings [{num, title, file}].
    every=True (2.29.0): every sentence is a row (requirements plus Context), so every sentence gets a decision."""
    sync_extra_words()
    items = []
    for name, chunks in files_chunks:
        items += read_items(chunks, name, file_kind(name), headings, every=every)
    # 2.56.1: a short line with no compliance words on two or more pages of a file (a PDF's running header or footer,
    # "Source Selection Information", the solicitation number) is kept once, like a Word header or footer
    pages_of = {}
    for it in items:
        if len(it["text"].split()) <= 10 and not _RUNNING_NO.search(it["text"]):
            m = re.search(r"\(p\.(\d+)", it.get("loc", "") or "")
            if m:
                pages_of.setdefault((it["file"], text_key(it["text"])), set()).add(m.group(1))
    running = {k for k, v in pages_of.items() if len(v) >= 2}
    seen, out = set(), []
    for it in items:
        # the same header, footer or page marking repeated on every page: keep once. 2.51.0: a sentence the RFP says
        # twice in its body (in two sections) is two rows, so nothing it says is lost
        k = (it["file"], text_key(it["text"]))
        repeat = bool(re.search(r"\((?:Header|Footer)\b|\bheader\b|\bfooter\b", it.get("loc", ""), re.I)) or bool(_BOILER.match(it["text"])) \
            or k in running
        if k in seen and repeat:
            continue
        seen.add(k)
        it.update({"key": text_key(it["text"]), "decision": "", "map": "", "owner": "", "status": "", "notes": "",
                   "response_page": "", "links": "", "new": "", "learned": ""})
        out.append(it)
    assign_ids(out)
    link_items(out)
    return out


def parse_header(text):
    """'L.4.2 Technical Approach' -> ('L.4.2', 'Technical Approach'); 'SECTION L - INSTRUCTIONS' -> ('L', 'Instructions');
    '3.1 Thermal' -> ('3.1', 'Thermal'); otherwise ('', the text)."""
    t = " ".join((text or "").split())
    m = _SEC_LETTER.match(t)
    if m:
        return f"{m.group(1)}.{m.group(2)}", m.group(3).strip(" .:-")
    m = _SEC_WORD.match(t)
    if m:
        return m.group(1).upper(), m.group(2).strip(" .:-").title() if m.group(2).isupper() else m.group(2).strip(" .:-")
    m = _SEC_NUM.match(t)
    if m:
        return m.group(1).rstrip("."), m.group(2).strip(" .:-")
    return "", t.strip(" .:")


def mark_header(items, iid, num, title):
    """2.51.0 Section header: the sentence is a heading of the RFP; the sentences after it in the same file take its
    section number and title, up to the next heading. Returns the ids renumbered."""
    i = next((k for k, it in enumerate(items) if it.get("id") == iid), None)
    if i is None:
        return []
    h = items[i]
    num, title = " ".join((num or "").split()), " ".join((title or "").split())
    h.update(type="Heading", section=num, section_title=title, decision=h.get("decision") or "Drop", header="Yes")
    if h.get("draft"):
        h["draft"] = ""
    done = []
    f = h.get("file", "")
    for it in items[i + 1:]:
        if it.get("file", "") != f:
            continue
        if it.get("type") == "Heading" or it.get("header") == "Yes":
            break
        it["section"], it["section_title"] = num, title
        done.append(it.get("id", ""))
    return done


def mark_format(items):
    """2.51.0 Format requirement: the sentences are formatting rules (font, margins, page limits, file format)."""
    for it in items:
        it["type"] = "Format"
        if it.get("strength") in ("", "Information", "Check"):
            it["strength"] = "Mandatory"
        if not it.get("decision"):
            it["decision"] = "Keep"
            it["draft"] = ""
    return len(items)


def progress(items):
    """2.52.0 how far the review is: every sentence decided (compliance requirement, support, the customer's or not
    used), per file and overall. {"total", "decided", "pct", "Keep", "Context", "Customer", "Drop", "needs_home",
    "files": [{"file", "total", "decided", "pct"}]}."""
    out = {"total": 0, "decided": 0, "Keep": 0, "Context": 0, "Customer": 0, "Drop": 0, "needs_home": 0, "files": []}
    per = {}
    for it in items:
        f = it.get("file", "") or "(no file)"
        d = per.setdefault(f, {"file": f, "total": 0, "decided": 0})
        d["total"] += 1
        out["total"] += 1
        dec = it.get("decision", "")
        if dec:
            d["decided"] += 1
            out["decided"] += 1
            out[dec] = out.get(dec, 0) + 1
            if dec in ("Keep", "Context") and not it.get("map") and not (dec == "Context" and it.get("context_for")):
                out["needs_home"] += 1
    for d in per.values():
        d["pct"] = round(100.0 * d["decided"] / d["total"]) if d["total"] else 100
        out["files"].append(d)
    out["pct"] = round(100.0 * out["decided"] / out["total"]) if out["total"] else 0
    return out


FOUND = "Found by a compliance word"
_REQ_FIELDS = ("actor", "type", "strength", "suggest", "modal", "xrefs", "compound")


def rescan_words(items, cfg=None):
    """2.50.0 Rescan with the compliance words: every sentence is checked again with the current list. A sentence not
    decided yet that now holds a compliance word becomes a requirement candidate (marked FOUND); one that was found
    only by a word since taken off the list goes back to Context. Decided sentences keep their decision; the ones
    that now hold a compliance word are counted so the person can look. Returns {"found", "back", "decided"} (ids)."""
    sync_extra_words(cfg)
    out = {"found": [], "back": [], "decided": []}
    for it in items:
        if it.get("type") == "Heading":
            continue
        new = classify(it.get("text", ""), it.get("lead", ""), (it.get("section", ""), it.get("section_title", "")),
                       file_kind(it.get("file", "")))
        was_req = it.get("type") != "Context"
        if it.get("decision") or it.get("draft"):
            if new is not None and not was_req:
                out["decided"].append(it.get("id", ""))
            continue
        if new is not None and not was_req:
            it.update({k: new[k] for k in _REQ_FIELDS})
            if not it.get("new"):
                it["new"] = FOUND
            out["found"].append(it.get("id", ""))
        elif new is None and was_req and it.get("new") == FOUND:
            ctx = context_item(it.get("text", ""), it.get("lead", ""), (it.get("section", ""), it.get("section_title", "")),
                               it.get("loc", ""))
            if ctx:
                it.update({k: ctx[k] for k in _REQ_FIELDS})
            it["new"] = ""
            out["back"].append(it.get("id", ""))
    return out


def merge_reread(old, new):
    """Re-reading the RFP (an amendment): decisions, mapping and notes carry over to identical sentences.
    New sentences are marked; sentences no longer present are kept and marked 'Not in the latest read'."""
    by_fk = {}
    for it in old:                                 # 2.51.0: a sentence said twice in a file is two rows
        by_fk.setdefault((it.get("file", ""), it["key"]), []).append(it)
    by_key = {}
    for it in old:
        by_key.setdefault(it["key"], []).append(it)
    used_ids = {it["id"] for it in old}
    out, taken = [], set()
    new_fk = {(n.get("file", ""), n["key"]) for n in new}

    def match(it):
        """The same sentence in the same file first (a sentence repeated in two files stays two rows, 2.29.0), then the
        same sentence anywhere (an amendment that renamed the file)."""
        o = next((x for x in by_fk.get((it.get("file", ""), it["key"]), []) if id(x) not in taken), None)
        if o is not None:
            return o
        return next((x for x in by_key.get(it["key"], []) if id(x) not in taken and
                     (x.get("file", ""), x["key"]) not in new_fk), None)
    for it in new:
        o = match(it)
        if o is not None:
            taken.add(id(o))
        if o:
            for f in ("id", "decision", "map", "owner", "status", "notes", "response_page", "text", "context_for", "draft", "draft_cf"):
                it[f] = o.get(f, "")
            it["new"] = ""
        else:
            it["new"] = "New in latest read"
            it["id"] = "" if it["id"] in used_ids else it["id"]
        out.append(it)
    for o in old:
        if id(o) not in taken:
            o["new"] = "Not in the latest read"
            out.append(o)
    assign_ids(out)
    link_items(out)
    return out


def _words(s):
    return {w[:-1] if len(w) > 4 and w.endswith("s") else w for w in re.findall(r"[a-z][a-z0-9\-]{2,}", s.lower())} - _STOPW


def link_items(items):
    """Link Evaluation (M) and SOW items to the Instruction items they concern: explicit section cross-references,
    then strong word overlap (Jaccard >= 0.3, best match only)."""
    instr = [it for it in items if it["type"] in ("Instruction", "Format")]
    for it in items:
        links = []
        refs = [r.strip() for r in it.get("xrefs", "").split(";") if r.strip()]
        for r in refs:
            num = re.sub(r"^(?:section|sec\.|paragraph|para\.?|¶)\s+", "", r, flags=re.I)
            if "." not in num:
                continue                                # "Section C" names a whole section: too broad to link
            for o in items:
                if o is not it and o["section"] and (o["section"] == num or o["section"].startswith(num + ".")):
                    links.append(o["id"])
        if it["type"] in ("Evaluation", "SOW") and not links and instr:
            w = _words(it["text"])
            best, score = None, 0.0
            for o in instr:
                ow = _words(o["text"])
                if w and ow:
                    j = len(w & ow) / len(w | ow)
                    if j > score:
                        best, score = o, j
            if best is not None and score >= 0.3:
                links.append(best["id"])
        it["links"] = "; ".join(dict.fromkeys(links))[:300]
    return items


def merge_items(items, ids):
    """Merge rows (in list order) into the first: texts joined, the rest removed. Returns the merged row or None."""
    rows = [it for it in items if it["id"] in ids]
    if len(rows) < 2:
        return None
    first = rows[0]
    if all(r.get("place") for r in rows):          # 2.57.0: the combined sentence keeps each part's place on the page
        import pagesent as PS
        first["place"] = PS.merge_places(rows)
    else:
        first["place"] = ""
    first["text"] = " ".join(r["text"].rstrip() for r in rows)
    first["key"] = text_key(first["text"])
    first["xrefs"] = "; ".join(dict.fromkeys(x for r in rows for x in r.get("xrefs", "").split("; ") if x))
    first["compound"] = True
    first["notes"] = ((first.get("notes") or "") + "; " if first.get("notes") else "") + "merged " + ", ".join(r["id"] for r in rows[1:])
    for r in rows[1:]:
        items.remove(r)
        remap_context_for(items, r["id"], first["id"])
    return first


# ------------------------------------------------------------------ Context linked to its requirement (2.33.0)
# 2.35.0: a Context sentence can explain several requirements ("Context for" holds IDs separated by "; "), and every
# decision is first a draft. The committed decision ("decision", "context_for") is what Place, the setup documents, the
# exports and the coverage report use; the draft ("draft", "draft_cf") is what Decide shows until it is committed.
CONTEXT_REACH = 12          # how many sentences before or after to look for the requirement
UNDECIDED = "Undecided"     # a draft that takes a decision back


def cf_ids(value):
    """'L.4-01; L.5-02' -> ['L.4-01', 'L.5-02']."""
    return [x.strip() for x in re.split(r"[;,]", value or "") if x.strip()]


def join_ids(ids):
    return "; ".join(dict.fromkeys(i for i in ids if i))


def eff(it):
    """The decision Decide shows: the draft when there is one, else the committed decision."""
    d = it.get("draft", "")
    return ("" if d == UNDECIDED else d) if d else it.get("decision", "")


def eff_cf(it):
    return it.get("draft_cf", "") if it.get("draft") else it.get("context_for", "")


def set_draft(it, decision, cf=None):
    """Pencil a decision (and for Context its requirements). Back to the committed state clears the draft.
    Returns True when the row changed."""
    decision = decision or ""
    if decision == "Context":
        new_cf = join_ids(cf) if cf is not None else (eff_cf(it) if eff(it) == "Context" else "")
    else:
        new_cf = ""
    before = (it.get("draft", ""), it.get("draft_cf", ""))
    committed_cf = it.get("context_for", "") if it.get("decision") == "Context" else ""
    if decision == it.get("decision", "") and new_cf == committed_cf:
        it["draft"], it["draft_cf"] = "", ""
    else:
        it["draft"], it["draft_cf"] = decision or UNDECIDED, new_cf
    return (it["draft"], it["draft_cf"]) != before


def drafts(items):
    return [it for it in items if it.get("draft")]


def commit(rows):
    """Make the drafts in rows count. Returns {"n", "Keep", "Context", "Drop", "Undecided", "links"}."""
    out = {"n": 0, "Keep": 0, "Context": 0, "Drop": 0, "Undecided": 0, "links": 0}
    for it in rows:
        if not it.get("draft"):
            continue
        d = eff(it)
        it["decision"] = d
        it["context_for"] = it.get("draft_cf", "") if d == "Context" else ""
        it["draft"], it["draft_cf"] = "", ""
        out["n"] += 1
        out[d or "Undecided"] = out.get(d or "Undecided", 0) + 1
        out["links"] += len(cf_ids(it["context_for"]))
    return out


def commit_text(r, left):
    """'Committed 23 decisions: 15 Keep, 6 Context with 9 links, 2 Drop. 41 drafts left.'"""
    if not r["n"]:
        return "Nothing to commit here: no drafts."
    bits = [f"{r['Keep']} Keep"] if r["Keep"] else []
    if r["Context"]:
        bits.append(f"{r['Context']} Context with {r['links']} link" + ("s" if r["links"] != 1 else ""))
    if r["Drop"]:
        bits.append(f"{r['Drop']} Drop")
    if r["Undecided"]:
        bits.append(f"{r['Undecided']} taken back to undecided")
    return (f"Committed {r['n']} decision" + ("s" if r["n"] != 1 else "") + ": " + ", ".join(bits) + ". "
            + (f"{left} draft" + ("s" if left != 1 else "") + " left." if left else "No drafts left."))


def default_context_for(items, it, dec=None):
    """The requirement a Context sentence explains: the nearest Keep before it in the same file and RFP section
    (the same paragraph), else the nearest Keep after it there. '' when there is none nearby. dec(row) gives the
    decision to look at (the committed one by default; Decide passes eff for drafts)."""
    dec = dec or (lambda x: x["decision"])
    try:
        i = next(k for k, x in enumerate(items) if x is it)
    except StopIteration:
        return ""
    same = lambda x: x.get("file", "") == it.get("file", "") and x.get("section", "") == it.get("section", "")
    for j in range(i - 1, max(-1, i - 1 - CONTEXT_REACH), -1):
        x = items[j]
        if not same(x):
            break
        if dec(x) == "Keep":
            return x["id"]
    for j in range(i + 1, min(len(items), i + 1 + CONTEXT_REACH)):
        x = items[j]
        if not same(x):
            break
        if dec(x) == "Keep":
            return x["id"]
    return ""


def link_contexts(items):
    """Give every committed Context sentence without a requirement its default one; clear the link on rows no
    longer Context. Returns the number of links made."""
    n = 0
    for it in items:
        if it["decision"] == "Context":
            if not it.get("context_for") and not it.get("map"):    # 2.36.0: a section's context keeps no link
                cf = default_context_for(items, it)
                if cf:
                    it["context_for"] = cf
                    n += 1
        elif it.get("context_for"):
            it["context_for"] = ""
    return n


# ------------------------------------------------------------------ Assign (2.36.0): one step, tied to the structure
def assign_requirement(items, rows, label):
    """Rows become compliance requirements in the outline section label. Context that goes with them (first link)
    follows them to the new section."""
    for it in rows:
        old = it.get("map", "")
        it.update(decision="Keep", map=label, context_for="", draft="", draft_cf="")
        for x in items:
            links = cf_ids(x.get("context_for", ""))
            if x.get("decision") == "Context" and links and links[0] == it["id"] and x.get("map", "") in ("", old):
                x["map"] = label
    return len(rows)


def assign_context(items, rows, req_ids=None, section=None):
    """Rows become Context: for the requirements req_ids (their section follows the first one), or for the outline
    section itself (no requirement)."""
    by = {it["id"]: it for it in items}
    for it in rows:
        it.update(decision="Context", draft="", draft_cf="")
        if section:
            it.update(context_for="", map=section)
        else:
            links = [r for r in (req_ids or []) if r != it["id"] and by.get(r, {}).get("decision") == "Keep"]
            it["context_for"] = join_ids(links)
            first = by.get(links[0]) if links else None
            it["map"] = first.get("map", "") if first is not None else ""
    return len(rows)


def assign_customer(items, rows):
    """2.48.0: the customer's own obligation ("The Government will evaluate..."): identified, never answered. No section,
    no links; Context that pointed at these sentences loses the link."""
    clear_assignment(items, rows)
    for it in rows:
        it["decision"] = "Customer"
    return len(rows)


def clear_assignment(items, rows):
    """Not used: no decision, no section, no links; Context that pointed at these requirements loses the link."""
    ids = {it["id"] for it in rows}
    for it in rows:
        it.update(decision="", map="", context_for="", draft="", draft_cf="")
    for x in items:
        links = cf_ids(x.get("context_for", ""))
        if links and ids & set(links):
            x["context_for"] = join_ids(i for i in links if i not in ids)
    return len(rows)


def remap_context_for(items, old_id, new_id):
    """A requirement was split or merged: its Context sentences (committed and drafted) follow the new ID."""
    for it in items:
        for f in ("context_for", "draft_cf"):
            ids = cf_ids(it.get(f, ""))
            if old_id in ids:
                it[f] = join_ids(new_id if i == old_id else i for i in ids)


def context_problems(items):
    """[(item, reason)] for committed Context sentences that are not with their requirements."""
    by = {it["id"]: it for it in items}
    out = []
    for it in items:
        if it["decision"] != "Context":
            continue
        ids = cf_ids(it.get("context_for", ""))
        if not ids:
            if not it.get("map"):                     # 2.36.0: context for a section needs only the section
                out.append((it, "not assigned to a requirement or a section"))
            continue
        bad = False
        for cf in ids:
            req = by.get(cf)
            if req is None:
                out.append((it, f"its requirement {cf} is no longer in the matrix")); bad = True
            elif req["decision"] != "Keep":
                out.append((it, f"its requirement {cf} is no longer Keep")); bad = True
        if bad or not it.get("map"):
            continue
        maps = [by[cf].get("map") for cf in ids if by.get(cf) is not None and by[cf].get("map")]
        if maps and it["map"] not in maps:
            out.append((it, f"placed in {it['map']}, its requirement" + ("s" if len(ids) > 1 else "") + f" in {', '.join(dict.fromkeys(maps))}"))
    return out


def _word_diff(old, new):
    """Plain-text word diff: [-removed-] {+added+}."""
    import difflib
    a, b = old.split(), new.split()
    out = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if op == "equal":
            out.append(" ".join(a[i1:i2]))
        if op in ("delete", "replace"):
            out.append("[-" + " ".join(a[i1:i2]) + "-]")
        if op in ("insert", "replace"):
            out.append("{+" + " ".join(b[j1:j2]) + "+}")
    return " ".join(out)


def amendment_changes(items, min_ratio=0.6):
    """After a re-read: pair each new sentence with the most similar removed one (same file kind of change, ratio
    >= 0.6) as Changed; the rest are Added or Removed. Returns [{kind, id, old_id, old, new, diff, section}]."""
    import difflib
    new = [it for it in items if it.get("new") == "New in latest read"]
    gone = [it for it in items if it.get("new") == "Not in the latest read"]
    used, out = set(), []
    for n in new:
        best, score = None, 0.0
        for g in gone:
            if g["id"] in used:
                continue
            r = difflib.SequenceMatcher(None, g["text"], n["text"], autojunk=False).ratio()
            if r > score:
                best, score = g, r
        if best is not None and score >= min_ratio:
            used.add(best["id"])
            out.append({"kind": "Changed", "id": n["id"], "old_id": best["id"], "old": best["text"], "new": n["text"],
                        "diff": _word_diff(best["text"], n["text"]), "section": n["section"] or best["section"],
                        "decision": best.get("decision", "")})
        else:
            out.append({"kind": "Added", "id": n["id"], "old_id": "", "old": "", "new": n["text"], "diff": "", "section": n["section"], "decision": ""})
    for g in gone:
        if g["id"] not in used:
            out.append({"kind": "Removed", "id": "", "old_id": g["id"], "old": g["text"], "new": "", "diff": "", "section": g["section"],
                        "decision": g.get("decision", "")})
    order = {"Changed": 0, "Added": 1, "Removed": 2}
    out.sort(key=lambda c: (order[c["kind"]], sec_key(c["section"] or "~")))
    return out


def write_changes_excel(changes, path, name=""):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "Amendment changes"
    ws.append(["Change", "New ID", "Old ID", "Section", "Old text", "New text", "Word differences", "Earlier decision"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for ch in changes:
        ws.append([ch["kind"], ch["id"], ch["old_id"], ch["section"], ch["old"], ch["new"], ch["diff"], ch["decision"]])
    for col, w in zip("ABCDEFGH", (10, 12, 12, 10, 60, 60, 70, 14)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    wb.save(path)


# ------------------------------------------------------------------ outline (C2)
_OUTLINE_LINE = re.compile(r"^(\s*)((?:Volume|Vol\.?)\s+[IVX\d]+|[A-M](?:\.\d+)+|\d+(?:\.\d+)*|[A-M]|[IVX]+)?[.)]?\s*[-–—:]?\s*(.*)$", re.I)


def parse_outline(text):
    """One section per line: '1.2 Management Approach', 'Volume I - Technical', or indented titles.
    Returns [{"num", "title", "level"}]."""
    out = []
    for raw in text.splitlines():
        if not raw.strip():
            continue
        line = raw.replace("\t", "    ")
        m = _OUTLINE_LINE.match(line)
        indent, num, title = len(m.group(1)), (m.group(2) or "").strip(), m.group(3).strip()
        if num and not title:
            title, num = num, ""
        if num.lower().startswith(("volume", "vol")):
            level = 1
        elif indent:                                 # 2.69.0: indentation says the level (sections under a volume)
            level = indent // 4 + 1
        elif num:
            level = num.count(".") + 1 if num[0].isdigit() else num.count(".") + 1
        else:
            level = indent // 4 + 1
        out.append({"num": num, "title": title or num, "level": max(1, min(level, 12))})     # 2.68.0: deeper than 6
    return out


def outline_text(outline):
    return "\n".join(("    " * (o["level"] - 1)) + (f"{o['num']} " if o["num"] else "") + o["title"] for o in outline)


def outline_label(o):
    return f"{o['num']} {o['title']}".strip()


def sec_key(num):
    """L.4.10 after L.4.2; letters before numbers."""
    return [(0, p) if not p.isdigit() else (1, int(p)) for p in re.split(r"[.]", num)]


def outline_from_headings(headings, letter="L"):
    """The RFP's numbered headings in one section (Section L by default), sorted by number: a starting outline."""
    seen, out = set(), []
    for h in sorted(headings, key=lambda h: sec_key(h["num"])):
        num = h["num"]
        if num in seen or (letter and not num.startswith(letter + ".")):
            continue
        seen.add(num)
        out.append({"num": num, "title": h["title"], "level": num.count(".") if num[0].isalpha() else num.count(".") + 1})
    return out


def outline_from_docx_xml(document_xml):
    """Headings of a Word outline (Heading 1-6 styles), from word/document.xml bytes (hardened parser)."""
    from lxml import etree
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False)
    root = etree.fromstring(document_xml, parser)
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    out = []
    for p in root.iter(W + "p"):
        st = p.find(f"{W}pPr/{W}pStyle")
        if st is None:
            continue
        m = re.match(r"^(?:Heading|heading)\s*(\d)$", st.get(W + "val", ""))
        if not m:
            continue
        import textcheck as _TC
        txt = _TC.w_text(p).strip()
        if txt:
            o = parse_outline(txt)[0] if parse_outline(txt) else {"num": "", "title": txt}
            out.append({"num": o["num"], "title": o["title"], "level": int(m.group(1))})
    return out


def word_df(items):
    """How many requirements use each word (rare words carry more meaning)."""
    df = {}
    for it in items:
        for w in _words(it["text"] + " " + it.get("lead", "")):
            df[w] = df.get(w, 0) + 1
    return df, max(1, len(items))


def _leaves(outline):
    """Sections without subsections: requirements map to the most specific place (not 'Volume I')."""
    out = []
    for i, o in enumerate(outline):
        nxt = outline[i + 1] if i + 1 < len(outline) else None
        if not (nxt and nxt["level"] > o["level"]):
            out.append(o)
    return out


def suggest_map(item, outline, stats=None):
    """Best outline section for an item: the same section number first; otherwise rare shared words, where most of
    the section title's meaning must match. Returns (label, score) or ("", 0)."""
    if not outline:
        return "", 0.0
    import math
    df, n = stats or ({}, 1)
    idf = lambda w: math.log(1 + n / (1 + df.get(w, 0)))
    sec = item["section"]
    best, score = None, 0.0
    for o in outline:
        if sec and o["num"] and (sec == o["num"] or sec.startswith(o["num"] + ".")):
            s = 2.0 + o["num"].count(".") * 0.01           # deepest matching number wins
            if s > score:
                best, score = o, s
        for ref in re.findall(r"\[([A-Z]\.[\d.]*\d)\]", o["title"]):   # "1.1 Technical Approach [L.4.1]" (2.28.0)
            if sec and (sec == ref or sec.startswith(ref + ".")):
                s = 2.5 + ref.count(".") * 0.01
                if s > score:
                    best, score = o, s
    if best is not None:
        return outline_label(best), round(score, 2)
    w = _words(item["text"] + " " + item.get("lead", ""))
    for o in _leaves(outline):
        ow = {x for x in _words(o["title"]) if df.get(x, 0) > 0 or not stats}
        hit = w & ow
        if not hit:
            continue
        if len(hit) < math.ceil(len(ow) * 0.75):
            continue                                    # most title words must appear: "Risk Management" needs both
        got = sum(idf(x) for x in hit)
        frac = got / sum(idf(x) for x in ow)
        s = frac * min(1.0, got / 2.0) if frac >= 0.6 else 0.0
        if s > score:
            best, score = o, s
    return (outline_label(best), round(score, 2)) if best is not None and score >= 0.3 else ("", 0.0)


# ------------------------------------------------------------------ the workbook
COLUMNS = ["ID", "Decision", "Suggested", "Learned", "Section", "Section title", "Requirement", "Lead-in", "Actor", "Type", "Strength",
           "Cross-references", "Linked items", "Outline section", "Context for", "Owner", "Status", "Response page", "Notes",
           "Source file", "Location", "Compound", "Change", "Draft", "Draft context for", "Context", "Key", "As read", "Place", "Lead place", "L/M"]
WIDTHS = [12, 10, 10, 10, 10, 24, 70, 30, 12, 13, 12, 22, 16, 28, 14, 14, 12, 12, 24, 26, 30, 9, 14, 10, 16, 60, 13, 40, 16, 16, 6]
FIELD = {"ID": "id", "Decision": "decision", "Suggested": "suggest", "Learned": "learned", "Section": "section", "Section title": "section_title",
         "Requirement": "text", "Lead-in": "lead", "Actor": "actor", "Type": "type", "Strength": "strength",
         "Cross-references": "xrefs", "Linked items": "links", "Outline section": "map", "Context for": "context_for",
         "Owner": "owner", "Status": "status",
         "Response page": "response_page", "Notes": "notes", "Source file": "file", "Location": "loc",
         "Compound": "compound", "Change": "new", "Draft": "draft", "Draft context for": "draft_cf", "Context": "context", "Key": "key",
         "As read": "orig_text",    # 2.55.0 Fix the text: the sentence as the reader read it (empty when not fixed)
         "Place": "place", "Lead place": "lead_place",    # 2.57.0: where the sentence is on its page ("1:117-218;2:0-30")
         "L/M": "lm"}                                     # 2.68.0: Instructions (L) or Evaluation (M), Shipley


class Locked(Exception):
    """The workbook is open in Excel."""


class Matrix:
    def __init__(self, path, name=""):
        self.path = path
        self.name = name or os.path.splitext(os.path.basename(path))[0]
        self.items, self.outline, self.sources, self.history = [], [], [], []
        self.headings = []
        self.setup = None          # Setup Builder (2.13.0): see setupdoc.py
        self.pending = False

    def load(self):
        self.items, self.outline, self.sources, self.history = [], [], [], []
        self.headings = []
        self.setup = None
        if not os.path.isfile(self.path):
            return self
        from openpyxl import load_workbook
        wb = load_workbook(self.path, read_only=True, data_only=True)
        try:
            ws = wb["Compliance Matrix"] if "Compliance Matrix" in wb.sheetnames else wb.worksheets[0]
            it = ws.iter_rows(values_only=True)
            head = [str(h or "").strip() for h in next(it, [])]
            idx = {h: i for i, h in enumerate(head)}
            for vals in it:
                vals = list(vals or [])
                if not any(v not in (None, "") for v in vals):
                    continue
                row = {}
                for col, f in FIELD.items():
                    v = vals[idx[col]] if col in idx and idx[col] < len(vals) else None
                    row[f] = "" if v is None else clean_chars(str(v)).strip()
                if not row["text"]:
                    continue
                row["compound"] = row["compound"].lower() in ("yes", "true", "1")
                row["decision"] = {"k": "Keep", "keep": "Keep", "yes": "Keep", "d": "Drop", "drop": "Drop", "no": "Drop",
                                   "c": "Context", "context": "Context", "customer": "Customer", "customer's": "Customer"}.get(row["decision"].lower(), "")
                row["draft"] = {"keep": "Keep", "context": "Context", "drop": "Drop", "undecided": UNDECIDED}.get(row["draft"].lower(), "")
                if not row["draft"]:
                    row["draft_cf"] = ""
                row["page"] = _page(row["loc"])
                row["key"] = row["key"] or text_key(row["text"])
                self.items.append(row)
            if "Outline" in wb.sheetnames:
                rows = list(wb["Outline"].iter_rows(values_only=True))[1:]
                for r in rows:
                    if r and len(r) >= 3 and (r[1] or r[0]):
                        try:
                            lvl = int(r[2] or 1)
                        except (TypeError, ValueError):
                            lvl = 1
                        o = {"num": str(r[0] or "").strip(), "title": str(r[1] or r[0]).strip(), "level": lvl}
                        if len(r) >= 5 and str(r[4] or "").strip().lower() == "auto":
                            o["auto"] = True           # 2.69.0: numbered by its place (made in Assign)
                        if len(r) >= 6 and r[5]:
                            o["from"] = str(r[5]).strip()   # 2.69.0: the RFP section it was copied from
                        self.outline.append(o)
            if "Sources" in wb.sheetnames:
                self.sources = [[("" if v is None else str(v)) for v in (list(r) + ["", "", "", ""])[:4]]
                                for r in list(wb["Sources"].iter_rows(values_only=True))[1:] if r and r[0]]
            if "RFP headings" in wb.sheetnames:
                self.headings = [{"num": str(r[0]), "title": str(r[1] or ""), "file": str(r[2] or "")}
                                 for r in list(wb["RFP headings"].iter_rows(values_only=True))[1:] if r and r[0]]
            if "History" in wb.sheetnames:
                self.history = [[("" if v is None else str(v)) for v in (list(r) + ["", "", "", ""])[:4]]
                                for r in list(wb["History"].iter_rows(values_only=True))[1:] if r and any(r)]
            if "About" in wb.sheetnames:
                for r in wb["About"].iter_rows(values_only=True):
                    if r and r[0] == "Matrix" and len(r) > 1 and r[1]:
                        self.name = str(r[1])
            if "Setup data" in wb.sheetnames:      # the Setup Builder's saved state (JSON in column A, split in parts)
                blob = "".join(str(r[0] or "") for r in wb["Setup data"].iter_rows(min_row=2, values_only=True) if r)
                try:
                    d = json.loads(blob) if blob else None
                    self.setup = d if isinstance(d, dict) else None
                except ValueError:
                    self.setup = None
        finally:
            wb.close()
        assign_ids(self.items)
        if self.setup and self.setup.get("owners"):  # 2.38.0: an owner equal to the section's lead was inherited, not set
            import setupdoc
            for it in self.items:
                if it.get("owner") and it["owner"] == setupdoc.section_people(self, it.get("map", ""))["lead"]:
                    it["owner"] = ""
        return self

    def _save_setup(self, wb, Font):
        """Setup sheet (readable: section, order, reference, text), Setup direction, and Setup data (the state)."""
        import setupdoc as SD
        S = SD.setup_of(self)
        P = SD.placed_in(self)
        B = SD.by_id(self)
        inh = SD.inherited_sections(self)
        ss = wb.create_sheet("Setup")
        ss.append(["Section", "Order", "ID", "Kind", "Reference", "Reference source", "RFP text", "Note (green)"])
        for lab, ids in P.items():
            for n, i in enumerate(ids, 1):
                it = B[i]
                ref, src = SD.ref_of(self, it, inh)
                ss.append([lab, n, i, SD.kind_of(it), ref, src, it["text"], S["notes"].get(i, "")])
        for c in ss[1]:
            c.font = Font(bold=True)
        for col, w in zip("ABCDEFGH", (34, 7, 14, 12, 14, 18, 90, 40)):
            ss.column_dimensions[col].width = w
        sd = wb.create_sheet("Setup direction")
        sd.append(["Section", "Management direction (green)"])
        for c in sd[1]:
            c.font = Font(bold=True)
        for lab in P:
            if S["mgmt"].get(lab, "").strip():
                sd.append([lab, S["mgmt"][lab]])
        sd.column_dimensions["A"].width = 34; sd.column_dimensions["B"].width = 100
        js = json.dumps(S, ensure_ascii=False, separators=(",", ":"))
        dd = wb.create_sheet("Setup data")
        dd.append(["Saved by Stationkeeper. Do not edit: change the setup in the app."])
        for k in range(0, len(js), 30000):
            dd.append([js[k:k + 30000]])
        dd.sheet_state = "hidden"

    def log(self, what, change):
        user = os.environ.get("USERNAME") or os.environ.get("USER") or "unknown"
        self.history.append([datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), user, what, change])
        self.pending = True

    def path_of(self, file_name):
        """Full path of an RFP file, from the Sources sheet (newest read first), if it is still there."""
        for row in reversed(self.sources):
            if row and row[0] == file_name and len(row) > 3 and row[3] and os.path.isfile(row[3]):
                return row[3]
        return ""

    def coverage(self):
        """Per source file (and a Total row): sentences, Keep, Context, Drop, undecided, placed, need a home (2.29.0)."""
        order = [s[0] for s in self.sources] + [it.get("file", "") for it in self.items]
        rows = {}
        for f in dict.fromkeys(order):
            rows[f] = {"file": f, "total": 0, "Keep": 0, "Context": 0, "Customer": 0, "Drop": 0, "": 0, "placed": 0, "needs_home": 0}
        for it in self.items:
            r = rows[it.get("file", "")]
            r["total"] += 1
            r[it["decision"]] = r.get(it["decision"], 0) + 1
            if it["decision"] in ("Keep", "Context"):
                r["placed" if it["map"] else "needs_home"] += 1
        out = [r for r in rows.values() if r["total"]]
        tot = {"file": "Total", **{k: sum(r[k] for r in out) for k in ("total", "Keep", "Context", "Customer", "Drop", "", "placed", "needs_home")}}
        return out + [tot]

    def counts(self):
        c = {"total": len(self.items), "Keep": 0, "Context": 0, "Customer": 0, "Drop": 0, "": 0, "mapped": 0, "kept_unmapped": 0,
             "context_mapped": 0, "context_unmapped": 0}
        for it in self.items:
            c[it["decision"]] = c.get(it["decision"], 0) + 1
            if it["decision"] == "Keep":
                c["mapped" if it["map"] else "kept_unmapped"] += 1
            elif it["decision"] == "Context":
                c["context_mapped" if it["map"] else "context_unmapped"] += 1
        c["decided"] = c["total"] - c[""]
        c["drafts"] = sum(1 for it in self.items if it.get("draft"))
        c["open"] = sum(1 for it in self.items if not eff(it))            # not decided even as a draft
        c["placed"] = c["mapped"] + c["context_mapped"]
        c["needs_home"] = c["kept_unmapped"] + c["context_unmapped"]
        return c

    def save(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.datavalidation import DataValidation
        d = os.path.dirname(os.path.abspath(self.path))
        os.makedirs(d, exist_ok=True)
        wb = Workbook()
        ws = wb.active; ws.title = "Compliance Matrix"
        ws.append(COLUMNS)
        for i, c in enumerate(ws[1], 1):
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1E3A8A")
            ws.column_dimensions[c.column_letter].width = WIDTHS[i - 1]
        fills = {"Keep": "E7F6EC", "Drop": "EEEEEE", "Customer": "F3E8FF"}
        import cwords
        cw = cwords.load()                          # 2.31.0: compliance words in the user's style (bold by default)
        rcol = COLUMNS.index("Requirement")
        owner = None
        if self.setup:                              # 2.38.0: the Owner column shows the section's lead when none is set
            import setupdoc
            owner = lambda it: setupdoc.owner_of(self, it)
        for it in self.items:
            ws.append([clean_chars(("Yes" if it.get(FIELD[c]) else "") if c == "Compound" else
                       (owner(it) if owner and c == "Owner" else it.get(FIELD[c], ""))) for c in COLUMNS])
            ws.cell(row=ws.max_row, column=rcol + 1).value = cwords.xlsx_value(clean_chars(it.get("text", "")), cw)
            f = fills.get(it["decision"])
            if f:
                for c in ws[ws.max_row]:
                    c.fill = PatternFill("solid", fgColor=f)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"
        last = ws.cell(row=1, column=len(COLUMNS)).column_letter
        ws.auto_filter.ref = f"A1:{last}{max(2, ws.max_row)}"
        dv = DataValidation(type="list", formula1='"Keep,Context,Drop"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"B2:B{max(2, ws.max_row) + 500}")
        for col in ("Context", "Key"):
            ws.column_dimensions[ws.cell(row=1, column=COLUMNS.index(col) + 1).column_letter].hidden = col in ("Key", "As read", "Place", "Lead place")
        ol = wb.create_sheet("Outline")
        ol.append(["Number", "Title", "Level", "Keep and Context rows placed", "Numbering", "From RFP section"])
        for c in ol[1]:
            c.font = Font(bold=True)
        for o in self.outline:
            lab = outline_label(o)
            ol.append([o["num"], o["title"], o["level"], sum(1 for it in self.items if it["decision"] in ("Keep", "Context") and it["map"] == lab),
                       "auto" if o.get("auto") else "", o.get("from", "")])
        for col, w in zip("ABCDEF", (12, 60, 8, 18, 12, 16)):
            ol.column_dimensions[col].width = w
        so = wb.create_sheet("Sources")
        so.append(["File", "SHA-256", "Read on", "Path on this PC"])
        for src in self.sources:
            so.append(src)
        for col, w in zip("ABCD", (50, 68, 18, 70)):
            so.column_dimensions[col].width = w
        if self.setup:
            self._save_setup(wb, Font)
        cv = wb.create_sheet("Coverage")             # 2.29.0: proof that every sentence was considered
        cv.append(["File", "Sentences", "Keep", "Context", "Customer", "Drop", "Undecided", "Placed", "Need a home", "Every sentence decided"])
        for c in cv[1]:
            c.font = Font(bold=True)
        for row in self.coverage():
            cv.append([row["file"], row["total"], row["Keep"], row["Context"], row["Customer"], row["Drop"], row[""], row["placed"],
                       row["needs_home"], "Yes" if row["total"] and not row[""] else "No"])
        for col, w in zip("ABCDEFGHIJ", (50, 11, 8, 9, 10, 8, 11, 9, 12, 22)):
            cv.column_dimensions[col].width = w
        rh = wb.create_sheet("RFP headings")
        rh.append(["Number", "Title", "File"])
        for h in self.headings:
            rh.append([h["num"], h["title"], h.get("file", "")])
        for col, w in zip("ABC", (12, 60, 40)):
            rh.column_dimensions[col].width = w
        hs = wb.create_sheet("History")
        hs.append(["Date", "User", "Item", "Change"])
        for h in self.history[-5000:]:
            hs.append(h)
        for col, w in zip("ABCD", (17, 14, 18, 70)):
            hs.column_dimensions[col].width = w
        ab = wb.create_sheet("About")
        c = self.counts()
        for r in (["Matrix", self.name], ["Saved", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")],
                  ["Sentences", c["total"]], ["Keep", c["Keep"]], ["Context", c["Context"]], ["Customer", c["Customer"]], ["Drop", c["Drop"]],
                  ["Undecided", c[""]], ["Placed (Keep or Context)", c["placed"]], ["Need a home", c["needs_home"]],
                  ["Note", "Decision: Keep (a requirement), Context (not a requirement, but it goes with the response), "
                           "Customer (the customer's own obligation: identified, not answered) or Drop; "
                           "blank = undecided. Edit here or in Stationkeeper; keep the column headings. "
                           "This workbook holds customer content: keep it on this PC or SharePoint."]):
            ab.append(r)
        ab.column_dimensions["A"].width = 18; ab.column_dimensions["B"].width = 90
        fd, tmp = tempfile.mkstemp(prefix=".matrix_", suffix=".xlsx", dir=d)
        os.close(fd)
        try:
            wb.save(tmp)
            try:
                import workspace as _W                  # 2.80.0: a sync tool holding the file for a moment is waited out
                _W.replace_retry(tmp, self.path)
            except PermissionError:
                raise Locked(self.path)
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)
        self.pending = False


# ------------------------------------------------------------------ Word outline export
def write_outline_docx(matrix, path, title=""):
    """Annotated outline: each outline section as a Word heading, with its kept requirements in a table
    (ID, RFP reference, requirement); then kept requirements not mapped yet."""
    import docxw, cwords
    cw = cwords.load()
    doc = docxw.Doc()
    doc.heading(title or f"{matrix.name}: proposal outline with compliance requirements", 1)
    doc.para(f"Built {datetime.date.today().isoformat()} from the compliance matrix. Kept requirements only.")
    kept = [it for it in matrix.items if it["decision"] == "Keep"]
    for o in matrix.outline:
        lab = outline_label(o)
        doc.heading(lab, min(o["level"] + 1, 6))
        rows = [it for it in kept if it["map"] == lab]
        if rows:
            doc.table(["ID", "RFP reference", "Requirement"],
                      [[it["id"], " ".join(x for x in (it["section"], it["section_title"]) if x) or it["file"],
                        cwords.docx_parts((it["lead"] + " " if it["lead"] else "") + it["text"], None, cw)] for it in rows],
                      [1300, 2000, 6060])
    rest = [it for it in kept if not it["map"]]
    if rest:
        doc.heading("Kept requirements not mapped to the outline", 2)
        doc.table(["ID", "RFP reference", "Requirement"],
                  [[it["id"], " ".join(x for x in (it["section"], it["section_title"]) if x) or it["file"],
                    cwords.docx_parts(it["text"], None, cw)] for it in rest],
                  [1300, 2000, 6060])
    doc.save(path)
    return {"sections": len(matrix.outline), "mapped": sum(1 for it in kept if it["map"]), "unmapped": len(rest)}


# ------------------------------------------------------------------ coverage report (2.29.0)
COVER_FILL = {"Keep": "D9F2E3", "Context": "DBEAFE", "Customer": "F3E8FF", "Drop": "F3F4F6", "": "F3F4F6"}
LABEL = {"Keep": "Compliance requirement", "Context": "Context", "Customer": "Customer's (not ours to answer)", "Drop": "Not used",
         "": "Not used"}


def context_label(it):
    """'Support for L.4-01; L.5-02', 'Support for section 1.2 Schedule' or 'Support' (2.53.0: the team's word; the
    decision is still stored as Context)."""
    if it.get("context_for"):
        return f"Support for {it['context_for']}"
    if it.get("map"):
        return f"Support for section {it['map']}"
    return "Support"


def write_coverage_docx(matrix, path, title=""):
    """Every sentence of the RFP in reading order, highlighted by how it was assigned (2.36.0: compliance
    requirement, context, or not used), with its ID and where it went: the proof that every line was read."""
    import docxw, cwords
    cw = cwords.load()
    doc = docxw.Doc()
    doc.props = {"StationkeeperMatrix": matrix.name or "", "StationkeeperCoverage": datetime.date.today().isoformat()}
    cov = matrix.coverage()
    tot = cov[-1]
    doc.heading(title or f"{matrix.name}: RFP coverage, every sentence", 1)
    unused = tot[""] + tot["Drop"]
    ok = tot["total"] and not tot["needs_home"]
    doc.rich([(f"Every sentence was read: {tot['Keep']} compliance requirement(s), {tot['Context']} context, "
               + (f"{tot['Customer']} the customer's, " if tot["Customer"] else "") + f"{unused} not used."
               + (f" {tot['needs_home']} assigned sentence(s) have no section yet." if tot["needs_home"] else
                  " Every assigned sentence has a section."), {"bold": True, "color": "15803D" if ok else "B91C1C"})])
    doc.para(f"Built {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')} from the compliance matrix {os.path.basename(matrix.path)}.")
    doc.table(["File", "Sentences", "Compliance", "Context", "Customer", "Not used", "In a section", "No section"],
              [[r["file"], r["total"], r["Keep"], r["Context"], r["Customer"], r[""] + r["Drop"], r["placed"], r["needs_home"]] for r in cov],
              [2560, 900, 1100, 900, 1000, 900, 1000, 900])
    doc.rich([("Key:  ", {"bold": True}), (" Compliance requirement ", {"fill": COVER_FILL["Keep"]}), ("  ", {}),
              (" Context ", {"fill": COVER_FILL["Context"]}), ("  ", {}),
              (" Customer's ", {"fill": COVER_FILL["Customer"]}), ("  ", {}),
              (" Not used ", {"fill": COVER_FILL[""], "color": "6B7280"})])
    by_file = {}
    for it in matrix.items:
        by_file.setdefault(it.get("file", ""), []).append(it)
    for r in cov[:-1]:
        f = r["file"]
        doc.heading(f or "(no file)", 2)
        sec = None
        for it in by_file.get(f, []):
            head = " ".join(x for x in (it.get("section", ""), it.get("section_title", "")) if x)
            if head and head != sec:
                doc.heading(head, 3)
                sec = head
            d = it["decision"]
            txt = ((it["lead"] + " ") if it.get("lead") else "") + it["text"]
            where = it["map"] or ("no section" if d in ("Keep", "Context") else "")
            dlab = context_label(it) if d == "Context" else LABEL.get(d, "Not used")
            if d == "Context" and not it.get("context_for"):
                where = "" if it["map"] else where
            tag = f"   [{it['id']} · {dlab}" + (f" · {where}" if where else "") + "]"
            doc.rich(cwords.docx_parts(txt, {"fill": COVER_FILL.get(d, COVER_FILL[""]), "color": "6B7280" if d in ("", "Drop") else None}, cw)
                     + [(tag, {"size": 16, "color": "B91C1C" if where == "no section" else "475569"})])
    doc.save(path)
    return {"sentences": tot["total"], "undecided": tot[""], "needs_home": tot["needs_home"], "files": len(cov) - 1}


# ------------------------------------------------------------------ source page (PDF)
def render_page(path, page_no, text, scale=2.0):
    """A PDF page as a picture with the requirement highlighted (PDFium text search on its first words).
    Returns a PIL image."""
    import pypdfium2 as pdfium
    from PIL import Image, ImageDraw
    pdf = pdfium.PdfDocument(path)
    try:
        page = pdf[page_no - 1]
        w, h = page.get_size()
        img = page.render(scale=scale).to_pil().convert("RGBA")
        tp = page.get_textpage()
        words = text.replace("(a) ", "").split()
        for n in (8, 5, 3):
            probe = " ".join(words[:n])
            if not probe:
                break
            found = tp.search(probe, match_case=False, match_whole_word=False).get_next()
            if found:
                idx, _cnt = found
                full = min(len(text) + 5, tp.count_chars() - idx)
                over = Image.new("RGBA", img.size, (0, 0, 0, 0))
                d = ImageDraw.Draw(over)
                for i in range(tp.count_rects(idx, full)):
                    l, b, r, t = tp.get_rect(i)
                    d.rectangle([l * scale - 2, (h - t) * scale - 2, r * scale + 2, (h - b) * scale + 2], fill=(253, 224, 71, 110),
                                outline=(202, 138, 4, 255), width=2)
                img = Image.alpha_composite(img, over)
                break
        return img.convert("RGB")
    finally:
        pdf.close()
