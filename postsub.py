"""The post-submission checklist (2.91.0): what the team does after the proposal goes in, kept with the effort.

Each item: what to do, who owns it, when it is due (days after submission, or a date set by hand), its status (Not
started, In progress, Done, Not needed), who finished it and when, and notes. A new effort's list starts from the team's
default (Save as default on this PC, carried by a team pack) or the built-in one below. Mark submitted records when the
proposal went in; due dates count from it (from the due date until then). Export writes a formatted Excel workbook for
the effort's folder. Kept in the effort's record (record.json "postsub"). Rules only (No AI); nothing leaves the PC.
"""
import datetime, getpass, json, os, uuid

STATUS = ("Not started", "In progress", "Done", "Not needed")
DEFAULT_FILE = "postsub_default.json"
BUILT_IN = (
    ("Confirm the customer received the submission; save the receipt or confirmation e-mail", "Proposal Manager", 0),
    ("Record the exact submission date and time, and how it was sent", "Proposal Manager", 0),
    ("Archive the final volumes as submitted (PDF and Word), with the as-submitted compliance matrix", "Production", 1),
    ("Save the final effort package to the approved SharePoint folder", "Production", 1),
    ("Send the team a thank-you and the submission summary", "Proposal Manager", 1),
    ("Watch for clarifications, discussions and requests for final proposal revisions", "Capture", 3),
    ("Plan for orals: team, dates and rehearsal schedule", "Capture", 5),
    ("Hold the lessons-learned meeting and record what to keep and change", "Proposal Manager", 7),
    ("Update the boilerplate and the past-performance library with what was written", "Production", 14),
    ("Clean up the team folder: drafts, duplicates and access", "Production", 14),
    ("Request a debrief when the award is announced", "Capture", 30),
)


def _now():
    return datetime.datetime.now().isoformat(timespec="minutes")


def default_items(config_dir=""):
    """[(text, owner, days)]: the team's default saved on this PC, else the built-in one."""
    p = os.path.join(config_dir, DEFAULT_FILE) if config_dir else ""
    try:
        with open(p, encoding="utf-8") as f:
            d = json.load(f)
        rows = [(str(x.get("text", "")).strip(), str(x.get("owner", "")), int(x.get("days", 0) or 0)) for x in d.get("items", [])]
        rows = [r for r in rows if r[0]]
        if rows:
            return rows
    except (OSError, ValueError, AttributeError, TypeError):
        pass
    return list(BUILT_IN)


def save_default(config_dir, items):
    """This effort's items (text, owner, days after submission) become the default for new efforts on this PC."""
    rows = [{"text": it["text"], "owner": it.get("owner", ""), "days": int(it.get("days", 0) or 0)} for it in items if it.get("text")]
    os.makedirs(config_dir, exist_ok=True)
    p = os.path.join(config_dir, DEFAULT_FILE)
    with open(p + ".tmp", "w", encoding="utf-8") as f:
        json.dump({"items": rows}, f, indent=1, ensure_ascii=False)
    os.replace(p + ".tmp", p)
    return len(rows)


def data(rec, config_dir=""):
    """The effort's checklist, made from the default the first time."""
    d = rec.data.get("postsub")
    if not isinstance(d, dict) or "items" not in d:
        d = {"submitted": "", "items": [_item(t, o, n) for t, o, n in default_items(config_dir)]}
        rec.data["postsub"] = d
        rec.save()
    return d


def _item(text, owner="", days=0):
    return {"id": uuid.uuid4().hex[:10], "text": " ".join(str(text).split())[:400], "owner": owner, "days": int(days or 0),
            "due": "", "status": "Not started", "done_by": "", "done_at": "", "notes": ""}


def base_date(rec):
    """The day due dates count from: when it was submitted, else the effort's due date; None without either."""
    import effortrec as R
    d = rec.data.get("postsub") or {}
    for v in (d.get("submitted", ""), rec.data.get("due", "")):
        if v:
            return R.parse_date(v[:10])
    return None


def due_of(rec, it):
    """The item's due date (a date set by hand wins), or None."""
    import effortrec as R
    if it.get("due"):
        return R.parse_date(it["due"])
    b = base_date(rec)
    return b + datetime.timedelta(days=int(it.get("days", 0) or 0)) if b else None


def late(rec, it, today=None):
    d = due_of(rec, it)
    return bool(d and it.get("status") not in ("Done", "Not needed") and d < (today or datetime.date.today()))


def add(rec, text, owner="", days=0, due="", who=None):
    d = data(rec)
    it = _item(text, owner, days)
    it["due"] = due or ""
    if not it["text"]:
        return None
    d["items"].append(it)
    rec.log("task", f"Post-submission: added {it['text'][:80]}", who=who, save=True)
    return it["id"]


def update(rec, iid, who=None, **fields):
    d = data(rec)
    it = next((x for x in d["items"] if x["id"] == iid), None)
    if it is None:
        return False
    for k, v in fields.items():
        if k in ("text", "owner", "due", "notes"):
            it[k] = " ".join(str(v or "").split()) if k != "notes" else str(v or "").strip()
        elif k == "days":
            it["days"] = int(v or 0)
        elif k == "status" and v in STATUS and v != it["status"]:
            it["status"] = v
            if v == "Done":
                it["done_by"], it["done_at"] = who or getpass.getuser(), _now()
                rec.log("task", f"Post-submission done: {it['text'][:80]}", who=who)
            else:
                it["done_by"] = it["done_at"] = ""
    rec.save()
    return True


def remove(rec, iid, who=None):
    d = data(rec)
    it = next((x for x in d["items"] if x["id"] == iid), None)
    if it is None:
        return False
    d["items"].remove(it)
    rec.log("task", f"Post-submission: removed {it['text'][:80]}", who=who, save=True)
    return True


def move(rec, iid, step):
    items = data(rec)["items"]
    i = next((k for k, x in enumerate(items) if x["id"] == iid), None)
    if i is None or not 0 <= i + step < len(items):
        return False
    items[i], items[i + step] = items[i + step], items[i]
    rec.save()
    return True


def mark_submitted(rec, when="", who=None):
    d = data(rec)
    d["submitted"] = when or _now()
    rec.log("effort", "Proposal submitted " + d["submitted"].replace("T", " "), who=who, save=True)
    return d["submitted"]


def progress(rec):
    """(done or not needed, total, late)."""
    items = data(rec)["items"]
    return (sum(1 for x in items if x["status"] in ("Done", "Not needed")), len(items), sum(1 for x in items if late(rec, x)))


def export_xlsx(rec, path, effort="", who=None):
    """A formatted workbook: the effort and its dates on top, one row per item, a Status drop-down, done rows in green
    and late ones in red, frozen headings, one page wide when printed."""
    import openpyxl
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    import effortrec as R
    d = data(rec)
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Post-submission"
    navy = "0F1E33"
    ws["A1"] = f"{effort} post-submission checklist" if effort else "Post-submission checklist"
    ws["A1"].font = Font(bold=True, size=16, color=navy)
    sub = d.get("submitted", "")
    due = rec.data.get("due", "")
    meta = [("Submitted", sub.replace("T", " ") if sub else "Not marked yet"),
            ("Proposal due", R.due_text(due, rec.data.get("zone", "")) if due else "Not set"),
            ("Exported", f"{datetime.datetime.now():%Y-%m-%d %H:%M} by {who or getpass.getuser()}")]
    done, total, nlate = progress(rec)
    meta.append(("Progress", f"{done} of {total} done" + (f", {nlate} late" if nlate else "")))
    for i, (k, v) in enumerate(meta, start=2):
        ws.cell(i, 1, f"{k}:  {v}").font = Font(color="4E5D71")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    head_row = len(meta) + 3
    heads = ("#", "Task", "Owner", "Due", "Status", "Done by", "Done on", "Notes")
    thin = Side(style="thin", color="D3DBE6")
    for c, h in enumerate(heads, start=1):
        cell = ws.cell(head_row, c, h)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F5FAD")
        cell.alignment = Alignment(vertical="center")
    for n, it in enumerate(d["items"], start=1):
        r = head_row + n
        dd = due_of(rec, it)
        vals = (n, it["text"], it.get("owner", ""), dd, it["status"], it.get("done_by", ""),
                (it.get("done_at", "") or "")[:10], it.get("notes", ""))
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(r, c, v)
            cell.alignment = Alignment(wrap_text=c in (2, 8), vertical="top", horizontal="center" if c == 1 else "left")
            cell.border = Border(bottom=thin)
            if c == 4 and dd:
                cell.number_format = "mmm d, yyyy"
        if n % 2 == 0:
            for c in range(1, len(heads) + 1):
                ws.cell(r, c).fill = PatternFill("solid", fgColor="F6F8FB")
    last = head_row + max(1, len(d["items"]))
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=False)
    ws.add_data_validation(dv); dv.add(f"E{head_row + 1}:E{last + 50}")
    rng = f"A{head_row + 1}:H{last + 50}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR($E{head_row + 1}="Done",$E{head_row + 1}="Not needed")'],
                                                   fill=PatternFill("solid", fgColor="DCF3E4"), font=Font(color="146C36")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($D{head_row + 1}<>"",$D{head_row + 1}<TODAY(),$E{head_row + 1}<>"Done",$E{head_row + 1}<>"Not needed")'],
                                                   fill=PatternFill("solid", fgColor="FBE2E2"), font=Font(color="B91C1C")))
    for col, wdt in zip("ABCDEFGH", (5, 62, 20, 14, 14, 16, 12, 40)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    ws.auto_filter.ref = f"A{head_row}:H{last}"
    ws.print_title_rows = f"{head_row}:{head_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Page &P of &N"
    tmp = path + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, path)
    return path
