"""My work, across efforts (3.1.0): for one person, everything assigned to them in every effort on this PC: the sections
they lead or write (with page budgets), the compliance requirements they answer (and those a customer version changed),
their tasks, the team's questions for them, the customer questions, deliverables and graphic needs they own, each with
its due date, soonest first. Read only: it never changes an effort. Rules only."""
import datetime, os
import setupdoc as SD

KIND_WORDS = {"task": "Task", "question": "Question for you", "custqa": "Customer question", "deliverable": "Deliverable",
              "graphic": "Graphic need", "section": "Section", "requirements": "Requirements", "stage": "Stage"}


def _is(person, names):
    p = (person or "").casefold()
    return bool(p) and any(n.casefold() == p for n in SD.split_names(names or ""))


def _when(text):
    import effortrec as R
    w = R.parse_when(text or "", end_of_day=True)
    return w


def for_effort(name, info, person, load_matrix=True):
    """One effort's work for person: {"effort", "due", "stage", "sections", "reqs", "awaiting", "items"}; None when the
    effort cannot be read. items: [{"kind", "text", "due" (iso), "id", "where"}]."""
    import workspace as W
    import effortrec as R
    path = (info or {}).get("path", "")
    if not path:
        return None
    root = W.data_dir_for(path)
    if not os.path.isdir(root):
        return None
    try:
        rec = R.Record(root)
    except Exception:
        return None
    out = {"effort": name, "due": rec.data.get("due", ""), "stage": (rec.stage() or {}).get("name", ""), "sections": [],
           "reqs": 0, "awaiting": 0, "items": []}
    items = out["items"]
    for t in rec.data.get("tasks", []):
        if _is(person, t.get("owner")) and t.get("status") != "Done":
            items.append({"kind": "task", "text": t["title"], "due": t.get("due", ""), "id": t["id"], "where": t.get("status", "")})
    for n in rec.data.get("notes", []):
        if n.get("question") and not n.get("resolved") and _is(person, n.get("assignee")):
            items.append({"kind": "question", "text": n.get("text", ""), "due": "", "id": n["id"], "where": n.get("who", "")})
    try:
        import custqa as Q
        for q in Q.items(rec):
            if _is(person, q.get("owner")) and q.get("status") in ("Draft", "Approved"):
                items.append({"kind": "custqa", "text": q.get("question", "") or q.get("quote", ""), "due": "", "id": q.get("id", ""),
                              "where": f"{q.get('status', '')}, {q.get('doc', '')}".strip(", ")})
    except Exception:
        pass
    try:
        import deliverables as DL
        for d in DL.items(rec):
            if _is(person, d.get("owner")) and d.get("status") not in ("Final", "Submitted"):
                items.append({"kind": "deliverable", "text": d.get("name", ""), "due": d.get("due", ""), "id": d.get("id", ""),
                              "where": d.get("status", "") + (f", limit {d['limit']}" if d.get("limit") else "")})
    except Exception:
        pass
    try:
        import graphics as GX
        for g in GX.items(rec):
            if _is(person, g.get("owner")) and g.get("status") != "Done":
                items.append({"kind": "graphic", "text": g.get("concept", ""), "due": "", "id": g["id"],
                              "where": " ".join(x for x in (g.get("status", ""), g.get("section", "")) if x)})
    except Exception:
        pass
    if load_matrix:
        mp = W.read_json(os.path.join(root, "settings.json"), {}).get("compliance_path") or ""
        if mp and os.path.isfile(mp):
            try:
                import compliance as C
                import pagebudget as PB
                M = C.Matrix(mp).load()
                for lab in SD.sections_of(M, person):
                    p = SD.section_people(M, lab)
                    out["sections"].append({"label": lab, "role": "lead" if p.get("lead") == person else "writer",
                                            "pages": PB.pages(M, lab)})
                mine = [it for it in M.items if it.get("decision") == "Keep" and _is(person, SD.owner_of(M, it))]
                out["reqs"] = len(mine)
                try:
                    import changereview as CR
                    wait = set(CR.Review(root).load().awaiting())
                    out["awaiting"] = sum(1 for it in mine if it["id"] in wait)
                except Exception:
                    pass
            except Exception:
                pass
    # the stage the person's writing is due in: the current stage's planned end, when one is set
    st = rec.stage()
    if (out["reqs"] or out["sections"]) and st is not None:
        items.append({"kind": "requirements", "text": f"{out['reqs']} compliance requirement(s) to answer" +
                      (f", {out['awaiting']} changed by a customer version" if out["awaiting"] else ""),
                      "due": st.get("planned_end", ""), "id": "", "where": f"{st['name']} ends" if st.get("planned_end") else st["name"]})
    return out


def collect(registry, person, load_matrix=True, names=None):
    """Every effort's work for person (archived efforts left out): [for_effort(...)] for those with anything assigned."""
    out = []
    for name in (names if names is not None else registry.names()):
        info = registry.get(name) or {}
        if info.get("archived"):
            continue
        got = for_effort(name, info, person, load_matrix)
        if got and (got["items"] or got["sections"]):
            out.append(got)
    return out


def soonest(work, today=None):
    """Every item across the efforts in one list: late first, then by due date, then those with no date (by the
    effort's due date). Each gets "effort", "late" and "due_date"."""
    today = today or datetime.datetime.now()
    rows = []
    for w in work:
        edue = _when(w.get("due"))
        for it in w["items"]:
            d = _when(it.get("due"))
            rows.append(dict(it, effort=w["effort"], due_date=d, late=bool(d and d < today), effort_due=edue))
    far = datetime.datetime(9999, 1, 1)
    rows.sort(key=lambda r: (r["due_date"] is None, r["due_date"] or r["effort_due"] or far, r["effort_due"] or far, r["kind"]))
    return rows



def export_excel(path, person, work):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "My work"
    ws.append(["Due", "Effort", "What", "Item", "Where"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1E3A8A")
    for r in soonest(work):
        ws.append([r["due_date"].strftime("%Y-%m-%d %H:%M").replace(" 23:59", "") if r["due_date"] else "", r["effort"],
                   KIND_WORDS.get(r["kind"], r["kind"]), r["text"], r["where"]])
        if r["late"]:
            for c in ws[ws.max_row]:
                c.font = Font(color="B91C1C")
    sh = wb.create_sheet("Sections")
    sh.append(["Effort", "Section", "Role", "Page budget"])
    for c in sh[1]:
        c.font = Font(bold=True)
    for w in work:
        for x in w["sections"]:
            sh.append([w["effort"], x["label"], x["role"], x["pages"] or ""])
    for s_, widths in ((ws, [16, 26, 20, 70, 30]), (sh, [26, 40, 10, 12])):
        for i, wd in enumerate(widths, 1):
            s_.column_dimensions[s_.cell(row=1, column=i).column_letter].width = wd
        for row in s_.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.properties.title = f"Work for {person}"
    wb.save(path)
    return path
