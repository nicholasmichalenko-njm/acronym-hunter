"""Customer document versions and compare (2.37.0). Rules and counts only (No AI); nothing leaves the PC.

Every customer document the effort receives is kept as a document with versions (Original, Amendment 1, Rev B...).
Each version records when it arrived, who added it, the file name, its SHA-256 fingerprint and a copy of the text
read from it (sentence by sentence, with the RFP section each sentence sits under). The copy lets any two versions
be compared later even if the file is moved, replaced or deleted, and the fingerprint shows the file is the one read.

Stored in the effort's folder (workspace): documents\\index.json and documents\\<version id>.json. Saved at once.

When a new file arrives the library suggests which document it is a new version of: by name (the name with version
words such as Amendment 2, Amd1, Rev B, v3, dates and copy marks removed) and by how much of the text is the same.
A person confirms every match.

Compare: sentences are lined up in reading order (difflib over the sentence lists); a removed sentence next to an
added one in the same place is paired as Changed when they are at least 50% alike, with a word diff. The result
says which changes touch submission requirements (format, page limits, due dates, files and copies).
"""
import datetime, difflib, getpass, hashlib, json, os, re, uuid

VERSION_WORDS = re.compile(
    r"(?<![a-z])(?:amend(?:ment)?|amd|mod(?:ification)?|ver(?:sion)?|v|r)[\s_.-]*(?:no\.?\s*)?\d{1,3}(?![a-z])|"
    r"(?<![a-z])rev(?:ision)?[\s_.-]*(?:\d{1,3}|[a-z])?(?![a-z])|"
    r"(?<![a-z])(?:draft|final|updated?|conformed|clean|redline|copy)(?![a-z])|\(\d+\)|"
    r"\b\d{4}[-_.]\d{2}[-_.]\d{2}\b|\b\d{1,2}[-_.]\d{1,2}[-_.]\d{2,4}\b|"
    r"(?<![a-z])(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*[\s_-]*\d{1,4}(?![a-z])", re.I)
PAGE_LIMIT = re.compile(r"\b(?:exceed|limit(?:ed)?|maximum|no more than|up to)\b[^.]{0,40}\b\d{1,3}\s*(?:-\s*)?pages?\b|"
                        r"\b\d{1,3}[- ]page (?:limit|maximum)\b", re.I)
STRONG, POSSIBLE = 0.6, 0.3
COMPARE_RULES = 1                     # 2.82.0: bump when compare() or summary_text() change; kept summaries are redone
_SHA = {}                             # (path, size, modified) -> fingerprint, for files already fingerprinted this session
_CHANGE = {}                          # (library folder, older version id, newer version id) -> what changed, this session


def sha256_cached(path):
    """sha256(path), worked out once per file until its size or time changes (2.82.0: a folder of large PDFs is not
    read again each time a page opens)."""
    st = os.stat(path)
    key = (os.path.abspath(path), st.st_size, st.st_mtime_ns)
    h = _SHA.get(key)
    if h is None:
        h = _SHA[key] = sha256(path)
    return h


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest().upper()


def short_sha(h):
    return f"{h[:4]} {h[4:8]} ... {h[-4:]}" if h else ""


def stem(name):
    """'Kestrel_RFP_CLM_Amd2 (1).docx' -> 'kestrel rfp clm'."""
    s = os.path.splitext(os.path.basename(name or ""))[0]
    s = VERSION_WORDS.sub(" ", s)
    s = re.sub(r"[_\-.]+", " ", s)
    return " ".join(w for w in s.lower().split() if w)


def display_name(name):
    """'Kestrel_RFP_CLM_Amd2 (1).docx' -> 'Kestrel RFP CLM' (the file's own capitals kept)."""
    s = os.path.splitext(os.path.basename(name or ""))[0]
    s = VERSION_WORDS.sub(" ", s)
    s = re.sub(r"[_]+", " ", s)
    return " ".join(s.split()).strip(" -.") or os.path.basename(name or "")


def name_similarity(a, b):
    sa, sb = stem(a), stem(b)
    if not sa or not sb:
        return 0.0
    if sa == sb:
        return 1.0
    return difflib.SequenceMatcher(None, sa, sb).ratio()


def _shingles(sents, n=5):
    out = set()
    for s in sents:
        w = re.findall(r"[a-z0-9]+", s.lower())
        for i in range(max(1, len(w) - n + 1)):
            out.add(" ".join(w[i:i + n]))
    return out


def text_similarity(a_sents, b_sents):
    """Share of word runs in common (0 to 1)."""
    A, B = _shingles(a_sents), _shingles(b_sents)
    if not A or not B:
        return 0.0
    return len(A & B) / max(1, min(len(A), len(B)))


# ------------------------------------------------------------------ reading a file
def snapshot(name, chunks):
    """[(location label, text)] from the file reader -> [{"s": section, "t": section title, "x": text, "l": where}]."""
    import compliance as C
    heads = []
    items = C.extract([(name, chunks)], heads)
    return [{"s": it.get("section", ""), "t": it.get("section_title", ""), "x": ((it.get("lead") + " ") if it.get("lead") else "") + it["text"],
             "l": it.get("loc", "")} for it in items]


# ------------------------------------------------------------------ what each customer document is (2.68.0)
DOC_TYPES = ("RFP", "Amendment", "SOW", "SOO", "PWS", "Section L", "Section M", "Q&A", "Attachment", "CDRL", "DD 254", "Contract",
             "Task order", "Delivery order", "Other")
TYPE_INITIAL = {"RFP": "R", "Amendment": "A", "SOW": "S", "SOO": "O", "PWS": "P", "Section L": "L", "Section M": "M", "Q&A": "Q",
                "Attachment": "J", "CDRL": "C", "DD 254": "D", "Contract": "CT", "Task order": "TO", "Delivery order": "DO", "Other": "X"}
ORDER_TYPES = ("Task order", "Delivery order")      # filed under the base contract or RFP, each with its own versions (mods)
BASE_TYPES = ("Contract", "RFP")


NUMBERED = {"Amendment", "Task order", "Delivery order"}            # 2.94.0: types numbered (A1, A2...; TO1, TO2...)
TYPE_ALIAS = {}                     # 2.94.0: an old type name -> its new name (Settings > Document types renamed it)
TYPE_WORDS = {}                     # 2.94.0: a team type -> words in a file name that suggest it
_BUILT_IN_GUESS = set(DOC_TYPES)


def set_types(types):
    """2.94.0: the team's document types (doctypes.py): names in order, initials, numbered, builds the matrix, file name
    words, old names."""
    global DOC_TYPES
    DOC_TYPES = tuple(t["name"] for t in types)
    TYPE_INITIAL.clear(); TYPE_INITIAL.update({t["name"]: t["initial"] for t in types})
    NUMBERED.clear(); NUMBERED.update(t["name"] for t in types if t.get("numbered"))
    MATRIX_TYPES.clear(); MATRIX_TYPES.update(t["name"] for t in types if t.get("matrix"))
    TYPE_ALIAS.clear(); TYPE_ALIAS.update({a: t["name"] for t in types for a in t.get("aliases", [])})
    TYPE_WORDS.clear(); TYPE_WORDS.update({t["name"]: t["words"] for t in types if t.get("words")})


def type_of(name):
    """A stored type name as it is called now (a renamed type follows its new name)."""
    return TYPE_ALIAS.get(name, name) if name else name


def clean_initial(text):
    """An initial: letters and digits only, upper case, at most 6 characters."""
    import re
    return re.sub(r"[^A-Za-z0-9]", "", text or "").upper()[:6]


def guess_type(name):
    """The document type a file name suggests (the person confirms it)."""
    import re
    n = (name or "").lower()
    for t, words in TYPE_WORDS.items():          # 2.94.0: the team's own types first, by their words
        for w in [x.strip().lower() for x in re.split(r"[,;]", words) if x.strip()]:
            if re.search(r"(?<![a-z0-9])" + re.escape(w) + r"(?![a-z0-9])", n.replace("_", " ")):
                return t
    for pat, t in ((r"\btask[\s_-]*order\b|\bto[\s_-]?\d{2,4}\b", "Task order"),
                   (r"\bdelivery[\s_-]*order\b|\bdo[\s_-]?\d{2,4}\b", "Delivery order"),
                   (r"\bidiq\b|\bbase[\s_-]*contract\b|\bcontract\b(?!.*\b(?:amend|mod)\b)", "Contract"),
                   (r"amend|\bamd\b|_amd|amd\d|\bmod\b", "Amendment"), (r"\bsow\b|statement[\s_-]*of[\s_-]*work", "SOW"),
                   (r"\bsoo\b|statement[\s_-]*of[\s_-]*objectives", "SOO"), (r"\bpws\b|performance[\s_-]*work", "PWS"),
                   (r"q\s*&\s*a|questions|\bqa\b", "Q&A"), (r"cdrl", "CDRL"), (r"dd[\s_-]*254", "DD 254"),
                   (r"attachment|\batt\b|exhibit|\bj-\d", "Attachment"), (r"section[\s_]*l\b", "Section L"),
                   (r"section[\s_]*m\b", "Section M"), (r"\brfp\b|solicitation|\brfq\b", "RFP")):
        if re.search(pat, n.replace("_", " ")) and t in DOC_TYPES:
            return t
    return "RFP" if "RFP" in DOC_TYPES else (DOC_TYPES[0] if DOC_TYPES else "Other")


# ------------------------------------------------------------------ the library
# ------------------------------------------------------------------ the document library (2.83.0)
CATEGORIES = (("customer", "Customer documents"), ("bidders", "Bidders library"), ("qa", "Q&A"), ("ours", "Our effort documents"))
ROLES = (("matrix", "Builds the matrix"), ("context", "Context"), ("reference", "Reference only"), ("not_needed", "Not needed"))
ROLE_NAMES = dict(ROLES)
MATRIX_TYPES = {"RFP", "Amendment", "SOW", "SOO", "PWS", "Section L", "Section M", "Task order", "Delivery order"}
_CAT_WORDS = (("qa", r"q\s*&\s*a|\bqa\b|questions?\b|answers?\b"),
              ("bidders", r"bidders?|library|reference|standards?|\bmil[\s_-]*(std|hdbk)|specs?\b|handbook"),
              ("ours", r"\bours?\b|proposal|draft|past[\s_-]*perf|resumes?|capture|win[\s_-]*theme|internal|cpars"),
              ("customer", r"customer|\brfp\b|solicitation|amendment|\bamd\b|section[\s_]*[lmc]\b|\bsow\b|\bpws\b|\bsoo\b|attachment|cdrl|dd[\s_-]*254|rfi\b"))
_TEXTS = {}                           # (library folder, version id) -> lower-case sentences, for search, this session


def suggest_role(category, dtype=""):
    """The role a document's category and type suggest (the person confirms or changes it)."""
    if category == "customer":
        return "matrix" if dtype in MATRIX_TYPES or not dtype else "context"
    return {"qa": "context", "bidders": "reference", "ours": "reference"}.get(category, "reference")


def guess_category(relpath, default="customer"):
    """The category a file's folder path and name suggest: the most specific folder decides."""
    parts = [x for x in re.split(r"[\\/]+", relpath or "") if x]
    for part in reversed(parts):
        low = part.lower().replace("_", " ")
        for cat, pat in _CAT_WORDS:
            if re.search(pat, low):
                return cat
    return default


def folder_tags(folder):
    """2.83.0: every folder name in the path is a tag (the person's own organizing, kept)."""
    out = []
    for x in re.split(r"[\\/]+", folder or ""):
        x = " ".join(x.split())
        if x and x.lower() not in {t.lower() for t in out}:
            out.append(x)
    return out


def clean_tags(tags):
    out = []
    for t in tags or []:
        t = " ".join(str(t).replace(",", " ").split())[:40]
        if t and t.lower() not in {x.lower() for x in out}:
            out.append(t)
    return out


class Library:
    def __init__(self, folder):
        self.folder = folder                       # <effort app data>\documents
        self.all = []                              # 2.83.0: every document; .docs is the customer documents
        self.extra_categories = []                 # categories the person added: [[key, name]]

    @property
    def docs(self):
        """The customer documents: what requirements, Compare, Impact review and Read RFP work on (2.83.0: the library
        also holds the bidders library, Q&A and our documents, which those leave alone)."""
        return [d for d in self.all if d.get("category", "customer") == "customer"]

    @property
    def data_dir(self):
        return os.path.dirname(self.folder)

    @property
    def index_path(self):
        return os.path.join(self.folder, "index.json")

    def load(self):
        import workspace as W
        d = W.read_json(self.index_path, None) or {}
        self.all = d.get("docs", [])
        for x in self.all:                          # 2.94.0: a renamed type: its documents follow the new name
            if x.get("type") in TYPE_ALIAS:
                x["type"] = TYPE_ALIAS[x["type"]]
        self.extra_categories = [c for c in d.get("categories", []) if isinstance(c, list) and len(c) == 2]
        return self

    def save(self):
        import workspace as W
        W.write_json(self.index_path, {"v": 2, "docs": self.all, "categories": self.extra_categories})

    def text(self, version):
        import workspace as W
        return W.read_json(os.path.join(self.folder, version["id"] + ".json"), None) or []

    def _write_text(self, vid, sents):
        import workspace as W
        W.write_json(os.path.join(self.folder, vid + ".json"), sents)

    def doc(self, doc_id):
        return next((d for d in self.all if d["id"] == doc_id), None)

    def change_text(self, d, i):
        """What changed from the version before (2.82.0): worked out once and kept with the version, because two
        versions' texts never change. Returns (text, newly worked out)."""
        if i <= 0 or i >= len(d["versions"]):
            return "", False
        v, prev = d["versions"][i], d["versions"][i - 1]
        c = v.get("change") or {}
        if c.get("prev") == prev["id"] and c.get("rules") == COMPARE_RULES and "text" in c:
            return c["text"], False
        key = (self.folder, prev["id"], v["id"])
        text = _CHANGE.get(key)
        if text is None:
            text = _CHANGE[key] = summary_text(summary(compare(self.text(prev), self.text(v))))
        v["change"] = {"prev": prev["id"], "text": text, "rules": COMPARE_RULES}
        return text, True

    def warm(self, stop=None):
        """Work out, in the background when an effort opens, what pages would otherwise work out on first view: what
        changed between versions and the submission requirements each version sets. Memory and per-version files only;
        the index is saved by the page that shows them (2.82.0). stop(): True to give up early."""
        import subreq as SR
        for d in list(self.docs):
            for i, v in enumerate(list(d["versions"])):
                if stop and stop():
                    return
                try:
                    SR._found_for(self, v)
                    if i:
                        prev = d["versions"][i - 1]
                        c = v.get("change") or {}
                        if not (c.get("prev") == prev["id"] and c.get("rules") == COMPARE_RULES):
                            key = (self.folder, prev["id"], v["id"])
                            if key not in _CHANGE:
                                _CHANGE[key] = summary_text(summary(compare(self.text(prev), self.text(v))))
                except Exception:
                    pass

    def find_sha(self, h):
        for d in self.all:
            for v in d["versions"]:
                if v.get("sha256") == h:
                    return d, v
        return None

    def files(self):
        return {v.get("file", "").lower() for d in self.all for v in d["versions"]}

    def latest(self, d):
        return d["versions"][-1] if d.get("versions") else None

    def suggest(self, name, sents):
        """[(doc, score, how)] best first: how the new file matches each document."""
        out = []
        for d in self.all:
            last = self.latest(d)
            if last is None:
                continue
            ns = max(name_similarity(name, last.get("file", "")), name_similarity(name, d["name"]))
            ts = text_similarity(sents, [s["x"] for s in self.text(last)]) if sents else 0.0
            score = max(ts, 0.5 * ns + 0.5 * ts, ns if ns >= 0.95 and ts >= 0.15 else 0)
            how = f"{int(round(ts * 100))}% of the text is the same" + ("; the name matches" if ns >= 0.9 else "")
            out.append((d, round(score, 3), how))
        out.sort(key=lambda x: -x[1])
        return out

    @staticmethod
    def band(score):
        return "strong" if score >= STRONG else "possible" if score >= POSSIBLE else ""

    def next_label(self, d):
        n = len(d["versions"])
        return f"Amendment {n}" if n else "Original"

    def add(self, path, sents, doc_id=None, doc_name="", label="", sha=None, when=None, who=None, initial="", dtype="",
            category="customer", folder="", tags=None, role=None, save=True, role_confirmed=True):
        """Add a file as a new version of doc_id, or as a new document (2.68.0: with its initial and type; 2.83.0: its
        category, folder, tags and role). Returns (doc, version)."""
        h = sha or sha256(path)
        d = self.doc(doc_id) if doc_id else None
        if d is None:
            cat = category if category in self.category_keys() else "customer"
            cust = cat == "customer"
            d = {"id": uuid.uuid4().hex[:12], "name": doc_name or display_name(path), "versions": [],
                 "initial": clean_initial(initial) if cust and clean_initial(initial) not in self.initials() else "",
                 "type": dtype if dtype in DOC_TYPES and cust else "", "category": cat, "folder": (folder or "").strip("/\\"),
                 "tags": clean_tags((tags if tags is not None else []) + folder_tags(folder))}
            r = role if role in ROLE_NAMES else suggest_role(cat, d["type"])
            d["role"] = r
            d["role_by"] = {"who": who or getpass.getuser(), "when": datetime.datetime.now().isoformat(timespec="minutes"),
                            "why": "", "suggested": role not in ROLE_NAMES or not role_confirmed}
            self.all.append(d)
        v = {"id": uuid.uuid4().hex[:12], "label": label or self.next_label(d), "file": os.path.basename(path),
             "added": when or datetime.datetime.now().isoformat(timespec="minutes"), "who": who or getpass.getuser(),
             "sha256": h, "sentences": len(sents), "reviewed": "", "path": os.path.abspath(path)}
        rel = self._rel(path)
        if rel:
            v["rel"] = rel                         # 2.83.0: a copy inside the effort: found again wherever the effort goes
        self._write_text(v["id"], sents)
        d["versions"].append(v)
        if save:
            self.save()
        return d, v

    def move_version(self, vid, to_doc_id, label=""):
        """'This file is a version of another document': move a version (keeps its text and fingerprint)."""
        src = next((d for d in self.all if any(v["id"] == vid for v in d["versions"])), None)
        dst = self.doc(to_doc_id)
        if src is None or dst is None or src is dst:
            return False
        v = next(v for v in src["versions"] if v["id"] == vid)
        src["versions"].remove(v)
        v["label"] = label or self.next_label(dst)
        dst["versions"].append(v)
        dst["versions"].sort(key=lambda x: x.get("added", ""))
        if not src["versions"]:
            self.all.remove(src)
        self.save()
        return True

    def rename(self, doc_id=None, vid=None, name=""):
        name = " ".join((name or "").split())
        if not name:
            return False
        for d in self.all:
            if doc_id and d["id"] == doc_id:
                d["name"] = name
            for v in d["versions"]:
                if vid and v["id"] == vid:
                    v["label"] = name
        self.save()
        return True

    # ------------------------------------------------------------ identity (2.68.0)
    def initials(self, but=None):
        return {d.get("initial", "") for d in self.docs if d.get("initial") and d["id"] != but}

    def suggest_initial(self, dtype, taken=()):
        """The type's letter, numbered when taken (R, R2...; amendments A1, A2...)."""
        base = TYPE_INITIAL.get(dtype, "X")
        used = self.initials() | set(taken)
        if dtype in NUMBERED:
            n = 1
            while f"{base}{n}" in used:
                n += 1
            return f"{base}{n}"
        if base not in used:
            return base
        n = 2
        while f"{base}{n}" in used:
            n += 1
        return f"{base}{n}"

    def check_identity(self, initial, dtype, doc_id=None):
        """'' when fine, else what is wrong (the person sees it)."""
        ini = clean_initial(initial)
        if not ini:
            return "Give the document an initial (letters and digits, up to 6)."
        if dtype not in DOC_TYPES:
            return "Pick the document's type."
        if ini in self.initials(but=doc_id):
            return f"Another document already uses {ini}."
        return ""

    def set_identity(self, doc_id, initial, dtype):
        """The document's initial and type. Returns '' when set, else why not."""
        d = self.doc(doc_id)
        if d is None:
            return "No such document."
        why = self.check_identity(initial, dtype, doc_id)
        if why:
            return why
        d["initial"], d["type"] = clean_initial(initial), dtype
        self.save()
        return ""

    def missing_identity(self):
        return [d for d in self.docs if not d.get("initial") or not d.get("type")]

    def set_under(self, doc_id, base_id, save=True):
        """A task or delivery order (any document) filed under a base document ('' takes it out). False when it would
        be filed under itself or one of its own."""
        d = self.doc(doc_id)
        if d is None or base_id == doc_id or (base_id and self.doc(base_id) is None):
            return False
        b = base_id
        while b:
            if b == doc_id:
                return False
            b = (self.doc(b) or {}).get("under", "")
        if base_id:
            d["under"] = base_id
        else:
            d.pop("under", None)
        if save:
            self.save()
        return True

    def children(self, base_id):
        return [d for d in self.docs if d.get("under") == base_id]

    def base_for_orders(self):
        """Where an order goes by default: the contract, else the RFP (the first customer document of that type)."""
        for t in BASE_TYPES:
            d = next((x for x in self.docs if x.get("type") == t and not x.get("under")), None)
            if d is not None:
                return d
        return None

    def doc_of_file(self, file_name):
        f = os.path.basename(file_name or "").lower()
        return next((d for d in self.docs if any(v.get("file", "").lower() == f for v in d["versions"])), None)

    def identity_map(self):
        """{file name: {"initial", "type", "doc"}} for every version of every document."""
        return {v.get("file", ""): {"initial": d.get("initial", ""), "type": d.get("type", ""), "doc": d["name"]}
                for d in self.docs for v in d["versions"] if v.get("file")}

    def mark_reviewed(self, vid, who=None):
        for d in self.all:
            for v in d["versions"]:
                if v["id"] == vid:
                    v["reviewed"] = f"{datetime.date.today().isoformat()} {who or getpass.getuser()}"
                    self.save()
                    return True
        return False

    # ------------------------------------------------------------ the library (2.83.0)
    def _rel(self, path):
        """path relative to the effort's data folder, when the file is inside it (a copy kept with the effort)."""
        try:
            base = os.path.normcase(os.path.abspath(self.data_dir))
            p = os.path.normcase(os.path.abspath(path))
            if p.startswith(base + os.sep):
                return os.path.relpath(os.path.abspath(path), os.path.abspath(self.data_dir)).replace(os.sep, "/")
        except (TypeError, ValueError):
            pass
        return ""

    def path_of(self, v):
        """Where a version's original file is now: the copy kept with the effort when there is one, else where it was."""
        if v.get("rel"):                           # the effort's own copy: only that counts, wherever the effort is
            return os.path.join(self.data_dir, *v["rel"].split("/"))
        return v.get("path", "")

    def has_original(self, v):
        p = self.path_of(v)
        return bool(p) and os.path.isfile(p)

    def categories(self):
        """[(key, name)]: the four, then the person's own."""
        return list(CATEGORIES) + [tuple(c) for c in self.extra_categories]

    def category_keys(self):
        return [k for k, _n in self.categories()]

    def category_name(self, key):
        return dict(self.categories()).get(key, key)

    def add_category(self, name):
        name = " ".join((name or "").split())[:40]
        if not name:
            return None
        for k, n in self.categories():
            if n.lower() == name.lower():
                return k
        key = "c_" + re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")[:24]
        self.extra_categories.append([key, name])
        self.save()
        return key

    def in_category(self, key):
        return [d for d in self.all if d.get("category", "customer") == key]

    def folders(self, key):
        """The folders used in a category, sorted, each with its parents."""
        out = set()
        for d in self.in_category(key):
            parts = [x for x in (d.get("folder") or "").split("/") if x]
            for i in range(1, len(parts) + 1):
                out.add("/".join(parts[:i]))
        return sorted(out, key=str.lower)

    def role_of(self, d):
        return d.get("role") or suggest_role(d.get("category", "customer"), d.get("type", ""))

    def set_role(self, doc_id, role, why="", who=None):
        d = self.doc(doc_id)
        if d is None or role not in ROLE_NAMES:
            return False
        d["role"] = role
        d["role_by"] = {"who": who or getpass.getuser(), "when": datetime.datetime.now().isoformat(timespec="minutes"),
                        "why": " ".join((why or "").split())[:200], "suggested": False}
        self.save()
        return True

    def set_category(self, doc_id, key, folder=None):
        """Move a document to another category (a customer document leaving keeps its initial for the matrix's sake)."""
        d = self.doc(doc_id)
        if d is None or key not in self.category_keys():
            return False
        d["category"] = key
        if folder is not None:
            d["folder"] = (folder or "").strip("/\\")
        self.save()
        return True

    def set_tags(self, doc_id, tags):
        d = self.doc(doc_id)
        if d is None:
            return False
        d["tags"] = clean_tags(tags)
        self.save()
        return True

    def all_tags(self):
        seen = {}
        for d in self.all:
            for t in d.get("tags", []):
                seen.setdefault(t.lower(), t)
        return sorted(seen.values(), key=str.lower)

    def undecided(self):
        """Documents whose role is only suggested (nobody confirmed it)."""
        return [d for d in self.all if (d.get("role_by") or {}).get("suggested", True)]

    def search(self, query, docs=None, limit=200):
        """Every document (latest version) whose name, tags or text holds all the words: [(doc, hits)] with hits =
        [(sentence index, sentence)] for up to 3 sentences that hold them all; name and tag matches first."""
        words = [w for w in re.findall(r"[\w&-]+", (query or "").lower()) if w]
        if not words:
            return []
        out = []
        for d in (docs if docs is not None else self.all):
            v = self.latest(d)
            meta = " ".join([d.get("name", ""), d.get("initial", ""), d.get("type", ""), d.get("folder", ""),
                             " ".join(d.get("tags", [])), " ".join(x.get("file", "") for x in d["versions"])]).lower()
            in_meta = all(w in meta for w in words)
            hits = []
            if v is not None:
                key = (self.folder, v["id"])
                low = _TEXTS.get(key)
                if low is None:
                    sents = self.text(v)
                    low = _TEXTS[key] = [(x.get("x") or "") for x in sents]
                for i, x in enumerate(low):
                    xl = x.lower()
                    if all(w in xl for w in words):
                        hits.append((i, x))
                        if len(hits) >= 3:
                            break
            if in_meta or hits:
                out.append((d, hits, in_meta))
            if len(out) >= limit:
                break
        out.sort(key=lambda t: (not t[2], t[0].get("name", "").lower()))
        return [(d, h) for d, h, _m in out]

    def counts_by_category(self):
        c = {k: 0 for k in self.category_keys()}
        for d in self.all:
            c[d.get("category", "customer")] = c.get(d.get("category", "customer"), 0) + 1
        return c

    def counts(self):
        return {"docs": len(self.docs), "versions": sum(len(d["versions"]) for d in self.docs),
                "unreviewed": sum(1 for d in self.docs for v in d["versions"][1:] if not v.get("reviewed"))}


# ------------------------------------------------------------------ compare
def _words(s):
    return s.split()


def word_diff(old, new):
    """[(kind, text)] with kind '=' same, '-' removed, '+' added."""
    a, b = _words(old), _words(new)
    out = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if op == "equal":
            out.append(("=", " ".join(a[i1:i2])))
        if op in ("delete", "replace"):
            out.append(("-", " ".join(a[i1:i2])))
        if op in ("insert", "replace"):
            out.append(("+", " ".join(b[j1:j2])))
    return out


def is_submission(text):
    """Does a sentence set a submission requirement (format, page limits, due dates, files and copies)?"""
    import compliance as C
    return bool(C._FORMAT.search(text or "") or PAGE_LIMIT.search(text or ""))


def compare(old, new, pair_ratio=0.5):
    """old, new: snapshots. Returns [{"kind": same|changed|added|removed, "old", "new", "section", "diff",
    "submission"}] in reading order."""
    A = [s["x"] for s in old]
    B = [s["x"] for s in new]
    rows = []

    def sec(s):
        return " ".join(x for x in (s.get("s", ""), s.get("t", "")) if x)

    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, A, B, autojunk=False).get_opcodes():
        if op == "equal":
            for k in range(i2 - i1):
                rows.append({"kind": "same", "old": A[i1 + k], "new": B[j1 + k], "section": sec(new[j1 + k]), "diff": [],
                             "oi": i1 + k, "nj": j1 + k})
            continue
        olds = list(range(i1, i2))
        news = list(range(j1, j2))
        used, block = set(), []
        for oi in olds:                            # pair each removed sentence with its closest added one
            best, score = None, 0.0
            for nj in news:
                if nj in used:
                    continue
                r = difflib.SequenceMatcher(None, A[oi], B[nj], autojunk=False).ratio()
                if r > score:
                    best, score = nj, r
            if best is not None and score >= pair_ratio:
                used.add(best)
                block.append((best, {"kind": "changed", "old": A[oi], "new": B[best], "section": sec(new[best]),
                                     "diff": word_diff(A[oi], B[best]), "oi": oi, "nj": best}))
            else:                                  # a removed sentence sits where it stood, before what follows it
                block.append((j1 + (oi - i1) - 0.5, {"kind": "removed", "old": A[oi], "new": "", "section": sec(old[oi]), "diff": [],
                                                     "oi": oi, "nj": None}))
        for nj in news:
            if nj not in used:
                block.append((nj, {"kind": "added", "old": "", "new": B[nj], "section": sec(new[nj]), "diff": [], "oi": None, "nj": nj}))
        rows += [r for _k, r in sorted(block, key=lambda x: x[0])]
    for r in rows:
        r["submission"] = r["kind"] != "same" and (is_submission(r["old"]) or is_submission(r["new"]))
    return rows


def summary(rows):
    c = {"changed": 0, "added": 0, "removed": 0, "same": 0, "submission": 0}
    for r in rows:
        c[r["kind"]] += 1
        c["submission"] += 1 if r.get("submission") else 0
    return c


def summary_text(c):
    return (f"{c['changed']} changed, {c['added']} added, {c['removed']} removed" +
            (f"; {c['submission']} {'touches' if c['submission'] == 1 else 'touch'} submission requirements" if c["submission"] else ""))


# ------------------------------------------------------------------ exports
def export_redline(path, doc_name, v_old, v_new, rows, only_changes=True):
    """A Word file: every sentence in reading order, removed words struck through in red, added words green."""
    import docxw
    d = docxw.Doc()
    d.props = {"StationkeeperCompare": f"{doc_name}: {v_old['label']} to {v_new['label']}"}
    d.heading(f"{doc_name}: {v_old['label']} compared with {v_new['label']}", 1)
    d.para(f"{v_old['label']}: {v_old['file']}, added {v_old['added']} by {v_old['who']}, SHA-256 {v_old['sha256']}")
    d.para(f"{v_new['label']}: {v_new['file']}, added {v_new['added']} by {v_new['who']}, SHA-256 {v_new['sha256']}")
    d.para(summary_text(summary(rows)) + f". Built {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')} by Stationkeeper.", bold=True)
    sec = None
    for r in rows:
        if only_changes and r["kind"] == "same":
            continue
        if r["section"] and r["section"] != sec:
            d.heading(r["section"], 3)
            sec = r["section"]
        tag = {"same": "", "changed": "Changed", "added": "Added", "removed": "Removed"}[r["kind"]]
        parts = []
        if r["kind"] == "changed":
            for k, t in r["diff"]:
                st = {"-": {"strike": True, "color": "B91C1C"}, "+": {"bold": True, "color": "15803D"}}.get(k, {})
                parts.append((t + " ", st))
        elif r["kind"] == "added":
            parts.append((r["new"], {"bold": True, "color": "15803D"}))
        elif r["kind"] == "removed":
            parts.append((r["old"], {"strike": True, "color": "B91C1C"}))
        else:
            parts.append((r["new"], {}))
        if tag:
            parts.append(("   [" + tag + (", submission requirement" if r.get("submission") else "") + "]", {"size": 16, "color": "475569"}))
        d.rich(parts)
    d.save(path)
    return path


def export_list(path, lib, doc=None, compares=()):
    """Excel: the version history (every document or one) and, for each (doc, old, new, rows) in compares, a sheet
    of the changes. The trail: dates, who, file names and fingerprints."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook()
    ws = wb.active; ws.title = "Version history"
    ws.append(["Document", "Version", "File", "Added", "Added by", "SHA-256", "Sentences", "Reviewed"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for d in ([doc] if doc else lib.docs):
        for v in d["versions"]:
            ws.append([d["name"], v["label"], v["file"], v["added"], v["who"], v["sha256"], v.get("sentences", ""), v.get("reviewed", "")])
    for col, w in zip("ABCDEFGH", (34, 16, 40, 18, 16, 70, 10, 22)):
        ws.column_dimensions[col].width = w
    for d, vo, vn, rows in compares:
        s = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "", f"{vo['label'][:12]} to {vn['label'][:12]}")[:31])
        s.append([f"{d['name']}: {vo['label']} ({vo['file']}) to {vn['label']} ({vn['file']})"])
        s.append(["Change", "Section", "Before", "After", "Words removed", "Words added", "Submission requirement"])
        for c in s[2]:
            c.font = Font(bold=True)
        for r in rows:
            if r["kind"] == "same":
                continue
            s.append([r["kind"].title(), r["section"], r["old"], r["new"],
                      " ".join(t for k, t in r["diff"] if k == "-"), " ".join(t for k, t in r["diff"] if k == "+"),
                      "Yes" if r.get("submission") else ""])
        for col, w in zip("ABCDEFG", (10, 26, 60, 60, 24, 24, 12)):
            s.column_dimensions[col].width = w
        for row in s.iter_rows(min_row=3):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return path
