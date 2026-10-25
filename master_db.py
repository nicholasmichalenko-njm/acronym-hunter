"""Master database (2.14.0): the company's database of every acronym, loaded from a file the team keeps.

- Load master database (first step of Start an effort, or File > Load master database): reads a CSV or Excel file
  (Term, Category, Definition, Action columns, or a plain two-column acronym list) and merges it into this PC's
  database. Rows the app or the user added on this PC stay. Rows from the previous master that the new version no
  longer has are removed; new rows are added; the same row is left alone. Each load is a numbered version with its
  date, file name, row count and SHA-256, kept in config\\master_database.json; a copy of the file is kept in
  config\\master.
- Known in: every result says where the term is known: the customer documents, the customer's list, the master
  database, or nowhere.
- Update the results: after a new version, the results are matched again against it (from the text already read
  when the session still has it, else the files are read again). Nothing is dropped: rows the update does not
  find again stay, changed rows say "Database update", and decisions, picked definitions and the effort list stay.
"""
import datetime, hashlib, json, os, shutil
import tkinter as tk
from tkinter import filedialog, messagebox
import customtkinter as ctk
import diag

ORIGIN_COL = "Origin"
ORIGIN_MASTER = "Master database"
SRC_DOCS, SRC_LIST, SRC_MASTER = "Customer docs", "Customer list", "Master database"
KNOWN_FILTERS = ["Known in: customer docs", "Known in: customer list", "Known in: master database only",
                 "Known in: no source"]
CHANGE_DB = "Database update"


def _eng():
    import app_gui
    return app_gui


def info_path():
    return os.path.join(_eng().CONFIG_DIR, "master_database.json")


def master_info():
    try:
        with open(info_path(), encoding="utf-8") as f:
            d = json.load(f)
        return d if isinstance(d, dict) else {}
    except (OSError, ValueError):
        return {}


def _sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def read_master_file(path):
    """[[term, category, definition, action, {}]] from a CSV or Excel database, or from a two-column list."""
    eng = _eng()
    ext = os.path.splitext(path)[1].lower()
    rows = []
    if ext == ".csv":
        for raw in eng.read_csv_rows(path):
            d = eng.db_row(raw)
            if d.get("Term"):
                rows.append([d["Term"].strip(), d.get("Category") or "Acronym", (d.get("Definition") or "").strip(),
                             d.get("Action") or "Include", {}])
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                it = ws.iter_rows(values_only=True)
                head = [str(h or "").strip() for h in next(it, [])]
                idx = {eng.db_field(h): i for i, h in enumerate(head) if eng.db_field(h)}
                if "Term" not in idx:
                    continue
                for vals in it:
                    vals = list(vals or [])
                    get = lambda f: str(vals[idx[f]] if f in idx and idx[f] < len(vals) and vals[idx[f]] is not None else "").strip()
                    if get("Term"):
                        rows.append([get("Term"), get("Category") or "Acronym", get("Definition"), get("Action") or "Include", {}])
        finally:
            wb.close()
    if not rows:                               # a plain acronym list (any format the client-list reader takes)
        rows = [[t, "Acronym", d, "Include", {}] for t, d in eng.parse_client_list(path) if t]
    return rows


def _key(r):
    eng = _eng()
    kind = eng._action_kind(r[3])
    return (eng.fold(r[0]).strip(), "" if kind == "exclude" else eng.norm_def(r[2]), kind)


def merge_master(new_rows):
    """Merge a master database into this PC's database. Returns {"added", "removed", "same", "total"}."""
    eng = _eng()
    cur = eng.read_db_rows()
    want = {}
    for r in new_rows:
        want.setdefault(_key(r), r)
    out, have, removed = [], set(), 0
    for r in cur:
        extra = r[4] if len(r) > 4 and isinstance(r[4], dict) else {}
        k = _key(r)
        if extra.get(ORIGIN_COL) == ORIGIN_MASTER and k not in want:
            removed += 1
            continue
        out.append(r); have.add(k)
    added = 0
    for k, r in want.items():
        if k in have:
            continue
        out.append([r[0], r[1], r[2], r[3], {ORIGIN_COL: ORIGIN_MASTER}])
        have.add(k); added += 1
    eng.write_db_rows(out)
    return {"added": added, "removed": removed, "same": len(want) - added, "total": len(want)}


def load_master(path):
    """Read, merge and record a new master version. Returns the version info with the merge counts."""
    rows = read_master_file(path)
    if not rows:
        raise ValueError("no terms found")
    counts = merge_master(rows)
    eng = _eng()
    old = master_info()
    ver = int(old.get("version", 0)) + 1
    keep = os.path.join(eng.CONFIG_DIR, "master")
    os.makedirs(keep, exist_ok=True)
    try:
        shutil.copy2(path, os.path.join(keep, f"v{ver}_{os.path.basename(path)}"))
    except OSError:
        diag.swallowed()
    info = {"version": ver, "file": os.path.basename(path), "loaded": datetime.datetime.now().isoformat(timespec="minutes"),
            "rows": counts["total"], "sha256": _sha256(path), **{k: counts[k] for k in ("added", "removed", "same")}}
    tmp = info_path() + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(info, f, indent=1)
    os.replace(tmp, info_path())
    return info


def info_text(info=None):
    info = master_info() if info is None else info
    if not info:
        return "No master database loaded on this PC yet."
    return (f"Master database version {info['version']}: {info.get('file', '')}, {info.get('rows', 0)} entries, "
            f"loaded {info.get('loaded', '').replace('T', ' ')}.")


def known_text(srcs):
    """Plain words for the Known in column."""
    s = set(srcs or [])
    if SRC_DOCS in s and SRC_LIST in s:
        t = "Customer docs and customer list"
    elif SRC_DOCS in s:
        t = "Customer docs"
    elif SRC_LIST in s:
        t = "Customer list"
    elif SRC_MASTER in s:
        return "Master database only"
    else:
        return "Not found in any source"
    return t + (" (also in master database)" if SRC_MASTER in s else "")


class MasterMixin:
    def db_terms(self, refresh=False):
        """Folded terms the database knows (include and omit rows), read again when the file changes."""
        eng = _eng()
        try:
            mt = os.path.getmtime(eng.DB_FILE)
        except OSError:
            mt = 0
        if refresh or getattr(self, "_db_terms_mt", None) != (eng.DB_FILE, mt):
            try:
                self._db_terms = set(eng.load_database()[0])
            except Exception:
                diag.swallowed(); self._db_terms = set()
            self._db_terms_mt = (eng.DB_FILE, mt)
        return self._db_terms

    def known_in(self, term, d):
        eng = _eng()
        out = []
        if d.get("customer_definitions"):
            out.append(SRC_DOCS)
        if d.get("client"):
            out.append(SRC_LIST)
        if eng.fold(term).strip() in self.db_terms():
            out.append(SRC_MASTER)
        return out

    def mark_known_in(self):
        for t, d in self.results.items():
            if not d.get("file_error"):
                d["known_in"] = self.known_in(t, d)

    def known_cell(self, term, d):
        if d.get("file_error"):
            return ""
        if "known_in" not in d:
            d["known_in"] = self.known_in(term, d)
        return known_text(d["known_in"])

    @staticmethod
    def known_matches(d, wanted):
        s = set(d.get("known_in") or [])
        return {"Known in: customer docs": SRC_DOCS in s,
                "Known in: customer list": SRC_LIST in s,
                "Known in: master database only": s == {SRC_MASTER},
                "Known in: no source": not s and not d.get("file_error")}.get(wanted)

    # ------------------------------------------------------------ load and update
    def master_load(self, path=None, ask=True):
        """Load a (new) master database. Returns the version info, or None."""
        path = path or filedialog.askopenfilename(title="Master database", filetypes=[("Database", "*.csv *.xlsx"),
                                                                                       ("Acronym lists", "*.docx *.pdf *.txt"),
                                                                                       ("All files", "*.*")])
        if not path:
            return None
        try:
            info = load_master(path)
        except Exception as ex:
            diag.swallowed()
            if ask:
                messagebox.showerror("Master database", f"Could not load {os.path.basename(path)} ({type(ex).__name__}). "
                                     "Is it open in Excel? Close it and try again.")
            return None
        diag.activity("master_loaded", rows=info["rows"], added=info["added"], removed=info["removed"])
        self.db_terms(refresh=True)
        self.update_master_label()
        text = (f"Master database version {info['version']} loaded: {info['added']} new, {info['removed']} removed, "
                f"{info['same']} already there.")
        self.status.configure(text=text)
        if ask and self.results and not self.busy() and messagebox.askyesno(
                "New master database", text + "\n\nUpdate the results against it now? Your results stay: rows the new "
                "database changes are updated and marked Database update, and your decisions and effort list are kept."):
            self.refresh_from_database()
        return info

    def update_master_label(self):
        if hasattr(self, "lbl_aistate") and hasattr(self, "ai_state_line"):
            self.ai_state_line()

    def refresh_from_database(self):
        """Match the current results again against the database without dropping any row."""
        eng = _eng()
        if self.busy():
            return False
        if self.project is not None:
            self._db_refresh = self._snapshot()
            self.rescan_project()
            return True
        last = getattr(self, "last_scan", None) or {}
        fps = [f for f in last.get("paths", []) if os.path.isfile(f)]
        if not fps:
            self.mark_known_in(); self.populate_table()
            self.status.configure(text="Database loaded. No scanned files to update; the next scan uses it.")
            return False
        self._db_refresh = self._snapshot()
        cache = last.get("cache") or {}
        self.queued_files = fps
        diag.activity("database_refresh", files=len(fps))

        def job(cb):
            return eng.refresh_batch(fps, cb, cache)
        self.run_job(job)
        return True

    def _snapshot(self):
        return {"results": dict(self.results),
                "keys": {t: (d.get("status"), d.get("selected_definition"), tuple(d.get("definitions", [])),
                             bool(d.get("in_db")), bool(d.get("excluded")), tuple(d.get("known_in") or ()))
                         for t, d in self.results.items()}}

    def after_db_refresh(self):
        """In scan_done after a database refresh: keep rows not found again, mark what the database changed."""
        snap = getattr(self, "_db_refresh", None)
        if not snap:
            return None
        self._db_refresh = None
        kept = changed = new = 0
        for t, d in snap["results"].items():
            if t not in self.results:
                self.results[t] = d; kept += 1
        for t, d in self.results.items():
            old = snap["keys"].get(t)
            k = (d.get("status"), d.get("selected_definition"), tuple(d.get("definitions", [])), bool(d.get("in_db")),
                 bool(d.get("excluded")), tuple(d.get("known_in") or ()))
            if old is None:
                new += 1; d["change"] = "New"
            elif old != k:
                changed += 1; d["change"] = CHANGE_DB
        diag.activity("database_refreshed", n=changed, added=new)
        return {"changed": changed, "new": new, "kept": kept}

    def master_step(self, body, win):
        """The first step of Start an effort: load the master database so every result can be checked against it."""
        ctk.CTkLabel(body, text="Load the master database first", anchor="w", font=("Segoe UI", 13, "bold")).pack(fill="x", pady=(6, 2))
        ctk.CTkLabel(body, text=("Every result is checked against the customer documents, the customer's own list and the "
                                 "company's master database, and says where it was found: customer docs, customer list, "
                                 "master database only, or not found in any source. The master database is shared by every "
                                 "effort on this PC."), anchor="w", justify="left", wraplength=660).pack(fill="x", pady=(0, 8))
        state = ctk.CTkLabel(body, text=info_text(), anchor="w", justify="left", wraplength=660, font=("Segoe UI", 13))
        state.pack(fill="x", pady=4)

        def load(path=None):
            info = self.master_load(path, ask=path is None)
            if info:
                state.configure(text=info_text(info) + f"  ({info['added']} new, {info['removed']} removed, {info['same']} already there)")
            return info
        ctk.CTkButton(body, text=("Load a newer master database..." if master_info() else "Load master database..."), width=230,
                      command=load).pack(anchor="w", pady=6)
        ctk.CTkLabel(body, text=("CSV or Excel (Term, Category, Definition, Action), or a plain acronym list. Load a newer "
                                 "version any time: File > Load master database. Rows added on this PC stay."),
                     anchor="w", justify="left", wraplength=660, text_color=("#475569", "#94A3B8")).pack(fill="x")
        win._load_master = load
