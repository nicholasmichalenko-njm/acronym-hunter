"""Effort start wizard and the effort record (2.14.0).

- Wizard (first start, top bar New effort..., Effort > New effort): effort and client (short name, solicitation number,
  internal number), the customer's own acronym lists (remembered by the effort and used on every scan), customer
  documents (copied into the effort folder and read for a jump start), how the list ships (whole effort, per volume,
  both; first-use rule), ready.
- Customer document intake any time (Effort > Add customer documents): copies the files into the effort's
  Customer documents folder, reads only new or changed files, and opens the review of new terms.
- Review new terms: one click adds the selected terms to the effort list as Proposed, with the documents they were
  found in; or puts them on the omit list, or marks them Not a term in this effort.
- Every scan updates the effort list: Found in, First found, Last found, Scan status (New in this scan, Found,
  No longer found in <file>); terms no longer found show in the results as rows to decide on.
- Activity: the effort list's History sheet, newest first.
- Exports: the effort's acronym list for the whole effort, per volume, or both (Word and Excel, with Found in), and
  every effort in one workbook.
"""
import datetime, os, re, shutil, threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import customtkinter as ctk
import diag
import effort as E

LIST_FORMATS = {"whole": "One list for the whole effort", "volume": "One list per volume", "both": "Both"}
CUSTOMER_DIR = "Customer documents"
DOC_EXTS = (".docx", ".pptx", ".xlsx", ".pdf")


def _eng():
    import app_gui
    return app_gui


def _safe(name):
    return re.sub(r"[^A-Za-z0-9 _.-]+", "", name).strip() or "Effort"


def files_of(d, file_of):
    """File names a result was found in."""
    return {file_of(l) for l in d.get("locations", []) if l and not l.startswith(("Added manually",))}


class WizardMixin:
    # ------------------------------------------------------------ wizard
    def open_wizard(self):
        diag.activity("wizard_opened")
        eng = _eng()
        win = ctk.CTkToplevel(self); win.title("Start an effort"); win.geometry("720x560"); win.transient(self)
        win.after(150, win.lift)
        diag.register(win, "wizard")
        st = {"i": 0, "lists": [], "docs": []}
        desk = os.path.join(os.path.expanduser("~"), "Desktop")
        base = desk if os.path.isdir(desk) else os.path.expanduser("~")
        head = ctk.CTkLabel(win, text="", font=("Segoe UI", 15, "bold"), anchor="w", justify="left", wraplength=680)
        head.pack(fill="x", padx=18, pady=(14, 2))
        body = ctk.CTkFrame(win, fg_color="transparent"); body.pack(fill="both", expand=True, padx=18, pady=4)
        msg = ctk.CTkLabel(win, text="", text_color="#B91C1C", anchor="w", wraplength=680, justify="left"); msg.pack(fill="x", padx=18)
        nav = ctk.CTkFrame(win, fg_color="transparent"); nav.pack(fill="x", padx=18, pady=12)
        b_back = ctk.CTkButton(nav, text="Back", width=90, fg_color="#475569"); b_back.pack(side="left")
        b_next = ctk.CTkButton(nav, text="Next", width=120); b_next.pack(side="right")
        ctk.CTkButton(nav, text="Cancel", width=90, fg_color="transparent", text_color=("#334155", "#CBD5E1"),
                      command=win.destroy).pack(side="right", padx=8)
        v = {k: tk.StringVar() for k in ("name", "client", "sol", "internal", "folder", "volumes", "due", "method", "contact", "team")}
        v["format"] = tk.StringVar(value="whole"); v["rule"] = tk.StringVar(value="volume"); v["scan"] = tk.BooleanVar(value=True)

        def entry(parent, label, var, hint=""):
            ctk.CTkLabel(parent, text=label, anchor="w").pack(fill="x", pady=(8, 0))
            e = ctk.CTkEntry(parent, textvariable=var, width=640, placeholder_text=hint); e.pack(anchor="w")
            return e

        def on_name(*_):
            n = _safe(v["name"].get())
            if not st.get("folder_set"):
                v["folder"].set(os.path.join(base, n) if v["name"].get().strip() else "")
        v["name"].trace_add("write", on_name)

        def step_master():
            head.configure(text="1 of 7   " + __import__("master_db").MASTER_NAME)
            self.master_step(body, win)

        def step0():
            head.configure(text="2 of 7   Effort and client")
            e = entry(body, "Effort name (required)", v["name"], "e.g. Kestrel Thermal")
            entry(body, "Client (short name)", v["client"], "e.g. USSF SSC")
            entry(body, "Solicitation or opportunity number", v["sol"], "e.g. FA8810-27-R-0042")
            entry(body, "Internal number (optional, your own numbering)", v["internal"])
            ctk.CTkLabel(body, text="Due date and time (optional; the top bar counts down to it)", anchor="w").pack(fill="x", pady=(8, 0))
            drow = ctk.CTkFrame(body, fg_color="transparent"); drow.pack(anchor="w")
            ctk.CTkEntry(drow, textvariable=v["due"], width=520, placeholder_text="Pick a date, then add the time: 2:00 p.m. MT").pack(side="left")

            def pick_due():                                  # 2.52.0: a calendar instead of a date format
                import datepick as DP, effortrec as R_
                cur = v["due"].get()
                rest = ", ".join(cur.split(", ")[2:]) if R_.parse_date(", ".join(cur.split(", ")[:2])) else ""
                def got(d):
                    v["due"].set(DP.nice(d.isoformat()) + (", " + rest if rest else ", 2:00 p.m."))
                return DP.CalendarPopup(drow, R_.parse_date(", ".join(cur.split(", ")[:2])) or __import__("datetime").date.today(), got)
            ctk.CTkButton(drow, text="Pick a date", width=100, command=pick_due).pack(side="left", padx=6)
            st["pick_due"] = pick_due
            entry(body, "How the proposal is submitted (optional; read from the customer documents too)", v["method"],
                  "e.g. DoD SAFE, PIEE, e-mail")                                                        # 2.49.0
            entry(body, "Submission contact (optional)", v["contact"], "e.g. Jane Doe, Contracting Officer, jane.doe@us.af.mil")
            entry(body, "Team (optional; names separated by commas; offered whenever you pick a lead, author or owner)", v["team"],
                  "e.g. J. Smith, K. Park, N. Michalenko")                                              # 2.52.0
            row = ctk.CTkFrame(body, fg_color="transparent"); row.pack(fill="x", pady=(8, 0))
            ctk.CTkLabel(row, text="Effort folder (holds the list, customer documents and saved work)", anchor="w").pack(fill="x")
            ctk.CTkEntry(row, textvariable=v["folder"], width=540).pack(side="left")
            def pick():
                d = filedialog.askdirectory(parent=win)
                if d:
                    st["folder_set"] = True; v["folder"].set(os.path.join(d, _safe(v["name"].get())) if v["name"].get().strip() else d)
            ctk.CTkButton(row, text="Browse...", width=90, command=pick).pack(side="left", padx=6)
            e.focus_set()

        def files_step(key, title, ask, note, exts):
            head.configure(text=title)
            ctk.CTkLabel(body, text=ask, anchor="w", font=("Segoe UI", 13, "bold"), wraplength=660, justify="left").pack(fill="x", pady=(6, 4))
            lb = tk.Listbox(body, height=9, font=("Segoe UI", 11), activestyle="none")
            lb.pack(fill="both", expand=True)
            def refresh():
                lb.delete(0, "end")
                for item in st[key]:
                    lb.insert("end", item["label"])
            def add():
                ps = filedialog.askopenfilenames(parent=win, filetypes=[("Documents", " ".join("*" + x for x in exts))])
                for p in ps:
                    if key == "lists":
                        try:
                            rows = eng.parse_client_list(p)
                        except Exception:
                            diag.swallowed(); rows = []
                        if not rows:
                            msg.configure(text=f"No terms found in {os.path.basename(p)}. Paste or type... takes a list copied from anywhere.")
                            continue
                        note = ("Read from pictures with the picture reader: check every line."
                                if os.path.splitext(p)[1].lower() in eng.CLIENT_IMAGE_EXTS else "")
                        rows = self._preview_client_list(os.path.basename(p), rows, note=note, parent=win)
                        if not rows:
                            continue
                        st[key].append({"path": p, "rows": rows, "label": f"{os.path.basename(p)}   ({len(rows)} terms)"})
                    else:
                        st[key].append({"path": p, "label": os.path.basename(p)})
                refresh()
            def remove():
                for i in reversed(lb.curselection()):
                    st[key].pop(i)
                refresh()
            row = ctk.CTkFrame(body, fg_color="transparent"); row.pack(fill="x", pady=6)
            ctk.CTkButton(row, text="Add files...", width=110, command=add).pack(side="left")
            if key == "lists":                   # paste or type a list, then confirm (2.28.0)
                def take(name, rows):
                    st[key].append({"path": None, "name": name, "rows": rows, "label": f"{name}   ({len(rows)} terms)"})
                    refresh()
                ctk.CTkButton(row, text="Paste or type...", width=130, fg_color="#6D28D9",
                              command=lambda: self.client_paste_window(take, parent=win)).pack(side="left", padx=6)
                win._paste_take = take
            ctk.CTkButton(row, text="Remove", width=90, fg_color="#475569", command=remove).pack(side="left", padx=6)
            ctk.CTkLabel(body, text=note, anchor="w", wraplength=660, justify="left", text_color=("#475569", "#94A3B8")).pack(fill="x")
            refresh()
            win._add_files = lambda paths: ([st[key].append({"path": p, "rows": eng.parse_client_list(p),
                                                             "label": os.path.basename(p)} if key == "lists" else
                                                            {"path": p, "label": os.path.basename(p)}) for p in paths], refresh())

        def step1():
            files_step("lists", "3 of 7   The customer's own lists",
                       "Do you have known or provided customer acronyms, abbreviations or units of measure?",
                       "Excel, Word, PDF, CSV or text. The effort remembers them and every scan uses them; the customer's "
                       "definition comes before the company database. Skip if you have none yet (Client Lists adds them later).",
                       (".xlsx", ".docx", ".pdf", ".csv", ".tsv", ".txt"))

        def step2():
            files_step("docs", "4 of 7   Customer documents",
                       "Add the customer documents you have now (RFP sections, SOW, CDRLs, attachments).",
                       "They are copied into the effort folder's Customer documents. Their own definitions are used on every "
                       "scan. Add more any time with Effort > Add customer documents; only new files are read.",
                       DOC_EXTS)
            ctk.CTkCheckBox(body, text="Read them when the effort is created and show the new terms for review",
                            variable=v["scan"]).pack(anchor="w", pady=(6, 0))

        def step_structure():
            """The response structure: suggested from the customer documents, or typed (2.28.0)."""
            head.configure(text="5 of 7   Response structure")
            ctk.CTkLabel(body, text="How will the response be organized?", anchor="w", font=("Segoe UI", 13, "bold")).pack(fill="x", pady=(6, 2))
            ctk.CTkLabel(body, anchor="w", justify="left", wraplength=660, text=(
                "One section per line: Volume I Technical, then 1.1 Technical Approach (indent or number the sections). "
                "Suggested from the customer documents when they say how to organize the proposal; edit it or type your own. "
                "The compliance matrix and Setup start from it.")).pack(fill="x")
            box = ctk.CTkTextbox(body, height=190, wrap="none", font=("Consolas", 12)); box.pack(fill="both", expand=True, pady=6)
            box.insert("1.0", st.get("outline", ""))
            notes = ctk.CTkLabel(body, text="\n".join(st.get("outline_notes", [])), anchor="w", justify="left", wraplength=660,
                                 text_color=("#0F766E", "#5EEAD4"))
            notes.pack(fill="x")

            def suggest():
                if not st["docs"]:
                    notes.configure(text="Add customer documents in step 4 to get a suggestion, or type the structure.")
                    return None
                import compliance as C, structure as S
                fc = []
                for item in st["docs"]:
                    try:
                        fc.append((os.path.basename(item["path"]), eng.typed_chunks(item["path"])))
                    except Exception:
                        diag.swallowed()
                heads = []
                try:
                    C.extract(fc, heads, every=False)
                except Exception:
                    diag.swallowed()
                res = S.suggest(fc, heads)
                st["suggested"] = True
                st["outline_notes"] = res["notes"]
                if res["text"]:
                    box.delete("1.0", "end"); box.insert("1.0", res["text"])
                notes.configure(text="\n".join(res["notes"]))
                diag.activity("structure_suggested", n=res["sections"], added=res["volumes"])
                return res
            row = ctk.CTkFrame(body, fg_color="transparent"); row.pack(fill="x")
            ctk.CTkButton(row, text="Suggest from the customer documents", width=260, command=suggest).pack(side="left")
            ctk.CTkButton(row, text="Clear", width=80, fg_color="#475569",
                          command=lambda: box.delete("1.0", "end")).pack(side="left", padx=6)
            st["_save"] = lambda: st.update(outline=box.get("1.0", "end").strip())
            win._suggest = suggest
            win._outline_box = box
            if st["docs"] and not st.get("suggested") and not st.get("outline"):
                suggest()

        def step3():
            head.configure(text="6 of 7   How the list ships")
            ctk.CTkLabel(body, text="Acronym list", anchor="w", font=("Segoe UI", 13, "bold")).pack(fill="x", pady=(6, 2))
            for k, lab in LIST_FORMATS.items():
                ctk.CTkRadioButton(body, text=lab, variable=v["format"], value=k).pack(anchor="w", pady=2)
            ctk.CTkLabel(body, text="First use", anchor="w", font=("Segoe UI", 13, "bold")).pack(fill="x", pady=(12, 2))
            for k, lab in E.FIRST_USE_RULES.items():
                ctk.CTkRadioButton(body, text=lab, variable=v["rule"], value=k).pack(anchor="w", pady=2)
            entry(body, "Volumes folder (optional, for checking deliverables)", v["volumes"])

        def step4():
            head.configure(text="7 of 7   Ready")
            lines = [f"Effort:  {v['name'].get().strip()}",
                     f"Client:  {v['client'].get().strip() or '-'}      Solicitation:  {v['sol'].get().strip() or '-'}"
                     + (f"      Internal:  {v['internal'].get().strip()}" if v['internal'].get().strip() else ""),
                     f"Folder:  {v['folder'].get().strip()}",
                     "Master:  " + __import__("master_db").info_text(),
                     f"Customer lists:  {len(st['lists'])} ({sum(len(x['rows']) for x in st['lists'])} terms)",
                     f"Customer documents:  {len(st['docs'])}" + ("  (read now, then review)" if st["docs"] and v["scan"].get() else ""),
                     "Response structure:  " + (f"{len([l for l in st.get('outline', '').splitlines() if l.strip()])} lines"
                                                if st.get("outline") else "none yet (add it later in the compliance matrix)"),
                     f"List:  {LIST_FORMATS[v['format'].get()]};  first use {E.FIRST_USE_RULES[v['rule'].get()].lower()}"]
            ctk.CTkLabel(body, text="\n".join(lines), anchor="w", justify="left", font=("Segoe UI", 13)).pack(fill="x", pady=8)
            ctk.CTkLabel(body, text="Everything here can be changed later in Effort > Effort settings.", anchor="w",
                         text_color=("#475569", "#94A3B8")).pack(fill="x")

        steps = [step_master, step0, step1, step2, step_structure, step3, step4]

        def show():
            st.pop("_save", lambda: None)()          # keep what was typed on the step being left
            for w in body.winfo_children():
                w.destroy()
            msg.configure(text="")
            steps[st["i"]]()
            b_back.configure(state="normal" if st["i"] else "disabled")
            b_next.configure(text="Create effort" if st["i"] == len(steps) - 1 else "Next")

        def nxt():
            if st["i"] == 1:
                name = v["name"].get().strip()
                if not name:
                    msg.configure(text="Give the effort a name."); return
                if name in self.efforts.names():
                    msg.configure(text="An effort with that name already exists. Pick it in the Effort box at the top."); return
                if not v["folder"].get().strip():
                    msg.configure(text="Pick a folder for the effort."); return
            if st["i"] < len(steps) - 1:
                st["i"] += 1; show(); return
            finish()

        def back():
            if st["i"]:
                st["i"] -= 1; show()

        def finish():
            name = v["name"].get().strip()
            folder = v["folder"].get().strip()
            try:
                os.makedirs(folder, exist_ok=True)
                cdir = os.path.join(folder, CUSTOMER_DIR)
                os.makedirs(cdir, exist_ok=True)
                path = os.path.join(folder, f"{_safe(name).replace(' ', '_')}_Effort_List.xlsx")
                L = E.EffortList(path, name, v["rule"].get(), client=v["client"].get().strip(), solicitation=v["sol"].get().strip(),
                                 internal_no=v["internal"].get().strip())
                if os.path.isfile(path):
                    L.load(); L.name, L.client, L.solicitation, L.internal_no = name, v["client"].get().strip(), v["sol"].get().strip(), v["internal"].get().strip()
                L.pending = True
                L.save()
            except E.ListLocked:
                msg.configure(text="The list file is open in Excel. Close it, then click Create effort again."); return
            except OSError as ex:
                diag.swallowed()
                msg.configure(text=f"Could not create the effort folder ({type(ex).__name__})."); return
            copied = []
            for item in st["docs"]:
                try:
                    dest = os.path.join(cdir, os.path.basename(item["path"]))
                    if os.path.abspath(dest) != os.path.abspath(item["path"]):
                        shutil.copy2(item["path"], dest)
                    copied.append(dest)
                except OSError:
                    diag.swallowed()
            self.efforts.put(name, path=path, first_use=v["rule"].get(), client=v["client"].get().strip(),
                             solicitation=v["sol"].get().strip(), internal_no=v["internal"].get().strip(),
                             list_format=v["format"].get(), customer_dir=cdir, volumes_dir=v["volumes"].get().strip())
            win.destroy()
            if not self.open_effort(name):
                return
            if hasattr(self, "effrec"):                  # 2.40.0: the effort record starts here
                try:
                    rec = self.effrec()
                    rec.log("effort", "Effort started: " + "; ".join(x for x in (v["client"].get().strip(), v["sol"].get().strip(),
                                                                                  v["internal"].get().strip()) if x))
                    iso, zone = __import__("effortrec").parse_due_text(v["due"].get()) if v["due"].get().strip() else ("", "")
                    if iso:
                        rec.set_due(iso, zone)
                    for nm in [x.strip() for x in v["team"].get().split(",") if x.strip()]:     # 2.52.0: the team
                        rec.add_member(nm)
                    rec.save()
                    reg = self.subreg() if hasattr(self, "subreg") else None         # 2.49.0: how and to whom, on Submission
                    if reg is not None:
                        if v["method"].get().strip():
                            reg.add("Delivery", "Delivery", v["method"].get())
                        if v["contact"].get().strip():
                            reg.add("Contact", "Submission contact", v["contact"].get())
                except Exception:
                    diag.swallowed()
            if st.get("outline"):                    # the response structure (2.28.0): the compliance matrix starts from it
                self.ws.set("outline_text", st["outline"])
                self.effort._hist("(structure)", f"response structure: {len(st['outline'].splitlines())} lines")
            if st["lists"]:
                lists = [{"name": x.get("name") or os.path.basename(x["path"]), "rows": x["rows"]} for x in st["lists"]]
                self.ws.set("client_lists", lists)
                self.set_client_lists(lists)
                self.effort._hist("(customer lists)", f"{sum(len(x['rows']) for x in lists)} customer terms from {len(lists)} list(s)")
                self._save_effort_list()
            if copied and hasattr(self, "sub_plan"):  # 2.51.0: the documents are in Customer Documents at once
                try:
                    self.sub_apply(self.sub_plan(copied))
                except Exception:
                    diag.swallowed()
            diag.activity("wizard_finished", files=len(copied), n=len(st["lists"]))
            if copied and v["scan"].get():
                self.intake(copied)
            else:
                self.status.configure(text=f"Effort {name} created. Add customer documents any time: Effort > Add customer documents.")

        b_back.configure(command=back); b_next.configure(command=nxt)
        show()
        win._api = {"v": v, "st": st, "next": nxt, "back": back, "finish": finish, "msg": msg}
        return win

    # ------------------------------------------------------------ customer documents, any time
    def add_customer_documents(self, paths=None):
        if self.effort is None:
            messagebox.showinfo("Customer documents", "Open or create an effort first.")
            return
        if hasattr(self, "sub_add_dialog") and paths is None:     # 2.51.0: customer documents are added in Customer Documents
            return self.sub_add_dialog()
        paths = paths or filedialog.askopenfilenames(title="Customer documents for this effort",
                                                     filetypes=[("Documents", " ".join("*" + x for x in DOC_EXTS))])
        if not paths:
            return
        info = self.efforts.get(self.effort.name) or {}
        cdir = info.get("customer_dir") or os.path.join(os.path.dirname(self.effort.path), CUSTOMER_DIR)
        os.makedirs(cdir, exist_ok=True)
        if info.get("customer_dir") != cdir:
            self.efforts.put(self.effort.name, customer_dir=cdir)
        copied = []
        for p in paths:
            dest = os.path.join(cdir, os.path.basename(p))
            try:
                if os.path.abspath(dest) != os.path.abspath(p):
                    shutil.copy2(p, dest)
                copied.append(dest)
            except OSError:
                diag.swallowed()
        diag.activity("customer_docs_added", files=len(copied))
        self.intake(copied)

    def _intake_new(self, paths):
        """Only files not read before (or changed since)."""
        seen = self.ws.get("intake", {}) or {}
        out = []
        for p in paths:
            try:
                st = os.stat(p)
            except OSError:
                continue
            key = os.path.basename(p)
            if seen.get(key) != [st.st_size, int(st.st_mtime)]:
                out.append(p)
        return out

    def intake(self, paths):
        """Read customer documents in the background, then open the review of terms not yet in the effort list."""
        eng = _eng()
        todo = self._intake_new(paths)
        if not todo:
            self.status.configure(text="Those customer documents were read before and have not changed.")
            return
        if self.busy():
            messagebox.showwarning("Busy", "Wait for the scan to finish, then add the documents again.")
            return
        for b in (self.btn_files, self.btn_folder, self.btn_project, self.btn_again):
            b.configure(state="disabled")
        self.status.configure(text=f"Reading {len(todo)} customer document(s)...")

        def work():
            try:
                res = eng.analyze_batch(todo, lambda m: self.after(0, lambda: self.status.configure(text=str(m)[:120])))
            except Exception:
                eng.log_error("Customer document intake failed")
                res = {}
            self.after(0, lambda: self._intake_done(todo, res))
        threading.Thread(target=work, daemon=True).start()

    def _intake_done(self, todo, res):
        eng = _eng()
        self.reset_ui()
        seen = self.ws.get("intake", {}) or {}
        for p in todo:
            try:
                st = os.stat(p); seen[os.path.basename(p)] = [st.st_size, int(st.st_mtime)]
            except OSError:
                pass
        self.ws.set("intake", seen)
        try:
            import custdefs
            defs = self.customer_definitions()
            custdefs.merge_into(res, defs, folded=lambda t: eng.fold(t).strip())
        except Exception:
            diag.swallowed()
        if self.effort is not None:
            found = {t: files_of(d, eng._file_of) for t, d in res.items() if not d.get("file_error")}
            self.effort.record_scan(found, [os.path.basename(p) for p in todo])
            self._save_effort_list()
        new = {t: d for t, d in res.items() if not d.get("file_error") and not d.get("excluded") and not d.get("manual")
               and self.effort is not None and t not in self.effort.rows}
        diag.activity("intake_done", files=len(todo), terms=len(new))
        self.status.configure(text=f"Read {len(todo)} customer document(s): {len(new)} term(s) not yet in the effort list.")
        if new:
            return self.review_new_terms(new, title=f"New terms from {len(todo)} customer document(s)")

    # ------------------------------------------------------------ one-click review
    def review_new_terms(self, new=None, title="New terms (not in the effort list)"):
        """Terms not in the effort list yet, with the documents they were found in. One click adds them as Proposed."""
        eng = _eng()
        if new is None:
            new = {t: d for t, d in self.results.items() if d.get("new_in_effort") and self.effort is not None and t not in self.effort.rows}
        if self.effort is None or not new:
            messagebox.showinfo("Review new terms", "No new terms: everything found is already in the effort list.")
            return None
        win = ctk.CTkToplevel(self); win.title(title); win.geometry("1100x620"); win.transient(self); win.after(150, win.lift)
        diag.register(win, "review_new")
        ctk.CTkLabel(win, text=f"{len(new)} term(s) found that the effort list does not have yet. Select, then add them as Proposed "
                               "(they keep the documents they came from), put them on the omit list, or mark them Not a term.",
                     anchor="w", wraplength=1060, justify="left").pack(fill="x", padx=12, pady=(10, 4))
        tv = ttk.Treeview(win, columns=("term", "type", "def", "src", "found"), show="headings", selectmode="extended")
        for c, h, w in (("term", "Term", 110), ("type", "Type", 120), ("def", "Definition", 330), ("src", "Definition from", 140),
                        ("found", "Found in", 380)):
            tv.heading(c, text=h); tv.column(c, width=w, stretch=(c in ("def", "found")))
        ys = ttk.Scrollbar(win, orient="vertical", command=tv.yview); tv.configure(yscrollcommand=ys.set)
        ys.pack(side="right", fill="y", padx=(0, 12)); tv.pack(fill="both", expand=True, padx=(12, 0))
        rows = {}
        for t in sorted(new, key=str.lower):
            d = new[t]
            defn = self._definition_for(d)
            rows[t] = d
            tv.insert("", "end", iid=t, values=(t, eng.base_type(d.get("category", "")), defn or "(needs a definition)",
                                                self._source_for(d, defn) or "-", "; ".join(sorted(files_of(d, eng._file_of)))))
        msg = ctk.CTkLabel(win, text="", anchor="w", text_color=("#047857", "#34D399")); msg.pack(fill="x", padx=12)

        def pick():
            return [t for t in tv.selection() if tv.exists(t)] or []

        def add(ts=None):
            ts = ts if ts is not None else pick()
            n = 0
            for t in ts:
                d = rows[t]
                defn = self._definition_for(d)
                docs = sorted(files_of(d, eng._file_of))
                fields = dict(Type=eng.base_type(d.get("category", "")), Source=self._source_for(d, defn),
                              **{"Found in": "; ".join(docs), "First found": E.today(), "Last found": E.today(),
                                 "Scan status": E.SCAN_NEW})
                if defn:
                    fields["Definition"] = defn
                cust = next((h for h in d.get("custdef_all", []) if h["definition"] == defn), None) or {}
                if cust:
                    fields["Customer document"] = cust.get("where", ""); fields["Customer wording"] = cust.get("sentence", "")[:500]
                self.effort.add_or_update(t, change_note=f"added (Proposed) from {', '.join(docs)[:120]}" + (f": {defn[:60]}" if defn else ""),
                                          **fields)
                tv.delete(t); n += 1
            if n:
                self._save_effort_list(); self.apply_effort(); self.update_effort_ui()
                diag.activity("review_added", added=n)
                msg.configure(text=f"{n} added to the effort list as Proposed.")
            return n

        def reject():
            ts = pick()
            for t in ts:
                self.effort.add_or_update(t, change_note="not a term in this effort (review)", Status="Rejected",
                                          Type=eng.base_type(rows[t].get("category", "")))
                tv.delete(t)
            if ts:
                self._save_effort_list(); self.apply_effort(); self.update_effort_ui()
                msg.configure(text=f"{len(ts)} marked Not a term in this effort.")

        def omit():
            ts = [t for t in pick() if t in self.results]
            if not ts:
                msg.configure(text="Scan the documents first (or use Edit > Omit list) to put these on the omit list."); return
            w = self.mark_known(ts)
            if w is not None:
                w.bind("<Destroy>", lambda _e: [tv.delete(t) for t in ts if tv.exists(t) and self.results.get(t, {}).get("known")])
        b = ctk.CTkFrame(win, fg_color="transparent"); b.pack(fill="x", padx=12, pady=10)
        ctk.CTkButton(b, text="Add selected as Proposed", command=add).pack(side="left")
        ctk.CTkButton(b, text="Add all as Proposed", command=lambda: add(list(tv.get_children()))).pack(side="left", padx=6)
        ctk.CTkButton(b, text="Omit list...", fg_color="#475569", command=omit).pack(side="left", padx=6)
        ctk.CTkButton(b, text="Not a term in this effort", fg_color="#B91C1C", command=reject).pack(side="left", padx=6)
        ctk.CTkButton(b, text="Close", fg_color="#475569", command=win.destroy).pack(side="right")
        win._api = {"tree": tv, "add": add, "reject": reject, "msg": msg}
        self._review_win = win
        return win

    # ------------------------------------------------------------ after every scan
    def effort_after_scan(self, scanned_paths):
        """Found in / First found / Last found / Scan status in the effort list; new and gone terms in the results."""
        eng = _eng()
        L = self.effort
        if L is None:
            return None
        names = [os.path.basename(p) for p in scanned_paths]
        found = {t: files_of(d, eng._file_of) for t, d in self.results.items() if not d.get("file_error") and not d.get("gone")}
        rep = L.record_scan(found, names)
        seen = set(self.ws.get("seen_terms", []) or [])
        new = 0
        for t, d in self.results.items():
            fresh = t not in L.rows and t not in seen and not d.get("excluded") and not d.get("file_error") and not d.get("known")
            d["new_in_effort"] = fresh
            new += fresh
        self.ws.set("seen_terms", sorted(seen | set(found)))
        for t in rep["gone"]:
            if t not in self.results:
                row = L.rows[t]
                self.results[t] = {"category": row.get("Type") or "Acronym", "count": 0, "locations": [], "definitions": [row.get("Definition", "")],
                                   "in_db": True, "excluded": False, "gone": True, "effort": row.get("Status", ""),
                                   "selected_definition": row.get("Definition", ""), "status": "No longer found: " + row.get("Scan status", "")}
        if rep["found"] or rep["gone"]:
            self._save_effort_list()
        return {"new": new, "gone": len(rep["gone"]), "found": rep["found"]}

    # ------------------------------------------------------------ activity
    def show_activity(self):
        if self.effort is None:
            return None
        win = ctk.CTkToplevel(self); win.title(f"Activity: {self.effort.name}"); win.geometry("900x520"); win.transient(self)
        win.after(150, win.lift)
        diag.register(win, "activity")
        tv = ttk.Treeview(win, columns=("when", "who", "term", "what"), show="headings")
        for c, h, w in (("when", "When", 130), ("who", "Who", 100), ("term", "Term", 120), ("what", "What happened", 520)):
            tv.heading(c, text=h); tv.column(c, width=w, stretch=(c == "what"))
        tv.pack(fill="both", expand=True, padx=12, pady=12)
        for h in reversed(self.effort.history[-1000:]):
            tv.insert("", "end", values=h[:4])
        ctk.CTkLabel(win, text="The same lines are in the effort list's History sheet.", anchor="w").pack(fill="x", padx=12, pady=(0, 10))
        win._tree = tv
        return win

    # ------------------------------------------------------------ exports
    def _list_rows(self, approved_only=False):
        eng = _eng()
        L = self.effort
        out = []
        for t in L.order:
            r = L.rows[t]
            if r["Status"] == "Rejected" or self.results.get(t, {}).get("known"):
                continue
            if approved_only and r["Status"] != "Approved":
                continue
            out.append(r)
        return sorted(out, key=lambda r: r["Term"].lower())

    def export_effort_lists(self, folder=None, scope=None, approved_only=False):
        """The effort's acronym list: whole effort, per volume (one per file in Found in), or both. Word + Excel."""
        import docxw
        from openpyxl import Workbook
        from openpyxl.styles import Font
        L = self.effort
        if L is None:
            return None
        info = self.efforts.get(L.name) or {}
        scope = scope or info.get("list_format", "whole")
        folder = folder or filedialog.askdirectory(title="Folder for the acronym lists", initialdir=os.path.dirname(L.path))
        if not folder:
            return None
        rows = self._list_rows(approved_only)
        stamp = datetime.date.today().isoformat()
        stem = _safe(L.name).replace(" ", "_")
        made = []

        def word(path, title, rs):
            d = docxw.Doc()
            d.props = {"StationkeeperEffort": L.name, "Client": L.client, "Solicitation": L.solicitation}
            d.heading(title, 1)
            d.table(["Term", "Definition"], [[r["Term"], r["Definition"]] for r in rs], [1800, 7560])
            d.save(path); made.append(path)

        if scope in ("whole", "both"):
            word(os.path.join(folder, f"{stem}_Acronym_List_{stamp}.docx"), f"Acronyms and Abbreviations: {L.name}", rows)
        if scope in ("volume", "both"):
            vols = sorted({f for r in rows for f in re.split(r"\s*;\s*", r.get("Found in", "")) if f}, key=str.lower)
            cust = set()
            cdir = info.get("customer_dir") or ""
            if cdir and os.path.isdir(cdir):
                cust = set(os.listdir(cdir))
            for vname in vols:
                if vname in cust:
                    continue                                   # customer documents are sources, not volumes
                rs = [r for r in rows if vname in re.split(r"\s*;\s*", r.get("Found in", ""))]
                word(os.path.join(folder, f"{stem}_Acronym_List_{_safe(os.path.splitext(vname)[0])}_{stamp}.docx"),
                     f"Acronyms and Abbreviations: {os.path.splitext(vname)[0]}", rs)
        wb = Workbook(); ws = wb.active; ws.title = "Acronym list"
        cols = ["Term", "Definition", "Type", "Status", "Found in", "First found", "Last found", "Scan status", "Effort", "Client", "Solicitation"]
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c, "") if c not in ("Effort", "Client", "Solicitation") else
                       {"Effort": L.name, "Client": L.client, "Solicitation": L.solicitation}[c] for c in cols])
        for col, w in zip("ABCDEFGHIJK", (14, 44, 14, 12, 50, 12, 12, 30, 22, 18, 18)):
            ws.column_dimensions[col].width = w
        xp = os.path.join(folder, f"{stem}_Acronym_List_{stamp}.xlsx")
        wb.save(xp); made.append(xp)
        diag.activity("effort_lists_exported", files=len(made), rows=len(rows))
        return made

    def export_all_efforts(self, path=None):
        """Every effort's list in one workbook, each row with its effort, client, solicitation and where found."""
        from openpyxl import Workbook
        from openpyxl.styles import Font
        path = path or filedialog.asksaveasfilename(title="All efforts", defaultextension=".xlsx",
                                                    initialfile=f"All_Efforts_Acronyms_{datetime.date.today().isoformat()}.xlsx",
                                                    filetypes=[("Excel", "*.xlsx")])
        if not path:
            return None
        wb = Workbook(); ws = wb.active; ws.title = "All efforts"
        cols = ["Term", "Definition", "Type", "Status", "Found in", "Effort", "Client", "Solicitation", "Internal number", "First found", "Last found"]
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        n = efforts = 0
        for name in self.efforts.names():
            info = self.efforts.get(name) or {}
            p = info.get("path", "")
            if not p or not os.path.isfile(p):
                continue
            try:
                L = E.EffortList(p, name).load()
            except Exception:
                diag.swallowed(); continue
            efforts += 1
            for t in L.order:
                r = L.rows[t]
                ws.append([r["Term"], r["Definition"], r["Type"], r["Status"], r.get("Found in", ""), L.name,
                           L.client or info.get("client", ""), L.solicitation or info.get("solicitation", ""),
                           L.internal_no or info.get("internal_no", ""), r.get("First found", ""), r.get("Last found", "")])
                n += 1
        ws.auto_filter.ref = f"A1:K{max(2, ws.max_row)}"
        ws.freeze_panes = "B2"
        for col, w in zip("ABCDEFGHIJK", (14, 44, 14, 12, 50, 24, 18, 18, 16, 12, 12)):
            ws.column_dimensions[col].width = w
        wb.save(path)
        diag.activity("all_efforts_exported", rows=n, n=efforts)
        return path, n, efforts
