"""Banned words and phrases for our deliverables (2.27.0 QA/QC).

The team's list lives on this PC in config\\banned_words.csv (Phrase, Level, Use instead, Why) and travels in Knowledge
Packs. Until someone saves a list, a starter list of common proposal style rules is used. QA/QC searches the Word
volumes we write and lists every candidate with where it is and the sentence around it; a person decides.

Matching: not case sensitive; whole words; a hyphen, space or line break between words all match ("state of the art"
finds "state-of-the-art"); a single word also finds its plural and -ed / -ing forms ("guarantee" finds "guaranteed");
* stands for any letters ("utiliz*" finds utilize, utilizes, utilization). Rules only, no AI.
"""
import csv, os, re

COLUMNS = ["Phrase", "Level", "Use instead", "Why"]
LEVELS = ["Banned", "Avoid"]          # Banned -> cross "Fix"; Avoid -> "! Check"

STARTER = [
    ("world-class", "Banned", "say what makes it better, with a number", "Unsupported claim"),
    ("best-in-class", "Banned", "a measured comparison", "Unsupported claim"),
    ("state of the art", "Avoid", "name the technology and its maturity (TRL)", "Unsupported claim"),
    ("cutting edge", "Avoid", "name the technology and its maturity (TRL)", "Unsupported claim"),
    ("unparalleled", "Banned", "a measured comparison", "Unsupported claim"),
    ("unique", "Avoid", "say what only we offer, with proof", "Overused"),
    ("seamless", "Avoid", "describe the interface", "Vague"),
    ("robust", "Avoid", "state the margin or the test", "Vague"),
    ("guarantee", "Banned", "commit to a measurable outcome", "Contract risk"),
    ("warrant", "Banned", "commit to a measurable outcome", "Contract risk"),
    ("best effort", "Banned", "a firm commitment", "Weak commitment"),
    ("will try", "Banned", "will", "Weak commitment"),
    ("attempt to", "Avoid", "will", "Weak commitment"),
    ("hopefully", "Banned", "will", "Weak commitment"),
    ("we believe", "Banned", "state the fact and its proof", "Weak commitment"),
    ("we feel", "Banned", "state the fact and its proof", "Weak commitment"),
    ("we think", "Banned", "state the fact and its proof", "Weak commitment"),
    ("if possible", "Avoid", "state the condition", "Weak commitment"),
    ("as appropriate", "Avoid", "say who decides and when", "Vague"),
    ("as needed", "Avoid", "say who decides and when", "Vague"),
    ("utiliz*", "Avoid", "use", "Plain language"),
    ("leverage", "Avoid", "use", "Plain language"),
    ("synergy", "Avoid", "describe the benefit", "Jargon"),
    ("and/or", "Avoid", "and, or, or both", "Ambiguous"),
    ("etc.", "Avoid", "list the items", "Vague"),
    ("very", "Avoid", "a number or a stronger word", "Filler"),
    ("TBD", "Banned", "the value", "Placeholder left in"),
    ("TBR", "Banned", "the value", "Placeholder left in"),
    ("XXX", "Banned", "the value", "Placeholder left in"),
    ("lorem ipsum", "Banned", "the text", "Placeholder left in"),
    ("[insert*", "Banned", "the text", "Placeholder left in"),
]


def _eng():
    import app_gui
    return app_gui


def path():
    return os.path.join(_eng().CONFIG_DIR, "banned_words.csv")


def starter():
    return [dict(zip(COLUMNS, r)) for r in STARTER]


def _clean(r):
    phrase = " ".join(str(r.get("Phrase") or "").split())
    level = str(r.get("Level") or "").strip().capitalize()
    return {"Phrase": phrase, "Level": level if level in LEVELS else "Banned",
            "Use instead": " ".join(str(r.get("Use instead") or "").split()),
            "Why": " ".join(str(r.get("Why") or "").split())}


def load():
    """The team's list; the starter list until one is saved."""
    p = path()
    if not os.path.isfile(p):
        return starter()
    out = []
    with open(p, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            r = _clean(r)
            if r["Phrase"]:
                out.append(r)
    return out


def save(items):
    p = path()
    os.makedirs(os.path.dirname(p), exist_ok=True)
    seen, rows = set(), []
    for it in items:
        it = _clean(it)
        k = it["Phrase"].casefold()
        if it["Phrase"] and k not in seen:
            seen.add(k); rows.append(it)
    with open(p + ".tmp", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        for it in sorted(rows, key=lambda x: x["Phrase"].casefold()):
            w.writerow(it)
    os.replace(p + ".tmp", p)
    return len(rows)


def read_file(fp):
    """Phrases from a CSV or Excel list (a Phrase column, or the first column), or a text file with one per line."""
    ext = os.path.splitext(fp)[1].lower()
    rows = []
    if ext == ".csv":
        with open(fp, encoding="utf-8-sig", newline="") as f:
            data = list(csv.reader(f))
        rows = _table(data)
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(fp, read_only=True, data_only=True)
        try:
            data = [["" if v is None else str(v) for v in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
        rows = _table(data)
    else:
        with open(fp, encoding="utf-8-sig", errors="replace") as f:
            rows = [{"Phrase": line.strip()} for line in f if line.strip() and not line.startswith("#")]
    return [r for r in (_clean(x) for x in rows) if r["Phrase"]]


def _table(data):
    if not data:
        return []
    head = [str(h or "").strip().casefold() for h in data[0]]
    alias = {"phrase": "Phrase", "word": "Phrase", "banned word": "Phrase", "banned words": "Phrase", "term": "Phrase",
             "level": "Level", "severity": "Level", "use instead": "Use instead", "replacement": "Use instead",
             "suggestion": "Use instead", "why": "Why", "reason": "Why", "note": "Why", "notes": "Why"}
    idx = {alias[h]: i for i, h in enumerate(head) if h in alias}
    if "Phrase" not in idx:                       # no header: the first column holds the phrases
        return [{"Phrase": r[0]} for r in data if r and str(r[0]).strip()]
    return [{c: (r[i] if i < len(r) else "") for c, i in idx.items()} for r in data[1:] if r]


def import_file(fp, keep_starter=True):
    """Add phrases from a file. Returns (added, already there). keep_starter=False (a team Knowledge Pack on a PC that
    never saved its own list): the team's list replaces the starter list instead of joining it."""
    cur = load() if keep_starter or os.path.isfile(path()) else []
    have = {x["Phrase"].casefold() for x in cur}
    added = 0
    for r in read_file(fp):
        if r["Phrase"].casefold() not in have:
            cur.append(r); have.add(r["Phrase"].casefold()); added += 1
    save(cur)
    return added, len(read_file(fp)) - added


def pattern(phrase):
    """The regular expression for one list entry."""
    words = phrase.split()
    parts = []
    for w in words:
        p = re.escape(w).replace(r"\*", r"[\w'-]*")
        parts.append(p)
    body = r"[\s\-‐‑–]+".join(parts)
    body = body.replace(r"\-", r"[\s\-‐‑–]*") if len(words) == 1 and "-" in phrase else body
    if len(words) == 1 and "*" not in phrase and phrase[-1:].isalpha() and len(phrase) >= 4 and not phrase.isupper():
        body += r"(?:s|es|d|ed|ing)?"
    start = r"(?<![\w])" if phrase[:1].isalnum() else ""
    end = r"(?![\w])" if phrase[-1:].isalnum() or phrase.endswith("*") else ""
    flags = 0 if phrase.isupper() and len(phrase) <= 5 else re.I    # TBD / XXX: capitals only
    return re.compile(start + body + end, flags)


def _sentence(text, a, b):
    s = max(text.rfind(". ", 0, a), text.rfind("\n", 0, a))
    e = min([x for x in (text.find(". ", b), text.find("\n", b)) if x != -1] or [len(text)])
    return " ".join(text[s + 1:e + 1].split())[:300]


def find(chunks, items=None, kind_of=None, skip_kinds=("notes",)):
    """chunks: [(location label, text)] of one volume. Returns [{"phrase", "found", "level", "use", "why", "loc",
    "kind", "sentence"}] in reading order."""
    items = load() if items is None else items
    kind_of = kind_of or (lambda loc: "text")
    pats = [(it, pattern(it["Phrase"])) for it in items if it.get("Phrase")]
    out = []
    for loc, text in chunks:
        kind = kind_of(loc)
        if kind in skip_kinds or not text:
            continue
        for it, rx in pats:
            for m in rx.finditer(text):
                out.append({"phrase": it["Phrase"], "found": m.group(0), "level": it["Level"], "use": it.get("Use instead", ""),
                            "why": it.get("Why", ""), "loc": loc, "kind": kind, "sentence": _sentence(text, m.start(), m.end())})
    return out
