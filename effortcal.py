"""The effort calendar as a file (2.88.0): a template people fill in, import with a preview, and export of the whole
calendar with its details.

The template (Excel; CSV reads the same way) has one row per date: Date, Start time, End time, Time zone, Kind, What,
Owner, Stage, Location, Details. Kind decides what a row does on import:
- Proposal due: the effort's due date and time, to the minute, with its time zone (the banner countdown follows it).
- Stage start / Stage end: the planned start or end of the stage named in Stage (matched by name).
- Task: a task with its owner and due date (Details go in its notes).
- Questions due, Review, Meeting, Orals, Milestone, Other: an entry on the effort's calendar with its times, owner,
  stage, location and details.
Import is add-or-update: a row matching one already there (same kind, date and wording; the same task; the same stage)
updates it instead of adding a copy, so a file can be imported again after it changes. Every import can be undone as a
whole. Outlook calendar files (.ics) import too: each event's category or wording gives its kind.
Export writes every date the effort has (due date, stage plans, tasks, calendar entries and the questions-due date the
customer documents give) in the template's columns, so an export can be edited and imported back, and as an .ics file
for Outlook. Files are only read and written here; nothing leaves the PC. Rules only (No AI).
"""
import csv, datetime, json, os, re, shutil, uuid

COLUMNS = ("Date", "Start time", "End time", "Time zone", "Kind", "What", "Owner", "Stage", "Location", "Details")
KINDS = ("Proposal due", "Questions due", "Stage start", "Stage end", "Review", "Meeting", "Orals", "Milestone", "Task", "Other")
EVENT_KINDS = ("Questions due", "Review", "Meeting", "Orals", "Milestone", "Other")
_HEAD = {"date": "Date", "day": "Date", "start": "Start time", "start time": "Start time", "time": "Start time", "from": "Start time",
         "end": "End time", "end time": "End time", "to": "End time", "time zone": "Time zone", "zone": "Time zone", "tz": "Time zone",
         "kind": "Kind", "type": "Kind", "category": "Kind", "what": "What", "title": "What", "event": "What", "subject": "What",
         "summary": "What", "name": "What", "owner": "Owner", "who": "Owner", "assigned to": "Owner", "stage": "Stage", "phase": "Stage",
         "location": "Location", "where": "Location", "room": "Location", "details": "Details", "notes": "Details",
         "description": "Details"}
SNAP_DIR = "calendar"
SNAP = "before_import.json"


# ------------------------------------------------------------------ reading values
def parse_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    import effortrec as R
    s = " ".join(str(v).split())
    d = R.parse_date(s)
    if d:
        return d
    for f in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%d %b %Y", "%b %d %Y", "%B %d %Y", "%A, %B %d, %Y", "%a %b %d %Y"):
        try:
            return datetime.datetime.strptime(s.replace(",", "") if "%A" not in f else s, f).date()
        except ValueError:
            continue
    return None


_TIME = re.compile(r"^(\d{1,2})(?::(\d{2}))?(?::\d{2})?\s*([ap])?\.?\s*m?\.?$", re.I)


def parse_time(v):
    """'2:00 PM', '14:00', '2 p.m.', '1400', a time or datetime cell -> 'HH:MM'; '' for nothing; None when not understood."""
    if v is None or v == "":
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%H:%M")
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    if isinstance(v, (int, float)) and 0 <= v < 1:                    # an Excel fraction of a day
        m = int(round(v * 24 * 60))
        return f"{m // 60:02d}:{m % 60:02d}"
    s = " ".join(str(v).split()).lower().replace("noon", "12:00 pm").replace("midnight", "12:00 am")
    if re.fullmatch(r"\d{3,4}", s):
        s = s[:-2] + ":" + s[-2:]
    m = _TIME.match(s)
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2) or 0)
    ap = (m.group(3) or "").lower()
    if ap == "p" and h < 12:
        h += 12
    if ap == "a" and h == 12:
        h = 0
    if not (0 <= h < 24 and 0 <= mi < 60):
        return None
    return f"{h:02d}:{mi:02d}"


def nice_time(hm):
    hm = parse_time(hm) if hm and not re.fullmatch(r"\d{2}:\d{2}", str(hm)) else hm
    if not hm:
        return ""
    h, m = int(hm[:2]), int(hm[3:5])
    return f"{h % 12 or 12}:{m:02d} {'a.m.' if h < 12 else 'p.m.'}"


def kind_of(text, title=""):
    t = " ".join(str(text or "").split()).casefold()
    for k in KINDS:
        if t == k.casefold():
            return k
    alias = {"due": "Proposal due", "proposal": "Proposal due", "submission due": "Proposal due", "questions": "Questions due",
             "start": "Stage start", "end": "Stage end", "color team": "Review", "meeting": "Meeting", "orals": "Orals",
             "milestone": "Milestone", "task": "Task", "to do": "Task"}
    if t in alias:
        return alias[t]
    w = " ".join(str(title or "").split()).casefold()
    if re.search(r"\b(proposal|submission|response)s? (is )?due\b", w):
        return "Proposal due"
    if re.search(r"\bquestions? (are )?due\b", w):
        return "Questions due"
    if "orals" in w or "oral presentation" in w:
        return "Orals"
    if re.search(r"\b(pink|red|gold|green|blue|white|black) team\b|\breview\b", w):
        return "Review"
    return "Meeting" if t in ("", "appointment", "event") and w else ("Other" if t else "Meeting")


def _norm(s, n=300):
    return " ".join(str(s if s is not None else "").split())[:n]


# ------------------------------------------------------------------ reading a file
def read_rows(path, effort=""):
    """Rows as dicts in the template's columns, each with its row number ("row"). effort: the effort's name, taken off
    the front of an event title our own .ics export wrote ("Kestrel: Red team review")."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".ics":
        return read_ics(path, effort)
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            table = [list(r) for r in csv.reader(f)]
    elif ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = next((s for s in wb.worksheets if s.title.lower().startswith("calendar")), wb.worksheets[0])
        table = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        raise ValueError("Calendar files are .xlsx, .csv or .ics")
    head_i, cols = None, {}
    for i, r in enumerate(table[:15]):
        got = {}
        for j, c in enumerate(r):
            k = _HEAD.get(_norm(c).casefold().rstrip(":"))
            if k and k not in got:
                got[k] = j
        if "Date" in got and ("What" in got or "Kind" in got):
            head_i, cols = i, got
            break
    if head_i is None:
        raise ValueError("No header row with Date and What (or Kind) was found")
    out = []
    for i, r in enumerate(table[head_i + 1:], start=head_i + 2):
        row = {c: (r[cols[c]] if c in cols and cols[c] < len(r) else "") for c in COLUMNS}
        if all(v in (None, "") for v in row.values()):
            continue
        if _norm(row["What"]).casefold().startswith("example"):
            continue
        row["row"] = i
        out.append(row)
    return out


def _unfold(text):
    return re.sub(r"\r?\n[ \t]", "", text).splitlines()


def _ics_unescape(s):
    return s.replace("\\n", "\n").replace("\\N", "\n").replace("\\,", ",").replace("\\;", ";").replace("\\\\", "\\")


def _ics_when(name, params, value):
    """(date, 'HH:MM' or '', zone)."""
    v = value.strip()
    zone = ""
    for p in params:
        if p.upper().startswith("TZID="):
            zone = p.split("=", 1)[1].strip('"')
    if re.fullmatch(r"\d{8}", v):
        return datetime.datetime.strptime(v, "%Y%m%d").date(), "", zone
    m = re.fullmatch(r"(\d{8})T(\d{2})(\d{2})(\d{2})?(Z?)", v)
    if not m:
        return None, "", zone
    d = datetime.datetime.strptime(m.group(1), "%Y%m%d").date()
    if m.group(5):
        zone = "UTC"
    return d, f"{m.group(2)}:{m.group(3)}", zone


def read_ics(path, effort=""):
    with open(path, encoding="utf-8-sig", errors="replace") as f:
        lines = _unfold(f.read())
    out, cur, n = [], None, 0
    for ln in lines:
        if ln.upper() == "BEGIN:VEVENT":
            cur = {c: "" for c in COLUMNS}; n += 1; cur["row"] = n
            continue
        if ln.upper() == "END:VEVENT":
            if cur is not None:
                cur["Kind"] = kind_of(cur.get("_cat", ""), cur["What"])
                if cur["Kind"] == "Other" and not cur.get("_cat"):
                    cur["Kind"] = "Meeting"
                cur.pop("_cat", None)
                if "_zone" in cur:                                   # our own export: the zone and details as written
                    cur["Time zone"] = cur.pop("_zone")
                if "_end" in cur:
                    cur["End time"] = cur.pop("_end")
                if cur.pop("_ours", False):                            # our export: Details as written, not the summary for Outlook
                    cur["Details"] = cur.pop("_details", "")
                out.append(cur)
            cur = None
            continue
        if cur is None or ":" not in ln:
            continue
        left, value = ln.split(":", 1)
        name, *params = left.split(";")
        name = name.upper()
        if name == "DTSTART":
            d, hm, zone = _ics_when(name, params, value)
            cur["Date"], cur["Start time"], cur["Time zone"] = d, hm, zone
        elif name == "DTEND":
            d, hm, _z = _ics_when(name, params, value)
            if hm:
                cur["End time"] = hm
        elif name == "SUMMARY":
            cur["What"] = _ics_unescape(value)
            if effort and cur["What"].startswith(effort + ": "):
                cur["What"] = cur["What"][len(effort) + 2:]
        elif name == "DESCRIPTION":
            cur["Details"] = _ics_unescape(value)
        elif name == "LOCATION":
            cur["Location"] = _ics_unescape(value)
        elif name == "CATEGORIES":
            cur["_cat"] = _ics_unescape(value).split(",")[0]
        elif name == "X-SK-OWNER":
            cur["Owner"] = _ics_unescape(value)
        elif name == "X-SK-STAGE":
            cur["Stage"] = _ics_unescape(value)
        elif name == "X-SK-KIND":
            cur["_cat"] = _ics_unescape(value); cur["_ours"] = True
        elif name == "X-SK-END":
            cur["_end"] = "" if value.strip() == "none" else value.strip()
        elif name == "X-SK-ZONE":
            cur["_zone"] = _ics_unescape(value)
        elif name == "X-SK-DETAILS":
            cur["_details"] = _ics_unescape(value)
    return out


# ------------------------------------------------------------------ the plan
def events(rec):
    return list(rec.data.get("events", []))


def _ekey(kind, date, title):
    return (kind, date, _norm(title).casefold())


def plan(rows, rec):
    """[{row, kind, date, start, end, zone, title, owner, stage, location, details, action, ok, why}]: what each row does.
    action: 'due', 'stage start', 'stage end', 'add task', 'update task', 'add event', 'update event', 'same' or 'skip'."""
    stages = {s["name"].casefold(): s for s in rec.stages}
    evs = {_ekey(e["kind"], e["date"], e["title"]): e for e in events(rec)}
    tasks = {(_norm(t["title"]).casefold()): t for t in rec.data.get("tasks", [])}
    out, seen = [], set()
    for r in rows:
        kind = kind_of(r.get("Kind"), r.get("What"))
        d = parse_date(r.get("Date"))
        st, en = parse_time(r.get("Start time")), parse_time(r.get("End time"))
        p = {"row": r.get("row"), "kind": kind, "date": d.isoformat() if d else "", "start": st or "", "end": en or "",
             "zone": _norm(r.get("Time zone"), 40), "title": _norm(r.get("What"), 200), "owner": _norm(r.get("Owner"), 80),
             "stage": _norm(r.get("Stage"), 80), "location": _norm(r.get("Location"), 200), "details": _norm(r.get("Details"), 2000),
             "action": "skip", "ok": False, "why": ""}
        if d is None:
            p["why"] = "No date, or a date that was not understood"
        elif st is None or en is None:
            p["why"] = "A time that was not understood (use 2:00 PM or 14:00)"
        elif kind == "Proposal due":
            iso = p["date"] + ("T" + p["start"] if p["start"] else "")
            p["iso"] = iso
            if "due" in seen:
                p["why"] = "A second Proposal due row; the first one counts"
            else:
                seen.add("due")
                p["ok"] = True
                p["action"] = "same" if rec.data.get("due") == iso and rec.data.get("zone", "") == p["zone"] else "due"
        elif kind in ("Stage start", "Stage end"):
            s = stages.get(p["stage"].casefold()) or stages.get(p["title"].casefold())
            if s is None:
                p["why"] = f"No stage named {p['stage'] or p['title'] or '(blank)'} in this effort (Effort Record > Stages)"
            else:
                p["stage"] = s["name"]
                field = "planned" if kind == "Stage start" else "planned_end"
                p["ok"], p["action"] = True, ("same" if (s.get(field) or "") == p["date"] else kind.lower())
        elif kind == "Task":
            if not p["title"]:
                p["why"] = "A task with no wording in What"
            else:
                t = tasks.get(p["title"].casefold())
                p["ok"] = True
                if t is None:
                    p["action"] = "add task"
                elif (t.get("due"), t.get("owner", "")) == (p["date"], p["owner"] or t.get("owner", "")):
                    p["action"] = "same"
                else:
                    p["action"] = "update task"
        else:
            if not p["title"]:
                p["title"] = kind
            k = _ekey(kind, p["date"], p["title"])
            if k in seen:
                p["why"] = "The same entry twice in the file"
            else:
                seen.add(k)
                e = evs.get(k)
                p["ok"] = True
                if e is None:
                    p["action"] = "add event"
                else:
                    same = all(e.get(f, "") == p[f] for f in ("start", "end", "zone", "owner", "stage", "location", "details"))
                    p["action"] = "same" if same else "update event"
        out.append(p)
    return out


def summary(pl):
    """'3 to add, 1 to update, 2 unchanged, 1 skipped'."""
    add = sum(1 for p in pl if p["action"] in ("add task", "add event"))
    upd = sum(1 for p in pl if p["action"] in ("due", "stage start", "stage end", "update task", "update event"))
    same = sum(1 for p in pl if p["action"] == "same")
    skip = sum(1 for p in pl if not p["ok"])
    return f"{add} to add, {upd} to update, {same} unchanged, {skip} skipped"


ACTION_WORDS = {"due": "Sets the due date", "stage start": "Sets the stage's planned start", "stage end": "Sets the stage's planned end",
                "add task": "Adds a task", "update task": "Updates the task", "add event": "Adds to the calendar",
                "update event": "Updates the calendar entry", "same": "Already there", "skip": "Skipped"}


# ------------------------------------------------------------------ applying it
def snapshot_path(data_dir):
    return os.path.join(data_dir, SNAP_DIR, SNAP)


def apply(pl, rec, source="", who=None):
    """Applies the plan's rows that do something. Keeps a copy of the record first (undo). Returns counts."""
    data_dir = os.path.dirname(rec.path)
    os.makedirs(os.path.join(data_dir, SNAP_DIR), exist_ok=True)
    with open(snapshot_path(data_dir), "w", encoding="utf-8") as f:
        json.dump({"at": datetime.datetime.now().isoformat(timespec="seconds"), "source": os.path.basename(source),
                   "record": rec.data}, f)
    n = {"due": 0, "stages": 0, "tasks": 0, "events": 0}
    stages = {s["name"]: s for s in rec.stages}
    evs = rec.data.setdefault("events", [])
    for p in pl:
        a = p["action"]
        if not p["ok"] or a == "same":
            continue
        if a == "due":
            rec.data["due"], rec.data["zone"] = p["iso"], p["zone"]
            n["due"] += 1
        elif a in ("stage start", "stage end"):
            stages[p["stage"]]["planned" if a == "stage start" else "planned_end"] = p["date"]
            n["stages"] += 1
        elif a == "add task":
            tid = rec.add_task(p["title"], owner=p["owner"], due=p["date"], notes=p["details"], who=who)
            if p["stage"] and tid:
                s = next((s for s in rec.stages if s["name"].casefold() == p["stage"].casefold()), None)
                if s:
                    rec.task(tid)["stage"] = s["id"]
            n["tasks"] += 1
        elif a == "update task":
            t = next(t for t in rec.data["tasks"] if _norm(t["title"]).casefold() == p["title"].casefold())
            t["due"] = p["date"]
            if p["owner"]:
                t["owner"] = p["owner"]
            if p["details"]:
                t["notes"] = p["details"]
            n["tasks"] += 1
        elif a in ("add event", "update event"):
            k = _ekey(p["kind"], p["date"], p["title"])
            e = next((e for e in evs if _ekey(e["kind"], e["date"], e["title"]) == k), None)
            if e is None:
                e = {"id": uuid.uuid4().hex[:10]}
                evs.append(e)
            e.update({f: p[f] for f in ("kind", "date", "start", "end", "zone", "title", "owner", "stage", "location", "details")})
            e["source"] = os.path.basename(source)
            n["events"] += 1
    evs.sort(key=lambda e: (e["date"], e.get("start", "")))
    said = [f"{v} {k}" for k, v in n.items() if v]
    rec.log("effort", f"Calendar imported from {os.path.basename(source) or 'a file'}: " + (", ".join(said) if said else "nothing new"),
            who=who)
    rec.save()
    return n


def can_undo(data_dir):
    return os.path.isfile(snapshot_path(data_dir))


def undo(rec, who=None):
    """Puts the record back as it was before the last import (what happened since then on the timeline stays)."""
    p = snapshot_path(os.path.dirname(rec.path))
    if not os.path.isfile(p):
        return False
    with open(p, encoding="utf-8") as f:
        snap = json.load(f)
    old = snap["record"]
    for k in ("due", "zone", "stages", "tasks", "events"):
        if k in old:
            rec.data[k] = old[k]
        else:
            rec.data.pop(k, None)
    rec.log("effort", f"Calendar import from {snap.get('source') or 'a file'} undone", who=who)
    rec.save()
    os.remove(p)
    return True


def remove_event(rec, eid, who=None):
    evs = rec.data.get("events", [])
    e = next((e for e in evs if e["id"] == eid), None)
    if e is None:
        return False
    evs.remove(e)
    rec.log("effort", f"Calendar: removed {e['kind'].lower()} {e['title']} ({e['date']})", who=who)
    rec.save()
    return True


# ------------------------------------------------------------------ export
def rows_of(rec, extra=()):
    """Every date the effort has, in the template's columns (date order)."""
    out = []
    due = rec.data.get("due", "")
    if due:
        out.append({"Date": due[:10], "Start time": due[11:16] if len(due) > 10 else "", "End time": "", "Time zone": rec.data.get("zone", ""),
                    "Kind": "Proposal due", "What": "Proposal due", "Owner": "", "Stage": "", "Location": "", "Details": ""})
    for s in rec.stages:
        for field, kind in (("planned", "Stage start"), ("planned_end", "Stage end")):
            if s.get(field):
                out.append({"Date": s[field], "Start time": "", "End time": "", "Time zone": "", "Kind": kind, "What": s["name"],
                            "Owner": "", "Stage": s["name"], "Location": "", "Details": ""})
    for t in rec.data.get("tasks", []):
        if t.get("due"):
            out.append({"Date": t["due"], "Start time": "", "End time": "", "Time zone": "", "Kind": "Task", "What": t["title"],
                        "Owner": t.get("owner", ""), "Stage": "", "Location": "", "Details": (t.get("notes") or "") +
                        (f" (status: {t.get('status')})" if t.get("status") and t.get("status") != "Not started" else "")})
    have = {(e["kind"], e["date"]) for e in events(rec)}
    for e in events(rec):
        out.append({"Date": e["date"], "Start time": e.get("start", ""), "End time": e.get("end", ""), "Time zone": e.get("zone", ""),
                    "Kind": e["kind"], "What": e["title"], "Owner": e.get("owner", ""), "Stage": e.get("stage", ""),
                    "Location": e.get("location", ""), "Details": e.get("details", "")})
    for d, kind, text in extra:
        if kind == "questions" and ("Questions due", d.isoformat()) not in have:
            out.append({"Date": d.isoformat(), "Start time": "", "End time": "", "Time zone": "", "Kind": "Questions due", "What": text,
                        "Owner": "", "Stage": "", "Location": "", "Details": "From the customer documents"})
    out.sort(key=lambda r: (r["Date"], r["Start time"] or "99"))
    return out


GUIDE = (("Date", "The day: 12/15/2026, 2026-12-15 or Dec 15, 2026. Required."),
         ("Start time", "2:00 PM or 14:00. Blank for an all-day entry. For Proposal due, the exact due time."),
         ("End time", "When a review, meeting or orals ends. Optional."),
         ("Time zone", "As the customer writes it: MT, Mountain Time, ET. Shown beside the time."),
         ("Kind", "Proposal due sets the effort's due date and time. Stage start and Stage end set the planned dates of the stage "
                  "named in Stage. Task adds a task. Questions due, Review, Meeting, Orals, Milestone and Other go on the calendar."),
         ("What", "The wording shown on the calendar (Red team review, Kickoff). For Task, the task."),
         ("Owner", "Who it belongs to; a name from the author list."),
         ("Stage", "For Stage start and Stage end: the stage's name exactly as in Effort Record > Stages. Optional otherwise."),
         ("Location", "A room, a Teams meeting name, a site."),
         ("Details", "Anything else: the agenda, what to bring, a phone bridge."),
         ("Importing again", "A row matching one already in the effort updates it; nothing is duplicated. Effort Home > Calendar > "
                             "Undo import puts back the calendar as it was before the last import."))


def write_xlsx(path, rows, effort="", template=False):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Calendar"
    ws.append([f"{effort} calendar" if effort else "Effort calendar"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["One row per date. Kind decides what the row does when it is imported (see the How to fill it in sheet)."])
    ws["A2"].font = Font(italic=True, color="4E5D71")
    ws.append(list(COLUMNS))
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F5FAD")
    if template and not rows:
        rows = [{"Date": "12/15/2026", "Start time": "2:00 PM", "End time": "", "Time zone": "MT", "Kind": "Proposal due",
                 "What": "Example: proposal due (delete the Example rows)", "Owner": "", "Stage": "", "Location": "", "Details": ""},
                {"Date": "11/20/2026", "Start time": "9:00 AM", "End time": "4:00 PM", "Time zone": "MT", "Kind": "Review",
                 "What": "Example: Red team review", "Owner": "", "Stage": "Red team", "Location": "Room 2B", "Details": "Bring the packet"}]
    for r in rows:
        d = parse_date(r["Date"])
        ws.append([d if d else r["Date"], nice_time(r["Start time"]) if r["Start time"] else "", nice_time(r["End time"]) if r["End time"] else "",
                   r["Time zone"], r["Kind"], r["What"], r["Owner"], r["Stage"], r["Location"], r["Details"]])
        if d:
            ws.cell(ws.max_row, 1).number_format = "mmm d, yyyy"
    for col, w in zip("ABCDEFGHIJ", (14, 11, 11, 11, 15, 42, 18, 18, 22, 50)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=4):
        row[9].alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"' + ",".join(KINDS) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"E4:E{max(ws.max_row, 4) + 300}")
    ws.freeze_panes = "A4"
    g = wb.create_sheet("How to fill it in")
    g.append(["Column", "What goes in it"])
    for c in g[1]:
        c.font = Font(bold=True)
    for a, b in GUIDE:
        g.append([a, b])
    g.column_dimensions["A"].width = 18; g.column_dimensions["B"].width = 110
    for row in g.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    tmp = path + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, path)
    return path


def _ics_escape(s):
    return str(s or "").replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def _fold(line):
    b = line.encode("utf-8")
    if len(b) <= 75:
        return line
    parts, cur = [], ""
    for ch in line:
        if len((cur + ch).encode("utf-8")) > (75 if not parts else 74):
            parts.append(cur); cur = ch
        else:
            cur += ch
    parts.append(cur)
    return "\r\n ".join(parts)


def write_ics(path, rows, effort=""):
    """An .ics file Outlook opens: every row an event (all-day when it has no time), in the PC's local time."""
    now = datetime.datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    out = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Stationkeeper//Effort calendar//EN", "CALSCALE:GREGORIAN",
           f"X-WR-CALNAME:{_ics_escape((effort + ' calendar').strip())}"]
    for i, r in enumerate(rows):
        d = parse_date(r["Date"])
        if d is None:
            continue
        uid = re.sub(r"[^A-Za-z0-9]", "", f"{effort}{r['Kind']}{d.isoformat()}{r['What']}")[:60] + f"-{i}@stationkeeper.local"
        out += ["BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{now}"]
        st = r["Start time"]
        if st:
            start = datetime.datetime.combine(d, datetime.time(int(st[:2]), int(st[3:5])))
            en = r["End time"]
            end = datetime.datetime.combine(d, datetime.time(int(en[:2]), int(en[3:5]))) if en else start + datetime.timedelta(minutes=30)
            out += [f"DTSTART:{start.strftime('%Y%m%dT%H%M%S')}", f"DTEND:{end.strftime('%Y%m%dT%H%M%S')}"]
        else:
            out += [f"DTSTART;VALUE=DATE:{d.strftime('%Y%m%d')}", f"DTEND;VALUE=DATE:{(d + datetime.timedelta(days=1)).strftime('%Y%m%d')}"]
        title = r["What"] if r["Kind"] in ("Proposal due", "Questions due") or not effort else f"{effort}: {r['What']}"
        desc = "\n".join(x for x in (r["Details"], f"Owner: {r['Owner']}" if r["Owner"] else "", f"Stage: {r['Stage']}" if r["Stage"] else "",
                                     f"Time zone as written: {r['Time zone']}" if r["Time zone"] else "") if x)
        out += [_fold("SUMMARY:" + _ics_escape(title)), _fold("CATEGORIES:" + _ics_escape(r["Kind"])),
                _fold("X-SK-KIND:" + _ics_escape(r["Kind"]))]
        if desc:
            out.append(_fold("DESCRIPTION:" + _ics_escape(desc)))
        if r["Location"]:
            out.append(_fold("LOCATION:" + _ics_escape(r["Location"])))
        if r["Owner"]:
            out.append(_fold("X-SK-OWNER:" + _ics_escape(r["Owner"])))
        if r["Stage"]:
            out.append(_fold("X-SK-STAGE:" + _ics_escape(r["Stage"])))
        if st:
            out.append("X-SK-END:" + (r["End time"] or "none"))
        if r["Time zone"]:
            out.append(_fold("X-SK-ZONE:" + _ics_escape(r["Time zone"])))
        if r["Details"]:
            out.append(_fold("X-SK-DETAILS:" + _ics_escape(r["Details"])))
        out.append("END:VEVENT")
    out.append("END:VCALENDAR")
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write("\r\n".join(out) + "\r\n")
    return path
